"""Masters: the brand's master for a design deck, uploaded masters imported, and masters kept through the port.

Three ways a deck gets its master:

* **From captured furniture** (`prepare_brand_master`): the brand's `capturedFurniture` (from
  `app.core.brand.extract`) is synthesised into a real master (`app.core.brand.synth_master`) and
  imported. This is what the old `/v1/jobs` did for every slide. Darwin's switches are honoured:
  - `apply_brand_layout=false` (**debrand**): the master keeps only the heading geometry (so the
    title still lands in its native placeholder) and carries no theme, furniture or background;
  - `layout_template` (the brand's master PNG) with `apply_template_bg`: the PNG is painted as the
    archetype layout's background, under the furniture (not in debrand mode).
  Both were ignored by Slide Studio's compat layer (ROADMAP).
* **Plain** (no furniture): a clean default master at 16:9.
* **Uploaded** (`import_uploaded_master`): a customer `.pptx` is checked and imported into the deck.
  A failed re-import puts the previous master back (Slide Studio `services/master_import`).

`persist_master` keeps an imported master through the storage port (`MasterPort`: the `.pptx`, the
manifest, one PNG per layout), so nothing is written to lasting storage any other way.

Import itself (manifest, layout backgrounds through the `Renderer` port) is
`app.core.engine_service.import_master`.
"""

from __future__ import annotations

import io
import logging
import shutil
import zipfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.oxml.ns import qn

from app.core import engine_service
from app.core.brand.synth_master import ARCHETYPE_NAME, synthesize_master
from app.core.design.deck import DesignDeck
from app.core.slides.spec import BrandOptions
from app.core.storage.models import CallerContext, Master, MasterCreate
from app.core.storage.ports import Storage
from app.engine.renderer import Renderer

_log = logging.getLogger(__name__)

JSON = dict[str, Any]

#: Default 16:9 slide in EMU (13.333 in x 7.5 in), for a deck with no captured furniture.
DEFAULT_W_EMU = 12192000
DEFAULT_H_EMU = 6858000
#: What an uploaded master may be.
MASTER_SUFFIXES = (".pptx", ".potx")
MAX_MASTER_BYTES = 50 * 1024 * 1024
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


class MasterError(ValueError):
    """The master is unusable; the message is safe to show."""


def plain_capture() -> JSON:
    """A minimal capturedFurniture: a clean default 16:9 master, no furniture, no theme."""
    return {"slideWidthEmu": DEFAULT_W_EMU, "slideHeightEmu": DEFAULT_H_EMU, "layouts": {}}


def debrand_capture(captured: JSON) -> JSON:
    """The furniture with the brand removed: canvas and heading geometry only (debrand mode)."""
    layouts: JSON = {}
    for arch, blob in (captured.get("layouts") or {}).items():
        if isinstance(blob, dict) and isinstance(blob.get("placeholders"), dict):
            layouts[arch] = {"placeholders": blob["placeholders"], "shapes": []}
    return {"slideWidthEmu": int(captured.get("slideWidthEmu") or DEFAULT_W_EMU),
            "slideHeightEmu": int(captured.get("slideHeightEmu") or DEFAULT_H_EMU), "layouts": layouts}


def capture_for(furniture: JSON | None, brand: BrandOptions) -> tuple[JSON, str]:
    """The capture to synthesise and which mode it is: "branded", "debranded" or "plain"."""
    if not furniture:
        return plain_capture(), "plain"
    if not brand.apply_brand_layout:
        return debrand_capture(furniture), "debranded"
    return furniture, "branded"


def apply_template_background(master: Path, png: bytes, layout_name: str | None) -> bool:
    """Paint `png` as the background of the layout named `layout_name` (else every layout). True when
    a layout took it. The picture stretches to the slide, under the layout's own shapes."""
    try:
        from PIL import Image

        with Image.open(io.BytesIO(png)) as image:
            image.verify()
    except Exception as exc:  # noqa: BLE001 - not an image: the template is refused, not half-applied
        raise MasterError("The layout template is not a readable image.") from exc
    prs: Any = Presentation(str(master))
    applied = False
    for layout in prs.slide_masters[0].slide_layouts:
        if layout_name and layout.name != layout_name:
            continue
        _image_part, rid = layout.part.get_or_add_image_part(io.BytesIO(png))
        c_sld = layout.element.find(qn("p:cSld"))
        if c_sld is None:
            continue
        old = c_sld.find(qn("p:bg"))
        if old is not None:
            c_sld.remove(old)
        from lxml import etree

        bg = etree.fromstring(
            '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/>'
            '<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
        c_sld.insert(0, bg)
        applied = True
    if applied:
        prs.save(str(master))
    return applied


def prepare_brand_master(deck: DesignDeck, furniture: JSON | None, brand: BrandOptions, archetype: str, *,
                         renderer: Renderer | None) -> JSON:
    """Synthesise the brand's master for `archetype`, import it into the deck, and return a summary:
    `{"mode", "layoutId", "templateBackground"}`."""
    captured, mode = capture_for(furniture, brand)
    work = deck.dir / ".master"
    work.mkdir(exist_ok=True)
    path = synthesize_master(captured, work / "master.pptx")
    template_bg = False
    if brand.layout_template and brand.apply_template_bg and mode != "debranded":
        name = ARCHETYPE_NAME.get(archetype) if mode == "branded" else None
        template_bg = apply_template_background(path, brand.layout_template, name)
    shutil.copyfile(path, deck.master_path)
    engine_service.import_master(deck, deck.master_path, "master.pptx", renderer=renderer,
                                 render=renderer is not None)
    return {"mode": mode, "layoutId": layout_for(deck, archetype), "templateBackground": template_bg}


def layout_for(deck: DesignDeck, archetype: str) -> str:
    """The deck's layout for an archetype: the one named Cover/Divider/Content, else the first."""
    want = ARCHETYPE_NAME.get(archetype, "Content")
    layouts = deck.master().get("layouts") or []
    for lay in layouts:
        if lay.get("name") == want:
            return str(lay["id"])
    if archetype == "content":
        named = next((lay for lay in layouts if "content" in str(lay.get("name") or "").lower()), None)
        if named:
            return str(named["id"])
    if layouts:
        return str(layouts[0]["id"])
    raise MasterError("The master has no layouts.")


# ------------------------------------------------------------------------------- uploaded masters
def check_master_bytes(data: bytes, filename: str | None) -> str:
    """The upload's suffix, or MasterError when it cannot be a PowerPoint master."""
    suffix = Path(filename or "").suffix.lower() or ".pptx"
    if suffix not in MASTER_SUFFIXES:
        raise MasterError("Upload a .pptx or .potx master deck.")
    if not data:
        raise MasterError("The file is empty.")
    if len(data) > MAX_MASTER_BYTES:
        raise MasterError("The master is too large.")
    if data[:2] != b"PK":
        raise MasterError("The file is not a .pptx.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "ppt/presentation.xml" not in archive.namelist():
                raise MasterError("The file is not a PowerPoint deck.")
    except zipfile.BadZipFile as exc:
        raise MasterError("The file is not a .pptx.") from exc
    return suffix


def import_uploaded_master(deck: DesignDeck, data: bytes, filename: str | None, *,
                           renderer: Renderer | None) -> JSON:
    """Make the upload the deck's master. A failed import keeps the master the deck had; slides on a
    layout the new master lacks move to its first layout (as new versions). Returns the master block
    plus `rehomed` (the slide ids moved)."""
    check_master_bytes(data, filename)
    previous_block = deck.master()
    previous_bytes = deck.master_path.read_bytes() if deck.master_path.is_file() else None
    incoming = deck.dir / ".incoming-master.pptx"
    incoming.write_bytes(data)
    try:
        block = engine_service.import_master(deck, incoming, Path(filename or "master.pptx").name, renderer=renderer,
                                             render=renderer is not None)
    except engine_service.MasterImportFailed:
        if previous_bytes is not None:
            deck.master_path.write_bytes(previous_bytes)
            deck.update(lambda m: m.update(master=previous_block or None))
        raise
    finally:
        incoming.unlink(missing_ok=True)
    moved = engine_service.rehome_slides(deck)
    return {**block, "rehomed": moved}


# ------------------------------------------------------------------------------ the storage port
async def persist_master(storage: Storage, ctx: CallerContext, brand_id: str, deck: DesignDeck, *, name: str,
                         idempotency_key: str | None = None) -> Master:
    """Keep the deck's imported master through the port: the record (with its manifest), the
    `.pptx`, and one PNG per layout (by layout index)."""
    block = deck.master()
    if not block.get("layouts"):
        raise MasterError("The deck has no imported master to keep.")
    manifest = {k: v for k, v in block.items() if k not in ("status", "filename", "importedAt")}
    master = await storage.masters.create(ctx, brand_id, MasterCreate(name=name[:200] or "Master", manifest=manifest),
                                          idempotency_key=idempotency_key)
    master = await storage.masters.put_file(ctx, master.id, "pptx", deck.master_path.read_bytes(), PPTX_TYPE)
    for index, layout in enumerate(block.get("layouts") or []):
        background = layout.get("background")
        png = deck.layouts_dir / Path(str(background)).name if background else None
        if png is not None and png.is_file():
            master = await storage.masters.put_file(ctx, master.id, "layout", png.read_bytes(), "image/png",
                                                    layout_index=index)
    return master
