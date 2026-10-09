"""Synthesise a PowerPoint master `.pptx` from `capturedFurniture` JSON.

Build a blank deck sized to the captured canvas, install the brand theme, and inject the captured
furniture onto the master (shared chrome) and three archetype layouts (Cover, Divider, Content), each
carrying its own furniture plus a heading placeholder (and, on the cover, a date placeholder) moved to
the captured box.

The result is a real master the engine's own `app.engine.importer.import_master` reads: there is no
bespoke manifest synthesis. A later pass can overwrite heading and date styles from the captured
resolved values; this module gets the geometry and structure in place.

Ported from Slide Studio `server/brand/synth_master.py` (migration plan §4.1, K).
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.core.brand import furniture_inject as fi

log = logging.getLogger(__name__)

#: Which blank-template layout each archetype repurposes, and the name to give it. Title Slide has a
#: centre title and a date (the cover's heading and date); Section Header and Title-and-Content both
#: carry a title placeholder for the divider and content heading.
ARCHETYPE_LAYOUT = {"cover": 0, "divider": 2, "content": 1}
ARCHETYPE_NAME = {"cover": "Cover", "divider": "Divider", "content": "Content"}

#: python-pptx placeholder type enum values (PP_PLACEHOLDER).
_TITLE_TYPES = {1, 3}  # TITLE, CENTER_TITLE
_DATE_TYPES = {16}  # DATE


def synthesize_master(captured: dict[str, Any], out_path: Path) -> Path:
    """Build a master `.pptx` at `out_path` from `captured`. Returns the path.

    Archetypes absent from the capture are skipped (their blank-template layouts keep the default
    names), so a capture with no layouts at all gives a clean default master at the captured size.
    """
    out_path = Path(out_path)
    w = int(captured["slideWidthEmu"])
    h = int(captured["slideHeightEmu"])

    prs: Any = Presentation()
    prs.slide_width = Emu(w)
    prs.slide_height = Emu(h)

    theme_ok = fi.apply_brand_theme(prs, captured)
    master = prs.slide_masters[0]
    media: dict[str, Any] = captured.get("media") or {}
    layouts: dict[str, Any] = captured.get("layouts") or {}

    master_done = False  # master-origin chrome is identical across archetypes: inject once
    summary: dict[str, dict[str, Any]] = {}
    for arch, layout_idx in ARCHETYPE_LAYOUT.items():
        blob = layouts.get(arch)
        if not blob:
            log.info("synth_master: archetype %r absent in capture; skipping", arch)
            continue
        layout = master.slide_layouts[layout_idx]
        _rename(layout, ARCHETYPE_NAME[arch])

        bg_ok = fi.apply_brand_background(prs, captured, layout, arch)

        shapes: list[dict[str, Any]] = blob.get("shapes") or []
        layout_shapes = [b for b in shapes if b.get("origin") != "master"]
        master_shapes = [b for b in shapes if b.get("origin") == "master"]

        n_layout = fi.inject_furniture_into_part(layout.part, layout.shapes._spTree, layout_shapes, media, w, h, w, h)
        n_master = 0
        if not master_done and master_shapes:
            n_master = fi.inject_furniture_into_part(
                master.part, master.shapes._spTree, master_shapes, media, w, h, w, h
            )
            master_done = True

        placed = _place_heading(layout, blob.get("placeholders") or {}, w, h)
        summary[arch] = {
            "layout_shapes": n_layout,
            "master_shapes": n_master,
            "background": bg_ok,
            "placeholders": placed,
        }

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    log.info("synth_master: wrote %s (theme=%s) %s", out_path.name, theme_ok, summary)
    return out_path


def _rename(layout: Any, name: str) -> None:
    c_sld = layout.element.find(qn("p:cSld"))
    if c_sld is not None:
        c_sld.set("name", name)


def _place_heading(layout: Any, placeholders: dict[str, Any], w: int, h: int) -> list[str]:
    """Move the layout's title (and cover date) placeholder to the captured box and style."""
    placed: list[str] = []
    head = placeholders.get("heading")
    if head:
        ph = _find_placeholder(layout, _TITLE_TYPES)
        if ph is not None and _style_placeholder(ph, head, w, h):
            placed.append("heading")
    date = placeholders.get("date")
    if date:
        ph = _find_placeholder(layout, _DATE_TYPES)
        if ph is not None and _style_placeholder(ph, date, w, h):
            placed.append("date")
    return placed


def _find_placeholder(layout: Any, type_values: set[int]) -> Any:
    for ph in layout.placeholders:
        kind = ph.placeholder_format.type
        if kind is not None and int(kind) in type_values:
            return ph
    return None


def _style_placeholder(ph: Any, box: dict[str, Any], w: int, h: int) -> bool:
    """Set a placeholder's geometry (fractions -> EMU) and default font, colour and size."""
    try:
        ph.left = Emu(int(round(float(box.get("left", 0)) * w)))
        ph.top = Emu(int(round(float(box.get("top", 0)) * h)))
        ph.width = Emu(int(round(float(box.get("width", 0)) * w)))
        ph.height = Emu(int(round(float(box.get("height", 0)) * h)))
        font = ph.text_frame.paragraphs[0].font
        if box.get("font"):
            font.name = str(box["font"])
        if box.get("size"):
            font.size = Pt(float(box["size"]))
        colour = box.get("color")
        if isinstance(colour, str) and colour.startswith("#") and len(colour) == 7:
            red, green, blue = bytes.fromhex(colour[1:])
            rgb: Any = RGBColor  # python-pptx leaves its constructor untyped
            font.color.rgb = rgb(red, green, blue)
        return True
    except Exception:  # noqa: BLE001 - an unstyled placeholder still leaves a usable master
        log.warning("synth_master: could not style placeholder", exc_info=True)
        return False
