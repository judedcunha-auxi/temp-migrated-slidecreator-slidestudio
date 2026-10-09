"""IR shapes, images and tables → native PowerPoint objects.

Native is the point: a rectangle is a rectangle with a fill you can change, a chevron is a
`custGeom` you can drag a vertex of, a donut is a `blockArc` with real adjustment handles. Nothing
here rasterises, and everything it cannot express exactly it records (`ctx.warn`) or flags as a
renderer gap (`ctx.gap`) rather than dropping in silence.

Three things in this file are easy to get wrong and are therefore pinned by the IR schema and by
`presetShapeDefinitions.xml`, not by taste:

* **Arc adjustments.** For `blockArc`/`pie`/`chord`, `adj1`/`adj2` are the start and end angle in
  60000ths of a degree clockwise from 3 o'clock, and `blockArc`'s `adj3` is the ring *thickness* as
  a fraction of `min(w, h)` — not the inner/outer radius ratio. python-pptx normalises adjustment
  values by 100000, so the angle goes in as `degrees × 0.6`.
* **Gradient angle.** DrawingML measures from 3 o'clock, CSS from 12: `ang = ((css − 90) mod 360) × 60000`.
* **Dashes.** `prstDash` lengths are multiples of the *line width*, so an IR dash in px is divided
  by the stroke width before it is compared with a preset.
"""
from __future__ import annotations

import math
from collections.abc import Callable, Sequence
from typing import Any, cast

from PIL import Image
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn

from app.engine.emit import text as text_engine
from app.engine.emit.draw import Draw
from app.engine.emit.template import (
    EmitContext,
    drop_shape_style,
    emu,
    emu_length,
    gradient_fill,
    hexval,
    inches,
    remove_all,
    set_alpha,
    sub,
)
from app.engine.ir import Box, Element

#: `prstDash` patterns in multiples of the line width (ECMA-376 §20.1.10.48).
PRESET_DASHES: dict[str, tuple[float, ...]] = {
    "dot": (1, 3),
    "sysDot": (1, 1),
    "dash": (4, 3),
    "sysDash": (3, 1),
    "lgDash": (8, 3),
    "dashDot": (4, 3, 1, 3),
    "sysDashDot": (3, 1, 1, 1),
    "lgDashDot": (8, 3, 1, 3),
    "lgDashDotDot": (8, 3, 1, 3, 1, 3),
    "sysDashDotDot": (3, 1, 1, 1, 1, 1),
}

#: How close a pattern must be to a preset to be written as one (master brief: 10 %).
DASH_TOLERANCE = 0.10

#: In the dash-length-only second pass, how far any other element may still be from the preset.
#: 2.0 = three times the preset's gap; past that the line reads as a different rhythm.
DASH_RHYTHM_LIMIT = 2.0

_CAPS = {"butt": "flat", "round": "rnd", "square": "sq"}
_JOINS = {"miter": "miter", "round": "round", "bevel": "bevel"}
_LINE_ENDS = {"triangle": "triangle", "stealth": "stealth", "oval": "oval",
              "diamond": "diamond", "arrow": "arrow", "none": "none"}

#: Bezier flattening resolution when a clip forces a curve into a polygon. 16 segments keeps the
#: error under a quarter pixel at slide scale, which the gate cannot see.
_FLATTEN_STEPS = 16

#: A quarter-circle's Bezier control-point distance, for corner rounding in custom geometry.
_KAPPA = 0.5522847498307936


# ------------------------------------------------------------------------------------ dispatch


def add_shape(container: Any, element: Element, ctx: EmitContext) -> Any | None:
    """One IR `shape` element → one PowerPoint shape (or connector). `None` when nothing was drawn."""
    # The clip can change the geometry (a clipped curve becomes a polygon), so the type is read
    # back afterwards rather than before.
    box, geometry, clipped = _apply_clip(element, dict(element.geometry or {}), ctx)
    if box is None:
        ctx.warn(element.id, "shape lies entirely outside its clip — dropped")
        return None
    gtype = geometry.get("type")

    if gtype == "line":
        shape = _add_line(container, element, geometry, box, ctx)
    elif gtype in ("rect", "roundRect", "ellipse", "blockArc", "pie", "chord"):
        shape = _add_preset(container, element, geometry, box, ctx)
    elif gtype in ("custom", "polyline"):
        shape = _add_custom(container, element, geometry, box, ctx)
    else:
        ctx.warn(element.id, f"unknown geometry type {gtype!r} — nothing emitted")
        return None
    if shape is None:
        return None

    drop_shape_style(shape)
    if gtype != "line":
        _apply_fill(shape, element, ctx)
    _apply_stroke(shape, element, ctx)
    _apply_effects(shape, element)
    _apply_transform(shape, element)
    name_shape(shape, element, ctx)
    if clipped:
        ctx.gap("a clipped curve was flattened to a polygon (overflow:hidden on a non-rect shape)")
    return shape


def name_shape(shape: Any, element: Element, ctx: EmitContext) -> None:
    """Name from `data-name`, else the source path, else the element id — never "Rectangle 7"."""
    if not ctx.options.name_shapes:
        return
    source = element.source or {}
    label = element.name or source.get("svg") or source.get("path") or element.id
    if label:
        shape.name = str(label)[:255]


# -------------------------------------------------------------------------------------- presets


def _add_preset(container: Any, element: Element, geometry: dict[str, Any], box: Box,
                ctx: EmitContext) -> Any:
    gtype = geometry["type"]
    preset = {
        "rect": MSO_SHAPE.RECTANGLE,
        "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
        "ellipse": MSO_SHAPE.OVAL,
        "blockArc": MSO_SHAPE.BLOCK_ARC,
        "pie": MSO_SHAPE.PIE,
        "chord": MSO_SHAPE.CHORD,
    }[gtype]

    if gtype == "roundRect" and not _radii_equal(geometry.get("radius")):
        return _add_custom(container, element, _rounded_rect_path(box, geometry.get("radius")), box, ctx)

    shape = container.add_shape(preset, emu_length(box.x), emu_length(box.y),
                                emu_length(max(box.w, 0.01)), emu_length(max(box.h, 0.01)))
    if gtype == "roundRect":
        radius = (geometry.get("radius") or {}).get("tl", 0.0)
        short = min(box.w, box.h)
        shape.adjustments[0] = max(0.0, min(0.5, float(radius) / short)) if short > 0 else 0.0
    elif gtype in ("blockArc", "pie", "chord"):
        _set_arc_adjustments(shape, gtype, geometry.get("arc") or {}, box, ctx, element)
    return shape


def _radii_equal(radius: dict[str, Any] | None) -> bool:
    if not radius:
        return True
    values = [float(radius.get(corner, 0.0) or 0.0) for corner in ("tl", "tr", "br", "bl")]
    return max(values) - min(values) <= 0.01


def _set_arc_adjustments(shape: Any, gtype: str, arc: dict[str, Any], box: Box,
                         ctx: EmitContext, element: Element) -> None:
    """`adj1`/`adj2` = start/end angle × 0.6; `blockArc`'s `adj3` = ring thickness / min(w, h)."""
    start = float(arc.get("startDeg") or 0.0) % 360.0
    end = float(arc.get("endDeg") or 0.0) % 360.0
    shape.adjustments[0] = start * 0.6
    shape.adjustments[1] = end * 0.6
    if gtype != "blockArc":
        return
    short = min(box.w, box.h)
    outer = float(arc.get("rOuter") or 0.0)
    inner = float(arc.get("rInner") or 0.0)
    if short <= 0 or outer <= 0:
        ctx.warn(element.id, "blockArc with no radius — ring thickness left at the preset default")
        return
    thickness = max(0.0, min(0.5, (outer - inner) / short))
    shape.adjustments[2] = thickness


# ---------------------------------------------------------------------------------- connectors


def _add_line(container: Any, element: Element, geometry: dict[str, Any], box: Box,
              ctx: EmitContext) -> Any | None:
    points = geometry.get("points") or []
    if len(points) < 2:
        ctx.warn(element.id, "line with fewer than 2 points — nothing emitted")
        return None
    (x1, y1), (x2, y2) = points[0], points[-1]
    return container.add_connector(
        MSO_CONNECTOR.STRAIGHT, emu_length(x1), emu_length(y1), emu_length(x2), emu_length(y2)
    )


# ------------------------------------------------------------------------------ custom geometry


def _add_custom(container: Any, element: Element, geometry: dict[str, Any], box: Box,
                ctx: EmitContext) -> Any | None:
    """A `custGeom` shape: a rectangle whose preset geometry is replaced by the IR's own path."""
    path = _geometry_path(geometry)
    if not path:
        ctx.warn(element.id, "custom geometry with no path — nothing emitted")
        return None
    shape = container.add_shape(MSO_SHAPE.RECTANGLE, emu_length(box.x), emu_length(box.y),
                                emu_length(max(box.w, 0.01)), emu_length(max(box.h, 0.01)))
    spPr = shape._element.spPr
    remove_all(spPr, "prstGeom", "custGeom")
    closed = _is_closed(path)
    spPr.insert(_geometry_position(spPr), _custom_geometry(
        spPr, path, box,
        even_odd=(geometry.get("fillRule") == "evenodd"),
        filled=closed and (element.fill or {}).get("type", "none") != "none",
    ))
    return shape


def _geometry_position(spPr: Any) -> int:
    """Geometry follows `a:xfrm` and precedes the fill — the schema's order, which PowerPoint enforces."""
    xfrm = spPr.find(qn("a:xfrm"))
    return list(spPr).index(xfrm) + 1 if xfrm is not None else 0


def _geometry_path(geometry: dict[str, Any]) -> list[list[Any]]:
    """Every geometry this module can express as a path, in absolute canvas px."""
    gtype = geometry.get("type")
    if gtype == "custom":
        return [list(segment) for segment in geometry.get("path") or []]
    if gtype == "polyline":
        points = geometry.get("points") or []
        if len(points) < 2:
            return []
        return [["M", float(points[0][0]), float(points[0][1])]] + [
            ["L", float(x), float(y)] for x, y in points[1:]
        ]
    return []


def _custom_geometry(parent: Any, path: list[list[Any]], box: Box, *, even_odd: bool,
                     filled: bool) -> Any:
    """`<a:custGeom>` for a path in absolute px, re-expressed in the box's own EMU space."""
    custGeom = parent.makeelement(qn("a:custGeom"), {})
    for tag in ("avLst", "gdLst", "ahLst", "cxnLst"):
        sub(custGeom, tag)
    width, height = max(emu(box.w), 1), max(emu(box.h), 1)
    sub(custGeom, "rect", l=0, t=0, r=width, b=height)
    pathLst = sub(custGeom, "pathLst")

    subpaths = _split_subpaths(path)
    # Even-odd is what makes a donut a donut: DrawingML applies it *within* one `a:path`, so all
    # subpaths share one element there, while nonzero subpaths each get their own.
    groups = [ [segment for sp in subpaths for segment in sp] ] if even_odd else subpaths
    for group in groups:
        if not group:
            continue
        element = sub(pathLst, "path", w=width, h=height,
                      fill="norm" if filled else "none", stroke=1, extrusionOk=0)
        for segment in group:
            _write_segment(element, segment, box)
    return custGeom


def _write_segment(path_element: Any, segment: Sequence[Any], box: Box) -> None:
    command = segment[0]
    coords = [float(v) for v in segment[1:]]

    def point(parent: Any, x: float, y: float) -> None:
        sub(parent, "pt", x=emu(x - box.x), y=emu(y - box.y))

    if command == "M":
        point(sub(path_element, "moveTo"), coords[0], coords[1])
    elif command == "L":
        point(sub(path_element, "lnTo"), coords[0], coords[1])
    elif command == "C":
        node = sub(path_element, "cubicBezTo")
        for index in range(0, 6, 2):
            point(node, coords[index], coords[index + 1])
    elif command == "Q":
        node = sub(path_element, "quadBezTo")
        for index in range(0, 4, 2):
            point(node, coords[index], coords[index + 1])
    elif command == "Z":
        sub(path_element, "close")


def _split_subpaths(path: list[list[Any]]) -> list[list[list[Any]]]:
    subpaths: list[list[list[Any]]] = []
    current: list[list[Any]] = []
    for segment in path:
        if segment[0] == "M" and current:
            subpaths.append(current)
            current = []
        current.append(list(segment))
    if current:
        subpaths.append(current)
    return subpaths


def _is_closed(path: list[list[Any]]) -> bool:
    return any(segment[0] == "Z" for segment in path)


def _rounded_rect_path(box: Box, radius: dict[str, Any] | None) -> dict[str, Any]:
    """A rectangle with four different corner radii, as a cubic path (no preset can express it)."""
    radius = radius or {}
    limit = min(box.w, box.h) / 2.0
    tl, tr, br, bl = (
        max(0.0, min(limit, float(radius.get(corner, 0.0) or 0.0)))
        for corner in ("tl", "tr", "br", "bl")
    )
    x, y, w, h = box.x, box.y, box.w, box.h
    k = _KAPPA
    path: list[list[Any]] = [["M", x + tl, y]]
    path.append(["L", x + w - tr, y])
    if tr:
        path.append(["C", x + w - tr + tr * k, y, x + w, y + tr - tr * k, x + w, y + tr])
    path.append(["L", x + w, y + h - br])
    if br:
        path.append(["C", x + w, y + h - br + br * k, x + w - br + br * k, y + h, x + w - br, y + h])
    path.append(["L", x + bl, y + h])
    if bl:
        path.append(["C", x + bl - bl * k, y + h, x, y + h - bl + bl * k, x, y + h - bl])
    path.append(["L", x, y + tl])
    if tl:
        path.append(["C", x, y + tl - tl * k, x + tl - tl * k, y, x + tl, y])
    path.append(["Z"])
    return {"type": "custom", "path": path}


# ------------------------------------------------------------------------------------- clipping


def _apply_clip(element: Element, geometry: dict[str, Any],
                ctx: EmitContext) -> tuple[Box | None, dict[str, Any], bool]:
    """Intersect the geometry with `element.clip`. Returns `(box, geometry, curves_flattened)`.

    A rectangle is clipped exactly by shrinking it. Anything else is flattened into a polygon and
    clipped against the rectangle, which is a real loss of curvature — hence the third return value
    and the renderer-gap note. Arcs keep their preset (a clipped `blockArc` with flattened edges is
    worse than a slightly oversized one) and say so.
    """
    box = element.box
    clip = element.clip
    if clip is None:
        return box, geometry, False
    inside = clip.x <= box.x + 0.01 and clip.y <= box.y + 0.01 and \
        clip.x2 >= box.x2 - 0.01 and clip.y2 >= box.y2 - 0.01
    if inside:
        return box, geometry, False

    intersection = box.intersect(clip)
    if intersection.w <= 0 or intersection.h <= 0:
        return None, geometry, False

    gtype = geometry.get("type")
    if gtype == "rect":
        return intersection, geometry, False
    if gtype in ("blockArc", "pie", "chord", "line"):
        ctx.warn(element.id, f"{gtype} overlaps a clip boundary — emitted unclipped (preset kept)")
        return box, geometry, False

    path = _geometry_path(geometry)
    if gtype == "roundRect":
        path = _rounded_rect_path(box, geometry.get("radius"))["path"]
    elif gtype == "ellipse":
        path = _ellipse_path(box)
    if not path:
        return intersection, geometry, False

    clipped: list[list[Any]] = []
    for subpath in _split_subpaths(path):
        polygon = _clip_polygon(_flatten(subpath), clip)
        if len(polygon) >= 3:
            clipped.append(["M", polygon[0][0], polygon[0][1]])
            clipped.extend(["L", x, y] for x, y in polygon[1:])
            clipped.append(["Z"])
    if not clipped:
        return None, geometry, False
    return intersection, {"type": "custom", "path": clipped,
                          "fillRule": geometry.get("fillRule")}, True


def _ellipse_path(box: Box) -> list[list[Any]]:
    rx, ry, cx, cy = box.w / 2.0, box.h / 2.0, box.cx, box.cy
    k = _KAPPA
    return [
        ["M", cx, cy - ry],
        ["C", cx + rx * k, cy - ry, cx + rx, cy - ry * k, cx + rx, cy],
        ["C", cx + rx, cy + ry * k, cx + rx * k, cy + ry, cx, cy + ry],
        ["C", cx - rx * k, cy + ry, cx - rx, cy + ry * k, cx - rx, cy],
        ["C", cx - rx, cy - ry * k, cx - rx * k, cy - ry, cx, cy - ry],
        ["Z"],
    ]


def _flatten(subpath: list[list[Any]]) -> list[tuple[float, float]]:
    """A subpath as a point list, curves sampled — the input to polygon clipping."""
    points: list[tuple[float, float]] = []
    cursor = (0.0, 0.0)
    for segment in subpath:
        command, coords = segment[0], [float(v) for v in segment[1:]]
        if command == "M":
            cursor = (coords[0], coords[1])
            points.append(cursor)
        elif command == "L":
            cursor = (coords[0], coords[1])
            points.append(cursor)
        elif command == "C":
            for step in range(1, _FLATTEN_STEPS + 1):
                t = step / _FLATTEN_STEPS
                points.append(_cubic(cursor, coords[0:2], coords[2:4], coords[4:6], t))
            cursor = (coords[4], coords[5])
        elif command == "Q":
            for step in range(1, _FLATTEN_STEPS + 1):
                t = step / _FLATTEN_STEPS
                control = (cursor[0] + 2 / 3 * (coords[0] - cursor[0]),
                           cursor[1] + 2 / 3 * (coords[1] - cursor[1]))
                second = (coords[2] + 2 / 3 * (coords[0] - coords[2]),
                          coords[3] + 2 / 3 * (coords[1] - coords[3]))
                points.append(_cubic(cursor, control, second, coords[2:4], t))
            cursor = (coords[2], coords[3])
    return points


def _cubic(p0: Sequence[float], p1: Sequence[float], p2: Sequence[float], p3: Sequence[float],
           t: float) -> tuple[float, float]:
    u = 1.0 - t
    a, b, c, d = u * u * u, 3 * u * u * t, 3 * u * t * t, t * t * t
    return (a * p0[0] + b * p1[0] + c * p2[0] + d * p3[0],
            a * p0[1] + b * p1[1] + c * p2[1] + d * p3[1])


def _clip_polygon(points: list[tuple[float, float]], clip: Box) -> list[tuple[float, float]]:
    """Sutherland–Hodgman against a rectangle (convex, so one polygon in, one polygon out)."""
    edges: tuple[tuple[Callable[[Any], bool], Callable[[Any, Any], tuple[float, float]]], ...] = (
        (lambda p: p[0] >= clip.x, lambda a, b: _cut(a, b, 0, clip.x)),
        (lambda p: p[0] <= clip.x2, lambda a, b: _cut(a, b, 0, clip.x2)),
        (lambda p: p[1] >= clip.y, lambda a, b: _cut(a, b, 1, clip.y)),
        (lambda p: p[1] <= clip.y2, lambda a, b: _cut(a, b, 1, clip.y2)),
    )
    polygon = list(points)
    for inside, intersect in edges:
        if not polygon:
            return []
        output: list[tuple[float, float]] = []
        previous = polygon[-1]
        for current in polygon:
            if inside(current):
                if not inside(previous):
                    output.append(intersect(previous, current))
                output.append(current)
            elif inside(previous):
                output.append(intersect(previous, current))
            previous = current
        polygon = output
    return polygon


def _cut(a: tuple[float, float], b: tuple[float, float], axis: int, value: float) -> tuple[float, float]:
    span = b[axis] - a[axis]
    t = 0.0 if abs(span) < 1e-9 else (value - a[axis]) / span
    return (a[0] + (b[0] - a[0]) * t, a[1] + (b[1] - a[1]) * t)


# ------------------------------------------------------------------------------- fill and stroke


def _apply_fill(shape: Any, element: Element, ctx: EmitContext) -> None:
    spPr = shape._element.spPr
    remove_all(spPr, "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill")
    fill = element.fill or {"type": "none"}
    opacity = float(element.opacity if element.opacity is not None else 1.0)
    position = _fill_position(spPr)

    if fill.get("type") == "solid":
        node = spPr.makeelement(qn("a:solidFill"), {})
        spPr.insert(position, node)
        colour = sub(node, "srgbClr", val=hexval(fill.get("color")))
        set_alpha(colour, float(fill.get("alpha", 1.0)) * opacity)
    elif fill.get("type") == "gradient":
        node = gradient_fill(spPr, fill, opacity)
        spPr.remove(node)
        spPr.insert(position, node)
    else:
        spPr.insert(position, spPr.makeelement(qn("a:noFill"), {}))


def _fill_position(spPr: Any) -> int:
    for tag in ("a:custGeom", "a:prstGeom", "a:xfrm"):
        found = spPr.find(qn(tag))
        if found is not None:
            return list(spPr).index(found) + 1
    return 0


def _apply_stroke(shape: Any, element: Element, ctx: EmitContext) -> None:
    """`<a:ln>` built by hand: its children are order-sensitive and python-pptx exposes half of them."""
    spPr = shape._element.spPr
    remove_all(spPr, "ln")
    stroke = element.stroke
    position = _line_position(spPr)
    line = spPr.makeelement(qn("a:ln"), {})
    spPr.insert(position, line)

    if not stroke or float(stroke.get("width") or 0) <= 0:
        sub(line, "noFill")
        return

    width_px = float(stroke["width"])
    line.set("w", str(max(1, emu(width_px))))
    cap = _CAPS.get(str(stroke.get("cap") or "butt"))
    if cap:
        line.set("cap", cap)
    colour = sub(sub(line, "solidFill"), "srgbClr", val=hexval(stroke.get("color")))
    # An opacity of 0 is 0, as in `_apply_fill` (r1e): `or 1.0` wrote an opacity-0 box's border opaque.
    opacity = 1.0 if element.opacity is None else float(element.opacity)
    set_alpha(colour, float(stroke.get("alpha", 1.0)) * opacity)

    _write_dash(line, stroke.get("dash"), width_px, ctx)
    join = _JOINS.get(str(stroke.get("join") or "miter"))
    if join:
        sub(line, join, **({"lim": 800000} if join == "miter" else {}))
    for side in ("headEnd", "tailEnd"):
        end = stroke.get(side)
        if end and _LINE_ENDS.get(str(end.get("type"))):
            size = str(end.get("size") or "med")
            sub(line, side, type=_LINE_ENDS[str(end["type"])], w=size, len=size)


def _line_position(spPr: Any) -> int:
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill",
                "a:custGeom", "a:prstGeom", "a:xfrm"):
        found = spPr.find(qn(tag))
        if found is not None:
            return list(spPr).index(found) + 1
    return 0


def _write_dash(line: Any, dash: Any, width_px: float, ctx: EmitContext) -> None:
    if not dash or dash == "solid":
        return
    if isinstance(dash, str):
        sub(line, "prstDash", val=dash)
        return
    pattern = [float(value) for value in dash if float(value) >= 0]
    if not pattern:
        return
    preset = match_preset_dash(pattern, width_px)
    if preset:
        sub(line, "prstDash", val=preset)
        return
    # `custDash` lengths are thousandths of a percent of the line width.
    custom = sub(line, "custDash")
    pairs = pattern + [pattern[-1]] if len(pattern) % 2 else pattern
    for index in range(0, len(pairs), 2):
        sub(custom, "ds",
            d=max(1, int(round(pairs[index] / max(width_px, 0.01) * 100000))),
            sp=max(1, int(round(pairs[index + 1] / max(width_px, 0.01) * 100000))))
    ctx.gap("custDash is drawn solid by PptxRender (renderer gap, the file is correct)")


def match_preset_dash(pattern: Sequence[float], width_px: float) -> str | None:
    """The `prstDash` this pattern is, or `None`.

    Two passes, because the eye reads a dashed line by its dash length long before its gap: first
    the whole pattern within 10 %, then — for patterns no preset matches as a whole — the dash
    length alone within 10 % among presets of the same shape. A 30 px dash at 1 px width matches
    nothing either way and becomes a `custDash`, which is right: it is not a preset.
    """
    width = max(float(width_px), 0.01)
    units = [value / width for value in pattern]
    if not units:
        return None

    best: tuple[float, str] | None = None
    for name, preset in PRESET_DASHES.items():
        if len(preset) != len(units):
            continue
        error = max(abs(u - p) / p for u, p in zip(units, preset, strict=False))
        if best is None or error < best[0]:
            best = (error, name)
    if best and best[0] <= DASH_TOLERANCE:
        return best[1]

    on_best: tuple[float, str] | None = None
    for name, preset in PRESET_DASHES.items():
        if len(preset) != len(units):
            continue
        # The gaps may differ — but only so far. Three times the preset's gap is a different
        # rhythm, not the same dash drawn loosely, and belongs in a custDash.
        if max(abs(u - p) / p for u, p in zip(units, preset, strict=False)) > DASH_RHYTHM_LIMIT:
            continue
        error = abs(units[0] - preset[0]) / preset[0]
        if on_best is None or error < on_best[0]:
            on_best = (error, name)
    if on_best and on_best[0] <= DASH_TOLERANCE:
        return on_best[1]
    return None


def _apply_effects(shape: Any, element: Element) -> None:
    """Exactly one `effectLst`, always — an empty one is what stops the theme's own shadow."""
    spPr = shape._element.spPr
    remove_all(spPr, "effectLst")
    effects = sub(spPr, "effectLst")
    shadow = element.shadow
    if not shadow:
        return
    dx, dy = float(shadow.get("dx") or 0.0), float(shadow.get("dy") or 0.0)
    node = sub(
        effects, "outerShdw",
        blurRad=max(0, emu(float(shadow.get("blur") or 0.0))),
        dist=max(0, emu(math.hypot(dx, dy))),
        dir=int(round(math.degrees(math.atan2(dy, dx)) % 360 * 60000)),
        rotWithShape=0,
    )
    # An opacity of 0 is 0, as in `_apply_fill` (r1e): `or 1.0` wrote an opacity-0 box's shadow opaque.
    opacity = 1.0 if element.opacity is None else float(element.opacity)
    set_alpha(sub(node, "srgbClr", val=hexval(shadow.get("color"), "000000")),
              float(shadow.get("alpha", 1.0)) * opacity)


def _apply_transform(shape: Any, element: Element) -> None:
    if element.rotation:
        shape.rotation = float(element.rotation)
    if not (element.flipH or element.flipV):
        return
    xfrm = shape._element.spPr.find(qn("a:xfrm"))
    if xfrm is None:
        return
    if element.flipH:
        xfrm.set("flipH", "1")
    if element.flipV:
        xfrm.set("flipV", "1")


# --------------------------------------------------------------------------------------- images


def add_image(container: Any, element: Element, ctx: EmitContext) -> Any | None:
    """An `image` (or a `raster`) as a real picture, cropped rather than resampled.

    The picture's `a:srcRect` is the IR's `crop`, written as given: the extractor measured it from
    the browser's own geometry (`page.js` `cutImage`: the rectangle the whole source is drawn into,
    cut by its frame), so the frame shows exactly the part of the source the browser shows. `fit`
    only fills in a crop the IR does not give (`crop` null — hand-written IR); fitting a crop that
    is already fitted cut a covering image twice — a portrait photo's middle sixth stretched over its
    frame, a `right bottom` crop taken from the middle (r3, plan §16 #30). A negative side is a band
    of the frame the source does not reach: PowerPoint's own Crop → Fit, drawn empty (measured).
    """
    if not element.src:
        ctx.warn(element.id, "image without a src — nothing emitted")
        return None
    box = element.box
    crop = dict(element.crop or {})
    intrinsic = _intrinsic_size(str(element.src)) if element.crop is None else None

    if element.fit in ("cover", "contain") and intrinsic:
        box, crop = _fit_box(box, crop, intrinsic, str(element.fit))
    if element.clip is not None:
        clipped, crop = _crop_to_clip(box, crop, element.clip)
        if clipped is None:
            ctx.warn(element.id, "image lies entirely outside its clip — dropped")
            return None
        box = clipped

    picture = container.add_picture(
        str(element.src), emu_length(box.x), emu_length(box.y),
        emu_length(max(box.w, 0.01)), emu_length(max(box.h, 0.01)),
    )
    for side in ("l", "t", "r", "b"):
        value = float(crop.get(side) or 0.0)
        if value:
            setattr(picture, {"l": "crop_left", "t": "crop_top", "r": "crop_right",
                              "b": "crop_bottom"}[side], value)

    if element.circle or (element.radius and element.radius > 0):
        geometry = picture._element.spPr.find(qn("a:prstGeom"))
        if geometry is not None:
            if element.circle:
                geometry.set("prst", "ellipse")
            else:
                geometry.set("prst", "roundRect")
                short = min(box.w, box.h)
                adjustments = geometry.find(qn("a:avLst"))
                remove_all(geometry, "avLst")
                adjustments = sub(geometry, "avLst")
                sub(adjustments, "gd", name="adj",
                    fmla=f"val {int(round(max(0.0, min(0.5, element.radius / short)) * 100000))}")
    if element.opacity is not None and element.opacity < 1.0:
        _picture_alpha(picture, float(element.opacity))
    if element.rotation:
        picture.rotation = float(element.rotation)
    flips = (element.extras or {}).get("x-wpf-flip") or {}
    xfrm = picture._element.spPr.find(qn("a:xfrm"))
    if xfrm is not None:
        for key in ("flipH", "flipV"):
            if flips.get(key):
                xfrm.set(key, "1")
    name_shape(picture, element, ctx)
    return picture


def _intrinsic_size(src: str) -> tuple[int, int] | None:
    try:
        with Image.open(src) as image:
            return image.size
    except Exception:
        return None


def _fit_box(box: Box, crop: dict[str, float], intrinsic: tuple[int, int],
             fit: str) -> tuple[Box, dict[str, float]]:
    """`object-fit` without resampling: `cover` becomes a crop, `contain` becomes a smaller box."""
    iw, ih = intrinsic
    if iw <= 0 or ih <= 0 or box.w <= 0 or box.h <= 0:
        return box, crop
    box_ratio, image_ratio = box.w / box.h, iw / ih
    if abs(box_ratio - image_ratio) < 1e-4:
        return box, crop
    if fit == "cover":
        if image_ratio > box_ratio:                       # image too wide: trim left and right
            keep = box_ratio / image_ratio
            side = (1.0 - keep) / 2.0
            crop = {**crop, "l": crop.get("l", 0.0) + side, "r": crop.get("r", 0.0) + side}
        else:
            keep = image_ratio / box_ratio
            side = (1.0 - keep) / 2.0
            crop = {**crop, "t": crop.get("t", 0.0) + side, "b": crop.get("b", 0.0) + side}
        return box, crop
    if image_ratio > box_ratio:                           # contain: letterbox inside the box
        height = box.w / image_ratio
        return Box(box.x, box.y + (box.h - height) / 2.0, box.w, height), crop
    width = box.h * image_ratio
    return Box(box.x + (box.w - width) / 2.0, box.y, width, box.h), crop


def _crop_to_clip(box: Box, crop: dict[str, float], clip: Box) -> tuple[Box | None, dict[str, float]]:
    """`overflow:hidden` on a picture is a crop, not a resize: the visible pixels keep their scale."""
    visible = box.intersect(clip)
    if visible.w <= 0 or visible.h <= 0:
        return None, crop
    if visible.w >= box.w - 0.01 and visible.h >= box.h - 0.01:
        return box, crop
    remaining_w = 1.0 - float(crop.get("l", 0.0)) - float(crop.get("r", 0.0))
    remaining_h = 1.0 - float(crop.get("t", 0.0)) - float(crop.get("b", 0.0))
    return visible, {
        "l": float(crop.get("l", 0.0)) + (visible.x - box.x) / box.w * remaining_w,
        "r": float(crop.get("r", 0.0)) + (box.x2 - visible.x2) / box.w * remaining_w,
        "t": float(crop.get("t", 0.0)) + (visible.y - box.y) / box.h * remaining_h,
        "b": float(crop.get("b", 0.0)) + (box.y2 - visible.y2) / box.h * remaining_h,
    }


def _picture_alpha(picture: Any, opacity: float) -> None:
    """`alphaModFix` inside the picture's `blipFill` — python-pptx has no opacity for images."""
    blip = picture._element.blipFill.find(qn("a:blip"))
    if blip is None:
        return
    remove_all(blip, "alphaModFix")
    sub(blip, "alphaModFix", amt=int(round(max(0.0, min(1.0, opacity)) * 100000)))


# --------------------------------------------------------------------------------------- tables


def add_table(slide: Any, element: Element, ctx: EmitContext) -> Any | None:
    """A real PowerPoint table: StageFlow's `Draw.table` for the grid, our text engine for the cells.

    `Draw.table` already knows the two things that are easy to get wrong — merging before writing,
    and "No Style, No Grid" so PowerPoint's banded blue default does not fight the design — but it
    writes one run per cell. Each cell is then finished here (WP-A, brief §6.1): dashed rules, the
    IR's own paragraphs through the text engine, the margins that put the first baseline where the
    browser drew it, and — for a cell whose content is all drawn over the table — a 1 pt empty
    paragraph, so an overlay-only cell never grows its row.
    """
    rows, cols = int(element.rows or 0), int(element.cols or 0)
    if rows <= 0 or cols <= 0 or not element.cells:
        ctx.warn(element.id, "table with no grid — nothing emitted")
        return None

    widths = list(element.colWidthsPx or [element.box.w / cols] * cols)
    heights = list(element.rowHeightsPx or [element.box.h / rows] * rows)
    cells = [_table_cell_spec(cell) for cell in element.cells]
    frame = Draw(slide, None).table(
        inches(element.box.x), inches(element.box.y), inches(element.box.w), inches(element.box.h),
        cells, rows, cols,
        [inches(w) for w in widths],
        [inches(h) for h in heights],
    )
    table = frame.table
    wpa = ((element.extras or {}).get("x-wpa") or {}).get("cells") or {}
    for cell, spec in zip(element.cells, cells, strict=False):
        native = table.cell(spec["r"], spec["c"])
        extras = wpa.get(f"{spec['r']}:{spec['c']}") or {}
        _cell_dashes(native._tc.get_or_add_tcPr(), cell, ctx)
        padding = cell.get("paddingPx") or {}
        slot = {side: max(0.0, float((extras.get("slot") or {}).get(side) or 0.0)) for side in ("t", "r", "b", "l")}
        span_width = sum(widths[spec["c"]:spec["c"] + spec["cspan"]])
        inner_width = (span_width - float(padding.get("l") or 0.0) - float(padding.get("r") or 0.0)
                       - slot["l"] - slot["r"])
        slot_top = float(element.box.y) + sum(heights[:spec["r"]])
        slot_height = sum(heights[spec["r"]:spec["r"] + spec["rspan"]])
        placement = {"slot_top_px": slot_top, "slot_height_px": slot_height,
                     "top_px": float(padding.get("t") or 0.0) + slot["t"],
                     "bottom_px": float(padding.get("b") or 0.0) + slot["b"],
                     "anchor": str(cell.get("valign") or "top")}
        layout = None
        if cell.get("paragraphs"):
            layout = text_engine.write_cell(native.text_frame, cell, element, ctx,
                                            inner_width_px=inner_width, extras=extras, placement=placement)
            for warning in layout.warnings:
                ctx.warn(element.id, warning)
        if layout is None or not layout.paragraphs:
            _empty_cell_paragraph(native.text_frame)
            layout = None
        residual = _cell_margins(native, cell, slot, layout, slot_top, slot_height)
        if abs(residual) > 0.5:
            ctx.warn(element.id, f"cell {spec['r']},{spec['c']}: the first baseline needs a margin beyond the "
                                 f"cell's insets; the text sits {abs(residual):.1f} px "
                                 f"{'low' if residual > 0 else 'high'}")
    name_shape(frame, element, ctx)
    return frame


def _cell_margins(native: Any, cell: dict[str, Any], slot: dict[str, float],
                  layout: text_engine.CellLayout | None, slot_top: float, slot_height: float) -> float:
    """The cell's insets: CSS padding + `x-wpa` slot (grid line → padding box), with the text moved
    up by Δ as `text.cell_margins_px` rules (given the cell's grid slot, so the measured first line —
    not a model of it — is the target). Returns the signed residual (positive: the text sits low).
    """
    padding = cell.get("paddingPx") or {}
    margin_top, margin_bottom, residual = text_engine.cell_margins_px(
        layout, slot_top_px=slot_top, slot_height_px=slot_height,
        top_px=float(padding.get("t") or 0.0) + slot["t"], bottom_px=float(padding.get("b") or 0.0) + slot["b"],
        anchor=str(cell.get("valign") or "top"))
    native.margin_left = emu_length(float(padding.get("l") or 0.0) + slot["l"])
    native.margin_right = emu_length(float(padding.get("r") or 0.0) + slot["r"])
    native.margin_top = emu_length(margin_top)
    native.margin_bottom = emu_length(margin_bottom)
    return residual


def _cell_dashes(tcPr: Any, cell: dict[str, Any], ctx: EmitContext) -> None:
    """A dashed or dotted cell rule stays dashed: `prstDash`/`custDash` inside the `ln*` Draw wrote.

    `lnL`…`lnB` are `CT_LineProperties` like `a:ln`, so the dash follows the fill child; the tcPr
    children are then put back in schema order.
    """
    from app.engine.emit.draw import _order_tc

    tags = {"left": "lnL", "right": "lnR", "top": "lnT", "bottom": "lnB"}
    for side, spec in (cell.get("borders") or {}).items():
        if not spec or spec.get("dash") in (None, "solid") or side not in tags:
            continue
        line = tcPr.find(qn(f"a:{tags[side]}"))
        if line is not None:
            _write_dash(line, spec.get("dash"), float(spec.get("width") or 1.0), ctx)
    _order_tc(tcPr)


def _empty_cell_paragraph(frame: Any) -> None:
    """A cell whose content is all drawn over the table: one empty 1 pt paragraph.

    `Draw.table` leaves an empty run at 11 pt, and PowerPoint grows a row to fit it. An explicit
    `endParaRPr sz="100"` with a 1 pt line is the only empty-cell size PptxRender grows a row for
    (1.2 × sz + margins, `SlideRenderer.Table.cs` 81–96), and PowerPoint kept a 10 px row declared
    this way at 10 px (probe (d), 2026-09-25) — so an overlay-only cell never grows its row.
    """
    body = frame._txBody
    for paragraph in list(body.findall(qn("a:p"))):
        body.remove(paragraph)
    paragraph = body.makeelement(qn("a:p"), {})
    body.append(paragraph)
    properties = sub(paragraph, "pPr")
    sub(sub(properties, "lnSpc"), "spcPts", val=100)
    sub(paragraph, "endParaRPr", sz=100)


def _table_cell_spec(cell: dict[str, Any]) -> dict[str, Any]:
    """One IR cell in the shape `Draw.table` expects (px → inches, pt for borders and text)."""
    padding = cell.get("paddingPx") or {}
    first_run = _first_run(cell)
    borders = {}
    for side, spec in (cell.get("borders") or {}).items():
        if spec and float(spec.get("width") or 0) > 0:
            borders[side] = {"w": float(spec["width"]) * 0.75, "hex": hexval(spec.get("color"))}
    fill = cell.get("fill") or {}
    return {
        "r": int(cell.get("r") or 0), "c": int(cell.get("c") or 0),
        "rspan": int(cell.get("rowSpan") or 1), "cspan": int(cell.get("colSpan") or 1),
        "padLIn": inches(float(padding.get("l") or 0.0)), "padRIn": inches(float(padding.get("r") or 0.0)),
        "padTIn": inches(float(padding.get("t") or 0.0)), "padBIn": inches(float(padding.get("b") or 0.0)),
        "valign": cell.get("valign") or "top",
        "fill": hexval(fill.get("color")) if fill.get("type") == "solid" else None,
        "fillAlpha": float(fill.get("alpha", 1.0)) if fill.get("type") == "solid" else 1.0,
        "align": (cell.get("paragraphs") or [{}])[0].get("align") or "left",
        "text": str(first_run.get("text") or ""),
        "sizePt": float(first_run.get("sizePx") or 14.667) * 0.75,
        "weight": int(first_run.get("weight") or 400),
        "italic": bool(first_run.get("italic")),
        "font": first_run.get("font"),
        "color": hexval(first_run.get("color"), "000000"),
        "bordersPt": borders,
    }


def _first_run(cell: dict[str, Any]) -> dict[str, Any]:
    for paragraph in cell.get("paragraphs") or []:
        for line in paragraph.get("lines") or []:
            for run in line.get("runs") or []:
                return cast(dict[str, Any], run)
    return {}
