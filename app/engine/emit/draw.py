"""A small Draw over python-pptx, plus the lxml patches for what it does not expose.

LICENCE-GATED (decision D3). This module is StageFlow's helper layer, copied verbatim (third-party
code whose licence terms have not been granted yet). It is ported as-is so the engine works, but it
must not ship to production until D3 is resolved: either terms are granted, or it is rewritten
(risk R7). See docs/licensing.md. Do not copy more StageFlow code into the service.

Deliberately per-ELEMENT, never per-line: a paragraph goes in as a paragraph and PowerPoint wraps it.
That is the whole difference from SlideForge's emitter, which pins every rendered line to its measured
rect. Reflowable here, pixel-exact there.
"""
# ---------------------------------------------------------------------------------------------
# Vendored from StageFlow (third-party; licence-gated, decision D3), 2026-09-18.
# Copied verbatim and imported as-is (it needs only lxml and python-pptx).
# ---------------------------------------------------------------------------------------------

from __future__ import annotations

from typing import Any, cast

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _hex(c: str) -> RGBColor:
    return cast(RGBColor, RGBColor.from_string((c or "000000").lstrip("#")[:6].upper()))  # type: ignore[no-untyped-call]


def _sub(parent: Any, tag: str, **attrs: Any) -> Any:
    el = etree.SubElement(parent, f"{{{A}}}{tag}")
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _order_tc(tcPr: Any) -> None:
    """CT_TableCellProperties: lnL, lnR, lnT, lnB, lnTlToBr, lnBlToTr, cell3D, then the fill."""
    order = ["lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D"]
    listed: list[Any] = []
    for n in order:
        listed += tcPr.findall(f"{{{A}}}{n}")
    rest = [c for c in tcPr if c not in listed]
    for c in list(tcPr):
        tcPr.remove(c)
    for c in listed + rest:
        tcPr.append(c)


ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
         "justify": PP_ALIGN.JUSTIFY, "start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


class Draw:
    """All geometry in INCHES — the mapper has already converted from stage px."""

    def __init__(self, slide: Any, mapper: Any) -> None:
        self.slide = slide
        self.m = mapper

    # ---------------------------------------------------------------- shapes
    def rect(self, x: float, y: float, w: float, h: float, fill: str | None = None, alpha: float = 1.0,
             radius: float = 0.0, line: str | None = None, line_w: float = 0.0, line_alpha: float = 1.0,
             shadow: bool = False) -> Any:
        shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE
        shp = self.slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.shadow.inherit = False

        if radius > 0 and min(w, h) > 0:
            # PowerPoint's rounded-rect adjustment is a fraction of the SHORT side, not an absolute.
            adj = max(0.0, min(0.5, radius / min(w, h)))
            shp.adjustments[0] = adj

        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = _hex(fill)
            if alpha < 1.0:
                self._alpha(shp.fill.fore_color._xFill.find(f"{{{A}}}srgbClr"), alpha)
        else:
            shp.fill.background()

        if line and line_w > 0:
            shp.line.color.rgb = _hex(line)
            shp.line.width = Emu(int(Inches(line_w)))
            if line_alpha < 1.0:
                self._alpha(shp.line._get_or_add_ln().find(f"{{{A}}}solidFill/{{{A}}}srgbClr"), line_alpha)
        else:
            shp.line.fill.background()

        if shadow:
            self.soft_shadow(shp)
        shp.text_frame.text = ""
        return shp

    def text(self, x: float, y: float, w: float, h: float, runs: Any, size_pt: float = 12.0,
             color: str = "000000", bold: bool = False, italic: bool = False, align: str = "left",
             anchor: str = "top", line_pt: float | None = None, font: str = "Calibri", wrap: bool = True,
             space_pt: float = 0.0) -> Any:
        """`runs` is a str, or a list of (text, overrides) for mixed formatting within one paragraph."""
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ANCHOR.get(anchor, MSO_ANCHOR.TOP)

        # A plain string with line breaks (<br>, stacked bullets) becomes one paragraph per line;
        # a literal "\n" inside a run is not a line break to PowerPoint.
        lines = [[(line, {})] for line in runs.split("\n")] if isinstance(runs, str) else [runs]
        for n, line in enumerate(lines):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
            if line_pt:
                p.line_spacing = Pt(line_pt)
            for text, over in line:
                r = p.add_run()
                r.text = text
                f = r.font
                f.size = Pt(over.get("size_pt", size_pt))
                f.bold = over.get("bold", bold)
                f.italic = over.get("italic", italic)
                f.name = over.get("font", font)
                f.color.rgb = _hex(over.get("color", color))
                sp = over.get("space_pt", space_pt)
                if sp:
                    r._r.get_or_add_rPr().set("spc", str(int(round(sp * 100))))
        return box

    # ---------------------------------------------------------------- table
    # "No Style, No Grid": PowerPoint's default table style paints its own banded blue look, which
    # would fight the deck. Start from nothing and apply only what the HTML actually specified.
    NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

    VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM,
              "baseline": MSO_ANCHOR.BOTTOM, "sub": MSO_ANCHOR.BOTTOM, "super": MSO_ANCHOR.TOP}

    def table(self, x: float, y: float, w: float, h: float, cells: Any, n_rows: int, n_cols: int,
              col_w: Any, row_h: Any, default_font: str = "Calibri") -> Any:
        """
        `cells` carry their grid coordinates and spans, so colspan/rowspan become real merged cells
        rather than the text landing in the wrong column.
        """
        gf = self.slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(h))
        tbl = gf.table

        tblPr = tbl._tbl.find(f"{{{A}}}tblPr")
        if tblPr is not None:
            tblPr.set("firstRow", "0")
            tblPr.set("bandRow", "0")
            for old in tblPr.findall(f"{{{A}}}tableStyleId"):
                tblPr.remove(old)
            etree.SubElement(tblPr, f"{{{A}}}tableStyleId").text = self.NO_STYLE

        for i, cw in enumerate(col_w[:n_cols]):
            tbl.columns[i].width = Emu(max(1, int(Inches(cw))))
        for i, rh in enumerate(row_h[:n_rows]):
            tbl.rows[i].height = Emu(max(1, int(Inches(rh))))

        # Merge first: merging clears the absorbed cells, so text written before would be lost.
        for c in cells:
            if c["cspan"] > 1 or c["rspan"] > 1:
                r0, c0 = c["r"], c["c"]
                r1 = min(n_rows - 1, r0 + c["rspan"] - 1)
                c1 = min(n_cols - 1, c0 + c["cspan"] - 1)
                if (r1, c1) != (r0, c0):
                    try:
                        tbl.cell(r0, c0).merge(tbl.cell(r1, c1))
                    except Exception:  # nosec B110
                        # ^ an overlapping merge python-pptx refuses stays unmerged
                        pass

        painted = set()
        for c in cells:
            cell = tbl.cell(c["r"], c["c"])
            painted.add((c["r"], c["c"]))
            cell.margin_left = Emu(int(Inches(c.get("padLIn", 0.06))))
            cell.margin_right = Emu(int(Inches(c.get("padRIn", 0.06))))
            cell.margin_top = Emu(int(Inches(c.get("padTIn", 0.03))))
            cell.margin_bottom = Emu(int(Inches(c.get("padBIn", 0.03))))
            cell.vertical_anchor = self.VALIGN.get(c.get("valign", "middle"), MSO_ANCHOR.MIDDLE)

            if c.get("fill"):
                cell.fill.solid()
                cell.fill.fore_color.rgb = _hex(c["fill"])
                if c.get("fillAlpha", 1.0) < 1.0:
                    self._alpha(cell.fill.fore_color._xFill.find(f"{{{A}}}srgbClr"), c["fillAlpha"])
            else:
                cell.fill.background()

            tf = cell.text_frame
            tf.word_wrap = bool(c.get("wrap", True))
            p = tf.paragraphs[0]
            p.alignment = ALIGN.get(c.get("align", "left"), PP_ALIGN.LEFT)
            run = p.add_run()
            run.text = c.get("text", "")
            f = run.font
            f.size = Pt(c["sizePt"])
            f.bold = c.get("weight", 400) >= 600 or c.get("head", False)
            f.italic = c.get("italic", False)
            f.name = c.get("font") or default_font
            f.color.rgb = _hex(c.get("color", "000000"))

            self._cell_borders(cell, c.get("bordersPt") or {})

        # Cells the grid has but the HTML never declared (ragged rows) still need a sane look.
        for r in range(n_rows):
            for cc in range(n_cols):
                if (r, cc) not in painted:
                    try:
                        tbl.cell(r, cc).fill.background()
                    except Exception:  # nosec B110
                        # ^ a spanned cell has no fill of its own to clear
                        pass
        return gf

    # tcPr children are order-sensitive: the four line elements come first, then the fill.
    BORDER_TAGS = {"left": "lnL", "right": "lnR", "top": "lnT", "bottom": "lnB"}

    def _cell_borders(self, cell: Any, borders: Any) -> None:
        """Per-side cell borders — python-pptx exposes none of them."""
        if not borders:
            return
        tcPr = cell._tc.get_or_add_tcPr()
        for side in ("left", "right", "top", "bottom"):
            tag = self.BORDER_TAGS[side]
            for old in tcPr.findall(f"{{{A}}}{tag}"):
                tcPr.remove(old)
            spec = borders.get(side)
            if not spec or spec.get("w", 0) <= 0:
                continue
            ln = etree.Element(f"{{{A}}}{tag}")
            ln.set("w", str(max(1, int(spec["w"] * 12700))))
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            ln.set("algn", "ctr")
            _sub(_sub(ln, "solidFill"), "srgbClr", val=spec["hex"].lstrip("#").upper()[:6])
            tcPr.append(ln)
        _order_tc(tcPr)

    def picture(self, path: Any, x: float, y: float, w: float, h: float, radius: float = 0.0) -> Any:
        pic = self.slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        if radius > 0:
            pic._element.spPr.find(f"{{{A}}}prstGeom").set("prst", "roundRect")
        return pic

    # ------------------------------------------------- what python-pptx will not do
    @staticmethod
    def _alpha(clr: Any, a: float) -> None:
        """Alpha on a colour element — python-pptx exposes no transparency at all."""
        if clr is None:
            return
        for old in clr.findall(f"{{{A}}}alpha"):
            clr.remove(old)
        _sub(clr, "alpha", val=int(round(max(0.0, min(1.0, a)) * 100000)))

    def gradient(self, shp: Any, stops: Any, angle_deg: float = 90.0, radial: bool = False) -> Any:
        """Linear or radial gradient fill. stops = [(pos 0..1, '#RRGGBB', alpha), …]."""
        spPr = shp._element.spPr
        for tag in ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"):
            for old in spPr.findall(f"{{{A}}}{tag}"):
                spPr.remove(old)
        grad = etree.SubElement(spPr, f"{{{A}}}gradFill")
        grad.set("rotWithShape", "1")
        lst = _sub(grad, "gsLst")
        for pos, hexc, a in stops:
            gs = _sub(lst, "gs", pos=int(round(pos * 100000)))
            clr = _sub(gs, "srgbClr", val=hexc.lstrip("#").upper()[:6])
            if a < 1.0:
                self._alpha(clr, a)
        if radial:
            _sub(grad, "path", path="circle")
        else:
            _sub(grad, "lin", ang=int(round(angle_deg * 60000)) % 21600000, scaled="0")
        spPr.insert(list(spPr).index(grad), grad)
        return shp

    def hatch(self, shp: Any, fg: str, bg: str, pattern: str = "ltUpDiag") -> Any:
        spPr = shp._element.spPr
        for tag in ("noFill", "solidFill", "gradFill", "pattFill"):
            for old in spPr.findall(f"{{{A}}}{tag}"):
                spPr.remove(old)
        pat = etree.SubElement(spPr, f"{{{A}}}pattFill")
        pat.set("prst", pattern)
        _sub(_sub(pat, "fgClr"), "srgbClr", val=fg.lstrip("#").upper()[:6])
        _sub(_sub(pat, "bgClr"), "srgbClr", val=bg.lstrip("#").upper()[:6])
        return shp

    def soft_shadow(self, shp: Any, blur_pt: float = 12.0, dist_pt: float = 3.0, angle_deg: float = 90.0,
                    color: str = "000000", alpha: float = 0.22) -> Any:
        """outerShdw via effectLst. Exactly one effectLst per spPr — a duplicate is a repair prompt."""
        spPr = shp._element.spPr
        for old in spPr.findall(f"{{{A}}}effectLst"):
            spPr.remove(old)
        eff = etree.SubElement(spPr, f"{{{A}}}effectLst")
        shdw = _sub(eff, "outerShdw",
                    blurRad=int(blur_pt * 12700), dist=int(dist_pt * 12700),
                    dir=int(angle_deg * 60000) % 21600000, rotWithShape="0")
        self._alpha(_sub(shdw, "srgbClr", val=color.lstrip("#").upper()[:6]), alpha)
        return shp

    def round_caps(self, shp: Any) -> Any:
        ln = shp.line._get_or_add_ln()
        ln.set("cap", "rnd")
        return shp
