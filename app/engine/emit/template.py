"""Template integration: the customer's master, its layouts and its placeholders.

Ported from StageFlow `deckbase.py` (the vendor copy is reference-only) and extended with the
things the acceptance gate asks for that StageFlow never needed:

* **The template must survive the export byte for byte.** python-pptx re-serialises every XML part
  it loaded, so `ppt/slideMasters/**`, `ppt/slideLayouts/**` and `ppt/theme/**` come back with the
  same meaning but different bytes (namespace declarations, attribute order, self-closing tags).
  Meaning-equal is not what master brief §9 asks for, and it is not what a customer's brand team
  wants either. `master_template_bytes` hands `pptx.finalize` the master's own bytes for those
  trees to write back into the saved package; we never change them, so restoring them is always
  correct, and `verify_template` proves it afterwards.
* **Placeholders are located by `idx`, not by position**, falling back to `type` on a legacy
  manifest that has no `idx` — and erroring, rather than guessing, when `type` is ambiguous.
* **Unused placeholders are deleted**, because an empty placeholder renders its prompt text in some
  viewers and counts as a defect in the gate. One that paints something in its layout (a fill or an
  outline, such as a cover's angled panel) is replaced by a plain shape with that geometry and paint,
  so the design survives without an empty placeholder. `dt`/`ftr`/`sldNum` are never cloned onto a slide by
  python-pptx in the first place (they stay inherited from the layout), which is exactly what
  "keep them if the layout shows them" means.

This module deliberately knows nothing about shapes or text: it is the frame the rest draws into.
"""
from __future__ import annotations

import copy
import hashlib
import posixpath
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from defusedxml import ElementTree
from defusedxml.ElementTree import ParseError
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu

from app.config import engine as config
from app.engine.ir import Canvas
from app.engine.manifest import PLACEHOLDER_ALIASES, Layout, Manifest
from app.engine.reports import EmitOptions, EmitReport

#: Package trees that must come out of an export identical to the master they went in on.
TEMPLATE_PREFIXES: tuple[str, ...] = ("ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/")

_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

_ANCHORS = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


# --------------------------------------------------------------------------------- emit context


@dataclass(slots=True)
class EmitContext:
    """What every emitter function needs to know, passed instead of six parameters.

    `report` is the single place a drop is recorded: `warn` for something we could not do at all,
    `gap` for something we emitted correctly but a renderer draws differently (master brief §10.9 —
    log what you drop, never silently).
    """

    manifest: Manifest
    layout: Layout
    slide_id: str
    slide_index: int
    canvas: Canvas
    options: EmitOptions
    report: EmitReport
    #: The *target* master's theme fonts — a run in one of them is written as `+mj-lt`/`+mn-lt`.
    theme_fonts: dict[str, str] = field(default_factory=dict)

    def warn(self, element_id: str | None, message: str) -> None:
        self.report.warnings.append(f"{self.slide_id}/{element_id or '-'}: {message}")

    def gap(self, message: str) -> None:
        """A renderer gap: correct in the file, drawn differently by PptxRender. Recorded once."""
        if message not in self.report.renderer_gaps:
            self.report.renderer_gaps.append(message)

    def count(self, kind: str, n: int = 1) -> None:
        self.report.shapes_by_kind[kind] = self.report.shapes_by_kind.get(kind, 0) + n


# --------------------------------------------------------------- shared DrawingML helpers
# These live here rather than in `shapes.py` so `text.py` can use them without importing
# `shapes.py`: the emit package is a straight line, template → text → shapes → pptx.

def sub(parent: Any, tag: str, **attrs: Any) -> Any:
    """`<a:tag …/>` appended to `parent`. Attribute values are stringified, ints included."""
    element = parent.makeelement(qn(f"a:{tag}"), {})
    for key, value in attrs.items():
        element.set(key, str(value))
    parent.append(element)
    return element


def hexval(color: str | None, default: str = "000000") -> str:
    """`#1a9afb` / `1A9AFB` → `1A9AFB`. The IR promises 6 hex digits; be forgiving anyway."""
    cleaned = (color or default).lstrip("#").strip()
    return (cleaned[:6] or default).upper()


def set_alpha(color_element: Any, alpha: float) -> None:
    """`<a:alpha val>` on a colour element — python-pptx exposes no transparency at all."""
    if color_element is None or alpha >= 1.0:
        return
    for old in color_element.findall(qn("a:alpha")):
        color_element.remove(old)
    sub(color_element, "alpha", val=int(round(max(0.0, min(1.0, alpha)) * 100000)))


def solid_fill(parent: Any, color: str, alpha: float = 1.0) -> Any:
    """`<a:solidFill><a:srgbClr val="…"><a:alpha…/></a:srgbClr></a:solidFill>`."""
    fill = sub(parent, "solidFill")
    clr = sub(fill, "srgbClr", val=hexval(color))
    set_alpha(clr, alpha)
    return fill


def gradient_fill(parent: Any, fill: dict[str, Any], opacity: float = 1.0) -> Any:
    """`<a:gradFill rotWithShape="1">` for an IR gradient fill, appended to `parent`.

    Shared by shape fills (`shapes._apply_fill`) and run fills (gradient text, `text._write_run`),
    which is why it lives here: `text.py` must not import `shapes.py`. Every stop's alpha is
    multiplied by the element's `opacity`. DrawingML measures the angle from 3 o'clock, CSS from
    12: `ang = ((css − 90) mod 360) × 60000`; a radial gradient is a centred circle path.
    """
    node = sub(parent, "gradFill", rotWithShape=1)
    stops = sub(node, "gsLst")
    for stop in fill.get("stops") or []:
        gs = sub(stops, "gs", pos=int(round(max(0.0, min(1.0, float(stop.get("pos", 0)))) * 100000)))
        colour = sub(gs, "srgbClr", val=hexval(stop.get("color")))
        set_alpha(colour, float(stop.get("alpha", 1.0)) * opacity)
    if fill.get("kind") == "radial":
        path = sub(node, "path", path="circle")
        sub(path, "fillToRect", l=50000, t=50000, r=50000, b=50000)
    else:
        css_angle = float(fill.get("angle", 180.0))
        sub(node, "lin", ang=int(round(((css_angle - 90.0) % 360.0) * 60000)), scaled=0)
    return node


def remove_all(parent: Any, *tags: str) -> None:
    for tag in tags:
        for old in parent.findall(qn(f"a:{tag}")):
            parent.remove(old)


def drop_shape_style(shape: Any) -> None:
    """Remove `<p:style>`.

    python-pptx's `add_shape` attaches a style block referencing the theme (`fillRef`, `lnRef`,
    **`effectRef`**). The effect reference alone can paint a shadow the design never asked for, and
    every visual property is set explicitly here, so the whole block is noise that only ever
    disagrees with the IR.
    """
    element = getattr(shape, "_element", None)
    if element is None:
        return
    for style in element.findall(qn("p:style")):
        element.remove(style)


# ------------------------------------------------------------------------------------ geometry


def emu(px: float) -> int:
    """Canvas px → EMU. One pixel is exactly 9525 EMU on every master (master brief §5)."""
    return int(round(float(px) * config.EMU_PER_PX))


def emu_length(px: float) -> Emu:
    return Emu(emu(px))


def inches(px: float) -> float:
    """Canvas px → inches, for the StageFlow-derived helpers that speak inches."""
    return float(px) / config.PX_PER_IN


# ---------------------------------------------------------------------------------- theme fonts


def theme_font_token(family: str | None, theme_fonts: dict[str, str]) -> str | None:
    """`+mj-lt` / `+mn-lt` when the run is in the target master's own theme font, else the name.

    Writing the token instead of the literal name is what makes the deck re-themeable: change the
    master's typeface and the exported text follows, exactly as native PowerPoint content does.
    Comparison is case-insensitive and ignores surrounding quotes, because a CSS font stack writes
    `"Arial"` where the theme says `Arial`.
    """
    if not family:
        return None
    cleaned = family.strip().strip("'\"").lower()
    if not cleaned:
        return None
    minor = (theme_fonts.get("minor") or "").strip().strip("'\"").lower()
    major = (theme_fonts.get("major") or "").strip().strip("'\"").lower()
    if minor and cleaned == minor:
        return "+mn-lt"
    if major and cleaned == major:
        return "+mj-lt"
    return family


# ---------------------------------------------------------------------------------- placeholders


def placeholder_type(shape: Any) -> str | None:
    """The OOXML `<p:ph type>` of a shape, exactly as the manifest spells it (`obj` when absent)."""
    element = getattr(shape, "_element", None)
    if element is None:
        return None
    ph = element.find(f".//{qn('p:ph')}")
    if ph is None:
        return None
    return ph.get("type") or "obj"


def placeholder_idx(shape: Any) -> int | None:
    element = getattr(shape, "_element", None)
    if element is None:
        return None
    ph = element.find(f".//{qn('p:ph')}")
    if ph is None:
        return None
    raw = ph.get("idx")
    return int(raw) if raw is not None else 0


def find_placeholder(slide: Any, spec: dict[str, Any]) -> tuple[Any | None, str | None]:
    """The slide's placeholder shape for `{"type": …, "idx": …}` — `(shape, problem)`.

    `idx` identifies a placeholder unambiguously and is used whenever the manifest has one. A legacy
    manifest has `idx: null` everywhere (master brief §6), so the fallback is the type — and when
    the layout has two placeholders of that type the answer is genuinely unknown, so this returns a
    problem instead of picking one and putting the body text in the second column.
    """
    wanted_type = spec.get("type")
    wanted_idx = spec.get("idx")
    shapes = list(slide.placeholders)

    if wanted_idx is not None:
        for shape in shapes:
            if placeholder_idx(shape) == int(wanted_idx):
                return shape, None
        return None, f"no placeholder with idx {wanted_idx} on this slide"

    allowed = PLACEHOLDER_ALIASES.get(str(wanted_type), (str(wanted_type),))
    matches = [s for s in shapes if placeholder_type(s) in allowed]
    if not matches:
        return None, f"no {wanted_type} placeholder on this slide"
    if len(matches) > 1:
        idxs = sorted(str(placeholder_idx(s)) for s in matches)
        return None, (
            f"{len(matches)} placeholders match type {wanted_type!r} (idx {', '.join(idxs)}) and the "
            f"manifest records no idx — re-import the master so the mapping is unambiguous"
        )
    return matches[0], None


def prune_placeholders(slide: Any, keep: Iterable[Any]) -> list[str]:
    """Delete every placeholder we did not fill; return what was removed (for the report).

    An empty placeholder is not free: PowerPoint draws its prompt text in some viewers, it sits in
    front of the drawing, and the acceptance gate counts "0 placeholder shapes with an empty text
    frame". `dt`/`ftr`/`sldNum` are not on the slide at all (python-pptx leaves them inherited), so
    the layout's own footer keeps showing.
    """
    # Compare the XML elements, not the wrappers: `slide.placeholders` builds a fresh proxy object
    # on every iteration, so `shape is shape` is false for the very shapes we just filled.
    kept = {id(shape._element) for shape in keep}
    removed: list[str] = []
    for shape in list(slide.placeholders):
        if id(shape._element) in kept:
            continue
        name = f"{placeholder_type(shape)}#{placeholder_idx(shape)}"
        art = _layout_art(slide, shape)
        if art is not None:
            shape._element.getparent().replace(shape._element, art)
            name += " (its layout artwork kept as a shape)"
        else:
            shape._element.getparent().remove(shape._element)
        removed.append(name)
    return removed


#: `spPr` children that paint something. A placeholder's fill or outline is part of the layout's design
#: (a cover's angled colour panel is often a picture placeholder's own fill), not a prompt.
_PAINT_TAGS = tuple(qn(f"a:{tag}") for tag in ("solidFill", "gradFill", "blipFill", "pattFill"))


def _paints(sp_pr: Any) -> bool:
    if sp_pr is None:
        return False
    if any(child.tag in _PAINT_TAGS for child in sp_pr):
        return True
    line = sp_pr.find(qn("a:ln"))
    return line is not None and any(child.tag in _PAINT_TAGS for child in line)


def _layout_art(slide: Any, shape: Any) -> Any:
    """A plain shape carrying what the layout (or master) draws for this unused placeholder, or None.

    Deleting the slide's placeholder also deletes what its layout placeholder paints, because an
    inherited placeholder is drawn only through the slide's copy. When that copy paints (fill or
    outline), the design needs it: keep the geometry and paint as an ordinary `p:sp` with no text
    and no `p:ph`, so the picture stays and no empty placeholder is left for the gate to count.
    """
    chain: list[Any] = []
    base = getattr(shape, "_base_placeholder", None)
    while base is not None and len(chain) < 3:
        chain.append(base._element)
        base = getattr(base, "_base_placeholder", None)
    source = next((element for element in chain if _paints(element.find(qn("p:spPr")))), None)
    if source is None or source.tag != qn("p:sp"):
        return None
    art = copy.deepcopy(source)
    nv_pr = art.find(qn("p:nvSpPr")).find(qn("p:nvPr"))
    for ph in nv_pr.findall(qn("p:ph")):
        nv_pr.remove(ph)
    for body in art.findall(qn("p:txBody")):
        art.remove(body)
    sp_pr = art.find(qn("p:spPr"))
    if sp_pr.find(qn("a:xfrm")) is None:
        # Position inherited from further up the chain: take the first ancestor that states it.
        xfrm = next((el.find(qn("p:spPr")).find(qn("a:xfrm")) for el in chain
                     if el.find(qn("p:spPr")) is not None and el.find(qn("p:spPr")).find(qn("a:xfrm")) is not None),
                    None)
        if xfrm is None:
            return None
        sp_pr.insert(0, copy.deepcopy(xfrm))
    c_nv_pr = art.find(qn("p:nvSpPr")).find(qn("p:cNvPr"))
    used = [int(el.get("id")) for el in slide.shapes._spTree.iter() if el.tag == qn("p:cNvPr") and el.get("id", "").isdigit()]
    c_nv_pr.set("id", str(max(used, default=1) + 1))
    c_nv_pr.set("name", f"{c_nv_pr.get('name', 'Placeholder')} (layout art)")
    return art


def prepare_placeholder_frame(shape: Any, anchor: str | None) -> Any:
    """A placeholder's text frame, stripped of the layout's own autofit/margins/anchor.

    Wrap is not decided here: `text.write_layout` writes it (`text._wraps`) on this same frame right
    after, so a filled placeholder carries the explicit `wrap="none"` a free text box does and the
    rule lives in one place. When `add_text` emits nothing ("no measurable lines") the frame keeps
    the layout's own wrap — harmless, because the placeholder is then never appended to
    `used_placeholders` and `prune_placeholders` deletes it.
    """
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = _ANCHORS.get(anchor or "top", MSO_ANCHOR.TOP)
    body = frame._txBody.bodyPr
    for tag in ("a:normAutofit", "a:spAutoFit", "a:noAutofit"):
        for old in body.findall(qn(tag)):
            body.remove(old)
    body.append(body.makeelement(qn("a:noAutofit"), {}))
    return frame


def set_notes(slide: Any, notes: str | None) -> None:
    """Speaker notes. Only touched when there are some — `notes_slide` creates a part on access."""
    if not notes:
        return
    slide.notes_slide.notes_text_frame.text = notes


# ------------------------------------------------------------------- template part preservation


def _resolve(part_name: str, target: str) -> str:
    """Resolve a relationship target against the part that declares it (`../media/x.png`)."""
    if target.startswith("/"):
        return target.lstrip("/")
    folder = posixpath.dirname(part_name.lstrip("/"))
    return posixpath.normpath(posixpath.join(folder, target)).replace("\\", "/")


def template_part_names(archive: zipfile.ZipFile) -> set[str]:
    """Every part that must be identical to the master: the three trees plus the media they use."""
    names = {n for n in archive.namelist() if n.startswith(TEMPLATE_PREFIXES)}
    media: set[str] = set()
    for name in sorted(names):
        if not name.endswith(".rels"):
            continue
        owner = name.replace("/_rels/", "/")[: -len(".rels")]
        try:
            root = ElementTree.fromstring(archive.read(name))
        except ParseError:
            continue
        for rel in root.findall(f"{{{_REL_NS}}}Relationship"):
            if rel.get("TargetMode") == "External":
                continue
            resolved = _resolve(owner, rel.get("Target") or "")
            if resolved.startswith("ppt/media/"):
                media.add(resolved)
    return names | {m for m in media if m in set(archive.namelist())}


def template_hashes(pptx: Path) -> dict[str, str]:
    """sha256 per template part — what WP5's suite compares and what `verify_template` uses."""
    with zipfile.ZipFile(pptx) as archive:
        names = template_part_names(archive)
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in sorted(names)}


def verify_template(out_pptx: Path, master_pptx: Path) -> list[str]:
    """Every template part of the master that the export lost or changed (empty list = clean).

    Only the master's own parts are checked. A part the *export* adds is not a violation: writing
    speaker notes onto a master that has no notes master makes python-pptx add one, and that notes
    master brings its own theme part. Nothing of the customer's was touched, which is the promise.
    """
    master = template_hashes(Path(master_pptx))
    produced = template_hashes(Path(out_pptx))
    problems: list[str] = []
    for name, digest in master.items():
        if name not in produced:
            problems.append(f"{name}: missing from the export")
        elif produced[name] != digest:
            problems.append(f"{name}: differs from the master")
    return problems


def added_template_parts(out_pptx: Path, master_pptx: Path) -> list[str]:
    """Template parts the export added (a notes master's theme, typically). Informational."""
    master = template_hashes(Path(master_pptx))
    return [name for name in template_hashes(Path(out_pptx)) if name not in master]


def master_template_bytes(master_pptx: Path) -> dict[str, bytes]:
    """The master's own bytes for every template part, ready to be written back into an export."""
    with zipfile.ZipFile(master_pptx) as archive:
        return {name: archive.read(name) for name in sorted(template_part_names(archive))}
