"""Native PowerPoint charts from a `data-chart` spec.

LICENCE-GATED (decision D3). This module is derived from StageFlow, third-party code whose licence
terms have not been granted yet. The chart-type tables, the derived-type rewriting (`_convert` and
friends) and the transparent-background helpers are StageFlow's, kept verbatim; the rest was written
on top of them. It is ported as-is so the engine works, but it must not ship to production until D3
is resolved: either terms are granted, or this module is rewritten (risk R7, about 3 engineer-weeks).
See docs/licensing.md. Do not copy more StageFlow code into the service.

Native means a real chart part with a real embedded workbook: double-click it in PowerPoint and the
data is there, the series recolour with the theme, and the whole thing stays editable. That is the
thing a rasterised chart image can never be.

python-pptx writes the chart XML (unlike pptxgenjs, whose chart output PowerPoint often refuses), so
the structure is sound; everything here is styling on top of it.

**The spec is read once, by `app/engine/chart_model.py`** (WP-C). `add` normalises the spec, lays it out
with `chart_model.layout`, and styles the chart from the model — the same model the validator, the
structural check and the app's preview (`chart-preview.js`, its line-for-line twin) read. Nothing in
this module reads a raw option any more.

Units, once, because mixing them is how a chart lands 40 px too low:

* `box_in` / `add(x, y, w, h)` are **inches** — the emitter converts the IR's canvas px once.
* `plotRect` and every `*_px` argument are **canvas px** (1 px = 1/96 in, master brief §5).
* `options.fontSize` and `style.sizePx` are **px** (the whole contract is px); `lineWidth` and
  `gridlines.width` are **points**, as in PowerPoint's own UI; `dash` arrays are px on the design.

What this module draws on the slide *besides* the chart frame: waterfall `outEnd` labels and
`referenceLines`, as real connectors/text boxes on top of the frame (PowerPoint has no native
reference line). They need the plot rectangle and the value-axis range, so a chart that carries them
is **pinned**: its plot rectangle is written as `c:manualLayout` (the author's `plotArea`, the IR's
measured `plotRect`, or `chart_model.layout`'s default), its category labels horizontal and
unwrapped and its legend where the layout put it. Every other chart carries no manual layout at all —
PowerPoint lays it out itself (WP-C §3.6).

A waterfall's **connectors are part of the chart** (WP-C §9's fallback, taken in C3 because
PowerPoint moved a pinned plot by up to 23 px once its category labels outgrew their slots): one
`c:scatterChart` series per connector on a deleted secondary axis pair laid over the primary one
(`_in_chart_connectors`), and its value axis is always written out so they sit at the right level.

**Never half-built** (WP-C §4.1): a spec the model cannot draw leaves a labelled placeholder, and a
failure while styling rolls the chart part back and leaves the same placeholder — never a chart
with default colours, no labels and an unpinned axis.
"""
# ---------------------------------------------------------------------------------------------
# Vendored from StageFlow (third-party; licence-gated, decision D3), 2026-09-18.
# The type tables, the derived-type rewriting (`_convert` and friends) and the transparent-background
# helpers are StageFlow's, kept verbatim. WP3a replaced the body of `add` with the authoring
# contract's option set and added waterfall, manual plot layout and the overlay geometry; WP-C moved
# every reading of the spec into `app/engine/chart_model.py`.
# ---------------------------------------------------------------------------------------------

from __future__ import annotations

import math
from collections.abc import Iterable, Sequence
from typing import Any, cast

from lxml import etree
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
    XL_TICK_LABEL_POSITION,
)
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from app.config import engine as config
from app.engine import chart_model

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"

X = XL_CHART_TYPE

# Types python-pptx can write directly (measured — every other XL_CHART_TYPE raises
# NotImplementedError from its chart-XML writer).
TYPES = {
    "column": X.COLUMN_CLUSTERED, "column_stacked": X.COLUMN_STACKED,
    "column_stacked_100": X.COLUMN_STACKED_100,
    "bar": X.BAR_CLUSTERED, "bar_stacked": X.BAR_STACKED, "bar_stacked_100": X.BAR_STACKED_100,
    "line": X.LINE, "line_markers": X.LINE_MARKERS,
    "line_stacked": X.LINE_STACKED, "line_stacked_100": X.LINE_STACKED_100,
    "line_markers_stacked": X.LINE_MARKERS_STACKED,
    "line_markers_stacked_100": X.LINE_MARKERS_STACKED_100,
    "area": X.AREA, "area_stacked": X.AREA_STACKED, "area_stacked_100": X.AREA_STACKED_100,
    "pie": X.PIE, "pie_exploded": X.PIE_EXPLODED,
    "doughnut": X.DOUGHNUT, "doughnut_exploded": X.DOUGHNUT_EXPLODED,
    "radar": X.RADAR, "radar_markers": X.RADAR_MARKERS, "radar_filled": X.RADAR_FILLED,
    "scatter": X.XY_SCATTER, "scatter_lines": X.XY_SCATTER_LINES,
    "scatter_lines_no_markers": X.XY_SCATTER_LINES_NO_MARKERS,
    "scatter_smooth": X.XY_SCATTER_SMOOTH,
    "scatter_smooth_no_markers": X.XY_SCATTER_SMOOTH_NO_MARKERS,
    "bubble": X.BUBBLE, "bubble_3d": X.BUBBLE_THREE_D_EFFECT,
}

# Everything else is reached by building a supported chart and rewriting its XML:
#   name -> (base type, converter kwargs)
DERIVED: dict[str, tuple[str, dict[str, Any]]] = {
    # 3-D bars and columns, plus the cone/cylinder/pyramid shapes, which are bar3D + <c:shape>.
    "column_3d":                 ("column",         dict(to="bar3DChart", shape="box")),
    "column_3d_clustered":       ("column",         dict(to="bar3DChart", shape="box")),
    "column_3d_stacked":         ("column_stacked", dict(to="bar3DChart", shape="box")),
    "column_3d_stacked_100":     ("column_stacked_100", dict(to="bar3DChart", shape="box")),
    "bar_3d_clustered":          ("bar",            dict(to="bar3DChart", shape="box")),
    "bar_3d_stacked":            ("bar_stacked",    dict(to="bar3DChart", shape="box")),
    "bar_3d_stacked_100":        ("bar_stacked_100", dict(to="bar3DChart", shape="box")),
    "cone_column":               ("column",         dict(to="bar3DChart", shape="cone")),
    "cone_column_clustered":     ("column",         dict(to="bar3DChart", shape="cone")),
    "cone_column_stacked":       ("column_stacked", dict(to="bar3DChart", shape="cone")),
    "cone_column_stacked_100":   ("column_stacked_100", dict(to="bar3DChart", shape="cone")),
    "cone_bar_clustered":        ("bar",            dict(to="bar3DChart", shape="cone")),
    "cone_bar_stacked":          ("bar_stacked",    dict(to="bar3DChart", shape="cone")),
    "cone_bar_stacked_100":      ("bar_stacked_100", dict(to="bar3DChart", shape="cone")),
    "cylinder_column":           ("column",         dict(to="bar3DChart", shape="cylinder")),
    "cylinder_column_clustered": ("column",         dict(to="bar3DChart", shape="cylinder")),
    "cylinder_column_stacked":   ("column_stacked", dict(to="bar3DChart", shape="cylinder")),
    "cylinder_column_stacked_100": ("column_stacked_100", dict(to="bar3DChart", shape="cylinder")),
    "cylinder_bar_clustered":    ("bar",            dict(to="bar3DChart", shape="cylinder")),
    "cylinder_bar_stacked":      ("bar_stacked",    dict(to="bar3DChart", shape="cylinder")),
    "cylinder_bar_stacked_100":  ("bar_stacked_100", dict(to="bar3DChart", shape="cylinder")),
    "pyramid_column":            ("column",         dict(to="bar3DChart", shape="pyramid")),
    "pyramid_column_clustered":  ("column",         dict(to="bar3DChart", shape="pyramid")),
    "pyramid_column_stacked":    ("column_stacked", dict(to="bar3DChart", shape="pyramid")),
    "pyramid_column_stacked_100": ("column_stacked_100", dict(to="bar3DChart", shape="pyramid")),
    "pyramid_bar_clustered":     ("bar",            dict(to="bar3DChart", shape="pyramid")),
    "pyramid_bar_stacked":       ("bar_stacked",    dict(to="bar3DChart", shape="pyramid")),
    "pyramid_bar_stacked_100":   ("bar_stacked_100", dict(to="bar3DChart", shape="pyramid")),
    # 3-D line / area / pie
    "line_3d":                   ("line",           dict(to="line3DChart")),
    "area_3d":                   ("area",           dict(to="area3DChart")),
    "area_3d_stacked":           ("area_stacked",   dict(to="area3DChart")),
    "area_3d_stacked_100":       ("area_stacked_100", dict(to="area3DChart")),
    "pie_3d":                    ("pie",            dict(to="pie3DChart")),
    "pie_3d_exploded":           ("pie_exploded",   dict(to="pie3DChart")),
    # surface
    "surface":                   ("column",         dict(to="surface3DChart")),
    "surface_wireframe":         ("column",         dict(to="surface3DChart", wireframe=True)),
    "surface_top_view":          ("column",         dict(to="surfaceChart")),
    "surface_top_view_wireframe": ("column",        dict(to="surfaceChart", wireframe=True)),
    # pie of pie / bar of pie
    "pie_of_pie":                ("pie",            dict(to="ofPieChart", of="pie")),
    "bar_of_pie":                ("pie",            dict(to="ofPieChart", of="bar")),
    # stock — series count must match the flavour (HLC 3, OHLC 4, VHLC 4, VOHLC 5)
    "stock_hlc":                 ("line",           dict(to="stockChart", need=3)),
    "stock_ohlc":                ("line",           dict(to="stockChart", need=4)),
    "stock_vhlc":                ("line",           dict(to="stockChart", need=4, volume=True)),
    "stock_vohlc":               ("line",           dict(to="stockChart", need=5, volume=True)),
}

#: Other spellings of a known type. One table, in `chart_model` (WP-C §2.5).
ALIASES = chart_model.ALIASES

STACKED = {k for k in TYPES if "stacked" in k} | {k for k in DERIVED if "stacked" in k}
NO_AXES = {"pie", "pie_exploded", "doughnut", "doughnut_exploded",
           "pie_3d", "pie_3d_exploded", "pie_of_pie", "bar_of_pie"}
XY_KINDS = {k for k in TYPES if k.startswith("scatter")}
BUBBLE_KINDS = {"bubble", "bubble_3d"}
THREE_D = {k for k in DERIVED if DERIVED[k][1]["to"] in
           ("bar3DChart", "line3DChart", "area3DChart", "pie3DChart", "surface3DChart")}

#: Types this engine adds on top of StageFlow's list, and the base chart each is built from.
#: `combo` is round 2 and the model refuses it, so only the waterfall reaches here.
ENGINE_BASE = {"waterfall": chart_model.ENGINE_BASES["waterfall"]}

#: Name of the invisible series a waterfall floats its bars on. Fixed so `expected_workbook` and
#: the structural check agree on what the embedded workbook should hold.
WATERFALL_BASE_SERIES = chart_model.BASE_SERIES

# Chart elements the PptxRender preview can actually draw. Verified in its source: ChartPainter only
# paints Bar / HorizontalBar / Line / Area / Pie / Doughnut / Scatter and treats the rest as
# Unsupported, i.e. a blank image. The FILES are valid either way — only the preview is limited.
PREVIEWABLE = {"barChart", "bar3DChart", "lineChart", "line3DChart", "areaChart", "area3DChart",
               "pieChart", "pie3DChart", "doughnutChart", "scatterChart"}


def resolve(kind: Any) -> tuple[str | None, str | None, dict[str, Any] | None]:
    """friendly name -> (canonical name, base name for python-pptx, converter kwargs or None).

    A wrapper over `chart_model.resolve_type_name`. A name the model refuses (unknown, Path-A-
    unsupported, round 2) is `(None, None, None)` — never a clustered column in disguise again.
    """
    name = chart_model.resolve_type_name(kind)["name"]
    if name is None:
        return None, None, None
    if name in ENGINE_BASE:
        return name, ENGINE_BASE[name], None
    if name in TYPES:
        return name, name, None
    base, conv = DERIVED[name]
    return name, base, conv


def target_element(kind: Any) -> str | None:
    """The plot-area element a type is written as (`barChart`, `ofPieChart`, …); None if refused."""
    k, resolved_base, conv = resolve(kind)
    if k is None:
        return None
    base = cast(str, resolved_base)  # resolve() names a base for every type it accepts
    if conv:
        return cast(str, conv["to"])
    return ("barChart" if base.startswith(("column", "bar")) else
            "lineChart" if base.startswith("line") else
            "areaChart" if base.startswith("area") else
            "pieChart" if base.startswith("pie") else
            "doughnutChart" if base.startswith("doughnut") else
            "radarChart" if base.startswith("radar") else
            "scatterChart" if base.startswith("scatter") else "bubbleChart")


def previewable(kind: Any) -> bool:
    return target_element(kind) in PREVIEWABLE


# ------------------------------------------------------------------ the contract's option values

#: `03-AUTHORING-CONTRACT.md` spelling -> (python-pptx enum, the OOXML `c:dLblPos` value).
LABEL_POSITIONS: dict[str, tuple[Any, str]] = {
    "center": (XL_LABEL_POSITION.CENTER, "ctr"),
    "inEnd": (XL_LABEL_POSITION.INSIDE_END, "inEnd"),
    "inBase": (XL_LABEL_POSITION.INSIDE_BASE, "inBase"),
    "outEnd": (XL_LABEL_POSITION.OUTSIDE_END, "outEnd"),
    "bestFit": (XL_LABEL_POSITION.BEST_FIT, "bestFit"),
    "left": (XL_LABEL_POSITION.LEFT, "l"),
    "right": (XL_LABEL_POSITION.RIGHT, "r"),
    "above": (XL_LABEL_POSITION.ABOVE, "t"),
    "below": (XL_LABEL_POSITION.BELOW, "b"),
}

#: Positions a stacked plot accepts; a label there sits ON a bar, so its ink contrasts with the bar.
INSIDE = {"center", "inEnd", "inBase"}

LEGEND_POSITIONS = {
    "bottom": XL_LEGEND_POSITION.BOTTOM, "top": XL_LEGEND_POSITION.TOP,
    "left": XL_LEGEND_POSITION.LEFT, "right": XL_LEGEND_POSITION.RIGHT,
    "topRight": XL_LEGEND_POSITION.CORNER,
}

MARKERS = {
    "none": XL_MARKER_STYLE.NONE, "auto": XL_MARKER_STYLE.AUTOMATIC,
    "circle": XL_MARKER_STYLE.CIRCLE, "square": XL_MARKER_STYLE.SQUARE,
    "diamond": XL_MARKER_STYLE.DIAMOND, "triangle": XL_MARKER_STYLE.TRIANGLE,
    "x": XL_MARKER_STYLE.X, "star": XL_MARKER_STYLE.STAR, "dash": XL_MARKER_STYLE.DASH,
    "dot": XL_MARKER_STYLE.DOT, "plus": XL_MARKER_STYLE.PLUS,
}

#: DrawingML preset dash names a `dash` string may use. A list (`[4, 3]`, px on the design) becomes
#: an `a:custDash`, whose lengths are percentages of the line width.
PRESET_DASHES = set(chart_model.DASH_NAMES)

#: Default colours for the three roles of a waterfall bar when `colors` does not name them.
WATERFALL_COLORS = chart_model.WATERFALL_COLORS

#: The failed-chart placeholder: a dashed red outline and the reason, in the chart's font.
PLACEHOLDER_COLOR = "E5484D"
#: Below this size the placeholder carries its message in its name only (nothing for fit to judge).
PLACEHOLDER_MIN_PX = (120.0, 40.0)


# ------------------------------------------------------------------------------- small utilities


def _c(tag: str) -> str:
    return f"{{{C}}}{tag}"


def _a(tag: str) -> str:
    return f"{{{A}}}{tag}"


def _hex(value: Any) -> str:
    return str(value).lstrip("#").upper()[:6]


def _rgb(h: Any) -> Any:
    return RGBColor.from_string(_hex(h))  # type: ignore[no-untyped-call]


def _luminance(hexc: Any) -> Any:
    h = _hex(hexc)
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    def f(c: float) -> float:
        return c / 12.92 if c <= 0.04045 else float(((c + 0.055) / 1.055) ** 2.4)

    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def _contrast(hexc: Any, dark: Any = "12233B", light: Any = "FFFFFF") -> Any:
    """
    Readable label colour for text sitting ON this fill. Picks whichever of the two gives the higher
    WCAG contrast ratio — a plain luminance threshold puts white on mid greys, where it is unreadable.
    """
    try:
        bg = _luminance(hexc)
    except ValueError:
        return dark
    def ratio(a: float, b: float) -> float:
        return (max(a, b) + 0.05) / (min(a, b) + 0.05)

    return dark if ratio(bg, _luminance(dark)) >= ratio(bg, _luminance(light)) else light


def _xml_safe(text: Any) -> str:
    """`text` with every character an XML part cannot hold dropped (C0 controls but tab, LF and CR;
    lone surrogates; U+FFFE/U+FFFF). python-pptx raises on them, and a name or message is written
    where nothing may raise (`_placeholder`, C1 review #13)."""
    return "".join(
        character for character in str(text)
        if (character >= " " or character in "\t\n\r")
        and not ("\ud800" <= character <= "\udfff") and character not in "\ufffe\uffff"
    )


def _number(value: Any) -> float | None:
    """`float(value)` for a real number, `None` for `None`/bool/anything else (a gap slot)."""
    if value is None or isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value)


def _int(value: float) -> int:
    """Half-up integer, as the model rounds (Python's `round` is half-even)."""
    return int(math.floor(float(value) + 0.5))


class _Diagnostics:
    """Where "I could not do that" goes. A list from the caller, else the module logger.

    `add_chart` keeps StageFlow's return type (a graphic frame), so the emitter passes a sink in
    `style["diagnostics"]` and gets the same dicts `IR.diagnostics` holds.
    """

    def __init__(self, sink: list[Any] | None, source: str) -> None:
        self.sink = sink if isinstance(sink, list) else None
        self.source = source
        self.entries: list[dict[str, str]] = []

    def add(self, level: str, message: str) -> None:
        entry = {"level": level, "source": self.source, "message": message}
        self.entries.append(entry)
        if self.sink is not None:
            self.sink.append(entry)
        else:
            import logging

            logging.getLogger(__name__).log(
                logging.WARNING if level in ("warn", "error") else logging.INFO,
                "%s: %s", self.source, message,
            )

    def info(self, message: str) -> None:
        self.add("info", message)

    def warn(self, message: str) -> None:
        self.add("warn", message)

    def error(self, message: str) -> None:
        self.add("error", message)


def _insert_after(parent: Any, child: Any, after_tags: Sequence[str]) -> None:
    """Put `child` straight after the last of `after_tags` present — OOXML sequences are ordered.

    PowerPoint silently "repairs" a chart whose children are out of sequence, and a repair drops
    shapes, so every element this module adds goes in through here.
    """
    index = 0
    for position, existing in enumerate(parent):
        if etree.QName(existing).localname in after_tags:
            index = position + 1
    parent.insert(index, child)


def _sub(parent: Any, tag: str, **attrs: Any) -> Any:
    element = etree.SubElement(parent, _c(tag))
    for key, value in attrs.items():
        element.set(key, str(value))
    return element


def _apply_dash(line_element: Any, dash: Any, width_px: float) -> None:
    """Write `a:prstDash` or `a:custDash` on an `a:ln`. `dash` is a preset name or px on/off pairs."""
    if dash is None:
        return
    for existing in line_element.findall(_a("prstDash")) + line_element.findall(_a("custDash")):
        line_element.remove(existing)
    if isinstance(dash, str):
        name = dash if dash in PRESET_DASHES else "dash"
        etree.SubElement(line_element, _a("prstDash")).set("val", name)
        return
    pairs = [float(v) for v in dash if _number(v) is not None and float(v) > 0]
    if not pairs:
        return
    if len(pairs) % 2:                      # "4" means 4 on, 4 off, as in SVG
        pairs = pairs + pairs
    unit = max(width_px, 0.1)               # custDash lengths are a % of the line width
    custom = etree.SubElement(line_element, _a("custDash"))
    for on, off in zip(pairs[::2], pairs[1::2], strict=False):
        etree.SubElement(custom, _a("ds")).attrib.update({
            "d": str(max(1000, int(round(on / unit * 100000)))),
            "sp": str(max(1000, int(round(off / unit * 100000)))),
        })


def _style_line(line_format: Any, color: Any, width_pt: float | None, dash: Any) -> None:
    """Colour, width and dash on a python-pptx `LineFormat`, in one place."""
    if color is not None:
        line_format.color.rgb = _rgb(color)
    if width_pt is not None:
        line_format.width = Pt(float(width_pt))
    if dash is not None:
        _apply_dash(line_format._get_or_add_ln(), dash, float(width_pt or 1.0) / config.PT_PER_PX)


def _nice_range(low: float, high: float) -> tuple[float, float, float]:
    """`chart_model.nice_range` — kept under its old name for callers of WP3a's module."""
    return chart_model.nice_range(low, high)


# ------------------------------------------------------------------------------------- waterfall


def waterfall_plan(spec: dict[str, Any]) -> dict[str, Any]:
    """The model's waterfall plan (`chart_model.normalise(spec)["waterfall"]`) for this spec.

    Also carries WP3a's two keys, `base` and `delta`: the cells of the invisible floor series and of
    the first visible one — the whole workbook of a single-series waterfall that never dips below
    zero. `{}` when the spec is not a drawable waterfall.
    """
    plan = dict(chart_model.normalise(spec).get("waterfall") or {})
    if plan:
        plan["base"] = list(plan["series"][0]["values"])
        plan["delta"] = list(plan["series"][1]["values"]) if len(plan["series"]) > 1 else []
    return plan


# ---------------------------------------------------------------------------------- overlay lines


def plot_area_from_rect(box_in: tuple[float, float, float, float],
                        plot_rect_px: Any) -> dict[str, float] | None:
    """`options.plotArea` fractions for a design plot rect (canvas px) inside this frame (inches).

    This is how the IR's `plotRect` reaches `c:manualLayout`: the chart's own plot rectangle lands
    where the design drew it instead of wherever PowerPoint's auto layout would put it.
    """
    rect = _rect_px(plot_rect_px)
    if rect is None:
        return None
    x, y, w, h = box_in
    if w <= 0 or h <= 0:
        return None
    px = float(config.PX_PER_IN)
    fractions = {
        "x": (rect[0] / px - x) / w,
        "y": (rect[1] / px - y) / h,
        "w": (rect[2] / px) / w,
        "h": (rect[3] / px) / h,
    }
    if any(not math.isfinite(v) for v in fractions.values()):
        return None
    return fractions


def _rect_px(rect: Any) -> tuple[float, float, float, float] | None:
    """Accept a Box, a dict or a 4-tuple of canvas px; `None` when there is nothing usable."""
    if rect is None:
        return None
    if isinstance(rect, dict):
        try:
            return (float(rect["x"]), float(rect["y"]), float(rect["w"]), float(rect["h"]))
        except (KeyError, TypeError, ValueError):
            return None
    if isinstance(rect, (tuple, list)) and len(rect) == 4:
        return tuple(float(v) for v in rect)            # type: ignore[return-value]
    for attribute in ("x", "y", "w", "h"):
        if not hasattr(rect, attribute):
            return None
    return (float(rect.x), float(rect.y), float(rect.w), float(rect.h))


def _with_plot_rect(spec: Any, box_in: tuple[float, float, float, float], plot_rect_px: Any) -> Any:
    """The spec with the IR's measured plot rect as its `options.plotArea` (4 dp), unless it has one.

    A measured plot rect *is* a plot area the author did not have to type; after this the model
    sees one spelling of it.
    """
    if not isinstance(spec, dict) or plot_rect_px is None:
        return spec
    options = spec.get("options")
    if options is None:
        options = {}
    if not isinstance(options, dict) or options.get("plotArea") is not None:
        return spec
    fractions = plot_area_from_rect(box_in, plot_rect_px)
    if not fractions:
        return spec
    options = dict(options, plotArea={key: chart_model._round_to(value, 4) for key, value in fractions.items()})
    return dict(spec, options=options)


def overlay_geometry(spec: dict[str, Any] | None, box_in: tuple[float, float, float, float], *,
                     plot_rect_px: Any = None,
                     value_range: tuple[float, float] | None = None,
                     size_px: float | None = None,
                     model: dict[str, Any] | None = None,
                     geometry: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    """What this module draws **on top of** the chart frame, in inches.

    Waterfall connectors between consecutive bars, waterfall `outEnd` labels, and `referenceLines`
    — PowerPoint has no native equivalent for any of them on the chart types Path A emits. All three
    are placed from the plot rectangle (`chart_model.layout`: the author's `plotArea`, the IR's
    `plotRect`, or the default rect a pinned chart writes) and the pinned value axis
    (`model["axis"]`, or `value_range`). Without an axis there is nothing to place: `[]`.
    """
    if model is None:
        model = chart_model.normalise(_with_plot_rect(spec, box_in, plot_rect_px))
    if model["type"] is None:
        return []
    x, y, w, h = box_in
    if geometry is None:
        geometry = chart_model.layout(model, w * config.PX_PER_IN, h * config.PX_PER_IN,
                                      size_px if size_px is not None else chart_model.DEFAULT_SIZE_PX)
    area = geometry["plotArea"]
    plot = (x + area["x"] * w, y + area["y"] * h, area["w"] * w, area["h"] * h)

    axis = model["axis"]
    if value_range is not None:
        low, high = float(value_range[0]), float(value_range[1])
    elif axis is not None:
        low, high = axis["min"], axis["max"]
    else:
        return []
    if high <= low:
        return []

    options = model["options"]
    horizontal = model["family"] == "bar"
    reverse = bool(options["valueAxis"]["reverse"])
    lines: list[dict[str, Any]] = []
    if model["family"] == "waterfall":
        lines.extend(_waterfall_connectors(model, plot, low, high, horizontal, reverse))
        lines.extend(_waterfall_labels(model, plot, low, high, horizontal, reverse))
    for index, line in enumerate(options["referenceLines"]):
        value = line["value"]
        if value < low or value > high:
            continue                                   # noted by the model; a clamped line would lie
        position = _value_to_position(value, low, high, plot, horizontal, reverse)
        if horizontal:
            geometry_ = {"x1": position, "y1": plot[1], "x2": position, "y2": plot[1] + plot[3]}
        else:
            geometry_ = {"x1": plot[0], "y1": position, "x2": plot[0] + plot[2], "y2": position}
        lines.append({
            "kind": "referenceLine", "index": index, "value": value,
            "color": line["color"], "width": line["width"], "dash": line["dash"],
            "label": line["label"], **geometry_,
        })
    return lines


def _value_to_position(value: float, low: float, high: float,
                       plot: tuple[float, float, float, float],
                       horizontal: bool, reverse: bool) -> float:
    fraction = (value - low) / (high - low)
    if reverse:
        fraction = 1.0 - fraction
    if horizontal:
        return plot[0] + fraction * plot[2]
    return plot[1] + (1.0 - fraction) * plot[3]


def _category_span(plot: Any, count: int, index: int, gap_width: float, horizontal: bool,
                   reverse: bool = False) -> tuple[float, float]:
    """(start, end) of one category's bar along the category axis, in inches.

    A stacked column puts one bar per category slot; the slot is `plot / count` and the bar fills it
    minus the gap, exactly as PowerPoint lays `c:gapWidth` out. A column chart's first category is
    on the left and a bar chart's at the bottom; `categoryAxis.reverse` mirrors the slot index
    (`count − 1 − i`), as PowerPoint and the preview do.
    """
    origin, extent = (plot[1], plot[3]) if horizontal else (plot[0], plot[2])
    count = max(count, 1)
    slot = extent / count
    bar = slot / (1.0 + max(gap_width, 0.0) / 100.0)
    if horizontal:
        position = index if reverse else count - 1 - index
    else:
        position = count - 1 - index if reverse else index
    centre = origin + (position + 0.5) * slot
    return centre - bar / 2.0, centre + bar / 2.0


def _waterfall_connectors(model: dict[str, Any], plot: Any, low: float, high: float, horizontal: bool,
                          reverse: bool) -> list[dict[str, Any]]:
    """The plan's connectors: from the edge of bar i facing bar i+1 to the edge of bar i+1 facing i."""
    plan = model["waterfall"]
    options = model["options"]
    style = options["connectors"]
    if style is False or not plan["connectors"]:
        return []
    count = len(plan["roles"])
    gap_width = options["gapWidth"]
    categories_reversed = bool(options["categoryAxis"]["reverse"])
    lines: list[dict[str, Any]] = []
    for connector in plan["connectors"]:
        index, level = connector["index"], connector["level"]
        position = _value_to_position(level, low, high, plot, horizontal, reverse)
        a0, a1 = _category_span(plot, count, index, gap_width, horizontal, categories_reversed)
        b0, b1 = _category_span(plot, count, index + 1, gap_width, horizontal, categories_reversed)
        start, end = (a1, b0) if b0 >= a1 else (a0, b1)
        if horizontal:
            geometry = {"x1": position, "y1": start, "x2": position, "y2": end}
        else:
            geometry = {"x1": start, "y1": position, "x2": end, "y2": position}
        lines.append({
            "kind": "connector", "index": index, "value": level,
            "color": style["color"], "width": style["width"], "dash": style["dash"],
            "label": None, **geometry,
        })
    return lines


def _waterfall_labels(model: dict[str, Any], plot: Any, low: float, high: float, horizontal: bool,
                      reverse: bool) -> list[dict[str, Any]]:
    """`labelPosition: "outEnd"` on a waterfall: one text box per bar, above a rise or a positive
    total and below a fall or a negative total (OOXML has no `outEnd` on a stacked bar)."""
    options = model["options"]
    if options["labelPosition"] != "outEnd" or horizontal:
        return []
    plan = model["waterfall"]
    labels_on = options["dataLabels"]
    count = len(plan["roles"])
    categories_reversed = bool(options["categoryAxis"]["reverse"])
    lines: list[dict[str, Any]] = []
    for index, role in enumerate(plan["roles"]):
        if role == "gap":
            continue
        present = [k for k, entry in enumerate(model["series"]) if entry["values"][index] is not None]
        if not any(labels_on[k] for k in present if k < len(labels_on)):
            continue
        value = plan["values"][index]
        own = plan["texts"][index] if index < len(plan["texts"]) else None
        label = own if own is not None else chart_model.format_number(value, plan["formats"][role])
        if label == "":
            continue
        above = role == "increase" or (role == "total" and value >= 0)
        anchor = plan["tops"][index] if above else plan["bottoms"][index]
        anchor = min(max(anchor, low), high)
        position = _value_to_position(anchor, low, high, plot, horizontal, reverse)
        start, end = _category_span(plot, count, index, options["gapWidth"], horizontal, categories_reversed)
        centre = (start + end) / 2.0
        lines.append({
            "kind": "label", "index": index, "value": value,
            "label": label,
            "place": "above" if above else "below",
            "color": None, "width": None, "dash": None,
            "x1": centre, "y1": position, "x2": centre, "y2": position,
        })
    return lines


def _draw_overlay(slide: Any, lines: Iterable[dict[str, Any]], *, font: str, size_pt: float,
                  color: str, name: str, drawn: list[Any] | None = None, label_color: str | None = None,
                  label_bold: bool = False) -> list[Any]:
    """Draw `overlay_geometry`'s items: connectors (plus a text box for a labelled line), and the
    waterfall's `outEnd` labels as text boxes — in `options.labelStyle`'s ink and weight, one
    paragraph per line of a text broken by `\\n`. Every shape is appended to `drawn` as it is made,
    so a failure half way can be rolled back."""
    drawn = drawn if drawn is not None else []
    gap_in = 2.0 / config.PX_PER_IN
    for line in lines:
        if line["kind"] == "label":
            text = str(line["label"])
            rows = chart_model.text_lines(text)
            widest = max(len(row) for row in rows)
            width_in = max(0.4, 0.55 * size_pt / 72.0 * widest + 0.08)
            height_in = size_pt / 72.0 * (1.6 + (len(rows) - 1) * chart_model.LINE_PITCH_EM)
            left = line["x1"] - width_in / 2.0
            top = line["y1"] - gap_in - height_in if line.get("place") == "above" else line["y1"] + gap_in
            box = slide.shapes.add_textbox(
                Emu(int(round(left * config.EMU_PER_IN))), Emu(int(round(top * config.EMU_PER_IN))),
                Emu(int(round(width_in * config.EMU_PER_IN))), Emu(int(round(height_in * config.EMU_PER_IN))),
            )
            drawn.append(box)
            box.name = f"{name} label {line['index'] + 1}"
            _fill_label(box, _xml_safe(text), font=font, size_pt=size_pt,
                        color=line.get("color") or label_color or color, centred=True, bold=label_bold)
            continue

        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Emu(int(round(line["x1"] * config.EMU_PER_IN))),
            Emu(int(round(line["y1"] * config.EMU_PER_IN))),
            Emu(int(round(line["x2"] * config.EMU_PER_IN))),
            Emu(int(round(line["y2"] * config.EMU_PER_IN))),
        )
        drawn.append(shape)
        shape.name = f"{name} {line['kind']} {line['index'] + 1}"
        _style_line(shape.line, line["color"], line["width"], line.get("dash"))

        if not line.get("label"):
            continue
        text = str(line["label"])
        # A rough advance width (0.55 em) is enough: the box only has to hold the run without
        # wrapping, and word_wrap is off, so a little slack costs nothing and a little short wraps.
        width_in = max(0.4, 0.55 * size_pt / 72.0 * len(text) + 0.08)
        height_in = size_pt / 72.0 * 1.6
        left = min(line["x1"], line["x2"])
        right = max(line["x1"], line["x2"])
        top = min(line["y1"], line["y2"])
        box = slide.shapes.add_textbox(
            Emu(int(round(max(left, right - width_in) * config.EMU_PER_IN))),
            Emu(int(round((top - height_in) * config.EMU_PER_IN))),
            Emu(int(round(width_in * config.EMU_PER_IN))),
            Emu(int(round(height_in * config.EMU_PER_IN))),
        )
        drawn.append(box)
        box.name = f"{name} {line['kind']} {line['index'] + 1} label"
        _fill_label(box, _xml_safe(text), font=font, size_pt=size_pt, color=line.get("color") or color,
                    centred=False)
    return drawn


def _fill_label(box: Any, text: str, *, font: str, size_pt: float, color: str, centred: bool,
                bold: bool = False) -> None:
    from pptx.enum.text import PP_ALIGN

    frame = box.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, row in enumerate(chart_model.text_lines(text)):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if centred:
            paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = row
        run.font.size = Pt(size_pt)
        run.font.name = font
        run.font.color.rgb = _rgb(color)
        if bold:
            run.font.bold = True


# ------------------------------------------------------------------------------ the public entry


def expected_workbook(spec: dict[str, Any]) -> dict[str, Any]:
    """What the emitted chart's embedded workbook holds, for a structural check (WP5).

    The model's own series — identity for every type but `waterfall`, which is emitted as the plan's
    invisible base series plus the visible steps (and their below-zero copies), so "the workbook
    equals the spec" has to be asked of *this*, not of the spec. Numeric strings arrive coerced.
    """
    model = chart_model.normalise(spec)
    if model["family"] == "scatter":
        return {"categories": None, "series": [
            {"name": s["name"], "x": list(s["x"] or []), "y": list(s["y"] or [])}
            for s in model["series"]
        ]}
    series = model["waterfall"]["series"] if model["waterfall"] else model["series"]
    if model["type"] is None:
        series = []
    return {"categories": list(model["categories"]), "series": [
        {"name": s["name"], "values": list(s["values"])} for s in series
    ]}


def expected_series(spec: dict[str, Any]) -> list[dict[str, Any]]:
    """Every `c:ser` the emitted chart holds, in document order: the workbook's series
    (`expected_workbook`), then a waterfall's in-chart connector series — literal x/y caches the
    workbook does not hold (`connector_points`). What the structural check compares caches with."""
    series = list(expected_workbook(spec)["series"])
    model = chart_model.normalise(spec)
    if model["family"] == "waterfall" and model["type"] is not None and model["axis"] is not None:
        for offset, (xs, ys) in enumerate(connector_points(model)):
            series.append({"name": f"{CONNECTOR_SERIES} {offset + 1}", "x": xs, "y": ys})
    return series


def add_chart(slide: Any, spec: dict[str, Any], box_in: tuple[float, float, float, float], style: dict[str, Any]) -> Any:
    """The engine's single entry point into this module: IR chart element -> GraphicFrame.

    `box_in` is (x, y, w, h) in **inches** — the emitter converts the IR's canvas px once, so this
    module keeps StageFlow's inch-based signature. `style` carries the design's text style for the
    chart's own labels: `{"font": "Arial", "sizePx": 11, "color": "516467"}`. Font size travels as
    px in the IR (everything else does) and PowerPoint wants points, hence the 0.75.

    Three optional keys on `style` let the emitter hand over what the IR knows and the spec does not
    (the signature is fixed by master brief §8, and none of these is required):

    * `plotRect` — the IR chart element's `plotRect`, in canvas px. Becomes `c:manualLayout`, so the
      plot lands where the design's plot rect was. `options.plotArea` wins when both are present.
    * `diagnostics` — a list to append `{"level", "source", "message"}` dicts to; without it the
      same messages go to the module logger. Nothing is ever dropped silently.
    * `name` — the PowerPoint shape name for the frame and for any overlay line drawn with it.

    Returns the chart's graphic frame, or — when the spec cannot be drawn or styling failed — the
    labelled placeholder shape (`shape.has_chart` is False; an `error` diagnostic says why).
    """
    x, y, w, h = box_in
    return add(
        slide,
        spec,
        x,
        y,
        w,
        h,
        font=style.get("font") or "Calibri",
        base_pt=float(style.get("sizePx") or chart_model.DEFAULT_SIZE_PX) * config.PT_PER_PX,
        color=style.get("color"),
        plot_rect_px=style.get("plotRect"),
        diagnostics=style.get("diagnostics"),
        name=style.get("name"),
    )


class _Text:
    """The chart's text style, resolved once from the model and the layout."""

    def __init__(self, model: dict[str, Any], geometry: dict[str, Any], font: str, color: Any,
                 accents: dict[int, str] | None = None) -> None:
        options = model["options"]
        #: The deck theme's accent1..6 as RRGGBB — what a `schemeClr accentN` fill looks like, so a
        #: label on it can take a readable ink (C1 review #14b).
        self.accents = dict(accents or {})
        self.base_pt = geometry["sizePx"] * config.PT_PER_PX
        # Data labels sit a touch below the chart's body text, as PowerPoint's own defaults do —
        # unless the author states a size, in which case that size is what they asked for.
        self.label_pt = self.base_pt if options["fontSize"] is not None else self.base_pt * 0.95
        self.font = options["font"] or font
        self.ink = options["fontColor"] or chart_model.normalise_colour(color)[0] or "12233B"
        # `options.labelStyle` (G24): the labels' own size, weight and ink. A stated colour is the
        # labels' ink everywhere — it replaces the contrast pick against each bar.
        style = options.get("labelStyle") or {}
        if style.get("fontSize") is not None:
            self.label_pt = chart_model.label_size(model, geometry["sizePx"]) * config.PT_PER_PX
        self.label_bold = bool(style.get("bold"))
        self.label_color: str | None = style.get("color")


def add(slide: Any, spec: dict[str, Any], x: Any, y: Any, w: Any, h: Any, font: Any = "Calibri", base_pt: Any = 10.0, *,
        color: str | None = None, plot_rect_px: Any = None,
        diagnostics: list[Any] | None = None, name: str | None = None) -> Any:
    """Draw the chart described by `spec` at this rect (inches). Returns the graphic frame.

    `spec` is the authoring contract's `data-chart` object, read by `chart_model.normalise`.
    Nothing that can raise runs after `slide.shapes.add_chart` outside a guard: a spec the model
    cannot draw, or a failure while styling (the chart part is rolled back), leaves a labelled
    placeholder rectangle and an `error` diagnostic instead of a half-built chart (WP-C §4.1).
    """
    box = (x, y, w, h)
    spec = _with_plot_rect(spec, box, plot_rect_px)
    model = chart_model.normalise(spec)
    name = _xml_safe(name) if name else name
    label = name or "Chart"
    log = _Diagnostics(diagnostics, name or f"chart[{model['type'] or model.get('requested') or '?'}]")
    for note in model["notes"]:
        log.add("warn" if note["kind"] == "warn" else "info", note["text"])
    size_px = float(base_pt) / config.PT_PER_PX
    font = model["options"].get("font") or font

    if model["type"] is None:
        log.error(model["reason"])
        return _placeholder(slide, box, model["reason"], font=font, size_pt=base_pt, name=label)
    kind, base = model["type"], model["base"]
    conv = DERIVED[kind][1] if kind in DERIVED else None
    if not previewable(kind):
        # The file is valid and PowerPoint draws it; PptxRender's ChartPainter does not, so the
        # pixel gate sees an empty frame. A renderer gap, recorded rather than discovered later.
        log.info(f"PptxRender cannot draw a {target_element(kind)} — it renders blank and warns; "
                 f"the emitted chart itself is valid (see 03-AUTHORING-CONTRACT.md §Charts)")

    geometry = chart_model.layout(model, w * config.PX_PER_IN, h * config.PX_PER_IN, size_px)
    _ink_the_totals(model, _theme_ink(slide))
    text = _Text(model, geometry, font, color, _theme_accents(slide))
    try:
        data = _chart_data(model)
        frame = slide.shapes.add_chart(TYPES[base], Inches(x), Inches(y), Inches(w), Inches(h), data)
    except Exception as error:                                   # noqa: BLE001 — see the docstring
        message = f"chart could not be built: {type(error).__name__}: {error}"
        log.error(message)
        return _placeholder(slide, box, message, font=font, size_pt=base_pt, name=label)

    drawn: list[Any] = []
    try:
        if name:
            frame.name = name
        _style(frame.chart, model, geometry, text)
        # Last: rewrite into a chart element python-pptx cannot write. Everything above styled the
        # supported base, and the conversion preserves what still applies to the target type.
        if conv:
            _convert(frame.chart, base, conv)
        _overlay(slide, model, geometry, box, text, label, log, drawn)
    except Exception as error:                                   # noqa: BLE001 — see the docstring
        _rollback(slide, frame, drawn)
        message = f"chart could not be built: {type(error).__name__}: {error}"
        log.error(message)
        return _placeholder(slide, box, message, font=font, size_pt=base_pt, name=label)
    return frame


def _theme_accents(slide: Any) -> dict[int, str]:
    """accent1..accent6 of the slide's theme as RRGGBB; `{}` when the theme cannot be read."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    try:
        theme = slide.slide_layout.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
    except Exception:                                            # noqa: BLE001 — a nicety, never fatal
        return {}
    accents: dict[int, str] = {}
    for number in range(1, 7):
        node = root.find(f".//{_a('clrScheme')}/{_a(f'accent{number}')}")
        if node is None:
            continue
        colour = node.find(_a("srgbClr"))
        value = colour.get("val") if colour is not None else None
        if value is None:
            system = node.find(_a("sysClr"))
            value = system.get("lastClr") if system is not None else None
        if value and len(value) == 6:
            accents[number] = value.upper()
    return accents


#: A theme text colour is "dark" — usable as the chart's ink on a light slide — when white text on
#: it would pass WCAG AA (contrast ≥ 4.5:1), i.e. its relative luminance is at most 1.05/4.5 − 0.05.
DARK_INK_LUMINANCE = 1.05 / 4.5 - 0.05


def _theme_ink(slide: Any) -> str | None:
    """The deck theme's second text colour as RRGGBB — `tx2` through the master's colour map
    (`p:clrMap`, normally `tx2="dk2"`) — when it is dark; None when it is light or unreadable."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    try:
        master = slide.slide_layout.slide_master
        theme = master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
        mapping = master._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}clrMap")
        slot = (mapping.get("tx2") if mapping is not None else None) or "dk2"
        node = root.find(f".//{_a('clrScheme')}/{_a(slot)}")
        if node is None:
            return None
        colour = node.find(_a("srgbClr"))
        value = colour.get("val") if colour is not None else None
        if value is None:
            system = node.find(_a("sysClr"))
            value = system.get("lastClr") if system is not None else None
        if not value or len(value) != 6 or _luminance(value) > DARK_INK_LUMINANCE:
            return None
        return cast(str, value).upper()
    except Exception:                                            # noqa: BLE001 — a nicety, never fatal
        return None


def _ink_the_totals(model: dict[str, Any], theme_ink: str | None) -> None:
    """A one-colour waterfall paints its totals in the chart's ink (render check #5a): the theme's
    dark `tx2` where the deck has one, else the model's `WATERFALL_INK`."""
    plan = model.get("waterfall")
    if not plan or not plan.get("inkRoles") or not plan.get("colors") or theme_ink is None:
        return
    for role in plan["inkRoles"]:
        plan["colors"][role] = theme_ink


def _chart_data(model: dict[str, Any]) -> Any:
    """python-pptx chart data from the model. The model guarantees the shape (names are strings,
    every values list is as long as the categories), so this cannot raise on a spec's shape."""
    number_format = model["options"]["numberFormat"] or "General"
    data: Any
    if model["family"] == "scatter":
        data = XyChartData()  # type: ignore[no-untyped-call]
        for entry in model["series"]:
            added = data.add_series(entry["name"], number_format)
            for px, py in zip(entry["x"] or [], entry["y"] or [], strict=False):
                if px is None or py is None:
                    continue
                added.add_data_point(px, py)
        return data
    data = CategoryChartData()  # type: ignore[no-untyped-call]
    data.categories = list(model["categories"])
    series = model["waterfall"]["series"] if model["waterfall"] else model["series"]
    for entry in series:
        data.add_series(entry["name"], list(entry["values"]), number_format)
    return data


def _rollback(slide: Any, frame: Any, drawn: list[Any]) -> None:
    """Take a half-styled chart back out: the frame, its relationship (so the chart part and its
    workbook are unreferenced and python-pptx does not write them), and every overlay shape."""
    element = frame._element
    rId = element.chart_rId
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)
    if rId:
        slide.part.drop_rel(rId)
    for shape in drawn:
        node = shape._element
        if node.getparent() is not None:
            node.getparent().remove(node)


def _placeholder(slide: Any, box_in: tuple[float, float, float, float], message: str, *, font: str,
                 size_pt: float, name: str) -> Any:
    """A dashed red rectangle where the chart would be, saying why it is not.

    No theme style (so no theme fill), no fill, a 1 pt dashed `E5484D` outline, the message in the
    chart's font at ≤ 10 pt. A box too small to hold the message (under 120 × 40 px, or the text
    would not fit) keeps an empty text frame and carries the message in its name only.
    """
    x, y, w, h = box_in
    message, name = _xml_safe(message), _xml_safe(name)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    shape.fill.background()
    shape.line.color.rgb = _rgb(PLACEHOLDER_COLOR)
    shape.line.width = Pt(1)
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0

    size = min(float(size_pt), 10.0)
    width_px, height_px = w * config.PX_PER_IN, h * config.PX_PER_IN
    size_px = size / config.PT_PER_PX
    per_line = max(1, int(width_px // (0.55 * size_px)))
    lines = max(1, math.ceil(len(message) / per_line))
    fits = (width_px >= PLACEHOLDER_MIN_PX[0] and height_px >= PLACEHOLDER_MIN_PX[1]
            and lines * size_px * 1.3 <= height_px)
    if fits:
        run = frame.paragraphs[0].add_run()
        run.text = message
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = _rgb(PLACEHOLDER_COLOR)
        shape.name = f"{name} (chart not built)"
    else:
        shape.name = f"{name} (chart not built: {message})"[:255]
    return shape


# ------------------------------------------------------------------------------- chart internals


def _style(chart: Any, model: dict[str, Any], geometry: dict[str, Any], text: _Text) -> None:
    """Everything on the chart part, from the model: fills, fonts, legend, series, labels, axes."""
    options = model["options"]
    family = model["family"]
    base = model["base"]
    waterfall = family == "waterfall"
    bar_family = family in ("column", "bar", "waterfall")
    no_axes = family in ("pie", "doughnut")

    # The deck's own background must show through, so nothing paints a white card behind the plot.
    _no_fill(chart._chartSpace)
    _no_fill_plot(chart._chartSpace)

    chart.font.size = Pt(text.base_pt)
    chart.font.name = text.font
    chart.font.color.rgb = _rgb(text.ink)

    title = options["title"]
    chart.has_title = bool(title)
    if title:
        chart.chart_title.text_frame.text = str(title)
        chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(text.base_pt * 1.15)

    # --- legend -----------------------------------------------------------------------------
    legend = options["legend"]
    chart.has_legend = bool(legend)
    if chart.has_legend:
        chart.legend.position = LEGEND_POSITIONS.get(legend, XL_LEGEND_POSITION.BOTTOM)
        chart.legend.include_in_layout = False

    # --- series and points ---------------------------------------------------------------------
    fills = _style_waterfall_series(chart, model) if waterfall else _style_series(chart, model)
    for series_index in sorted(fills):
        _color_points(chart, series_index, fills[series_index])
    if bar_family:
        for plot_series in chart.series:
            if hasattr(plot_series, "invert_if_negative"):
                # A negative bar is drawn hollow unless the series says otherwise (WP-C §5).
                plot_series.invert_if_negative = False

    # --- plot-level options -------------------------------------------------------------------
    for plot in chart.plots:
        if bar_family:
            plot.gap_width = _int(options["gapWidth"])
            plot.overlap = _int(options["overlap"])
        try:
            vary = options["varyColors"]
            plot.vary_by_categories = bool(vary if vary is not None else no_axes)
            # Same reason as `c:delete`: python-pptx omits @val when it equals the schema default,
            # and a reader that wants the attribute then reads the opposite of what we meant.
            element = plot._element.find(_c("varyColors"))
            if element is not None:
                element.set("val", "1" if plot.vary_by_categories else "0")
        except (AttributeError, NotImplementedError):
            pass

    if waterfall:
        _apply_waterfall_labels(chart, model, fills, text)
    else:
        _apply_data_labels(chart, model, fills, text)

    # --- axes -----------------------------------------------------------------------------------
    if not no_axes and base not in BUBBLE_KINDS:
        # On an XY chart python-pptx's `category_axis` is the x value axis, so it takes the same
        # numeric treatment as the value axis — `categoryAxis.min` there really is a number.
        _apply_axis(chart.category_axis, options["categoryAxis"], text.font, text.base_pt, text.ink,
                    gridlines=False, default_format=None, category=family != "scatter")
        # A category name broken by `\n` keeps its break in the cached text (`c:pt/c:v`) and the
        # workbook, and PowerPoint draws it on two lines — nothing to write here.
        wanted = dict(options["valueAxis"])
        if waterfall and model["axis"] is not None:
            # A waterfall's value range is pinned: its connectors and labels are placed from it.
            for key in ("min", "max", "majorUnit"):
                wanted[key] = model["axis"][key]
        tick_format = chart_model.value_axis_format(model)
        if chart_model.is_percent_axis(family, model["stacked"]):
            # The spec states a 100 % axis in percent (the model reads a fraction as one); a
            # `percentStacked` axis is scaled 0–1, so `max: 100` written as is would draw every bar
            # as a 1 % sliver (render check #3, d2_s07).
            for key in ("min", "max", "majorUnit", "minorUnit"):
                if wanted.get(key) is not None:
                    wanted[key] = chart_model._round_to(wanted[key] / 100.0, 12)
        _apply_axis(chart.value_axis, wanted, text.font, text.base_pt, text.ink,
                    gridlines=options["gridlines"],
                    default_format=None if tick_format == "General" else tick_format, category=False)
        if bar_family and not options["valueAxis"]["reverse"]:
            shown = chart_model.axis_range(model)
            if shown is not None and shown["min"] < 0:
                # The category labels go to the plot's edge, not onto the zero line inside it.
                chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        if family != "scatter" and options["categoryAxis"]["reverse"]:
            # Reversed categories put the first one at the far end, and the value axis crosses at the
            # first category, so it would jump to the other side (C1 review #10).
            _value_axis_crosses_at_max(chart)

    if family == "doughnut" and options["holeSize"] is not None:
        _set_hole_size(chart, options["holeSize"])

    # --- waterfall connectors, in the chart (WP-C §9's fallback, taken on C3's measurement) -----
    if waterfall:
        connectors = _in_chart_connectors(chart, model)
        if chart.has_legend:
            plan_series = model["waterfall"]["series"]
            hidden = [index for index, entry in enumerate(plan_series)
                      if not entry["visible"] or entry["side"] == "-"]
            _legend_entries(chart, hidden + connectors)

    # --- plot rectangle -------------------------------------------------------------------------
    if geometry["pinned"]:
        _manual_plot_layout(chart, geometry["plotArea"])
        _pin_furniture(chart, model, geometry)


def _style_series(chart: Any, model: dict[str, Any]) -> dict[int, dict[int, str]]:
    """Colours, lines, markers and explosion per series; returns the per-point fills to write."""
    options = model["options"]
    family, base = model["family"], model["base"]
    line_like = family in ("line", "scatter")
    colors = model["colors"]
    usable = [c for c in colors if c is not None]
    fills: dict[int, dict[int, str]] = {}
    for index, plot_series in enumerate(chart.series):
        entry = model["series"][index] if index < len(model["series"]) else {}
        colour = colors[index % len(colors)] if colors else None
        width = entry.get("lineWidth") if entry.get("lineWidth") is not None else options["lineWidth"]
        dash = entry.get("dash") if entry.get("dash") is not None else options["dash"]
        if colour is not None:
            if line_like:
                _style_line(plot_series.format.line, colour, width if width is not None else 2.25, dash)
            else:
                plot_series.format.fill.solid()
                plot_series.format.fill.fore_color.rgb = _rgb(colour)
        elif line_like:
            _style_line(plot_series.format.line, None, width, dash)

        if base.startswith("line") and hasattr(plot_series, "smooth"):
            smooth = entry.get("smooth") if entry.get("smooth") is not None else options["smooth"]
            plot_series.smooth = bool(smooth)

        # `scatter` is Excel's markers-only flavour; `scatter_lines*` is the one with a line. Left
        # to itself PowerPoint joins the points of both, which turns eight companies into a zigzag.
        if base == "scatter" and not entry.get("line"):
            plot_series.format.line.fill.background()

        if base in XY_KINDS or base.startswith("line_markers"):
            default_marker = "none" if base.endswith("_no_markers") else "circle"
            wanted = entry.get("marker") if entry.get("marker") is not None else options["markers"]
            size = entry.get("markerSize") if entry.get("markerSize") is not None else options["markerSize"]
            _apply_marker(plot_series, default_marker if wanted is None else wanted, size, colour)

        explosion = entry.get("explosion") if entry.get("explosion") is not None else options["explosion"]
        if explosion is not None and family in ("pie", "doughnut"):
            # One `c:explosion` per series: python-pptx's exploded templates already write one
            # (`val="25"`), and a second is a schema error PowerPoint answers with a repair prompt
            # (C1 review #2).
            element = plot_series._element.find(_c("explosion"))
            if element is None:
                element = etree.Element(_c("explosion"))
                _insert_after(plot_series._element, element, ("spPr", "tx", "order", "idx"))
            element.set("val", str(_int(explosion)))

    # Per-point colours: a pie/doughnut always, a bar row when the spec names points.
    if usable and len(chart.series) == 1 and family in ("pie", "doughnut"):
        count = len(model["categories"])
        fills[0] = {i: colors[i % len(colors)] for i in range(count) if colors[i % len(colors)] is not None}
    for series_key, points in options["pointColors"].items():
        fills.setdefault(int(series_key), {}).update({int(p): c for p, c in points.items()})
    return fills


def _style_waterfall_series(chart: Any, model: dict[str, Any]) -> dict[int, dict[int, str]]:
    """The invisible bases, the visible series' colours, and every visible cell's fill.

    Single series: `colors` is [increase, decrease, total] per bar (`plan["colors"]`), and
    `pointColors` of series 0 override per point on every visible series carrying that point (both
    sides of zero). Several series: `colors[k]` is series k's colour (its theme accent when absent)
    on both `S_k` and `S_k−`, and `pointColors[k]` overrides per point.
    """
    plan = model["waterfall"]
    point_colors = model["options"]["pointColors"]
    single = len(model["series"]) == 1
    colors = model["colors"]
    fills: dict[int, dict[int, str]] = {}
    for index, plot_series in enumerate(chart.series):
        entry = plan["series"][index]
        if not entry["visible"]:
            plot_series.format.fill.background()
            plot_series.format.line.fill.background()
            continue
        source = entry["source"]
        overrides = point_colors.get(str(source), {})
        plot_series.format.fill.solid()
        if single:
            # Every visible cell carries its role colour in a `c:dPt`; the series fill is what a
            # legend key shows, so it is the rise colour, not the theme's accent (C1 review #14a).
            plot_series.format.fill.fore_color.rgb = _rgb(plan["colors"]["increase"])
        else:
            colour = colors[source] if source < len(colors) else None
            if colour is not None:
                plot_series.format.fill.fore_color.rgb = _rgb(colour)
            else:
                accent = getattr(MSO_THEME_COLOR, f"ACCENT_{source % 6 + 1}")
                plot_series.format.fill.fore_color.theme_color = accent
        for point, cell in enumerate(entry["values"]):
            if cell is None:
                continue
            chosen = overrides.get(str(point))
            if chosen is None and single:
                chosen = plan["colors"][plan["roles"][point]]
            if chosen is not None:
                fills.setdefault(index, {})[point] = chosen
    return fills


def _legend_entries(chart: Any, hidden: list[int]) -> None:
    """`c:legendEntry` deletions for these series' `c:idx` — python-pptx has no API for them.

    CT_Legend is `legendPos?, legendEntry*, layout?, overlay?, spPr?, txPr?, extLst?`, so each goes
    straight after `c:legendPos` and any entry already there.

    `c:legendEntry/c:idx` counts the entries **as the legend displays them**, read in PowerPoint (C1
    review #4; C3, `out/c3/ppt3/legends.json`): a horizontal legend (`t`, `b`) lists every series in
    document order; a vertical one (`r`, `l`, `tr` — `r` when `c:legendPos` carries no value) lists
    the stacked column's series top of the stack first — reversed, position `p` of the `n` bar
    series is entry `n − 1 − p` — and then the connectors' scatter group in its own order (entry
    `p`).
    """
    legend = chart._chartSpace.find(f".//{_c('legend')}")
    if legend is None:
        return
    count = len(chart._chartSpace.findall(f".//{_c('ser')}"))
    first_group = next((element for element in chart._chartSpace.find(f".//{_c('plotArea')}")
                        if etree.QName(element).localname.endswith("Chart")), None)
    bars = len(first_group.findall(_c("ser"))) if first_group is not None else count
    position_element = legend.find(_c("legendPos"))
    legend_position = "r" if position_element is None else (position_element.get("val") or "r")
    reversed_order = legend_position in ("r", "l", "tr")
    entries = sorted(bars - 1 - position if reversed_order and position < bars else position
                     for position in hidden if 0 <= position < count)
    for index in entries:
        entry = etree.Element(_c("legendEntry"))
        _sub(entry, "idx", val=str(index))
        _sub(entry, "delete", val="1")
        _insert_after(legend, entry, ("legendPos", "legendEntry"))


def _apply_marker(plot_series: Any, wanted: Any, size: Any, colour: str | None) -> None:
    """Marker style/size/colour. Without this PowerPoint assigns a different symbol per series."""
    if wanted is True:
        wanted = "circle"
    if wanted is False:
        wanted = "none"
    try:
        plot_series.marker.style = MARKERS.get(str(wanted).lower(), XL_MARKER_STYLE.CIRCLE)
        if str(wanted).lower() not in ("none", "auto"):
            plot_series.marker.size = _int(size if size is not None else 7)
            if colour is not None:
                marker_format = plot_series.marker.format
                marker_format.fill.solid()
                marker_format.fill.fore_color.rgb = _rgb(colour)
                marker_format.line.color.rgb = _rgb(colour)
    except (AttributeError, NotImplementedError, ValueError):
        pass


def _label_ink(plot_series: Any, series_index: int, fills: dict[int, dict[int, str]], ink: str,
               inside: bool, accents: dict[int, str] | None = None) -> tuple[str, dict[int, str]]:
    """The series' label colour and the points that differ from it.

    A label sitting ON a coloured bar has to contrast with the bar, not with the slide. The series
    colour takes the reading that suits the most points and only the minority gets a per-point
    override: PowerPoint honours both, PptxRender reads only the series one, so the majority reading
    is the one that survives into the gate's render.
    """
    if not inside:
        return ink, {}
    # Every point's fill, not just the overridden ones: a series of blue bars with one pale point
    # must not take its reading from that one point.
    colours: dict[int, str] = {}
    series_fill = _series_fill_colour(plot_series) or _scheme_fill_colour(plot_series, accents or {})
    if series_fill:
        colours = {point: series_fill for point in _drawn_points(plot_series)}
    colours.update(fills.get(series_index) or {})
    readings = {point: _contrast(colour, dark=ink) for point, colour in colours.items()}
    if not readings:
        return ink, {}
    series_colour = _majority(readings.values())
    return series_colour, {point: colour for point, colour in readings.items() if colour != series_colour}


def _apply_data_labels(chart: Any, model: dict[str, Any], fills: dict[int, dict[int, str]], text: _Text) -> None:
    """Data labels, per series, in the contract's vocabulary (every family but the waterfall).

    A series-level `<c:dLbls>` *replaces* the plot's rather than inheriting from it, so every flag
    is restated per series; that is also what makes `dataLabels: [true, false]` possible at all.
    """
    options = model["options"]
    labels = options["dataLabels"]
    if not any(labels):
        return
    number_format = options["labelFormat"] or options["numberFormat"]
    if number_format == "General" and not options["labelFormat"]:
        number_format = None
    position = options["labelPosition"]
    enum_position, xml_position = LABEL_POSITIONS[position] if position else (None, None)
    no_axes = model["family"] in ("pie", "doughnut")

    for index, plot_series in enumerate(chart.series):
        if not (labels[index] if index < len(labels) else False):
            _hide_series_labels(plot_series)
            continue
        entry = model["series"][index] if index < len(model["series"]) else {}
        texts = _text_overrides(entry.get("labels"))
        if not hasattr(plot_series, "data_labels"):
            # XY and bubble series have no plot-level `dLbls` in the schema and python-pptx offers
            # no API for the series-level one, so it is written directly.
            ink = text.label_color or text.ink
            _raw_series_labels(plot_series._element, show_value=True, number_format=number_format,
                               xml_position=xml_position, font=text.font, size_pt=text.label_pt,
                               color=ink, bold=text.label_bold)
            if texts:
                _point_labels(plot_series, texts, font=text.font, size_pt=text.label_pt,
                              xml_position=xml_position, default_format=number_format,
                              default_color=ink, bold=text.label_bold)
            continue
        series_labels = plot_series.data_labels
        series_labels.show_value = True
        series_labels.show_category_name = False
        series_labels.show_series_name = False
        series_labels.show_legend_key = False
        series_labels.show_percentage = False
        if number_format:
            series_labels.number_format = number_format
            series_labels.number_format_is_linked = False
        if enum_position is not None:
            try:
                series_labels.position = enum_position
            except (ValueError, NotImplementedError):
                pass
        series_labels.font.size = Pt(text.label_pt)
        series_labels.font.name = text.font
        if text.label_bold:
            series_labels.font.bold = True

        inside = position in INSIDE or (no_axes and position in ("bestFit", None))
        if text.label_color:
            series_colour, minority = text.label_color, cast(dict[int, str], {})
        else:
            series_colour, minority = _label_ink(plot_series, index, fills, text.ink, inside)
        series_labels.font.color.rgb = _rgb(series_colour)
        overrides = {point: {"color": colour} for point, colour in minority.items()}
        for point, settings in texts.items():
            overrides.setdefault(point, {}).update(settings)
        if overrides:
            _point_labels(plot_series, overrides, font=text.font, size_pt=text.label_pt,
                          xml_position=xml_position, default_format=number_format,
                          default_color=series_colour, bold=text.label_bold)


def _text_overrides(texts: list[Any] | None) -> dict[int, dict[str, Any]]:
    """`series[k].dataLabels` as `_point_labels` settings: a text is frozen rich text on that
    point's label (render check #6), `""` deletes the point's label, `None` leaves the value."""
    out: dict[int, dict[str, Any]] = {}
    for point, value in enumerate(texts or []):
        if value is None:
            continue
        out[point] = {"delete": True} if value == "" else {"text": _xml_safe(value)}
    return out


def _apply_waterfall_labels(chart: Any, model: dict[str, Any], fills: dict[int, dict[int, str]], text: _Text) -> None:
    """Native labels inside the bars, frozen where a segment crosses zero (WP-C §3.4).

    Each visible cell belongs to one segment of the plan: the part holding a segment's label shows
    its cell in the role's format (`"+"0;"+"0` for a rise, `"-"0;"-"0` for a fall, the author's for
    a total), or — for a segment split across zero, where neither cell is the step — frozen rich
    text; the other part's label is deleted. `outEnd` deletes every native label (the text boxes
    of `_waterfall_labels` take their place).
    """
    options = model["options"]
    plan = model["waterfall"]
    labels = options["dataLabels"]
    position = options["labelPosition"] or "center"
    series_format = plan["formats"]["total"]
    per_series: dict[int, dict[int, dict[str, Any]]] = {}
    for segment in plan["segments"]:
        for series_index, point in segment["hide"]:
            per_series.setdefault(series_index, {})[point] = {"delete": True}
        label = segment["label"]
        settings: dict[str, Any] = {}
        if label["text"] == "":
            settings = {"delete": True}                  # the author's `""`: no label on this bar
        elif label["text"] is not None:
            settings = {"text": _xml_safe(label["text"])}
        if label["text"] is None and label["format"] != series_format:
            settings["numFmt"] = label["format"]
        if segment["lo"] == segment["hi"]:
            # A zero-height bar's label sits on the slide, not on a bar: the chart's own ink.
            settings["color"] = text.label_color or text.ink
            settings["source"] = segment["source"]
        per_series.setdefault(label["series"], {})[label["point"]] = settings

    enum_position, xml_position = LABEL_POSITIONS.get(position, LABEL_POSITIONS["center"])
    for index, plot_series in enumerate(chart.series):
        entry = plan["series"][index]
        source = entry["source"]
        if not entry["visible"]:
            _hide_series_labels(plot_series)
            # A zero-height bar below zero has no cell of its own (a 0 would stack at the zero line):
            # its label is frozen text at the inside end of `Base−`'s bar, i.e. at the level
            # (C1 review #9).
            frozen = {point: settings for point, settings in (per_series.get(index) or {}).items()
                      if settings.get("text") is not None
                      and (labels[settings["source"]] if settings["source"] < len(labels) else False)}
            if frozen and position != "outEnd":
                _point_labels(plot_series, frozen, font=text.font, size_pt=text.label_pt,
                              xml_position="inEnd", default_format=None,
                              default_color=text.label_color or text.ink, bold=text.label_bold)
            continue
        if not (labels[source] if source < len(labels) else False):
            _hide_series_labels(plot_series)
            continue
        if position == "outEnd":
            _delete_series_labels(plot_series)
            continue
        series_labels = plot_series.data_labels
        series_labels.show_value = True
        series_labels.show_category_name = False
        series_labels.show_series_name = False
        series_labels.show_legend_key = False
        series_labels.show_percentage = False
        series_labels.number_format = series_format
        series_labels.number_format_is_linked = False
        try:
            series_labels.position = enum_position
        except (ValueError, NotImplementedError):
            pass
        series_labels.font.size = Pt(text.label_pt)
        series_labels.font.name = text.font
        if text.label_bold:
            series_labels.font.bold = True
        if text.label_color:
            series_colour, minority = text.label_color, cast(dict[int, str], {})
        else:
            series_colour, minority = _label_ink(plot_series, index, fills, text.ink, inside=True,
                                                 accents=text.accents)
        series_labels.font.color.rgb = _rgb(series_colour)

        overrides: dict[int, dict[str, Any]] = {}
        for point, settings in (per_series.get(index) or {}).items():
            settings = {key: value for key, value in settings.items() if key != "source"}
            if settings:
                overrides[point] = settings
        for point, colour in minority.items():
            if not (overrides.get(point) or {}).get("delete"):
                overrides.setdefault(point, {}).setdefault("color", colour)
        # A frozen label carries its own run colour: the bar's contrast ink.
        for settings in overrides.values():
            if settings.get("text") is not None and "color" not in settings:
                settings["color"] = series_colour
        if overrides:
            _point_labels(plot_series, overrides, font=text.font, size_pt=text.label_pt,
                          xml_position=xml_position, default_format=series_format,
                          default_color=series_colour, bold=text.label_bold)


def _majority(values: Iterable[str]) -> str:
    """The most common value, ties broken by sort order — a vote has to be deterministic."""
    counts: dict[str, int] = {}
    for value in values:
        counts[value] = counts.get(value, 0) + 1
    return sorted(counts, key=lambda value: (-counts[value], value))[0]


#: What precedes `c:dLbls` in a series (CT_BarSer / CT_LineSer / CT_PieSer / CT_ScatterSer).
_BEFORE_DLBLS = ("idx", "order", "tx", "spPr", "invertIfNegative", "pictureOptions", "marker",
                 "explosion", "dPt")


def _hide_series_labels(plot_series: Any) -> None:
    if not hasattr(plot_series, "data_labels"):
        holder = plot_series._element.find(_c("dLbls"))
        if holder is None:
            holder = etree.Element(_c("dLbls"))
            _insert_after(plot_series._element, holder, _BEFORE_DLBLS)
        _sub(holder, "delete", val="1")
        return
    labels = plot_series.data_labels
    labels.show_value = False
    labels.show_category_name = False
    labels.show_series_name = False
    labels.show_legend_key = False
    labels.show_percentage = False


def _delete_series_labels(plot_series: Any) -> None:
    """`<c:dLbls><c:delete val="1"/></c:dLbls>`: no native label on any point of this series."""
    element = plot_series._element
    for existing in element.findall(_c("dLbls")):
        element.remove(existing)
    holder = etree.Element(_c("dLbls"))
    _sub(holder, "delete", val="1")
    _insert_after(element, holder, _BEFORE_DLBLS)


def _raw_series_labels(series_element: Any, *, show_value: bool, number_format: str | None,
                       xml_position: str | None, font: str, size_pt: float, color: str,
                       bold: bool = False) -> None:
    """A series-level `<c:dLbls>` written by hand, for the series types python-pptx cannot label."""
    holder = series_element.find(_c("dLbls"))
    if holder is not None:
        series_element.remove(holder)
    holder = etree.Element(_c("dLbls"))
    if number_format:
        element = etree.SubElement(holder, _c("numFmt"))
        element.set("formatCode", number_format)
        element.set("sourceLinked", "0")
    shape_properties = etree.SubElement(holder, _c("spPr"))
    etree.SubElement(shape_properties, _a("noFill"))
    etree.SubElement(etree.SubElement(shape_properties, _a("ln")), _a("noFill"))
    text_properties = etree.SubElement(holder, _c("txPr"))
    etree.SubElement(text_properties, _a("bodyPr"))
    etree.SubElement(text_properties, _a("lstStyle"))
    paragraph = etree.SubElement(text_properties, _a("p"))
    properties = etree.SubElement(etree.SubElement(paragraph, _a("pPr")), _a("defRPr"))
    properties.set("sz", str(int(round(size_pt * 100))))
    if bold:
        properties.set("b", "1")
    fill = etree.SubElement(properties, _a("solidFill"))
    etree.SubElement(fill, _a("srgbClr")).set("val", _hex(color))
    etree.SubElement(properties, _a("latin")).set("typeface", font)
    etree.SubElement(paragraph, _a("endParaRPr")).set("lang", "en-US")
    if xml_position:
        etree.SubElement(holder, _c("dLblPos")).set("val", xml_position)
    for tag, value in (("showLegendKey", "0"), ("showVal", "1" if show_value else "0"),
                       ("showCatName", "0"), ("showSerName", "0"), ("showPercent", "0"),
                       ("showBubbleSize", "0")):
        etree.SubElement(holder, _c(tag)).set("val", value)
    _insert_after(series_element, holder, _BEFORE_DLBLS)


def _point_count(plot_series: Any) -> int:
    try:
        return len(list(plot_series.values))
    except (AttributeError, TypeError):
        return 0


def _drawn_points(plot_series: Any) -> list[int]:
    """The points that carry a value — a blank cell draws no bar and needs no label ink."""
    try:
        return [point for point, value in enumerate(plot_series.values) if value is not None]
    except (AttributeError, TypeError):
        return []


def _scheme_fill_colour(plot_series: Any, accents: dict[int, str]) -> str | None:
    """The RGB behind a series' `schemeClr accentN` fill, from the deck's theme (or None)."""
    node = plot_series._element.find(f"{_c('spPr')}/{_a('solidFill')}/{_a('schemeClr')}")
    if node is None:
        return None
    name = node.get("val") or ""
    if name.startswith("accent") and name[6:].isdigit():
        return accents.get(int(name[6:]))
    return None


def _series_fill_colour(plot_series: Any) -> str | None:
    fill = plot_series.format.fill
    try:
        if fill.type is not None and fill.fore_color.type is not None:
            return str(fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _run_properties(parent: Any, tag: str, *, size_pt: float, color: str, font: str, lang: bool = False,
                    bold: bool = False) -> Any:
    """`<a:defRPr>`/`<a:rPr>` with size, a solid fill and a latin face — children in schema order —
    and `b="1"` for `options.labelStyle.bold`."""
    properties = etree.SubElement(parent, _a(tag))
    if lang:
        properties.set("lang", "en-US")
    properties.set("sz", str(int(round(size_pt * 100))))
    if bold:
        properties.set("b", "1")
    fill = etree.SubElement(properties, _a("solidFill"))
    etree.SubElement(fill, _a("srgbClr")).set("val", _hex(color))
    etree.SubElement(properties, _a("latin")).set("typeface", font)
    return properties


def _point_labels(plot_series: Any, overrides: dict[int, dict[str, Any]], *, font: str,
                  size_pt: float, xml_position: str | None, default_format: str | None,
                  default_color: str = "12233B", bold: bool = False) -> None:
    """One `<c:dLbl>` per point that needs its own colour, number format, text, or no label.

    python-pptx has no API for this, and `CT_DLbl` is order-sensitive:
    `idx, (delete | layout?, tx?, numFmt?, spPr?, txPr?, dLblPos?, showLegendKey?, showVal?,
    showCatName?, showSerName?, showPercent?, showBubbleSize?, separator?), extLst?`. Out of order,
    PowerPoint repairs the file and drops content. Settings per point:

    * `delete` — `<c:idx/><c:delete val="1"/>` and nothing else (PptxRender reads it into
      `HiddenLabels`);
    * `text` — frozen rich text (`c:tx/c:rich`), which carries its own run properties, so no
      `numFmt` and no `txPr`; a `\\n` in it starts a new `a:p`, as PowerPoint writes a label
      the user breaks with Enter;
    * `numFmt`, `color` — the label's own format and ink.

    `bold` (`options.labelStyle.bold`) writes `b="1"` on every run property this writes.
    """
    holder = plot_series._element.find(_c("dLbls"))
    if holder is None:
        return
    for point in sorted(overrides):
        settings = overrides[point]
        label = etree.Element(_c("dLbl"))
        etree.SubElement(label, _c("idx")).set("val", str(point))
        if settings.get("delete"):
            etree.SubElement(label, _c("delete")).set("val", "1")
            _insert_after(holder, label, ("dLbl",))
            continue
        colour = settings.get("color") or default_color
        frozen = settings.get("text")
        if frozen is not None:
            rich = etree.SubElement(etree.SubElement(label, _c("tx")), _c("rich"))
            etree.SubElement(rich, _a("bodyPr"))
            etree.SubElement(rich, _a("lstStyle"))
            for line in chart_model.text_lines(str(frozen)):
                paragraph = etree.SubElement(rich, _a("p"))
                _run_properties(etree.SubElement(paragraph, _a("pPr")), "defRPr",
                                size_pt=size_pt, color=colour, font=font, bold=bold)
                run = etree.SubElement(paragraph, _a("r"))
                _run_properties(run, "rPr", size_pt=size_pt, color=colour, font=font, lang=True, bold=bold)
                etree.SubElement(run, _a("t")).text = line
        else:
            number_format = settings.get("numFmt") or default_format
            if number_format:
                element = etree.SubElement(label, _c("numFmt"))
                element.set("formatCode", number_format)
                element.set("sourceLinked", "0")
        shape_properties = etree.SubElement(label, _c("spPr"))
        etree.SubElement(shape_properties, _a("noFill"))
        etree.SubElement(etree.SubElement(shape_properties, _a("ln")), _a("noFill"))

        if frozen is None:
            text_properties = etree.SubElement(label, _c("txPr"))
            etree.SubElement(text_properties, _a("bodyPr"))
            etree.SubElement(text_properties, _a("lstStyle"))
            paragraph = etree.SubElement(text_properties, _a("p"))
            _run_properties(etree.SubElement(paragraph, _a("pPr")), "defRPr",
                            size_pt=size_pt, color=colour, font=font, bold=bold)
            etree.SubElement(paragraph, _a("endParaRPr")).set("lang", "en-US")

        if xml_position:
            etree.SubElement(label, _c("dLblPos")).set("val", xml_position)
        for tag, value in (("showLegendKey", "0"), ("showVal", "1"), ("showCatName", "0"),
                           ("showSerName", "0"), ("showPercent", "0"), ("showBubbleSize", "0")):
            etree.SubElement(label, _c(tag)).set("val", value)
        _insert_after(holder, label, ("dLbl",))


def _color_points(chart: Any, series_index: int, points: dict[int, str]) -> None:
    """One `<c:dPt>` per coloured point — python-pptx has no API for per-point fills."""
    all_series = chart._chartSpace.findall(f".//{_c('ser')}")
    if series_index >= len(all_series):
        return
    series = all_series[series_index]
    # PowerPoint reads a bar point's `dPt` without `c:invertIfNegative` as inverted — whatever the
    # series says — and draws a negative point hollow (C1 review #1). CT_DPt: idx, invertIfNegative?,
    # marker?, bubble3D?, explosion?, spPr?.
    bar_series = etree.QName(series.getparent()).localname in ("barChart", "bar3DChart")
    for point in sorted(points):
        data_point = etree.Element(_c("dPt"))
        _sub(data_point, "idx", val=str(point))
        if bar_series:
            _sub(data_point, "invertIfNegative", val="0")
        _sub(data_point, "bubble3D", val="0")
        shape_properties = etree.SubElement(data_point, _c("spPr"))
        fill = etree.SubElement(shape_properties, _a("solidFill"))
        etree.SubElement(fill, _a("srgbClr")).set("val", _hex(points[point]))
        # dPt follows idx/order/tx/spPr/invertIfNegative/explosion/marker and precedes dLbls/cat/val.
        _insert_after(series, data_point,
                      ("idx", "order", "tx", "spPr", "invertIfNegative", "pictureOptions", "explosion",
                       "marker", "dPt"))


def _apply_axis(axis: Any, wanted: Any, font: str, base_pt: float, ink: str, *, gridlines: Any,
                default_format: str | None, category: bool) -> None:
    """`{min,max,majorUnit,minorUnit,visible,format,reverse,title,bold}` onto a python-pptx axis."""
    wanted = dict(wanted or {}) if isinstance(wanted, dict) else ({} if wanted is not False else {"visible": False})
    visible = wanted.get("visible", True)
    axis.visible = bool(visible)
    # python-pptx writes a bare `<c:delete/>` for a hidden axis, because CT_Boolean's @val defaults
    # to true. Valid, but a reader that looks for the attribute (PptxRender does) then draws the
    # axis anyway, so state it: an axis the design hid must be hidden in the render too.
    delete = axis._element.find(_c("delete"))
    if delete is not None:
        delete.set("val", "0" if visible else "1")

    axis.has_minor_gridlines = False
    axis.has_major_gridlines = bool(gridlines)
    if gridlines:
        settings = gridlines if isinstance(gridlines, dict) else {}
        _style_line(axis.major_gridlines.format.line,
                    settings.get("color") or "E4E9F0", settings.get("width") or 0.75,
                    settings.get("dash"))

    if visible:
        axis.tick_labels.font.size = Pt(base_pt)
        axis.tick_labels.font.name = font
        axis.tick_labels.font.color.rgb = _rgb(ink)
        if wanted.get("bold"):
            axis.tick_labels.font.bold = True          # `valueAxis.bold` / `categoryAxis.bold` (G24)
        number_format = wanted.get("format") or default_format
        if number_format:
            axis.tick_labels.number_format = number_format
            axis.tick_labels.number_format_is_linked = False
        _style_line(axis.format.line, "C9D3E0", 0.75, None)

    scaling = axis._element.find(_c("scaling"))
    if wanted.get("reverse") and scaling is not None:
        orientation = scaling.find(_c("orientation"))
        if orientation is None:
            orientation = etree.Element(_c("orientation"))
            scaling.insert(0, orientation)
        orientation.set("val", "maxMin")

    if category:
        # A category axis has no numeric scale; `majorUnit` is how many categories to skip between
        # labels, and min/max are category positions — both live in the XML, not in python-pptx.
        unit = _number(wanted.get("majorUnit"))
        if unit is not None and unit >= 1:
            # tickLblSkip precedes tickMarkSkip in CT_CatAx, and both come after lblOffset.
            for tag in ("tickLblSkip", "tickMarkSkip"):
                element = axis._element.find(_c(tag))
                if element is None:
                    element = etree.Element(_c(tag))
                    _insert_after(axis._element, element,
                                  ("crosses", "crossesAt", "crossAx", "auto", "lblAlgn",
                                   "lblOffset", "tickLblSkip"))
                element.set("val", str(int(unit)))
        for key, tag in (("min", "min"), ("max", "max")):
            value = _number(wanted.get(key))
            if value is not None and scaling is not None:
                element = scaling.find(_c(tag))
                if element is None:
                    element = etree.SubElement(scaling, _c(tag))
                element.set("val", str(value))
    else:
        for key, attribute in (("min", "minimum_scale"), ("max", "maximum_scale"),
                               ("majorUnit", "major_unit"), ("minorUnit", "minor_unit")):
            value = _number(wanted.get(key))
            if value is not None:
                setattr(axis, attribute, value)

    title = wanted.get("title")
    axis.has_title = bool(title)
    if title:
        axis.axis_title.text_frame.text = str(title)
        for paragraph in axis.axis_title.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(base_pt)
                run.font.name = font
                run.font.color.rgb = _rgb(ink)


def _value_axis_crosses_at_max(chart: Any) -> None:
    """`c:valAx/c:crosses val="max"`: the value axis crosses the category axis at its last
    category — which, reversed, is the near one, so the axis stays where the preview and the pinned
    gutters put it. python-pptx has no API for it (`crosses` exists on its ValueAxis only, and there
    writes the *category* axis's element). CT_ValAx: … crossAx, (crosses | crossesAt)?, crossBetween?…"""
    value_axis = chart.value_axis._element
    for existing in value_axis.findall(_c("crossesAt")):
        value_axis.remove(existing)
    crosses = value_axis.find(_c("crosses"))
    if crosses is None:
        crosses = etree.Element(_c("crosses"))
        _insert_after(value_axis, crosses, ("crossAx",))
    crosses.set("val", "max")


def _set_hole_size(chart: Any, hole: Any) -> None:
    for doughnut in chart._chartSpace.findall(f".//{_c('doughnutChart')}"):
        element = doughnut.find(_c("holeSize"))
        if element is None:
            element = etree.Element(_c("holeSize"))
            _insert_after(doughnut, element, ("firstSliceAng", "dLbls", "ser", "varyColors"))
        element.set("val", str(_int(hole)))


def _manual_plot_layout(chart: Any, fractions: dict[str, Any]) -> None:
    """`c:manualLayout` for the **inner** plot rect — the axes rectangle, excluding tick labels.

    Written only for a pinned chart (WP-C §3.6). Fractions of the frame, which is exactly what
    `plotRect / box` gives, so the emitted plot lands on the design's plot rect (the structural
    check reads it back from here and checks it within 4 px).
    """
    plot_area = chart._chartSpace.find(f".//{_c('plotArea')}")
    if plot_area is None:
        return
    layout = plot_area.find(_c("layout"))
    if layout is None:
        layout = etree.Element(_c("layout"))
        plot_area.insert(0, layout)
    for existing in layout.findall(_c("manualLayout")):
        layout.remove(existing)
    manual = etree.SubElement(layout, _c("manualLayout"))
    _sub(manual, "layoutTarget", val="inner")
    _sub(manual, "xMode", val="edge")
    _sub(manual, "yMode", val="edge")
    for tag in ("x", "y", "w", "h"):
        _sub(manual, tag, val=repr(round(float(fractions[tag]), 6)))


def _pin_furniture(chart: Any, model: dict[str, Any], geometry: dict[str, Any]) -> None:
    """What a pinned chart needs besides its plot rect for PowerPoint to keep that rect (WP-C C3).

    Measured with PowerPoint COM (`chart_model` AXIS_LINE_EM and its neighbours): PowerPoint keeps a
    written inner rect to 0.1 px when the text around it fits the gutters, and otherwise moves or
    shrinks the plot — which PptxRender never does, so the overlay lands beside the bars in
    PowerPoint only (C1 review #7). Two things it would otherwise decide itself are fixed here:

    * the category labels of a column, line or area chart stay horizontal (`a:bodyPr rot="0"
      vert="horz"`): auto-rotated labels need a gutter nobody can predict — PowerPoint skips labels
      that do not fit instead, as PptxRender does — and do not wrap (`wrap="none"`), because every
      wrapped line takes a line's height off the plot whatever the gutter;
    * the legend sits in `chart_model.legend_box` (`c:legend/c:layout/c:manualLayout`, edge mode),
      below the category labels rather than on them.
    """
    family = model["family"]
    if family in ("column", "waterfall", "line", "area") and model["options"]["categoryAxis"]["visible"]:
        category_axis = chart._chartSpace.find(f".//{_c('catAx')}")
        body = category_axis.find(f"{_c('txPr')}/{_a('bodyPr')}") if category_axis is not None else None
        if body is not None:
            body.set("rot", "0")
            body.set("vert", "horz")
            body.set("wrap", "none")
    box = geometry.get("legend")
    legend = chart._chartSpace.find(f".//{_c('legend')}")
    if box is None or legend is None:
        return
    for existing in legend.findall(_c("layout")):
        legend.remove(existing)
    layout = etree.Element(_c("layout"))
    manual = etree.SubElement(layout, _c("manualLayout"))
    _sub(manual, "xMode", val="edge")
    _sub(manual, "yMode", val="edge")
    for tag in ("x", "y", "w", "h"):
        _sub(manual, tag, val=repr(round(float(box[tag]), 6)))
    # CT_Legend: legendPos?, legendEntry*, layout?, overlay?, spPr?, txPr?
    _insert_after(legend, layout, ("legendPos", "legendEntry"))


#: The secondary axis pair the in-chart connectors are plotted on (xs:unsignedInt, clear of
#: python-pptx's template ids, which are negative).
CONNECTOR_AXIS_IDS = (500100001, 500100002)
#: The name of the connector series, numbered (a literal: not data, so the workbook does not hold it).
CONNECTOR_SERIES = "Connector"


def connector_points(model: dict[str, Any]) -> list[tuple[list[float], list[float]]]:
    """The in-chart connectors, one `(x, y)` pair of two-point literals each. x is in category units
    on an axis running 0..count — slot i spans [i, i + 1] exactly as the category axis lays it out —
    so a connector runs from bar i's edge, `i + 0.5 + bar/2`, to bar i + 1's, `i + 1.5 − bar/2`,
    with `bar = 1 / (1 + gapWidth/100)`; y is the level on the pinned value axis. One series per
    connector, not one broken by blanks: PowerPoint plots a scatter series beside a category chart
    only up to the category count's points (C3, `out/c3/ppt3`: D drew 2 of its 5). `[]` when
    nothing is drawn."""
    plan = model.get("waterfall") or {}
    if model["options"].get("connectors") is False or not plan.get("connectors"):
        return []
    half = 0.5 / (1.0 + max(model["options"]["gapWidth"], 0.0) / 100.0)
    out: list[tuple[list[float], list[float]]] = []
    for connector in plan["connectors"]:
        index, level = connector["index"], connector["level"]
        out.append(([round(index + 0.5 + half, 6), round(index + 1.5 - half, 6)], [level, level]))
    return out


def _num_lit(parent: Any, tag: str, values: Sequence[float | None]) -> None:
    """`<c:{tag}><c:numLit>` — formatCode, ptCount, one c:pt per value; a None is a blank (no pt)."""
    literal = etree.SubElement(etree.SubElement(parent, _c(tag)), _c("numLit"))
    etree.SubElement(literal, _c("formatCode")).text = "General"
    _sub(literal, "ptCount", val=str(len(values)))
    for index, value in enumerate(values):
        if value is None:
            continue
        point = _sub(literal, "pt", idx=str(index))
        etree.SubElement(point, _c("v")).text = repr(float(value))


def _connector_axis(ax_id: int, cross_id: int, position: str, low: float, high: float, reverse: bool) -> Any:
    """A deleted `c:valAx` of the connector pair (CT_ValAx order: axId, scaling, delete, axPos,
    majorTickMark, minorTickMark, tickLblPos, crossAx, crosses, crossBetween)."""
    axis = etree.Element(_c("valAx"))
    _sub(axis, "axId", val=str(ax_id))
    scaling = _sub(axis, "scaling")
    _sub(scaling, "orientation", val="maxMin" if reverse else "minMax")
    _sub(scaling, "max", val=repr(float(high)))
    _sub(scaling, "min", val=repr(float(low)))
    _sub(axis, "delete", val="1")
    _sub(axis, "axPos", val=position)
    _sub(axis, "majorTickMark", val="none")
    _sub(axis, "minorTickMark", val="none")
    _sub(axis, "tickLblPos", val="none")
    _sub(axis, "crossAx", val=str(cross_id))
    _sub(axis, "crosses", val="max")
    _sub(axis, "crossBetween", val="midCat")
    return axis


def _in_chart_connectors(chart: Any, model: dict[str, Any]) -> list[int]:
    """A waterfall's connectors as a series of the chart itself: a `c:scatterChart` group in the
    same plot area, on a secondary axis pair laid over the primary one (x 0..count in category
    units, y the pinned value axis), so they sit on the bars wherever PowerPoint lays the plot out.

    WP-C §9's fallback, taken because PowerPoint's own layout measured up to 23 px off a pinned
    overlay's bars once category labels outgrow their slots (C3, `out/c3/ppt2`): PowerPoint moves,
    widens or shrinks a pinned plot then, PptxRender never does. Returns the new series' position
    among the chart's `c:ser` elements (for their legend entries); `[]` when there is nothing to draw.
    PptxRender draws none of these connectors (C3's renders of charts-hardening, charts-negative,
    p05 and p11: bars and labels drawn, no connector) — a renderer gap inside the masked frame;
    `measure_overlays` on its PNG reports each pinned chart's connectors as missing.
    """
    pairs = connector_points(model)
    axis = model["axis"]
    if not pairs or axis is None:
        return []
    style = model["options"]["connectors"]
    space = chart._chartSpace
    plot_area = space.find(f".//{_c('plotArea')}")
    existing = space.findall(f".//{_c('ser')}")
    index = max([int(e.find(_c("idx")).get("val")) for e in existing] + [-1]) + 1
    count = len(model["categories"])
    x_id, y_id = CONNECTOR_AXIS_IDS

    group = etree.Element(_c("scatterChart"))
    _sub(group, "scatterStyle", val="lineMarker")
    _sub(group, "varyColors", val="0")
    for offset, (xs, ys) in enumerate(pairs):
        series = _sub(group, "ser")
        _sub(series, "idx", val=str(index + offset))
        _sub(series, "order", val=str(index + offset))
        etree.SubElement(_sub(series, "tx"), _c("v")).text = f"{CONNECTOR_SERIES} {offset + 1}"
        shape_properties = _sub(series, "spPr")
        line = etree.SubElement(shape_properties, _a("ln"))
        line.set("w", str(int(round(float(style["width"]) * 12700))))
        line.set("cap", "flat")
        etree.SubElement(etree.SubElement(line, _a("solidFill")), _a("srgbClr")).set("val", _hex(style["color"]))
        _apply_dash(line, style["dash"], float(style["width"] or 1.0) / config.PT_PER_PX)
        _sub(_sub(series, "marker"), "symbol", val="none")
        _sub(_sub(series, "dLbls"), "delete", val="1")
        _num_lit(series, "xVal", xs)
        _num_lit(series, "yVal", ys)
        _sub(series, "smooth", val="0")
    _sub(group, "axId", val=str(x_id))
    _sub(group, "axId", val=str(y_id))
    # CT_PlotArea: layout?, the chart groups, then the axes, then dTable?, spPr?.
    _insert_after(plot_area, group, ("layout", "barChart", "bar3DChart", "lineChart", "areaChart"))
    options = model["options"]
    x_axis = _connector_axis(x_id, y_id, "t", 0.0, float(count), bool(options["categoryAxis"]["reverse"]))
    y_axis = _connector_axis(y_id, x_id, "r", axis["min"], axis["max"], bool(options["valueAxis"]["reverse"]))
    _insert_after(plot_area, x_axis, ("layout", "barChart", "bar3DChart", "lineChart", "areaChart",
                                      "scatterChart", "catAx", "dateAx", "valAx", "serAx"))
    _insert_after(plot_area, y_axis, ("layout", "barChart", "bar3DChart", "lineChart", "areaChart",
                                      "scatterChart", "catAx", "dateAx", "valAx", "serAx"))
    return [len(existing) + offset for offset in range(len(pairs))]


def _overlay(slide: Any, model: dict[str, Any], geometry: dict[str, Any], box: tuple[float, float, float, float], text: _Text,
             label: str, log: _Diagnostics, drawn: list[Any]) -> None:
    """Waterfall `outEnd` labels and reference lines on top of the frame (connectors are in the
    chart: `_in_chart_connectors`)."""
    options = model["options"]
    if not model["overlay"]:
        if options["referenceLines"] and model["family"] not in ("pie", "doughnut") and model["axis"] is None:
            log.warn("overlay lines dropped: no value-axis range — set options.valueAxis.min and max so "
                     "the position is measured, not guessed")
        return
    lines = [line for line in overlay_geometry(None, box, model=model, geometry=geometry)
             if line["kind"] != "connector"]
    _draw_overlay(slide, lines, font=text.font, size_pt=text.label_pt, color=text.ink, name=label,
                  drawn=drawn, label_color=text.label_color, label_bold=text.label_bold)


# --------------------------------------------------------------------------- derived chart types
# python-pptx writes 29 of the 73 XL_CHART_TYPE values. The rest are the same data in a different
# chart element, so build a supported one and rewrite it. Each conversion has to respect the element
# order in ECMA-376: PowerPoint silently "repairs" a chart whose children are out of sequence.

def _order(parent: Any, names: Any) -> None:
    """Reorder parent's children to match `names`, keeping anything unlisted at the end."""
    listed = []
    for n in names:
        listed += parent.findall(f"{{{C}}}{n}")
    rest = [c for c in parent if c not in listed]
    for c in list(parent):
        parent.remove(c)
    for c in listed + rest:
        parent.append(c)


def _strip(el: Any, *names: Any) -> None:
    for n in names:
        for x in el.findall(f"{{{C}}}{n}"):
            el.remove(x)


def _view3d(cs: Any, rot_x: Any = 15, rot_y: Any = 20, depth: Any = 100, right_angle: Any = True) -> None:
    chart = cs.find(f"{{{C}}}chart")
    if chart is None or chart.find(f"{{{C}}}view3D") is not None:
        return
    v = etree.Element(f"{{{C}}}view3D")
    etree.SubElement(v, f"{{{C}}}rotX").set("val", str(rot_x))
    etree.SubElement(v, f"{{{C}}}rotY").set("val", str(rot_y))
    etree.SubElement(v, f"{{{C}}}depthPercent").set("val", str(depth))
    etree.SubElement(v, f"{{{C}}}rAngAx").set("val", "1" if right_angle else "0")
    # CT_Chart: title?, autoTitleDeleted?, pivotFmts?, view3D?, floor?, …, plotArea — so view3D goes
    # after the first three, not first (C1 review #15: `insert(0)` put it ahead of autoTitleDeleted).
    _insert_after(chart, v, ("title", "autoTitleDeleted", "pivotFmts"))


def _add_ser_ax(cs: Any, node: Any) -> None:
    """A 3-D plot needs a third (series) axis, and the chart element needs its axId."""
    plot = cs.find(f".//{{{C}}}plotArea")
    ids = [int(a.get("val")) for a in cs.findall(f".//{{{C}}}axId")]
    new_id = (max(ids) if ids else 100000000) + 1
    cat_id = node.findall(f"{{{C}}}axId")[0].get("val")

    etree.SubElement(node, f"{{{C}}}axId").set("val", str(new_id))

    ax = etree.SubElement(plot, f"{{{C}}}serAx")
    etree.SubElement(ax, f"{{{C}}}axId").set("val", str(new_id))
    sc = etree.SubElement(ax, f"{{{C}}}scaling")
    etree.SubElement(sc, f"{{{C}}}orientation").set("val", "minMax")
    etree.SubElement(ax, f"{{{C}}}delete").set("val", "1")
    etree.SubElement(ax, f"{{{C}}}axPos").set("val", "b")
    etree.SubElement(ax, f"{{{C}}}crossAx").set("val", cat_id)
    # serAx must follow the value axis in the plot area.
    _order(plot, ["layout", "barChart", "bar3DChart", "lineChart", "line3DChart", "areaChart",
                  "area3DChart", "pieChart", "pie3DChart", "doughnutChart", "ofPieChart",
                  "radarChart", "scatterChart", "bubbleChart", "surfaceChart", "surface3DChart",
                  "stockChart", "catAx", "dateAx", "valAx", "serAx", "dTable", "spPr"])


def _convert(chart: Any, base_kind: Any, conv: Any) -> None:
    """Rewrite a built chart into one of the types python-pptx cannot write."""
    cs = chart._chartSpace
    plot = cs.find(f".//{{{C}}}plotArea")
    src = None
    for tag in ("barChart", "lineChart", "areaChart", "pieChart", "doughnutChart"):
        src = plot.find(f"{{{C}}}{tag}")
        if src is not None:
            break
    if src is None:
        return
    dst = conv["to"]
    src.tag = f"{{{C}}}{dst}"

    # A 3-D, surface or stock plot has no label position in the schema; one left behind by the
    # 2-D base chart is exactly the kind of thing PowerPoint answers with a repair prompt.
    if dst in ("bar3DChart", "line3DChart", "area3DChart", "pie3DChart",
               "surfaceChart", "surface3DChart", "stockChart"):
        for position in src.findall(f".//{{{C}}}dLblPos"):
            position.getparent().remove(position)

    if dst == "bar3DChart":
        _strip(src, "overlap", "serLines")
        for ser in src.findall(f"{{{C}}}ser"):
            _strip(ser, "shape")
        sh = etree.SubElement(src, f"{{{C}}}shape")
        sh.set("val", conv.get("shape", "box"))
        _order(src, ["barDir", "grouping", "varyColors", "ser", "dLbls",
                     "gapWidth", "gapDepth", "shape", "axId"])
        _add_ser_ax(cs, src)
        _view3d(cs)

    elif dst in ("line3DChart", "area3DChart"):
        for ser in src.findall(f"{{{C}}}ser"):
            _strip(ser, "marker", "smooth")
        _strip(src, "marker", "overlap", "gapWidth", "smooth")      # CT_Line3DChart has no marker/smooth
        _order(src, ["grouping", "varyColors", "ser", "dLbls", "dropLines", "gapDepth", "axId"])
        _add_ser_ax(cs, src)
        _view3d(cs)

    elif dst == "pie3DChart":
        _strip(src, "firstSliceAng", "holeSize")
        _order(src, ["varyColors", "ser", "dLbls"])
        _view3d(cs, rot_x=30, rot_y=0, right_angle=False)

    elif dst in ("surfaceChart", "surface3DChart"):
        # CT_SurfaceSer is a bare series: no labels, no per-point formatting, no gaps.
        _strip(src, "gapWidth", "overlap", "dLbls", "varyColors", "grouping", "barDir", "shape")
        for ser in src.findall(f"{{{C}}}ser"):
            _strip(ser, "dLbls", "dPt", "invertIfNegative", "shape", "marker", "smooth")
            _order(ser, ["idx", "order", "tx", "spPr", "cat", "val"])
        wf = etree.Element(f"{{{C}}}wireframe")
        wf.set("val", "1" if conv.get("wireframe") else "0")
        src.insert(0, wf)
        _order(src, ["wireframe", "ser", "bandFmts", "axId"])
        _add_ser_ax(cs, src)
        _view3d(cs, rot_x=15 if dst == "surface3DChart" else 90, rot_y=20,
                right_angle=dst != "surface3DChart")

    elif dst == "ofPieChart":
        _strip(src, "firstSliceAng", "holeSize")
        t = etree.Element(f"{{{C}}}ofPieType")
        t.set("val", conv.get("of", "pie"))
        src.insert(0, t)
        gw = etree.SubElement(src, f"{{{C}}}gapWidth")
        gw.set("val", "100")
        st = etree.SubElement(src, f"{{{C}}}splitType")
        st.set("val", "pos")
        sp = etree.SubElement(src, f"{{{C}}}splitPos")
        sp.set("val", "2")
        sz = etree.SubElement(src, f"{{{C}}}secondPieSize")
        sz.set("val", "75")
        sl = etree.SubElement(src, f"{{{C}}}serLines")
        # CT_ChartLines is `spPr?`: the line lives in c:spPr/a:ln, not directly under serLines.
        etree.SubElement(etree.SubElement(sl, f"{{{C}}}spPr"), f"{{{A}}}ln")
        _order(src, ["ofPieType", "varyColors", "ser", "dLbls", "gapWidth", "splitType",
                     "splitPos", "custSplit", "secondPieSize", "serLines"])

    elif dst == "stockChart":
        _strip(src, "grouping", "marker", "varyColors")
        for ser in src.findall(f"{{{C}}}ser"):
            _strip(ser, "smooth", "dLbls")
            # A stock series draws no line; only the hi-low lines and up-down bars are shown.
            spPr = ser.find(f"{{{A}}}spPr") or ser.find(f"{{{C}}}spPr")
            if spPr is None:
                spPr = etree.SubElement(ser, f"{{{C}}}spPr")
            _strip(spPr, "ln")
            etree.SubElement(etree.SubElement(spPr, f"{{{A}}}ln"), f"{{{A}}}noFill")
            _order(ser, ["idx", "order", "tx", "spPr", "marker", "dPt", "dLbls",
                         "trendline", "errBars", "cat", "val", "smooth"])
        hl = etree.SubElement(src, f"{{{C}}}hiLowLines")
        etree.SubElement(etree.SubElement(hl, f"{{{C}}}spPr"), f"{{{A}}}ln")      # CT_ChartLines
        ud = etree.SubElement(src, f"{{{C}}}upDownBars")
        etree.SubElement(ud, f"{{{C}}}gapWidth").set("val", "150")
        etree.SubElement(ud, f"{{{C}}}upBars")
        etree.SubElement(ud, f"{{{C}}}downBars")
        _order(src, ["ser", "dLbls", "dropLines", "hiLowLines", "upDownBars", "axId"])


def _no_fill(chartSpace: Any) -> None:
    for old in chartSpace.findall(f"{{{C}}}spPr"):
        chartSpace.remove(old)
    spPr = etree.Element(f"{{{C}}}spPr")
    etree.SubElement(spPr, f"{{{A}}}noFill")
    ln = etree.SubElement(spPr, f"{{{A}}}ln")
    etree.SubElement(ln, f"{{{A}}}noFill")
    # spPr follows <c:chart> in the chartSpace schema.
    chart = chartSpace.find(f"{{{C}}}chart")
    chartSpace.insert(list(chartSpace).index(chart) + 1 if chart is not None else 0, spPr)


def _no_fill_plot(chartSpace: Any) -> None:
    plot = chartSpace.find(f".//{{{C}}}plotArea")
    if plot is None:
        return
    for old in plot.findall(f"{{{C}}}spPr"):
        plot.remove(old)
    spPr = etree.SubElement(plot, f"{{{C}}}spPr")
    etree.SubElement(spPr, f"{{{A}}}noFill")
    ln = etree.SubElement(spPr, f"{{{A}}}ln")
    etree.SubElement(ln, f"{{{A}}}noFill")
