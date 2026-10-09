"""Text fit prediction on the emitted deck — what PowerPoint will do, before anyone opens it.

Ported from StageFlow `score.fit_report` (the vendored copy is reference-only), with the four
changes WP5's brief asks for:

1. **Fonts are resolved through `config.FONT_DIRS` with a fontTools name-table scan**, not by
   globbing `%LOCALAPPDATA%` for `Arial*.ttf`. A file name is not a family name: `ARIALN.TTF` is
   Arial Narrow and would otherwise be measured as Arial, and a family whose file is called
   something else entirely would silently degrade to "no metrics, no check". The index and the
   face-matching rule are the emitter's own (`engine.emit.text.face_for_exact`, plan D4), so the
   predictor measures every run with the file the emitter measured it with.
2. **Every slide is iterated**, not just the first.
3. **Lines are predicted per `a:br`-delimited segment** and the total compared with the IR's own
   line count. The emitter locks the browser's line breaks; if PowerPoint would break a line
   somewhere else, this says so before the pixel gate has to.
4. **The server's structural check was retired with the Claude export.** It was ported here for the
   Path B export loop to read back to the model; that loop went with decision #11 (Claude only
   designs), and nothing calls it any more. Its rules survive where they are still used:
   `_check_shape` keeps the zero-width, narrow-box and off-slide checks for the suite's fit row.

Everything is measured in canvas px (master brief §5). A prediction is only a prediction, so the
report carries `details` per text box: a disagreement can be read, not just counted.
"""
from __future__ import annotations

import re
from collections.abc import Sequence
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, cast

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from app.config import engine as config
from app.engine.emit.text import METRIC_ALIASES, Face, face_for_exact, face_width_px
from app.engine.ir import IR, Box, Element
from app.engine.reports import FitIssue, FitReport

#: PowerPoint's own line spacing when a paragraph declares none.
_DEFAULT_LINE_SPACING = 1.2

#: How much a box may be overfull before it is reported, in px (0.02 in — StageFlow's slack).
OVERFLOW_SLACK_PX: float = 0.02 * config.PX_PER_IN

#: How far a line in a frame that does not wrap may be predicted past its box before it is an
#: `overhang`, in px: the same 0.02 in. The emitter sizes every frame to its widest predicted line
#: × `config.TEXT_BOX_WIDTH_SLACK` and measures with the faces this module resolves (plan D4), so
#: on its own output this fires only where the fonts differ from the ones the export measured with,
#: or where `text._place_box` regressed (fidelity WP-D). Widen it only with a measurement.
OVERHANG_SLACK_PX: float = OVERFLOW_SLACK_PX

#: A text box narrower than this cannot hold a word; PowerPoint stacks one character per line.
NARROW_BOX_IN: float = 0.2

#: How far a shape may hang off the canvas before it is called off-slide (0.1 in, as the server's
#: check used). A centred or right-aligned label legitimately overhangs; its *text anchor* does not.
OFF_SLIDE_SLACK_PX: float = 0.1 * config.PX_PER_IN


# ------------------------------------------------------------------------------------- font faces


def resolve_face(family: str, bold: bool, italic: bool) -> Face | None:
    """The face PowerPoint draws for a run naming `family` with `b`/`i` — or None when it is absent.

    One index and one rule for the emitter and the predictor (plan D4): `text.face_for_exact` at
    700 for `b="1"` and 400 otherwise. The private index this replaced filed each file under one
    name ID and preferred the shortest subfamily, so it measured Arial as Arial Narrow, Calibri as
    Calibri Light and found no "Calibri Light" at all — every client run was predicted narrower than
    PowerPoint draws it.

    When the family itself is not installed (the Linux container has no Microsoft fonts), its
    metric-compatible stand-in is used (`text.METRIC_ALIASES`: Calibri -> Carlito), which predicts
    the same widths.
    """
    weight = 700 if bold else 400
    face = face_for_exact(family, weight, italic)
    if face is None and family:
        for alias in METRIC_ALIASES.get(family.strip().strip("'\"").lower(), ()):
            face = face_for_exact(alias, weight, italic)
            if face is not None:
                break
    return face


def text_width_px(
    text: str,
    family: str,
    size_px: float,
    *,
    bold: bool = False,
    italic: bool = False,
    letter_spacing_px: float = 0.0,
) -> float | None:
    """Advance width of `text` in px, or None when the family is not installed on this machine."""
    if not text:
        return 0.0
    face = resolve_face(family or "Arial", bold, italic)
    if face is None:
        return None
    try:
        width = face_width_px(face, text, size_px)
    except Exception:  # noqa: BLE001 — an unusable face is the same as a missing one
        return None
    return width + letter_spacing_px * len(text)


# ------------------------------------------------------------------------------- deck inspection


@dataclass(frozen=True, slots=True)
class _Piece:
    """One measurable fragment of a line: some text in one run's font."""

    text: str
    family: str
    size_px: float
    bold: bool
    italic: bool
    spacing_px: float

    @property
    def width(self) -> float | None:
        return text_width_px(self.text, self.family, self.size_px, bold=self.bold,
                             italic=self.italic, letter_spacing_px=self.spacing_px)

    def scaled(self, factor: float) -> _Piece:
        return _Piece(self.text, self.family, self.size_px * factor, self.bold, self.italic,
                      self.spacing_px * factor)


def _font_from_rpr(rpr: Any, fonts: dict[str, str] | None = None) -> dict[str, Any]:
    """`a:rPr` / `a:defRPr` → the five things measurement needs.

    `+mj-*` / `+mn-*` are theme references: with `fonts` (the deck's own theme fonts,
    `_theme_fonts_of`) they resolve to the major and minor family — what PowerPoint draws — so a
    title is measured in the major face. Without it they leave the caller's default in place.
    """
    out: dict[str, Any] = {}
    size = rpr.get("sz")
    if size:
        out["size_px"] = int(size) / 100.0 / config.PT_PER_PX
    if rpr.get("b") is not None:
        out["bold"] = rpr.get("b") in ("1", "true")
    if rpr.get("i") is not None:
        out["italic"] = rpr.get("i") in ("1", "true")
    spacing = rpr.get("spc")
    if spacing:
        out["spacing_px"] = int(spacing) / 100.0 / config.PT_PER_PX
    latin = rpr.find(qn("a:latin"))
    typeface = latin.get("typeface") if latin is not None else None
    if typeface and not typeface.startswith("+"):
        out["family"] = typeface
    elif typeface and fonts:
        if typeface.startswith("+mj"):
            out["family"] = fonts["major"]
        elif typeface.startswith("+mn"):
            out["family"] = fonts["minor"]
    return out


def _shape_defaults(shape: Any, fonts: dict[str, str]) -> dict[str, Any]:
    """The font a run inherits when it says nothing: the shape's list style, then the minor font."""
    defaults: dict[str, Any] = {
        "family": fonts["minor"],
        "size_px": 18.0 / config.PT_PER_PX,
        "bold": False,
        "italic": False,
        "spacing_px": 0.0,
    }
    if not shape.has_text_frame:
        return defaults
    list_style = shape.text_frame._txBody.find(qn("a:lstStyle"))
    if list_style is not None:
        level = list_style.find(qn("a:lvl1pPr"))
        if level is not None:
            default_rpr = level.find(qn("a:defRPr"))
            if default_rpr is not None:
                defaults.update(_font_from_rpr(default_rpr, fonts))
    return defaults


def _paragraph_defaults(
    paragraph: Any, defaults: dict[str, Any], fonts: dict[str, str] | None = None
) -> dict[str, Any]:
    out = dict(defaults)
    properties = paragraph._p.find(qn("a:pPr"))
    if properties is not None:
        default_rpr = properties.find(qn("a:defRPr"))
        if default_rpr is not None:
            out.update(_font_from_rpr(default_rpr, fonts))
    return out


def _segments(
    paragraph: Any, defaults: dict[str, Any], fonts: dict[str, str] | None = None
) -> list[list[_Piece]]:
    """The paragraph split at every `a:br` — PowerPoint wraps each segment independently.

    `fonts` resolves theme tokens (`_font_from_rpr`); `_check_table` calls without it.
    """
    segments: list[list[_Piece]] = [[]]
    for child in paragraph._p.iterchildren():
        if child.tag == qn("a:br"):
            segments.append([])
        elif child.tag in (qn("a:r"), qn("a:fld")):
            text_element = child.find(qn("a:t"))
            text = text_element.text if text_element is not None and text_element.text else ""
            font = dict(defaults)
            rpr = child.find(qn("a:rPr"))
            if rpr is not None:
                font.update(_font_from_rpr(rpr, fonts))
            segments[-1].append(
                _Piece(text=text, family=str(font["family"]), size_px=float(font["size_px"]),
                       bold=bool(font["bold"]), italic=bool(font["italic"]),
                       spacing_px=float(font["spacing_px"]))
            )
    return segments


#: Where PowerPoint may break a line: a run of white space, except the no-break spaces (UAX #14 class
#: GL: U+00A0, U+2007, U+202F), which bind the words either side of them as a letter would. The
#: extractor writes U+00A0 into the run as the browser drew it (r1b), so a no-break space is part
#: of its word here: measured with it, never a place to break.
_BREAK = re.compile(r"([^\S\u00a0\u2007\u202f]+)")


def _tokens(pieces: Sequence[_Piece]) -> list[tuple[str, list[_Piece]]]:
    """The tokens between break opportunities, each `(gap, fragments)`: `gap` is the white space in
    front of the token and `fragments` its text (a word may span two runs)."""
    tokens: list[tuple[str, list[_Piece]]] = []
    current: list[_Piece] = []
    gap = pending = ""
    for piece in pieces:
        for part in _BREAK.split(piece.text):
            if not part:
                continue
            if _BREAK.fullmatch(part):
                if current:
                    tokens.append((gap, current))
                    current = []
                pending += part
                continue
            if not current:
                gap, pending = pending, ""
            current.append(_Piece(part, piece.family, piece.size_px, piece.bold, piece.italic,
                                  piece.spacing_px))
    if current:
        tokens.append((gap, current))
    return tokens


def _space_width(piece: _Piece, gap: str = " ") -> float:
    """The advance of the white space `gap` (one space by default) in `piece`'s font, `spc` left out.

    A break is usually one space; the extractor also keeps an em space or a thin space as the
    browser drew it (r1b), and that gap is as wide as those characters, not as one space.
    """
    width = text_width_px(gap or " ", piece.family, piece.size_px, bold=piece.bold, italic=piece.italic)
    return width if width is not None else piece.size_px * 0.28 * max(1, len(gap))


def _predict_lines(pieces: Sequence[_Piece], width_px: float, wrap: bool) -> tuple[int, bool]:
    """Greedy word wrap at `width_px` — PowerPoint's own algorithm. Returns (lines, measured)."""
    tokens = _tokens(pieces)
    if not tokens:
        return 1, True
    if not wrap:
        return 1, True

    measured = True
    lines, current = 1, 0.0
    for gap_text, token in tokens:
        width = 0.0
        for fragment in token:
            piece_width = fragment.width
            if piece_width is None:
                measured = False
                piece_width = len(fragment.text) * fragment.size_px * 0.5  # crude, and flagged
            width += piece_width
        if current <= 0.0:
            current = width
            continue
        gap = _space_width(token[0], gap_text)
        if current + gap + width <= width_px:
            current += gap + width
        else:
            lines += 1
            current = width
    return lines, measured


def _segment_width(pieces: Sequence[_Piece]) -> tuple[float, bool, list[str], list[str]]:
    """How wide one `a:br` segment draws on a line of its own: `(px, measured, missing, synthetic)`.

    The sum of the pieces' advances (`spc` is already folded in by `_font_from_rpr`), less the
    segment's trailing whitespace, which hangs past the margin and is never drawn. `measured` is
    False when a family is not installed here, and `missing` names those families. `synthetic`
    names the families a bold piece resolves to a face lighter than 600 in: PowerPoint synthesises
    that bold, and both the emitter and this check measure it with the regular file — exactly the
    blind spot a reader of the detail should be told about.
    """
    trimmed = list(pieces)
    while trimmed:
        last = trimmed[-1]
        kept = last.text.rstrip()
        if kept:
            if kept != last.text:
                trimmed[-1] = _Piece(kept, last.family, last.size_px, last.bold, last.italic,
                                     last.spacing_px)
            break
        trimmed.pop()
    width, measured = 0.0, True
    missing: set[str] = set()
    synthetic: set[str] = set()
    for piece in trimmed:
        if not piece.text:
            continue
        face = resolve_face(piece.family or "Arial", piece.bold, piece.italic)
        piece_width = piece.width if face is not None else None
        if piece_width is None:
            measured = False
            missing.add(piece.family)
            continue
        if piece.bold and face is not None and face.weight < 600:
            synthetic.add(piece.family)
        width += piece_width
    return width, measured, sorted(missing), sorted(synthetic)


def _line_measure(paragraph: Any) -> tuple[str, float, float, float]:
    """`(algn, first line start, later lines' start, right indent)` in px, from the paragraph's `pPr`.

    With a hanging bullet (`indent < 0` and a `buChar`/`buAutoNum`) the text of every line starts
    at `marL`; otherwise the first line starts at `marL + indent`. Only the paragraph's own `pPr` is
    read — the emitter always writes it; a foreign placeholder's inherited indents are not resolved.
    """
    properties = paragraph._p.find(qn("a:pPr"))
    if properties is None:
        return "l", 0.0, 0.0, 0.0

    def px(name: str) -> float:
        value = properties.get(name)
        return int(value) / config.EMU_PER_PX if value else 0.0

    margin_left, margin_right, indent = px("marL"), px("marR"), px("indent")
    bulleted = any(properties.find(qn(tag)) is not None for tag in ("a:buChar", "a:buAutoNum", "a:buBlip"))
    first = margin_left if (bulleted and indent < 0) else max(0.0, margin_left + indent)
    return properties.get("algn") or "l", first, margin_left, margin_right


def _line_height_px(paragraph: Any, pieces: Sequence[_Piece]) -> float:
    """Line height in px: an explicit `lnSpc`, else 1.2 × the largest run (PowerPoint's default).

    `Length` is a subclass of `int` in python-pptx, so the "is it a multiple or a measurement?"
    test has to ask for `.pt` first — an `isinstance(spacing, int)` branch reads 330,200 EMU as a
    line-spacing multiple of 330,200 and every box then overflows by a few million pixels.
    """
    biggest = max((p.size_px for p in pieces), default=18.0 / config.PT_PER_PX)
    spacing = paragraph.line_spacing
    if spacing is None:
        return biggest * _DEFAULT_LINE_SPACING
    if hasattr(spacing, "pt"):
        return float(spacing.pt) / config.PT_PER_PX  # a Length: exact points per line
    return biggest * float(spacing)


def _space_px(length: Any) -> float:
    """Space before/after a paragraph, in px. Same `Length`-is-an-`int` trap as above."""
    if length is None:
        return 0.0
    if hasattr(length, "pt"):
        return float(length.pt) / config.PT_PER_PX
    return float(length) / config.PT_PER_PX


def _autofit_of(shape: Any) -> tuple[str, float]:
    """`("none"|"shape"|"text", font scale)` — PowerPoint grows or shrinks before it overflows."""
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    if body is None:
        return "none", 1.0
    if body.find(qn("a:spAutoFit")) is not None:
        return "shape", 1.0
    normal = body.find(qn("a:normAutofit"))
    if normal is not None:
        scale = normal.get("fontScale")
        return "text", (int(scale) / 100000.0 if scale else 1.0)
    return "none", 1.0


def _insets_px(shape: Any) -> tuple[float, float, float, float]:
    """`(left, top, right, bottom)` internal margins in px, with PowerPoint's defaults when unset."""
    frame = shape.text_frame
    defaults = (0.1 * config.PX_PER_IN, 0.05 * config.PX_PER_IN,
                0.1 * config.PX_PER_IN, 0.05 * config.PX_PER_IN)
    values = (frame.margin_left, frame.margin_top, frame.margin_right, frame.margin_bottom)
    return tuple(  # type: ignore[return-value]
        default if value is None else value / config.EMU_PER_PX
        for value, default in zip(values, defaults, strict=False)
    )


def _alignment_of(shape: Any) -> str:
    paragraphs = shape.text_frame.paragraphs
    alignment = paragraphs[0].alignment if paragraphs else None
    return str(alignment or "").split(".")[-1].split(" ")[0].upper()


# ---------------------------------------------------------------------------------- IR matching


def _ir_line_count(element: Element) -> int:
    return sum(len(paragraph.get("lines") or []) for paragraph in (element.paragraphs or []))


def _iou(a: Box, b: Box) -> float:
    overlap = a.intersect(b).area
    union = a.area + b.area - overlap
    return overlap / union if union > 0 else 0.0


def _match_element(shape_box: Box, name: str, candidates: list[Element], used: set[int]) -> Element | None:
    """Pair an emitted shape with the IR element it came from — by name first, then by overlap.

    The emitter names shapes from `data-name`/the source path, so the name is the reliable join; the
    geometry fallback exists because a Path B deck (Claude's) has no such convention at all.
    """
    for index, element in enumerate(candidates):
        if index not in used and name and name in {element.name, element.id}:
            used.add(index)
            return element
    best_index, best_score = -1, 0.0
    for index, element in enumerate(candidates):
        if index in used:
            continue
        score = _iou(shape_box, element.box)
        if score > best_score:
            best_index, best_score = index, score
    if best_index >= 0 and best_score >= 0.4:
        used.add(best_index)
        return candidates[best_index]
    return None


def _theme_fonts_of(slide: Any, fallback: dict[str, str] | None = None) -> dict[str, str]:
    """`{"major", "minor"}` from the deck's own theme part — what PowerPoint resolves `+mj-lt` to.

    The slide's master's theme (`a:fontScheme/a:majorFont|a:minorFont/a:latin@typeface`), falling
    back per slot to `fallback` (the IR's `fonts.forced`: what measurement forced) and then Arial.
    The deck's theme, not the manifest, because it is what the file will be drawn with: on
    `test-16x9` the theme part says Calibri for the major font while the manifest says Calibri
    Light (a fixture inconsistency reported to `make_test_masters.py`'s owner, plan D4).
    """
    theme: dict[str, str] = {}
    try:
        master_part = slide.slide_layout.slide_master.part
        theme = dict(_fonts_in_theme(master_part.part_related_by(RT.THEME).blob))
    except (KeyError, AttributeError, ValueError):
        theme = {}
    fallback = fallback or {}
    out: dict[str, str] = {}
    for slot in ("major", "minor"):
        out[slot] = theme.get(slot) or str(fallback.get(slot) or "") or "Arial"
    return out


@lru_cache(maxsize=32)
def _fonts_in_theme(blob: bytes) -> tuple[tuple[str, str], ...]:
    """The latin major/minor typefaces of one theme part, parsed once per distinct part."""
    from lxml import etree

    try:
        root = etree.fromstring(blob)
    except etree.XMLSyntaxError:
        return ()
    scheme = root.find(f"{qn('a:themeElements')}/{qn('a:fontScheme')}")
    if scheme is None:
        return ()
    found: list[tuple[str, str]] = []
    for slot, tag in (("major", "a:majorFont"), ("minor", "a:minorFont")):
        latin = scheme.find(f"{qn(tag)}/{qn('a:latin')}")
        typeface = (latin.get("typeface") or "").strip() if latin is not None else ""
        if typeface:
            found.append((slot, typeface))
    return tuple(found)


# ------------------------------------------------------------------------------------------ fit


def fit(pptx: Path, irs: list[IR]) -> FitReport:
    """Predict the text problems in this deck. Never opens a browser and never renders.

    `irs` may be shorter than the deck, or empty: the geometry checks run on every slide either way,
    and the line-count comparison only runs where there is an IR to compare against.
    """
    report = FitReport()
    pptx = Path(pptx)
    if not pptx.exists():
        report.issues.append(
            FitIssue(slide=0, shape="(deck)", kind="off-slide", detail=f"no such file: {pptx}")
        )
        return report

    presentation = Presentation(str(pptx))
    canvas_w = cast(int, presentation.slide_width) / config.EMU_PER_PX
    canvas_h = cast(int, presentation.slide_height) / config.EMU_PER_PX
    forced = next((dict(ir.fonts.forced) for ir in irs if ir.fonts.forced), {})

    for index, slide in enumerate(presentation.slides, start=1):
        ir = irs[index - 1] if index - 1 < len(irs) else None
        fonts = _theme_fonts_of(slide, (ir.fonts.forced if ir is not None else None) or forced)
        candidates = [] if ir is None else [e for e in ir.elements if e.kind == "text"]
        used: set[int] = set()
        for shape in slide.shapes:
            _check_shape(report, index, shape, canvas_w, canvas_h, fonts, candidates, used)
        report.slides_checked += 1

    return report


#: Preset geometries that are a stroke rather than an area, so zero extent across is correct.
_LINE_PRESETS: frozenset[str] = frozenset({
    "line", "straightConnector1", "bentConnector2", "bentConnector3", "bentConnector4",
    "bentConnector5", "curvedConnector2", "curvedConnector3", "curvedConnector4", "curvedConnector5",
})


def _is_line(shape: Any) -> bool:
    """True when this shape is a connector or a `line` preset — something with no area by design."""
    element = getattr(shape, "_element", None)
    if element is None:
        return False
    if element.tag.rpartition("}")[2] == "cxnSp":
        return True
    preset = element.find(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom"
    )
    return preset is not None and preset.get("prst") in _LINE_PRESETS


def _wraps(shape: Any) -> bool:
    """Whether PowerPoint may re-wrap this text box (`wrap="none"` means it may not)."""
    if not shape.has_text_frame:
        return False
    # python-pptx reports None when the attribute is absent, and the OOXML default is to wrap.
    return shape.text_frame.word_wrap is not False


def _check_shape(
    report: FitReport,
    slide_index: int,
    shape: Any,
    canvas_w: float,
    canvas_h: float,
    fonts: dict[str, str],
    candidates: list[Element],
    used: set[int],
) -> None:
    """Every fit rule, applied to one shape (recursing into groups)."""
    if shape.shape_type is not None and str(shape.shape_type).startswith("GROUP"):
        for child in shape.shapes:
            _check_shape(report, slide_index, child, canvas_w, canvas_h, fonts, candidates, used)
        return
    if shape.left is None or shape.top is None:
        return

    box = Box(
        shape.left / config.EMU_PER_PX,
        shape.top / config.EMU_PER_PX,
        (shape.width or 0) / config.EMU_PER_PX,
        (shape.height or 0) / config.EMU_PER_PX,
    )
    has_text = bool(shape.has_text_frame and shape.text_frame.text.strip())
    text = shape.text_frame.text.strip() if has_text else ""
    label = f"{shape.name!r}" + (f" holding {text[:40]!r}" if text else "")

    # -- geometry ---------------------------------------------------------------------------------
    if has_text and box.w <= 0:
        report.issues.append(FitIssue(
            slide=slide_index, shape=label, kind="zero-width",
            detail=f"the text box has zero width, so PowerPoint stacks its {len(text)} characters "
                   f"one per line; give it the box from the measurements",
        ))
        return
    if box.w <= 0 or box.h <= 0:
        # A horizontal rule *is* zero-height and a vertical one *is* zero-width; that is the shape
        # being correct, not a defect. Judged by what the shape is (a connector, or a `line` preset)
        # rather than by its name: the emitter names shapes after the design, so a rule extracted
        # from SVG is called `line#1` and a name test for "Line" misses it 31 times on one deck.
        if has_text or not _is_line(shape):
            report.issues.append(FitIssue(
                slide=slide_index, shape=label, kind="zero-width",
                detail=f"zero {'width' if box.w <= 0 else 'height'} ({box.w:.1f}x{box.h:.1f}px)",
            ))
        return
    # Without word wrap PowerPoint cannot stack characters, however narrow the box is, so the rule
    # only applies to boxes that wrap.
    if has_text and _wraps(shape) and box.w < NARROW_BOX_IN * config.PX_PER_IN and len(text) > 1:
        report.issues.append(FitIssue(
            slide=slide_index, shape=label, kind="narrow-box",
            detail=f"only {box.w / config.PX_PER_IN:.2f}in wide — PowerPoint will stack its text "
                   f"one character per line",
        ))
        return

    off_slide = _check_off_slide(report, slide_index, shape, box, canvas_w, canvas_h, has_text, label)

    if getattr(shape, "has_table", False):
        _check_table(report, slide_index, shape, fonts["minor"], label)
        return
    if not has_text:
        return

    # -- the prediction ----------------------------------------------------------------------------
    frame = shape.text_frame
    autofit, font_scale = _autofit_of(shape)
    margin_l, margin_t, margin_r, margin_b = _insets_px(shape)
    inner_width = max(1.0, box.w - margin_l - margin_r)
    wrap = frame.word_wrap is not False

    defaults = _shape_defaults(shape, fonts)
    needed = 0.0
    predicted = 0
    measured = True                   # the line count
    widths_measured = True            # every segment's width (a wrap-off count measures nothing)
    missing: set[str] = set()
    synthetic: set[str] = set()
    # The segment with the least room on its line — the widest one wherever every line starts at the
    # same indent: (1-based number over all segments, width, headroom, its paragraph's algn, the
    # line's start inside the frame, the paragraph's right indent).
    tightest: tuple[int, float, float, str, float, float] | None = None
    number = 0
    for paragraph in frame.paragraphs:
        segments = _segments(paragraph, _paragraph_defaults(paragraph, defaults, fonts), fonts)
        if font_scale != 1.0:
            segments = [[piece.scaled(font_scale) for piece in segment] for segment in segments]
        pieces = [piece for segment in segments for piece in segment]
        algn, first_start, later_start, margin_right = _line_measure(paragraph)
        lines = 0
        for position, segment in enumerate(segments):
            count, ok = _predict_lines(segment, inner_width, wrap)
            lines += count
            measured = measured and ok
            number += 1
            width, ok, absent, synthesised = _segment_width(segment)
            widths_measured = widths_measured and ok
            missing.update(absent)
            synthetic.update(synthesised)
            start = first_start if position == 0 else later_start
            headroom = inner_width - start - margin_right - width
            if tightest is None or headroom < tightest[2]:
                tightest = (number, width, headroom, algn, start, margin_right)
        predicted += lines
        needed += lines * _line_height_px(paragraph, pieces)
        needed += _space_px(paragraph.space_before) + _space_px(paragraph.space_after)

    available = max(0.0, box.h - margin_t - margin_b)
    detail: dict[str, Any] = {
        "slide": slide_index,
        "shape": shape.name,
        "predictedLines": predicted,
        "neededPx": round(needed, 2),
        "availablePx": round(available, 2),
        "autofit": autofit,
        "measured": measured and widths_measured,
        "wrap": wrap,
        "widestLine": tightest[0] if tightest else 0,
        "widestLinePx": round(tightest[1], 2) if tightest else 0.0,
        "headroomPx": round(tightest[2], 2) if tightest else round(inner_width, 2),
        "missingFonts": sorted(missing),
        "syntheticBold": sorted(synthetic),
    }

    element = _match_element(box, shape.name, candidates, used)
    if element is not None:
        detail["irElement"] = element.id
        detail["irLines"] = _ir_line_count(element)
        if measured and detail["irLines"] and predicted != detail["irLines"]:
            report.issues.append(FitIssue(
                slide=slide_index, shape=label, kind="line-count",
                detail=f"PowerPoint will lay this out in {predicted} line(s) but the browser "
                       f"measured {detail['irLines']}; the design's line breaks will move",
            ))

    if autofit == "none" and needed > available + OVERFLOW_SLACK_PX:
        report.issues.append(FitIssue(
            slide=slide_index, shape=label, kind="overflow",
            detail=f"needs {needed:.0f}px, the box holds {available:.0f}px "
                   f"({needed - available:+.0f}px over)",
        ))
        detail["overflowPx"] = round(needed - available, 2)

    # A frame that does not wrap cannot re-break a line that is too wide: the line overhangs its box
    # in the direction the paragraph is aligned to (fidelity WP-D). Past the canvas edge that is an
    # `off-slide`; otherwise an `overhang`. Judged only on measured widths, never on vertical text
    # (`vert`: its line runs along the box's height), and — one issue per box — not again for a box
    # whose text anchor `_check_off_slide` already put off the slide.
    vertical = (frame._txBody.bodyPr.get("vert") or "horz") != "horz"
    if (tightest is not None and not wrap and widths_measured and not vertical and not off_slide
            and tightest[2] < -OVERHANG_SLACK_PX):
        number, width, headroom, algn, start, margin_right = tightest
        left_edge = box.x + margin_l + start
        right_edge = box.x2 - margin_r - margin_right
        if algn == "ctr":
            line_left = (left_edge + right_edge - width) / 2.0
        elif algn == "r":
            line_left = right_edge - width
        else:                                           # l, just, dist, …: grows to the right
            line_left = left_edge
        past_right = line_left + width - canvas_w
        past_left = -line_left
        rotated = abs(float(getattr(shape, "rotation", 0.0) or 0.0)) > 1e-6
        if not rotated and max(past_right, past_left) > OFF_SLIDE_SLACK_PX:
            edge, over = ("right", past_right) if past_right >= past_left else ("left", past_left)
            report.issues.append(FitIssue(
                slide=slide_index, shape=label, kind="off-slide",
                detail=f"line {number} is predicted {width:.0f}px wide and runs {over:.0f}px past "
                       f"the {edge} edge of the slide",
            ))
        else:
            measure = width + headroom
            report.issues.append(FitIssue(
                slide=slide_index, shape=label, kind="overhang",
                detail=f"line {number} is predicted {width:.0f}px wide, the box is {measure:.0f}px "
                       f"({-headroom:+.0f}px past its edge)",
            ))

    report.details.append(detail)


def _check_off_slide(
    report: FitReport,
    slide_index: int,
    shape: Any,
    box: Box,
    canvas_w: float,
    canvas_h: float,
    has_text: bool,
    label: str,
) -> bool:
    """Off-slide, judged the way the server judged it: by where the *text* lands, not the box.

    A wide box hanging past the edge is normal for a centred or right-aligned label — the text still
    sits on the slide. A plain shape is only reported when none of it is visible. Returns whether it
    reported, so the line-level rule in `_check_shape` does not report the same box twice.
    """
    if has_text:
        alignment = _alignment_of(shape)
        anchor = box.x2 if alignment == "RIGHT" else (box.cx if alignment == "CENTER" else box.x)
        off_horizontally = not (-OFF_SLIDE_SLACK_PX <= anchor <= canvas_w + OFF_SLIDE_SLACK_PX)
        if off_horizontally or box.y2 <= 0 or box.y >= canvas_h:
            report.issues.append(FitIssue(
                slide=slide_index, shape=label, kind="off-slide",
                detail=f"its text lands off the {canvas_w:.0f}x{canvas_h:.0f}px canvas "
                       f"(at {box.x:.0f}, {box.y:.0f})",
            ))
            return True
    elif box.x2 <= 0 or box.y2 <= 0 or box.x >= canvas_w or box.y >= canvas_h:
        report.issues.append(FitIssue(
            slide=slide_index, shape=label, kind="off-slide",
            detail=f"entirely off the {canvas_w:.0f}x{canvas_h:.0f}px canvas "
                   f"(at {box.x:.0f}, {box.y:.0f})",
        ))
        return True
    return False


def _check_table(report: FitReport, slide_index: int, shape: Any, theme_font: str, label: str) -> None:
    """A PowerPoint row height is a *minimum*: a cell that needs two lines pushes the table down.

    The declared height comes from the HTML, where the same text fitted on one line, so this is the
    defect that keeps appearing — and it is predictable without rendering anything. A cell run's
    `+mj-lt` / `+mn-lt` resolves through the deck's own theme fonts (G-1's `fonts=`, plan §16 #11),
    read from the table's slide as `_check_shape` reads them, `theme_font` standing in for both.
    """
    table = shape.table
    columns = [c.width / config.EMU_PER_PX for c in table.columns]
    rows = [r.height / config.EMU_PER_PX for r in table.rows]
    defaults = {"family": theme_font, "size_px": 12.0 / config.PT_PER_PX,
                "bold": False, "italic": False, "spacing_px": 0.0}
    fonts = _theme_fonts_of(getattr(shape.part, "slide", None), {"major": theme_font, "minor": theme_font})
    grown = 0.0
    for row_index in range(len(rows)):
        needed_row, worst = 0.0, ""
        for column_index in range(len(columns)):
            cell = table.cell(row_index, column_index)
            if cell.is_spanned:
                continue
            text = cell.text_frame.text.strip()
            if not text:
                continue
            span = getattr(cell, "span_width", 1) or 1
            cell_width = sum(columns[column_index:column_index + span])
            inner = max(1.0, cell_width - (cell.margin_left + cell.margin_right) / config.EMU_PER_PX)
            wrap = cell.text_frame.word_wrap is not False
            needed_cell = 0.0
            paragraphs = list(cell.text_frame.paragraphs)
            for index, paragraph in enumerate(paragraphs):
                segments = _segments(paragraph, _paragraph_defaults(paragraph, defaults, fonts), fonts)
                pieces = [piece for segment in segments for piece in segment]
                if not pieces:
                    continue
                # The paragraph's `marL` narrows every segment (text in a grid column, under a
                # margin, beside a bullet) and its first-line `indent` the first one unless it is
                # bulleted (the bullet hangs in the indent; the text starts at `marL`). PowerPoint
                # drops a frame's first `spcBef` and its last `spcAft` (`spcFirstLastPara` defaults false).
                properties = paragraph._p.find(qn("a:pPr"))
                pPr = properties.attrib if properties is not None else {}
                indent, left = (int(pPr.get(name) or 0) / config.EMU_PER_PX for name in ("indent", "marL"))
                bulleted = properties is not None and (properties.find(qn("a:buChar")) is not None
                                                       or properties.find(qn("a:buAutoNum")) is not None)
                first_indent = 0.0 if bulleted else indent
                lines = sum(_predict_lines(segment, max(1.0, inner - max(0.0, left)
                                                        - (first_indent if number == 0 else 0.0)), wrap)[0]
                            for number, segment in enumerate(segments))
                needed_cell += lines * _line_height_px(paragraph, pieces)
                if index > 0:
                    needed_cell += _space_px(paragraph.space_before)
                if index < len(paragraphs) - 1:
                    needed_cell += _space_px(paragraph.space_after)
            needed_cell += (cell.margin_top + cell.margin_bottom) / config.EMU_PER_PX
            if needed_cell > needed_row:
                needed_row, worst = needed_cell, text[:40]
        if needed_row > rows[row_index] + 1.0:
            grown += needed_row - rows[row_index]
            report.issues.append(FitIssue(
                slide=slide_index, shape=f"{label} row {row_index + 1}", kind="overflow",
                detail=f"row {row_index + 1} needs {needed_row:.0f}px but is declared "
                       f"{rows[row_index]:.0f}px ({worst!r})",
            ))
    if grown > OVERFLOW_SLACK_PX:
        report.details.append({"slide": slide_index, "shape": shape.name,
                               "tableGrowthPx": round(grown, 2)})

