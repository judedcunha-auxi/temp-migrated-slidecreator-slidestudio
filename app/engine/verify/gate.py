"""The pixel gate: visible differences between the browser composite and PptxRender's render.

`gate(pptx, references, masks, out_dir, renderer=...)` runs `renderer.verify` (blur 2 / ΔE 8 / ≥ 60 px² /
2 px slack — `config.GATE_*`) and reports the defect area per slide, twice:

* `GateSlide.defect_area` — everything, masks included. The honest total.
* `GateSlide.non_chart_area` — the number the acceptance gate reads, with the masked regions taken
  out of both images.

**How a mask is applied to the candidate.** The masks have to be painted on *both* pictures or the
comparison reports the mask itself as a defect, and the candidate is not a picture we own: the CLI
renders it from the deck. So the masked run scores a **copy of the deck with a flat rectangle added
over each mask box** against references with the same rectangles painted in the same colour. One
canvas pixel is exactly 9525 EMU (master brief §5) and the renderer draws at the reference's width,
so the two rectangles land on the same pixels.

Two masks exist, for two different reasons (master brief §9):

* **chart frames** — the reference draws a chart with `chart-preview.js`, the candidate with
  PptxRender's chart painter. They will never agree pixel for pixel, which is why charts are
  verified structurally instead (`engine.verify.coverage`).
* **`sldNum` / `dt` / `ftr` placeholder boxes** — the layout PNG under the reference shows the
  layout's static page number; a real slide shows its own.

`masks` is either one list of boxes applied to every slide, or one list per slide (index 0 = slide
1). The per-slide form is what the suite uses: masking slide 2's chart frame on slide 1 as well
would hide real defects there.
"""
from __future__ import annotations

import math
import shutil
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from app.config import engine as config
from app.engine.renderer import Renderer
from app.engine.reports import Box, GateReport, GateSlide

#: The colour both sides are painted with. Flat, saturated and nowhere in a corporate palette, so a
#: mask that somehow fails to land on one side shows up as a defect instead of blending in.
MASK_RGB: tuple[int, int, int] = (255, 0, 255)

#: Grown by this many px on every side before painting. A chart frame lands on a half pixel as often
#: as not (294.4999 px here), and the renderer then anti-aliases the rectangle's edge while PIL draws
#: it hard — which the gate reads as a 600 × 4 px defect strip along the mask boundary.
MASK_INFLATE_PX: int = 1


def _snap(box: Box) -> Box:
    """A mask box on whole pixels, grown by `MASK_INFLATE_PX`, so both sides paint the same pixels."""
    x0 = math.floor(box.x) - MASK_INFLATE_PX
    y0 = math.floor(box.y) - MASK_INFLATE_PX
    x1 = math.ceil(box.x + box.w) + MASK_INFLATE_PX
    y1 = math.ceil(box.y + box.h) + MASK_INFLATE_PX
    return Box(float(x0), float(y0), float(x1 - x0), float(y1 - y0))


def _per_slide_masks(masks: Sequence[Any] | None, slides: int) -> list[list[Box]]:
    """Normalise `masks` to one list per slide, accepting both shapes of the argument."""
    if not masks:
        return [[] for _ in range(slides)]
    if all(isinstance(entry, (list, tuple)) for entry in masks):
        per_slide = [list(entry) for entry in masks]
        per_slide += [[] for _ in range(max(0, slides - len(per_slide)))]
        return per_slide[:slides]
    return [list(masks) for _ in range(slides)]  # type: ignore[arg-type]


def _paint(image_path: Path, boxes: Sequence[Box], out_path: Path) -> Path:
    """Copy an image with the mask boxes filled in. Returns the new path (a copy even with no boxes)."""
    from PIL import Image, ImageDraw

    with Image.open(image_path) as image:
        painted = image.convert("RGB")
        if boxes:
            draw = ImageDraw.Draw(painted)
            for box in (_snap(b) for b in boxes):
                # PIL's rectangle is inclusive of both corners, so the far edge is x2 - 1.
                draw.rectangle(
                    [int(box.x), int(box.y), int(box.x + box.w) - 1, int(box.y + box.h) - 1],
                    fill=MASK_RGB,
                )
        out_path.parent.mkdir(parents=True, exist_ok=True)
        painted.save(out_path)
    return out_path


def masked_deck(pptx: Path, per_slide: Sequence[Sequence[Box]], out_path: Path) -> Path:
    """A copy of the deck with a flat rectangle over every mask box — the candidate side of a mask."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(pptx, out_path)
    presentation = Presentation(str(out_path))
    for index, slide in enumerate(presentation.slides):
        for box in (_snap(b) for b in (per_slide[index] if index < len(per_slide) else [])):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Emu(round(box.x * config.EMU_PER_PX)),
                Emu(round(box.y * config.EMU_PER_PX)),
                Emu(round(box.w * config.EMU_PER_PX)),
                Emu(round(box.h * config.EMU_PER_PX)),
            )
            shape.name = "engine-gate-mask"
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*MASK_RGB)  # type: ignore[no-untyped-call]
            shape.line.fill.background()
            shape.shadow.inherit = False
            # `shadow.inherit = False` writes an empty `<a:effectLst/>`, which PptxRender still
            # overrides from the autoshape's `<p:style><a:effectRef>` — the theme's soft shadow then
            # bleeds a grey band below and right of every mask and the gate calls it a defect
            # (600 × 4 px on a client chart frame). Dropping the style reference removes it.
            style = shape._element.find(qn("p:style"))
            if style is not None:
                shape._element.remove(style)
            if shape.has_text_frame:                      # an autoshape brings an empty text frame
                shape.text_frame.word_wrap = False
    presentation.save(str(out_path))
    return out_path


def compose_compare(reference: Path, render: Path, annotated: Path | None, out_path: Path) -> Path:
    """`compare-NN.png`: the browser reference, PptxRender's render and the annotated diff, side by side."""
    from PIL import Image, ImageDraw

    panels: list[tuple[str, Path]] = [("reference (browser)", reference), ("render (PptxRender)", render)]
    if annotated is not None and Path(annotated).exists():
        panels.append(("defects", Path(annotated)))

    images = [Image.open(path).convert("RGB") for _, path in panels]
    try:
        gutter, header = 8, 16
        width = sum(image.width for image in images) + gutter * (len(images) - 1)
        height = max(image.height for image in images) + header
        sheet = Image.new("RGB", (width, height), (255, 255, 255))
        draw = ImageDraw.Draw(sheet)
        x = 0
        for (label, _), image in zip(panels, images, strict=False):
            sheet.paste(image, (x, header))
            draw.text((x + 4, 3), label, fill=(60, 60, 70))
            x += image.width + gutter
        out_path.parent.mkdir(parents=True, exist_ok=True)
        sheet.save(out_path)
    finally:
        for image in images:
            image.close()
    return out_path


def gate(
    pptx: Path, references: list[Path], masks: Sequence[Any], out_dir: Path, *, renderer: Renderer
) -> GateReport:
    """Score `pptx` against the browser composites, with and without the masks.

    `renderer` must serve `render` and `verify` (a local PptxRender build; the deployed service
    serves `render-layouts` only, see `app.engine.renderer`).

    Raises `renderer.RendererError` when the renderer checked a different number of slides than
    there are references — the CLI silently skips a slide whose render size differs from its
    reference, and a "clean" run over half the deck is worse than a failure.
    """
    pptx, out_dir = Path(pptx), Path(out_dir)
    references = [Path(r) for r in references]
    out_dir.mkdir(parents=True, exist_ok=True)
    if not references:
        return GateReport(slides=[], slides_checked=0, source="none",
                          warnings=["no reference images: the gate did not run"])

    from PIL import Image

    with Image.open(references[0]) as probe:
        width = probe.width

    # The candidate, as pixels: the CLI renders it again inside `verify`, but the compare sheet and
    # any later triage need the render itself, and it costs ~0.4 s a slide.
    renders = renderer.render(pptx, width, out_dir / "render")

    raw = renderer.verify(pptx, references, out_dir / "gate")
    per_slide = _per_slide_masks(masks, len(references))
    masked = None
    if any(per_slide):
        masked_references = [
            _paint(reference, per_slide[index], out_dir / "masked" / f"reference-{index + 1:02d}.png")
            for index, reference in enumerate(references)
        ]
        deck = masked_deck(pptx, per_slide, out_dir / "masked" / f"{pptx.stem}-masked.pptx")
        masked = renderer.verify(deck, masked_references, out_dir / "gate-masked")

    scored = masked or raw
    slides: list[GateSlide] = []
    for index in range(1, len(references) + 1):
        raw_slide = next((s for s in raw.slides if s.index == index), GateSlide(index=index))
        scored_slide = next((s for s in scored.slides if s.index == index), GateSlide(index=index))
        diff = scored_slide.diff_png
        compare = None
        if index - 1 < len(renders):
            compare = compose_compare(
                references[index - 1], renders[index - 1], diff, out_dir / f"compare-{index:02d}.png"
            )
        slides.append(
            GateSlide(
                index=index,
                defect_area=raw_slide.defect_area,
                clean=scored_slide.defect_area == 0,
                components=list(scored_slide.components),
                diff_png=compare or diff,
                non_chart_area=scored_slide.defect_area,
            )
        )

    report = GateReport(
        slides=slides,
        slides_checked=scored.slides_checked,
        total_area=raw.total_area,
        settings={
            **raw.settings,
            "masks": sum(len(boxes) for boxes in per_slide),
            "maskedTotalArea": sum(s.non_chart_area or 0 for s in slides),
            "width": width,
        },
        # Whatever transport the two passes actually used — hardcoding this made a sidecar-backed
        # report claim the CLI produced it, and a number nobody can trace to a renderer is not
        # evidence (engine/renderer.py's own words).
        source=scored.source or raw.source,
        warnings=sorted(set(raw.warnings) | set(scored.warnings)),
    )
    return report
