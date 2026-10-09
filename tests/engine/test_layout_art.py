"""An unused placeholder whose layout paints something keeps that paint as a plain shape.

The render check's cover layout draws its angled colour panel as a picture placeholder's own fill; deleting the
unused placeholder deleted the panel, and the white cover title landed on white (render check 2026-09-29).
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

from app.engine.emit.template import prune_placeholders

MASTER = Path(__file__).resolve().parent / "fixtures" / "masters" / "test-16x9.pptx"


def _slide_with_painted_placeholder(paint: bool):
    prs = Presentation(str(MASTER))
    layout = next(lay for lay in prs.slide_layouts if len(lay.placeholders) >= 2)
    target = [p for p in layout.placeholders if p.placeholder_format.idx != 0][0]
    sp_pr = target._element.find(qn("p:spPr"))
    for old in sp_pr.findall(qn("a:solidFill")):
        sp_pr.remove(old)
    if paint:
        fill = sp_pr.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:schemeClr"), {"val": "accent1"})
        fill.append(clr)
        geometry = sp_pr.find(qn("a:prstGeom"))
        (geometry.addnext(fill) if geometry is not None else sp_pr.append(fill))
    slide = prs.slides.add_slide(layout)
    return slide, target.placeholder_format.idx


def test_a_painting_placeholder_becomes_a_plain_shape_with_its_paint():
    slide, idx = _slide_with_painted_placeholder(paint=True)
    removed = prune_placeholders(slide, [])
    assert list(slide.placeholders) == []
    art = [s for s in slide.shapes if s.name.endswith("(layout art)")]
    assert len(art) == 1
    sp = art[0]._element
    assert sp.find(qn("p:nvSpPr")).find(qn("p:nvPr")).find(qn("p:ph")) is None
    assert sp.find(qn("p:txBody")) is None
    assert sp.find(qn("p:spPr")).find(qn("a:solidFill")) is not None
    assert sp.find(qn("p:spPr")).find(qn("a:xfrm")) is not None
    assert any("layout artwork kept" in r and f"#{idx}" in r for r in removed)
    ids = [el.get("id") for el in slide.shapes._spTree.iter(qn("p:cNvPr"))]
    assert len(ids) == len(set(ids))


def test_a_placeholder_that_paints_nothing_is_simply_deleted():
    slide, _ = _slide_with_painted_placeholder(paint=False)
    prune_placeholders(slide, [])
    assert list(slide.placeholders) == []
    assert not [s for s in slide.shapes if s.name.endswith("(layout art)")]
