"""WP-B — side-by-side text: the `rows` torture family, the promoted probe families `flexgrid` (p01),
`layout-misc` (p09) and `text-misc` (p10), and the text row's checks.

`docs/archive/engine/fidelity/11-WPB-side-by-side.md`. The browser lays text out in two dimensions; these
tests hold the extractor to it: a flex or grid row is one element per item, a flex column that
stacks is one element, a chip with its own geometry is its own element over its own shape, text
beside an icon or a swatch starts where the browser put it, and alignment that came from flexbox or
grid survives as `align`/`anchor`. Every number is re-derivable from the fixture's CSS.

The deck-side checks run on the emitted decks through `engine.verify.text_deck.text_report`, the
text row's one dict: G-1's `collisions` and `exportOverlaps`, and WP-B's `rows` (no two paragraphs of
one element side by side), `positions` (V3: every paragraph's first line starts where the browser put
it, from the emitter's measured `marL`/`marR`/`indent`, §3.6) and `seams` (no glued-runs diagnostic).
The B1 critique's cases (fidelity-reports/b1-review.md) are pinned here too, each by its finding's number.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from app.config import engine as config
from app.engine.classify.charts import recognise_charts
from app.engine.classify.placeholders import map_placeholders
from app.engine.emit import pptx as emitter
from app.engine.extract.html import measuring_page
from app.engine.ir import IR, Box, Canvas, Element, Slide, assign_ids_and_z
from app.engine.manifest import Manifest
from app.engine.reports import EmitOptions
from app.engine.verify.fit import fit
from app.engine.verify.fixtures import context_for
from app.engine.verify.text_deck import text_report, text_rows
from tests.engine.helpers import extract_html

ROWS = config.TORTURE_FIXTURE / "rows.html"
#: The fidelity probes p01, p09 and p10, promoted to torture families by WP-G (G-2, plan 16 #16).
PROBES = config.TORTURE_FIXTURE
PROBE_STEMS = ("flexgrid", "layout-misc", "text-misc")


# ------------------------------------------------------------------------------------- fixtures


def _extract(html: Path, page) -> IR:
    context = context_for(html)
    return extract_html(html, context.manifest, context.layout_id, context.assets_dir,
                        slide_id=html.stem, page=page)


@pytest.fixture(scope="module")
def irs() -> dict[str, IR]:
    """`rows.html` and the three promoted probe families, extracted once on one page."""
    for html in [ROWS] + [PROBES / f"{stem}.html" for stem in PROBE_STEMS]:
        assert html.exists(), f"{html} is missing"
    with measuring_page(Canvas(1280, 720)) as page:
        found = {"rows": _extract(ROWS, page)}
        for stem in PROBE_STEMS:
            found[stem] = _extract(PROBES / f"{stem}.html", page)
    return found


@pytest.fixture(scope="module")
def decks(irs: dict[str, IR], tmp_path_factory) -> dict[str, tuple[Path, IR]]:
    """Each IR classified and emitted on its own master, as the pipeline does."""
    directory = tmp_path_factory.mktemp("rows")
    built: dict[str, tuple[Path, IR]] = {}
    for stem, ir in irs.items():
        html = ROWS if stem == "rows" else PROBES / f"{stem}.html"
        context = context_for(html)
        classified = map_placeholders(recognise_charts(ir), context.layout)
        deck = directory / f"{stem}.pptx"
        emitter.emit([classified], context.manifest, context.master, deck, EmitOptions())
        built[stem] = (deck, classified)
    return built


# -------------------------------------------------------------------------------------- helpers


def _texts(ir: IR, name: str | None = None) -> list[Element]:
    return [e for e in ir.elements if e.kind == "text" and (name is None or e.name == name)]


def _paragraphs(element: Element) -> list[str]:
    """Each paragraph's text, its lines joined by a space."""
    return [" ".join("".join(r["text"] for r in line["runs"]) for line in p["lines"])
            for p in element.paragraphs or []]


def _text(element: Element) -> str:
    return " | ".join(_paragraphs(element))


def _one(ir: IR, text: str, *, name: str | None = None) -> Element:
    found = [e for e in _texts(ir, name) if _text(e) == text]
    assert len(found) == 1, f"{text!r}: {len(found)} elements; have {[_text(e) for e in _texts(ir)]}"
    return found[0]


def _shapes_in(ir: IR, band: tuple[float, float], *, width: tuple[float, float] | None = None) -> list[Element]:
    """Shapes whose top lies in the vertical band `[y0, y1)`, optionally of a width in `[w0, w1]`."""
    shapes = [e for e in ir.elements if e.kind == "shape" and band[0] <= e.box.y < band[1]]
    if width is not None:
        shapes = [s for s in shapes if width[0] <= s.box.w <= width[1]]
    return sorted(shapes, key=lambda s: s.box.x)


def _first_line(element: Element, paragraph: int = 0, line: int = 0) -> dict:
    return element.paragraphs[paragraph]["lines"][line]["box"]


def _glued(ir: IR, element: Element) -> list[str]:
    return [d.message for d in ir.diagnostics if d.elementId == element.id and d.message.startswith("glued runs:")]


# ----------------------------------------------------------------------------------- the family


def test_rows_meets_its_expectation(irs: dict[str, IR]):
    """The designed counts (brief §4: 50 text / 30 shape / 1 image), measured, not copied."""
    ir = irs["rows"]
    expect = json.loads(ROWS.with_suffix(".expect.json").read_text(encoding="utf-8"))
    assert ir.validate() == []
    assert ir.counts() == expect["kinds"]
    for wanted in expect["expectedDiagnostics"]:
        assert any(d.level == wanted["level"] and wanted["match"] in d.message for d in ir.diagnostics), wanted


def test_label_value_rows_split_with_the_value_right_of_its_label(irs: dict[str, IR]):
    """R1: flex items side by side are not stacked, so each is its own element."""
    ir = irs["rows"]
    for label, value in (("Revenue 2023", "$15.2B"), ("EBITDA margin", "21.4%")):
        left, right = _one(ir, label), _one(ir, value)
        assert right.box.x >= left.box.x2, f"{value} starts inside {label}"
        assert abs(_first_line(left)["y"] - _first_line(right)["y"]) < 0.5
    loose = sorted(_texts(ir, "kv-loose"), key=lambda e: e.box.x)
    assert [_text(e) for e in loose] == ["Headcount", "5.1M"], "a block beside loose text splits too"
    assert loose[1].box.x >= loose[0].box.x2
    assert loose[1].box.x2 == pytest.approx(420.0, abs=0.5), "the loose run hugs its own line at the right"


def test_a_grid_is_one_element_per_cell(irs: dict[str, IR]):
    cells = _texts(irs["rows"], "dl")
    assert len(cells) == 8
    assert all(len(cell.paragraphs) == 1 and len(cell.paragraphs[0]["lines"]) == 1 for cell in cells)


def test_a_centring_flex_column_is_one_element_of_centred_paragraphs(irs: dict[str, IR]):
    """R4's tie-break: `.n` fills the union box, so the column's `align-items: center` decides."""
    kpi = _one(irs["rows"], "$3.8B | Deal value", name="kpi")
    assert [p["align"] for p in kpi.paragraphs] == ["center", "center"]
    assert kpi.anchor == "top", "a column of auto-height items has no vertical slack to derive from"


def test_a_badge_is_centred_both_ways_and_its_text_stacks(irs: dict[str, IR]):
    ir = irs["rows"]
    badge = _one(ir, "1", name="badge")
    assert (badge.paragraphs[0]["align"], badge.anchor) == ("center", "middle")
    txt = _one(ir, "Key takeaway | Concentrate capital on two platforms.")
    assert len(txt.paragraphs) == 2, "the <b> block and the loose text after it stack under each other"
    assert txt.box.x >= badge.box.x2


@pytest.mark.parametrize("label", ["On track", "Applied AI"])
def test_a_pill_s_text_starts_after_its_dot(irs: dict[str, IR], label: str):
    """R2: the run shares the pill with the dot, so it hugs its own line instead of the pill's box."""
    ir = irs["rows"]
    text = _one(ir, label)
    dots = [s for s in _shapes_in(ir, (160, 190), width=(7.5, 8.5)) if s.box.x2 <= text.box.x]
    assert dots, f"no dot left of {label!r}"
    assert text.box.x >= dots[-1].box.x2 + 5, "the pill's 6 px gap separates the dot from the text"


def test_a_legend_label_starts_after_its_swatch(irs: dict[str, IR]):
    ir = irs["rows"]
    swatches = _shapes_in(ir, (200, 215), width=(9.5, 10.5))
    labels = sorted((_one(ir, t) for t in ("UAE", "KSA", "Qatar")), key=lambda e: e.box.x)
    assert len(swatches) == 3
    for swatch, label in zip(swatches, labels, strict=False):
        assert label.box.x >= swatch.box.x2 + 5


def test_display_contents_cells_are_walked_and_painted(irs: dict[str, IR]):
    """R5: the wrapper has no box; its two cells are flex items like the other two."""
    ir = irs["rows"]
    cells = sorted(_texts(ir, "contents"), key=lambda e: e.box.x)
    assert [_text(c) for c in cells] == ["Sources", "Uses", "Valuation", "Comps"]
    fills = _shapes_in(ir, (230, 240))
    assert len(fills) == 4 and all((f.fill or {}).get("color") == "E8F1FB" for f in fills)
    for cell, fill in zip(cells, fills, strict=False):
        assert fill.box.x <= cell.box.x and cell.box.x2 <= fill.box.x2


@pytest.mark.parametrize("name", ["centred-chip", "centred-circle"])
def test_flex_and_grid_centred_text_is_center_and_middle(irs: dict[str, IR], name: str):
    [element] = _texts(irs["rows"], name)
    assert (element.paragraphs[0]["align"], element.anchor) == ("center", "middle")


def test_text_around_an_inline_image_hugs_each_side(irs: dict[str, IR]):
    ir = irs["rows"]
    [image] = [e for e in ir.elements if e.kind == "image"]
    before, after = sorted(_texts(ir, "icon-inline"), key=lambda e: e.box.x)
    assert (_text(before), _text(after)) == ("Rated", "by analysts")
    assert before.box.x2 <= image.box.x + 0.5 and after.box.x >= image.box.x2 - 0.5


def test_a_right_aligned_flex_column_is_right_on_both_paragraphs(irs: dict[str, IR]):
    [column] = _texts(irs["rows"], "stacked-right")
    assert [p["align"] for p in column.paragraphs] == ["right", "right"]
    first, second = _first_line(column, 0), _first_line(column, 1)
    assert first["x"] + first["w"] == pytest.approx(second["x"] + second["w"], abs=0.5)


def test_a_swatch_first_wrapping_item_is_split_after_its_first_line(irs: dict[str, IR]):
    """R8: one bulleted paragraph would put the continuation where the text after the swatch starts."""
    ir = irs["rows"]
    [element] = _texts(ir, "swatch-li")
    assert [bool(p.get("bullet")) for p in element.paragraphs] == [True, True, False]
    swatches = _shapes_in(ir, (60, 110), width=(9.5, 10.5))
    item_left = swatches[0].box.x
    assert _first_line(element, 2)["x"] == pytest.approx(item_left, abs=0.5), "the continuation is at the item's edge"
    assert _first_line(element, 1)["x"] > item_left + 10, "the first line starts after the swatch"
    assert any("split after its first line" in d.message for d in ir.diagnostics if d.elementId == element.id)


def test_text_after_a_swatch_in_a_sentence_stays_left_aligned(irs: dict[str, IR]):
    """The swatch is part of the line: the shrink-wrapped block ends where the text does, but the
    text is not right-aligned for that."""
    [element] = _texts(irs["rows"], "swatch-inline")
    assert element.paragraphs[0]["align"] == "left"
    assert _first_line(element)["x"] > element.box.x + 10


def test_a_glued_chip_leaves_its_heading(irs: dict[str, IR]):
    """R3 (a): the chip's margin glues it to 'compound'; it becomes its own text over its own shape."""
    ir = irs["rows"]
    heading = _one(ir, "Size and spend compound", name="chip-tail")
    chip = _one(ir, "ILLUSTRATIVE")
    assert chip.box.x >= heading.box.x2
    shapes = [s for s in ir.elements if s.kind == "shape" and (s.fill or {}).get("color") == "FFE600"
              and s.box.y < 300]
    assert len(shapes) == 1, "painted once, by the element walk"
    assert shapes[0].box.x <= chip.box.x and chip.box.x2 <= shapes[0].box.x2
    assert not _glued(ir, heading) and not _glued(ir, chip)


def test_a_spaced_mid_line_chip_leaves_its_paragraph_and_its_part_line_is_a_paragraph(irs: dict[str, IR]):
    """R3 whatever the spaces: a run cannot carry the chip's padding. The text after it is one
    element of two paragraphs: the part-line that resumes after the chip, and the whole line under it
    at the paragraph's left edge (B1 critique #2 — a part-line does not share the paragraph's
    alignment geometry, so it is placed by margins of its own, not by a first-line indent). The
    separating spaces stay in the gaps: each side's text starts at its first glyph (critique #1).
    The chip paints and reads between the two texts (critique #3)."""
    ir = irs["rows"]
    texts = _texts(ir, "chip-mid")
    before = next(e for e in texts if _text(e) == "The mid-line chip")
    after = next(e for e in texts if e is not before)
    chip = _one(ir, "FORECAST")
    shapes = [s for s in _shapes_in(ir, (440, 460)) if s.box.x <= chip.box.x and chip.box.x2 <= s.box.x2]
    assert len(shapes) == 1
    shape = shapes[0]
    assert [len(p["lines"]) for p in after.paragraphs] == [1, 1]
    first, second = _first_line(after, 0, 0), _first_line(after, 1, 0)
    assert first["x"] >= shape.box.x2 + 2.5, "a 15 px Calibri space (~3.4 px) lies between the chip and 'sits'"
    assert second["x"] == pytest.approx(_first_line(before)["x"], abs=0.5)
    edge = _first_line(before)
    assert edge["x"] + edge["w"] <= shape.box.x - 2.5, "'chip' ends a space before the chip's margin"
    order = [e.id for e in ir.elements]
    assert order.index(before.id) < order.index(shape.id) < order.index(chip.id) < order.index(after.id)


def test_a_centred_part_line_after_a_chip_is_its_own_paragraph(irs: dict[str, IR]):
    """B1 critique #2: in a centred paragraph opening with a chip, the text after the chip is one
    element of two centred paragraphs — the part-line after the chip and the whole line under it —
    because the two do not share a centre (one alignment could not place both)."""
    ir = irs["rows"]
    [element] = _texts(ir, "chip-centre-wrap")
    chip = _one(ir, "NEW")
    assert [p["align"] for p in element.paragraphs] == ["center", "center"]
    first, second = _first_line(element, 0), _first_line(element, 1)
    assert first["x"] >= chip.box.x2
    centres = [line["x"] + line["w"] / 2 for line in (first, second)]
    assert abs(centres[0] - centres[1]) > 5, "the part-line is not centred on the paragraph"


def test_a_right_aligned_part_line_before_a_chip_is_its_own_paragraph(irs: dict[str, IR]):
    """B1 critique #2: a right-aligned paragraph ending with a chip. Its last line ends at the chip,
    not at the paragraph's edge, so it is a right-aligned paragraph of its own."""
    ir = irs["rows"]
    [element] = _texts(ir, "chip-right-end")
    chip = _one(ir, "END")
    assert [p["align"] for p in element.paragraphs] == ["right", "right"]
    whole, part = _first_line(element, 0), _first_line(element, 1)
    assert whole["x"] + whole["w"] == pytest.approx(element.box.x2, abs=0.5)
    assert part["x"] + part["w"] <= chip.box.x - 5


def test_a_standfirst_pulled_into_the_heading_is_its_own_element(irs: dict[str, IR]):
    """B1 critique #8: the emitter cannot write a negative gap, so an overlap beyond STACK_TOL (15 %
    of the shorter line) splits rather than merges and pushes the standfirst down."""
    ir = irs["rows"]
    heading, standfirst = _one(ir, "Tight heading"), _one(ir, "A standfirst pulled up under it")
    assert standfirst.box.y < heading.box.y2, "the fixture must overlap the two line boxes"


def test_a_mark_around_a_chip_paints_under_it_and_the_chip_reads_between_the_texts(irs: dict[str, IR]):
    """B1 critique #3: the chip is walked after the text before it, so the mark's fill (painted with
    that text) lies under the chip and the three texts read in the browser's order."""
    ir = irs["rows"]
    before, after = _one(ir, "Plain highlight with"), _one(ir, "inside after")
    chip = _one(ir, "CHIP")
    mark = [s for s in _shapes_in(ir, (600, 640)) if (s.fill or {}).get("color") == "FFE9A8"]
    assert len(mark) == 1 and mark[0].box.x <= chip.box.x and chip.box.x2 <= mark[0].box.x2
    chip_fill = [s for s in ir.elements if s.kind == "shape" and s.box.x <= chip.box.x
                 and chip.box.x2 <= s.box.x2 and (s.fill or {}).get("color") == "FFE600"]
    assert len(chip_fill) == 1
    order = [e.id for e in ir.elements]
    assert (order.index(mark[0].id) < order.index(before.id) < order.index(chip_fill[0].id)
            < order.index(chip.id) < order.index(after.id))


def test_an_inline_flex_that_lays_out_two_items_is_a_layout(irs: dict[str, IR]):
    """B1 critique #6: the gap between '6.2' and '$B' is inside the inline-flex box's text union, so
    the width test cannot see it; the box lays two items with text out, so it is walked as a layout."""
    ir = irs["rows"]
    row = sorted((e for e in _texts(ir) if 640 < e.box.y < 660), key=lambda e: e.box.x)
    assert [_text(e) for e in row] == ["Total", "6.2", "$B", "in 2030"]
    total, number, unit, tail = row
    assert total.box.x2 < number.box.x < number.box.x2 < unit.box.x < unit.box.x2 < tail.box.x
    assert unit.box.x - number.box.x2 == pytest.approx(8, abs=0.5)
    for element in (total, number, unit, tail):
        assert not _glued(ir, element)


def test_flex_wrapped_tags_are_a_row_not_a_stack(irs: dict[str, IR]):
    """B1 critique #9: 'Energy transition' wrapped onto the second flex line under 'AI'. Items of a
    flex row never stack, so it is its own element."""
    ir = irs["rows"]
    tags = _texts(ir, "wrap-tags")
    assert sorted(_text(e) for e in tags) == ["AI", "Energy transition", "Sovereign cloud"]


def test_a_list_item_of_chips_keeps_its_bullet_on_the_first_chip(irs: dict[str, IR]):
    """B1 critique #11: an item whose only text is in detached chips gives its marker to the first
    chip's text, at the list's edge; the second chip carries none."""
    ir = irs["rows"]
    ai = [e for e in _texts(ir) if _text(e) == "AI" and e.box.y > 670]
    cloud = [e for e in _texts(ir) if _text(e) == "Cloud"]
    assert len(ai) == 1 and len(cloud) == 1
    assert ai[0].paragraphs[0]["bullet"] is not None
    assert cloud[0].paragraphs[0]["bullet"] is None
    assert ai[0].box.x < _first_line(ai[0])["x"], "the box reaches back to the list's edge, where the bullet is"


def test_two_spaced_chips_are_two_elements_and_the_space_is_nothing(irs: dict[str, IR]):
    ir = irs["rows"]
    baseline, target = _one(ir, "Baseline"), _one(ir, "Target")
    assert target.box.x > baseline.box.x2
    fills = [s for s in _shapes_in(ir, (195, 225)) if (s.fill or {}).get("color") in ("E8F4FE", "E6F6EC")]
    assert len(fills) == 2


def test_a_label_column_and_a_text_column_are_two_elements(irs: dict[str, IR]):
    ir = irs["rows"]
    label = _one(ir, "INSIGHT")
    column = [e for e in _texts(ir) if _text(e).startswith("Two columns on one line box")]
    assert len(column) == 1 and len(column[0].paragraphs[0]["lines"]) == 2
    assert label.box.w == pytest.approx(90.0, abs=0.01), "the label column keeps its own width"
    assert column[0].box.x >= label.box.x2


def test_a_number_and_its_unit_on_one_baseline_are_two_elements(irs: dict[str, IR]):
    number, unit = sorted(_texts(irs["rows"], "unit"), key=lambda e: e.box.x)
    assert (_text(number), _text(unit)) == ("6.2", "$B SOM")
    assert unit.box.x >= number.box.x2


def test_style_changes_inside_words_stay_runs_with_no_seam(irs: dict[str, IR]):
    ir = irs["rows"]
    [element] = _texts(ir, "styled-word")
    assert len(element.paragraphs) == 1
    runs = [r["text"] for line in element.paragraphs[0]["lines"] for r in line["runs"]]
    assert len(runs) >= 6 and "Studio" in runs
    assert not _glued(ir, element)


def test_a_padded_inline_span_between_spaces_stays_inline_and_its_spaces_carry_the_padding(irs: dict[str, IR]):
    """r1a: the span stays inline, and the 6 px padding either side of its text — which used to
    shift every word after it 12 px left (b.md's known gap) — rides on the space beside it, split
    off its run and letter-spaced by the padding. The text reads the same."""
    ir = irs["rows"]
    [element] = _texts(ir, "padded-inline")
    assert _text(element) == "Words and spaced and more words"
    runs = _runs(element)
    assert [r["text"] for r in runs] == ["Words and", " ", "spaced", " ", "and more words"]
    assert [round(runs[k]["letterSpacingPx"], 1) for k in (1, 3)] == [6.0, 6.0]
    assert not _glued(ir, element)


def _runs(element: Element, paragraph: int = 0, line: int = 0) -> list[dict]:
    return element.paragraphs[paragraph]["lines"][line]["runs"]


def _space_px(run: dict) -> float:
    """What a run of one space advances in the file: the face's space + its letter spacing (`spc`)."""
    from app.engine.emit import text as text_engine
    from tests.engine.helpers import measuring_face

    # The face itself, or its metric-compatible stand-in (Carlito for Calibri on Linux): the browser
    # draws the stand-in there, and the space advance is identical by design.
    face = measuring_face(run["font"], int(run.get("weight") or 400), bool(run.get("italic")))
    assert face is not None, f"{run['font']} (or a metric-compatible stand-in) is not installed"
    return text_engine.face_width_px(face, " ", float(run["sizePx"])) + float(run.get("letterSpacingPx") or 0.0)


#: The r1a cases: each a paragraph whose texts an inline box's margin or padding, or an empty
#: inline box, keeps apart — with the texts the file reads, spaces included.
SEAMS = {
    "seam-bullet": ["•", " ", "Severe leakage:", " net price fell 4.2% on discounting"],
    "seam-label": ["Target:", " ", "Business, events, premium upsell"],
    "seam-formula": ["GROWTH", " ", "+", " ", "LONGER TERMS", " ", "=", " ", "IMPACT"],
    "seam-legend": ["Residual exposure:", " ", "High", " ", "Low"],
    "seam-unit": ["70M", " ", "intl"],
    "seam-centred": ["PLAN FY2030", " ", "—", " ", "THE MANDATE"],
}


@pytest.mark.parametrize("name", sorted(SEAMS))
def test_a_layout_gap_between_two_texts_is_carried_by_a_space_sized_to_it(irs: dict[str, IR], name: str):
    """r1a (plan §16 #25; the survey's 66 collisions at r1): a gap the browser made with layout — a
    hand-made bullet's `margin-right`, a label's `margin-left`, padded operators, legend swatches, a
    unit's margin, a centred line's dash — used to glue its two texts in the file (`•Severe`,
    `Target:Business`). A space inserted between them, in the smaller of their two styles, is sized
    to the gap by its letter spacing: its run box is the gap itself, and its advance in the face the
    file names plus that spacing equals it. Nothing is glued, so R6 reports nothing."""
    ir = irs["rows"]
    [element] = _texts(ir, name)
    assert len(element.paragraphs) == 1 and len(element.paragraphs[0]["lines"]) == 1
    runs = _runs(element)
    assert [r["text"] for r in runs] == SEAMS[name]
    boxes = element.extras["x-run-boxes"][0][0]
    spaces = [k for k, r in enumerate(runs) if r["text"] == " "]
    for k in spaces:
        gap = boxes[k + 1][0] - boxes[k - 1][1]
        assert gap >= 3.0, f"{name}: the design's gap is at least 3 px"
        assert (boxes[k][0], boxes[k][1]) == (boxes[k - 1][1], boxes[k + 1][0])
        assert _space_px(runs[k]) == pytest.approx(gap, abs=0.05), (name, k)
        smaller = min(runs[k - 1]["sizePx"], runs[k + 1]["sizePx"])
        assert runs[k]["sizePx"] == smaller and runs[k]["underline"] is False
    assert not _glued(ir, element)


def test_a_wrapping_line_keeps_its_bullets_gap_and_its_second_line(irs: dict[str, IR]):
    """The case split boxes cannot place (r1a's measurement, 5 px off in both renderers): a hand-made
    bullet whose text wraps. The space carries the bullet's 4 px margin on line 1; line 2 starts at
    the paragraph's edge, under the bullet, where the browser put it."""
    ir = irs["rows"]
    [element] = _texts(ir, "seam-wrap")
    lines = element.paragraphs[0]["lines"]
    assert len(lines) == 2
    assert [r["text"] for r in lines[0]["runs"]][:3] == ["•", " ", "Low use:"]
    assert lines[1]["box"]["x"] == pytest.approx(lines[0]["box"]["x"], abs=0.5)
    boxes = element.extras["x-run-boxes"][0][0]
    assert _space_px(lines[0]["runs"][1]) == pytest.approx(boxes[2][0] - boxes[0][1], abs=0.05)


def test_the_seams_are_written_as_spaces_with_their_spacing(decks: dict[str, tuple[Path, IR]]):
    """The file carries what the IR says: each inserted space is an `a:r` of one space whose `spc`
    (hundredths of a point) is its letter spacing — so PowerPoint and PptxRender advance it by the
    gap (measured: within 1.3 px of the browser on 46 seams, r1a.md) — and the text row finds no join."""
    from pptx import Presentation
    from pptx.oxml.ns import qn

    deck, ir = decks["rows"]
    [legend] = _texts(ir, "seam-legend")
    expected = [round(r["letterSpacingPx"] * config.PT_PER_PX * 100) for r in _runs(legend) if r["text"] == " "]
    shape = next(s for s in Presentation(str(deck)).slides[0].shapes if s.name == "seam-legend")
    written = [(r.find(qn("a:t")).text, r.find(qn("a:rPr")).get("spc"))
               for r in shape.text_frame._txBody.iter(qn("a:r"))]
    assert [t for t, _ in written] == SEAMS["seam-legend"]
    # The width lock may add its per-character correction (a few hundredths of a point) on top.
    got = [int(spc) for t, spc in written if t == " "]
    assert len(got) == len(expected) and all(abs(g - e) <= 10 for g, e in zip(got, expected, strict=False)), (got, expected)
    assert text_report(deck, [ir])["collisions"] == []


def test_a_text_blocks_own_opacity_is_on_its_text_and_the_box_it_detaches(irs: dict[str, IR]):
    """r1a (plan §16 #25): an in-flow text block is pushed with its container's context, so its own
    `opacity` was lost on its text and on the box it detaches. Now the text carries it as run alpha —
    one element can hold paragraphs of different opacity (flexgrid's cards) — and the detached box as
    element opacity (E's `pseudoContext`, as a pseudo copy gets it); the card around it is untouched."""
    ir = irs["rows"]
    card = next(e for e in ir.elements if e.kind == "shape" and e.name == "faded-card")
    assert (card.opacity or 1.0) == pytest.approx(1.0)
    texts = sorted((e for e in _texts(ir) if 400 < e.box.y < 412 and e.box.x < 420), key=lambda e: e.box.x)
    assert [_text(e) for e in texts] == ["A faded label with a", "NOTE", "inside it"]
    before, box, after = texts
    for segment in (before, after):
        assert (segment.opacity or 1.0) == pytest.approx(1.0)
        assert {round(r["alpha"], 3) for r in _runs(segment)} == {0.6}
    assert box.opacity == pytest.approx(0.6) and _runs(box)[0]["alpha"] == pytest.approx(1.0)


def test_flexgrid_labels_keep_their_opacity(decks: dict[str, tuple[Path, IR]]):
    """G-2's m3 (flexgrid `.card .lbl{opacity:.85}`): each card is one element of two paragraphs, the
    number at full alpha and the label at 0.85 — in the IR and in the file (`a:alpha` 85000)."""
    from pptx import Presentation
    from pptx.oxml.ns import qn

    deck, ir = decks["flexgrid"]
    cards = [e for e in _texts(ir) if len(e.paragraphs or []) == 2
             and _runs(e, 1)[0]["sizePx"] == pytest.approx(11.0)]
    assert len(cards) == 6
    for card in cards:
        assert (card.opacity or 1.0) == pytest.approx(1.0)
        assert [r["alpha"] for r in _runs(card, 0)] == [pytest.approx(1.0)]
        assert [r["alpha"] for r in _runs(card, 1)] == [pytest.approx(0.85)]
    alphas = []
    for shape in Presentation(str(deck)).slides[0].shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size is not None and round(run.font.size.pt * 4) == round(11 * 0.75 * 4):
                    node = run._r.find(qn("a:rPr")).find(".//" + qn("a:alpha"))
                    alphas.append(None if node is None else int(node.get("val")))
    assert alphas == [85000] * 6


def test_every_line_carries_its_run_boxes(irs: dict[str, IR]):
    """`extras["x-run-boxes"]`: per paragraph, per line, per run `[left, right]`, inside the line box."""
    for stem, ir in irs.items():
        for element in _texts(ir):
            boxes = (element.extras or {}).get("x-run-boxes")
            assert boxes is not None, f"{stem}/{element.id} has no run boxes"
            assert len(boxes) == len(element.paragraphs), f"{stem}/{element.id}"
            for paragraph, per_line in zip(element.paragraphs, boxes, strict=False):
                assert len(per_line) == len(paragraph["lines"])
                for line, per_run in zip(paragraph["lines"], per_line, strict=False):
                    assert len(per_run) == len(line["runs"]), f"{stem}/{element.id}: {per_run} vs {line['runs']}"
                    box = line["box"]
                    for left, right in per_run:
                        assert left <= right + 0.01
                        assert box["x"] - 1 <= left and right <= box["x"] + box["w"] + 1


GLYPHS_JS = r"""
() => {
  const glyphs = [];
  const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
  const range = document.createRange();
  let node;
  while ((node = walker.nextNode())) {
    const parent = node.parentElement;
    if (!parent || parent.closest('aside.notes')) continue;
    const cs = getComputedStyle(parent);
    if (cs.display === 'none' || cs.visibility === 'hidden') continue;
    for (let i = 0; i < node.data.length; i++) {
      if (/\s/.test(node.data[i])) continue;
      range.setStart(node, i); range.setEnd(node, i + 1);
      for (const q of range.getClientRects()) {
        if (q.width > 0) { glyphs.push([q.left, q.right, (q.top + q.bottom) / 2]); break; }
      }
    }
  }
  return glyphs;
}
"""


@pytest.fixture(scope="module")
def glyph_pages() -> dict[str, tuple[IR, list]]:
    """`rows.html` and `text.html`, each extracted and then asked for every glyph's rect."""
    found = {}
    with measuring_page(Canvas(1280, 720)) as page:
        for html in (ROWS, config.TORTURE_FIXTURE / "text.html"):
            ir = _extract(html, page)
            found[html.stem] = (ir, page.evaluate(GLYPHS_JS))
    return found


@pytest.mark.parametrize("stem", ["rows", "text"])
def test_every_line_starts_and_ends_at_a_glyph(glyph_pages: dict[str, tuple[IR, list]], stem: str):
    """B1 critique #1: a line's box spans the ink of the characters it carries. A space the browser
    did not collapse (after a chip, a badge, an inline image) belongs to the gap, not to the text
    beside it: counted in, it put the text one space toward the chip, and E1 writes the line's x."""
    ir, glyphs = glyph_pages[stem]
    wrong = []
    for element in _texts(ir):
        if element.rotation:
            continue
        for p, paragraph in enumerate(element.paragraphs or []):
            for index, line in enumerate(paragraph["lines"]):
                box = line["box"]
                inside = [g for g in glyphs if box["y"] <= g[2] <= box["y"] + box["h"]
                          and g[0] >= box["x"] - 0.5 and g[1] <= box["x"] + box["w"] + 0.5]
                if not inside:
                    continue
                left, right = min(g[0] for g in inside), max(g[1] for g in inside)
                if abs(left - box["x"]) > 0.5 or abs(box["x"] + box["w"] - right) > 0.5:
                    text = "".join(r["text"] for r in line["runs"])
                    wrong.append((element.id, p, index, text[:30], round(left - box["x"], 2),
                                  round(box["x"] + box["w"] - right, 2)))
    assert wrong == []


# ------------------------------------------------------------------------------ the text row


@pytest.mark.parametrize("stem", ("rows",) + PROBE_STEMS)
def test_no_element_keeps_two_paragraphs_side_by_side(irs: dict[str, IR], stem: str):
    assert text_rows([irs[stem]]) == []


def _findings(report: dict) -> dict[str, list]:
    return {key: value for key, value in report.items() if isinstance(value, list) and value}


def test_the_rows_deck_passes_the_text_row_and_fits(decks: dict[str, tuple[Path, IR]]):
    """Every check of the text row on the emitted rows deck — G-1's joins and export overlaps, and
    WP-B's rows, positions (V3: every paragraph's first line starts where the browser put it, from
    the emitter's measured `marL`/`marR`/`indent`) and seams — plus fit."""
    deck, ir = decks["rows"]
    report = text_report(deck, [ir])
    assert report["checked"] and {"rows", "positions", "seams"} <= set(report)
    assert _findings(report) == {}
    fitted = fit(deck, [ir])
    assert fitted.clean, [(i.kind, i.shape, i.detail) for i in fitted.issues]


@pytest.mark.parametrize("stem", ("flexgrid", "layout-misc"))
def test_the_flex_families_pass_the_text_row_and_fit(decks: dict[str, tuple[Path, IR]], stem: str):
    """The text row and fit on the whole family (brief §4). As probes, p01's shrink-to-fit bold heading
    overhung under WP-D's rule (drawn in Calibri Light by the browser, written `+mj-lt` `b=1`, Calibri
    Bold through test-16x9's theme part — the masters' face mismatch) and was let through here; the
    promoted `flexgrid` gives its heading a width, as `charts.html` does, so nothing is let through."""
    deck, ir = decks[stem]
    assert _findings(text_report(deck, [ir])) == {}
    fitted = fit(deck, [ir])
    assert fitted.clean, [(i.kind, i.shape, i.detail) for i in fitted.issues]


def test_text_misc_passes_the_text_row_but_for_its_css_columns(decks: dict[str, tuple[Path, IR]]):
    """text-misc (p10) row B is CSS multi-column text (package F's): before F its two columns were one
    paragraph whose lines held both columns, glued (a seam, and joins); anything the text row still
    reports must be there. Rows C, D, E are WP-B's."""
    deck, ir = decks["text-misc"]
    columns = {e.id for e in _texts(ir) if "sponsor" in _text(e)}
    names = {e.id: e.name for e in ir.elements}
    assert len(columns) == 1
    report = _findings(text_report(deck, [ir]))
    for key, entries in report.items():
        for entry in entries:
            where = entry.get("element") or next((i for i, n in names.items() if n == entry.get("shape")), None)
            assert where in columns, (key, entry)


# ---------------------------------------------------------------------- the promoted probe families


def test_flexgrid_rows_cards_badges_kpis_and_grid(irs: dict[str, IR]):
    ir = irs["flexgrid"]
    # A: five label | value rows, two elements each, the value right of its label.
    for label, value in (("Revenue 2023", "$15.2B"), ("EBITDA margin", "21.4%"), ("Headcount", "5.1M"),
                         ("Market share", "34%"), ("Deals closed", "412")):
        assert _one(ir, value).box.x >= _one(ir, label).box.x2
    # B: six stat cards, each one element of number + label.
    cards = [e for e in _texts(ir) if len(e.paragraphs) == 2 and e.paragraphs[0]["lines"][0]["runs"][0]["sizePx"] == 24]
    assert len(cards) == 6
    # C: the badges' numbers, centred both ways in their circles.
    for number in ("1", "2", "3"):
        badge = _one(ir, number)
        assert (badge.paragraphs[0]["align"], badge.anchor) == ("center", "middle")
    # D: five KPIs, each two centred paragraphs.
    kpis = [e for e in _texts(ir) if len(e.paragraphs) == 2 and e.paragraphs[0]["lines"][0]["runs"][0]["sizePx"] == 28]
    assert len(kpis) == 5
    assert all([p["align"] for p in kpi.paragraphs] == ["center", "center"] for kpi in kpis)
    # E: the label | description grid, eight cells.
    grid = [e for e in _texts(ir) if 560 <= e.box.y < 620]
    assert len(grid) == 8


def test_layout_misc_pills_chips_and_circles(irs: dict[str, IR]):
    ir = irs["layout-misc"]
    # C: each pill's text starts right of its dot.
    dots = _shapes_in(ir, (300, 320), width=(7.5, 8.5))
    pills = sorted((_one(ir, t) for t in ("On track", "Sovereign cloud", "Applied AI")), key=lambda e: e.box.x)
    assert len(dots) == 3
    for dot, pill in zip(dots, pills, strict=False):
        assert pill.box.x >= dot.box.x2 + 5
    # F: flex-centred chips and grid-centred circles.
    for label in ("Leader", "Challenger", "Emerging", "1", "2", "3"):
        element = _one(ir, label)
        assert (element.paragraphs[0]["align"], element.anchor) == ("center", "middle"), label


def test_text_misc_contents_unit_and_legend(irs: dict[str, IR]):
    ir = irs["text-misc"]
    # C: four cells over four fills.
    cells = sorted((_one(ir, t) for t in ("Sources", "Uses", "Valuation", "Comps")), key=lambda e: e.box.x)
    fills = _shapes_in(ir, (325, 340))
    assert len(fills) == 4
    for cell, fill in zip(cells, fills, strict=False):
        assert fill.box.x <= cell.box.x and cell.box.x2 <= fill.box.x2
    # D: the number and its unit, two elements, the unit right of the number.
    assert _one(ir, "$B SOM").box.x >= _one(ir, "6.2").box.x2
    # E: four labels, each right of its swatch.
    swatches = _shapes_in(ir, (425, 450), width=(9.5, 10.5))
    labels = sorted((_one(ir, t) for t in ("UAE", "KSA", "Qatar", "Bahrain")), key=lambda e: e.box.x)
    assert len(swatches) == 4
    for swatch, label in zip(swatches, labels, strict=False):
        assert label.box.x >= swatch.box.x2 + 5


# ------------------------------------------------------------------------------- placeholders


def _placeholder_text(name: str, box: Box, size: float) -> Element:
    line = {"box": box.to_json(), "runs": [{"text": name, "font": "Calibri", "sizePx": size,
                                            "weight": 400, "italic": False, "color": "000000", "alpha": 1}]}
    return Element(kind="text", box=box, name=name, source={"path": name},
                   paragraphs=[{"align": "left", "lineHeightPx": box.h, "spaceBeforePx": 0,
                                "spaceAfterPx": 0, "bullet": None, "lines": [line]}])


def _title_zone_of_layout_06():
    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")
    layout = manifest.layout("layout-06")
    return layout, next(p for p in layout.placeholders if p.type == "title")


def _mapped(layout, elements: list[Element]) -> dict[str, dict | None]:
    ir = IR(canvas=Canvas(1280, 720), slide=Slide(id="s", layoutId="layout-06"),
            elements=assign_ids_and_z(elements))
    return {e.name: e.placeholder for e in map_placeholders(ir, layout).elements}


def test_the_title_placeholder_goes_to_the_heading_not_the_chip_inside_it():
    """Pass 2: a chip wholly inside the title zone (100 %) must not take the title from the heading
    it sits on (the heading covers less of *itself* with the zone): the larger type wins."""
    layout, zone = _title_zone_of_layout_06()
    # The heading runs 5 % past the zone's right edge; the chip sits inside it.
    heading = _placeholder_text("heading", Box(zone.x + 4, zone.y + 4, zone.w * 1.0, zone.h - 8), 28)
    chip = _placeholder_text("chip", Box(zone.x + zone.w * 0.7, zone.y + 8, 80, zone.h - 16), 10)
    by_name = _mapped(layout, [chip, heading])
    assert by_name["heading"] == {"type": "title", "idx": zone.idx}
    assert by_name["chip"] is None


def test_the_title_placeholder_goes_to_the_title_not_a_larger_standfirst():
    """B1 critique #7: a short title wholly inside the title zone against a long standfirst 75 %
    inside it. Shared area alone handed the standfirst the title (48,000 px² against 8,800); the
    title's larger type keeps it."""
    layout, zone = _title_zone_of_layout_06()
    title = _placeholder_text("title", Box(zone.x + 4, zone.y + 6, 220, 40), 28)
    standfirst = _placeholder_text("standfirst", Box(zone.x + 4, zone.y + 60, 800, 80), 16)
    assert standfirst.box.intersect(zone.box).area > title.box.intersect(zone.box).area
    assert standfirst.box.overlap_fraction(zone.box) >= 0.6
    by_name = _mapped(layout, [title, standfirst])
    assert by_name["title"] == {"type": "title", "idx": zone.idx}
    assert by_name["standfirst"] is None


# ------------------------------------------------ WP-E's pseudo copies beside WP-B's segments


PSEUDO_SLIDE = """<!doctype html><html><head><meta charset="utf-8"><style>
html,body{{margin:0;width:1280px;height:720px;overflow:hidden;background:transparent;
  font-family:Arial,sans-serif;color:#1d2433}}
.chip{{display:inline-block;padding:2px 8px;background:#E8F1FB;color:#0B5CAD;font-size:14px}}
.accent{{position:relative}}
.accent::after{{content:"";position:absolute;left:0;bottom:-6px;width:60px;height:3px;background:#D64545}}
{css}
</style></head><body>
{body}
</body></html>"""


def _pseudo_slide(tmp_path: Path, body: str, css: str = "") -> IR:
    """One slide on the `test-16x9` master (as the torture families are), extracted with WP-E's pre-pass."""
    html = tmp_path / "slide.html"
    html.write_text(PSEUDO_SLIDE.format(css=css, body=body), encoding="utf-8")
    (tmp_path / "slide.expect.json").write_text('{"master": "test-16x9", "layoutId": "layout-06"}',
                                               encoding="utf-8")
    context = context_for(html)
    return extract_html(html, context.manifest, context.layout_id, context.assets_dir, slide_id="slide")


def _accents(ir: IR) -> list[Element]:
    return [e for e in ir.elements if e.kind == "shape" and (e.fill or {}).get("color") == "D64545"]


def _undrawn(ir: IR) -> list[str]:
    return [d.message for d in ir.diagnostics if "::after" in d.message or "::before" in d.message]


def test_a_segmented_heading_keeps_its_pseudo_accent_once_after_its_text(tmp_path):
    """B1 critique #4, at the E rebase: a heading with an absolute `::after` accent and a glued chip.
    The chip splits the heading into segments; the accent rides on the last text segment and is
    drawn exactly once, after the heading's text and the chip (positioned boxes paint last)."""
    ir = _pseudo_slide(tmp_path, """
<div style="position:absolute;left:40px;top:40px;width:600px"><div class="accent" style="font-size:24px">Market outlook<span class="chip">NEW</span> for 2026</div></div>
""")
    (accent,) = _accents(ir)
    assert (accent.box.x, accent.box.w, accent.box.h) == (40, 60, 3)
    texts = [e for e in ir.elements if e.kind == "text"]
    assert {t for e in texts for t in _paragraphs(e)} >= {"Market outlook", "NEW"}
    assert all(accent.z > e.z for e in texts), "the accent paints after every text of its heading"
    assert not _undrawn(ir)


def test_a_block_of_chips_only_keeps_its_pseudo_accent(tmp_path):
    """No text of the block outside its detached boxes: the accent still paints, once, after them,
    and in its host's context (WP-E's `walkDeferred`: the host's opacity is the copy's too)."""
    ir = _pseudo_slide(tmp_path, """
<div style="position:absolute;left:40px;top:120px;width:600px"><div class="accent" style="opacity:0.5"><span class="chip">Alpha</span> <span class="chip">Beta</span></div></div>
""")
    (accent,) = _accents(ir)
    chips = [e for e in ir.elements if e.kind == "text" and _paragraphs(e)[0] in ("Alpha", "Beta")]
    assert len(chips) == 2 and all(accent.z > chip.z for chip in chips)
    assert accent.opacity == pytest.approx(0.5)
    assert not _undrawn(ir)


@pytest.mark.parametrize("text", [
    '<span style="visibility:hidden">Hidden</span>',     # no text node to measure: no segment at all
    '<span style="font-size:0">Zero</span>',             # a segment whose text measures no line
])
def test_a_block_whose_text_measures_no_line_keeps_its_pseudo_accent(tmp_path, text: str):
    """A text block with no measurable line never enters the run; its visible accent is WP-E's copy
    all the same — drawn once, through `walkDeferred` (in its host's context), never reported."""
    ir = _pseudo_slide(tmp_path, f"""
<div style="position:absolute;left:40px;top:200px;width:600px"><div class="accent" style="height:30px;font-size:16px;opacity:0.5">{text}</div></div>
""")
    (accent,) = _accents(ir)
    assert accent.opacity == pytest.approx(0.5)
    assert not [e for e in ir.elements if e.kind == "text"]
    assert not _undrawn(ir)


def test_a_text_less_dot_before_a_label_lands_the_label_after_it(tmp_path):
    """E3 review M4, closed by WP-B's line geometry and E1 margins: a legend dot — a `::before` copy
    or a real `<span>`, inline-block and text-less — stays inline paint, and the label's first line
    starts after it, in the IR (58 px: 40 + 10 + 8) and in the emitted file (V3 finds nothing)."""
    ir = _pseudo_slide(tmp_path, """
<div style="position:absolute;left:40px;top:40px;width:300px;font-size:16px">
  <div class="it">North</div><div class="it">East</div><div><span class="dot"></span>Real span</div>
</div>
""", ".it::before,.dot{content:'';display:inline-block;width:10px;height:10px;border-radius:50%;"
     "margin-right:8px;background:#0B5CAD}")
    dots = sorted((e for e in ir.elements if e.kind == "shape" and (e.fill or {}).get("color") == "0B5CAD"),
                  key=lambda e: e.box.y)
    labels = sorted((e for e in ir.elements if e.kind == "text"), key=lambda e: e.box.y)
    assert len(dots) == 3 and [p for e in labels for p in _paragraphs(e)] == ["North", "East", "Real span"]
    starts = [line["box"]["x"] for e in labels for p in e.paragraphs for line in p["lines"][:1]]
    assert starts == [pytest.approx(58, abs=0.5)] * 3
    assert all(dot.box.x2 <= 58 + 0.5 for dot in dots)
    context = context_for(tmp_path / "slide.html")
    deck = tmp_path / "slide.pptx"
    emitter.emit([ir], context.manifest, context.master, deck, EmitOptions())
    report = text_report(deck, [ir])
    assert report["checked"] and _findings(report) == {}


def test_a_display_contents_host_in_a_flex_row_is_no_glue(tmp_path):
    """E3 review m5, closed by WP-B's R5: a `display: contents` host's in-flow `::before` text is a
    flex item of the row like its host's children, so the three are three elements where the browser
    put them — no glued seam."""
    ir = _pseudo_slide(tmp_path, """
<div style="position:absolute;left:40px;top:40px;display:flex;gap:8px;font-size:16px"><div class="dc"><div>Inside</div></div><div>After</div></div>
""", ".dc{display:contents} .dc::before{content:'DC';color:#D64545}")
    texts = sorted((e for e in ir.elements if e.kind == "text"), key=lambda e: e.box.x)
    assert [_text(e) for e in texts] == ["DC", "Inside", "After"]
    assert texts[1].box.x >= texts[0].box.x2 + 7.5 and texts[2].box.x >= texts[1].box.x2 + 7.5
    assert texts[0].paragraphs[0]["lines"][0]["runs"][0]["color"] == "D64545"
    assert not [d for d in ir.diagnostics if d.message.startswith("glued runs:")]
    assert not _undrawn(ir)
