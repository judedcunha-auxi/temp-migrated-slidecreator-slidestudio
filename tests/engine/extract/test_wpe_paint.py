"""WP-E — paint what the browser paints (14-WPE-paint): colours, gradient text, `::before`/`::after`,
CSS border joins.

Every number below is re-derivable from the HTML in the test or from the named family with a ruler
and a colour table: `color-mix(in srgb, #1E9E5A 12%, white)` is `E4F3EB` because 0.12 × 0x1E + 0.88
× 0xFF = 228.0 = 0xE4, and the `pseudo` (probe p03) triangle starts at x = 272 because `.flow` sits at 40 and `.tri`
at 232 inside it. The pixel-agreement test does not trust the normaliser's own arithmetic at all: it
asks the browser to paint each colour over white and over black and reads the pixels back.

Inputs that are missing make a test fail, never skip (master brief §10.7).
"""
from __future__ import annotations

import io
import re
from pathlib import Path

import pytest
from PIL import Image

from app.config import engine as config
from app.engine.extract import html as html_extract
from app.engine.extract.html import COLOR_JS
from app.engine.ir import IR, Canvas, Element
from app.engine.verify.fixtures import context_for
from tests.engine.helpers import expand_svg, extract_html

#: The fidelity probes p03, p04 and p10 are the torture families `pseudo`, `colours` and `text-misc`
#: since WP-G (G-2, plan 16 #16).
PROBES = config.TORTURE_FIXTURE
#: The readiness panel's classifier of measured findings: Slide Studio's `web/src/lib/words.ts`
#: `MEASURED` table (first match wins, the last row is the catch-all), copied here because the
#: frontend is not part of this service. Engine messages must keep landing on the right row. When
#: the classifier is ported (Phase 3 / 7c), point these tests at it instead of this copy.
MEASURED: tuple[tuple[str, str, str], ...] = (
    (r"^glued runs:", "i", "glued-runs"),
    (r"rasterised: remote image", "i", "remote-raster"),
    (r"rasterised at its displayed size", "i", "raster-svg"),
    (r"^rasterised:", "i", "raster-css"),
    (r"possible chart", "i", "hand-drawn-chart"),
    (r"is not installed here", "i", "font-not-installed"),
    (r"was requested but is not in", "i", "img-not-asset"),
    (r"image did not load", "i", "image-missing"),
    (r"^box-shadow:", "i", "shadow-dropped"),
    (r"data-chart is not valid JSON", "i", "chart-spec"),
    (r"placeholder", "i", "placeholder-mismatch"),
    (r"text is clipped", "i", "text-clipped"),
    (r"outside the canvas", "i", "outside-canvas"),
    (r"looping animation", "i", "animation"),
    (r"animation\(s\) still running", "i", "measure-failed"),
    (r"overflowed into .* column", "i", "text-overflow"),
    (r"column-rule|fragmented across columns", "i", "approximated"),
    (r"remote resource blocked by policy|blocked remote resource", "i", "remote-resource"),
    (r"measuring page is scrolled|isolated capture failed|no expansion came back|no element matches"
     r"|unknown record kind", "i", "measure-failed"),
    (r"has no area; dropped|fewer than two points; dropped|empty d; dropped", "i", "invisible-shape"),
    (r"colour .* could not be read", "i", "unreadable-colour"),
    (r"::(before|after)", "i", "pseudo-dropped"),
    (r"^not exported:|was not painted|not expanded|references missing id|does not start with a move"
     r"|dropped$", "i", "dropped-element"),
    (r"arrowhead of .* line-end size", "i", "approximated"),
    (r"gradient|roundRect|stroke-linecap|text-transform|skewed rect|clipPath rect widened"
     r"|fragments; merged|plain container|text stroke dropped|counter style|counter\(list-item\)"
     r"|quotes for lang", "i", "approximated"),
    (r".", "", "extract"),
)

SLIDE = """<!doctype html><html><head><meta charset="utf-8"><style>
html,body{{margin:0;width:1280px;height:720px;overflow:hidden;background:transparent;
  font-family:Arial,sans-serif;color:#1d2433}}
{css}
</style></head><body>
{body}
</body></html>"""


def _slide(tmp_path: Path, body: str, css: str = "", name: str = "slide") -> Path:
    path = tmp_path / f"{name}.html"
    path.write_text(SLIDE.format(css=css, body=body), encoding="utf-8")
    return path


def _probe_ir(stem: str) -> IR:
    html = PROBES / f"{stem}.html"
    assert html.exists(), f"{html} is missing"
    context = context_for(html)
    return extract_html(html, context.manifest, context.layout_id, context.assets_dir, slide_id=stem, title=stem)


def _named(ir: IR, name: str) -> list[Element]:
    return [e for e in ir.elements if e.name == name]


def _runs(element: Element) -> list[dict]:
    return [run for p in element.paragraphs or [] for line in p["lines"] for run in line["runs"]]


def _hex(value: str) -> tuple[int, int, int]:
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


# ============================================================================================ colours

#: Every syntax the normaliser must read. In-gamut only: a wide-gamut colour is clamped here while
#: the browser gamut-maps it (the `colour-gamut` lint), which is not an error the ±2 check should see.
PIXEL_CASES = [
    "color-mix(in srgb, #1E9E5A 12%, white)",
    "color-mix(in oklab, #0B5CAD 10%, transparent)",
    "color-mix(in oklab, #0B5CAD 60%, #1E9E5A)",
    "oklch(0.95 0.04 150)",
    "oklch(0.55 0.12 250 / 0.6)",
    "oklab(0.6 -0.05 0.04)",
    "hsl(150 60% 94%)",
    "hsl(210 70% 40% / 0.5)",
    "hwb(150 10% 20%)",
    "#1234",
    "#1E9E5A1F",
    "rgb(11 92 173 / 8%)",
    "color(srgb 0.2 0.4 0.6)",
    "color(display-p3 0.9 0.97 0.92)",
    "lab(95 -5 3)",
    "lch(50 30 120)",
    "rebeccapurple",
    "color(srgb none 0.5 0.2)",
]


@pytest.fixture
def colour_page(page):
    """A blank page with `color.js` evaluated — the normaliser alone, no `page.js`."""
    page.set_content("<!doctype html><html><body style='margin:0;background:#fff'></body></html>")
    page.add_script_tag(path=str(COLOR_JS))
    return page


def _normalise(page, text: str, selector: str | None = None):
    return page.evaluate(
        "([t, s]) => window.__engineColor.normalise(t, s ? document.querySelector(s) : null)", [text, selector]
    )


def test_every_syntax_agrees_with_the_pixels_the_browser_paints(colour_page):
    """Paint each colour over white and over black; the normaliser must predict both pixels.

    Recovering the colour from the two pixels divides by alpha, which turns one unit of rounding into
    twelve at alpha 0.08 — so the RGB comparison is made where the eye makes it, on the composited
    pixels (±2 per channel), and alpha is recovered directly (±0.01).
    """
    cells = "".join(
        f"<div style='position:absolute;left:{i * 20}px;top:0;width:20px;height:40px;background:#fff'>"
        f"<div style='width:20px;height:20px;background:{text}'></div></div>"
        f"<div style='position:absolute;left:{i * 20}px;top:40px;width:20px;height:20px;background:#000'>"
        f"<div style='width:20px;height:20px;background:{text}'></div></div>"
        for i, text in enumerate(PIXEL_CASES)
    )
    colour_page.set_content(f"<!doctype html><html><body style='margin:0;background:#fff'>{cells}</body></html>")
    colour_page.add_script_tag(path=str(COLOR_JS))
    shot = Image.open(io.BytesIO(colour_page.screenshot(clip={"x": 0, "y": 0, "width": 20 * len(PIXEL_CASES),
                                                                "height": 60}))).convert("RGB")
    for i, text in enumerate(PIXEL_CASES):
        read = _normalise(colour_page, text)
        assert read is not None, f"{text} could not be read"
        white = shot.getpixel((i * 20 + 10, 10))
        black = shot.getpixel((i * 20 + 10, 50))
        alpha = 1 - sum(w - b for w, b in zip(white, black, strict=False)) / (3 * 255)
        assert read["alpha"] == pytest.approx(alpha, abs=0.01), f"{text}: alpha {read['alpha']} vs painted {alpha:.3f}"
        predicted = _hex(read["color"])
        for ground, painted in ((255, white), (0, black)):
            expected = [c * read["alpha"] + ground * (1 - read["alpha"]) for c in predicted]
            assert all(abs(e - p) <= 2 for e, p in zip(expected, painted, strict=False)), (
                f"{text}: normalised {read} composites to {[round(e) for e in expected]} over {ground}, "
                f"the browser painted {painted}"
            )


@pytest.mark.parametrize(
    ("text", "colour", "alpha"),
    [
        ("color-mix(in srgb,#1E9E5A 12%,white)", "E4F3EB", 1.0),
        ("hsl(150 60% 94%)", "E7F9F0", 1.0),
        ("#1E9E5A1F", "1E9E5A", 31 / 255),
        ("rgb(11 92 173 / 8%)", "0B5CAD", 0.08),
        ("rgba(11, 92, 173, .5)", "0B5CAD", 0.5),
        ("#1234", "112233", 0x44 / 255),
        ("transparent", "000000", 0.0),
        ("color(srgb none 0.5 0.2)", "008033", 1.0),
    ],
)
def test_exact_arithmetic(colour_page, text: str, colour: str, alpha: float):
    read = _normalise(colour_page, text)
    assert read["color"] == colour
    assert read["alpha"] == pytest.approx(alpha, abs=1e-6)


def test_a_wide_gamut_colour_is_clamped_not_dropped(colour_page):
    read = _normalise(colour_page, "color(display-p3 1 0 0)")
    assert read["color"][:2] == "FF" and read["alpha"] == 1


def test_an_unreadable_string_is_null_and_recorded_once(colour_page):
    assert _normalise(colour_page, "garbage(1 2 3)") is None
    assert _normalise(colour_page, "garbage(1 2 3)") is None
    # asking quietly ("is this a colour?") records nothing: a gradient's direction is not a failure
    assert colour_page.evaluate("() => window.__engineColor.normalise('to right', null, true)") is None
    assert colour_page.evaluate("() => window.__engineColor.unparsed()") == ["garbage(1 2 3)"]


def test_currentcolor_resolves_against_the_element_it_is_read_from(colour_page):
    colour_page.evaluate(
        "() => { const d = document.createElement('div'); d.id = 'c'; d.style.color = '#0B5CAD';"
        " document.body.appendChild(d); }"
    )
    read = _normalise(colour_page, "color-mix(in srgb, currentcolor 20%, white)", "#c")
    assert read["color"] == "CEDEEF"
    assert _normalise(colour_page, "currentColor", "#c")["color"] == "0B5CAD"


def test_the_probe_lives_outside_the_slide(colour_page):
    """One span under <html>, hidden and off-screen: nothing under <body> can see it."""
    _normalise(colour_page, "oklch(0.5 0.1 150)")
    where = colour_page.evaluate(
        "() => [...document.querySelectorAll('[data-engine-probe]')].map(p => p.parentElement.tagName)"
    )
    assert where == ["HTML"]


def test_tokens_split_at_depth_zero_outside_quotes(colour_page):
    tokens = colour_page.evaluate(
        "() => window.__engineColor.tokens('oklch(0.2 0.1 240 / 0.3) 0px  4px 12px \"a b\" x')"
    )
    assert tokens == ["oklch(0.2 0.1 240 / 0.3)", "0px", "4px", "12px", '"a b"', "x"]


def test_p04_every_swatch_keeps_its_fill_and_text_colour():
    ir = _probe_ir("colours")
    # twelve swatches and the banner, all rounded boxes; the banner's left edge is the 14th shape
    swatches = [e for e in ir.elements if e.kind == "shape" and e.geometry["type"] == "roundRect"]
    filled = [e for e in swatches if (e.fill or {}).get("type") in ("solid", "gradient")]
    assert len(filled) == 13, [(e.name, e.fill) for e in swatches]
    assert not [d for d in ir.diagnostics if "could not be read" in d.message]

    by_left = {round(e.box.x): e for e in swatches if round(e.box.y) == 90}
    assert by_left[940].fill["color"] == "0B5CAD" and by_left[940].fill["alpha"] == pytest.approx(0.1, abs=1e-3)
    assert by_left[640].fill["color"] == "E4F3EB"
    gradient = next(e for e in swatches if (e.fill or {}).get("type") == "gradient").fill
    assert gradient["stops"][0]["color"] == "CEDEEF" and gradient["angle"] == 90

    texts = [e for e in ir.elements if e.kind == "text"]
    swatch_11 = next(t for t in texts if "11" in _runs(t)[0]["text"])
    swatch_12 = next(t for t in texts if "12" in _runs(t)[0]["text"])
    for run in _runs(swatch_11):
        assert all(abs(a - b) <= 1 for a, b in zip(_hex(run["color"]), _hex("DADEE3"), strict=False)), run["color"]
    assert all(run["color"] != "000000" for run in _runs(swatch_12))


def test_a_colour_nobody_can_read_is_one_warning_and_no_shape(page, tmp_path, sample_manifest, sample_dir):
    """The browser resolves every colour it accepts, so an unreadable one is simulated: the page's
    normaliser is wrapped (before any script runs) to refuse one colour, the way a future syntax the
    probe cannot convert would be refused. What is under test is the walk's handling of the refusal."""
    page.add_init_script("""
      Object.defineProperty(window, '__engineColor', { configurable: true, get() { return undefined; },
        set(real) {
          const refused = new Set();
          const wrapped = Object.assign({}, real, {
            normalise(t, el, quiet) {
              if (/^rgb\\(1, 2, 3\\)$/.test(String(t).trim())) { if (!quiet) refused.add(String(t).trim()); return null; }
              return real.normalise(t, el, quiet);
            },
            unparsed() { return real.unparsed().concat([...refused]).sort(); },
          });
          Object.defineProperty(window, '__engineColor', { value: wrapped, writable: true, configurable: true });
        } });
    """)
    html = _slide(tmp_path, "<div data-name='a' style='position:absolute;left:10px;top:10px;width:50px;"
                            "height:50px;background:rgb(1,2,3)'></div>"
                            "<div data-name='b' style='position:absolute;left:100px;top:10px;width:50px;"
                            "height:50px;background:rgb(1,2,3)'></div>")
    ir = extract_html(html, sample_manifest, sample_manifest.layouts[0].id, sample_dir / "assets", page=page)
    assert not [e for e in ir.elements if e.kind == "shape"]
    warns = [d for d in ir.diagnostics if "could not be read" in d.message]
    assert len(warns) == 1 and warns[0].level == "warn" and "rgb(1, 2, 3)" in warns[0].message


def test_shadows_read_the_colour_first_serialisation(tmp_path, ir_of):
    html = _slide(tmp_path, "<div data-name='s' style='position:absolute;left:40px;top:40px;width:100px;"
                            "height:60px;background:#fff;box-shadow:0 4px 12px oklch(0.2 0.1 240 / 0.3)'></div>")
    shadow = _named(ir_of(html), "s")[0].shadow
    assert shadow["dx"] == 0 and shadow["dy"] == 4 and shadow["blur"] == 12
    assert shadow["alpha"] == pytest.approx(0.3, abs=1e-6)


def test_gradient_interpolation_head_is_read_and_reported(tmp_path, ir_of):
    html = _slide(tmp_path, "<div data-name='g' style='position:absolute;left:40px;top:40px;width:200px;"
                            "height:60px;background:linear-gradient(90deg in oklab, #0B5CAD, #1E9E5A)'></div>"
                            "<div data-name='h' style='position:absolute;left:40px;top:140px;width:200px;"
                            "height:60px;background:linear-gradient(90deg, #0B5CAD, #1E9E5A)'></div>")
    ir = ir_of(html)
    fill = _named(ir, "g")[0].fill
    assert fill["angle"] == 90 and [s["color"] for s in fill["stops"]] == ["0B5CAD", "1E9E5A"]
    notes = [d for d in ir.diagnostics if "interpolated in" in d.message]
    assert len(notes) == 1 and notes[0].level == "info" and "oklab" in notes[0].message
    assert notes[0].source.endswith("div:nth-child(1)")  # the legacy-only gradient says nothing


def test_currentcolor_in_a_gradient_stop_is_the_elements_colour(tmp_path, ir_of):
    html = _slide(tmp_path, "<div data-name='g' style='position:absolute;left:40px;top:40px;width:200px;"
                            "height:60px;color:#0B5CAD;background:linear-gradient(90deg, currentcolor, white)'></div>")
    fill = _named(ir_of(html), "g")[0].fill
    assert fill["stops"][0]["color"] == "0B5CAD" and fill["stops"][1]["color"] == "FFFFFF"


SVG_PAGE = """<!doctype html><html><body style="margin:0">
<svg width="400" height="200" viewBox="0 0 400 200">
  <rect id="r" x="10" y="10" width="100" height="50" style="fill: oklch(0.95 0.04 150)"/>
  <rect id="m" x="150" y="10" width="100" height="50" style="fill: color-mix(in srgb, #1E9E5A 12%, white);
        filter: drop-shadow(0 2px 4px oklch(0.2 0.1 240 / 0.3))"/>
  <linearGradient id="lg"><stop offset="0" stop-color="oklch(0.95 0.04 150)"/><stop offset="1" stop-color="#1234"/></linearGradient>
  <rect id="g" x="10" y="100" width="100" height="50" fill="url(#lg)"/>
</svg></body></html>"""


def test_svg_paint_uses_the_normaliser_where_page_js_never_ran(browser, tmp_path):
    html = tmp_path / "svg.html"
    html.write_text(SVG_PAGE, encoding="utf-8")
    context = browser.new_context(viewport={"width": 1280, "height": 720}, device_scale_factor=1)
    page = context.new_page()
    try:
        page.goto(html.resolve().as_uri())
        assert page.evaluate("() => window.__engineColor === undefined")
        page.evaluate("() => document.querySelector('svg').setAttribute('data-engine-svg', 'svg#1')")
        expansion = expand_svg(page, ["svg#1"], Canvas(1280, 720), tmp_path)["svg#1"]
    finally:
        page.close()
        context.close()
    shapes = [e for e in expansion.elements if e.kind == "shape"]
    assert len(shapes) == 3, [d.message for d in expansion.diagnostics]
    oklch, mixed, gradient = shapes
    assert oklch.fill["type"] == "solid" and oklch.fill["color"] == "DCF7E1"
    assert mixed.fill["color"] == "E4F3EB"
    assert mixed.shadow["dx"] == 0 and mixed.shadow["dy"] == 2 and mixed.shadow["blur"] == 4
    assert mixed.shadow["alpha"] == pytest.approx(0.3, abs=1e-3)
    assert gradient.fill["type"] == "gradient"
    assert [s["color"] for s in gradient.fill["stops"]] == ["DCF7E1", "112233"]


def test_the_probe_code_moved_into_color_js():
    page_js = (config.ENGINE_DIR / "extract" / "page.js").read_text(encoding="utf-8")
    svg_js = (config.ENGINE_DIR / "extract" / "svg.js").read_text(encoding="utf-8")
    assert "probe.style.color" not in page_js
    assert "canvas" not in COLOR_JS.read_text(encoding="utf-8")
    assert svg_js.count("__engineColor") >= 3


# ================================================================================ lint and words


def _measured_rows() -> list[tuple[re.Pattern[str], str]]:
    """`MEASURED` in Python's regex engine: first match wins, as in the panel's `ruleOf`."""
    return [(re.compile(p, re.I if "i" in f else 0), rule) for p, f, rule in MEASURED]


def _rule_of(message: str) -> str:
    return next(rule for pattern, rule in _measured_rows() if pattern.search(message))


def test_the_colour_warning_has_its_own_words_not_dropped_element():
    message = "colour rgb(1, 2, 3) could not be read; the paint that used it is dropped"
    assert _rule_of(message) == "unreadable-colour"


def _lint_levels(body: str) -> dict[str, str]:
    from app.engine.verify.lint import lint

    report = lint(f"<!doctype html><html><body>{body}</body></html>")
    return {finding.rule: finding.level for finding in report.findings}


@pytest.mark.parametrize(
    "value",
    ["oklch(0.9 0.1 150)", "oklab(0.9 0.1 0.1)", "lab(90 5 5)", "lch(90 5 5)", "color(display-p3 1 0 0)",
     "color(rec2020 1 0 0)", "color(xyz 0.5 0.5 0.5)"],
)
def test_colour_gamut_is_info_for_every_wide_gamut_syntax(value: str):
    assert _lint_levels(f'<div style="background:{value}">x</div>').get("colour-gamut") == "info"


@pytest.mark.parametrize(
    "value",
    ["color-mix(in srgb, #1E9E5A 12%, white)", "color-mix(in oklab, #0B5CAD 10%, transparent)",
     "hsl(150 60% 94%)", "hwb(150 10% 20%)", "color(srgb 0.2 0.4 0.6)", "#1E9E5A1F"],
)
def test_exact_srgb_syntaxes_say_nothing(value: str):
    assert "colour-gamut" not in _lint_levels(f'<div style="background:{value}">x</div>')


# ======================================================================================= gradient text

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _context_slide(tmp_path: Path, body: str, css: str = "", name: str = "slide") -> Path:
    """A slide on the `test-16x9` master, the way the torture families and probes are authored."""
    html = _slide(tmp_path, body, css, name)
    (tmp_path / f"{name}.expect.json").write_text('{"master": "test-16x9", "layoutId": "layout-06"}',
                                                  encoding="utf-8")
    return html


def _extract(html: Path) -> IR:
    context = context_for(html)
    return extract_html(html, context.manifest, context.layout_id, context.assets_dir,
                        slide_id=html.stem, title=html.stem)


def _emit(ir: IR, html: Path, out: Path):
    from app.engine.emit import pptx as emitter
    from app.engine.reports import EmitOptions

    context = context_for(html)
    deck = out / f"{html.stem}.pptx"
    report = emitter.emit([ir], context.manifest, context.master, deck, EmitOptions())
    return deck, report


def _slide_shapes(deck: Path):
    from pptx import Presentation

    return Presentation(str(deck)).slides[0].shapes


def _run_rprs(deck: Path, text: str):
    """Every `a:rPr` of a run whose text is `text`, in the first slide of the deck."""
    found = []
    for shape in _slide_shapes(deck):
        for run in shape._element.iter(f"{A_NS}r"):
            t = run.find(f"{A_NS}t")
            if t is not None and t.text == text:
                found.append(run.find(f"{A_NS}rPr"))
    return found


def _overlap(a, b) -> float:
    return max(0.0, min(a.x + a.w, b.x + b.w) - max(a.x, b.x)) * max(0.0, min(a.y + a.h, b.y + b.h) - max(a.y, b.y))


def test_p10_gradient_text_is_a_run_fill_and_no_slab(tmp_path):
    html = PROBES / "text-misc.html"
    ir = _probe_ir("text-misc")
    grad = next(e for e in ir.elements if e.kind == "text" and "$12.4B" in "".join(r["text"] for r in _runs(e)))
    # no shape anywhere under the text: the gradient is not painted as a box
    covering = [s for s in ir.elements if s.kind == "shape" and _overlap(s.box, grad.box) > 1]
    assert not covering, [(s.name, s.box) for s in covering]
    fills = grad.extras["x-wpe-run-fills"]
    assert list(fills) == ["0/0/0"]
    assert fills["0/0/0"]["angle"] == 90
    assert [(s["pos"], s["color"]) for s in fills["0/0/0"]["stops"]] == [(0, "0B5CAD"), (1, "1E9E5A")]
    run = _runs(grad)[0]
    assert all(abs(a - b) <= 1 for a, b in zip(_hex(run["color"]), _hex("157D84"), strict=False)), run["color"]
    assert "_fill" not in run

    deck, report = _emit(ir, html, tmp_path)
    (rpr,) = _run_rprs(deck, "$12.4B")
    grad_fill = rpr.find(f"{A_NS}gradFill")
    assert grad_fill is not None and rpr.find(f"{A_NS}solidFill") is None
    stops = grad_fill.findall(f"{A_NS}gsLst/{A_NS}gs")
    assert [gs.get("pos") for gs in stops] == ["0", "100000"]
    assert [gs.find(f"{A_NS}srgbClr").get("val") for gs in stops] == ["0B5CAD", "1E9E5A"]
    assert grad_fill.find(f"{A_NS}lin").get("ang") == "0"
    # the fill sits where a fill belongs: before the typeface elements
    children = [child.tag.replace(A_NS, "") for child in rpr]
    assert children.index("gradFill") < children.index("latin")
    assert any("gradient text fill" in gap for gap in report.renderer_gaps)


GRADIENT_TEXT_CSS = """
.g{position:absolute;font-size:40px;font-weight:700;color:transparent;-webkit-background-clip:text;background-clip:text}
"""


def test_gradient_text_variants(tmp_path):
    html = _context_slide(tmp_path, """
<div class="g" data-name="solid" style="left:40px;top:40px;background-color:#0B5CAD">$12.4B</div>
<div data-name="span" style="position:absolute;left:40px;top:120px;font-size:20px;color:#1d2433">Growth of
  <span style="background:linear-gradient(90deg,#0B5CAD,#1E9E5A);-webkit-background-clip:text;background-clip:text;color:transparent">42%</span> this year</div>
<div class="g" data-name="three" style="left:40px;top:180px;background-image:linear-gradient(to right,#0B5CAD,#FFFFFF 50%,#1E9E5A)">Mid</div>
<div data-name="twin" style="position:absolute;left:40px;top:260px;font-size:30px;font-weight:700"><span
  style="background:linear-gradient(90deg,#0B5CAD,#1E9E5A);-webkit-background-clip:text;color:transparent">AB</span><span
  style="background:linear-gradient(90deg,#1E9E5A,#0B5CAD);-webkit-background-clip:text;color:transparent">CD</span></div>
<div class="g" data-name="fill-colour" style="left:40px;top:340px;color:#000;-webkit-text-fill-color:transparent;
  background-image:linear-gradient(90deg,#0B5CAD,#1E9E5A)">Fill</div>
""", GRADIENT_TEXT_CSS)
    ir = _extract(html)
    shapes = [e for e in ir.elements if e.kind == "shape"]
    assert not shapes, [(s.name, s.box, s.fill) for s in shapes]          # not one slab

    solid = _named(ir, "solid")[0]
    assert [r["color"] for r in _runs(solid)] == ["0B5CAD"] and not solid.extras.get("x-wpe-run-fills")

    span = _named(ir, "span")[0]
    runs = _runs(span)
    coloured = [r for r in runs if r["text"].strip() == "42%"]
    assert coloured and coloured[0]["color"] == "157D84"
    assert all(r["color"] == "1D2433" for r in runs if r["text"].strip() != "42%")
    assert len(span.extras["x-wpe-run-fills"]) == 1

    three = _named(ir, "three")[0]
    assert _runs(three)[0]["color"] == "FFFFFF"

    twin = _named(ir, "twin")[0]
    assert [r["text"] for r in _runs(twin)] == ["AB", "CD"]           # same mid-colour, two fills
    assert _runs(twin)[0]["color"] == _runs(twin)[1]["color"]
    assert sorted(twin.extras["x-wpe-run-fills"]) == ["0/0/0", "0/0/1"]

    assert _runs(_named(ir, "fill-colour")[0])[0]["color"] == "157D84"
    notes = [d for d in ir.diagnostics if d.message.startswith("gradient text:")]
    assert len(notes) == 4 and {d.level for d in notes} == {"info"}


def test_gradient_text_over_an_image_is_a_declared_raster(tmp_path):
    (tmp_path / "assets").mkdir()
    Image.new("RGB", (40, 30), (11, 92, 173)).save(tmp_path / "assets" / "photo.png")
    html = _context_slide(tmp_path, """
<div class="g" data-name="img" style="left:40px;top:40px;background-image:url(/api/projects/p/assets/photo.png)">$9</div>
""", GRADIENT_TEXT_CSS)
    ir = _extract(html)
    rasters = [e for e in ir.elements if e.kind == "raster"]
    assert len(rasters) == 1 and rasters[0].reason == "background-clip:text over an image"
    assert [d.level for d in ir.diagnostics if "background-clip:text over an image" in d.message] == ["warn"]


# ================================================================================== border joins


def _path_points(shape: Element) -> list[tuple[float, float]]:
    return [(round(seg[1], 2), round(seg[2], 2)) for seg in shape.geometry["path"] if seg[0] in ("M", "L")]


def test_p03_triangles_are_exact_join_polygons(tmp_path):
    html = PROBES / "pseudo.html"
    ir = _probe_ir("pseudo")
    tris = [e for e in ir.elements if e.kind == "shape" and e.geometry["type"] == "custom"
            and (e.fill or {}).get("color") == "0B5CAD" and e.box.w == 14]
    assert [_path_points(t) for t in tris] == [
        [(272, 343), (272, 327), (286, 335)],
        [(532, 343), (532, 327), (546, 335)],
    ]
    assert all(t.geometry["path"][-1] == ["Z"] and t.stroke is None for t in tris)
    assert [(t.box.x, t.box.y, t.box.w, t.box.h) for t in tris] == [(272, 327, 14, 16), (532, 327, 14, 16)]
    # no 14 x 16 rectangle anywhere: the old reading drew the triangle's bounding box
    assert not [e for e in ir.elements if e.kind == "shape" and e.geometry["type"] == "rect"
                and (round(e.box.w), round(e.box.h)) == (14, 16)]

    deck, _ = _emit(ir, html, tmp_path)
    shape = next(s for s in _slide_shapes(deck) if s.name == tris[0].name)
    (path,) = shape._element.iter(f"{A_NS}path")
    assert [child.tag.replace(A_NS, "") for child in path] == ["moveTo", "lnTo", "lnTo", "close"]
    assert path.get("fill") == "norm"


TRIANGLES = """
<div data-name="tri-right" style="position:absolute;left:40px;top:530px;width:0;height:0;
  border-left:14px solid #0B5CAD;border-top:8px solid transparent;border-bottom:8px solid transparent"></div>
<div data-name="tri-up" style="position:absolute;left:235px;top:530px;width:0;height:0;
  border-bottom:12px solid #1E9E5A;border-left:10px solid transparent;border-right:10px solid transparent"></div>
<div data-name="two-sides-zero" style="position:absolute;left:625px;top:530px;width:0;height:0;
  border-top:10px solid #D64545;border-left:10px solid #0B5CAD"></div>
<div data-name="two-sides-box" style="position:absolute;left:820px;top:530px;width:160px;height:110px;
  box-sizing:border-box;border-top:6px solid #D64545;border-left:10px solid #0B5CAD;background:#F7F7FA"></div>
<div data-name="tri-oklch" style="position:absolute;left:1015px;top:530px;width:0;height:0;
  border-left:14px solid #0B5CAD;border-top:8px solid oklch(0 0 0 / 0);border-bottom:8px solid oklch(0 0 0 / 0)"></div>
<div data-name="hair" style="position:absolute;left:40px;top:40px;width:100px;height:40px;
  border-top:0.4px solid #0B5CAD;border-left:6px solid #D64545"></div>
<p data-name="para" style="position:absolute;left:40px;top:300px;margin:0;font-size:16px">Text with a
  <span data-name="chip" style="display:inline-block;border-left:3px solid #D64545;padding:0 6px;background:#EEF">chip</span> inside</p>
"""


def test_css_border_triangles_and_joins(tmp_path):
    ir = _extract(_context_slide(tmp_path, TRIANGLES))

    def sides(name):
        return {e.name[len(name) + 1:].split(" ")[0]: e for e in ir.elements
                if e.kind == "shape" and (e.name or "").startswith(name + " ") and e.name.endswith(" border")}

    assert _path_points(sides("tri-right")["left"]) == [(40, 546), (40, 530), (54, 538)]
    assert _path_points(sides("tri-up")["bottom"]) == [(255, 542), (235, 542), (245, 530)]
    zero = sides("two-sides-zero")
    assert _path_points(zero["top"]) == [(625, 530), (635, 530), (635, 540)]
    assert _path_points(zero["left"]) == [(625, 540), (625, 530), (635, 540)]
    box = sides("two-sides-box")
    assert _path_points(box["top"]) == [(820, 530), (980, 530), (980, 536), (830, 536)]
    assert _path_points(box["left"]) == [(820, 640), (820, 530), (830, 536), (830, 640)]
    assert (box["top"].box.x, box["top"].box.y, box["top"].box.w, box["top"].box.h) == (820, 530, 160, 6)
    assert _path_points(sides("tri-oklch")["left"]) == [(1015, 546), (1015, 530), (1029, 538)]
    for shape in list(zero.values()) + list(box.values()):
        assert shape.fill["type"] == "solid" and shape.stroke is None and shape.geometry["fillRule"] == "nonzero"

    # Chromium rounds a border to whole device pixels (0.4px paints 1px), so the hairline is a 1 px join
    hair = sides("hair")
    assert hair["top"].geometry["type"] == "custom" and hair["top"].box.h == 1

    # an inline-block chip with one painted side gets that side as a 3 px shape
    chip = sides("chip")
    assert list(chip) == ["left"] and chip["left"].box.w == 3 and chip["left"].geometry["type"] == "rect"


def test_boxes_per_side_borders_keep_their_count_and_boxes(torture_dir):
    context = context_for(torture_dir / "boxes.html")
    ir = extract_html(torture_dir / "boxes.html", context.manifest, context.layout_id, context.assets_dir,
                      slide_id="boxes", title="boxes")
    shapes = [e for e in ir.elements if e.kind == "shape" and (e.name or "").startswith("per-side-borders")]
    assert len(shapes) == 5
    by_side = {(e.name.split(" ")[1] if " " in e.name else "box"): e for e in shapes}
    assert by_side["top"].geometry["type"] == "custom" and by_side["left"].geometry["type"] == "custom"
    assert by_side["right"].geometry["type"] == "line" and by_side["bottom"].geometry["type"] == "line"
    top, left = by_side["top"].box, by_side["left"].box
    assert (top.x, top.y, top.w, top.h) == (430, 50, 160, 4)
    assert (left.x, left.y, left.w, left.h) == (430, 50, 1, 110)


def test_a_hairline_join_polygon_survives_keep():
    """`_keep` drops sub-half-pixel geometry except border sides, and a join polygon is a side."""
    side = {"kind": "shape", "box": {"x": 0, "y": 0, "w": 100, "h": 0.4},
            "geometry": {"type": "custom", "path": [["M", 0, 0], ["L", 100, 0], ["L", 99, 0.4], ["Z"]]}}
    assert html_extract._keep(side)
    assert not html_extract._keep({**side, "box": {"x": 0, "y": 0, "w": 0.4, "h": 0.4}})


# ================================================================================= pseudo-elements


def _pseudo(ir: IR, suffix: str = "") -> list[Element]:
    """Elements rebuilt from a `::before`/`::after` (their source path says so)."""
    return [e for e in ir.elements if "::" in (e.source or {}).get("path", "") and
            (e.source or {}).get("path", "").endswith(suffix)]


def _texts_of(element: Element) -> list[str]:
    return ["".join(r["text"] for line in p["lines"] for r in line["runs"]) for p in element.paragraphs or []]


def _box(e: Element) -> tuple[float, float, float, float]:
    return (round(e.box.x, 2), round(e.box.y, 2), round(e.box.w, 2), round(e.box.h, 2))


def test_p03_pseudo_elements_are_rebuilt():
    ir = _probe_ir("pseudo")
    notes = [d for d in ir.diagnostics if d.source == "pseudo-elements"]
    assert [(d.level, d.message) for d in notes] == [("info", "13 pseudo-elements rebuilt as shapes and text, 0 dropped")]
    assert not [d for d in ir.diagnostics if "::" in d.message]

    # A: the timeline axis and four numbered milestone dots
    (axis,) = [e for e in _pseudo(ir, "div.tl:nth-child(2)::before")]
    assert axis.kind == "shape" and _box(axis) == (40, 140, 1200, 3) and axis.fill["color"] == "0B5CAD"
    dots = [e for e in _pseudo(ir, "::before") if e.kind == "shape" and e.geometry["type"] == "roundRect"
            and e.box.w == 26]
    assert [_box(d) for d in dots] == [(x, 128, 26, 26) for x in (40, 340, 640, 940)]
    assert all(d.geometry["radius"] == {"tl": 13, "tr": 13, "br": 13, "bl": 13} for d in dots)
    numbers = [e for e in _pseudo(ir, "::before") if e.kind == "text" and e.box.y < 200]
    assert [_texts_of(n) for n in numbers] == [["1"], ["2"], ["3"], ["4"]]
    for number in numbers:
        (run,) = _runs(number)
        assert run["color"] == "FFFFFF" and run["sizePx"] == 13 and run["weight"] == 700

    # B: the `.box.arrow::after` triangle. Its host has a 1 px border, so the pseudo's containing
    # block is the padding box: right:-24px puts its right edge at 779 + 24 = 803, top:26px at 327.
    (arrow,) = [e for e in _pseudo(ir, "div.box:nth-child(5)::after")]
    assert _path_points(arrow) == [(789, 343), (789, 327), (803, 335)] and arrow.fill["color"] == "D64545"

    # C: the chevron tips, 28 x 56 triangles at the right edge of each 210 px host
    tips = [e for e in _pseudo(ir, "::after") if (e.fill or {}).get("color") == "1E9E5A"]
    assert [_path_points(t) for t in tips] == [[(x, 566), (x, 510), (x + 28, 538)] for x in (250, 500, 750)]

    # D: `.e2::before` is the 6 px colour edge spanning the host's padding box
    (edge,) = [e for e in _pseudo(ir, "div.edge:nth-child(11)::before")]
    assert _box(edge) == (341, 601, 6, 58) and edge.fill["color"] == "D64545"

    # E: the bullets stay ONE text element of three paragraphs, their dots three shapes after it
    (bullets,) = [e for e in ir.elements if e.kind == "text" and "Bullet dot via ::before" in _texts_of(e)[0]]
    assert _texts_of(bullets) == ["Bullet dot via ::before", "Second bullet", "Third bullet"]
    bullet_dots = [e for e in _pseudo(ir, "::before") if e.kind == "shape" and e.box.w == 7]
    assert [(d.box.x, d.box.w, d.box.h) for d in bullet_dots] == [(1010, 7, 7)] * 3
    assert all(d.z > bullets.z for d in bullet_dots)                     # painted after the text


def test_the_pre_pass_reports_what_it_rebuilt(page, tmp_path):
    """§16 #12: `{materialised, dropped}` with `<cssPath>::before|::after` paths."""
    html = _context_slide(tmp_path, """
<div id="a" style="position:absolute;left:40px;top:40px;width:100px;height:40px"></div>
<table style="position:absolute;left:40px;top:200px"><tr class="r"><td>1</td></tr></table>
""", "#a::before{content:'x';color:#0B5CAD} #a::after{content:'';} .r::before{content:'row'}")
    context = context_for(html)
    from app.engine.extract.html import PAGE_JS, prepare_page

    prepare_page(page, html, context.manifest, context.assets_dir)
    page.add_script_tag(path=str(COLOR_JS))
    page.add_script_tag(path=str(PAGE_JS))
    result = page.evaluate("() => window.__engineMaterialisePseudo()")
    assert result == {"materialised": ["div#a::before"],
                      "dropped": ["body:nth-child(2) > table:nth-child(2) > tbody:nth-child(1) > tr.r:nth-child(1)::before"]}
    warnings = page.evaluate("() => window.__enginePrepass")
    assert [w["message"] for w in warnings] == ["a ::before on a table row is not drawn by the table export; dropped"]


def _family_irs(names: list[str], materialise: bool, monkeypatch) -> dict[str, IR]:
    monkeypatch.setattr(html_extract, "MATERIALISE_PSEUDO", materialise)
    irs = {}
    for name in names:
        html = config.TORTURE_FIXTURE / f"{name}.html"
        irs[name] = _extract(html)
    return irs


def _geometry_of(ir: IR) -> list[tuple]:
    """Every non-pseudo element as (kind, path, box, line boxes and texts, everything else it paints),
    boxes to 3 decimals. "Everything else" (E3 review m6) is the name, the fill, the geometry type and
    radius, the flips and rotation, each paragraph's bullet, the clip, the stroke and the run styles:
    a rebuild that restyles an original element without moving it must show up here too."""
    import json

    out = []
    for e in ir.elements:
        path = (e.source or {}).get("path", "")
        if "::" in path:
            continue
        lines = tuple(
            (round(ln["box"]["x"], 3), round(ln["box"]["y"], 3), round(ln["box"]["w"], 3), round(ln["box"]["h"], 3),
             "".join(r["text"] for r in ln["runs"]))
            for p in e.paragraphs or [] for ln in p["lines"]
        )
        geometry = getattr(e, "geometry", None) or {}
        paint = json.dumps({
            "name": e.name, "fill": getattr(e, "fill", None), "stroke": getattr(e, "stroke", None),
            "geometry": geometry.get("type"), "radius": geometry.get("radius"),
            "flip": (getattr(e, "flipH", None), getattr(e, "flipV", None)), "rotation": getattr(e, "rotation", None),
            "bullets": [p.get("bullet") for p in e.paragraphs or []],
            "clip": None if e.clip is None else [round(v, 3) for v in (e.clip.x, e.clip.y, e.clip.w, e.clip.h)],
            "runs": [{k: v for k, v in r.items() if k != "text"} for p in e.paragraphs or [] for ln in p["lines"]
                     for r in ln["runs"]],
        }, sort_keys=True, default=str)
        out.append((e.kind, path, (round(e.box.x, 3), round(e.box.y, 3), round(e.box.w, 3), round(e.box.h, 3)),
                    lines, paint))
    return out


def test_layout_invariance_with_and_without_the_pre_pass(monkeypatch):
    """Every family but `paint` (`pseudo` is probe p03), pre-pass on vs off: every non-pseudo box,
    line box and source path — HTML and SVG-derived — agrees to 3 decimals. The rebuild moves nothing;
    it only adds."""
    names = [n for n in sorted(p.stem for p in config.TORTURE_FIXTURE.glob("*.html")) if n != "paint"]
    assert "pseudo" in names
    on = _family_irs(names, True, monkeypatch)
    off = _family_irs(names, False, monkeypatch)
    worst = 0.0
    for name in names:
        a, b = _geometry_of(on[name]), _geometry_of(off[name])
        assert [x[:2] for x in a] == [x[:2] for x in b], f"{name}: elements or source paths differ"
        for x, y in zip(a, b, strict=False):
            deltas = [abs(p - q) for p, q in zip(x[2], y[2], strict=False)]
            deltas += [abs(p - q) for lx, ly in zip(x[3], y[3], strict=False) for p, q in zip(lx[:4], ly[:4], strict=False)]
            worst = max([worst] + deltas)
            assert [ln[4] for ln in x[3]] == [ln[4] for ln in y[3]], f"{name} {x[1]}: text differs"
            assert x[4] == y[4], f"{name} {x[1]}: name or paint differs"
    assert worst == 0.0, f"max |delta| {worst}"
    assert len(_pseudo(on["pseudo"])) > 0 and not _pseudo(off["pseudo"])


def test_an_author_layer_that_wins_the_cascade_is_reverted_and_reported(tmp_path):
    html = _context_slide(tmp_path, """
<div id="h" style="position:absolute;left:40px;top:40px;font-size:20px">Host text</div>
""", "@layer a { #h::before { content: 'x ' !important; color: #D64545 } } #h::before { content: 'y ' }")
    ir = _extract(html)
    assert not _pseudo(ir)
    (warn,) = [d for d in ir.diagnostics if "::before" in d.message]
    assert warn.level == "warn" and warn.message.endswith("dropped") and warn.source == "div#h::before"
    assert "1 dropped" in next(d.message for d in ir.diagnostics if d.source == "pseudo-elements")


def test_a_rebuild_that_repaints_a_sibling_is_reverted(tmp_path):
    """`.row > :first-child` stops matching the span once a child is inserted before it: nothing
    moves, but the span would repaint — check (iii) catches it and the copy is taken back."""
    html = _context_slide(tmp_path, """
<div class="row" style="position:absolute;left:40px;top:40px;font-size:20px"><span>First</span> rest</div>
""", ".row > :first-child { color: #0B5CAD } .row::before { content: '→ ' }")
    ir = _extract(html)
    assert not _pseudo(ir)
    assert [d.level for d in ir.diagnostics if "could not be rebuilt" in d.message] == ["warn"]
    runs = _runs(next(e for e in ir.elements if e.kind == "text"))
    assert [r["text"] for r in runs] == ["First", " rest"] and runs[0]["color"] == "0B5CAD"


COUNTERS = """
<ol class="outer" style="position:absolute;left:40px;top:40px;list-style:none;margin:0;padding:0;font-size:14px">
  <li>A<ol style="list-style:none;padding:0"><li>A1</li><li>A2</li></ol></li>
  <li>B</li>
</ol>
<div class="roman" style="position:absolute;left:400px;top:40px;font-size:14px">
  <div>i</div><div>ii</div><div>iii</div><div>iv</div>
</div>
<div class="greek" style="position:absolute;left:600px;top:40px;font-size:14px"><div>g</div></div>
<div class="cjk" style="position:absolute;left:600px;top:140px;font-size:14px"><div>c</div></div>
<div class="set" style="position:absolute;left:800px;top:40px;font-size:14px"><div>x</div><div class="jump">y</div><div>z</div></div>
"""
COUNTERS_CSS = """
ol.outer, ol.outer ol { counter-reset: item }
ol.outer li { counter-increment: item }
ol.outer li::before { content: counters(item, ".") " "; }
.roman { counter-reset: r } .roman div { counter-increment: r } .roman div::before { content: counter(r, upper-roman) ". " }
.greek div::before { content: counter(g, lower-greek) ") " } .greek { counter-reset: g 2 }
.cjk div::before { content: counter(c, cjk-decimal) ") " } .cjk { counter-reset: c 2 }
.set { counter-reset: s } .set div { counter-increment: s } .set .jump { counter-set: s 10 }
.set div::before { content: counter(s, decimal-leading-zero) " " }
"""


def test_counters_follow_css_lists_scoping(tmp_path):
    ir = _extract(_context_slide(tmp_path, COUNTERS, COUNTERS_CSS))
    texts = [t for e in ir.elements if e.kind == "text" for t in _texts_of(e)]
    assert {"1 A", "1.1 A1", "1.2 A2", "2 B"} <= set(texts), texts
    assert {"I. i", "II. ii", "III. iii", "IV. iv"} <= set(texts), texts
    assert "β) g" in texts
    assert {"01 x", "11 y", "12 z"} <= set(texts), texts          # set, then increment
    # An unsupported style is written in decimal and reported. Here "2" is narrower than "二", so
    # the rebuild would move the text after it: it is reverted, and that is reported too.
    (warn,) = [d for d in ir.diagnostics if d.message.startswith("counter style")]
    assert warn.level == "warn" and "cjk-decimal" in warn.message and warn.message.endswith("approximated")
    assert "2) c" in texts or [d for d in ir.diagnostics if d.source == warn.source and "could not be rebuilt" in d.message]


def test_attr_and_quotes_and_nesting(tmp_path):
    ir = _extract(_context_slide(tmp_path, """
<div class="tag" data-x="Q3" style="position:absolute;left:40px;top:40px;font-size:16px">Revenue</div>
<div class="q" style="position:absolute;left:40px;top:100px;font-size:16px">Outer <span class="q">inner</span> end</div>
<div class="missing" style="position:absolute;left:40px;top:160px;font-size:16px">none</div>
""", ".tag::before{content:'→ ' attr(data-x) ' '} .q::before{content:open-quote} .q::after{content:close-quote}"
       " .missing::before{content:'[' attr(data-nothing) ']'}"))
    texts = [t for e in ir.elements if e.kind == "text" for t in _texts_of(e)]
    assert "→ Q3 Revenue" in texts
    assert "“Outer ‘inner’ end”" in texts
    assert "[]none" in texts


def test_empty_pseudos_leave_nothing_and_move_nothing(tmp_path):
    ir = _extract(_context_slide(tmp_path, """
<div class="none" style="position:absolute;left:40px;top:40px">Plain</div>
<div class="cf" style="position:absolute;left:40px;top:100px;width:200px"><div style="float:left;width:60px;height:30px;background:#E8F1FB"></div></div>
""", ".none::before{content:none} .cf::after{content:'';display:table;clear:both}"))
    assert not _pseudo(ir)
    assert not [d for d in ir.diagnostics if d.source == "pseudo-elements"]


def test_a_display_contents_host_is_never_silent(tmp_path):
    ir = _extract(_context_slide(tmp_path, """
<div style="position:absolute;left:40px;top:40px;display:flex;gap:8px;font-size:16px"><div class="dc"><div>Inside</div></div><div>After</div></div>
""", ".dc{display:contents} .dc::before{content:'DC';color:#D64545}"))
    rebuilt = [e for e in ir.elements if any("DC" in t for t in _texts_of(e))]
    dropped = [d for d in ir.diagnostics if "::before" in d.message]
    assert rebuilt or dropped
    if rebuilt:
        assert _runs(rebuilt[0])[0]["color"] == "D64545"


def test_list_style_none_draws_no_bullet_and_keeps_the_pseudo_tick(tmp_path):
    ir = _extract(_context_slide(tmp_path, """
<ul style="position:absolute;left:40px;top:40px;list-style:none;margin:0;padding:0;font-size:16px">
  <li>First</li><li>Second</li>
</ul>
<ol style="position:absolute;left:400px;top:40px;list-style:none;font-size:16px"><li>Numbered</li></ol>
""", "ul li::before{content:'✓ ';color:#1E9E5A}"))
    (ticks,) = [e for e in ir.elements if e.kind == "text" and "First" in "".join(_texts_of(e))]
    assert _texts_of(ticks) == ["✓ First", "✓ Second"]
    assert all(p["bullet"] is None for p in ticks.paragraphs)
    (numbered,) = [e for e in ir.elements if e.kind == "text" and "Numbered" in "".join(_texts_of(e))]
    assert numbered.paragraphs[0]["bullet"] is None


def _cell_texts(table: Element) -> list[str]:
    return ["".join(r["text"] for p in c["paragraphs"] for ln in p["lines"] for r in ln["runs"]) for c in table.cells]


def _pseudo_warnings(ir: IR) -> list:
    return [d for d in ir.diagnostics if d.level != "info" and ("::before" in d.message or "::after" in d.message
                                                                or "::before" in (d.source or "")
                                                                or "::after" in (d.source or ""))]


def test_a_cell_pseudo_joins_the_cell_text_and_a_paint_only_one_is_drawn(tmp_path):
    """Plan §16 #22: cells are walked (WP-A's partition), so a paint-only `::before` on a cell is
    rebuilt like any other and drawn by the cell walk — the absolute dot as a shape over the table,
    where the pre-pass used to drop it with a 'table cell' warn. A text one still joins the cell."""
    ir = _extract(_context_slide(tmp_path, """
<table style="position:absolute;left:40px;top:40px;border-collapse:collapse;font-size:14px">
  <tr><td class="money">125</td><td class="dot" style="position:relative;padding-left:20px">Live</td></tr>
</table>
""", "td.money::before{content:'$'} td.dot::before{content:'';position:absolute;left:4px;top:6px;width:8px;"
     "height:8px;border-radius:50%;background:#1E9E5A}"))
    (table,) = [e for e in ir.elements if e.kind == "table"]
    assert _cell_texts(table) == ["$125", "Live"]
    (dot,) = _pseudo(ir, "td.dot:nth-child(2)::before")
    assert dot.kind == "shape" and dot.fill["color"] == "1E9E5A" and dot.z > table.z
    cell = next(c for c in table.cells if c["c"] == 1)
    # the dot sits in the cell's left padding, 4 px in and 6 px down from the cell's border box
    assert dot.box.w == pytest.approx(8, abs=0.5) and dot.box.h == pytest.approx(8, abs=0.5)
    assert dot.extras["x-wpa"]["cell"] == [cell["r"], cell["c"]]
    assert not _pseudo_warnings(ir), _pseudo_warnings(ir)
    assert "2 pseudo-elements rebuilt as shapes and text, 0 dropped" in [d.message for d in ir.diagnostics]


def test_a_cell_swatch_and_badge_are_drawn_once_and_an_undrawable_one_is_reported(tmp_path):
    """Plan §16 #22 (A's measurement on its head): `td.ok::before`, an inline-block swatch, was dropped
    by the pre-pass although the cell text already left its room; `td.tag::after`, a "NEW" badge, was
    drawn by the cell walk and still warned 'dropped'. Both are drawn now, each exactly once — one
    shape, and for the badge one text box, never also inside the cell's own paragraph — and counted
    once as rebuilt. A copy the export truly cannot draw (an outline is not a shape the IR has) is
    still taken back and reported by E's sweep, and counted as dropped."""
    ir = _extract(_context_slide(tmp_path, """
<table style="position:absolute;left:40px;top:40px;border-collapse:collapse;font-size:14px">
  <tr><td class="ok">Approved</td><td class="tag">Pricing</td><td class="o">Outline</td></tr>
</table>
""", "td{padding:6px 10px} "
     "td.ok::before{content:'';display:inline-block;width:8px;height:8px;background:green;margin-right:6px} "
     "td.tag::after{content:'NEW';background:#FFE600;color:#1d2433;font-size:10px;font-weight:700;"
     "padding:1px 6px;border-radius:8px;margin-left:6px} "
     "td.o::before{content:'';display:inline-block;width:10px;height:10px;outline:2px solid red;margin-right:6px}"))
    (table,) = [e for e in ir.elements if e.kind == "table"]
    assert _cell_texts(table) == ["Approved", "Pricing", "Outline"]
    cells = table.extras["x-wpa"]["cells"]

    (swatch,) = _pseudo(ir, "td.ok:nth-child(1)::before")
    assert swatch.kind == "shape" and swatch.fill["color"] == "008000" and swatch.extras["x-wpa"]["role"] == "chip"
    assert (swatch.box.w, swatch.box.h) == (pytest.approx(8, abs=0.5), pytest.approx(8, abs=0.5))
    # the cell's text starts where the browser drew it: after the swatch and its 6 px margin
    assert cells["0:0"]["indentPx"] == [pytest.approx(14, abs=0.5)] and cells["0:0"]["overlays"] == 1
    assert cells["0:0"]["runBoxes"][0][0][0][0] == pytest.approx(swatch.box.x + 14, abs=0.5)

    badge = _pseudo(ir, "td.tag:nth-child(2)::after")
    assert sorted(e.kind for e in badge) == ["shape", "text"]
    (chip,) = [e for e in badge if e.kind == "shape"]
    (words,) = [e for e in badge if e.kind == "text"]
    assert chip.fill["color"] == "FFE600" and _texts_of(words) == ["NEW"]
    assert sum(_texts_of(e).count("NEW") for e in ir.elements if e.kind == "text") == 1
    assert cells["0:1"]["overlays"] == 2

    # the outline swatch: rebuilt, not drawn by the walk, so taken back and said so — never silent
    assert not _pseudo(ir, "td.o:nth-child(3)::before") and cells["0:2"]["overlays"] == 0
    (warn,) = _pseudo_warnings(ir)
    assert warn.level == "warn" and warn.source.endswith("td.o:nth-child(3)::before")
    assert warn.message == "a ::before was rebuilt but the export cannot draw it in its place; dropped"
    assert not [d for d in ir.diagnostics if "table cell" in d.message]
    assert "2 pseudo-elements rebuilt as shapes and text, 1 dropped" in [d.message for d in ir.diagnostics]


def test_image_content_is_reported(tmp_path):
    (tmp_path / "assets").mkdir()
    Image.new("RGB", (16, 16), (11, 92, 173)).save(tmp_path / "assets" / "icon.png")
    ir = _extract(_context_slide(tmp_path, """
<div class="i" style="position:absolute;left:40px;top:40px;font-size:16px">Label</div>
""", ".i::before{content:url(/api/projects/p/assets/icon.png)}"))
    (warn,) = [d for d in ir.diagnostics if "image content" in d.message]
    assert warn.level == "warn" and warn.message == "::before image content dropped"


def test_negative_z_pseudos_paint_under_the_text_and_the_rest_over_it(tmp_path):
    ir = _extract(_context_slide(tmp_path, """
<div class="p" style="position:absolute;left:40px;top:40px;width:300px;font-size:16px;z-index:0">
  <div class="item" style="position:relative">One</div><div class="item" style="position:relative">Two</div>
</div>
""", ".item::before{content:'';position:absolute;left:-10px;top:0;width:4px;height:16px;background:#D64545;z-index:-1}"
     " .item::after{content:'';position:absolute;right:0;top:0;width:4px;height:16px;background:#0B5CAD}"))
    (text,) = [e for e in ir.elements if e.kind == "text"]
    assert _texts_of(text) == ["One", "Two"]                    # one element: the rule holds both ways
    below = [e for e in _pseudo(ir, "::before")]
    above = [e for e in _pseudo(ir, "::after")]
    assert len(below) == 2 and len(above) == 2
    assert all(e.z < text.z for e in below) and all(e.z > text.z for e in above)


def test_pseudo_extraction_is_deterministic():
    first, second = _probe_ir("pseudo"), _probe_ir("pseudo")
    assert first.dumps() == second.dumps()


def test_the_reference_render_never_sees_the_pre_pass(tmp_path, page):
    """`render_reference` never runs the pre-pass (it has no hook for it), and extracting on a page
    first changes nothing about the reference rendered next: two renders are byte-identical."""
    import inspect

    from app.engine.extract.html import render_reference

    assert "__engineMaterialisePseudo" not in inspect.getsource(render_reference)
    html = PROBES / "pseudo.html"
    context = context_for(html)
    fonts = context.manifest.fonts        # G-1 (D5): the reference takes the master's fonts, required
    first = render_reference(html, context.layout_png, tmp_path / "a.png", assets_dir=context.assets_dir,
                             fonts=fonts)
    extract_html(html, context.manifest, context.layout_id, context.assets_dir, page=page)
    second = render_reference(html, context.layout_png, tmp_path / "b.png", assets_dir=context.assets_dir,
                              fonts=fonts, page=page)
    assert first.read_bytes() == second.read_bytes()


def test_the_pseudo_warnings_have_their_own_words():
    for message in ("a ::before could not be rebuilt without moving the layout; dropped",
                    "a ::after on a table row is not drawn by the table export; dropped",
                    "a ::before was rebuilt but the export cannot draw it in its place; dropped",
                    "::after image content dropped"):
        assert _rule_of(message) == "pseudo-dropped", message
    assert _rule_of("counter style lower-greek is not supported; decimal used, approximated") == "approximated"
    assert _rule_of("gradient text: the gradient is stretched over the text box, approximated") == "approximated"


def test_the_three_lint_rules_have_their_levels():
    levels = _lint_levels(
        '<style>.x::before{content:url(/api/projects/p/assets/a.png)} .y:after{content:"a" url(b.png)}'
        ' .z::before{content:"plain"}</style>'
        '<div class="x" style="background:oklch(0.9 0.1 150)">x</div>'
        '<div class="y" style="background:linear-gradient(red,blue);-webkit-background-clip:text;color:transparent">y</div>'
        '<div class="z">z</div>'
    )
    assert levels["colour-gamut"] == "info"
    assert levels["gradient-text"] == "info"
    assert levels["pseudo-content-image"] == "warn"


def test_pseudo_content_image_is_reported_once_per_rule_whatever_it_matches():
    from app.engine.verify.lint import lint

    report = lint('<!doctype html><html><head><style>.x::after{content:url(/api/projects/p/assets/a.png)}</style>'
                  '</head><body><div class="x">1</div><div class="x">2</div><div class="x">3</div></body></html>')
    findings = [f for f in report.findings if f.rule == "pseudo-content-image"]
    assert len(findings) == 1 and "::after" in findings[0].message


def test_a_repaint_behind_a_transition_is_still_caught(tmp_path):
    """At t = 0 a transition still computes to the old value, so a sibling the rebuild flips would pass
    check (iii) and be read later in its new colour. Transitions are frozen while the rebuild runs."""
    html = _context_slide(tmp_path, """
<div class="row" style="position:absolute;left:40px;top:40px;font-size:20px"><span>First</span> rest</div>
""", ".row > span { transition: color 5s linear } .row > :first-child { color: #0B5CAD }"
     " .row::before { content: '→ ' }")
    ir = _extract(html)
    assert not _pseudo(ir)
    assert [d.level for d in ir.diagnostics if "could not be rebuilt" in d.message] == ["warn"]
    assert _runs(next(e for e in ir.elements if e.kind == "text"))[0]["color"] == "0B5CAD"


def test_a_shrink_to_fit_pseudo_is_rebuilt_despite_sub_pixel_serialisation(tmp_path):
    """An absolute bullet's used width is 1/64 px units printed to 6 digits; setting it back lands one
    unit lower. Six real slides lost every bullet to that before `sameComputed` (2026-09-25)."""
    html = _context_slide(tmp_path, """
<div style="position:absolute;left:40px;top:40px;width:400px;font-size:13px">
  <div class="b" style="position:relative;padding-left:14px">Reach full operating scale</div>
  <div class="b" style="position:relative;padding-left:14px">Establish public sector leadership</div>
</div>
""", ".b::before { content: '•'; position: absolute; left: 0; color: #747480 }")
    ir = _extract(html)
    assert len(_pseudo(ir)) == 2 and not [d for d in ir.diagnostics if "could not be rebuilt" in d.message]


def test_one_bad_rebuild_does_not_cost_the_others(tmp_path):
    """The batch fails as a whole, then each pseudo is rebuilt alone: only the one that repaints a
    sibling is dropped, the unrelated bullet dot is kept."""
    html = _context_slide(tmp_path, """
<div class="row" style="position:absolute;left:40px;top:40px;font-size:20px"><span>First</span> rest</div>
<div class="dot" style="position:absolute;left:40px;top:120px;padding-left:14px;font-size:16px">Kept</div>
""", ".row > :first-child { color: #0B5CAD } .row::before { content: '→ ' }"
     " .dot::before { content: ''; position: absolute; left: 0; top: 6px; width: 7px; height: 7px;"
     " border-radius: 50%; background: #D64545 }")
    ir = _extract(html)
    kept = _pseudo(ir)
    assert [e.source["path"].rsplit(" > ", 1)[-1] for e in kept] == ["div.dot:nth-child(2)::before"]
    dropped = [d for d in ir.diagnostics if "could not be rebuilt" in d.message]
    assert [d.source.rsplit(" > ", 1)[-1] for d in dropped] == ["div.row:nth-child(1)::before"]


# ================================================================ the E3 review's findings (max pass)
#
# `fidelity-reports/e3-review.md`: two blockers, four majors, six minors against the pre-pass as it
# stood at f0955b3. Each case below is the review's own, reduced; the test says what the browser draws.

#: `page.js`'s `cssPath`, for reading the page after an extraction (a copy is its host's path + `::kind`).
_CSS_PATH_JS = """
  function cssPath(el) {
    var parts = [];
    for (var node = el; node && node.nodeType === 1 && node !== document.documentElement; node = node.parentElement) {
      var piece = node.tagName.toLowerCase();
      if (node.id) { piece += "#" + node.id; parts.unshift(piece); break; }
      if (node.classList.length) piece += "." + node.classList[0];
      var index = 1, sibling = node;
      while ((sibling = sibling.previousElementSibling)) if (!sibling.hasAttribute("data-engine-pseudo")) index++;
      parts.unshift(piece + ":nth-child(" + index + ")");
    }
    return parts.join(" > ");
  }
"""

B1_NESTED = ("""
<p class="t" style="position:absolute;left:40px;top:40px;width:700px;font-size:24px;margin:0">Revenue grew <span class="hl">42%</span> this year and <span class="new">Beta</span> launched</p>
<div class="t2" style="position:absolute;left:40px;top:140px;width:700px;font-size:20px">Plain host <b class="dot">bold</b> text</div>
<div class="c1" style="position:absolute;left:40px;top:240px;width:500px;font-size:20px">Run with <span class="hl">mark</span> text<div>Block sibling</div></div>
<p class="c2" style="position:absolute;left:40px;top:340px;width:500px;font-size:20px;margin:0">Inline <span class="bb">block-in-inline</span> end</p>
<div class="c" data-chart='{"type":"bar","data":{"labels":["A","B"],"datasets":[{"label":"S","data":[1,2]}]}}' style="position:absolute;left:700px;top:300px;width:400px;height:240px"></div>
<div class="h" style="position:absolute;left:40px;top:600px;font-size:18px;visibility:hidden">Hidden host</div>
""", """
.hl{position:relative}
.hl::after{content:"";position:absolute;left:0;right:0;bottom:2px;height:8px;background:rgba(255,200,0,.5)}
.new{position:relative}
.new::after{content:"NEW";position:absolute;top:-10px;right:-34px;font-size:9px;background:#D64545;color:#fff;padding:1px 3px}
.dot{position:relative}
.dot::before{content:"";position:absolute;left:-10px;top:8px;width:6px;height:6px;border-radius:50%;background:#1E9E5A}
.bb::before{content:'';display:block;height:3px;width:40px;background:#D64545}
.c::before{content:'Revenue ($M)';position:absolute;left:0;top:-24px;font-size:14px;color:#0B5CAD}
.h::before{content:'Shown';visibility:visible;color:#D64545}
""")


def test_every_rebuilt_pseudo_is_drawn_or_reported(page, tmp_path):
    """B1: a copy the walk never reads was counted as rebuilt. Now every pseudo-element ends in one
    of three states — drawn (an element with its path, or its text in a run), reverted before the walk
    with a warn, or taken back after it with a warn — and only the drawn ones are counted as rebuilt.
    Nested out-of-flow copies (a highlight under a word, a badge on a span, a dot on a `<b>`, the
    same inside an anonymous run) and a block-in-inline bar are drawn; a chart host's label and a
    visible pseudo on a hidden host are reported."""
    html = _context_slide(tmp_path, *B1_NESTED)
    context = context_for(html)
    ir = extract_html(html, context.manifest, context.layout_id, context.assets_dir, slide_id="b1", page=page)
    left = page.evaluate("() => {" + _CSS_PATH_JS + """
      return [...document.querySelectorAll('[data-engine-pseudo]')].map(c =>
        [cssPath(c.parentElement) + '::' + c.getAttribute('data-engine-pseudo'), c.textContent]); }""")
    undrawn = page.evaluate("() => window.__enginePseudoUndrawn")
    paths = {(e.source or {}).get("path") for e in ir.elements}
    words = {text for e in ir.elements if e.kind == "text" for text in _texts_of(e)}
    for path, text in left:
        assert path in paths or (text and any(text in w for w in words)), f"{path} is neither drawn nor reported"
    warned = {d.source for d in ir.diagnostics if d.level == "warn" and _rule_of(d.message) == "pseudo-dropped"}
    assert set(undrawn) <= warned and len(undrawn) == 2
    assert {u.rsplit(" > ", 1)[-1] for u in undrawn} == {"div.c:nth-child(5)::before", "div.h:nth-child(6)::before"}
    note = next(d.message for d in ir.diagnostics if d.source == "pseudo-elements")
    assert note == f"{len(left)} pseudo-elements rebuilt as shapes and text, 2 dropped" and len(left) == 5

    shapes = {e.source["path"].rsplit(" > ", 1)[-1]: e for e in _pseudo(ir) if e.kind == "shape"}
    assert shapes["span.hl:nth-child(1)::after"].fill["color"] == "FFC800"
    assert shapes["span.new:nth-child(2)::after"].fill["color"] == "D64545"
    assert shapes["b.dot:nth-child(1)::before"].geometry["type"] == "roundRect"
    assert _box(shapes["span.bb:nth-child(1)::before"])[2:] == (40, 3)
    (badge,) = [e for e in _pseudo(ir, "span.new:nth-child(2)::after") if e.kind == "text"]
    assert _runs(badge)[0]["text"] == "NEW" and _runs(badge)[0]["color"] == "FFFFFF"
    # the highlight is positioned, so it paints after the words it marks (CSS paint order)
    (sentence,) = [e for e in ir.elements if e.kind == "text" and "Revenue grew" in _texts_of(e)[0]]
    assert shapes["span.hl:nth-child(1)::after"].z > sentence.z
    # the chart host's label is not in the chart, and the live page shows it again (the reference does)
    assert page.evaluate("() => getComputedStyle(document.querySelector('.c'), '::before').content") != "none"


#: B2: the structural rules the review found the old 22-property signature could not see. Each copy
#: flips a style that moves nothing: a radius, a transform, a z-index, `:empty`, a `::marker` colour.
B2_CASES = {
    "radius": ("""
<div class="seg" style="position:absolute;left:40px;top:40px;display:flex;font-size:16px"><span>Week</span><span>Month</span><span>Year</span></div>
""", """.seg>span{background:#E8F1FB;padding:6px 14px;border:1px solid #0B5CAD}
.seg>:first-child{border-radius:8px 0 0 8px} .seg>:last-child{border-radius:0 8px 8px 0}
.seg::after{content:"";position:absolute;left:0;right:0;bottom:-6px;height:2px;background:#0B5CAD}"""),
    "transform": ("""
<div class="arrows" style="position:absolute;left:40px;top:40px;display:flex;gap:20px">
  <div style="width:60px;height:40px;background:linear-gradient(90deg,#0B5CAD,#1E9E5A)"></div>
  <div style="width:60px;height:40px;background:linear-gradient(90deg,#0B5CAD,#1E9E5A)"></div>
</div>""", """.arrows>:last-child{transform:scaleX(-1)}
.arrows::after{content:"";position:absolute;left:0;right:0;bottom:-8px;height:2px;background:#D64545}"""),
    "z-index": ("""
<div class="stack" style="position:absolute;left:40px;top:40px;width:300px;height:200px">
  <div style="position:absolute;left:0;top:0;width:150px;height:100px;background:#D64545"></div>
  <div style="position:absolute;left:50px;top:50px;width:150px;height:100px;background:#0B5CAD"></div>
</div>""", """.stack>:first-child{z-index:2}
.stack::before{content:"";position:absolute;left:0;top:-10px;width:300px;height:4px;background:#1E9E5A}"""),
    "empty": ("""
<div style="position:absolute;left:40px;top:40px;display:flex;gap:10px;align-items:center;font-size:16px">
  <span class="sw"></span><span>Series A</span>
</div>""", """.sw{display:inline-block;width:14px;height:14px;background:#0B5CAD} .sw:empty{border-radius:50%}
.sw::before{content:'';display:block;width:6px;height:6px;margin:4px;background:#fff}"""),
    "marker": ("""
<ul class="u" style="position:absolute;left:60px;top:40px;width:400px;font-size:18px;margin:0"><li>First</li><li>Second</li></ul>
""", """.u::before{content:'';position:absolute;left:-20px;top:0;bottom:0;width:3px;background:#0B5CAD}
li:first-child::marker{color:#D64545}"""),
}


def _everything_of(ir: IR) -> list[tuple]:
    """Every element as the invariance test sees it, plus z order (the paint order)."""
    return [(e.z,) + row for e, row in zip([e for e in ir.elements if "::" not in (e.source or {}).get("path", "")],
                                           _geometry_of(ir), strict=False)]


@pytest.mark.parametrize("case", sorted(B2_CASES))
def test_a_rebuild_that_restyles_anything_is_reverted(tmp_path, monkeypatch, case):
    """B2: check (iii) compares every computed longhand of every element (and a list item's
    `::marker`, and every pseudo-element still live), so a copy that flips a radius, a mirror, a z
    order or a marker colour is taken back with a warn, and the export is the pre-pass-off export."""
    html = _context_slide(tmp_path, *B2_CASES[case])
    on = _extract(html)
    monkeypatch.setattr(html_extract, "MATERIALISE_PSEUDO", False)
    off = _extract(html)
    assert not _pseudo(on)
    (warn,) = [d for d in on.diagnostics if d.level == "warn"]
    assert "could not be rebuilt" in warn.message and _rule_of(warn.message) == "pseudo-dropped"
    assert _everything_of(on) == _everything_of(off)


NEGATIVE_Z = """
<div class="wrap" style="position:absolute;left:40px;top:40px;width:400px;display:flex;flex-direction:column;gap:12px;padding:10px">
  <div class="card">Gradient border card</div>
</div>
<div class="card abs" style="position:absolute;left:600px;top:50px;width:300px">Absolute card</div>
<div class="wrap2" style="position:absolute;left:40px;top:300px;width:400px;padding:10px;background:#EEF1F5">
  <div class="card">Under a painted wrapper</div>
</div>
<div class="wrap3" style="position:absolute;left:600px;top:300px;width:400px;padding:10px">
  <div class="card sc">Stacking-context card</div>
</div>
"""
NEGATIVE_Z_CSS = """
.card{position:relative;background:#fff;border-radius:12px;padding:20px;font-size:18px}
.card::before{content:"";position:absolute;inset:-3px;border-radius:15px;background:linear-gradient(90deg,#0B5CAD,#1E9E5A);z-index:-1}
.card.sc{z-index:0}
"""


def test_a_negative_z_copy_paints_under_its_hosts_background(tmp_path):
    """M1: `z-index:-1` under a host that is not a stacking context paints under the host's own
    background (the gradient-border idiom) — for a text-block host and a walked host alike. Under a
    host that is a stacking context it paints over it. Under a *painted wrapper* (no stacking context
    between) the browser hides it under the wrapper's background: the export cannot draw it there,
    and says so rather than drawing it over the wrapper."""
    ir = _extract(_context_slide(tmp_path, NEGATIVE_Z, NEGATIVE_Z_CSS))

    def pair(prefix):
        card = next(e for e in ir.elements if e.kind == "shape" and (e.source or {}).get("path", "").endswith(prefix))
        rims = [e for e in _pseudo(ir, prefix + "::before") if e.kind == "shape"]
        return card, rims

    card, (rim,) = pair("div.wrap:nth-child(1) > div.card:nth-child(1)")
    assert rim.z < card.z and rim.fill["type"] == "gradient"
    card, (rim,) = pair("div.card:nth-child(2)")
    assert rim.z < card.z
    card, (rim,) = pair("div.wrap3:nth-child(4) > div.card:nth-child(1)")
    assert rim.z > card.z
    card, rims = pair("div.wrap2:nth-child(3) > div.card:nth-child(1)")
    assert not rims
    (warn,) = [d for d in ir.diagnostics if d.level == "warn"]
    assert warn.source.endswith("div.wrap2:nth-child(3) > div.card:nth-child(1)::before")
    assert "cannot draw it in its place" in warn.message and _rule_of(warn.message) == "pseudo-dropped"


def test_a_deferred_copy_keeps_its_hosts_clip(tmp_path):
    """M2: a copy walked beside its text block's paragraph takes the block's own `overflow` clip
    (and opacity), as a child walked as an element does: the KPI card's corner circle is cut."""
    ir = _extract(_context_slide(tmp_path, """
<div class="col" style="position:absolute;left:40px;top:40px;width:400px;display:flex;flex-direction:column;gap:12px">
  <div class="kpi">$12.4B</div>
</div>""", """.kpi{position:relative;overflow:hidden;padding:20px;background:#0B2545;color:#fff;border-radius:12px;font-size:28px;opacity:.8}
.kpi::before{content:"";position:absolute;top:-30px;right:-30px;width:100px;height:100px;border-radius:50%;background:#1E9E5A}"""))
    (circle,) = _pseudo(ir)
    card = next(e for e in ir.elements if e.kind == "shape" and e is not circle)
    assert _box(circle) == (370, 10, 100, 100)
    clip = circle.clip
    assert clip is not None and (clip.x, clip.y, clip.w, clip.h) == (370, 40, 70, 70)   # the card's padding box
    assert circle.clip.x + circle.clip.w <= card.box.x + card.box.w + 0.01
    assert circle.opacity == pytest.approx(0.8)


def test_no_bullet_where_the_browser_draws_no_marker(tmp_path):
    """M3: a marker is drawn only for `display: list-item` with a `::marker` that has content. A flex
    or block `li`, and `::marker{content:"" | none}`, draw only their own `::before` — once."""
    ir = _extract(_context_slide(tmp_path, """
<ul class="l" style="position:absolute;left:40px;top:40px;width:400px;margin:0;padding:0;font-size:18px"><li>First point</li><li>Second point</li></ul>
<ul class="m" style="position:absolute;left:500px;top:40px;width:400px;margin:0;padding:0;font-size:18px"><li>Marker empty</li></ul>
<ul class="n" style="position:absolute;left:40px;top:200px;width:400px;margin:0;padding:0;font-size:18px"><li>Block li</li></ul>
<ul class="o" style="position:absolute;left:500px;top:200px;width:400px;margin:0;padding:0 0 0 20px;font-size:18px"><li>Marker none</li></ul>
<ul class="p" style="position:absolute;left:40px;top:400px;width:400px;margin:0;padding:0 0 0 20px;font-size:18px"><li>Real marker</li></ul>
""", """.l li{display:flex;align-items:center;gap:8px}
.l li::before{content:"";width:8px;height:8px;border-radius:50%;background:#1E9E5A;flex:none}
.m li::marker{content:""} .m li::before{content:"✓ ";color:#1E9E5A}
.n li{display:block} .n li::before{content:"→ ";color:#0B5CAD}
.o li::marker{content:none} .o li::before{content:"✓ ";color:#0B5CAD}"""))
    texts = [e for e in ir.elements if e.kind == "text"]
    bullets = {_texts_of(e)[0]: [p.get("bullet") for p in e.paragraphs] for e in texts}
    assert bullets["First point"] == [None] and bullets["Second point"] == [None]
    assert bullets["✓ Marker empty"] == [None] and bullets["→ Block li"] == [None]
    assert bullets["✓ Marker none"] == [None]
    assert bullets["Real marker"][0]["char"] == "•"                  # an ordinary list keeps its marker
    dots = [e for e in _pseudo(ir, "::before") if e.kind == "shape"]
    assert len(dots) == 2 and all(e.geometry["type"] == "roundRect" for e in dots)


def test_counters_follow_the_layout_tree_and_say_where_they_cannot(tmp_path):
    """m1: a `display: contents` wrapper's own counter properties do nothing and its children count
    as its parent's (Chromium keeps counters on boxes): three wrapped steps are I, II, III — the
    browser's own widths prove it, since check (i) passes. `counter(list-item)` in a reversed list
    and `quotes: auto` in French cannot be matched: each says so (approximated), never silently."""
    ir = _extract(_context_slide(tmp_path, """
<div class="w" style="position:absolute;left:40px;top:40px;font-size:20px">
  <div class="dc"><div class="s">One</div><div class="s">Two</div></div>
  <div class="dc"><div class="s">Three</div></div>
</div>
<div class="k" style="position:absolute;left:400px;top:40px;font-size:18px;width:300px">
  <div class="dk"><div class="c">One</div><div class="c">Two</div></div>
  <div class="dk"><div class="c">Three</div></div>
</div>
<ol reversed style="position:absolute;left:800px;top:40px;list-style:none;margin:0;padding:0;font-size:18px;width:300px">
  <li class="r">Alpha</li><li class="r">Beta</li>
</ol>
<blockquote lang="fr" style="position:absolute;left:40px;top:300px;width:600px;margin:0;font-size:24px">La croissance est forte</blockquote>
""", """.dc{display:contents;counter-reset:k} .s{counter-increment:k} .s::before{content:counter(k, upper-roman) '. '}
.dk{display:contents;counter-reset:q} .c{counter-increment:q;position:relative;padding-left:30px;height:26px}
.c::before{content:counter(q);position:absolute;left:0;top:0;width:22px;height:22px;border-radius:50%;background:#0B5CAD;color:#fff;font-size:12px;line-height:22px;text-align:center}
.r{position:relative;padding-left:30px}
.r::before{content:counter(list-item);position:absolute;left:0;top:0;width:22px;height:22px;background:#0B5CAD;color:#fff;font-size:12px;text-align:center}
blockquote::after{content:close-quote} blockquote::before{content:no-open-quote}"""))
    (roman,) = [e for e in ir.elements if e.kind == "text" and _texts_of(e)[0].startswith("I.")]
    lines = ["".join(r["text"] for r in line["runs"]) for p in roman.paragraphs for line in p["lines"]]
    assert lines == ["I. One", "II. Two", "III. Three"]
    circles = [e for e in _pseudo(ir, "::before") if e.kind == "text" and "div.c:" in e.source["path"]]
    assert [_texts_of(c) for c in circles] == [["1"], ["2"], ["3"]]
    warns = [d.message for d in ir.diagnostics if d.level == "warn"]
    assert warns == ["counter(list-item) in a reversed or renumbered list is resolved in document order, approximated",
                     "quotes for lang fr are drawn as English quotes, approximated"]
    assert all(_rule_of(w) == "approximated" for w in warns)
    assert "0 dropped" in next(d.message for d in ir.diagnostics if d.source == "pseudo-elements")


def test_a_rollback_that_does_not_restore_the_page_costs_no_other_pseudo(tmp_path):
    """m2: Chromium renumbers a reversed list's `counter(list-item)` after any change to the page, so
    taking a failed batch back no longer restores the first layout. The rest are judged against the
    page as it now is (said so, info) — the unrelated dot and the `ol start` items are kept."""
    ir = _extract(_context_slide(tmp_path, """
<ol class="r" reversed style="position:absolute;left:40px;top:40px;list-style:none;margin:0;padding:0;font-size:18px;width:300px">
  <li>Alpha</li><li>Beta</li><li>Gamma</li></ol>
<ol class="s" start="4" style="position:absolute;left:800px;top:40px;list-style:none;margin:0;padding:0;font-size:18px;width:300px">
  <li>Four</li><li>Five</li></ol>
<div class="bl" style="position:absolute;left:40px;top:300px;width:300px;font-size:16px"><div>Unrelated bullet</div></div>
""", """ol li::before{content:counter(list-item) '. ';color:#0B5CAD}
.bl div{position:relative;padding-left:16px}
.bl div::before{content:'';position:absolute;left:0;top:6px;width:7px;height:7px;background:#D64545;border-radius:50%}"""))
    assert not [d for d in ir.diagnostics if "could not be rebuilt" in d.message]
    (starts,) = [e for e in ir.elements if e.kind == "text" and _texts_of(e)[0].endswith("Four")]
    assert _texts_of(starts) == ["4. Four", "5. Five"]
    assert [e.kind for e in _pseudo(ir, "div:nth-child(1)::before")] == ["shape"]
    assert next(d.message for d in ir.diagnostics if d.source == "pseudo-elements"
                and d.message[0].isdigit()) == "6 pseudo-elements rebuilt as shapes and text, 0 dropped"
    infos = [d.message for d in ir.diagnostics if d.level == "info" and d.message.startswith("the page did not return")]
    assert len(infos) <= 1       # said once when Chromium re-lays the list out, never per pseudo
    assert [d.message.split(" in ")[0] for d in ir.diagnostics if d.level == "warn"] == ["counter(list-item)"]


def test_every_painted_pseudo_is_built_or_reported(tmp_path):
    """m4: the page's own `html::before`, a custom checkbox tick (`appearance:none`) and a pseudo that
    paints only an inset ring used to vanish without a word. The first two cannot hold a copy and are
    reported; the ring is built (and rasterised like any inset shadow)."""
    ir = _extract(_context_slide(tmp_path, """
<label style="position:absolute;left:40px;top:120px;font-size:18px;display:flex;align-items:center;gap:8px"><input type="checkbox" class="cb" checked> Done item</label>
<div class="ring" style="position:absolute;left:400px;top:120px;width:120px;height:60px;font-size:16px">Ring</div>
""", """html::before{content:'';position:fixed;left:0;top:0;width:1280px;height:8px;background:#0B5CAD}
.cb{appearance:none;width:16px;height:16px;border:2px solid #0B5CAD;border-radius:3px;position:relative;margin:0}
.cb::after{content:'';position:absolute;left:3px;top:0;width:5px;height:9px;border:solid #0B5CAD;border-width:0 2px 2px 0;transform:rotate(45deg)}
.ring::after{content:'';position:absolute;inset:0;border-radius:8px;box-shadow:inset 0 0 0 3px #D64545}"""))
    pseudo_warns = {d.source.rsplit(" > ", 1)[-1]: d.message for d in ir.diagnostics
                    if d.level == "warn" and _rule_of(d.message) == "pseudo-dropped"}
    assert pseudo_warns == {"html::before": "a ::before on the page itself is not drawn; dropped",
                            "input.cb:nth-child(1)::after": "a ::after on a form control is not drawn; dropped"}
    assert [e.kind for e in _pseudo(ir, "div.ring:nth-child(2)::after")] == ["raster"]


def test_the_new_pseudo_words_are_reached():
    for message in ("a ::after was rebuilt but the export cannot draw it in its place; dropped",
                    "a ::before on the page itself is not drawn; dropped",
                    "a ::after on a form control is not drawn; dropped"):
        assert _rule_of(message) == "pseudo-dropped", message
    for message in ("counter(list-item) in a reversed or renumbered list is resolved in document order, approximated",
                    "quotes for lang de are drawn as English quotes, approximated"):
        assert _rule_of(message) == "approximated", message


def test_the_pre_pass_stays_fast_with_the_full_signature(page, tmp_path):
    """The full-longhand check (iii) on a busy slide: 40 hosts, 80 pseudo-elements, one layout check
    for the batch — well inside the brief's 0.5 s per slide (the 36-pseudo app slide takes ~0.25 s)."""
    import time

    rows = "".join(f'<div class="r" style="top:{20 + i * 17}px">Row {i}</div>' for i in range(40))
    html = _context_slide(tmp_path, rows, """.r{position:absolute;left:60px;font-size:12px}
.r::before{content:'';position:absolute;left:-12px;top:4px;width:6px;height:6px;border-radius:50%;background:#0B5CAD}
.r::after{content:' →';color:#1E9E5A}""")
    context = context_for(html)
    from app.engine.extract.html import PAGE_JS, prepare_page

    prepare_page(page, html, context.manifest, context.assets_dir)
    page.add_script_tag(path=str(COLOR_JS))
    page.add_script_tag(path=str(PAGE_JS))
    start = time.perf_counter()
    result = page.evaluate("() => window.__engineMaterialisePseudo()")
    elapsed = time.perf_counter() - start
    assert len(result["materialised"]) == 80 and not result["dropped"]
    assert elapsed < 2.0, f"{elapsed:.2f} s"          # generous for a loaded machine; ~0.3 s idle


# ============================================================================== the `paint` family


def test_the_paint_family_holds_its_numbers(torture_dir):
    """`fixtures/torture/paint.html`: one assertion per cell, the numbers in paint.expect.json."""
    import json

    html = torture_dir / "paint.html"
    ir = _extract(html)
    expect = json.loads((torture_dir / "paint.expect.json").read_text(encoding="utf-8"))
    assert ir.counts() == expect["kinds"]

    def one(name):
        (element,) = [e for e in ir.elements if e.name == name]
        return element

    # row 1: colours
    assert one("mix-srgb").fill == {"type": "solid", "color": "E4F3EB", "alpha": 1}
    assert one("oklch-fill").fill["color"] == "DCF7E1"
    assert one("p3-fill").fill["color"] == "E1F8EA"
    assert one("lab-fill").fill["color"] == "E8F4EB"
    gradient = one("gradient-modern-stops").fill
    assert len(gradient["stops"]) == 2 and gradient["stops"][0]["color"] == "CEDEEF"
    dark = [e for e in ir.elements if e.name == "dark-oklch-text"]
    box, text = (dark[0], dark[1]) if dark[0].kind == "shape" else (dark[1], dark[0])
    assert box.fill["color"] == "0B2545" and box.stroke["width"] == 2
    assert box.shadow["dy"] == 4 and box.shadow["blur"] == 12 and box.shadow["alpha"] == pytest.approx(0.3)
    assert {r["color"] for r in _runs(text)} == {"F8F8F8"}
    spaces = [d for d in ir.diagnostics if "interpolated in" in d.message]
    assert len(spaces) == 1 and "oklab" in spaces[0].message

    # row 2: gradient text — no slab anywhere in the row, the run fills and mid-colours
    assert not [e for e in ir.elements if e.kind == "shape" and 210 <= e.box.y < 320]
    assert _runs(one("grad-text"))[0]["color"] == "157D84" and list(one("grad-text").extras["x-wpe-run-fills"]) == ["0/0/0"]
    # WP-B: every text element carries `x-run-boxes`; a solid clip still has no run fills.
    assert _runs(one("solid-clip-text"))[0]["color"] == "0B5CAD" and not one("solid-clip-text").extras.get("x-wpe-run-fills")
    assert list(one("span-clip-text").extras["x-wpe-run-fills"]) == ["0/0/1"]
    assert one("clip-text-image").kind == "raster"
    assert _runs(one("grad-text-padding"))[0]["color"] == "157D84"
    assert _runs(one("grad-text-3-stops"))[0]["color"] == "FFFFFF"

    # row 3: pseudo-elements
    assert _box(one("axis-dot::before")) == (40, 390, 160, 3)
    assert _box(one("axis-dot::after")) == (107, 378, 26, 26)
    assert _texts_of(one("counter-circles")) == ["Step one", "Step two", "Step three"]
    circles = [e for e in _pseudo(ir, "::before") if e.kind == "text" and 370 <= e.box.y < 460]
    assert [_texts_of(c) for c in circles] == [["1"], ["2"], ["3"]]
    assert _texts_of(one("attr-inline")) == ["→ Q3 Revenue"]
    assert _texts_of(one("quotes")) == ["“Quoted”"]
    assert _texts_of(one("none-and-clearfix")) == ["No pseudo"]
    assert _box(one("flex-block-before::before")) == (1015, 413, 24, 24)
    note = next(d for d in ir.diagnostics if d.source == "pseudo-elements")
    assert note.message == "10 pseudo-elements rebuilt as shapes and text, 0 dropped"

    # row 4: joins
    assert _path_points(one("tri-right left border")) == [(40, 546), (40, 530), (54, 538)]
    assert _path_points(one("tri-up bottom border")) == [(255, 542), (235, 542), (245, 530)]
    assert _path_points(one("chevron-after::after left border")) == [(550, 570), (550, 530), (570, 550)]
    assert _path_points(one("two-sides-zero top border"))[-1] == _path_points(one("two-sides-zero left border"))[-1] == (635, 540)
    assert _path_points(one("two-sides-box top border")) == [(820, 530), (980, 530), (980, 536), (830, 536)]
    assert _path_points(one("tri-oklch-transparent left border")) == [(1015, 546), (1015, 530), (1029, 538)]

    # row 5: opacity (r1d, r1e) — an element at 0, a contents wrapper that fades nothing, an inline box's paint
    assert one("opacity-0-text").opacity == 0 and _texts_of(one("opacity-0-text")) == ["Invisible words"]
    row5 = [e for e in ir.elements if 320 <= e.box.y < 368]           # the band between rows 2 and 3
    chips = [e for e in row5 if e.kind == "shape" and e.box.x < 600]
    assert [round(e.opacity, 3) for e in chips] == [1.0]                  # under the contents wrapper at 0.2
    slab = [e for e in row5 if e.kind == "shape" and 600 <= e.box.x < 900]
    assert [round(e.opacity, 3) for e in slab] == [0.5]
    assert [r["alpha"] for r in _runs(one("inline-half")) if r["text"] == "half"] == [0.5]
    around = sorted((e for e in row5 if e.kind == "shape" and e.box.x >= 900), key=lambda e: e.box.x)
    assert [(e.fill["color"], round(e.opacity, 3)) for e in around] == [("FFE600", 0.5), ("0B5CAD", 0.5)]
    assert not [d for d in ir.diagnostics if d.source == "text-overflow"]    # judged against the paragraph

    warns = [d for d in ir.diagnostics if d.level == "warn"]
    assert [d.message for d in warns] == ["rasterised: background-clip:text over an image"]


# ================================================================== r1d: opacity (plan §16 #27)
#
# Three rules r1a found (r1a.md, findings 4 and 5, known gap 2), each measured against the browser on
# `fidelity-reports/r1d-evidence/probe.html` before the change: an element's opacity of 0 is 0 in the
# file, a `display: contents` element's opacity fades nothing, and an inline box's own opacity is on
# its own paint as it is on its text.


def _alpha_of(colour) -> str | None:
    """`a:alpha@val` of a colour element, None when it has none (fully opaque)."""
    alpha = colour.find(f"{A_NS}alpha") if colour is not None else None
    return alpha.get("val") if alpha is not None else None


def _shape_fill(deck: Path, colour: str):
    """The `srgbClr` of the one shape in the deck whose own solid fill is `colour`."""
    (found,) = [s._element.spPr.find(f"{A_NS}solidFill/{A_NS}srgbClr") for s in _slide_shapes(deck)
                if s._element.spPr.find(f"{A_NS}solidFill/{A_NS}srgbClr[@val='{colour}']") is not None]
    return found


def test_an_element_at_opacity_0_writes_its_text_at_alpha_0(tmp_path):
    """`_write_run` read `element.opacity or 1.0`, so an opacity of 0 was written as 1: a positioned
    `<div style="opacity:0">` exported its words in opaque ink (r1a finding 5) and a gradient text
    inside one drew its gradient. The file now carries what the browser shows — nothing — the way it
    already carries an opacity-0 box (fill alpha 0, `shapes._apply_fill`, unchanged) and a
    `color: transparent` run (alpha 0): the element stays, invisible and editable."""
    html = _context_slide(tmp_path, """
<div style="position:absolute;left:40px;top:40px;opacity:0;font-size:28px;color:#0B2545">Invisible words</div>
<div style="position:absolute;left:340px;top:40px;opacity:0;font-size:36px;font-weight:bold"><span class="g">$12.4B</span></div>
<div style="position:absolute;left:640px;top:40px;width:200px;height:60px;opacity:0;background:#0B5CAD"></div>
<div style="position:absolute;left:40px;top:200px;opacity:.5;font-size:28px;color:#0B2545">Half words</div>""",
                           ".g{background:linear-gradient(90deg,#0B5CAD,#1E9E5A);-webkit-background-clip:text;"
                           "background-clip:text;color:transparent}")
    ir = _extract(html)
    texts = {"".join(r["text"] for r in _runs(e)): e for e in ir.elements if e.kind == "text"}
    assert texts["Invisible words"].opacity == 0 and texts["$12.4B"].opacity == 0
    deck, _ = _emit(ir, html, tmp_path)

    (rpr,) = _run_rprs(deck, "Invisible words")
    assert _alpha_of(rpr.find(f"{A_NS}solidFill/{A_NS}srgbClr")) == "0"
    (rpr,) = _run_rprs(deck, "$12.4B")
    stops = rpr.findall(f"{A_NS}gradFill/{A_NS}gsLst/{A_NS}gs")
    assert len(stops) == 2 and [_alpha_of(gs.find(f"{A_NS}srgbClr")) for gs in stops] == ["0", "0"]
    (rpr,) = _run_rprs(deck, "Half words")
    assert _alpha_of(rpr.find(f"{A_NS}solidFill/{A_NS}srgbClr")) == "50000"
    assert _alpha_of(_shape_fill(deck, "0B5CAD")) == "0"      # the shape path was right, and stays so

CONTENTS_CSS = """
.chip{display:inline-block;padding:4px 10px;background:#0B5CAD;color:#fff}
.flag{position:relative}
.flag::after{content:"";position:absolute;left:0;right:0;bottom:-6px;height:5px;background:#D64545}
"""


def test_a_display_contents_ancestors_opacity_fades_no_copy_and_no_detached_box(tmp_path):
    """E's `pseudoContext` multiplied in the opacity of every ancestor between a copy (or a detached
    box) and its walk's container, `display: contents` ones included — but such an element generates
    no box, so its opacity paints nothing: Edge draws the chip and the flag bars below at full colour
    (measured, r1d-evidence `probe.png` b1–b3; r1a measured the same for text). Skipped, as
    `opacityBelow` skips it for text. A real inline wrapper's opacity still fades its copy."""
    html = _context_slide(tmp_path, """
<p style="position:absolute;left:40px;top:40px;width:330px;font-size:22px;margin:0">Pay <span style="display:contents;opacity:.2"><span class="chip">CHIP</span></span> now</p>
<p style="position:absolute;left:420px;top:40px;width:330px;font-size:22px;margin:0">Note <span style="display:contents;opacity:.2"><b class="flag">flagged</b></span> here</p>
<div style="position:absolute;left:800px;top:40px;width:400px"><p style="font-size:22px;margin:0">Mark <span style="display:contents;opacity:.2"><b class="flag">this</b></span> word</p></div>
<p style="position:absolute;left:40px;top:200px;width:330px;font-size:22px;margin:0">Real <span style="opacity:.2"><b class="flag">faded</b></span> wrapper</p>""",
                           CONTENTS_CSS)
    ir = _extract(html)
    (chip,) = [e for e in ir.elements if e.kind == "shape" and (e.fill or {}).get("color") == "0B5CAD"]
    (chip_text,) = [e for e in ir.elements if e.kind == "text" and _texts_of(e) == ["CHIP"]]
    assert chip.opacity == pytest.approx(1.0) and chip_text.opacity == pytest.approx(1.0)
    bars = sorted(_pseudo(ir, "::after"), key=lambda e: (e.box.y, e.box.x))
    assert [e.kind for e in bars] == ["shape"] * 3
    assert [round(e.opacity, 3) for e in bars] == [1.0, 1.0, 0.2]   # anonymous run, text block, real wrapper


def test_an_inline_boxs_own_opacity_is_on_its_own_paint(tmp_path):
    """`emitInlineDecoration` pushed an inline box's fill and sides in its paragraph's context, so
    `<span style="opacity:.5;background:…">` exported an opaque slab under text r1a had faded to 0.5
    (r1a known gap 2). The box's paint now takes the opacity of the box and of the inline boxes
    between it and its paragraph — r1a's rule for its text (`opacityBelow`), contents boxes skipped —
    in a text block, an anonymous run, and around a detached box (`walkDetached`)."""
    html = _context_slide(tmp_path, """
<p style="position:absolute;left:40px;top:40px;font-size:26px;margin:0">Plain <span style="opacity:.5;background:#0B5CAD;color:#fff;padding:2px 6px">half</span> text</p>
<p style="position:absolute;left:420px;top:40px;font-size:26px;margin:0">Say <span style="opacity:.5"><mark style="background:#FFE600;color:#000">marked</mark></span> twice</p>
<p style="position:absolute;left:800px;top:40px;font-size:26px;margin:0">Side <span style="opacity:.4;border-left:8px solid #D64545;background:#1E9E5A;color:#fff;padding:2px 6px">left</span> end</p>
<p style="position:absolute;left:40px;top:160px;font-size:26px;margin:0">Text <mark style="opacity:.5;background:#F2C14E;padding:6px 0">around <span style="display:inline-block;background:#7A3E9D;color:#fff;padding:2px 8px">CHIP</span></mark> after</p>
<div style="position:absolute;left:560px;top:160px;width:400px;font-size:26px"><span style="opacity:.5;background:#2A9D8F;color:#fff;padding:2px 6px">anon</span> run<div style="font-size:20px">block below</div></div>
<p style="position:absolute;left:40px;top:300px;font-size:26px;margin:0">Kept <span style="display:contents;opacity:.2"><mark style="background:#E76F51">whole</mark></span> here</p>""")
    ir = _extract(html)

    def paint(colour: str) -> list[float]:
        return [round(e.opacity, 3) for e in ir.elements
                if e.kind == "shape" and (e.fill or {}).get("color") == colour]

    assert paint("0B5CAD") == [0.5]                                # its own opacity, in a text block
    assert paint("FFE600") == [0.5]                                # an inline ancestor's
    assert paint("1E9E5A") == [0.4] and paint("D64545") == [0.4]   # the fill and its side
    assert paint("F2C14E") == [0.5]                                # the mark painted under a detached chip
    assert paint("7A3E9D") == [0.5]                                # the chip itself, as before (pseudoContext)
    assert paint("2A9D8F") == [0.5]                                # an anonymous run's inline box
    assert paint("E76F51") == [1.0]                                # a contents box fades nothing
    deck, _ = _emit(ir, html, tmp_path)
    assert _alpha_of(_shape_fill(deck, "0B5CAD")) == "50000"


# ============================================================= r1e: r1d's leftovers (plan §16 #28)
#
# Four rules r1d measured and left to their owners (r1d.md, known gaps 1-4), each measured again on
# `fidelity-reports/r1e-evidence/probe/probe.html` (r1d's probe plus four list cases) before the change:
# an opacity-0 box's border and shadow, an opacity-0 list's bullet, an in-flow block pseudo copy's own
# opacity, and a text-overflow warn judged against a `display: contents` wrapper's zero rect.


def _stroke_and_shadow(deck: Path, colour: str):
    """The line colour and the outer-shadow colour of the one shape whose own solid fill is `colour`."""
    (shape,) = [s for s in _slide_shapes(deck)
                if s._element.spPr.find(f"{A_NS}solidFill/{A_NS}srgbClr[@val='{colour}']") is not None]
    spPr = shape._element.spPr
    return (spPr.find(f"{A_NS}ln/{A_NS}solidFill/{A_NS}srgbClr"),
            spPr.find(f"{A_NS}effectLst/{A_NS}outerShdw/{A_NS}srgbClr"))


def test_an_opacity_0_box_writes_its_border_and_shadow_at_alpha_0(tmp_path):
    """`_apply_stroke` and `_apply_effects` multiplied by `element.opacity or 1.0`, so an opacity-0 box
    exported its border and its shadow opaque while its fill was already at alpha 0 (r1d probe a3:
    2,144 px of red border where the browser draws nothing). They now read the opacity as
    `_apply_fill` does: 0 is 0. A box at 0.5 keeps 0.5 on all three."""
    html = _context_slide(tmp_path, """
<div style="position:absolute;left:40px;top:40px;width:200px;height:60px;opacity:0;background:#0B5CAD;border:4px solid #D64545;box-shadow:0 8px 0 #1E9E5A"></div>
<div style="position:absolute;left:340px;top:40px;width:200px;height:60px;opacity:.5;background:#F2C14E;border:4px solid #7A3E9D;box-shadow:0 8px 0 #2A9D8F"></div>""")
    ir = _extract(html)
    deck, _ = _emit(ir, html, tmp_path)
    line, shadow = _stroke_and_shadow(deck, "0B5CAD")
    assert (line.get("val"), shadow.get("val")) == ("D64545", "1E9E5A")
    assert [_alpha_of(line), _alpha_of(shadow), _alpha_of(_shape_fill(deck, "0B5CAD"))] == ["0", "0", "0"]
    line, shadow = _stroke_and_shadow(deck, "F2C14E")
    assert [_alpha_of(line), _alpha_of(shadow), _alpha_of(_shape_fill(deck, "F2C14E"))] == ["50000"] * 3


def _bullets(deck: Path) -> dict[str, object]:
    """Per paragraph text: its `buClr` colour element, or `"buNone"`, or None when it has neither."""
    out: dict[str, object] = {}
    for shape in _slide_shapes(deck):
        for p in shape._element.iter(f"{A_NS}p"):
            text = "".join(t.text or "" for t in p.iter(f"{A_NS}t"))
            pPr = p.find(f"{A_NS}pPr")
            if not text.strip() or pPr is None:
                continue
            colour = pPr.find(f"{A_NS}buClr/{A_NS}srgbClr")
            out[text] = colour if colour is not None else ("buNone" if pPr.find(f"{A_NS}buNone") is not None else None)
    return out


def test_a_bullet_is_as_opaque_as_its_paragraph(tmp_path):
    """`_set_bullet` wrote `a:buClr` with no alpha, so an opacity-0 list kept its bullet (r1d probe a4)
    — and so did a list at 0.5 (at full colour), an in-flow item at opacity 0 (r1a's fold puts that on
    its runs) and an item in a transparent colour. The browser paints a marker in its item's colour
    under its item's opacity; the paragraph's most opaque run carries both (an inline box can only fade
    a run further), so the bullet takes that alpha times its element's opacity. Measured first
    (r1e-evidence/bullets): PowerPoint and PptxRender both draw `a:alpha` in `buClr` — 50000 as a
    half-tone bullet, 0 as none — while `buNone` would move the first line to the bullet's column."""
    html = _context_slide(tmp_path, """
<ul style="position:absolute;left:40px;top:40px;margin:0;padding:0 0 0 28px;opacity:0;font-size:24px;color:#0B2545"><li>Gone list</li></ul>
<ul style="position:absolute;left:340px;top:40px;margin:0;padding:0 0 0 28px;opacity:.5;font-size:24px;color:#0B2545"><li>Half list</li></ul>
<div style="position:absolute;left:640px;top:40px;font-size:24px;color:#0B2545"><ul style="margin:0;padding:0 0 0 28px"><li style="opacity:0">Gone item</li></ul></div>
<ul style="position:absolute;left:940px;top:40px;margin:0;padding:0 0 0 28px;font-size:24px;color:transparent"><li>Clear item</li></ul>
<ul style="position:absolute;left:40px;top:200px;margin:0;padding:0 0 0 28px;font-size:24px;color:#0B2545"><li><span style="opacity:.5">Faded</span> word</li></ul>""")
    ir = _extract(html)
    deck, _ = _emit(ir, html, tmp_path)
    bullets = _bullets(deck)
    assert set(bullets) == {"Gone list", "Half list", "Gone item", "Clear item", "Faded word"}
    assert all(b is not None and not isinstance(b, str) for b in bullets.values()), bullets   # never buNone
    assert {text: _alpha_of(colour) for text, colour in bullets.items()} == {
        "Gone list": "0", "Half list": "50000", "Gone item": "0", "Clear item": "0", "Faded word": None}


def test_an_in_flow_block_pseudo_copy_paints_under_its_own_opacity(tmp_path):
    """`walkDeferred` drew an in-flow block copy (a bar in a block-in-inline) with `emitDecoration` in
    `pseudoContext`, which folds the opacity of the copy's ancestors but not the copy's own, so a
    `::before{display:block;opacity:.4}` bar exported opaque (r1d probe d1: 999 px). Its own opacity
    is now folded in, as `walkElement` folds it for an out-of-flow copy; its text was already faded."""
    css = """
.bar::before{content:"";display:block;width:80px;height:8px;background:#D64545;opacity:.4}
.solid::before{content:"";display:block;width:80px;height:8px;background:#1E9E5A}
"""
    html = _context_slide(tmp_path, """
<p style="position:absolute;left:40px;top:40px;width:400px;font-size:22px;margin:0">Lead <span class="bar">word</span> end</p>
<p style="position:absolute;left:40px;top:160px;width:400px;font-size:22px;margin:0">Lead <span style="opacity:.5"><span class="bar">word</span></span> end</p>
<p style="position:absolute;left:40px;top:280px;width:400px;font-size:22px;margin:0">Lead <span class="solid">word</span> end</p>""", css)
    ir = _extract(html)

    def paint(colour: str) -> list[float]:
        return [round(e.opacity, 3) for e in sorted(ir.elements, key=lambda e: e.box.y)
                if e.kind == "shape" and (e.fill or {}).get("color") == colour]

    assert paint("D64545") == [0.4, 0.2]      # its own; its own times an inline ancestor's
    assert paint("1E9E5A") == [1.0]


def test_a_chip_under_a_contents_wrapper_is_judged_against_a_box_that_exists(tmp_path):
    """`checkTextOverflow` walked from a detached chip up through its ancestors and judged each as
    the box the text must fit — a `display: contents` wrapper included, whose rect is all zeros: the
    chip's text 'runs 231 px below its box … the box is 0 px tall' (r1d known gap 4). A contents
    element generates no box, so the chain skips it and judges the nearest ancestor that has one:
    nothing for a chip in an auto-height paragraph, the real 12 px box for one that does overflow."""
    css = ".chip{display:inline-block;padding:4px 10px;background:#0B5CAD;color:#fff}"
    html = _context_slide(tmp_path, """
<p style="position:absolute;left:40px;top:40px;width:330px;font-size:22px;margin:0">Pay <span style="display:contents;opacity:.2"><span class="chip">CHIP</span></span> now</p>
<div style="position:absolute;left:40px;top:300px;width:400px;height:12px;font-size:22px;white-space:nowrap">Due <span style="display:contents"><span class="chip">LATE</span></span> soon</div>""", css)
    ir = _extract(html)
    overflow = [d.message for d in ir.diagnostics if d.source == "text-overflow"]
    assert not [m for m in overflow if "CHIP" in m], overflow
    late = [m for m in overflow if "LATE" in m]
    assert len(late) == 1 and late[0].endswith("the box is 12 px tall"), overflow
