"""r3 — a picture shows exactly the part of its source the browser shows (plan §16 #30).

The PowerPoint arbitration of the torture `images` family found PowerPoint and PptxRender agreeing with each
other and not with the browser: the emitter fitted every `cover` picture a second time on top of the crop
the extractor had measured, and `page.js` could not read a position written as lengths, `calc()` or the
four-value offset syntax (the image was dropped). The rule now: `page.js` builds the rectangle the whole
source is drawn into from the browser's own geometry — the frame, the natural size, the computed
`object-fit` / `background-size` and the computed positions, whose arithmetic the browser does — cuts it by
the frame, and the emitter writes that cut as `a:srcRect`, fitting nothing again.

The probe is `tests/engine/fixtures/verify_cases/crop_evidence.py`: 65 frames on three slides, each drawing a
coordinate-coded source (red = x, green = y), so a screenshot says which source pixel shows where.

1. Per crop mode, the IR's box and crop and the file's frame and `a:srcRect` are the values the CSS
   defines (hand-computed below, with Blink's pixel snapping of an `<img>` where it applies).
2. The window the IR describes is the window the browser shows (its own pixels, fitted), ± 0.5 px.
3. The window PptxRender draws from the file is the window the browser shows, ± 0.5 px; the cases with
   no single window (tiles, a rotation, an SVG) differ in only a few pixels.
4. The emitter writes a given crop as given (a hand-written IR without one is still fitted), a negative
   side as the empty band, and a mirrored picture's flip.
"""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

import pytest

pytest.importorskip("numpy")  # the probe measures pictures with numpy (requirements-dev.txt)

from pptx import Presentation  # noqa: E402
from pptx.util import Emu

from app.engine.emit import pptx as emitter
from app.engine.emit.shapes import _crop_to_clip, add_image
from app.engine.extract.html import measuring_page, render_reference
from app.engine.ir import IR, Box, Canvas, Element
from app.engine.reports import EmitOptions
from app.engine.verify.fixtures import context_for
from tests.engine import helpers as renderer
from tests.engine.fixtures.verify_cases import crop_evidence as probe
from tests.engine.helpers import extract_html

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

#: (page, index in the page, case tuple) for every probe case.
CASES = [(page, index, case) for page, (_, cases) in enumerate(probe.PAGES) for index, case in enumerate(cases)]
BY_NAME = {case[0]: (page, index, case) for page, index, case in CASES}


def _origin(index: int) -> tuple[int, int]:
    return probe.COLUMNS[index % len(probe.COLUMNS)], probe.ROWS[index // len(probe.COLUMNS)]


@pytest.fixture(scope="module")
def probed(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Any]:
    """Extract, reference-render and emit (unrounded and rounded IR) the three probe slides.

    PptxRender's renders are `probed_renders`, apart, so the tests that read only the IR and the
    file run without the render service."""
    root = tmp_path_factory.mktemp("crops")
    pages = probe.build(root)
    irs, references = [], []
    for number, html in enumerate(pages, start=1):
        context = context_for(html)
        with measuring_page(Canvas(1280, 720)) as page:
            irs.append(extract_html(html, context.manifest, context.layout_id, context.assets_dir,
                                    slide_id=f"r3-{html.stem}", title=html.stem, page=page))
        reference = root / "browser" / f"slide-{number:02d}.png"
        render_reference(html, context.layout_png, reference, assets_dir=context.assets_dir,
                         fonts=dict(context.manifest.fonts))
        references.append(reference)
    context = context_for(pages[0])
    exact = root / "exact.pptx"            # the in-memory IR, as the export pipeline emits it
    emitter.emit(irs, context.manifest, context.master, exact, EmitOptions())
    saved = root / "saved.pptx"            # the IR through its JSON (3 decimals), as the torture runner emits it
    emitter.emit([IR.from_json(ir.to_json()) for ir in irs], context.manifest, context.master, saved,
                 EmitOptions())
    return {"irs": irs, "references": references, "exact": exact, "saved": saved, "root": root}


@pytest.fixture(scope="module")
def probed_renders(probed: dict[str, Any]) -> list[Path]:
    """PptxRender's rendering of the saved (JSON round-tripped) deck — needs the render service."""
    return renderer.render(probed["saved"], 1280, probed["root"] / "pptxrender")


def _images(ir: IR, name: str) -> list[Element]:
    return [e for e in ir.elements if e.kind == "image" and e.name == name]


def _pictures(deck: Path, page: int, name: str) -> list[Any]:
    slide = Presentation(str(deck)).slides[page]
    return [s for s in slide.shapes if s.shape_type == 13 and s.name == name]


# ------------------------------------------------------------------------ 1. the values the CSS defines

#: name → [(dx, dy, w, h) of the IR box from the frame's origin, crop (l, t, r, b)], one entry per picture.
#: Frames are 180 × 140 at whole pixels; sources L 400 × 300, P 300 × 400, S 120 × 80, T 60 × 40.
EXPECTED: dict[str, list[tuple[tuple[float, ...], tuple[float, ...]]]] = {
    # object-fit — an <img>'s destination and content rects snapped to whole pixels, as Blink paints them
    "cover-L": [((0, 0, 180, 140), (3 / 186, 0, 3 / 186, 0))],              # 186.67 wide from −3.33 → 186 from −3
    "cover-P": [((0, 0, 180, 140), (0, 50 / 240, 0, 50 / 240))],
    "cover-L-left-top": [((0, 0, 180, 140), (0, 0, 7 / 187, 0))],            # 0 → 186.67 paints 0 → 187
    "cover-P-right-bottom": [((0, 0, 180, 140), (0, 100 / 240, 0, 0))],      # the family's row 1 frame 4 twin
    "cover-P-50-20": [((0, 0, 180, 140), (0, 20 / 240, 0, 80 / 240))],
    "cover-P-lengths": [((0, 0, 180, 140), (0, 10 / 240, 0, 90 / 240))],
    "cover-P-offsets": [((10, 0, 170, 120), (0, 120 / 240, 10 / 180, 0))],  # left 10px bottom 20px
    "cover-L-calc": [((0, 0, 175, 140), (12 / 187, 0, 0, 0))],               # calc(100% − 5px) of −6.67
    "cover-L-120pct": [((0, 0, 179, 140), (8 / 187, 0, 0, 0))],              # the frame shows past the image
    "contain-L": [((0, 3, 180, 135), (0, 0, 0, 0))],                         # 2.5 rounds half up
    "contain-P-left-top": [((0, 0, 105, 140), (0, 0, 0, 0))],
    "contain-P-offsets": [((65, 0, 105, 140), (0, 0, 0, 0))],                # right 10px
    "contain-L-30-80": [((0, 4, 180, 135), (0, 0, 0, 0))],
    "none-L": [((0, 0, 180, 140), (110 / 400, 80 / 300, 110 / 400, 80 / 300))],
    "none-L-lengths": [((0, 0, 180, 140), (20 / 400, 30 / 300, 200 / 400, 130 / 300))],
    "none-S": [((30, 30, 120, 80), (0, 0, 0, 0))],
    "none-S-offsets": [((55, 55, 120, 80), (0, 0, 0, 0))],
    "scale-down-L": [((0, 3, 180, 135), (0, 0, 0, 0))],
    "scale-down-S": [((0, 60, 120, 80), (0, 0, 0, 0))],
    "fill-P": [((0, 0, 180, 140), (0, 0, 0, 0))],
    "cover-L-padded": [((14, 14, 152, 112), (0, 1 / 114, 0, 1 / 114))],     # the content box is the frame
    "cover-P-radius": [((0, 0, 180, 140), (0, 50 / 240, 0, 50 / 240))],
    "cover-P-circle": [((0, 0, 140, 140), (0, 23 / 186, 0, 23 / 186))],
    "bleed-P": [((-40, -30, 300, 225), (0, 175 / 400, 0, 0))],               # clipped by its frame: below
    # background-size / -position — fractional tiles, the painting area snapped
    "bg-cover-P-right-bottom": [((0, 0, 180, 140), (0, 100 / 240, 0, 0))],  # the family's row 2 frame 4 twin
    "bg-cover-L-25-75": [((0, 0, 180, 140), (1 / 112, 0, 3 / 112, 0))],
    "bg-cover-P-offsets": [((10, 0, 170, 120), (0, 120 / 240, 10 / 180, 0))],
    "bg-cover-P-lengths": [((0, 0, 180, 140), (0, 30 / 240, 0, 70 / 240))],
    "bg-cover-L-calc": [((0, 10, 170, 130), (5 / 56, 0, 0, 10 / 140))],
    "bg-contain-L-right": [((0, 3, 180, 135), (0, 0, 0, 0))],
    "bg-contain-P-offsets": [((10, 0, 105, 140), (0, 0, 0, 0))],
    "bg-100pct": [((0, 0, 180, 140), (0, 0, 0, 0))],
    "bg-120px-auto": [((50, 40, 120, 90), (0, 0, 0, 0))],
    "bg-auto-200px": [((15, 0, 150, 140), (0, 30 / 200, 0, 30 / 200))],
    "bg-natural-negative": [((0, 0, 180, 140), (50 / 400, 40 / 300, 170 / 400, 120 / 300))],
    "bg-50pct": [((45, 35, 90, 70), (0, 0, 0, 0))],
    "bg-origin-content": [((20, 20, 140, 120), (0, 0, 0, 5 / 14))],
    "bg-clip-padding": [((10, 10, 160, 120), (0, 0, 0, 0))],
    "bg-clip-content": [((15, 15, 150, 110), (45 / 560, 15 / 140, 65 / 560, 15 / 140))],  # at 0% 0% of the padding box
    "bg-radius": [((0, 0, 180, 140), (0, 0, 0, 100 / 240))],                # background-position 0% 0%
    "tile-contain-P": [((0, 0, 37.5, 140), (67.5 / 105, 0, 0, 0)), ((37.5, 0, 105, 140), (0, 0, 0, 0)),
                       ((142.5, 0, 37.5, 140), (0, 0, 67.5 / 105, 0))],
    "tile-repeat-x": [((0, 0, 120, 80), (0, 0, 0, 0)), ((120, 0, 60, 80), (0, 0, 60 / 120, 0))],
    # a rounded frame keeps its box; the bands the source does not reach are negative
    "contain-L-radius": [((0, 0, 180, 140), (0, -3 / 135, 0, -2 / 135))],
    "contain-P-circle": [((0, 0, 140, 140), (-18 / 105, 0, -17 / 105, 0))],
    "none-S-radius": [((0, 0, 180, 140), (-30 / 120, -30 / 80, -30 / 120, -30 / 80))],
    "bg-contain-radius": [((0, 0, 180, 140), (-38 / 105, 0, -37 / 105, 0))],
    "bg-small-circle": [((0, 0, 140, 140), (-10 / 120, -30 / 80, -10 / 120, -30 / 80))],
    "img-bordered": [((4, 4, 172, 132), (2 / 176, 0, 2 / 176, 0))],
    "flip-cover": [((0, 0, 180, 140), (0, 0, 7 / 187, 0))],
    "bleed-canvas": [((0, 0, 300, 140), (0, 182 / 400, 0, 78 / 400))],
    # SVG <image>: preserveAspectRatio's alignment, framed by the viewport
    "svgimg-slice-xMinYMax": [((0, 0, 180, 140), (0, 100 / 240, 0, 0))],
    "svgimg-meet-xMaxYMid": [((0, 0, 180, 140), (-75 / 105, 0, 0, 0))],
    "svgimg-meet-xMinYMin": [((0, 0, 180, 140), (0, 0, 0, -5 / 135))],
    "svgimg-slice-xMidYMin": [((0, 0, 180, 140), (10 / 560, 0, 10 / 560, 0))],
    # a border-box background runs under a side that paints nothing, and stops at one that paints
    "under-transparent": [((10, 10, 160, 130), (0, 0, 0, 250 / 640))],
    "under-dashed": [((6, 6, 168, 128), (0, 0, 0, 96 / 224))],
}

#: Tile count and the area the tiles cover: `repeat` and `round` fill the 180 x 140 painting area once
#: (`round` with 3 x 4 tiles of 60 x 35), `space` leaves 10 px gaps between its 3 x 3 whole tiles.
TILE_COUNTS = {"tile-T": (12, 25200), "tile-round": (12, 25200), "tile-space": (9, 9 * 60 * 40),
               "tile-offset": (20, 25200)}


@pytest.mark.parametrize("name", sorted(EXPECTED))
def test_each_crop_mode_is_the_window_its_css_defines(probed: dict[str, Any], name: str) -> None:
    page, index, _ = BY_NAME[name]
    x, y = _origin(index)
    images = _images(probed["irs"][page], name)
    assert len(images) == len(EXPECTED[name]), f"{name}: {len(images)} pictures"
    for image, ((dx, dy, w, h), crop) in zip(images, EXPECTED[name], strict=False):
        assert (image.box.x, image.box.y, image.box.w, image.box.h) == pytest.approx((x + dx, y + dy, w, h),
                                                                                     abs=1e-3), name
        assert [image.crop[k] for k in "ltrb"] == pytest.approx(list(crop), abs=1e-6), name


@pytest.mark.parametrize("name", sorted(EXPECTED))
def test_the_file_writes_the_measured_crop_as_its_srcrect(probed: dict[str, Any], name: str) -> None:
    """`a:srcRect` is the IR's crop — after the clip only — and the frame is the IR's box: nothing is fitted twice."""
    page, _, _ = BY_NAME[name]
    images = _images(probed["irs"][page], name)
    pictures = _pictures(probed["exact"], page, name)
    assert len(pictures) == len(images)
    for image, picture in zip(images, pictures, strict=False):
        box, crop = image.box, dict(image.crop)
        if image.clip is not None:
            box, crop = _crop_to_clip(box, crop, image.clip)
        written = (picture.crop_left, picture.crop_top, picture.crop_right, picture.crop_bottom)
        assert written == pytest.approx([crop[k] for k in "ltrb"], abs=2e-5), name
        frame = [Emu(v).pt / 0.75 for v in (picture.left, picture.top, picture.width, picture.height)]
        assert frame == pytest.approx([box.x, box.y, box.w, box.h], abs=0.01), name


def test_the_bleeding_image_is_cut_by_its_frame_in_the_file(probed: dict[str, Any]) -> None:
    [picture] = _pictures(probed["exact"], 0, "bleed-P")
    x, y = _origin(BY_NAME["bleed-P"][1])
    assert [Emu(v).pt / 0.75 for v in (picture.left, picture.top, picture.width, picture.height)] == \
        pytest.approx([x, y, 180, 140], abs=0.01)
    assert (picture.crop_left, picture.crop_top, picture.crop_right, picture.crop_bottom) == \
        pytest.approx((40 / 300, 205 / 400, 80 / 300, 55 / 400), abs=2e-5)


@pytest.mark.parametrize("name, expected", sorted(TILE_COUNTS.items()))
def test_background_repeat_draws_one_picture_per_visible_tile(probed: dict[str, Any], name: str,
                                                              expected: tuple[int, int]) -> None:
    page, index, _ = BY_NAME[name]
    x, y = _origin(index)
    images = _images(probed["irs"][page], name)
    assert len(images) == expected[0]
    frame = Box(x, y, probe.FRAME_W, probe.FRAME_H)
    assert sum(image.box.area for image in images) == pytest.approx(expected[1], rel=1e-6)
    assert all(frame.intersect(image.box).area == pytest.approx(image.box.area) for image in images)
    if name == "tile-round":
        assert all((image.box.w, image.box.h) == pytest.approx((60, 35)) for image in images)


def test_too_many_tiles_or_tiles_under_rounded_corners_are_one_captured_picture(probed: dict[str, Any]) -> None:
    ir = probed["irs"][2]
    for name, words in (("tile-many", "draws 1050 tiles"), ("tile-radius", "tiles under rounded corners"),
                        ("tile-dust", "draws 25200 tiles")):
        [image] = _images(ir, name)
        assert image.fit == "fill" and image.crop is None and image.radius == 0
        assert Path(image.src).name.startswith("background-")
        found = [d for d in ir.diagnostics if d.level == "warn" and d.message.startswith("rasterised: ")
                 and words in d.message]
        assert len(found) == 1, name


def test_an_img_keeps_its_own_border_and_an_element_its_own_decoration(probed: dict[str, Any]) -> None:
    """The UA's `overflow: clip` on an <img> clipped its own border away; now it frames only the picture."""
    ir = probed["irs"][2]
    [shape] = [e for e in ir.elements if e.kind == "shape" and e.name == "img-bordered"]
    x, y = _origin(BY_NAME["img-bordered"][1])
    assert shape.stroke["width"] == 4 and shape.clip is None
    assert (shape.box.x, shape.box.y, shape.box.w, shape.box.h) == pytest.approx((x + 2, y + 2, 176, 136))
    assert shape.shadow is not None


def test_a_mirrored_image_keeps_its_flip(probed: dict[str, Any], sources: Path) -> None:
    """`scaleX(-1)` on an <img>: the flip travels in `x-wpf-flip` and lands on the picture's `a:xfrm`."""
    [image] = _images(probed["irs"][2], "flip-cover")
    assert image.extras.get("x-wpf-flip") == {"flipH": True}
    element = Element(kind="image", box=Box(10, 10, 100, 50), src=str(sources / "coords-400x300.png"), fit="cover",
                      crop={"l": 0, "t": 0, "r": 0.25, "b": 0}, extras={"x-wpf-flip": {"flipH": True}})
    picture = _emit_one(element)
    xfrm = picture._element.spPr.find(f"{{{A}}}xfrm")
    assert xfrm.get("flipH") == "1" and xfrm.get("flipV") is None


# ------------------------------------------------------------------ 2 and 3. the browser's own pixels


def _measurable() -> list[str]:
    return [case[0] for _, _, case in CASES
            if case[3] != "repeat" and not case[0].startswith(probe.DIFF_ONLY)]


def _predicted(image: Element, source: str) -> tuple[list[float], list[float]]:
    """The rect of the source's pixels and the window they show (source px), from the IR alone."""
    box, crop = image.box, dict(image.crop)
    if image.clip is not None:
        box, crop = _crop_to_clip(box, crop, image.clip)
    _, width, height = probe.SOURCES[source]
    cl, ct, cr, cb = (crop[k] for k in "ltrb")
    span_x, span_y = 1 - cl - cr, 1 - ct - cb
    x0 = box.x + max(0.0, -cl) / span_x * box.w
    x1 = box.x2 - max(0.0, -cr) / span_x * box.w
    y0 = box.y + max(0.0, -ct) / span_y * box.h
    y1 = box.y2 - max(0.0, -cb) / span_y * box.h
    window = [max(cl, 0.0) * width, max(ct, 0.0) * height, (1 - max(cr, 0.0)) * width, (1 - max(cb, 0.0)) * height]
    if (image.extras or {}).get("x-wpf-flip", {}).get("flipH"):
        window[0], window[2] = window[2], window[0]
    return [x0, y0, x1, y1], window


@pytest.mark.xfail(
    sys.platform == "win32", strict=False,
    reason="Windows Chromium resamples image edges 0.51-0.65 px off the fitted window (Linux Chromium, "
           "which CI and the container run, holds 0.5 px); the tolerance stays 0.5 where it is measured",
)
@pytest.mark.parametrize("name", _measurable())
def test_the_ir_describes_the_window_the_browser_shows(probed: dict[str, Any], name: str) -> None:
    """Fitted to the browser's own pixels: its source coordinate at each predicted edge is the IR's, ± 0.5 px."""
    page, index, (_, source, _, _, _) = BY_NAME[name]
    seen = probe.window(probed["references"][page], index, source)
    assert seen is not None, name
    [image] = _images(probed["irs"][page], name)
    rect, window = _predicted(image, source)
    kx, cx, ky, cy = seen["map"]
    browser = [kx * rect[0] + cx, ky * rect[1] + cy, kx * rect[2] + cx, ky * rect[3] + cy]
    errors = [abs(browser[i] - window[i]) / seen["scale"][i % 2] for i in range(4)]
    assert max(errors) <= 0.5, f"{name}: edges {errors} px off (browser {browser}, IR {window})"
    if not image.radius and not image.circle:              # a rounded frame's own rows are partial
        left, top, right, bottom = probe.cell_of(index)     # the search region cuts a bleed, as it cut `seen`
        expected = [max(rect[0], left), max(rect[1], top), min(rect[2], right), min(rect[3], bottom)]
        measured = [seen["rect"][0], seen["rect"][1], seen["rect"][0] + seen["rect"][2],
                    seen["rect"][1] + seen["rect"][3]]
        assert measured == pytest.approx(expected, abs=0.51), name


#: Differing pixels a case without a single window may keep: tile seams at half pixels (the browser
#: blends two texels across one, the file draws two picture edges), the rotated frame's rounded edge.
DIFF_BUDGET = {"tile-contain-P": 300,
               # the picture stops at the dashed side's padding edge so the dashes stay drawn (the browser
               # also shows the photo between the bottom strip's dashes), and PptxRender draws a dashed
               # line solid — plan §5's named renderer gap: every gap between two dashes (1,392 px measured)
               "under-dashed": 1500}


@pytest.mark.renderer_full
@pytest.mark.parametrize("name", [case[0] for _, _, case in CASES])
def test_the_file_shows_the_window_the_browser_shows(probed: dict[str, Any], probed_renders: list[Path],
                                                     name: str) -> None:
    """PptxRender's rendering of the file (IR through its JSON, as the torture runner) against the browser."""
    page, index, (_, source, _, frame_css, _) = BY_NAME[name]
    reference, render = probed["references"][page], probed_renders[page]
    if name in _measurable():
        truth, seen = probe.window(reference, index, source), probe.window(render, index, source)
        assert truth is not None and seen is not None, name
        rect = max(abs(a - b) for a, b in zip(seen["rect"], truth["rect"], strict=False))
        edges = max(abs(seen["window"][i] - truth["window"][i]) / truth["scale"][i % 2] for i in range(4))
        assert rect <= 0.5 and edges <= 0.5, f"{name}: rect {rect} px, window {edges} px"
    if name != "img-bordered":             # its box-shadow: a no-fill shape's shadow (E's decoration)
        assert probe.ink_difference(reference, render, index) <= DIFF_BUDGET.get(name, 30), name


# ------------------------------------------------------------------------------- 4. the emitter alone


@pytest.fixture
def sources(tmp_path: Path) -> Path:
    probe.write_sources(tmp_path)
    return tmp_path


def _emit_one(element: Element) -> Any:
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    class _Ctx:
        options = EmitOptions()

        def warn(self, *args: Any) -> None:
            raise AssertionError(f"unexpected emit warning {args}")

    element.id = element.id or "e1"
    return add_image(slide.shapes, element, _Ctx())


def test_a_given_crop_is_written_as_given_and_never_fitted_again(sources: Path) -> None:
    """The r3 defect: a `cover` crop the extractor measured, fitted a second time by the emitter."""
    element = Element(kind="image", box=Box(655, 70, 180, 140), src=str(sources / "coords-300x400.png"),
                      fit="cover", crop={"l": 0, "t": 100 / 240, "r": 0, "b": 0})
    picture = _emit_one(element)
    assert (picture.crop_left, picture.crop_top, picture.crop_right, picture.crop_bottom) == \
        pytest.approx((0, 100 / 240, 0, 0), abs=2e-5)          # was t 0.625, b 0.208


def test_a_hand_written_ir_without_a_crop_is_still_fitted(sources: Path) -> None:
    element = Element(kind="image", box=Box(0, 0, 180, 140), src=str(sources / "coords-300x400.png"), fit="cover")
    picture = _emit_one(element)
    assert (picture.crop_top, picture.crop_bottom) == pytest.approx((50 / 240, 50 / 240), abs=2e-5)
    contain = _emit_one(Element(kind="image", box=Box(0, 0, 180, 140), src=str(sources / "coords-300x400.png"),
                                fit="contain"))
    assert (Emu(contain.width).pt / 0.75, Emu(contain.height).pt / 0.75) == pytest.approx((105, 140), abs=0.01)


def test_a_negative_crop_is_written_as_the_empty_band(sources: Path) -> None:
    """PowerPoint's Crop → Fit: the frame keeps its box (and its rounded corners), the band stays empty."""
    element = Element(kind="image", box=Box(0, 0, 180, 140), src=str(sources / "coords-400x300.png"),
                      fit="contain", crop={"l": 0, "t": -3 / 135, "r": 0, "b": -2 / 135}, radius=24.0)
    picture = _emit_one(element)
    source_rect = picture._element.blipFill.find(f"{{{A}}}srcRect")
    assert abs(int(source_rect.get("t")) + round(3 / 135 * 100000)) <= 1
    assert abs(int(source_rect.get("b")) + round(2 / 135 * 100000)) <= 1
    assert picture._element.spPr.find(f"{{{A}}}prstGeom").get("prst") == "roundRect"
    assert (Emu(picture.width).pt / 0.75, Emu(picture.height).pt / 0.75) == pytest.approx((180, 140), abs=0.01)
