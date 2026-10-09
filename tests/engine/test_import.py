"""The deterministic importer: a master `.pptx` in, the v2 manifest (plus backgrounds and assets) out.

Every master here is synthetic. The client master the importer was built against is not in this
repository (plan R6), so the properties that made it hard are rebuilt with python-pptx in
`client_shaped_master`: two slide masters with their own themes, an inverted colour scheme (dk1
white, lt1 charcoal, so `bg1` resolves through `<p:clrMap>` to charcoal), layout relationships in the
reverse of id-list order (so PptxRender's enumeration disagrees with python-pptx's), every layout name
used twice, a dark-backgrounded first layout, and a PNG plus an EMF image on the master.

Repo inputs fail rather than skip; only tests that judge PptxRender's pixels need the service.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Emu

from app.config import engine as config
from app.engine.importer import MasterImportError, Palette, _usage_of, import_master, ingest
from app.engine.manifest import Manifest, ManifestError, Placeholder
from app.engine.renderer import (
    BlankLayoutsRenderer,
    LayoutIdentity,
    RendererError,
    layout_identities,
    map_layouts_by_filename,
)

#: The synthetic brand masters `tests/engine/fixtures/masters/extra/` is committed with (written by
#: `make_test_masters.py --extra-only`), and what they prove.
COMMITTED_EXTRA_MASTERS: dict[str, tuple[int, int]] = {
    "synthetic-89-layouts": (1280, 720),       # 89 layouts on one master, Georgia/Arial theme
    "synthetic-14x8.5in": (1344, 816),         # 14 x 8.5 in, not 16:9; theme font not installed
    "synthetic-portrait-a4": (720, 1040),      # portrait; its only layout inherits the title box
    "synthetic-widescreen": (1280, 720),
}

_LAYOUT_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
_DARK = "1A1A24"


def mean_rgb(png: Path) -> tuple[int, int, int]:
    assert png.exists(), f"{png} was not written"
    with Image.open(png) as image:
        return tuple(int(round(c)) for c in image.convert("RGB").resize((1, 1)).getpixel((0, 0)))  # type: ignore[return-value]


# ------------------------------------------------------------------------- the synthetic masters


def _scheme(theme_xml: str, dk1: str, lt1: str) -> str:
    theme_xml = re.sub(r"<a:dk1>.*?</a:dk1>", f'<a:dk1><a:srgbClr val="{dk1}"/></a:dk1>', theme_xml, flags=re.S)
    return re.sub(r"<a:lt1>.*?</a:lt1>", f'<a:lt1><a:srgbClr val="{lt1}"/></a:lt1>', theme_xml, flags=re.S)


def client_shaped_master(path: Path) -> Path:
    """A 16:9 master with the properties of a real corporate template (see the module docstring)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    master = prs.slide_master.part
    theme = master.part_related_by(RT.THEME)
    theme._blob = _scheme(theme.blob.decode("utf-8"), "FFFFFF", "2E2E38").encode("utf-8")
    png = io.BytesIO()
    Image.new("RGB", (40, 20), "#1A9AFA").save(png, format="PNG")
    for name, content_type, blob in (("logo.png", "image/png", png.getvalue()),
                                     ("logo.emf", "image/x-emf", b"\x01\x00\x00\x00synthetic-emf")):
        master.relate_to(Part(PackURI(f"/ppt/media/{name}"), content_type, prs.part.package, blob), RT.IMAGE)
    buffer = io.BytesIO()
    prs.save(buffer)

    package = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    files = {name: package.read(name) for name in package.namelist()}
    text = {name: data.decode("utf-8") for name, data in files.items() if name.endswith((".xml", ".rels"))}
    numbers = sorted(int(m.group(1)) for name in files
                     if (m := re.match(r"ppt/slideLayouts/slideLayout(\d+)\.xml$", name)))
    count = len(numbers)

    # The first layout ("Title Slide") paints a dark background: a picture that tells layouts apart.
    first = "ppt/slideLayouts/slideLayout1.xml"
    text[first] = re.sub(
        r"(<p:cSld[^>]*>)",
        rf'\1<p:bg><p:bgPr><a:solidFill><a:srgbClr val="{_DARK}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>',
        text[first], count=1)

    # Master 2: a copy of master 1 with its own (light) theme and its own copies of every layout.
    next_id = 2147483648 + count + 1
    master2_id = next_id
    ids = iter(range(next_id + 1, next_id + 1 + count))
    text["ppt/slideMasters/slideMaster2.xml"] = re.sub(
        r'<p:sldLayoutId id="\d+"', lambda _m: f'<p:sldLayoutId id="{next(ids)}"',
        text["ppt/slideMasters/slideMaster1.xml"])
    rels1 = text["ppt/slideMasters/_rels/slideMaster1.xml.rels"]
    rels2 = re.sub(r"slideLayout(\d+)\.xml", lambda m: f"slideLayout{int(m.group(1)) + count}.xml", rels1)
    text["ppt/slideMasters/_rels/slideMaster2.xml.rels"] = rels2.replace("theme1.xml", "theme2.xml")
    text["ppt/theme/theme2.xml"] = _scheme(text["ppt/theme/theme1.xml"], _DARK, "FFFFFF")
    for number in numbers:
        text[f"ppt/slideLayouts/slideLayout{number + count}.xml"] = text[f"ppt/slideLayouts/slideLayout{number}.xml"]
        text[f"ppt/slideLayouts/_rels/slideLayout{number + count}.xml.rels"] = text[
            f"ppt/slideLayouts/_rels/slideLayout{number}.xml.rels"].replace("slideMaster1.xml", "slideMaster2.xml")

    # Master 1's layout relationships in the reverse of id-list order.
    relationships = re.findall(r"<Relationship [^>]*/>", rels1)
    to_layouts = [r for r in relationships if f'Type="{_LAYOUT_REL}"' in r]
    others = [r for r in relationships if r not in to_layouts]
    text["ppt/slideMasters/_rels/slideMaster1.xml.rels"] = (
        rels1[: rels1.index("<Relationship ")] + "".join(to_layouts[::-1] + others) + "</Relationships>")

    pml = "application/vnd.openxmlformats-officedocument"
    overrides = [
        f'<Override PartName="/ppt/slideMasters/slideMaster2.xml" ContentType="{pml}.presentationml.slideMaster+xml"/>',
        f'<Override PartName="/ppt/theme/theme2.xml" ContentType="{pml}.theme+xml"/>',
    ] + [f'<Override PartName="/ppt/slideLayouts/slideLayout{n + count}.xml" '
         f'ContentType="{pml}.presentationml.slideLayout+xml"/>' for n in numbers]
    text["[Content_Types].xml"] = text["[Content_Types].xml"].replace("</Types>", "".join(overrides) + "</Types>")
    text["ppt/_rels/presentation.xml.rels"] = text["ppt/_rels/presentation.xml.rels"].replace(
        "</Relationships>",
        '<Relationship Id="rIdMaster2" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/slideMaster" Target="slideMasters/slideMaster2.xml"/></Relationships>')
    text["ppt/presentation.xml"] = text["ppt/presentation.xml"].replace(
        "</p:sldMasterIdLst>", f'<p:sldMasterId id="{master2_id}" r:id="rIdMaster2"/></p:sldMasterIdLst>')

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name in sorted(set(files) | set(text)):
            out.writestr(name, text[name].encode("utf-8") if name in text else files[name])
    return path


class ColourByRenderIndex:
    """An offline renderer whose layout PNG `NN` is the colour (NN*10, 0, 0): a background joined to
    the wrong layout is then visible without PptxRender."""

    def render_layouts(self, pptx: Path, width: int, out_dir: Path) -> dict[str, Path]:
        out_dir.mkdir(parents=True, exist_ok=True)
        identities = layout_identities(pptx)
        produced = []
        for identity in identities:
            png = out_dir / f"layout-{identity.renderIndex:02d}-x.png"
            Image.new("RGB", (width, round(width * 9 / 16)), (identity.renderIndex * 10, 0, 0)).save(png)
            produced.append(png)
        return map_layouts_by_filename(pptx, produced, identities, out_dir)

    def render(self, pptx: Path, width: int, out_dir: Path) -> list[Path]:
        raise RendererError("not used")

    def verify(self, pptx: Path, references: list[Path], out_dir: Path):  # type: ignore[no-untyped-def]
        raise RendererError("not used")


def _import(master: Path, out_dir: Path) -> Manifest:
    """`import_master` with offline white backgrounds: these tests need a background per layout,
    not a picture of it (the `renderer` tests judge real pixels)."""
    return import_master(master, out_dir, renderer=BlankLayoutsRenderer())


# ------------------------------------------------------------------------------------ fixtures


@pytest.fixture(scope="session")
def client_master(tmp_path_factory: pytest.TempPathFactory) -> Path:
    return client_shaped_master(tmp_path_factory.mktemp("client-shaped") / "master.pptx")


@pytest.fixture(scope="session")
def client_import(tmp_path_factory: pytest.TempPathFactory, client_master: Path) -> tuple[Manifest, Path]:
    out_dir = tmp_path_factory.mktemp("import-client-shaped")
    return import_master(client_master, out_dir, renderer=ColourByRenderIndex()), out_dir


@pytest.fixture(scope="session")
def small_import(tmp_path_factory: pytest.TempPathFactory, masters_dir: Path) -> tuple[Manifest, Path]:
    """`test-4x3.pptx`: the 960x720 canvas, and 40 placeholders that inherit their box."""
    master = masters_dir / "test-4x3.pptx"
    assert master.exists(), f"{master} is missing; the generated test masters are tracked"
    out_dir = tmp_path_factory.mktemp("import-4x3")
    return _import(master, out_dir), out_dir


# ------------------------------------------------------------------------ client-shaped import


def test_client_shaped_import_shape(client_import: tuple[Manifest, Path]) -> None:
    manifest, _ = client_import
    assert manifest.version == 2
    assert manifest.importer == "deterministic"
    assert manifest.sourceHash and len(manifest.sourceHash) == 64
    assert (manifest.slideWidthEmu, manifest.slideHeightEmu) == (12192000, 6858000)
    assert (manifest.canvas_w, manifest.canvas_h) == (1280, 720)
    assert manifest.canvas["pxPerIn"] == config.PX_PER_IN
    assert len(manifest.layouts) == 22
    assert manifest.validate() == []


def test_ids_follow_the_id_lists_not_the_relationship_order(client_import, client_master: Path) -> None:
    """layout-NN is numbered in python-pptx (id-list) order across masters; names repeat across them."""
    manifest, _ = client_import
    assert [(lay.masterIndex, lay.layoutIndex) for lay in manifest.layouts] == (
        [(0, i) for i in range(11)] + [(1, i) for i in range(11)])
    assert [lay.id for lay in manifest.layouts] == [f"layout-{n:02d}" for n in range(1, 23)]
    names = [lay.name for lay in manifest.layouts]
    assert names[:11] == names[11:] and len(set(names)) == 11, "every name is used by two layouts"
    # PptxRender enumerates master 1's layouts backwards: its first render is the last layout.
    first_render = layout_identities(client_master)[0]
    assert (first_render.masterIndex, first_render.layoutIndex) == (0, 10)


def test_a_reimport_keeps_the_ids_it_is_given(client_import, client_master: Path, tmp_path: Path) -> None:
    """Re-importing a project's master adopts its layout ids, so `slide.layoutId` still means something."""
    manifest, _ = client_import
    adopted = {(lay.masterIndex, lay.layoutIndex): lay.id for lay in manifest.layouts}
    adopted[(1, 0)], adopted[(1, 1)] = adopted[(1, 1)], adopted[(1, 0)]  # a project that had them swapped
    again = import_master(client_master, tmp_path, render=False, layout_ids=adopted)
    assert {(lay.masterIndex, lay.layoutIndex): lay.id for lay in again.layouts} == adopted
    assert [(lay.id, lay.name) for lay in again.layouts if lay.masterIndex == 0] == \
        [(lay.id, lay.name) for lay in manifest.layouts if lay.masterIndex == 0]


def test_each_master_has_its_own_theme(client_import) -> None:
    manifest, _ = client_import
    assert [m.index for m in manifest.masters] == [0, 1]
    for master in manifest.masters:
        assert master.partName and master.partName.startswith("/ppt/slideMasters/")
        colors = master.theme["colors"]
        assert set(colors) >= {"dk1", "lt1", "accent1", "bg1", "tx1"}
        assert all(value.startswith("#") and len(value) == 7 for value in colors.values())
        assert master.theme["fonts"]["major"] and master.theme["fonts"]["minor"]
    # The masters genuinely differ: master 0 is the inverted dark scheme, master 1 the light one.
    assert manifest.master(0).theme["colors"]["dk1"] == "#FFFFFF"
    assert manifest.master(1).theme["colors"]["dk1"] == f"#{_DARK}"
    assert manifest.theme == manifest.master(0).theme


def test_colour_map_is_applied(client_import) -> None:
    """`<p:clrMap>` sends bg1 to the theme's lt1; this scheme inverts, so bg1 must be charcoal not white."""
    manifest, _ = client_import
    colors = manifest.master(0).theme["colors"]
    assert colors["lt1"] == "#2E2E38" and colors["dk1"] == "#FFFFFF"
    assert colors["bg1"] == "#2E2E38", "bg1 must resolve through clrMap to lt1"
    assert colors["tx1"] == "#FFFFFF", "tx1 must resolve through clrMap to dk1"


def test_title_placeholder_geometry_and_style(client_import) -> None:
    manifest, _ = client_import
    layout = manifest.layout("layout-06")
    assert layout.name == "Title Only"
    title = layout.placeholder("title")
    assert title is not None and title.type == "title"
    assert title.idx == 0, "the deterministic import must record <p:ph idx>"
    assert title.emu == {"x": 457200, "y": 274638, "w": 8229600, "h": 1143000}
    for got, want in zip((title.x, title.y, title.w, title.h), (48, 28.833, 864, 120), strict=True):
        assert abs(got - want) <= 0.001
    assert title.emu["x"] == round(title.x * config.EMU_PER_PX)
    style = title.style or {}
    assert style["sizePt"] == pytest.approx(44.0)
    assert style["color"] == "#FFFFFF", "tx1 through the inverted scheme"
    assert style["align"] == "left", "a title-only layout's title defaults left"
    assert style["anchor"] == "middle" and style["autofit"] == "shrink"
    assert style["insetsPx"] == {"l": 9.6, "t": 4.8, "r": 9.6, "b": 4.8}
    assert style["boxFrom"] == "master:title[0]", "the layout declares no xfrm: inherited from the master"


def test_backgrounds_are_joined_on_partname(client_import, client_master: Path) -> None:
    """The renderer's first layout is master 1's last; only the partName join puts it on layout-11."""
    manifest, out_dir = client_import
    by_part = {i.partName: i for i in layout_identities(client_master)}
    for layout in manifest.layouts:
        assert layout.background == f"{layout.id}.png"
        identity: LayoutIdentity = by_part[layout.partName or ""]
        assert mean_rgb(out_dir / "layouts" / layout.background)[0] == identity.renderIndex * 10, layout.id
    assert mean_rgb(out_dir / "layouts" / "layout-11.png")[0] == 10, "render 1 belongs to layout-11"
    assert mean_rgb(out_dir / "layouts" / "layout-01.png")[0] == 110, "a position join would give 10"
    index = json.loads((out_dir / "layouts" / "index.json").read_text(encoding="utf-8"))
    assert [entry["id"] for entry in index] == [lay.id for lay in manifest.layouts]


@pytest.mark.renderer
def test_backgrounds_are_joined_on_partname_through_pptxrender(client_master: Path, pptx_renderer, tmp_path: Path) -> None:
    """The same join with PptxRender's real pictures: the dark first layout lands on layout-01 and -12."""
    manifest = import_master(client_master, tmp_path, renderer=pptx_renderer)
    layouts = tmp_path / "layouts"
    assert max(mean_rgb(layouts / "layout-01.png")) < 80, "layout-01 paints the dark background"
    assert max(mean_rgb(layouts / "layout-12.png")) < 80, "so does its copy on master 2"
    assert min(mean_rgb(layouts / "layout-02.png")) > 200, "the next layout does not"
    for layout in manifest.layouts:
        assert layout.background
        with Image.open(layouts / layout.background) as image:
            assert image.size == (manifest.canvas_w, manifest.canvas_h)


def test_text_styles_and_bullets(client_import) -> None:
    manifest, _ = client_import
    assert set(manifest.textStyles) == {"title", "body", "other"}
    title_lvl1 = manifest.textStyles["title"]["levels"][0]
    assert title_lvl1["level"] == 0 and title_lvl1["align"] == "center"
    assert title_lvl1["color"] == "#FFFFFF", "tx1 in titleStyle resolves through the colour map"
    assert len(manifest.bullets) >= 5
    first = manifest.bullets[0]
    assert first["level"] == 0 and first["kind"] == "char" and first["font"] == "Arial"
    assert first["marLEmu"] == 342900 and first["indentEmu"] == -342900


def test_assets_and_report(client_import) -> None:
    manifest, out_dir = client_import
    report = json.loads((out_dir / "import-report.json").read_text(encoding="utf-8"))
    assert report["layouts"] == 22 and report["masters"] == 2
    assert report["validation"] == [] and report["orphanLayouts"] == []
    assert report["elapsedS"] < 20

    written = {entry["file"] for entry in report["assets"]}
    assert written == {"asset-logo.png", "asset-logo.emf"}
    for name in written:
        assert (out_dir / "assets" / name).exists()
    # Only browser-usable formats reach the design prompt; an EMF logo must not.
    assert manifest.assets == sorted(entry["file"] for entry in report["assets"] if entry.get("webRenderable"))
    assert manifest.assets == ["asset-logo.png"]


def test_round_trips_through_disk(client_import, tmp_path: Path) -> None:
    manifest, _ = client_import
    path = manifest.save(tmp_path / "manifest.json")
    reloaded = Manifest.load(path)
    assert reloaded.to_json() == manifest.to_json()
    title = reloaded.layout("layout-06").placeholder("title")
    assert title is not None and title.idx == 0
    assert reloaded.layout_by_part("/ppt/slideLayouts/slideLayout6.xml").id == "layout-06"
    assert reloaded.layout_by_part("/ppt/slideLayouts/slideLayout17.xml").id == "layout-17"


def test_a_v1_manifest_is_refused(tmp_path: Path) -> None:
    """The legacy (pre-v2) conversion was dropped in the port (D14): an old manifest is an error."""
    path = tmp_path / "project.json"
    path.write_text(json.dumps({"master": {"canvas": {"w": 1280, "h": 720}, "layouts": []}}), encoding="utf-8")
    with pytest.raises(ManifestError, match="not a v2 manifest"):
        Manifest.load(path)


# ------------------------------------------------------------------------------- other masters


def test_4x3_canvas(small_import) -> None:
    manifest, _ = small_import
    assert (manifest.slideWidthEmu, manifest.slideHeightEmu) == (9144000, 6858000)
    assert (manifest.canvas_w, manifest.canvas_h) == (960, 720)
    assert manifest.validate() == []


def test_placeholders_inherit_their_box_from_the_master(small_import) -> None:
    """The generated masters leave 40 of 58 layout placeholders without an `<a:xfrm>`."""
    manifest, _ = small_import
    everything = [ph for lay in manifest.layouts for ph in lay.placeholders]
    inherited = [ph for ph in everything if (ph.style or {})["boxFrom"].startswith("master")]
    assert len(inherited) == 40
    assert not [ph for ph in everything if (ph.style or {})["boxFrom"] == "missing"]
    for placeholder in inherited:
        assert placeholder.w > 0 and placeholder.h > 0, f"{placeholder.type} inherited an empty box"
        assert 0 <= placeholder.x <= manifest.canvas_w
        assert 0 <= placeholder.y <= manifest.canvas_h
    # A layout numbers date/footer/slide-number 10/11/12 where the master numbers them 2/3/4;
    # matching on idx alone would leave them at 0x0 (or steal another zone's box).
    furniture = [ph for ph in inherited if ph.type == "sldNum"]
    assert furniture and all(ph.idx == 12 for ph in furniture)
    assert all((ph.style or {})["boxFrom"] == "master:sldNum[4]" for ph in furniture)


def test_every_committed_extra_master_is_present() -> None:
    found = {path.stem for path in config.extra_masters()}
    missing = sorted(set(COMMITTED_EXTRA_MASTERS) - found)
    assert not missing, f"{config.EXTRA_MASTERS} is missing committed masters: {missing}"


#: Layouts per committed extra master — one slide master each.
EXTRA_MASTER_LAYOUTS = {"synthetic-89-layouts": 89, "synthetic-14x8.5in": 8,
                        "synthetic-portrait-a4": 1, "synthetic-widescreen": 10}


def test_committed_extra_masters_keep_their_shape() -> None:
    """The properties the synthetic masters stand in for, so a regenerated set cannot quietly lose one."""
    for stem, layouts in EXTRA_MASTER_LAYOUTS.items():
        deck = Presentation(str(config.EXTRA_MASTERS / f"{stem}.pptx"))
        assert len(deck.slide_masters) == 1, stem
        assert len(deck.slide_layouts) == layouts, stem
        assert deck.core_properties.author == "Slide Studio test fixtures", stem


@pytest.mark.parametrize("deck", config.extra_masters(), ids=lambda path: path.stem)
def test_extra_master_imports_and_validates(deck: Path, tmp_path: Path) -> None:
    """"It should be able to do a great job on any slide master": the four synthetic brand masters."""
    manifest = _import(deck, tmp_path)
    assert manifest.validate(project_dir=tmp_path) == []
    assert manifest.importer == "deterministic" and manifest.has_part_names
    assert manifest.layouts, f"{deck.name} produced no layouts"
    assert manifest.fonts["major"] and manifest.fonts["minor"]

    expected = COMMITTED_EXTRA_MASTERS.get(deck.stem)
    if expected:
        assert (manifest.canvas_w, manifest.canvas_h) == expected

    identities = {identity.partName for identity in layout_identities(deck)}
    assert {lay.partName for lay in manifest.layouts} <= identities
    for layout in manifest.layouts:
        assert layout.background and (tmp_path / "layouts" / layout.background).exists()
        for placeholder in layout.placeholders:
            assert placeholder.idx is not None
            assert placeholder.emu is not None
            # px is the EMU box / 9525 rounded to 3 decimals, so the round trip is exact to
            # within half a thousandth of a pixel: about five EMU.
            for key, value in (("x", placeholder.x), ("y", placeholder.y),
                               ("w", placeholder.w), ("h", placeholder.h)):
                assert abs(placeholder.emu[key] - value * config.EMU_PER_PX) <= 5
            assert (placeholder.style or {}).get("boxFrom")

    report = json.loads((tmp_path / "import-report.json").read_text(encoding="utf-8"))
    assert report["validation"] == []


def test_a_missing_master_fails_loudly(tmp_path: Path) -> None:
    with pytest.raises(MasterImportError):
        import_master(tmp_path / "nope.pptx", tmp_path / "out", renderer=BlankLayoutsRenderer())


def test_rendering_without_a_renderer_is_refused_not_skipped(masters_dir: Path, tmp_path: Path) -> None:
    """No process-wide renderer to fall back on: `render=True` with none configured raises."""
    with pytest.raises(RendererError, match="needs a renderer"):
        import_master(masters_dir / "test-16x9.pptx", tmp_path)
    structure_only = import_master(masters_dir / "test-16x9.pptx", tmp_path / "s", render=False)
    assert structure_only.layouts and all(lay.background is None for lay in structure_only.layouts)


# --------------------------------------------------------------------------- derived fields


@pytest.mark.parametrize(
    "types, expected",
    [
        ([], "blank"),
        (["title"], "title only"),
        (["title", "body"], "title + body"),
        (["body", "title"], "title + body"),
        (["ctrTitle", "subTitle"], "centred title + subtitle"),
        (["title", "body", "body"], "title + 2 body"),
        (["title", "dt", "ftr", "sldNum"], "title only"),
        (["title", "pic", "tbl"], "title + picture + table"),
    ],
)
def test_usage_is_derived_from_the_zones(types: list[str], expected: str) -> None:
    assert _usage_of([Placeholder(type=t) for t in types]) == expected


def test_palette_applies_luminance_transforms() -> None:
    palette = Palette({"accent1": "#4F81BD", "lt1": "#FFFFFF"}, {"bg1": "lt1"})
    from lxml import etree

    plain = etree.fromstring('<schemeClr xmlns="x" val="accent1"/>')
    assert palette.resolve(plain) == "#4F81BD"
    lighter = etree.fromstring(
        '<schemeClr xmlns="x" val="accent1"><lumMod xmlns="x" val="60000"/>'
        '<lumOff xmlns="x" val="40000"/></schemeClr>'
    )
    assert palette.resolve(lighter) != "#4F81BD"
    assert palette.resolve(etree.fromstring('<schemeClr xmlns="x" val="phClr"/>')) is None
    assert palette.resolve(etree.fromstring('<srgbClr xmlns="x" val="1A9AFA"/>')) == "#1A9AFA"
    assert palette.resolve(etree.fromstring('<sysClr xmlns="x" val="window" lastClr="FFFFFF"/>')) == "#FFFFFF"


# ------------------------------------------------------------------------------------- ingest


def test_ingest_lands_an_importers_output_in_a_project(client_import, tmp_path: Path) -> None:
    manifest, out_dir = client_import
    project_dir = tmp_path / "prj"
    landed = ingest(project_dir, json.loads(json.dumps(manifest.to_json())), out_dir, importer="reimport")
    assert landed.importer == "reimport"
    assert landed.assets == ["asset-logo.png"], \
        "the EMF part is copied but must not reach the design prompt's asset list"
    assert (project_dir / "manifest.json").exists()
    assert all((project_dir / "assets" / name).exists() for name in landed.assets)
    for layout in landed.layouts:
        assert layout.background
        background = project_dir / "layouts" / layout.background
        with Image.open(background) as image:
            assert image.size == (landed.canvas_w, landed.canvas_h)
    assert Manifest.load(project_dir / "manifest.json").layout("layout-06").placeholder("title") is not None


def test_ingest_refuses_a_manifest_whose_background_is_missing(client_import, tmp_path: Path) -> None:
    """An import that lost a layout render fails here, not as a blank backdrop later."""
    manifest, out_dir = client_import
    data = json.loads(json.dumps(manifest.to_json()))
    data["layouts"][0]["background"] = "not-supplied.png"
    with pytest.raises(ManifestError, match="not-supplied.png"):
        ingest(tmp_path / "prj", data, [out_dir / "layouts", out_dir / "assets"])


def test_content_and_title_only_titles_default_left_over_a_centred_title_style(tmp_path: Path) -> None:
    """A master whose `titleStyle` centres titles: content and title-only layouts that set no
    alignment of their own get a left title; a layout's own alignment, and other layouts, keep theirs."""
    from lxml import etree

    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    p = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    prs = Presentation()
    prs.slide_master.element.find(f"{p}txStyles/{p}titleStyle/{a}lvl1pPr").set("algn", "ctr")
    layouts = {layout.element.get("type"): layout for layout in prs.slide_layouts}
    for layout in layouts.values():  # start every layout title from "says nothing"
        for lst in layout.element.iter(f"{a}lstStyle"):
            for lvl in lst.iter(f"{a}lvl1pPr"):
                lvl.attrib.pop("algn", None)
    body = layouts["titleOnly"].placeholders[0].text_frame._txBody
    lst = body.find(f"{a}lstStyle")
    if lst is None:
        lst = etree.SubElement(body, f"{a}lstStyle")
        body.insert(1, lst)
    etree.SubElement(lst, f"{a}lvl1pPr").set("algn", "r")
    prs.save(tmp_path / "master.pptx")

    manifest = _import(tmp_path / "master.pptx", tmp_path / "out")
    align = {lay.name: (lay.placeholder("title").style or {})["align"]  # type: ignore[union-attr]
             for lay in manifest.layouts if lay.placeholder("title")}
    assert align["Title and Content"] == "left", "content layout with no alignment of its own"
    assert align["Title Only"] == "right", "a layout's own alignment wins"
    assert align["Two Content"] == "center", "other layouts still follow titleStyle"
