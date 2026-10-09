"""Render-check package R3 (2026-09-29): the theme's major font off the headings, and the title.

The render check's master is major Georgia / minor Arial. Its decks set card headers, big numbers and banners
in Georgia on plain `div`s; `font_forcing_css` gave the major font only to `h1–h6` and the title
placeholder, so every one of them was measured and exported in Arial (defect 7). And a title left in
normal flow landed at x = 0 across the canvas and never reached its placeholder; the render check's
divider titles came out in capitals, which a master style with `cap="all"` does to any run that does
not say `cap="none"` (defect 8). Plan: `docs/engine/rendercheck/00-PLAN.md`.

These tests run on the `test-16x9` master with the manifest's theme fonts set to Georgia / Arial, as
the render-check master's are; both faces ship with Windows.
"""
from __future__ import annotations

import copy
import re
import zipfile
from pathlib import Path
from typing import Any

import pytest
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.util import Emu

from app.config import engine as config
from app.engine.classify.placeholders import map_placeholders
from app.engine.emit import pptx as emitter
from app.engine.emit.text import face_for_exact
from app.engine.extract.html import render_reference, title_zone_of
from app.engine.ir import IR, Box, Canvas, Element, Slide
from app.engine.manifest import Manifest
from app.engine.reports import EmitOptions
from app.engine.verify.lint import lint
from tests.engine.helpers import extract_html

#: The tests that read which family the browser drew need Georgia itself installed (Windows and
#: macOS have it; the Linux CI image has no metric twin for it, so the browser draws a fallback).
needs_georgia = pytest.mark.skipif(face_for_exact("Georgia", 400, False) is None,
                                   reason="needs the Georgia font installed (not in the Linux CI font set)")

LAYOUT = "layout-06"                     # Title Only: title zone (48, 28.8, 864, 120)
PAGE = ("<!doctype html><html><head><meta charset='utf-8'><style>"
        "html,body{{margin:0;width:1280px;height:720px;overflow:hidden;background:transparent;"
        "font-family:Arial,sans-serif;color:#021D44}}{css}</style></head><body>{body}</body></html>")


@pytest.fixture(scope="module")
def georgia(masters_dir: Path) -> Manifest:
    """test-16x9's manifest with the render-check master's theme fonts, deck-wide and on every master."""
    manifest = copy.deepcopy(Manifest.load(masters_dir / "test-16x9.manifest.json"))
    fonts = {"major": "Georgia", "minor": "Arial"}
    manifest.theme = {**(manifest.theme or {}), "fonts": dict(fonts)}
    for master in manifest.masters:
        master.theme = {**(master.theme or {}), "fonts": dict(fonts)}
    return manifest


@pytest.fixture(scope="module")
def shouting_master(masters_dir: Path, tmp_path_factory) -> Path:
    """test-16x9.pptx whose title style says bold, capitals, underline, strike and wide tracking —
    what a master can set on a title that the design never asked for."""
    source = masters_dir / "test-16x9.pptx"
    target = tmp_path_factory.mktemp("master") / "shouting.pptx"
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/slideMasters/slideMaster1.xml":
                xml = data.decode("utf-8")
                start = xml.index("<p:titleStyle>")
                head, tail = xml[:start], xml[start:]
                tail = tail.replace('<a:defRPr sz="4400" kern="1200">',
                                    '<a:defRPr sz="4400" kern="1200" b="1" u="sng" strike="sngStrike" '
                                    'cap="all" spc="600">', 1)
                assert 'cap="all"' in tail, "the master's title style moved; update the patch"
                data = (head + tail).encode("utf-8")
            zout.writestr(item, data)
    return target


def _write(tmp_path: Path, name: str, body: str, css: str = "") -> Path:
    html = tmp_path / f"{name}.html"
    html.write_text(PAGE.format(css=css, body=body), encoding="utf-8")
    (tmp_path / "assets").mkdir(exist_ok=True)
    return html


def _extract(html: Path, manifest: Manifest, layout_id: str = LAYOUT) -> IR:
    ir = extract_html(html, manifest, layout_id, html.parent / "assets", slide_id=html.stem, title=html.stem)
    return map_placeholders(ir, manifest.layout(layout_id))


def _text(element: Element) -> str:
    return "".join(run.get("text", "") for paragraph in element.paragraphs or []
                   for line in paragraph.get("lines") or [] for run in line.get("runs") or [])


def _runs(element: Element) -> list[dict]:
    return [run for paragraph in element.paragraphs or [] for line in paragraph.get("lines") or []
            for run in line.get("runs") or []]


def _by_text(ir: IR, starts: str) -> Element:
    found = [e for e in ir.elements if e.kind == "text" and _text(e).startswith(starts)]
    assert len(found) == 1, [_text(e) for e in ir.elements if e.kind == "text"]
    return found[0]


def _emit(ir: IR, manifest: Manifest, master: Path, deck: Path) -> Path:
    report = emitter.emit([ir], manifest, master, deck, EmitOptions())
    assert deck.exists(), report
    return deck


def _run_props(deck: Path, text: str) -> list[dict[str, str]]:
    """Every `a:rPr` (attributes plus its `latin` typeface) of the runs whose text starts `text`."""
    xml = zipfile.ZipFile(deck).read("ppt/slides/slide1.xml").decode("utf-8")
    out = []
    for run in re.findall(r"<a:r>(.*?)</a:r>", xml, re.S):
        said = "".join(re.findall(r"<a:t>([^<]*)</a:t>", run))
        if not said.lower().startswith(text.lower()):
            continue
        attributes = dict(re.findall(r'(\w+)="([^"]*)"', re.search(r"<a:rPr\b[^>]*>", run).group(0)))
        latin = re.search(r'<a:latin typeface="([^"]*)"', run)
        attributes["latin"] = latin.group(1) if latin else ""
        out.append(attributes)
    assert out, f"no run starting {text!r} in the slide"
    return out


# ------------------------------------------------------------------------------ defect 7: the fonts


@needs_georgia
def test_a_div_in_the_major_font_is_measured_and_exported_in_it(tmp_path: Path, georgia: Manifest,
                                                                masters_dir: Path):
    """A Georgia card header is Georgia in the IR and `+mj-lt` in the file; an Arial child inside it,
    a body line, a second-family Georgia and a made-up first family stay minor. The variable counts."""
    css = (".card{position:absolute;left:80px;top:200px;width:520px;font-family:Georgia,'Times New Roman',serif;"
           "font-size:24px;font-weight:700}.card .note{font-family:Arial,sans-serif;font-size:14px;font-weight:400}"
           ".plain{position:absolute;left:80px;top:420px;width:520px;font-size:18px}"
           ".var{position:absolute;left:700px;top:200px;width:500px;font-family:var(--engine-font-major),serif;font-size:20px}"
           ".second{position:absolute;left:700px;top:300px;width:500px;font-family:'No Such Face 4471',Georgia;font-size:20px}")
    html = _write(tmp_path, "major", css=css, body=(
        '<div class="card">Card header in Georgia<div class="note">A note in Arial inside it</div></div>'
        '<div class="plain">Body copy in the minor font</div>'
        '<div class="var">Through the variable</div>'
        '<div class="second">Georgia only as the fallback</div>'))
    ir = _extract(html, georgia)
    runs = [run for element in ir.elements if element.kind == "text" for run in _runs(element)]
    fonts = {starts: {run.get("font") for run in runs if run.get("text", "").startswith(starts)}
             for starts in ("Card header", "A note", "Body copy", "Through the", "Georgia only")}
    assert fonts == {"Card header": {"Georgia"}, "A note": {"Arial"}, "Body copy": {"Arial"},
                     "Through the": {"Georgia"}, "Georgia only": {"Arial"}}, fonts

    deck = _emit(ir, georgia, masters_dir / "test-16x9.pptx", tmp_path / "major.pptx")
    assert {p["latin"] for p in _run_props(deck, "Card header")} == {"+mj-lt"}
    assert {p["latin"] for p in _run_props(deck, "Through the")} == {"+mj-lt"}
    assert {p["latin"] for p in _run_props(deck, "A note")} == {"+mn-lt"}
    assert {p["latin"] for p in _run_props(deck, "Body copy")} == {"+mn-lt"}


@needs_georgia
def test_a_pseudo_element_in_the_major_font_keeps_it(tmp_path: Path, georgia: Manifest):
    """`::before` inherits Georgia from its host and is rebuilt in it; the host's Arial sibling is not."""
    css = (".h{position:absolute;left:80px;top:200px;width:600px;font-family:Georgia,serif;font-size:22px}"
           ".h::before{content:'01 '}.b{position:absolute;left:80px;top:300px;width:600px;font-size:22px}"
           ".b::before{content:'02 '}")
    html = _write(tmp_path, "pseudo", css=css, body='<div class="h">Header</div><div class="b">Body</div>')
    ir = _extract(html, georgia)
    texts = {_text(e): {run.get("font") for run in _runs(e)} for e in ir.elements if e.kind == "text"}
    assert texts == {"01 Header": {"Georgia"}, "02 Body": {"Arial"}}, texts


# ---------------------------------------------------------------------------- defect 8: the title


@needs_georgia
def test_a_static_title_is_pinned_to_the_zone_and_mapped(tmp_path: Path, georgia: Manifest, masters_dir: Path):
    """An `h1[data-placeholder=title]` left in flow — with the UA's margins, after an unhidden notes
    block — is moved to the title zone before measurement, maps to the title placeholder by its
    attribute, and the placeholder is written at the zone's x, y and width."""
    zone = georgia.layout(LAYOUT).placeholder("title").box
    html = _write(tmp_path, "static", css="h1{font-family:Georgia,serif;font-size:28px;line-height:34px}", body=(
        '<aside class="notes">Speaker notes someone forgot to hide, long enough to push the title down.</aside>'
        '<h1 data-placeholder="title">A title nobody positioned</h1>'
        '<div style="position:absolute;left:48px;top:400px;width:600px">Body</div>'))
    ir = _extract(html, georgia)
    title = _by_text(ir, "A title")
    assert title.placeholder == {"type": "title", "idx": 0}
    assert title.box.x == pytest.approx(zone.x, abs=0.5) and title.box.y == pytest.approx(zone.y, abs=0.5)
    assert title.box.w == pytest.approx(zone.w, abs=1.0)
    assert {run.get("font") for run in _runs(title)} == {"Georgia"}
    assert any("not positioned; placed at layout-06's title zone" in d.message for d in ir.diagnostics)

    deck = _emit(ir, georgia, masters_dir / "test-16x9.pptx", tmp_path / "static.pptx")
    slide = Presentation(str(deck)).slides[0]
    [shape] = [s for s in slide.placeholders if s.placeholder_format.type is not None
               and s.placeholder_format.idx == 0]
    assert Emu(shape.left).inches * 96 == pytest.approx(zone.x, abs=1)
    # The emitter places the frame so PowerPoint's first baseline lands on the browser's (`_place_box`):
    # a few px from the measured top, as for every text box.
    assert Emu(shape.top).inches * 96 == pytest.approx(zone.y, abs=4)
    # …and wide enough never to re-wrap: the zone's width plus the emitter's slack (`TEXT_BOX_WIDTH_SLACK`).
    assert zone.w - 1 <= Emu(shape.width).inches * 96 <= zone.w * config.TEXT_BOX_WIDTH_SLACK + 1


def test_a_positioned_title_keeps_its_rect(tmp_path: Path, georgia: Manifest, masters_dir: Path):
    """The author placed it: nothing is pinned, the attribute still maps it, the rect is the design's."""
    html = _write(tmp_path, "positioned", body=(
        '<h1 data-placeholder="title" style="position:absolute;left:300px;top:380px;width:500px;margin:0;'
        'font-size:30px">Placed by the author</h1>'))
    ir = _extract(html, georgia)
    title = _by_text(ir, "Placed by")
    assert title.placeholder == {"type": "title", "idx": 0}
    assert (round(title.box.x), round(title.box.y), round(title.box.w)) == (300, 380, 500)
    assert not any("title zone" in d.message for d in ir.diagnostics)

    deck = _emit(ir, georgia, masters_dir / "test-16x9.pptx", tmp_path / "positioned.pptx")
    [shape] = [s for s in Presentation(str(deck)).slides[0].placeholders if s.placeholder_format.idx == 0]
    assert round(Emu(shape.left).inches * 96) == 300 and Emu(shape.top).inches * 96 == pytest.approx(380, abs=4)


def test_a_title_inside_a_positioned_card_is_left_alone(tmp_path: Path, georgia: Manifest):
    html = _write(tmp_path, "carded", body=(
        '<div style="position:absolute;left:600px;top:300px;width:400px">'
        '<h2 data-placeholder="title" style="margin:0;font-size:24px">Inside a card</h2></div>'))
    title = _by_text(_extract(html, georgia), "Inside a card")
    assert (round(title.box.x), round(title.box.y)) == (600, 300)


def test_the_reference_pins_the_title_where_the_export_does(tmp_path: Path, georgia: Manifest,
                                                             masters_dir: Path):
    """`render_reference(title_zone=…)` draws a static title exactly as if it had been positioned at
    the zone — the gate compares the export with the same picture."""
    zone = title_zone_of(georgia, LAYOUT)
    layout_png = masters_dir / "test-16x9" / "layouts" / f"{LAYOUT}.png"
    css = "h1{font-family:Georgia,serif;font-size:28px;line-height:34px;margin:0}"
    static = _write(tmp_path, "ref-static", css=css, body='<h1 data-placeholder="title">Same place</h1>')
    placed = _write(tmp_path, "ref-placed", css=css, body=(
        f'<h1 data-placeholder="title" style="position:absolute;left:{round(zone.x, 3)}px;'
        f'top:{round(zone.y, 3)}px;width:{round(zone.w, 3)}px">Same place</h1>'))
    fonts = dict(georgia.fonts)
    pinned = render_reference(static, layout_png, tmp_path / "pinned.png", assets_dir=tmp_path / "assets",
                              fonts=fonts, title_zone=zone)
    wanted = render_reference(placed, layout_png, tmp_path / "wanted.png", assets_dir=tmp_path / "assets",
                              fonts=fonts, title_zone=zone)
    loose = render_reference(static, layout_png, tmp_path / "loose.png", assets_dir=tmp_path / "assets",
                             fonts=fonts)
    with Image.open(pinned) as a, Image.open(wanted) as b, Image.open(loose) as c:
        assert ImageChops.difference(a.convert("RGB"), b.convert("RGB")).getbbox() is None
        assert ImageChops.difference(a.convert("RGB"), c.convert("RGB")).getbbox() is not None


def test_runs_state_every_property_a_title_style_could_set(tmp_path: Path, georgia: Manifest,
                                                           shouting_master: Path):
    """Under a title style of bold + capitals + underline + strike + tracking, a plain run writes
    `b="0" u="none" strike="noStrike" cap="none" spc="0"`; an uppercase one still says `cap="all"`."""
    html = _write(tmp_path, "caps", body=(
        '<h1 data-placeholder="title" style="position:absolute;left:48px;top:29px;width:864px;margin:0;'
        'font-size:30px;font-weight:400">Mixed case title</h1>'
        '<div style="position:absolute;left:48px;top:400px;width:600px;text-transform:uppercase">shouted</div>'))
    ir = _extract(html, georgia)
    deck = _emit(ir, georgia, shouting_master, tmp_path / "caps.pptx")
    master = zipfile.ZipFile(deck).read("ppt/slideMasters/slideMaster1.xml").decode("utf-8")
    assert 'cap="all"' in master[master.index("<p:titleStyle>"):], "the style the runs must override"
    [title] = _run_props(deck, "Mixed case")
    assert {k: title.get(k) for k in ("b", "u", "strike", "cap")} == {
        "b": "0", "u": "none", "strike": "noStrike", "cap": "none"}, title
    # `spc` is always written: the width lock's own few hundredths of a point, or "0", never the style's 600.
    assert abs(int(title["spc"])) < 50, title
    assert {p["cap"] for p in _run_props(deck, "shouted")} == {"all"}


# ------------------------------------------------------------------------------------ the classifier


def _text_element(text: str, box: Box, size: float, **extras) -> Element:
    run = {"text": text, "font": "Arial", "sizePx": size, "weight": 400, "italic": False, "color": "000000",
           "alpha": 1, "underline": False, "strike": False, "letterSpacingPx": 0}
    return Element(kind="text", box=box, paragraphs=[{"lines": [{"box": box.to_json(), "runs": [run]}]}],
                   extras=dict(extras))


def test_one_title_host_with_a_chip_is_one_request(georgia: Manifest):
    """The words and the chip of one `data-placeholder` element: the larger type takes the title, the
    chip stays a text box, and nothing warns. A second host asking for the title is refused aloud."""
    host = {"x-wp4-placeholder": "title", "x-wp4-placeholder-host": "body > h1"}
    chip = _text_element("+12%", Box(700, 40, 60, 20), 12, **host)
    words = _text_element("Revenue grew", Box(48, 36, 600, 40), 30, **host)
    other = _text_element("Another title", Box(48, 300, 600, 40), 30,
                          **{"x-wp4-placeholder": "title", "x-wp4-placeholder-host": "body > div"})
    ir = IR(canvas=Canvas(1280, 720), slide=Slide(id="s", title="s", layoutId=LAYOUT),
            elements=[chip, words, other])
    mapped = map_placeholders(ir, georgia.layout(LAYOUT))
    placed = [(_text(e), e.placeholder) for e in mapped.elements]
    assert placed == [("+12%", None), ("Revenue grew", {"type": "title", "idx": 0}), ("Another title", None)]
    warned = [d.message for d in mapped.diagnostics if d.level == "warn"]
    assert len(warned) == 1 and "no free title placeholder" in warned[0]


# ------------------------------------------------------------------------------------------- lint


def test_lint_names_the_zone_of_an_unpositioned_title(georgia: Manifest):
    static = PAGE.format(css="", body='<h1 data-placeholder="title">Loose</h1>')
    (finding,) = [f for f in lint(static, georgia, layout_id=LAYOUT).findings if f.rule == "title-unpositioned"]
    assert finding.level == "warn" and "layout-06's title zone (x 48, y 28.83, width 864 px" in finding.message
    assert finding.line == 1
    assert "the layout's title zone" in next(f.message for f in lint(static).findings
                                             if f.rule == "title-unpositioned")
    placed = PAGE.format(css="h1{position:absolute;left:48px;top:29px}", body='<h1 data-placeholder="title">T</h1>')
    carded = PAGE.format(css=".c{position:absolute;left:9px;top:9px}",
                         body='<div class="c"><h1 data-placeholder="title">T</h1></div>')
    for html in (placed, carded):
        assert "title-unpositioned" not in {f.rule for f in lint(html, georgia, layout_id=LAYOUT).findings}


def _synthetic_font(path: Path, family: str, characters: str) -> Path:
    """A tiny TrueType font named `family` whose cmap holds exactly `characters` (empty glyphs)."""
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.ttGlyphPen import TTGlyphPen

    names = [".notdef"] + [f"g{ord(c):04X}" for c in characters]
    builder = FontBuilder(1000, isTTF=True)
    builder.setupGlyphOrder(names)
    builder.setupCharacterMap({ord(c): f"g{ord(c):04X}" for c in characters})
    empty = TTGlyphPen(None).glyph()
    builder.setupGlyf({name: empty for name in names})
    builder.setupHorizontalMetrics({name: (500, 0) for name in names})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable({"familyName": family, "styleName": "Regular"})
    builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200)
    builder.setupPost()
    builder.save(str(path))
    return path


def test_the_glyph_lint_judges_a_major_font_div_by_the_major_font(georgia: Manifest, monkeypatch, tmp_path: Path):
    """A div in the major font is set in it, so its characters are judged against that font's cmap:
    here a synthetic major with no arrow (U+2192) over a synthetic minor that has one."""
    from app.engine.emit import text as text_engine
    from app.engine.extract import html as html_extract

    _synthetic_font(tmp_path / "synth-major.ttf", "Synth Major", "Next ")
    _synthetic_font(tmp_path / "synth-minor.ttf", "Synth Minor", "Next →")
    monkeypatch.setattr(config, "FONT_DIRS", (tmp_path,))
    text_engine.forget_fonts()
    html_extract.forget_installed_families()
    try:
        manifest = copy.deepcopy(georgia)
        manifest.theme = {**manifest.theme, "fonts": {"major": "Synth Major", "minor": "Synth Minor"}}
        in_major = PAGE.format(css=".k{font-family:'Synth Major',sans-serif}",
                               body='<div class="k">Next →</div>')
        in_minor = PAGE.format(css="", body="<div>Next →</div>")
        child = PAGE.format(css=".k{font-family:'Synth Major'}.n{font-family:'Synth Minor'}",
                            body='<div class="k"><span class="n">Next →</span></div>')

        def glyphs(html: str) -> list[Any]:
            return [f for f in lint(html, manifest).findings if f.rule == "glyph-missing"]

        (finding,) = glyphs(in_major)
        assert "Synth Major" in finding.message and "U+2192" in finding.message
        assert glyphs(in_minor) == [] and glyphs(child) == []
    finally:
        text_engine.forget_fonts()
        html_extract.forget_installed_families()
