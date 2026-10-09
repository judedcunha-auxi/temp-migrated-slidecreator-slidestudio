"""`border-collapse: separate` with a `border-spacing`: the gaps are spacer rows and columns.

A PowerPoint table has no cell spacing, so before this the native table's cell fills ran over the
white gaps of a scorecard. `page.js: spacerGrid` makes every gap between cells that paint an empty,
unfilled, borderless row or column of the native table (the torture family `tables-spacing`, which
the gate also renders). PowerPoint lays out no row or column under 2 pt (2.667 px, measured through
COM on 2026-09-29): a gap from 1 pt up is widened to that, one under 1 pt is closed with a warn.
Collapse tables and tables whose spacing is 0 keep the logical grid, byte for byte.
"""
from __future__ import annotations

import shutil
from dataclasses import dataclass
from functools import cached_property
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from app.config import engine as config
from app.engine.classify.placeholders import map_placeholders
from app.engine.emit import pptx as emitter
from app.engine.ir import IR, Element
from app.engine.manifest import Manifest
from app.engine.reports import EmitOptions, EmitReport
from app.engine.verify.fit import fit
from app.engine.verify.fixtures import context_for
from tests.engine.helpers import extract_html, render

FAMILY = config.TORTURE_FIXTURE / "tables-spacing.html"
REFERENCE = config.TORTURE_FIXTURE / "reference" / "tables-spacing.png"

#: PowerPoint's floor for a row height or a column width (2 pt), in px.
FLOOR_PX = 8 / 3
EMU_PER_PX = 9525


@dataclass
class Spaced:
    ir: IR
    deck: Path
    report: EmitReport
    out: Path
    canvas_w: int
    cells: dict[str, list[dict]]        # per table name, the browser's own th/td rects in DOM order

    @cached_property
    def render(self) -> Path:
        """PptxRender's picture of the deck, drawn on first use: only the `renderer` tests need it."""
        pngs = render(self.deck, self.canvas_w, self.out / "render")
        rendered = self.out / "render.png"
        shutil.copy(pngs[0], rendered)
        return rendered


@pytest.fixture(scope="module")
def spaced(tmp_path_factory, browser) -> Spaced:
    context = context_for(FAMILY)
    out = tmp_path_factory.mktemp("spacing")
    page = browser.new_page(viewport={"width": 1280, "height": 720})
    try:
        page.goto(FAMILY.as_uri())
        cells = page.evaluate("""() => {
          const out = {};
          document.querySelectorAll('table[data-name]').forEach(t => {
            out[t.dataset.name] = [...t.querySelectorAll('th,td')].map(c => {
              const r = c.getBoundingClientRect();
              return {text: c.textContent.trim(), x: r.x, y: r.y, w: r.width, h: r.height};
            });
          });
          return out;
        }""")
    finally:
        page.close()
    ir = extract_html(FAMILY, context.manifest, context.layout_id, context.assets_dir,
                      slide_id=FAMILY.stem, title=FAMILY.stem)
    ir = map_placeholders(ir, context.layout)
    deck = out / "tables-spacing.pptx"
    report = emitter.emit([ir], context.manifest, context.master, deck, EmitOptions())
    return Spaced(ir=ir, deck=deck, report=report, out=out, canvas_w=context.manifest.canvas_w,
                  cells=cells)


def _table(ir: IR, name: str) -> Element:
    return next(e for e in ir.elements if e.kind == "table" and e.name == name)


def _lines(sizes: list[float], start: float) -> list[float]:
    out = [start]
    for size in sizes:
        out.append(out[-1] + size)
    return out


def _rect(table: Element, cell: dict) -> tuple[float, float, float, float]:
    """A cell's physical rect in the native table: its grid lines, spans included."""
    xs, ys = _lines(table.colWidthsPx, table.box.x), _lines(table.rowHeightsPx, table.box.y)
    return (xs[cell["c"]], ys[cell["r"]], xs[cell["c"] + cell["colSpan"]], ys[cell["r"] + cell["rowSpan"]])


def _text(cell: dict) -> str:
    return " ".join("".join(run["text"] for run in line["runs"])
                    for paragraph in cell["paragraphs"] for line in paragraph["lines"])


def _spacer(table: Element, r: int, c: int) -> bool:
    return bool(table.extras["x-wpa"]["cells"][f"{r}:{c}"].get("spacer"))


# ------------------------------------------------------------------------------------ extraction


def test_scorecard_is_a_spacer_grid(spaced: Spaced):
    """2 px between columns and 4 px between rows: a 5 x 5 table becomes 9 x 9, every other track a
    spacer; the 4 px rows are exact, the 2 px columns are PowerPoint's 2 pt floor."""
    table = _table(spaced.ir, "scorecard")
    wpa = table.extras["x-wpa"]
    assert (table.rows, table.cols) == (9, 9)
    assert wpa["spacingPx"] == {"h": 2, "v": 4} and wpa["spacers"] == {"rows": [1, 3, 5, 7], "cols": [1, 3, 5, 7]}
    assert [table.rowHeightsPx[k] for k in (1, 3, 5, 7)] == pytest.approx([4] * 4, abs=1e-6)
    assert [table.colWidthsPx[k] for k in (1, 3, 5, 7)] == pytest.approx([FLOOR_PX] * 4, abs=1e-6)
    assert sum(table.rowHeightsPx) == pytest.approx(table.box.h, abs=0.01)
    assert sum(table.colWidthsPx) == pytest.approx(table.box.w, abs=0.01)
    spacers = [c for c in table.cells if _spacer(table, c["r"], c["c"])]
    assert spacers and all(c["fill"] == {"type": "none"} and not c["paragraphs"] and
                           c["paddingPx"] == {"t": 0, "r": 0, "b": 0, "l": 0} and
                           not any(c["borders"].values()) for c in spacers)
    assert all((c["r"] in (1, 3, 5, 7) or c["c"] in (1, 3, 5, 7)) == _spacer(table, c["r"], c["c"])
               for c in table.cells)


def test_scorecard_cells_land_where_the_browser_drew_them(spaced: Spaced):
    """Each content cell's rect in the native table is the browser's th/td rect: rows exactly, columns
    within the third of a pixel each neighbour gives up to widen a 2 px gap to 2 pt."""
    table = _table(spaced.ir, "scorecard")
    content = [c for c in table.cells if not _spacer(table, c["r"], c["c"])]
    browser = spaced.cells["scorecard"]
    assert [_text(c) for c in content] == [b["text"] for b in browser]
    for cell, drawn in zip(content, browser, strict=False):
        x0, y0, x1, y1 = _rect(table, cell)
        assert (y0, y1) == pytest.approx((drawn["y"], drawn["y"] + drawn["h"]), abs=0.01), drawn["text"]
        give = (FLOOR_PX - 2) / 2 + 0.01
        assert abs(x0 - drawn["x"]) <= give and abs(x1 - (drawn["x"] + drawn["w"])) <= give, drawn["text"]


def test_a_rowspan_spans_the_spacer_row_between_its_rows(spaced: Spaced):
    table = _table(spaced.ir, "scorecard")
    growth = next(c for c in table.cells if _text(c) == "Growth")
    assert (growth["r"], growth["c"], growth["rowSpan"], growth["colSpan"]) == (2, 0, 3, 1)
    assert not [c for c in table.cells if (c["r"], c["c"]) == (3, 0)], "the merged cell covers the spacer slot"
    assert _spacer(table, 3, 2) and _spacer(table, 3, 1)


def test_scorecard_says_what_it_did(spaced: Spaced):
    messages = [(d.level, d.message) for d in spaced.ir.diagnostics]
    assert ("info", "border-spacing 2 4 px: the gaps between painted cells are 4 spacer row(s) and 4 spacer "
                    "column(s) of the native table (empty, unfilled)") in messages
    assert any(level == "info" and m.startswith("border-spacing 2 px is drawn 2.67 px wide") for level, m in messages)
    assert not any("cell fills extend into the spacing" in m for _, m in messages)


def test_cards_rules_move_onto_the_spacer_cells(spaced: Spaced):
    """PowerPoint draws neither a lower cell's `lnT` nor a right cell's `lnL` on an interior line
    (measured): a cell's top and left rules become the `lnB` / `lnR` of the spacer cells above and to
    its left. Grid lines sit half the widest rule inside each band, so 6 px + two half borders = 7."""
    table = _table(spaced.ir, "cards")
    cells = {(c["r"], c["c"]): c for c in table.cells}
    assert (table.rows, table.cols) == (5, 5)
    assert table.rowHeightsPx[1] == table.rowHeightsPx[3] == 7 and table.colWidthsPx[1] == table.colWidthsPx[3] == 7
    rule = {"width": 1, "color": "C9D2DF", "dash": "solid"}
    interviews = cells[(2, 2)]
    assert interviews["borders"]["top"] is None and interviews["borders"]["left"] is None
    assert interviews["borders"]["right"] == rule and interviews["borders"]["bottom"] == rule
    assert cells[(1, 2)]["borders"]["bottom"] == rule, "Interviews' top rule, on the spacer above"
    assert cells[(2, 1)]["borders"]["right"] == rule, "Interviews' left rule, on the spacer to its left"
    assert cells[(0, 0)]["borders"]["top"] is not None and cells[(0, 0)]["borders"]["left"] is not None, "outer edges stay"


def test_cards_merged_cells_span_the_spacers_and_keep_their_rules(spaced: Spaced):
    """The colspan header covers physical columns 0-2, the rowspan accent cell rows 2-4. PowerPoint
    draws a merged cell's interior rule from the origin only along its first column / row and takes
    the rest from the neighbour's `lnT` / `lnL` (measured; it writes them so itself): so the spacer
    cells under the header's continuation columns carry its bottom rule, and those beside the accent
    cell's continuation rows its right rule."""
    table = _table(spaced.ir, "cards")
    cells = {(c["r"], c["c"]): c for c in table.cells}
    header, accent = cells[(0, 0)], cells[(2, 0)]
    assert (header["rowSpan"], header["colSpan"], _text(header)) == (1, 3, "Workstream")
    assert (accent["rowSpan"], accent["colSpan"], _text(accent)) == (3, 1, "Discovery and build")
    blue = {"width": 1, "color": "0B5CAD", "dash": "solid"}
    assert header["borders"]["bottom"] == blue
    assert cells[(1, 1)]["borders"]["top"] == blue and cells[(1, 2)]["borders"]["top"] == blue
    assert cells[(1, 0)]["borders"]["top"] is None, "the origin's own column needs no copy"
    grey = {"width": 1, "color": "C9D2DF", "dash": "solid"}
    assert accent["borders"]["right"] == grey and accent["borders"]["left"]["width"] == 4
    assert cells[(3, 1)]["borders"]["left"] == grey and cells[(4, 1)]["borders"]["left"] == grey


def test_a_gap_under_one_point_is_closed_and_said(spaced: Spaced):
    """1 px between filled cells: widening it to 2 pt would misplace more pixels than closing it."""
    table = _table(spaced.ir, "hairline")
    assert (table.rows, table.cols) == (2, 3) and table.extras["x-wpa"]["spacers"] == {"rows": [], "cols": []}
    drawn = spaced.cells["hairline"]
    assert table.box.x == pytest.approx(drawn[0]["x"]) and table.box.y == pytest.approx(drawn[0]["y"]), \
        "the table's edge moves to the painted cells' edge"
    warns = [d.message for d in spaced.ir.diagnostics if d.level == "warn" and "border-spacing" in d.message]
    assert warns == ["border-spacing 1 px between painted cells is narrower than 1 pt: PowerPoint draws no row or "
                     "column under 2 pt, so the cell fills close the gap (use a spacing of 2 px or more, or "
                     "border-collapse with a border in the gap's colour)"]


def test_a_one_pixel_gap_between_bordered_cells_is_kept(tmp_path: Path, sample_manifest: Manifest, sample_dir: Path):
    """Half of each 1 px rule lies in the spacer, so its track is 2 px: widening it to 2 pt misdraws
    0.67 px, closing it 1 px — the spacer stays (info, no warn)."""
    html = tmp_path / "bordered.html"
    html.write_text(
        '<!doctype html><html><body style="margin:0;font-family:Arial;font-size:12px">'
        '<table data-name="t" style="position:absolute;left:40px;top:40px;border-spacing:1px">'
        '<tr><td style="border:1px solid #999;background:#eee;padding:4px">A</td>'
        '<td style="border:1px solid #999;background:#eee;padding:4px">B</td></tr></table></body></html>',
        encoding="utf-8")
    ir = extract_html(html, sample_manifest, "layout-05", sample_dir / "assets", slide_id="bordered")
    table = _table(ir, "t")
    assert table.extras["x-wpa"]["spacers"] == {"rows": [], "cols": [1]}
    assert table.colWidthsPx[1] == pytest.approx(FLOOR_PX, abs=1e-6)
    assert not [d for d in ir.diagnostics if d.level == "warn" and "border-spacing" in d.message]
    assert [d for d in ir.diagnostics if d.level == "info" and d.message.startswith("border-spacing 1 px is drawn 2.67 px")]


def test_a_table_that_paints_nothing_keeps_its_grid(spaced: Spaced):
    """The default 2 px spacing around cells with no fill and no border is invisible: no spacer."""
    table = _table(spaced.ir, "plain")
    wpa = table.extras["x-wpa"]
    assert (table.rows, table.cols) == (2, 3) and "spacers" not in wpa
    assert wpa["cells"]["0:0"]["slot"] == {"t": 2, "r": 2, "b": 2, "l": 2}
    assert not any(c.get("spacer") for c in wpa["cells"].values())


def test_collapse_and_zero_spacing_keep_the_logical_grid(tmp_path: Path, sample_manifest: Manifest, sample_dir: Path):
    html = tmp_path / "grid.html"
    html.write_text(
        '<!doctype html><html><body style="margin:0;font-family:Arial;font-size:12px">'
        '<table data-name="collapse" style="position:absolute;left:40px;top:40px;border-collapse:collapse;'
        'border-spacing:6px"><tr><td style="background:#0B2545;color:#fff;padding:4px" rowspan="2">A</td>'
        '<td style="background:#EEF1F6;padding:4px">B</td></tr><tr><td style="background:#EEF1F6;padding:4px">C</td></tr></table>'
        '<table data-name="zero" style="position:absolute;left:40px;top:200px;border-spacing:0">'
        '<tr><td style="background:#0B2545;color:#fff;padding:4px">A</td><td style="border:1px solid #999">B</td></tr>'
        '</table></body></html>', encoding="utf-8")
    ir = extract_html(html, sample_manifest, "layout-05", sample_dir / "assets", slide_id="grid")
    for name, shape in (("collapse", (2, 2)), ("zero", (1, 2))):
        table = _table(ir, name)
        assert (table.rows, table.cols) == shape and "spacers" not in table.extras["x-wpa"], name
        assert not any(c.get("spacer") for c in table.extras["x-wpa"]["cells"].values())
    assert not [d for d in ir.diagnostics if "border-spacing" in d.message]


# ------------------------------------------------------------------------------------ emission


def test_spacer_cells_are_empty_1pt_cells_with_no_inset(spaced: Spaced):
    """An empty cell's end paragraph at 1 pt and no margins: nothing but PowerPoint's own 2 pt floor
    holds the row open, and the spacers are declared at or above it."""
    table = next(s for s in Presentation(str(spaced.deck)).slides[0].shapes
                 if s.has_table and s.name == "scorecard").table
    assert [round(r.height / EMU_PER_PX, 3) for r in table.rows][1::2] == [4.0] * 4
    assert [c.width for c in table.columns][1::2] == pytest.approx([2 * 12700] * 4, abs=1)
    spacer = table.cell(1, 1)._tc
    tcPr = spacer.find(qn("a:tcPr"))
    assert all(tcPr.get(side) == "0" for side in ("marL", "marR", "marT", "marB"))
    assert tcPr.find(qn("a:noFill")) is not None and tcPr.find(qn("a:solidFill")) is None
    paragraphs = spacer.findall(f"{qn('a:txBody')}/{qn('a:p')}")
    assert len(paragraphs) == 1 and paragraphs[0].find(qn("a:endParaRPr")).get("sz") == "100"
    growth = table.cell(2, 0)._tc
    assert growth.get("rowSpan") == "3" and table.cell(3, 0)._tc.get("vMerge") == "1"


def test_cards_rules_reach_the_file(spaced: Spaced):
    table = next(s for s in Presentation(str(spaced.deck)).slides[0].shapes
                 if s.has_table and s.name == "cards").table

    def lines(r: int, c: int) -> set[str]:
        tcPr = table.cell(r, c)._tc.find(qn("a:tcPr"))
        return {child.tag.split("}")[1] for child in tcPr if child.tag.split("}")[1] in ("lnL", "lnR", "lnT", "lnB")}

    assert table.cell(0, 0)._tc.get("gridSpan") == "3"
    assert lines(1, 2) == {"lnT", "lnB"} and lines(2, 1) == {"lnR"} and lines(2, 2) == {"lnR", "lnB"}


@pytest.mark.renderer_full
def test_rows_keep_their_height_and_the_gaps_show(spaced: Spaced):
    """fit() predicts no row growth, and PptxRender paints the page, not a cell, in the gaps: the
    render agrees with the browser's reference at every scorecard gap centre and cell centre."""
    report = fit(spaced.deck, [spaced.ir])
    assert not [i for i in report.issues if i.kind == "overflow" and " row " in i.shape], report.issues
    ours = [w for w in spaced.report.warnings if "cell" in w]
    assert ours == [], ours
    np = pytest.importorskip("numpy")
    reference = np.asarray(Image.open(REFERENCE).convert("RGB"), dtype=np.int16)
    rendered = np.asarray(Image.open(spaced.render).convert("RGB"), dtype=np.int16)
    browser = spaced.cells["scorecard"]
    header = browser[:5]
    points = []
    for left, _right in zip(header, header[1:], strict=False):             # column gaps in the header row
        points.append((left["x"] + left["w"] + 1, left["y"] + left["h"] / 2, True))
    measure = browser[1]                                      # the row gap under the 'Measure' header
    points.append((measure["x"] + measure["w"] / 2, measure["y"] + measure["h"] + 2, True))
    points += [(c["x"] + c["w"] / 2, c["y"] + 3, False) for c in browser]
    for x, y, gap in points:
        want, got = reference[int(y), int(x)], rendered[int(y), int(x)]
        assert np.abs(want - got).max() <= 24, (x, y, gap, want.tolist(), got.tolist())
        if gap:
            assert got.min() >= 235, ("a gap should show the page", x, y, got.tolist())
