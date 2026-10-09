"""WP4 — the emitter, measured rather than eyeballed.

Inputs are `tests/engine/fixtures/torture/ir/emit-*.json`: small hand-built IRs, one per construct
family, that exercise the emitter before WP1/WP2 can feed it real slides. They are deliberately
*this* package's fixtures (`emit-` prefix) so they can never collide with the extraction fixtures
WP0b writes beside them.

A missing input fails the test; nothing here skips (master brief §10.7).
"""
from __future__ import annotations

import contextlib
import hashlib
import json
import math
import subprocess
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image, ImageChops, ImageFont
from pptx import Presentation

from app.config import engine as config
from app.engine.classify.placeholders import map_placeholders
from app.engine.emit import pptx as emitter
from app.engine.emit import shapes as shape_emitter
from app.engine.emit import template as template_module
from app.engine.emit import text as text_engine
from app.engine.ir import IR, Box, Canvas, Element, Slide, assign_ids_and_z
from app.engine.manifest import Manifest
from app.engine.reports import EmitOptions
from tests.engine import helpers as renderer
from tests.engine.helpers import available, require_faces

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

#: Which generated master a fixture belongs to, decided by its canvas (master brief §5: the canvas
#: *is* the slide size in inches × 96, so it identifies the master without a side-car file).
MASTER_BY_CANVAS = {(1280, 720): "test-16x9", (960, 720): "test-4x3"}

WIDE_FIXTURES = ("emit-shapes", "emit-text", "emit-placeholders", "emit-images",
                 "emit-tables", "emit-charts", "emit-tables-rich")
ALL_FIXTURES = WIDE_FIXTURES + ("emit-sizes-4x3",)


# ----------------------------------------------------------------------------------- fixtures


def _ir_dir() -> Path:
    directory = config.TORTURE_FIXTURE / "ir"
    assert directory.is_dir(), f"emitter IR fixtures missing at {directory}"
    return directory


def load_ir(name: str) -> IR:
    """One emitter fixture, with relative asset paths resolved against the fixture directory.

    The JSON keeps `src` relative on purpose: an absolute Windows path in a tracked fixture is a
    fixture that only works on one machine.
    """
    path = _ir_dir() / f"{name}.json"
    assert path.exists(), f"{path} is missing"
    data = json.loads(path.read_text(encoding="utf-8"))
    for element in data.get("elements") or []:
        src = element.get("src")
        if src and not Path(src).is_absolute():
            element["src"] = str((path.parent / src).resolve())
    ir = IR.from_json(data)
    assert ir.validate() == [], f"{name}: {ir.validate()}"
    return ir


def master_of(ir: IR) -> tuple[Path, Manifest]:
    stem = MASTER_BY_CANVAS[(ir.canvas.w, ir.canvas.h)]
    master = config.MASTERS_FIXTURE / f"{stem}.pptx"
    manifest = config.MASTERS_FIXTURE / f"{stem}.manifest.json"
    assert master.exists() and manifest.exists(), f"{stem} fixture master is missing"
    return master, Manifest.load(manifest)


def emit_fixture(name: str, out_dir: Path, *, options: EmitOptions | None = None):
    """Classify + emit one fixture; returns `(pptx path, EmitReport, IR, master path, manifest)`."""
    ir = load_ir(name)
    master, manifest = master_of(ir)
    ir = map_placeholders(ir, manifest.layout(ir.slide.layoutId))
    out = Path(out_dir) / f"{name}.pptx"
    report = emitter.emit([ir], manifest, master, out, options or EmitOptions())
    return out, report, ir, master, manifest


@pytest.fixture(scope="session")
def emitted(tmp_path_factory) -> dict[str, tuple[Path, object, IR, Path, Manifest]]:
    """Every fixture emitted once — the decks the assertions below read back."""
    directory = tmp_path_factory.mktemp("emit")
    return {name: emit_fixture(name, directory) for name in ALL_FIXTURES}


def shapes_of(pptx: Path, index: int = 0):
    return list(Presentation(str(pptx)).slides[index].shapes)


def by_name(pptx: Path, name: str):
    for shape in shapes_of(pptx):
        if shape.name == name:
            return shape
        if shape.shape_type is not None and hasattr(shape, "shapes"):
            for child in shape.shapes:
                if child.name == name:
                    return child
    raise AssertionError(f"no shape named {name!r} in {pptx.name}")


# ------------------------------------------------------------------------------ the fixtures load


def test_every_fixture_is_a_valid_ir():
    for name in ALL_FIXTURES:
        ir = load_ir(name)
        assert ir.elements, f"{name} has no elements"
        assert ir.validate() == []


def test_every_fixture_emits_without_warnings(emitted):
    for name, (out, report, ir, _, _) in emitted.items():
        assert out.exists(), f"{name} produced no file"
        assert report.warnings == [], f"{name}: {report.warnings}"
        assert sum(report.shapes_by_kind.values()) >= len(ir.elements) - len(ir.groups)


# ------------------------------------------------------------------------------------- template


def test_template_parts_are_byte_identical_to_the_master(emitted):
    """The customer's masters, layouts and themes must come out exactly as they went in."""
    for name, (out, _, _, master, _) in emitted.items():
        assert template_module.verify_template(out, master) == [], name


def test_the_only_template_part_an_export_adds_is_a_notes_master_theme(emitted):
    """Writing notes onto a master that has none makes python-pptx add a notes master.

    That notes master brings a theme part of its own, which is an *addition*, not a change to
    anything of the customer's — WP5's suite should read it the same way.
    """
    added = {name: template_module.added_template_parts(out, master)
             for name, (out, _, _, master, _) in emitted.items()}
    assert added["emit-text"] == ["ppt/theme/theme2.xml"], added["emit-text"]
    for name, parts in added.items():
        if name != "emit-text":                # emit-text is the only fixture carrying notes
            assert parts == [], f"{name} added {parts}"


def test_the_master_keeps_its_own_slide_count(emitted):
    out, _, _, master, _ = emitted["emit-shapes"]
    assert len(Presentation(str(master)).slides) == 0
    assert len(Presentation(str(out)).slides) == 1


def test_speaker_notes_ride_along(emitted):
    out, _, ir, _, _ = emitted["emit-text"]
    slide = Presentation(str(out)).slides[0]
    assert slide.has_notes_slide
    assert ir.slide.notes in slide.notes_slide.notes_text_frame.text


# --------------------------------------------------------------------------------------- shapes


def test_preset_shapes_and_counts(emitted):
    out, report, _, _, _ = emitted["emit-shapes"]
    presets = {}
    for shape in shapes_of(out):
        geometry = shape._element.find(f".//{{{A}}}prstGeom")
        if geometry is not None:
            presets.setdefault(geometry.get("prst"), []).append(shape.name)
    assert "roundRect" in presets and "roundrect-equal" in presets["roundRect"]
    assert presets["blockArc"] == ["donut"]
    assert presets["pie"] == ["pie"]
    assert presets["chord"] == ["chord"]
    assert "ellipse" in presets
    assert report.shapes_by_kind["shape"] == 19
    assert report.shapes_by_kind["group"] == 2


def test_block_arc_adjustments_are_angles_and_ring_thickness(emitted):
    """`adj1`/`adj2` are start/end angle × 60000; `adj3` is (rOuter − rInner) / min(w, h)."""
    out, _, _, _, _ = emitted["emit-shapes"]
    donut = by_name(out, "donut")
    adjustments = {gd.get("name"): int(gd.get("fmla").split()[1])
                   for gd in donut._element.findall(f".//{{{A}}}gd")}
    assert adjustments["adj1"] == 270 * 60000
    assert adjustments["adj2"] == 90 * 60000
    assert adjustments["adj3"] == round((70 - 42) / 140 * 100000)   # 20 %

    pie = by_name(out, "pie")
    pie_adjustments = {gd.get("name"): int(gd.get("fmla").split()[1])
                       for gd in pie._element.findall(f".//{{{A}}}gd")}
    assert pie_adjustments == {"adj1": 0, "adj2": 120 * 60000}


@pytest.mark.renderer_full
def test_block_arc_renders_as_a_ring_of_the_right_radius(emitted, tmp_path):
    """The adjustment semantics are verified in pixels, not only in XML.

    A `blockArc` from 270° to 90° over the right half of the box: the ring must be painted at the
    outer radius, empty at the inner radius, and empty on the left half.
    """
    out, _, _, _, _ = emitted["emit-shapes"]
    png = renderer.render(out, 1280, tmp_path / "arc")[0]
    image = Image.open(png).convert("RGB")
    centre_x, centre_y = 110, 270

    def is_blue(x: int, y: int) -> bool:
        r, g, b = image.getpixel((x, y))
        return b > 150 and r < 120

    assert is_blue(centre_x + 56, centre_y), "no ink in the ring band on the right"
    assert not is_blue(centre_x + 20, centre_y), "the hole is filled"
    assert not is_blue(centre_x - 56, centre_y), "the unswept half is painted"
    assert is_blue(centre_x, centre_y - 56), "the 12 o'clock edge of the sweep is missing"


def test_unequal_corner_radii_become_custom_geometry(emitted):
    out, _, _, _, _ = emitted["emit-shapes"]
    shape = by_name(out, "roundrect-unequal")
    assert shape._element.find(f".//{{{A}}}custGeom") is not None
    commands = [child.tag.split("}")[1]
                for child in shape._element.find(f".//{{{A}}}path")]
    assert commands.count("cubicBezTo") == 2          # only two corners are rounded
    assert commands[0] == "moveTo" and commands[-1] == "close"


def test_custom_path_commands_survive(emitted):
    out, _, _, _, _ = emitted["emit-shapes"]
    ring = by_name(out, "ring-evenodd")
    paths = ring._element.findall(f".//{{{A}}}path")
    # Even-odd is expressed by keeping both subpaths inside ONE a:path, which is how DrawingML
    # makes a hole. Two a:path elements would paint a filled square over a filled square.
    assert len(paths) == 1
    commands = [child.tag.split("}")[1] for child in paths[0]]
    assert commands.count("moveTo") == 2
    assert commands.count("close") == 2
    assert commands.count("lnTo") == 6


def test_polyline_is_an_open_unfilled_path(emitted):
    out, _, _, _, _ = emitted["emit-shapes"]
    path = by_name(out, "polyline")._element.find(f".//{{{A}}}path")
    assert path.get("fill") == "none"
    assert [child.tag.split("}")[1] for child in path] == ["moveTo"] + ["lnTo"] * 4


def test_line_becomes_a_connector_with_arrow_heads(emitted):
    out, _, _, _, _ = emitted["emit-shapes"]
    arrow = by_name(out, "arrow")
    assert arrow._element.tag.endswith("}cxnSp")
    line = arrow._element.find(f".//{{{A}}}ln")
    assert line.find(f"{{{A}}}headEnd").get("type") == "oval"
    assert line.find(f"{{{A}}}tailEnd").get("type") == "triangle"
    assert line.get("cap") == "rnd"


@pytest.mark.parametrize(
    "pattern, width, expected",
    [
        ([4, 3], 1.0, "dash"),          # exactly the preset
        ([3, 2], 1.0, "sysDash"),       # the dash length is the preset's; the gap is not
        ([1, 1], 1.0, "sysDot"),
        ([1, 3], 1.0, "dot"),
        ([8, 3], 1.0, "lgDash"),
        ([8, 6], 2.0, "dash"),          # same pattern at twice the line width
        ([30, 20], 1.0, None),          # a long dash is not a preset — custDash
        ([2, 2], 0.5, "dash"),          # four line-widths on, four off: `dash` within the rhythm
        ([5, 5], 1.0, None),            # no preset has a 5-line-width dash
        ([4, 40], 1.0, None),           # the dash length matches `dash`; the rhythm does not
    ],
)
def test_dash_pattern_mapping(pattern, width, expected):
    assert shape_emitter.match_preset_dash(pattern, width) == expected


def test_dashes_reach_the_file(emitted):
    out, report, _, _, _ = emitted["emit-shapes"]
    def dash_of(name):
        line = by_name(out, name)._element.find(f".//{{{A}}}ln")
        preset = line.find(f"{{{A}}}prstDash")
        custom = line.find(f"{{{A}}}custDash")
        return preset.get("val") if preset is not None else ("custDash" if custom is not None else None)

    assert dash_of("roundrect-equal") == "dash"
    assert dash_of("roundrect-unequal") == "sysDash"
    assert dash_of("dash-sysdot") == "sysDot"
    assert dash_of("dash-custom") == "custDash"
    ds = by_name(out, "dash-custom")._element.find(f".//{{{A}}}ds")
    assert (int(ds.get("d")), int(ds.get("sp"))) == (3000000, 2000000)   # 30 px / 1 px line width
    assert any("custDash" in gap for gap in report.renderer_gaps)


def test_gradient_angle_and_stops(emitted):
    out, _, _, _, _ = emitted["emit-shapes"]
    gradient = by_name(out, "ellipse-gradient")._element.find(f".//{{{A}}}gradFill")
    assert gradient.find(f"{{{A}}}lin").get("ang") == str(((90 - 90) % 360) * 60000)
    stops = gradient.findall(f".//{{{A}}}gs")
    assert [int(gs.get("pos")) for gs in stops] == [0, 100000]
    assert stops[1].find(f"{{{A}}}srgbClr/{{{A}}}alpha").get("val") == "40000"

    radial = by_name(out, "rect-radial")._element.find(f".//{{{A}}}gradFill")
    assert radial.find(f"{{{A}}}path").get("path") == "circle"


def test_shadow_direction_and_distance(emitted):
    out, _, _, _, _ = emitted["emit-shapes"]
    shadow = by_name(out, "rect-radial")._element.find(f".//{{{A}}}outerShdw")
    assert int(shadow.get("blurRad")) == 12 * config.EMU_PER_PX
    assert int(shadow.get("dist")) == round((4 ** 2 + 6 ** 2) ** 0.5 * config.EMU_PER_PX)
    assert 56 * 60000 < int(shadow.get("dir")) < 57 * 60000            # atan2(6, 4) ≈ 56.3°
    assert shadow.find(f"{{{A}}}srgbClr/{{{A}}}alpha").get("val") == "25000"


def test_every_shape_carries_an_effect_list(emitted):
    """An empty `effectLst` is what stops a theme's effect style adding a shadow we never asked for."""
    out, _, _, _, _ = emitted["emit-shapes"]
    for shape in shapes_of(out):
        if shape._element.tag.endswith("}sp"):
            assert shape._element.spPr.find(f"{{{A}}}effectLst") is not None, shape.name
            assert shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style") is None


def test_rotation_and_flip(emitted):
    out, _, _, _, _ = emitted["emit-shapes"]
    xfrm = by_name(out, "rect-rotated")._element.spPr.find(f"{{{A}}}xfrm")
    assert xfrm.get("rot") == str(15 * 60000)
    assert xfrm.get("flipH") == "1"


def test_clip_shrinks_a_rectangle_and_flattens_a_curve(emitted):
    out, report, _, _, _ = emitted["emit-shapes"]
    rect = by_name(out, "rect-clipped")
    assert (rect.width, rect.height) == (120 * config.EMU_PER_PX, 80 * config.EMU_PER_PX)
    ellipse = by_name(out, "ellipse-clipped")
    assert ellipse._element.find(f".//{{{A}}}custGeom") is not None
    assert ellipse.width == 140 * config.EMU_PER_PX
    assert any("flattened" in gap for gap in report.renderer_gaps)


def test_groups_are_real_groups_with_correct_extents(emitted):
    out, _, _, _, _ = emitted["emit-shapes"]
    groups = [s for s in shapes_of(out) if s.shape_type is not None and hasattr(s, "shapes")]
    assert [g.name for g in groups] == ["value-chain"]
    outer = groups[0]
    # Two 120 px children at x=40 and x=180 plus a nested group at x=320 → 40..400 wide.
    assert outer.left == 40 * config.EMU_PER_PX
    assert outer.width == 360 * config.EMU_PER_PX
    inner = [s for s in outer.shapes if hasattr(s, "shapes")]
    assert [g.name for g in inner] == ["badge"]
    assert inner[0].left == 320 * config.EMU_PER_PX and inner[0].width == 80 * config.EMU_PER_PX
    xfrm = outer._element.find(f".//{{{A}}}xfrm")
    assert xfrm.find(f"{{{A}}}chOff").get("x") == xfrm.find(f"{{{A}}}off").get("x")


# ----------------------------------------------------------------------------------------- text


def test_line_breaks_are_locked_not_re_wrapped(emitted):
    out, _, _, _, _ = emitted["emit-text"]
    frame = by_name(out, "paragraphs").text_frame
    first = frame.paragraphs[0]._p
    assert len(first.findall(f"{{{A}}}br")) == 1
    assert [r.text for r in frame.paragraphs[0].runs] == [
        "A paragraph the browser broke across", "exactly two lines."]
    assert frame._txBody.bodyPr.find(f"{{{A}}}noAutofit") is not None
    assert frame.margin_left == 0 and frame.margin_top == 0


def test_font_size_keeps_hundredth_of_a_point_precision(emitted):
    out, _, _, _, _ = emitted["emit-text"]
    rPr = by_name(out, "paragraphs").text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    assert rPr.get("sz") == str(round(20.0 * config.FONT_SZ_PER_PX))     # 1500 = 15.00 pt


def test_run_properties_round_trip(emitted):
    out, _, _, _, _ = emitted["emit-text"]
    runs = by_name(out, "mixed-runs").text_frame.paragraphs[0].runs
    props = {r.text: r._r.get_or_add_rPr() for r in runs}
    assert props["bold"].get("b") == "1"
    assert props["italic"].get("i") == "1"
    assert props["underline"].get("u") == "sng"
    assert props["strike"].get("strike") == "sngStrike"
    assert int(props["wide"].get("spc")) == round(2.5 * config.PT_PER_PX * 100)
    supers = [r for r in runs if r._r.get_or_add_rPr().get("baseline") == "30000"]
    subs = [r for r in runs if r._r.get_or_add_rPr().get("baseline") == "-30000"]
    assert len(supers) == 1 and len(subs) == 1
    assert props["link"].find(f"{{{A}}}hlinkClick") is not None


def test_theme_fonts_become_tokens_and_others_stay_literal(emitted, tmp_path):
    """A run in the target master's own theme font is written as `+mn-lt`, so a re-theme follows."""
    out, _, _, _, _ = emitted["emit-text"]
    latin = by_name(out, "paragraphs").text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr() \
        .find(f"{{{A}}}latin")
    assert latin.get("typeface") == "Arial"        # the test master's theme is Calibri

    def typeface_when(fonts: dict[str, str], where: Path) -> str:
        ir = load_ir("emit-text")
        master, manifest = master_of(ir)
        manifest.masters[0].theme["fonts"] = dict(fonts)
        manifest.theme["fonts"] = dict(fonts)
        emitter.emit([map_placeholders(ir, manifest.layout(ir.slide.layoutId))], manifest, master,
                     where, EmitOptions())
        return by_name(where, "paragraphs").text_frame.paragraphs[0].runs[0] \
            ._r.get_or_add_rPr().find(f"{{{A}}}latin").get("typeface")

    # Body text is minor by default, so when a master uses one family for both, `+mn-lt` wins.
    assert typeface_when({"major": "Arial", "minor": "Arial"}, tmp_path / "one.pptx") == "+mn-lt"
    assert typeface_when({"major": "Arial", "minor": "Calibri"}, tmp_path / "two.pptx") == "+mj-lt"


# ------------------------------------------------------------------ weights → installed faces (WP-G)


#: Measured with kerning off in both engines: Pillow's basic layout does not kern GPOS fonts, and
#: kerning is not what tells a Light face from a Regular one.
FACE_SAMPLE = "The quick brown fox jumps over the lazy dog 0123456789"


def test_face_choice_matches_the_measuring_browser():
    """`face_for_exact` picks the face the measuring browser draws, at every weight (16-WPG §2.2).

    Edge's canvas `measureText` at 40 px, per family and weight 100–900, against Pillow's width of
    every upright face of the family: the nearest must be the one `face_for_exact` chose. This is
    what pins the tie rule (Segoe UI 500 → Semibold, Arial 800 → Black) to the browser that runs
    the suite rather than to a table measured once. "Calibri Light" at 700 is the synthetic-bold
    fact: the browser emboldens the Light face and its advances do not change.
    """
    from app.engine.extract.html import measuring_page

    # Segoe UI and Calibri ship only with Windows/Office; on Linux CI the rule is checked on Arial
    # (ttf-mscorefonts-installer), which carries the Black face the 800 tie rule needs.
    families = available(["Segoe UI", "Arial", "Calibri", "Calibri Light"])
    assert "Arial" in families, "Arial is not installed (CI installs ttf-mscorefonts-installer)"
    with measuring_page(Canvas(320, 200)) as page:
        page.set_content("<html><body></body></html>")
        widths = page.evaluate(
            """([families, text]) => {
                const context = document.createElement('canvas').getContext('2d');
                context.fontKerning = 'none';
                const out = {};
                for (const family of families) {
                    out[family] = [];
                    for (let weight = 100; weight <= 900; weight += 100) {
                        context.font = weight + ' 40px "' + family + '"';
                        out[family].push(context.measureText(text).width);
                    }
                }
                return out;
            }""",
            [families, FACE_SAMPLE],
        )
    problems = []
    for family in families:
        upright = {face.path: face for face in text_engine.font_index()[family.lower()] if not face.italic}
        assert upright, f"{family} is not installed; the face rule cannot be checked"
        for position, weight in enumerate(range(100, 1000, 100)):
            browser = widths[family][position]
            chosen = text_engine.face_for_exact(family, weight)
            nearest = min(upright.values(),
                          key=lambda face: abs(text_engine.face_width_px(face, FACE_SAMPLE, 40.0) - browser))
            ours = text_engine.face_width_px(chosen, FACE_SAMPLE, 40.0)
            if nearest.path != chosen.path or abs(ours - browser) / browser > 0.005:
                problems.append(f"{family} {weight}: browser {browser:.2f}px ~ {nearest.path.name}, "
                                f"face_for_exact chose {chosen.path.name} ({ours:.2f}px)")
    assert not problems, "\n".join(problems)
    if "Calibri Light" in widths:
        assert widths["Calibri Light"][6] == pytest.approx(widths["Calibri Light"][3], abs=0.01), \
            "the browser no longer synthesises bold on Calibri Light at 700"


def test_the_browser_emboldens_calibri_light_at_700():
    """The ink half of the synthetic-bold fact the test above checks by width alone (G-2 review M2).

    Edge draws "Calibri Light" at 700 with far more ink than at 300 (1.55 x, measured 2026-09-25 at 30 px on
    this sample) on the same advances. A width test cannot see that, and the promoted families' notes once
    said the opposite. It is why a 700 heading on a Calibri-Light major keeps `b="1"` (`emitted_face` rule
    (b)): the browser drew it bold, and `b="0"` would export every such heading Light.
    """
    import io

    import numpy as np
    from PIL import Image

    from app.engine.extract.html import measuring_page

    require_faces("Calibri Light")
    sample = "Why balance-sheet banks are losing share"
    weights = (300, 700)
    html = "<html><body style='margin:0;background:#fff'>" + "".join(
        f"<div style=\"position:absolute;left:10px;top:{10 + 50 * i}px;font:{weight} 30px 'Calibri Light';"
        f"color:#000;white-space:nowrap\">{sample}</div>" for i, weight in enumerate(weights)) + "</body></html>"
    with measuring_page(Canvas(640, 120)) as page:
        page.set_content(html)
        widths = page.evaluate("Array.from(document.querySelectorAll('div')).map(d => d.getBoundingClientRect().width)")
        shot = page.screenshot()
    grey = np.asarray(Image.open(io.BytesIO(shot)).convert("L")).astype(float)
    ink = [float((255 - grey[10 + 50 * i: 55 + 50 * i, :]).sum() / 255) for i in range(len(weights))]
    assert widths[1] == pytest.approx(widths[0], abs=0.01), f"Calibri Light's advances changed at 700: {widths}"
    assert ink[1] >= 1.3 * ink[0], f"Edge no longer emboldens Calibri Light at 700: ink {ink}"


def test_browser_centres_glyphs_with_the_metrics_the_calibration_assumes():
    """`browser_ascent_descent` is the measuring browser's own rule, checked on the live browser.

    The Pillow calibration below is browser-independent only if its expected baseline is where the
    browser really puts it: 100 px text in a 200 px line box, baseline read from a zero-height
    inline-block, against (ascent − descent) / 2 below the centre.
    """
    from app.engine.extract.html import measuring_page

    with measuring_page(Canvas(800, 400)) as page:
        page.set_content("<html><body style='margin:0'></body></html>")
        families = available(("Calibri", "Calibri Light", "Arial", "Segoe UI", "Times New Roman"))
        assert {"Arial", "Times New Roman"} <= set(families), "CI installs ttf-mscorefonts-installer"
        for family in families:
            top = page.evaluate(
                """(family) => {
                    document.body.replaceChildren();
                    const line = document.createElement('div');
                    Object.assign(line.style, {position: 'absolute', top: '0', left: '0',
                        fontFamily: JSON.stringify(family), fontSize: '100px', lineHeight: '200px'});
                    line.append('Hxp');
                    const mark = document.createElement('span');
                    Object.assign(mark.style, {display: 'inline-block', width: '1px', height: '0'});
                    line.append(mark);
                    document.body.append(line);
                    return mark.getBoundingClientRect().top;
                }""",
                family,
            )
            ascent, descent = browser_ascent_descent(text_engine.face_for_exact(family))
            assert (top - 100.0) / 100.0 == pytest.approx((ascent - descent) / 2, abs=0.01), family


def _run(text: str, family: str, weight: int = 400, italic: bool = False, size: float = 20.0) -> dict:
    return {"text": text, "font": family, "sizePx": size, "weight": weight, "italic": italic,
            "color": "1D2433", "alpha": 1, "underline": False, "strike": False,
            "letterSpacingPx": 0, "baseline": "normal", "baselineShiftPx": 0, "caps": False,
            "href": None}


def _one_line(name: str, run: dict, x: float, y: float, *, placeholder: str | None = None,
              line_height: float = 30.0) -> Element:
    width = text_engine.text_width_px(run["text"], run["font"], run["sizePx"], run["weight"], run["italic"])
    return Element(
        kind="text", box=Box(x, y, width, line_height), name=name,
        paragraphs=[{"align": "left", "lineHeightPx": line_height, "spaceBeforePx": 0,
                     "spaceAfterPx": 0, "bullet": None,
                     "lines": [{"box": {"x": x, "y": y, "w": width, "h": line_height}, "runs": [run]}]}],
        anchor="top", writingMode="horizontal", wrap=True,
        extras={"x-wp4-placeholder": placeholder} if placeholder else {})


#: (shape name, family, weight, italic) → (typeface, b, i) on test-16x9 (major "Calibri Light",
#: minor "Calibri"), with theme tokens on and off. 16-WPG §2.3's matrix.
WEIGHT_MATRIX = [
    ("calibri-300", "Calibri", 300, False, ("Calibri Light", "0", "0"), ("Calibri Light", "0", "0")),
    ("calibri-400", "Calibri", 400, False, ("+mn-lt", "0", "0"), ("Calibri", "0", "0")),
    ("calibri-500", "Calibri", 500, False, ("+mn-lt", "0", "0"), ("Calibri", "0", "0")),
    ("calibri-600", "Calibri", 600, False, ("+mn-lt", "1", "0"), ("Calibri", "1", "0")),
    ("calibri-700", "Calibri", 700, False, ("+mn-lt", "1", "0"), ("Calibri", "1", "0")),
    ("light-400", "Calibri Light", 400, False, ("+mj-lt", "0", "0"), ("Calibri Light", "0", "0")),
    ("light-700-h4", "Calibri Light", 700, False, ("+mj-lt", "1", "0"), ("Calibri Light", "1", "0")),
    ("calibri-300-italic", "Calibri", 300, True, ("Calibri Light", "0", "1"), ("Calibri Light", "0", "1")),
    ("segoe-300", "Segoe UI", 300, False, ("Segoe UI Light", "0", "0"), ("Segoe UI Light", "0", "0")),
    ("segoe-500", "Segoe UI", 500, False, ("Segoe UI Semibold", "0", "0"), ("Segoe UI Semibold", "0", "0")),
    ("segoe-600", "Segoe UI", 600, False, ("Segoe UI Semibold", "0", "0"), ("Segoe UI Semibold", "0", "0")),
    ("segoe-700", "Segoe UI", 700, False, ("Segoe UI", "1", "0"), ("Segoe UI", "1", "0")),
    ("segoe-800", "Segoe UI", 800, False, ("Segoe UI Black", "0", "0"), ("Segoe UI Black", "0", "0")),
    ("segoe-900", "Segoe UI", 900, False, ("Segoe UI Black", "0", "0"), ("Segoe UI Black", "0", "0")),
    ("arial-300", "Arial", 300, False, ("Arial", "0", "0"), ("Arial", "0", "0")),
    ("arial-800", "Arial", 800, False, ("Arial Black", "0", "0"), ("Arial Black", "0", "0")),
    ("arial-900", "Arial", 900, False, ("Arial Black", "0", "0"), ("Arial Black", "0", "0")),
    ("rounded-400", "Arial Rounded MT Bold", 400, False,
     ("Arial Rounded MT Bold", "0", "0"), ("Arial Rounded MT Bold", "0", "0")),
    ("variable-700", "Bahnschrift", 700, False, ("Bahnschrift", "1", "0"), ("Bahnschrift", "1", "0")),
    ("absent-600", "Lexend", 600, False, ("Lexend", "1", "0"), ("Lexend", "1", "0")),
]


def _run_properties(pptx: Path, name: str) -> tuple[str | None, str | None, str | None]:
    rPr = by_name(pptx, name).text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    latin = rPr.find(f"{{{A}}}latin")
    return (latin.get("typeface") if latin is not None else None), rPr.get("b"), rPr.get("i")


def test_weights_become_installed_faces(tmp_path):
    """Each weight names the face the browser drew, with `b` only for a Bold member or synthetic bold."""
    # The matrix is of Windows/Office faces (Segoe UI's five weights, Calibri, a variable Bahnschrift).
    require_faces("Segoe UI", "Calibri", "Calibri Light", "Arial", "Arial Rounded MT Bold", "Bahnschrift")
    assert text_engine.face_for_exact("Bahnschrift").variable
    assert text_engine.face_for_exact("Lexend") is None, "Lexend must stay absent here (Peter #10)"

    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")
    master = config.MASTERS_FIXTURE / "test-16x9.pptx"
    elements = [_one_line(name, _run(f"{name} Hamburgefonstiv", family, weight, italic), 40,
                          120 + 26 * index, line_height=26.0)
                for index, (name, family, weight, italic, _, _) in enumerate(WEIGHT_MATRIX)]
    # The 700 title on the Calibri-Light major: what `font_forcing_css` makes of a bold title.
    elements.append(_one_line("title-700", _run("Why balance-sheet banks", "Calibri Light", 700,
                                                 size=30.0), 48, 30, placeholder="title", line_height=40.0))
    assign_ids_and_z(elements)
    ir = IR(canvas=Canvas(1280, 720), slide=Slide(id="weights", layoutId="layout-02"), elements=elements)
    ir = map_placeholders(ir, manifest.layout("layout-02"))
    assert any((e.placeholder or {}).get("type") == "title" for e in ir.elements)

    tokens, literal = tmp_path / "tokens.pptx", tmp_path / "literal.pptx"
    emitter.emit([ir], manifest, master, tokens, EmitOptions())
    emitter.emit([ir], manifest, master, literal, EmitOptions(theme_fonts=False))

    wrong = []
    for name, family, weight, italic, with_tokens, without in WEIGHT_MATRIX:
        for deck, expected in ((tokens, with_tokens), (literal, without)):
            got = _run_properties(deck, name)
            if got != expected:
                wrong.append(f"{deck.stem} {name} ({family} {weight}{' italic' if italic else ''}): "
                             f"{got} != {expected}")
    title = next(s for s in shapes_of(tokens) if s.is_placeholder and s.has_text_frame
                 and "Why balance" in s.text_frame.text)
    rPr = title.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    assert (rPr.find(f"{{{A}}}latin").get("typeface"), rPr.get("b"), rPr.get("i")) == ("+mj-lt", "1", "0")
    assert not wrong, "\n".join(wrong)


def test_bold_is_explicit_both_ways(sample_master, tmp_path):
    """A regular run in a placeholder whose layout title style is bold must not inherit the bold."""
    from pptx import Presentation as _Presentation

    from app.engine.importer import import_master

    master = _Presentation(str(sample_master))
    title_style = master.slide_master.element.find(f"{{{P}}}txStyles/{{{P}}}titleStyle/{{{A}}}lvl1pPr")
    title_style.find(f"{{{A}}}defRPr").set("b", "1")
    bold_master = tmp_path / "bold-title-master.pptx"
    master.save(str(bold_master))
    manifest = import_master(bold_master, tmp_path / "import", render=False)
    layout = next(lay for lay in manifest.layouts if lay.placeholder("title"))
    zone = layout.placeholder("title")
    assert zone is not None and (zone.style or {}).get("bold") is True, "the precondition: a bold title style"

    element = _one_line("regular-title", _run("A regular-weight title", "Arial", 400, size=28.0),
                        zone.x, zone.y, placeholder="title", line_height=36.0)
    assign_ids_and_z([element])
    ir = IR(canvas=Canvas(manifest.canvas_w, manifest.canvas_h),
            slide=Slide(id="regular-title", layoutId=layout.id), elements=[element])
    ir = map_placeholders(ir, layout)
    deck = tmp_path / "regular-title.pptx"
    emitter.emit([ir], manifest, bold_master, deck, EmitOptions())
    title = next(s for s in shapes_of(deck) if s.is_placeholder and s.has_text_frame
                 and "regular-weight" in s.text_frame.text)
    rPr = title.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    assert (rPr.get("b"), rPr.get("i")) == ("0", "0")


def _ink(image: Image.Image, box: Box) -> float:
    """Mean darkness of one row's crop (0 = white): heavier faces put more ink in the same text."""
    import numpy as np

    crop = np.asarray(image.convert("L").crop((int(box.x), int(box.y), int(box.x2), int(box.y2))),
                      dtype=np.float64)
    return float((255.0 - crop).mean())


@pytest.mark.renderer_full
def test_named_faces_render_at_their_weight(tmp_path):
    """PptxRender draws the named faces at their weight: Segoe UI 400 < 600 < 700, Arial 400 < 900.

    One measured exception, a renderer gap and not the file: PptxRender draws `"Calibri Light"` as
    Calibri **Regular**. `FontResolver.MatchFamily` consults its alias table (`["Calibri Light"] =
    {"Carlito", "Calibri"}`) before `TryWeightSuffix`, so the name lands on the Calibri family at
    the requested weight, 400 — and reports it under `fontsSubstituted`. PowerPoint draws the Light
    face (the bench's PowerPoint panel). The assertion pins the gap so that a fixed renderer fails
    here and the check can tighten to Calibri 300 < 400.
    """
    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")
    rows = [("Calibri", 300), ("Calibri", 400), ("Calibri", 700),
            ("Segoe UI", 400), ("Segoe UI", 600), ("Segoe UI", 700), ("Arial", 400), ("Arial", 900)]
    text = "The quick brown fox jumps over the lazy dog"
    elements = [_one_line(f"{family}-{weight}", _run(text, family, weight, size=28.0), 40,
                          60 + 80 * index, line_height=40.0)
                for index, (family, weight) in enumerate(rows)]
    assign_ids_and_z(elements)
    ir = IR(canvas=Canvas(1280, 720), slide=Slide(id="ink", layoutId="layout-07"), elements=elements)
    deck = tmp_path / "ink.pptx"
    emitter.emit([ir], manifest, config.MASTERS_FIXTURE / "test-16x9.pptx", deck, EmitOptions())
    [png] = renderer.render(deck, 1280, tmp_path / "png")
    with Image.open(png) as image:
        ink = {(family, weight): _ink(image, Box(20, 60 + 80 * index - 10, 900, 60))
               for index, (family, weight) in enumerate(rows)}
    assert ink[("Segoe UI", 400)] < ink[("Segoe UI", 600)] < ink[("Segoe UI", 700)], ink
    assert ink[("Arial", 400)] < ink[("Arial", 900)], ink
    assert ink[("Calibri", 400)] < ink[("Calibri", 700)], ink
    substituted = json.loads((tmp_path / "png" / "warnings.json").read_text(encoding="utf-8"))
    assert "Calibri Light" in (substituted.get("fontsSubstituted") or []), substituted
    assert ink[("Calibri", 300)] == pytest.approx(ink[("Calibri", 400)], rel=0.01), (
        f"PptxRender now draws 'Calibri Light' lighter than Calibri: tighten this to 300 < 400 ({ink})")


def test_bullets_are_real_bullets(emitted):
    out, _, _, _, _ = emitted["emit-text"]
    bulleted = by_name(out, "paragraphs").text_frame.paragraphs[1]
    pPr = bulleted._p.get_or_add_pPr()
    assert pPr.find(f"{{{A}}}buChar").get("char") == "•"
    assert pPr.find(f"{{{A}}}buClr/{{{A}}}srgbClr").get("val") == "1A9AFB"
    assert int(pPr.get("marL")) == 18 * config.EMU_PER_PX
    assert int(pPr.get("indent")) == -18 * config.EMU_PER_PX
    assert "•" not in "".join(r.text for r in bulleted.runs)

    numbered = by_name(out, "numbered").text_frame.paragraphs[0]._p.get_or_add_pPr()
    assert numbered.find(f"{{{A}}}buAutoNum").get("type") == "arabicPeriod"


def test_paragraphs_without_a_bullet_say_so(emitted):
    """A body placeholder inherits the master's list style; silence would grow a bullet."""
    out, _, _, _, _ = emitted["emit-placeholders"]
    body = by_name(out, "body-by-overlap").text_frame
    for paragraph in body.paragraphs:
        assert paragraph._p.get_or_add_pPr().find(f"{{{A}}}buNone") is not None


def test_exact_line_spacing_and_paragraph_gap(emitted):
    out, _, _, _, _ = emitted["emit-text"]
    frame = by_name(out, "paragraphs").text_frame
    first = frame.paragraphs[0]._p.get_or_add_pPr()
    assert int(first.find(f"{{{A}}}lnSpc/{{{A}}}spcPts").get("val")) == round(30 * 0.75 * 100)
    second = frame.paragraphs[1]._p.get_or_add_pPr()
    # Line centres are 42 px apart across the paragraph break, 30 px of which is the line itself: its
    # line top is 12 px below the previous line's bottom. A one-line paragraph at 1.5 em is written at
    # its glyphs' height (r2: `_fit_lone_line` fits every one-line paragraph), shortened by whole pixels
    # the space before it takes back — so gap + spacing, and with them the baseline, stay 12 + 30 px.
    gap = int(second.find(f"{{{A}}}spcBef/{{{A}}}spcPts").get("val"))
    spacing = int(second.find(f"{{{A}}}lnSpc/{{{A}}}spcPts").get("val"))
    assert gap + spacing == round(12 * 0.75 * 100) + round(30 * 0.75 * 100)
    assert spacing < text_engine.POWERPOINT_SINGLE_EM * 20 * 0.75 * 100 and (30 * 75 - spacing) % 75 == 0


def test_uppercase_is_a_capital_style_not_a_rewritten_string(emitted):
    out, _, _, _, _ = emitted["emit-text"]
    shouty = by_name(out, "uppercase").text_frame.paragraphs[0].runs[0]
    assert shouty._r.get_or_add_rPr().get("cap") == "all"
    assert shouty.text == "shouty small print"          # the text itself is never mangled


def test_vertical_text_is_a_quarter_turned_box(emitted):
    """Vertical text keeps every calibrated number by being a rotated horizontal box."""
    out, _, ir, _, _ = emitted["emit-text"]
    element = next(e for e in ir.elements if e.name == "vertical")
    shape = by_name(out, "vertical")
    assert shape.rotation == 270
    assert shape._element.find(f".//{{{A}}}bodyPr").get("vert") is None
    # Un-rotated the box is wide and short; rotation about the centre stands it up in place.
    assert shape.width > shape.height
    centre_x = (shape.left + shape.width / 2) / config.EMU_PER_PX
    centre_y = (shape.top + shape.height / 2) / config.EMU_PER_PX
    assert centre_x == pytest.approx(element.box.cx, abs=3.0)
    assert centre_y == pytest.approx(element.box.cy, abs=8.0)


def test_text_box_is_wide_enough_never_to_rewrap(emitted):
    out, _, ir, _, _ = emitted["emit-text"]
    for element in ir.elements:
        if element.kind != "text" or element.writingMode == "vertical":
            continue
        shape = by_name(out, element.name)
        widest = max(line["box"]["w"] for p in element.paragraphs for line in p["lines"])
        assert shape.width >= widest * config.EMU_PER_PX, element.name


def test_no_text_box_hangs_off_the_slide_or_is_a_letter_wide(emitted):
    """Two fit-check rows the emitter can simply never fail (master brief §9)."""
    for name, (out, _, ir, _, _) in emitted.items():
        for shape in shapes_of(out):
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            assert shape.left >= 0, f"{name}/{shape.name} starts off the left edge"
            assert shape.left + shape.width <= ir.canvas.w * config.EMU_PER_PX + 1, \
                f"{name}/{shape.name} runs off the right edge"
            if len(shape.text_frame.text.strip()) > 1:
                assert shape.width >= 0.2 * config.EMU_PER_IN, f"{name}/{shape.name} is a sliver"


def test_alignment_decides_which_way_the_slack_grows(emitted):
    """Widening a right-aligned box to the right would move the text; it grows left instead, and
    its right edge is pinned where the browser ended the line (WP-B E2) — in this hand-written
    fixture 13 px past the element's box, which the old rule clamped the text back into."""
    out, _, ir, _, _ = emitted["emit-text"]
    right = by_name(out, "line-height-1")
    element = next(e for e in ir.elements if e.name == "line-height-1")
    line = element.paragraphs[0]["lines"][0]["box"]
    assert right.left < element.box.x * config.EMU_PER_PX
    assert right.left + right.width == pytest.approx(
        (line["x"] + line["w"]) * config.EMU_PER_PX, abs=0.5 * config.EMU_PER_PX)
    pPr = right.text_frame.paragraphs[0]._p.get_or_add_pPr()
    assert int(pPr.get("marR")) == 0

    left = by_name(out, "paragraphs")
    left_element = next(e for e in ir.elements if e.name == "paragraphs")
    assert left.left == left_element.box.x * config.EMU_PER_PX


def test_width_lock_pulls_a_line_onto_its_measured_width_both_ways():
    """A line PowerPoint would draw too wide is squeezed; one it would draw too narrow is spread."""
    face = text_engine.face_for("Arial")
    assert face is not None
    predicted = text_engine.text_width_px("Squeeze me", "Arial", 20.0, 400, False)
    runs = [{"text": "Squeeze me", "font": "Arial", "sizePx": 20.0, "weight": 400}]  # 10 characters

    narrow = text_engine.LineLayout(runs=runs, ascent=0.0, descent=0.0, centre=0.0,
                                    width=predicted - 5.0, predicted=predicted)
    assert text_engine._tracking_for(narrow, []) == pytest.approx(-5.0 / 10 * config.WIDTH_LOCK_FACTOR)

    wide = text_engine.LineLayout(runs=runs, ascent=0.0, descent=0.0, centre=0.0,
                                  width=predicted + 5.0, predicted=predicted)
    assert text_engine._tracking_for(wide, []) == pytest.approx(+5.0 / 10 * config.WIDTH_LOCK_FACTOR)

    on = text_engine.LineLayout(runs=runs, ascent=0.0, descent=0.0, centre=0.0,
                                width=predicted, predicted=predicted)
    assert text_engine._tracking_for(on, []) == 0.0


def test_width_lock_is_bounded_by_the_per_character_ratio_in_both_directions():
    """A huge mismatch (a substituted font) is capped, not obeyed — and it is reported."""
    runs = [{"text": "abcd", "font": "Arial", "sizePx": 20.0, "weight": 400}]
    limit = text_engine._MAX_TRACKING_RATIO * 20.0
    warnings: list[str] = []
    crushed = text_engine.LineLayout(runs=runs, ascent=0.0, descent=0.0, centre=0.0,
                                     width=1.0, predicted=1000.0)
    assert text_engine._tracking_for(crushed, warnings) == pytest.approx(-limit)
    gappy = text_engine.LineLayout(runs=runs, ascent=0.0, descent=0.0, centre=0.0,
                                   width=1000.0, predicted=1.0)
    assert text_engine._tracking_for(gappy, warnings) == pytest.approx(+limit)
    assert len(warnings) == 2


def test_first_baseline_residual_falls_back_from_face_to_family():
    """`_dy_factor` prefers an exact face key, then the family, then the empty fallback."""
    table = {"": -0.02, "Calibri": -0.07, "Calibri|700": -0.05}
    assert text_engine._dy_factor(table, "Calibri", 700, False) == -0.05   # exact face
    assert text_engine._dy_factor(table, "Calibri", 400, False) == -0.07   # family fallback
    assert text_engine._dy_factor(table, "Nunito", 400, False) == -0.02    # empty fallback
    assert text_engine._dy_factor({}, "Calibri", 400, False) == 0.0        # nothing at all


def _forget_fonts() -> None:
    """Drop every cached font scan, so the next lookup reads `config.FONT_DIRS` as it is now."""
    from app.engine.extract import html as html_extract

    text_engine.forget_fonts()           # the one index: fit's `resolve_face` reads it too (plan D4)
    html_extract.forget_installed_families()


@contextlib.contextmanager
def _blind_fonts(monkeypatch, tmp_path: Path):
    """An export that cannot see a single font file, while the renderers still see every one.

    `config.FONT_DIRS` → an empty directory: `face_for` returns None for every family, so the
    emitter keeps `predicted = measured`, adds no tracking, falls back to 0.8/0.2 em metrics and
    warns once per element that the family is "not installed — measured with a fallback". PptxRender
    and PowerPoint read the machine's own fonts, so they draw the real face: this is how the export
    is made blind to the width the target will draw (fidelity F4, the `emit-wrap` fixture). Fonts
    are restored — and every cached lookup forgotten — on the way out, whatever happened inside.
    """
    empty = tmp_path / "no-fonts"
    empty.mkdir(exist_ok=True)
    try:
        with monkeypatch.context() as patch:
            patch.setattr(config, "FONT_DIRS", (empty,))
            _forget_fonts()
            assert text_engine.face_for("Arial") is None, "the blind-fonts route still sees Arial"
            yield empty
    finally:
        _forget_fonts()


def _ink_bands(png: Path, y0: int, y1: int, x0: int = 20, x1: int = 1279) -> list[tuple[int, int]]:
    """Rows in `y0..y1` with any pixel < 250 across `x0..x1`, merged by adjacency — one per text line.

    The scan `test_first_baseline_calibration` uses for the first ink row, kept for every row.
    """
    image = Image.open(png).convert("L")
    pixels = image.load()
    bands: list[list[int]] = []
    for y in range(y0, min(y1, image.height - 1) + 1):
        if any(pixels[x, y] < 250 for x in range(x0, min(x1, image.width - 1) + 1)):
            if bands and bands[-1][1] == y - 1:
                bands[-1][1] = y
            else:
                bands.append([y, y])
    return [(top, bottom) for top, bottom in bands]


def test_fonts_in_subfolders_are_found_by_every_scan(tmp_path, monkeypatch):
    """Linux keeps fonts in subfolders (`/usr/share/fonts/truetype/msttcorefonts/Arial.ttf`).

    A flat listing found no font there at all, so the Docker image measured every line with the
    last-resort average advance and the width lock squeezed the words into each other.
    """
    from app.engine.extract import html as html_extract
    from app.engine.verify import fit as fit_module

    arial = text_engine.face_for("Arial")
    assert arial is not None
    nested = tmp_path / "truetype" / "msttcorefonts"
    nested.mkdir(parents=True)
    (nested / arial.path.name).write_bytes(arial.path.read_bytes())
    monkeypatch.setattr(config, "FONT_DIRS", (tmp_path,))
    _forget_fonts()
    try:
        assert config.font_files() == [nested / arial.path.name]
        assert text_engine.face_for("Arial") is not None
        assert fit_module.resolve_face("Arial", bold=False, italic=False) is not None
        assert html_extract.installed_families(["Arial"]) == {"Arial": True}
    finally:
        _forget_fonts()


def test_width_lock_does_not_squeeze_toward_a_guess(tmp_path, monkeypatch):
    """With no font installed the prediction is an average advance, and the browser's width stands."""
    text = "production features, executed not just"
    measured = text_engine.text_width_px(text, "Arial", 10.0, 400, False)
    element = Element(
        kind="text", box=Box(40, 100, 600, 13), name="caption",
        paragraphs=[{"align": "left", "lineHeightPx": 13, "spaceBeforePx": 0, "spaceAfterPx": 0,
                     "bullet": None,
                     "lines": [{"box": {"x": 40, "y": 100, "w": measured, "h": 13},
                                "runs": [{"text": text, "font": "Arial", "sizePx": 10.0,
                                          "weight": 400, "italic": False}]}]}],
    )
    ctx = SimpleNamespace(options=EmitOptions(), canvas=SimpleNamespace(w=1280))
    monkeypatch.setattr(config, "FONT_DIRS", (tmp_path,))
    _forget_fonts()
    try:
        assert text_engine.face_for("Arial") is None
        guess = text_engine.text_width_px(text, "Arial", 10.0, 400, False)
        assert guess > measured * 1.1, "the average advance should overshoot Arial here"
        layout = text_engine.plan_text(element, ctx)
        assert [line.tracking_px for p in layout.paragraphs for line in p.lines] == [0.0]
    finally:
        _forget_fonts()


# ------------------------------------------------------------------ WP-B: horizontal placement (§3.6)


def _placed_paragraph(lines: list[tuple[float, float, str]], y: float, *, align: str = "left",
                      size: float = 20.0, line_height: float = 30.0, bullet: dict | None = None) -> dict:
    """One IR paragraph whose lines sit at `(x, width, text)`, one under the other from `y`."""
    return {"align": align, "lineHeightPx": line_height, "spaceBeforePx": 0, "spaceAfterPx": 0,
            "bullet": bullet,
            "lines": [{"box": {"x": x, "y": y + i * line_height, "w": w, "h": line_height},
                       "runs": [_run(text, "Arial", size=size)]}
                      for i, (x, w, text) in enumerate(lines)]}


def _plan(element: Element) -> text_engine.TextLayout:
    ctx = SimpleNamespace(options=EmitOptions(), canvas=SimpleNamespace(w=1280, h=720))
    return text_engine.plan_text(element, ctx)


def _arial(text: str, size: float = 20.0) -> float:
    return text_engine.text_width_px(text, "Arial", size, 400, False)


def _band(layout: text_engine.TextLayout, index: int) -> tuple[float, float]:
    """Where PowerPoint lays paragraph `index` out: `[x + marL, x2 − marR]` of the placed box."""
    paragraph = layout.paragraphs[index]
    return layout.box.x + paragraph.mar_left_px, layout.box.x2 - paragraph.mar_right_px


def test_a_mid_line_start_becomes_a_first_line_indent():
    """Line 1 starts 160 px right of line 2 (text after a swatch, a `text-indent`): `marL` from
    line 2, `indent` = the first line's offset from it — written, and back where the browser put it."""
    one, two = "first line starts late", "second line at the edge"
    element = Element(kind="text", box=Box(40, 100, 600, 60), name="indented",
                      paragraphs=[_placed_paragraph([(200, _arial(one), one), (40, _arial(two), two)], 100)])
    layout = _plan(element)
    [paragraph] = layout.paragraphs
    assert layout.box.x == pytest.approx(40)
    assert (paragraph.mar_left_px, paragraph.indent_px, paragraph.mar_right_px) == pytest.approx((0.0, 160.0, 0.0))


def test_a_hanging_first_line_gets_a_negative_indent():
    """B1 critique #14: a hand-made bullet — line 1 at the edge, lines 2.. 16 px in. `marL` comes from
    the continuation lines and the first line hangs left of it."""
    one, two, three = "- a hand-made bullet whose", "second line hangs under", "the text, not the dot"
    element = Element(kind="text", box=Box(40, 100, 600, 90), name="hanging",
                      paragraphs=[_placed_paragraph([(40, _arial(one), one), (56, _arial(two), two),
                                                     (56, _arial(three), three)], 100)])
    [paragraph] = _plan(element).paragraphs
    assert (paragraph.mar_left_px, paragraph.indent_px) == pytest.approx((16.0, -16.0))


def test_a_right_aligned_paragraph_narrower_than_its_element_gets_marR():
    """A left paragraph fills the element; a right-aligned one under it ends 40 px short of it. The
    box grows right (the first paragraph is left-aligned), and `marR` puts the right paragraph's end
    back where the browser put it."""
    wide, short = "A left-aligned line across the element", "ends short"
    right_end = 40 + _arial(wide) - 40
    element = Element(kind="text", box=Box(40, 100, _arial(wide), 60), name="mixed-right",
                      paragraphs=[_placed_paragraph([(40, _arial(wide), wide)], 100),
                                  _placed_paragraph([(right_end - _arial(short), _arial(short), short)], 130,
                                                    align="right")])
    layout = _plan(element)
    assert layout.box.x == pytest.approx(40), "the first paragraph is left-aligned: its edge stays"
    assert _band(layout, 1)[1] == pytest.approx(right_end, abs=0.01)
    assert layout.paragraphs[1].mar_right_px > 40


def test_a_right_aligned_element_keeps_its_right_edge_and_grows_left():
    """E2: the aligned edge is pinned. A right-aligned line whose prediction is wider than the
    browser's width grows the box to the left; its end stays where the browser ended it."""
    text = "Right-aligned value"
    width = _arial(text)
    element = Element(kind="text", box=Box(900 - width, 100, width, 30), name="right",
                      paragraphs=[_placed_paragraph([(900 - width, width, text)], 100, align="right")])
    layout = _plan(element)
    assert layout.box.x2 == pytest.approx(900, abs=0.01)
    assert layout.box.x < 900 - width
    assert layout.paragraphs[0].mar_right_px == pytest.approx(0.0, abs=0.01)


def test_a_centred_paragraph_under_a_wider_left_one_lands_within_half_a_pixel():
    """§3.6 E1: the centred line's centre is `(marL − marR)/2` off the placed box's centre, so it
    lands where the browser centred it (inside its own 200 px column, not the element's)."""
    wide, label = "A left-aligned heading that spans the whole element", "Centred"
    centre = 40 + 100
    element = Element(kind="text", box=Box(40, 100, _arial(wide), 60), name="mixed-centre",
                      paragraphs=[_placed_paragraph([(40, _arial(wide), wide)], 100),
                                  _placed_paragraph([(centre - _arial(label) / 2, _arial(label), label)], 130,
                                                    align="center")])
    layout = _plan(element)
    left, right = _band(layout, 1)
    assert (left + right) / 2 == pytest.approx(centre, abs=0.5)
    assert right - left >= _arial(label), "room for the line inside the centring band"


def test_growth_never_squeezes_a_mixed_element(monkeypatch):
    """A right-aligned line that fills the element, predicted wider than the browser drew it (a
    substituted font, beyond what the width lock may take back), under a short left-aligned label.
    The first paragraph is left-aligned, so the slack grows right — but the right-aligned line needs
    its extra width on the *left*, where it grows from its end. The box covers it (the brief's single
    `offset + width` formula did not: it sized the box right and left the line 90 px short, to re-wrap)."""
    real = text_engine.text_width_px
    monkeypatch.setattr(text_engine, "text_width_px",
                        lambda text, *args, **kwargs: real(text, *args, **kwargs) * 1.5)
    label, value = "Label", "Revenue in the first year of the programme"
    width = real(value, "Arial", 20.0, 400, False)
    element = Element(kind="text", box=Box(300, 100, width, 60), name="squeeze",
                      paragraphs=[_placed_paragraph([(300, real(label, "Arial", 20.0, 400, False), label)], 100),
                                  _placed_paragraph([(300, width, value)], 130, align="right")])
    layout = _plan(element)
    line = layout.paragraphs[1].lines[0]
    assert line.drawn() > width + 50, "the width lock cannot take the whole overshoot back"
    left, right = _band(layout, 1)
    assert right == pytest.approx(300 + width, abs=0.01), "the value still ends where the browser ended it"
    assert right - left >= line.drawn() - 0.01
    assert layout.box.x < 300 and _band(layout, 0)[0] == pytest.approx(300, abs=0.01)


def test_a_right_aligned_element_predicted_wider_does_not_rewrap(tmp_path, monkeypatch):
    """§4: a right-aligned two-line element whose font PowerPoint draws 40 % wider than the browser
    did (the emitter's prediction and fit's both patched, as a substituted font would): the box grows
    left from the pinned right edge by what the width lock cannot take back, and fit predicts the
    browser's two lines — no re-wrap. (fit reads the box's inner width, not `marL`/`marR`: WP-D's
    region, proposed.)"""
    from app.engine.verify import fit as fit_module

    real_emit, real_fit = text_engine.text_width_px, fit_module.face_width_px
    monkeypatch.setattr(text_engine, "text_width_px", lambda *a, **k: real_emit(*a, **k) * 1.4)
    monkeypatch.setattr(fit_module, "face_width_px", lambda *a, **k: real_fit(*a, **k) * 1.4)
    # One token per line: fit's `_space_width` leaves the run's `spc` out of every space (a finding
    # for fit's owner), which alone would predict a wrap once the width lock tracks a line tighter.
    one, two = "Right-aligned-statement-whose-font", "draws-wider-in-PowerPoint"
    w1, w2 = real_emit(one, "Arial", 20.0, 400, False), real_emit(two, "Arial", 20.0, 400, False)
    element = Element(kind="text", box=Box(1000 - w1, 100, w1, 60), name="wider", anchor="top",
                      writingMode="horizontal", wrap=True,
                      paragraphs=[_placed_paragraph([(1000 - w1, w1, one), (1000 - w2, w2, two)], 100,
                                                    align="right")])
    assign_ids_and_z([element])
    ir = IR(canvas=Canvas(1280, 720), slide=Slide(id="w", layoutId="layout-07"), elements=[element])
    deck = tmp_path / "wider.pptx"
    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")
    emitter.emit([ir], manifest, config.MASTERS_FIXTURE / "test-16x9.pptx", deck, EmitOptions())
    shape = next(s for s in Presentation(str(deck)).slides[0].shapes if s.name == "wider")
    assert (shape.left + shape.width) / config.EMU_PER_PX == pytest.approx(1000, abs=0.01)
    assert shape.left / config.EMU_PER_PX < 1000 - w1
    report = fit_module.fit(deck, [ir])
    assert [issue for issue in report.issues if issue.kind in ("line-count", "overflow")] == []


def test_a_bullet_hangs_from_the_measured_text_start():
    """The measured hanging indent: text at `marL` (where the item's text starts), bullet at the
    box's left edge (the list's) — today's `indentPx` for an ordinary item."""
    text = "A bulleted item"
    marker = {"type": "char", "char": "•", "level": 0, "indentPx": 24, "color": "000000", "start": 1}
    element = Element(kind="text", box=Box(40, 100, 300, 30), name="bullet",
                      paragraphs=[_placed_paragraph([(64, _arial(text), text)], 100, bullet=marker)])
    [paragraph] = _plan(element).paragraphs
    assert (paragraph.mar_left_px, paragraph.indent_px) == pytest.approx((24.0, -24.0))


def test_two_paragraphs_side_by_side_are_reported_not_stacked_silently():
    """E3: a defence against an extractor regression — a row merged into one element."""
    element = Element(kind="text", box=Box(40, 100, 600, 30), name="row",
                      paragraphs=[_placed_paragraph([(40, _arial("Label"), "Label")], 100),
                                  _placed_paragraph([(400, _arial("Value"), "Value")], 100)])
    warnings = _plan(element).warnings
    assert any("is not below paragraph 1" in warning for warning in warnings), warnings


def test_margins_are_written_on_every_paragraph(tmp_path):
    """`_set_indent` writes `marL`, `marR` and `indent` on every paragraph, so a placeholder's list
    style cannot leak one in; the values are the plan's."""
    one, two = "first line starts late", "second line at the edge"
    element = Element(kind="text", box=Box(40, 100, 600, 60), name="indented",
                      paragraphs=[_placed_paragraph([(200, _arial(one), one), (40, _arial(two), two)], 100)],
                      anchor="top", writingMode="horizontal", wrap=True)
    assign_ids_and_z([element])
    ir = IR(canvas=Canvas(1280, 720), slide=Slide(id="m", layoutId="layout-07"), elements=[element])
    deck = tmp_path / "margins.pptx"
    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")
    emitter.emit([ir], manifest, config.MASTERS_FIXTURE / "test-16x9.pptx", deck, EmitOptions())
    shape = next(s for s in Presentation(str(deck)).slides[0].shapes if s.name == "indented")
    pPr = shape.text_frame.paragraphs[0]._p.get_or_add_pPr()
    assert (pPr.get("marL"), pPr.get("marR"), pPr.get("indent")) == ("0", "0", str(160 * config.EMU_PER_PX))


#: §3.9, measured 2026-09-25 on this machine: the deck below, rendered by PowerPoint 365 through
#: `engine/tools/render_pptx.ps1` (dev-time, once), puts the first ink of each line at the IR's x + 1 px
#: (the 'H' side bearing) — line 2 of `indent+160` at 41, not 201; the hanging case's line 1 at 41 and
#: lines 2–3 at 57; `indent+80x3`'s lines 2 and 3 at 41. `indent` moves the first line only; every line
#: after an `a:br` starts at `marL`. PptxRender measured the same pixels (below): no renderer gap.
INDENT_AFTER_BREAK_POWERPOINT = {
    "indent+160": [201, 41], "hang-16": [41, 57, 57], "indent+80x3": [121, 41, 41],
}


@pytest.mark.renderer_full
def test_indent_after_break_matches_powerpoint(tmp_path):
    """§3.9: E1 relies on continuation lines after a soft break starting at `marL`, not at
    `marL + indent` — PowerPoint's behaviour, measured (above). PptxRender, which the gate uses, must
    agree: ink x of every line within 1 px of PowerPoint's."""
    height = 30.0
    cases = [
        ("indent+160", Box(40, 60, 700, 2 * height), [(200, "HHHH first line starts late"),
                                                       (40, "HHHH second line at marL")]),
        ("hang-16", Box(40, 200, 700, 3 * height), [(40, "HHHH first line hangs left"),
                                                    (56, "HHHH second line at marL"),
                                                    (56, "HHHH third line at marL")]),
        ("indent+80x3", Box(40, 340, 700, 3 * height), [(120, "HHHH first line starts late"),
                                                        (40, "HHHH second line at marL"),
                                                        (40, "HHHH third line at marL")]),
    ]
    elements = [Element(kind="text", box=box, name=name, anchor="top", writingMode="horizontal", wrap=True,
                        paragraphs=[_placed_paragraph([(x, _arial(t), t) for x, t in lines], box.y,
                                                      line_height=height)])
                for name, box, lines in cases]
    assign_ids_and_z(elements)
    ir = IR(canvas=Canvas(1280, 720), slide=Slide(id="indent", layoutId="layout-07"), elements=elements)
    deck = tmp_path / "indent.pptx"
    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")
    emitter.emit([ir], manifest, config.MASTERS_FIXTURE / "test-16x9.pptx", deck, EmitOptions())
    written = {s.name: s.text_frame.paragraphs[0]._p.get_or_add_pPr()
               for s in Presentation(str(deck)).slides[0].shapes}
    assert written["indent+160"].get("indent") == str(160 * config.EMU_PER_PX)
    assert (written["hang-16"].get("marL"), written["hang-16"].get("indent")) == (
        str(16 * config.EMU_PER_PX), str(-16 * config.EMU_PER_PX))
    [png] = renderer.render(deck, 1280, tmp_path / "png")
    with Image.open(png) as image:
        grey = image.convert("L")
        pixels = grey.load()

        def ink_left(y0: float, y1: float) -> int | None:
            for x in range(0, 1280):
                if any(pixels[x, y] < 128 for y in range(int(y0), int(y1))):
                    return x
            return None

        measured = {name: [ink_left(box.y + i * height + 4, box.y + (i + 1) * height - 4)
                           for i in range(len(lines))]
                    for name, box, lines in cases}
    for name, expected in INDENT_AFTER_BREAK_POWERPOINT.items():
        for got, want in zip(measured[name], expected, strict=False):
            assert got is not None and abs(got - want) <= 1, (name, measured[name], expected)


def browser_ascent_descent(face) -> tuple[float, float]:
    """The ascent/descent (em) the measuring browser centres a line's glyphs with.

    Edge on Windows lays text out with DirectWrite, whose font metrics are the OS/2 **win**
    ascent/descent unless the font sets USE_TYPO_METRICS (then the typo pair). For Arial, Segoe UI
    and Times New Roman that equals `hhea`, which is what `text.Face` carries; for Calibri it does
    not — win 0.952 / 0.269 against hhea 0.750 / 0.250 — and Edge puts Calibri's baseline 0.340 em
    below the line-box centre (measured 2026-09-25, 100 px text, 200 px line height), not the 0.250
    the `hhea` pair predicts. The emitter's own model still uses `hhea`; `TEXT_DY_BY_FONT` carries
    the difference (see its comment), which is why the expected position here must not.
    """
    from fontTools.ttLib import TTCollection, TTFont

    font = (TTCollection(str(face.path)).fonts[face.index]
            if face.path.suffix.lower() in (".ttc", ".otc") else TTFont(str(face.path), lazy=True))
    units = float(font["head"].unitsPerEm)
    os2 = font["OS/2"]
    if int(os2.fsSelection) & 0x80:                                   # USE_TYPO_METRICS
        return os2.sTypoAscender / units, abs(os2.sTypoDescender) / units
    return os2.usWinAscent / units, os2.usWinDescent / units


@pytest.mark.renderer_full
def test_first_baseline_calibration(tmp_path):
    """The measurement behind `config.TEXT_DY_BY_FONT`, run as a test so it cannot rot.

    For each case the expected ink top is computed from the browser's placement (line-box centre,
    the ascent/descent the browser centres glyphs with — `browser_ascent_descent`) plus Pillow's own
    raster of the same string, then compared with where PptxRender actually put the ink.
    Thresholding both rasters the same way makes the antialiasing bias cancel instead of being
    absorbed into the constant.
    """
    sizes = ((14, 21), (20, 20), (20, 60), (33, 33), (40, 46))
    # (Segoe UI, 600) is written as the literal "Segoe UI Semibold" face and keyed to Segoe UI's
    # residual through the face chain; "Lexend" is absent here (Peter #10), so it is measured with
    # the fallback face and keeps its own key — the two cases WP-G's face rule touches.
    cases = [(family, weight, size, line_height)
             for family, weight in (("Arial", 400), ("Calibri", 400), ("Times New Roman", 400),
                                    ("Segoe UI", 400), ("Segoe UI", 600), ("Lexend", 400))
             for size, line_height in sizes]
    sample, top = "HEIGHT hxp", 200.0
    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")

    irs, expected = [], []
    assert text_engine.face_for_exact("Lexend") is None, "the absent-family case needs Lexend absent"
    for index, (family, weight, size, line_height) in enumerate(cases):
        face = text_engine.face_for(family, weight)
        assert face is not None, f"{family} is not installed; the calibration cannot be measured"
        font = ImageFont.truetype(str(face.path), size=size, index=face.index)
        ascent, descent = font.getmetrics()
        canvas = Image.new("L", (int(font.getlength(sample)) + 20, ascent + descent + 20), 255)
        from PIL import ImageDraw
        ImageDraw.Draw(canvas).text((5, 5), sample, font=font, fill=0)
        pixels = canvas.load()
        ink_rel = next(y for y in range(canvas.height)
                       if any(pixels[x, y] < 250 for x in range(canvas.width))) - (5 + ascent)

        browser_ascent, browser_descent = browser_ascent_descent(face)
        baseline = (top + line_height / 2) + (browser_ascent * size - browser_descent * size) / 2
        element = Element(
            kind="text", box=Box(40, top, 1100, line_height), name=f"case{index}",
            paragraphs=[{"align": "left", "lineHeightPx": line_height, "spaceBeforePx": 0,
                         "spaceAfterPx": 0, "bullet": None,
                         "lines": [{"box": {"x": 40, "y": top, "w": font.getlength(sample),
                                            "h": line_height},
                                    "runs": [{"text": sample, "font": family, "sizePx": float(size),
                                              "weight": weight, "italic": False, "color": "000000",
                                              "alpha": 1, "underline": False, "strike": False,
                                              "letterSpacingPx": 0, "baseline": "normal",
                                              "baselineShiftPx": 0, "caps": False, "href": None}]}]}],
            anchor="top", writingMode="horizontal", wrap=True)
        assign_ids_and_z([element])
        irs.append(IR(canvas=Canvas(1280, 720), slide=Slide(id=f"c{index:02d}", layoutId="layout-07"),
                      elements=[element]))
        expected.append((f"{family} {weight}", size, baseline + ink_rel))

    deck = tmp_path / "calibration.pptx"
    emitter.emit(irs, manifest, config.MASTERS_FIXTURE / "test-16x9.pptx", deck, EmitOptions())
    pngs = renderer.render(deck, 1280, tmp_path / "png")
    assert len(pngs) == len(cases)

    residuals = []
    for png, (family, size, expected_top) in zip(pngs, expected, strict=False):
        image = Image.open(png).convert("L")
        pixels = image.load()
        ink = next((y for y in range(image.size[1])
                    if any(pixels[x, y] < 250 for x in range(20, 1000))), None)
        assert ink is not None, f"{family} {size}px rendered no ink"
        residuals.append((family, size, ink - expected_top))

    worst = max(abs(residual) for _, _, residual in residuals)
    assert worst <= 2.0, f"first baseline drifts by {worst:.2f} px: {residuals}"


# ------------------------------------------------------ a lone line is written at its own height (r1c)


def _lone_element(family: str, weight: int, size: float, line_height: float, *, lines: int = 1,
                  paragraphs: int = 1, y: float = 112.0, anchor: str = "top") -> Element:
    """`paragraphs` × `lines` lines of `HHxH` at `line_height` px, each line box where a browser puts it."""
    width = text_engine.text_width_px("HHxH", family, size, weight, False)
    blocks = [{"align": "left", "lineHeightPx": line_height, "spaceBeforePx": 0, "spaceAfterPx": 0, "bullet": None,
               "lines": [{"box": {"x": 300.0, "y": y + (p * lines + k) * line_height, "w": width, "h": line_height},
                          "runs": [_run("HHxH", family, weight=weight, size=size)]} for k in range(lines)]}
              for p in range(paragraphs)]
    return Element(kind="text", box=Box(300.0, y, width, line_height * lines * paragraphs), name="lone",
                   anchor=anchor, writingMode="horizontal", wrap=True, paragraphs=blocks)


def _plan_with(element: Element, *, rule: bool = True, line_lock: bool = True) -> text_engine.TextLayout:
    ctx = SimpleNamespace(options=EmitOptions(line_lock=line_lock), canvas=SimpleNamespace(w=1280, h=720))
    if rule:
        return text_engine.plan_text(element, ctx)
    original = text_engine._fit_lone_line
    text_engine._fit_lone_line = lambda planned: None
    try:
        return text_engine.plan_text(element, ctx)
    finally:
        text_engine._fit_lone_line = original


def test_a_lone_tall_line_is_written_at_its_own_height(tmp_path):
    """a client slide's operators: `×` in 20 px bold Arial, `height:60px; line-height:60px`. PowerPoint drew
    them 10 px above the browser and PptxRender (r1's arbitration), because an exact spacing at or above
    1.2 em puts its baseline about 0.75 of the line down in PowerPoint and at the line less the descent
    in PptxRender. The line is written at its glyphs' own height — shortened by whole pixels, the frame
    moved down by as many — so the model's baseline (frame top + spacing − descent) does not move, and
    the file says so: `a:spcPts`, the frame's `off`/`ext`, `wrap="none"`, zero insets, the anchor kept."""
    element = _lone_element("Arial", 700, 20.0, 60.0)
    layout, before = _plan_with(element), _plan_with(element, rule=False)
    [paragraph] = layout.paragraphs
    line = paragraph.lines[0]
    natural = line.ascent + line.descent
    assert natural - 1.0 < paragraph.spacing_px <= natural, (paragraph.spacing_px, natural)
    shortened = 60.0 - paragraph.spacing_px
    assert shortened >= 1.0 and shortened == pytest.approx(round(shortened), abs=1e-9)
    assert paragraph.spacing_px < text_engine.POWERPOINT_SINGLE_EM * 20.0
    assert layout.box.h == pytest.approx(paragraph.spacing_px)
    assert layout.box.y == pytest.approx(before.box.y + shortened)
    assert layout.box.y + paragraph.spacing_px == pytest.approx(before.box.y + 60.0), "the baseline stays put"
    assert (layout.box.x, layout.box.w) == (before.box.x, before.box.w)

    assign_ids_and_z([element])
    ir = IR(canvas=Canvas(1280, 720), slide=Slide(id="lone", layoutId="layout-07"), elements=[element])
    deck = tmp_path / "lone.pptx"
    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")
    emitter.emit([ir], manifest, config.MASTERS_FIXTURE / "test-16x9.pptx", deck, EmitOptions())
    [shape] = [s for s in shapes_of(deck) if s.has_text_frame and s.text_frame.text == "HHxH"]
    assert (shape.top, shape.height) == (text_engine.emu(layout.box.y), text_engine.emu(paragraph.spacing_px))
    spacing = shape.text_frame.paragraphs[0]._p.find(f"{{{A}}}pPr/{{{A}}}lnSpc/{{{A}}}spcPts")
    assert spacing is not None and int(spacing.get("val")) == round(paragraph.spacing_px * config.PT_PER_PX * 100)
    body = _body_properties(shape)
    assert {key: body.get(key) for key in LOCKED_BODY} == LOCKED_BODY and body.get("anchor") == "t", body


@pytest.mark.parametrize("case", ["under-the-switch", "two-lines", "unlocked"])
def test_the_lone_line_rule_leaves_other_text_alone(case):
    """Only a paragraph of one line at or above PowerPoint's single spacing is rewritten: a line at
    1.15 em is already where both renderers agree; a paragraph of more lines keeps the browser's pitch;
    and the A/B arm without line locking lays its paragraphs out itself. (An element of more paragraphs
    is fitted paragraph by paragraph since r2: `test_every_one_line_paragraph_is_fitted`.)"""
    element = {
        "under-the-switch": _lone_element("Arial", 400, 20.0, 23.0),
        "two-lines": _lone_element("Arial", 700, 20.0, 60.0, lines=2),
        "unlocked": _lone_element("Arial", 700, 20.0, 60.0),
    }[case]
    line_lock = case != "unlocked"
    layout = _plan_with(element, line_lock=line_lock)
    before = _plan_with(element, rule=False, line_lock=line_lock)
    assert [p.spacing_px for p in layout.paragraphs] == [p.spacing_px for p in before.paragraphs]
    assert layout.box == before.box


@pytest.mark.parametrize("family, weight", [("Arial", 400), ("Arial", 700), ("Calibri", 400), ("Calibri Light", 400),
                                            ("Segoe UI", 400), ("Segoe UI", 600), ("Times New Roman", 400)])
def test_a_lone_line_is_written_below_powerpoints_single_spacing(family, weight):
    """Every face, size and tall line-height: the written spacing is under 1.2 em (where PowerPoint's
    placement switches), at most the glyphs' height and `LONE_LINE_MAX_EM` (Segoe UI's own height is
    1.33 em), less than a pixel short of that, and the CSS spacing less a whole number of pixels."""
    for size in (9.0, 11.0, 13.0, 16.0, 20.0, 28.0, 40.0, 64.0):
        for ratio in (1.2, 1.25, 1.35, 1.5, 2.0, 3.0, 4.5):
            element = _lone_element(family, weight, size, ratio * size)
            [paragraph] = _plan_with(element).paragraphs
            line = paragraph.lines[0]
            target = min(line.ascent + line.descent, text_engine.LONE_LINE_MAX_EM * size)
            written = paragraph.spacing_px
            assert target - 1.0 < written <= target + 1e-9, (family, size, ratio, written, target)
            assert written < text_engine.POWERPOINT_SINGLE_EM * size
            shortened = ratio * size - written
            assert shortened >= 1.0 and shortened == pytest.approx(round(shortened), abs=1e-9)


#: PowerPoint's ink rows (first, last) for `lone_line_evidence.LONE_CASES` as the emitter writes them
#: now — measured 2026-09-25 (r1c) with `python engine/fixtures/fidelity/lone_line_evidence.py cases
#: --powerpoint`, PowerPoint through `render_pptx.ps1`. PptxRender drew every case on the same rows with
#: `text._fit_lone_line` and without it (arial-op 52–66, calibri-chevron 272–281, times 486–502, …);
#: PowerPoint, without it, 1 to 12 px higher (arial-op 42–56, arial-op-middle 47–61, arial-plus 155–166
#: against 167–178, calibri-chevron 262–271, times 480–496, segoe-counter 366–375, calibri-body 254–263),
#: and with it 0 to 2 px lower — the within-a-pixel agreement below 1.2 em.
LONE_LINES_POWERPOINT = {
    "arial-op": (53, 67), "arial-op-middle": (53, 67), "arial-small": (152, 159), "arial-plus": (168, 179),
    "calibri-chevron": (273, 282), "calibri-body": (256, 265), "calibri-light": (373, 391),
    "segoe-counter": (370, 379), "segoe-large": (480, 504), "times": (487, 503), "times-12": (586, 597),
}


@pytest.mark.renderer_full
def test_a_lone_tall_line_lands_where_powerpoint_draws_it(tmp_path):
    """The gate's renderer draws a lone line on exactly the pixels it drew at the CSS spacing (the frame
    moves by whole pixels), and those pixels are within 2 px of PowerPoint's (recorded above)."""
    from app.engine.fixtures.fidelity import lone_line_evidence as evidence

    rows = {}
    for label, rule in (("on", True), ("off", False)):
        deck = evidence.emit_cases(tmp_path / f"{label}.pptx", rule=rule)
        [png] = renderer.render(deck, 1280, tmp_path / label)
        rows[label] = evidence.ink_rows(png)
        if rule:
            # The file PowerPoint was measured on: every spacing under its single spacing (1.2 × sz).
            for shape in shapes_of(deck):
                if shape.has_text_frame and shape.text_frame.text.strip():
                    paragraph = shape.text_frame.paragraphs[0]
                    spacing = int(paragraph._p.find(f"{{{A}}}pPr/{{{A}}}lnSpc/{{{A}}}spcPts").get("val"))
                    size = int(paragraph.runs[0]._r.get_or_add_rPr().get("sz"))
                    assert spacing < text_engine.POWERPOINT_SINGLE_EM * size, (shape.name, spacing, size)
    assert rows["on"] == rows["off"], "PptxRender must draw a lone line where it drew it before"
    assert set(rows["on"]) == set(LONE_LINES_POWERPOINT)
    for name, (top, bottom) in LONE_LINES_POWERPOINT.items():
        got = rows["on"][name]
        assert got is not None and abs(got[0] - top) <= 2 and abs(got[1] - bottom) <= 2, (name, got, (top, bottom))


# ------------------------------------------- the gap under a change of leading (r2, plan §16 #29)


#: `(family, weight, size px, line-height px, lines, margin before px)` per paragraph: stacks with a change
#: of leading between paragraphs — a client deck's KPI (a 34 px number at 36 px over an 11.5 px label at 14 px
#: over a 9.5 px note), the torture `text` family's head (a 30 px Calibri Light h1 at 1.15 over 15 px body
#: at 1.35), a Segoe UI and a Times New Roman heading over body text.
LEADING_STACKS = {
    "kpi-arial": [("Arial", 700, 34.0, 36.0, 1, 0.0), ("Arial", 700, 11.5, 14.0, 1, 6.0),
                  ("Arial", 400, 9.5, 12.0, 2, 4.0)],
    "head-calibri": [("Calibri Light", 400, 30.0, 34.5, 1, 0.0), ("Calibri", 400, 15.0, 20.25, 2, 6.0)],
    "segoe": [("Segoe UI", 600, 24.0, 28.0, 1, 0.0), ("Segoe UI", 400, 13.0, 18.2, 3, 8.0)],
    "times": [("Times New Roman", 700, 28.0, 30.0, 1, 0.0), ("Times New Roman", 400, 12.0, 16.0, 2, 5.0)],
}

#: What Edge drew for those runs: `(ascent above the baseline, descent below it)` of the text's content
#: area, px — `round(OS/2 win metric × size)`, measured 2026-09-25 (r2) with Range rects in the measuring
#: browser and re-measured on every run by `test_the_browser_rounds_ascent_and_descent_to_whole_pixels`.
EDGE_CONTENT_AREAS = {
    ("Arial", 700, 34.0): (31, 7), ("Arial", 700, 11.5): (10, 2), ("Arial", 400, 9.5): (9, 2),
    ("Arial", 400, 11.0): (10, 2), ("Arial", 700, 9.0): (8, 2), ("Arial", 700, 13.0): (12, 3),
    ("Calibri Light", 400, 30.0): (29, 8), ("Calibri", 400, 15.0): (14, 4),
    ("Calibri", 400, 13.0): (12, 3), ("Segoe UI", 600, 24.0): (26, 6), ("Segoe UI", 400, 13.0): (14, 3),
    ("Times New Roman", 700, 28.0): (25, 6), ("Times New Roman", 400, 12.0): (11, 3),
}


def _stacked_element(blocks, *, x: float = 300.0, top: float = 100.3, name: str = "stack") -> Element:
    """`blocks` laid out the way the browser stacks them: each block's line boxes at its line-height,
    its margin before it (negative: the flex column overlapped the boxes)."""
    y, paragraphs = top, []
    for family, weight, size, line_height, lines, margin in blocks:
        y += margin
        width = text_engine.text_width_px("HHxH", family, size, weight, False)
        paragraphs.append({"align": "left", "lineHeightPx": line_height, "spaceBeforePx": 0, "spaceAfterPx": 0,
                           "bullet": None,
                           "lines": [{"box": {"x": x, "y": y + k * line_height, "w": width, "h": line_height},
                                      "runs": [_run("HHxH", family, weight=weight, size=size)]}
                                     for k in range(lines)]})
        y += lines * line_height
    return Element(kind="text", box=Box(x, top, 200.0, y - top), name=name, anchor="top",
                   writingMode="horizontal", wrap=True, paragraphs=paragraphs)


def _browser_baselines(element: Element) -> list[float]:
    """Every line's baseline where Edge drew it: its line box's centre plus half of (ascent − descent)
    of the content area Edge laid it out on (`EDGE_CONTENT_AREAS`)."""
    out = []
    for paragraph in element.paragraphs:
        for line in paragraph["lines"]:
            run = line["runs"][0]
            ascent, descent = EDGE_CONTENT_AREAS[(run["font"], run["weight"], run["sizePx"])]
            out.append(line["box"]["y"] + line["box"]["h"] / 2.0 + (ascent - descent) / 2.0)
    return out


def _renderer_baselines(layout: text_engine.TextLayout) -> list[float]:
    """Every line's baseline as PptxRender sets an exact-spaced frame: the frame top, each paragraph's
    space before it (`paragraph_gap`, none before the first), and a baseline `drop` above its line's
    bottom (`text._baseline_metrics`)."""
    y, out = layout.box.y, []
    for index, paragraph in enumerate(layout.paragraphs):
        y += text_engine.paragraph_gap(layout.paragraphs, index)
        for line in paragraph.lines:
            y += paragraph.spacing_px
            out.append(y - text_engine._baseline_metrics(line)[1])
    return out


@contextlib.contextmanager
def _old_gap():
    """`paragraph_gap` as it was before r2: `Δc − (L₁ + L₂)/2`, blind to a change of leading."""
    original = text_engine.paragraph_gap
    text_engine.paragraph_gap = lambda planned, index: 0.0 if index == 0 else max(
        0.0, planned[index].lines[0].centre - planned[index - 1].lines[-1].centre
        - (planned[index - 1].spacing_px + planned[index].spacing_px) / 2.0)
    try:
        yield
    finally:
        text_engine.paragraph_gap = original


@pytest.mark.parametrize("stack", sorted(LEADING_STACKS))
def test_the_gap_keeps_every_baseline_where_the_browser_drew_it(stack):
    """Under a change of face, size or leading between paragraphs, every line after the first sits the
    browser's distance below it — the first line's own placement residual (`_first_line_dy`) is common to
    the frame. The old gap, `Δc − (L₁ + L₂)/2`, put a client deck's KPI label 2.6 px low in both renderers."""
    require_faces(*sorted({block[0] for block in LEADING_STACKS[stack]}))
    element = _stacked_element(LEADING_STACKS[stack])
    layout = _plan_with(element)
    browser, renderer_rows = _browser_baselines(element), _renderer_baselines(layout)
    for k in range(1, len(browser)):
        assert renderer_rows[k] - renderer_rows[0] == pytest.approx(browser[k] - browser[0], abs=1e-6), (stack, k)
    with _old_gap():
        old = _renderer_baselines(_plan_with(element))
    assert max(old[k] - old[0] - (browser[k] - browser[0]) for k in range(1, len(browser))) > 1.5, (stack, old)


def test_the_gap_is_the_old_rule_without_a_change_of_leading():
    """Two paragraphs of one face, size and leading: the gap is still `Δc − L`, to the micron."""
    element = _stacked_element([("Calibri", 400, 15.0, 17.25, 2, 0.0), ("Calibri", 400, 15.0, 17.25, 2, 9.0)])
    planned = _plan_with(element).paragraphs
    measured = planned[1].lines[0].centre - planned[0].lines[-1].centre
    assert [p.spacing_px for p in planned] == [17.25, 17.25]
    assert text_engine.paragraph_gap(planned, 1) == pytest.approx(measured - 17.25, abs=1e-9) == pytest.approx(9.0)


def test_the_gap_falls_back_to_the_model_for_a_face_that_is_not_installed():
    """No installed face, no DirectWrite metrics: the gap is r1c's model on the fallback face's metrics —
    `Δc − L₂ + (N₂ − N₁)/2 + dy₁ − dy₂` — and a change of leading still moves the label up, not down."""
    element = _stacked_element([("No Such Face 7", 700, 34.0, 36.0, 1, 0.0),
                                ("No Such Face 7", 700, 11.5, 14.0, 2, 6.0)])
    planned = _plan_with(element).paragraphs
    above, below = planned[0].lines[-1], planned[1].lines[0]
    assert text_engine._baseline_metrics(above) is None
    expected = (below.centre - above.centre - planned[1].spacing_px
                + ((below.ascent + below.descent) - (above.ascent + above.descent)) / 2.0
                + text_engine._first_line_dy(above) - text_engine._first_line_dy(below))
    assert text_engine._exact_gap(planned, 1) == pytest.approx(expected, abs=1e-9)
    with _old_gap():
        old = text_engine.paragraph_gap(planned, 1)
    assert text_engine.paragraph_gap(planned, 1) < old - 1.0


def test_the_browser_rounds_ascent_and_descent_to_whole_pixels():
    """What `text._baseline_metrics` assumes about the measuring browser, measured on every run: Edge lays
    a line out on a content area of exactly `round(ascent × size)` above the baseline and `round(descent
    × size)` below it, with DirectWrite's (OS/2 win) metrics — not hhea's (Calibri: 0.952 / 0.269 against
    0.75 / 0.25), and not unrounded (Arial 11 px: 9.96 / 2.33 → 10 / 2)."""
    from app.engine.extract.html import measuring_page

    # Only the faces installed here: Calibri and Segoe UI ship with Windows/Office only.
    families = available(("Arial", "Calibri", "Calibri Light", "Segoe UI", "Times New Roman"))
    assert {"Arial", "Times New Roman"} <= set(families), "CI installs ttf-mscorefonts-installer"
    cases = [case for case in sorted(EDGE_CONTENT_AREAS) if case[0] in families] + [
        (family, 400, size) for family in families if family != "Calibri Light"
        for size in (7.5, 8.5, 10.5, 13.0, 16.0, 25.0, 34.0)]
    with measuring_page(Canvas(640, 200)) as page:
        page.set_content("<html><body></body></html>")
        measured = page.evaluate(
            """(cases) => cases.map(([family, weight, size]) => {
                const block = document.createElement('div');
                block.style.cssText = `position:absolute;left:10px;top:10.3px;white-space:nowrap;font-family:"${family}";`
                    + `font-weight:${weight};font-size:${size}px;line-height:${size * 1.6}px`;
                block.innerHTML = 'HHxg<span style="display:inline-block;width:0;height:0"></span>';
                document.body.appendChild(block);
                const range = document.createRange();
                range.selectNodeContents(block.firstChild);
                const rect = range.getClientRects()[0];
                const baseline = block.lastChild.getBoundingClientRect().top;
                block.remove();
                return [baseline - rect.top, rect.bottom - baseline];
            })""",
            [list(case) for case in cases],
        )
    problems = []
    for (family, weight, size), (above, below) in zip(cases, measured, strict=False):
        line = text_engine.LineLayout([_run("HHxg", family, weight=weight, size=size)], 0.0, 0.0, 0.0, 0.0, 0.0)
        offset, _ = text_engine._baseline_metrics(line)
        if abs((above - below) / 2.0 - offset) > 1e-3 or above != round(above) or below != round(below):
            problems.append(f"{family} {weight} {size}px: Edge {above:.3f} / {below:.3f}, model offset {offset}")
        if (family, weight, size) in EDGE_CONTENT_AREAS and (above, below) != EDGE_CONTENT_AREAS[(family, weight, size)]:
            problems.append(f"{family} {weight} {size}px: Edge {above} / {below}, recorded "
                            f"{EDGE_CONTENT_AREAS[(family, weight, size)]}")
    assert not problems, "\n".join(problems)


def test_every_one_line_paragraph_is_fitted():
    """Two one-line paragraphs at 3 em (a client deck's operators, stacked): each is written below PowerPoint's single
    spacing, shortened by whole pixels; the first by moving the frame down, the second by the space before
    it growing as much — every baseline, and the frame's bottom, where they were."""
    element = _lone_element("Arial", 700, 20.0, 60.0, paragraphs=2)
    layout, before = _plan_with(element), _plan_with(element, rule=False)
    assert all(p.spacing_px < text_engine.POWERPOINT_SINGLE_EM * 20.0 for p in layout.paragraphs)
    assert all((60.0 - p.spacing_px) == pytest.approx(round(60.0 - p.spacing_px), abs=1e-9) for p in layout.paragraphs)
    assert _renderer_baselines(layout) == pytest.approx(_renderer_baselines(before), abs=1e-9)
    assert layout.box.y + layout.box.h == pytest.approx(before.box.y + before.box.h, abs=1e-9)


def test_a_one_line_paragraph_set_closer_than_its_leading_is_landed():
    """a client slide: a 9 px label at 10 px over a 13 px pillar name at 16 px, the name's line box
    starting half a pixel inside the label's. `a:spcBef` cannot be negative, so at its own spacing the
    name sat a pixel low; shortened by whole pixels, the space before it lands it."""
    element = _stacked_element([("Arial", 700, 9.0, 10.0, 1, 0.0), ("Arial", 700, 13.0, 16.0, 1, -0.5)])
    layout, before = _plan_with(element), _plan_with(element, rule=False)
    assert text_engine._exact_gap(before.paragraphs, 1) < -0.5          # the clamp that left it low
    assert text_engine.paragraph_gap(before.paragraphs, 1) == 0.0
    assert layout.paragraphs[1].spacing_px == 14.0 and text_engine.paragraph_gap(layout.paragraphs, 1) > 0.0
    browser, rows = _browser_baselines(element), _renderer_baselines(layout)
    assert rows[1] - rows[0] == pytest.approx(browser[1] - browser[0], abs=1e-6)


@pytest.mark.renderer_full
def test_a_change_of_leading_lands_where_the_browser_drew_it_in_pptxrender(tmp_path):
    """End to end through the gate's renderer: every line of the four `LEADING_STACKS` has its last ink
    row (the baseline, for `HHxH`) within a pixel of the browser's; with the old gap every later
    paragraph came out 2–3 px low."""
    elements = [_stacked_element(blocks, x=100.0 + 300.0 * index, name=name)
                for index, (name, blocks) in enumerate(sorted(LEADING_STACKS.items()))]
    assign_ids_and_z(elements)
    ir = IR(canvas=Canvas(1280, 720), slide=Slide(id="leading", layoutId="layout-07"), elements=elements)
    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")

    def last_ink_rows(deck: Path, out: Path) -> dict[str, list[int]]:
        [png] = renderer.render(deck, 1280, out)
        with Image.open(png) as image:
            grey = image.convert("L")
            pixels = grey.load()
            rows = {}
            for element in elements:
                x0 = int(element.box.x) - 2
                ink = [y for y in range(60, 260) if any(pixels[x, y] < 140 for x in range(x0, x0 + 120))]
                bands = [y for k, y in enumerate(ink) if k + 1 == len(ink) or ink[k + 1] != y + 1]
                rows[element.name] = bands
            return rows

    deck = tmp_path / "leading.pptx"
    emitter.emit([ir], manifest, config.MASTERS_FIXTURE / "test-16x9.pptx", deck, EmitOptions())
    rendered = last_ink_rows(deck, tmp_path / "now")
    old_deck = tmp_path / "old.pptx"
    with _old_gap():
        emitter.emit([ir], manifest, config.MASTERS_FIXTURE / "test-16x9.pptx", old_deck, EmitOptions())
    old = last_ink_rows(old_deck, tmp_path / "old")
    for element in elements:
        expected = [math.floor(b + 0.5) - 1 for b in _browser_baselines(element)]
        got = rendered[element.name]
        assert len(got) == len(expected) and all(abs(g - e) <= 1 for g, e in zip(got, expected, strict=False)), \
            (element.name, got, expected)
        assert all(o - e >= 2 for o, e in zip(old[element.name][1:], expected[1:], strict=False)), \
            (element.name, old[element.name], expected)


@pytest.mark.renderer_full
def test_pptxrender_sets_an_exact_line_descent_plus_line_gap_above_its_bottom(tmp_path):
    """The renderer half of `_exact_gap`: PptxRender puts an exact-spaced line's baseline `drop` above the
    line's bottom — DirectWrite's descent plus line gap (`text._baseline_metrics`), not hhea's (Calibri's
    hhea descent + lineGap is 0.47 em against 0.27, 4 px at 20 px). One `HHHH` line per frame, frames at
    whole pixels with whole-pixel heights (so PptxRender's rect snapping is a no-op), four faces, 9–40 px:
    the H's baseline row is the model rounded, every one within half a pixel (measured 2026-09-25)."""
    from app.engine.fixtures.fidelity.lone_line_evidence import _blank_deck, _textbox

    faces = (("Arial", "Arial"), ("Calibri", "Calibri"), ("Segoe UI", "Segoe UI"), ("Times New Roman", "Times New Roman"))
    sizes = (9.0, 10.0, 11.0, 12.0, 13.0, 14.0, 16.0, 18.0, 20.0, 24.0, 30.0, 40.0)
    presentation = _blank_deck()
    cases = []
    for number, (typeface, family) in enumerate(faces):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for k, size in enumerate(sizes):
            spacing = round(size * 1.1)                           # whole px: 1.1 em, below single spacing
            x, y = 20 + (k % 6) * 200, 40 + (k // 6) * 200
            _textbox(slide, x, y, 190, spacing, typeface, False, size, ("spcPts", int(round(spacing * 75))),
                     text="HHHH")
            line = text_engine.LineLayout([_run("HHHH", family, size=size)], 0.0, 0.0, 0.0, 0.0, 0.0)
            cases.append((number, typeface, size, x, y, y + spacing - text_engine._baseline_metrics(line)[1]))
    deck = tmp_path / "drop.pptx"
    presentation.save(str(deck))
    pngs = renderer.render(deck, 1280, tmp_path / "render")
    problems = []
    for number, typeface, size, x, y, model in cases:
        with Image.open(pngs[number]) as image:
            grey = image.convert("L")
            pixels = grey.load()
            ink = [row for row in range(int(y) - 10, int(model) + 12) if any(pixels[col, row] < 128
                                                                             for col in range(int(x), int(x) + 150))]
        baseline = ink[-1] + 1 if ink else None
        if baseline is None or abs(baseline - model) > 0.6:
            problems.append(f"{typeface} {size}px: baseline row {baseline}, model {model:.3f}")
    assert not problems, "\n".join(problems)


# ------------------------------------------------------------- a locked line never wraps (WP-D)


def _text_shapes(pptx: Path):
    """Every shape on the first slide holding text, group children included."""
    for shape in shapes_of(pptx):
        for candidate in (list(shape.shapes) if hasattr(shape, "shapes") else [shape]):
            if candidate.has_text_frame and candidate.text_frame.text.strip():
                yield candidate


def _body_properties(shape) -> dict[str, str]:
    """The slide part's own `a:bodyPr` attributes — not python-pptx's view inherited from the layout."""
    body = shape._element.find(f".//{{{A}}}bodyPr")
    assert body is not None, f"{shape.name} has no bodyPr"
    return {**dict(body.attrib), "noAutofit": str(body.find(f"{{{A}}}noAutofit") is not None)}


LOCKED_BODY = {"wrap": "none", "lIns": "0", "tIns": "0", "rIns": "0", "bIns": "0", "noAutofit": "True"}


def test_line_locked_frames_never_wrap(emitted):
    """Under line locking every frame is `wrap="none"`, and every browser break is still an `a:br`.

    A wrap PowerPoint adds can only break a locked line a second time (fidelity F4): a wider metric
    must overhang instead (`text._wraps`).
    """
    checked = 0
    for name, (out, _, ir, _, _) in emitted.items():
        elements = {element.name: element for element in ir.elements if element.kind == "text"}
        for shape in _text_shapes(out):
            body = _body_properties(shape)
            assert {key: body.get(key) for key in LOCKED_BODY} == LOCKED_BODY, f"{name}/{shape.name}"
            element = elements[shape.name]
            lines = [len(paragraph.get("lines") or []) for paragraph in element.paragraphs]
            breaks = [len(p._p.findall(f"{{{A}}}br")) for p in shape.text_frame.paragraphs]
            assert breaks == [count - 1 for count in lines], f"{name}/{shape.name}"
            checked += 1
    assert checked == sum(1 for _, _, ir, _, _ in emitted.values()
                          for element in ir.elements if element.kind == "text")


def test_a_filled_placeholder_writes_its_own_body_properties(emitted):
    """A title in the layout's placeholder overrides `titleStyle`'s wrap, insets, anchor and autofit.

    `prepare_placeholder_frame` and `write_layout` work on the same frame; the wrap decision is
    `write_layout`'s alone, and it lands on the slide part itself, not by inheritance.
    """
    for fixture, name in (("emit-placeholders", "title"), ("emit-sizes-4x3", "title-4x3")):
        shape = by_name(emitted[fixture][0], name)
        assert shape.is_placeholder, f"{fixture}/{name} is not in the layout's placeholder"
        body = _body_properties(shape)
        assert {key: body.get(key) for key in LOCKED_BODY} == LOCKED_BODY, f"{fixture}/{name}: {body}"
        assert body.get("anchor") == "t"
        assert shape._element.find(f".//{{{A}}}bodyPr/{{{A}}}normAutofit") is None


def _wrap_before_wp_d(pptx: Path, out: Path) -> Path:
    """The same deck with the emitter's rule before WP-D (bad2f9b): every frame of 2+ lines wraps."""
    presentation = Presentation(str(pptx))
    for shape in presentation.slides[0].shapes:
        if shape.has_text_frame:
            lines = sum(1 + len(p._p.findall(f"{{{A}}}br")) for p in shape.text_frame.paragraphs)
            if lines > 1:
                shape.text_frame.word_wrap = True
    presentation.save(str(out))
    return out


#: `emit-wrap`'s ordinary elements, as crops around their boxes: they must render the same whether
#: their frame wraps or not.
WRAP_ORDINARY = {"body-column": (54, 190, 380, 330), "centred-card": (690, 190, 990, 250),
                 "chip": (990, 390, 1050, 425)}


@pytest.mark.renderer_full
def test_a_locked_line_survives_wider_metrics(tmp_path, monkeypatch):
    """Benchmark slide 51 (fidelity F4), reproduced: the export cannot see Arial, PptxRender can.

    Blind, the emitter keeps the browser's 948 px for line 1 and boxes the title from x 24 to the
    canvas edge (1256 px); PptxRender draws Arial Bold 31 px at 1316 px. With `wrap="square"` (the
    rule before WP-D) line 1 breaks again at "versus" and "balance-" gets a line of its own — three
    lines. With `wrap="none"` it keeps the browser's two and overhangs instead. PowerPoint does the
    same (`docs/engine/fidelity/evidence/wp-d-slide51-mechanism.png`, `wrap_evidence.py emit-wrap`).
    """
    with _blind_fonts(monkeypatch, tmp_path):
        out, report, ir, _, _ = emit_fixture("emit-wrap", tmp_path / "fix")
    note = ": font 'Arial' is not installed — measured with a fallback"
    assert sorted(report.warnings) == sorted(f"emit-wrap/{e.id}{note}" for e in ir.elements), \
        report.warnings

    control = _wrap_before_wp_d(out, tmp_path / "control.pptx")
    fixed = renderer.render(out, 1280, tmp_path / "png-fix")[0]
    before = renderer.render(control, 1280, tmp_path / "png-control")[0]
    title = _ink_bands(fixed, 0, 170)
    assert len(title) == 2, f"the locked title was broken again: {title}"
    assert len(_ink_bands(before, 0, 170)) >= 3, "the control no longer re-wraps: the fixture proves nothing"

    image = Image.open(fixed).convert("L")
    pixels = image.load()
    line_two = [x for x in range(image.width)
                if any(pixels[x, y] < 250 for y in range(title[1][0], title[1][1] + 1))]
    assert abs(line_two[0] - 24) <= 3, f"'sheet banks' starts at x {line_two[0]}, not at the box's 24"

    # The ordinary elements render the same either way. Pixel-identical in both renderers on the
    # evidence run; one PptxRender run once differed by a single antialiasing pixel (16 levels) in
    # the centred card, so the ink geometry is compared exactly and the pixels nearly so.
    fixed_rgb, before_rgb = Image.open(fixed).convert("RGB"), Image.open(before).convert("RGB")
    for name, (x0, y0, x1, y1) in WRAP_ORDINARY.items():
        assert _ink_bands(fixed, y0, y1, x0, x1) == _ink_bands(before, y0, y1, x0, x1), name
        red, green, blue = ImageChops.difference(fixed_rgb.crop((x0, y0, x1, y1)),
                                                 before_rgb.crop((x0, y0, x1, y1))).split()
        changed = sum(ImageChops.lighter(ImageChops.lighter(red, green), blue).histogram()[1:])
        assert changed <= 0.001 * (x1 - x0) * (y1 - y0), f"{name}: {changed} pixels differ"


@pytest.mark.validator
def test_emit_wrap_fixture_is_valid_deterministic_and_validates(tmp_path, monkeypatch):
    """`emit-wrap` under blind fonts: reproducible, valid OOXML, and nothing the renderer objects to."""
    with _blind_fonts(monkeypatch, tmp_path):
        first, _, _, master, manifest = emit_fixture("emit-wrap", tmp_path / "a")
        second = emit_fixture("emit-wrap", tmp_path / "b")[0]
    assert _sha256(first) == _sha256(second), "emit-wrap is not reproducible"

    assert config.VALIDATE_PY.exists(), f"validator missing at {config.VALIDATE_PY}"
    completed = subprocess.run(
        [sys.executable, str(config.VALIDATE_PY), str(first), "--original", str(master)],
        capture_output=True, text=True, timeout=300,
    )
    assert completed.returncode == 0, f"{completed.stdout}\n{completed.stderr}"

    baseline = tmp_path / "baseline.pptx"
    emitter.emit([IR(canvas=Canvas(1280, 720), slide=Slide(id="empty", layoutId="layout-06"))],
                 manifest, master, baseline, EmitOptions())
    renderer.render(baseline, 1280, tmp_path / "baseline-png")
    expected = set(json.loads((tmp_path / "baseline-png" / "warnings.json")
                              .read_text(encoding="utf-8"))["warnings"])
    renderer.render(first, 1280, tmp_path / "png")
    warnings = json.loads((tmp_path / "png" / "warnings.json").read_text(encoding="utf-8"))
    assert set(warnings["warnings"]) - expected == set(), warnings["warnings"]
    assert not warnings["truncated"]
    assert warnings["fontsSubstituted"] == []


# --------------------------------------------------------------------------------- placeholders


def test_explicit_placeholder_wins_and_overlap_fills_the_rest():
    ir = load_ir("emit-placeholders")
    _, manifest = master_of(ir)
    mapped = map_placeholders(ir, manifest.layout("layout-02"))
    by_label = {e.name: e.placeholder for e in mapped.elements if e.kind == "text"}
    assert by_label["title"] == {"type": "title", "idx": 0}
    assert by_label["body-by-overlap"] == {"type": "obj", "idx": 1}
    assert by_label["free-caption"] is None


def test_placeholder_text_lands_in_the_placeholder(emitted):
    out, report, _, _, _ = emitted["emit-placeholders"]
    slide = Presentation(str(out)).slides[0]
    placeholders = list(slide.placeholders)
    assert {p.placeholder_format.idx for p in placeholders} == {0, 1}
    assert "Explicitly the title" in placeholders[0].text_frame.text
    assert report.placeholders_used == [
        {"slide": 1, "type": "title", "idx": 0}, {"slide": 1, "type": "obj", "idx": 1}]


def test_unused_placeholders_are_deleted(emitted):
    """Nothing may be left with an empty text frame — the gate counts those."""
    for name, (out, _, _, _, _) in emitted.items():
        for shape in Presentation(str(out)).slides[0].placeholders:
            assert shape.text_frame.text.strip(), f"{name}: empty {shape.name}"


def test_a_filled_placeholder_is_moved_into_paint_order(emitted):
    """python-pptx clones placeholders first; the title must still sit above the rule behind it."""
    out, _, _, _, _ = emitted["emit-placeholders"]
    names = [s.name for s in shapes_of(out)]
    assert names.index("rule") < names.index("title")


def test_legacy_manifest_maps_by_type_when_idx_is_missing(tmp_path):
    """A legacy import records no `idx`; the type still identifies a unique placeholder."""
    ir = load_ir("emit-placeholders")
    master, manifest = master_of(ir)
    for layout in manifest.layouts:
        for placeholder in layout.placeholders:
            placeholder.idx = None
    mapped = map_placeholders(ir, manifest.layout("layout-02"))
    title = next(e for e in mapped.elements if e.name == "title")
    assert title.placeholder == {"type": "title", "idx": None}

    out = tmp_path / "legacy.pptx"
    report = emitter.emit([mapped], manifest, master, out, EmitOptions())
    assert report.warnings == []
    assert report.placeholders_used[0]["type"] == "title"
    assert "Explicitly the title" in Presentation(str(out)).slides[0].placeholders[0].text_frame.text


def test_ambiguous_legacy_placeholder_is_an_error_not_a_guess():
    """Two `obj` zones and no `idx`: the answer is unknown, so it is reported, not invented."""
    ir = load_ir("emit-placeholders")
    _, manifest = master_of(ir)
    layout = manifest.layout("layout-04")        # Two Content: two `obj` placeholders
    for placeholder in layout.placeholders:
        placeholder.idx = None
    # Sit the body text squarely inside the left content zone so it maps by overlap.
    left = next(p for p in layout.placeholders if p.type == "obj")
    ir.elements = [e.replace(box=Box(left.x + 10, left.y + 10, left.w - 20, left.h - 20))
                   if e.name == "body-by-overlap" else e for e in ir.elements]

    mapped = map_placeholders(ir, layout)
    assert any(d.level == "error" and "idx" in d.message for d in mapped.diagnostics)
    assert all((e.placeholder or {}).get("type") != "obj"
               for e in mapped.elements if e.kind == "text")


# --------------------------------------------------------------------------------------- images


def test_images_are_cropped_not_squashed(emitted):
    out, _, _, _, _ = emitted["emit-images"]
    cover = by_name(out, "photo-cover")
    # The synthetic 400x300 photo into a 300x200 box: the image is relatively taller, so the top
    # and bottom go.
    assert cover.crop_top > 0 and cover.crop_top == pytest.approx(cover.crop_bottom)
    assert cover.crop_left == 0 and cover.width == 300 * config.EMU_PER_PX

    contain = by_name(out, "photo-contain")
    assert contain.crop_top == 0
    # The image is relatively taller than the box, so `contain` pillarboxes it: the height fills
    # and the width shrinks to the image's own aspect ratio. Never stretched.
    assert contain.height == 200 * config.EMU_PER_PX
    assert contain.width < 300 * config.EMU_PER_PX
    assert contain.width / contain.height == pytest.approx(400 / 300, rel=1e-3)


def test_image_geometry_crop_clip_and_alpha(emitted):
    out, _, _, _, _ = emitted["emit-images"]
    assert by_name(out, "logo-circle")._element.find(f".//{{{A}}}prstGeom").get("prst") == "ellipse"
    rounded = by_name(out, "logo-rounded")._element
    assert rounded.find(f".//{{{A}}}prstGeom").get("prst") == "roundRect"
    assert rounded.find(f".//{{{A}}}alphaModFix").get("amt") == "60000"

    cropped = by_name(out, "logo-cropped")
    assert cropped.crop_left == pytest.approx(0.1) and cropped.crop_right == pytest.approx(0.1)

    clipped = by_name(out, "logo-clipped")
    assert clipped.width == 160 * config.EMU_PER_PX
    assert clipped.crop_right == pytest.approx(1 - 160 / 314, abs=1e-3)


def test_rasters_are_reported_with_their_reason(emitted):
    out, report, _, _, _ = emitted["emit-images"]
    assert report.rasters == [{"id": "e7", "reason": "svg filter: feGaussianBlur"}]
    assert report.shapes_by_kind["image"] == 6 and report.shapes_by_kind["raster"] == 1


# --------------------------------------------------------------------------------------- tables


def test_table_grid_merges_and_cell_text(emitted):
    out, _, _, _, _ = emitted["emit-tables"]
    table = by_name(out, "summary").table
    assert len(table.rows) == 3 and len(table.columns) == 3
    assert table.cell(1, 0).is_merge_origin and table.cell(1, 0).span_width == 2
    assert table.cell(0, 0).text == "Market"
    assert table.cell(1, 0).text == "Internal spans two columns"
    assert table.columns[0].width == 320 * config.EMU_PER_PX
    header = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    assert header._r.get_or_add_rPr().get("b") == "1"
    borders = table.cell(0, 0)._tc.find(f"{{{A}}}tcPr").find(f"{{{A}}}lnB")
    assert borders is not None


# ------------------------------------------------------------------ tables-rich (WP-A, §6, §8.5)


def _rich_cell(emitted, r: int, c: int):
    out, _, _, _, _ = emitted["emit-tables-rich"]
    return by_name(out, "rich-cells").table.cell(r, c)


def _rich_table() -> Element:
    return next(e for e in load_ir("emit-tables-rich").elements if e.kind == "table")


def _rich_ir_cell(r: int, c: int) -> dict:
    return next(cell for cell in _rich_table().cells if (cell["r"], cell["c"]) == (r, c))


def test_overlay_only_cell_needs_one_point(emitted):
    """A cell whose content is all drawn over the table holds one 1 pt empty paragraph, no run."""
    body = _rich_cell(emitted, 0, 0).text_frame._txBody
    paragraphs = body.findall(f"{{{A}}}p")
    assert len(paragraphs) == 1
    assert paragraphs[0].findall(f"{{{A}}}r") == []
    assert paragraphs[0].find(f"{{{A}}}endParaRPr").get("sz") == "100"
    assert paragraphs[0].find(f"{{{A}}}pPr/{{{A}}}lnSpc/{{{A}}}spcPts").get("val") == "100"


def test_cell_first_line_indent(emitted):
    """The measured leading offset (an icon before the first word) is a positive first-line indent."""
    pPr = _rich_cell(emitted, 0, 1).text_frame.paragraphs[0]._p.find(f"{{{A}}}pPr")
    assert pPr.get("marL") == "0"
    assert int(pPr.get("indent")) == round(22.5 * config.EMU_PER_PX)


def test_cell_bullets_and_hanging_indent(emitted):
    paragraphs = _rich_cell(emitted, 0, 2).text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Sandbox", "Talent visas"]
    for paragraph in paragraphs:
        pPr = paragraph._p.find(f"{{{A}}}pPr")
        assert pPr.find(f"{{{A}}}buChar").get("char") == "•"
        assert int(pPr.get("marL")) == 18 * config.EMU_PER_PX == -int(pPr.get("indent"))
        assert pPr.find(f"{{{A}}}buClr/{{{A}}}srgbClr").get("val") == "2E2E38"


def test_cell_measured_margin_and_indent(emitted):
    """`x-wpa` `marginPx`/`indentPx` are the paragraph's `marL`/`indent` (the kinds review's findings 2
    and 4): text in a grid column or under a margin starts at `marL` on every line; a bullet whose
    text starts after a chip keeps its text at `marL` and its bullet at the list edge (`indent` =
    −`marL`). No `marR` but package B's `_set_indent` 0 is written: PptxRender does not read it."""
    column = _rich_cell(emitted, 2, 1).text_frame.paragraphs[0]._p.find(f"{{{A}}}pPr")
    assert int(column.get("marL")) == round(60 * config.EMU_PER_PX) and column.get("indent") == "0"
    chip_first = _rich_cell(emitted, 2, 3).text_frame.paragraphs[0]._p.find(f"{{{A}}}pPr")
    assert chip_first.find(f"{{{A}}}buChar").get("char") == "•"
    assert int(chip_first.get("marL")) == round(40 * config.EMU_PER_PX) == -int(chip_first.get("indent"))
    out, _, _, _, _ = emitted["emit-tables-rich"]
    table = by_name(out, "rich-cells").table
    assert not [p for row in table.rows for cell in row.cells for p in cell.text_frame.paragraphs
                if p._p.find(f"{{{A}}}pPr") is not None and p._p.find(f"{{{A}}}pPr").get("marR") not in (None, "0")]


def test_cell_dashed_border(emitted):
    """A dashed rule stays dashed: `prstDash` inside `lnB`, and the lines still precede the fill."""
    tcPr = _rich_cell(emitted, 1, 1)._tc.find(f"{{{A}}}tcPr")
    lnB = tcPr.find(f"{{{A}}}lnB")
    assert lnB.find(f"{{{A}}}prstDash").get("val") == "dash"
    assert [child.tag.rpartition("}")[2] for child in lnB] == ["solidFill", "prstDash"]
    tags = [child.tag.rpartition("}")[2] for child in tcPr]
    fills = [i for i, tag in enumerate(tags) if tag in ("noFill", "solidFill")]
    lines = [i for i, tag in enumerate(tags) if tag.startswith("ln")]
    assert lines and fills and max(lines) < min(fills), tags


def test_cell_width_lock_tracking(emitted):
    """The line measured 6 px narrower than predicted is tracked in; the other cells are not."""
    runs = _rich_cell(emitted, 0, 3).text_frame.paragraphs[0].runs
    assert int(runs[0]._r.get_or_add_rPr().get("spc")) < 0
    for r, c in ((0, 1), (1, 1), (2, 1)):
        for run in _rich_cell(emitted, r, c).text_frame.paragraphs[0].runs:
            # R3: `spc` is always written, "0" when there is no tracking (a style must not show through).
            assert run._r.get_or_add_rPr().get("spc") == "0", (r, c)


def test_cell_margins_shifted_by_the_vertical_model(emitted):
    """`marT + marB` is the browser's inset (padding + slot); the top gives way to Δ = (L − N)/2 + dy."""
    wpa = _rich_table().extras["x-wpa"]["cells"]
    for r, c in ((0, 1), (1, 3), (2, 0)):
        cell, native, slot = _rich_ir_cell(r, c), _rich_cell(emitted, r, c), wpa[f"{r}:{c}"]["slot"]
        pad = cell["paddingPx"]
        inset = (pad["t"] + pad["b"] + slot["t"] + slot["b"]) * config.EMU_PER_PX
        assert abs(native.margin_top + native.margin_bottom - inset) <= 2, (r, c)
        assert native.margin_left == round((pad["l"] + slot["l"]) * config.EMU_PER_PX)
        # A 22 px line of 14.67 px Calibri: L > N, so the text moves up and marT < padT + slot.t.
        assert native.margin_top < (pad["t"] + slot["t"]) * config.EMU_PER_PX, (r, c)
    layout = text_engine.plan_cell(_rich_ir_cell(0, 1), {}, SimpleNamespace(options=EmitOptions()))
    first = layout.paragraphs[0].lines[0]
    expected = (22.0 - (first.ascent + first.descent)) / 2 + text_engine._first_line_dy(first)
    assert text_engine.cell_margin_shift_px(layout) == pytest.approx(expected)
    assert _rich_cell(emitted, 0, 1).margin_top == round((8 - expected) * config.EMU_PER_PX)


def test_cell_margin_shift_targets_the_measured_first_line():
    """A first line the browser drew lower than its CSS line-height says (a taller inline box on
    it) moves the text down by exactly that much, under every anchor."""
    cell = _rich_ir_cell(0, 1)
    ctx = SimpleNamespace(options=EmitOptions())
    layout = text_engine.plan_cell(cell, {}, ctx)
    model = text_engine.cell_margin_shift_px(layout)
    at_rest = text_engine.cell_margin_shift_px(layout, slot_top_px=120.0, slot_height_px=60.0,
                                               top_px=8.0, bottom_px=8.0, anchor="top")
    assert at_rest == pytest.approx(model, abs=1e-6)
    lowered = {**cell, "paragraphs": [{**p, "lines": [{**line, "box": {**line["box"], "y": line["box"]["y"] + 2}}
                                                       for line in p["lines"]]} for p in cell["paragraphs"]]}
    moved = text_engine.cell_margin_shift_px(text_engine.plan_cell(lowered, {}, ctx), slot_top_px=120.0,
                                             slot_height_px=60.0, top_px=8.0, bottom_px=8.0)
    assert moved == pytest.approx(model - 2.0, abs=1e-6)
    middle = text_engine.cell_margin_shift_px(layout, slot_top_px=120.0, slot_height_px=60.0,
                                              top_px=8.0, bottom_px=8.0, anchor="middle")
    assert middle == pytest.approx(model + (60.0 - 16.0 - 22.0) / 2, abs=1e-6)
    outside = text_engine.cell_margin_shift_px(layout, slot_top_px=400.0, slot_height_px=60.0)
    assert outside == pytest.approx(model), "a line outside its slot is no measurement"


def _tall_line_cell(lines: int = 1) -> dict:
    """The fixture's "Singapore" cell as the heat-map idiom: padding 0, `line-height` = the 60 px row."""
    cell = _rich_ir_cell(0, 1)
    paragraph = cell["paragraphs"][0]
    drawn = [{**paragraph["lines"][0], "box": {**paragraph["lines"][0]["box"], "y": 120.0 + 60.0 * k, "h": 60.0}}
             for k in range(lines)]
    return {**cell, "paddingPx": {"t": 0, "r": 0, "b": 0, "l": 0},
            "paragraphs": [{**paragraph, "lineHeightPx": 60.0, "lines": drawn}]}


@pytest.mark.parametrize("anchor", ["top", "middle", "bottom"])
def test_a_one_line_cell_with_a_tall_line_height_is_fitted(anchor):
    """`height:60px; line-height:60px; padding:0`: PowerPoint puts the exact line's leading above the
    glyphs, so the text would need a negative margin to land (an app heat map sat 14 px low), and
    PowerPoint and PptxRender part on where so tall a line's baseline is. The spacing becomes the
    glyphs' height: no residual, the row not grown, and — top-anchored — the model's baseline
    (marT + L − descent + dy) on the browser's (centre + (ascent − descent)/2)."""
    ctx = SimpleNamespace(options=EmitOptions())
    layout = text_engine.plan_cell(_tall_line_cell(), {}, ctx)
    place = {"slot_top_px": 120.0, "slot_height_px": 60.0, "top_px": 0.0, "bottom_px": 0.0, "anchor": anchor}
    assert text_engine.cell_margins_px(layout, **place)[2] > 5.0, "unfitted, the text sits low"
    spacing = text_engine.fit_cell_spacing(layout, **place)
    line = layout.paragraphs[0].lines[0]
    assert spacing == pytest.approx(line.ascent + line.descent) and layout.paragraphs[0].spacing_px == spacing
    top, bottom, residual = text_engine.cell_margins_px(layout, **place)
    assert abs(residual) <= 0.01 and min(top, bottom) >= 0.0
    assert top + spacing + bottom <= 60.0 + 0.01, "the row keeps its height"
    if anchor == "top":
        powerpoint = 120.0 + top + spacing - line.descent + text_engine._first_line_dy(line)
        assert powerpoint == pytest.approx(line.centre + (line.ascent - line.descent) / 2.0, abs=0.05)


def test_a_cell_of_more_lines_keeps_its_spacing():
    """Two 60 px lines: the pitch is the design's, so the spacing stays and the residual is reported."""
    ctx = SimpleNamespace(options=EmitOptions())
    layout = text_engine.plan_cell(_tall_line_cell(lines=2), {}, ctx)
    place = {"slot_top_px": 120.0, "slot_height_px": 120.0, "top_px": 0.0, "bottom_px": 0.0, "anchor": "top"}
    assert text_engine.fit_cell_spacing(layout, **place) == pytest.approx(60.0)
    assert text_engine.cell_margins_px(layout, **place)[2] > 5.0


def test_a_cell_line_flush_with_its_column_is_tracked_in():
    """A cell cannot be widened the way a text box is (`TEXT_BOX_WIDTH_SLACK`), and the browser fills a
    column to its last fraction (a client slide: a 103.609 px line in a 103.609 px column wrapped in
    PptxRender). Such a line is tracked in to leave the text-box allowance; a line with room is not."""
    ctx = SimpleNamespace(options=EmitOptions())
    layout = text_engine.plan_cell(_rich_ir_cell(1, 1), {}, ctx)
    line = layout.paragraphs[0].lines[0]
    width, characters = max(line.predicted, line.width), text_engine._chars(line)
    text_engine._leave_cell_slack(layout, width + 40.0)
    assert line.tracking_px == 0.0, "40 px of room: nothing to do"
    text_engine._leave_cell_slack(layout, width)
    allowance = max(text_engine.CELL_LINE_SLACK_PX, width * (1.0 - 1.0 / config.TEXT_BOX_WIDTH_SLACK))
    assert line.tracking_px < 0.0
    assert width + line.tracking_px * characters == pytest.approx(width - allowance, abs=1e-6)


def test_cell_margins_keep_the_side_that_places_the_block():
    """A bottom-anchored block is placed by `marB` alone: when `top − Δ` would be negative, `marB` still
    gets `bottom + Δ` as long as the row has the room (the top inset takes what is left)."""
    ctx = SimpleNamespace(options=EmitOptions())
    cell = _rich_ir_cell(0, 1)
    layout = text_engine.plan_cell(cell, {}, ctx)
    place = {"slot_top_px": 120.0, "slot_height_px": 60.0, "top_px": 0.0, "bottom_px": 20.0, "anchor": "bottom"}
    delta = text_engine.cell_margin_shift_px(layout, **place)
    assert delta > 0.0
    top, bottom, residual = text_engine.cell_margins_px(layout, **place)
    assert bottom == pytest.approx(20.0 + delta) and top >= 0.0 and residual == 0.0


def test_cell_opacity_reaches_the_runs(emitted):
    """Cell runs are written against the table element, so its opacity reaches their colour."""
    rPr = _rich_cell(emitted, 1, 0).text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    assert rPr.find(f"{{{A}}}solidFill/{{{A}}}srgbClr/{{{A}}}alpha").get("val") == "90000"


def test_cell_paragraph_gap_is_space_before(emitted):
    """The measured 6 px gap between two cell paragraphs is the second one's `spcBef`."""
    second = _rich_cell(emitted, 1, 0).text_frame.paragraphs[1]._p.find(f"{{{A}}}pPr")
    assert int(second.find(f"{{{A}}}spcBef/{{{A}}}spcPts").get("val")) == round(6 * config.PT_PER_PX * 100)


def test_overlays_are_above_the_table(emitted):
    """Every overlay follows the table's graphicFrame in `spTree`; the gradient cell precedes it."""
    out, _, _, _, _ = emitted["emit-tables-rich"]
    names = [shape.name for shape in shapes_of(out)]
    table = names.index("rich-cells")
    assert names.index("rich-cells r2c2 cellGradient") < table
    assert names.index("rich-cells r0c0 chip") > table
    assert names.index("rich-cells r0c0 chipText") > names.index("rich-cells r0c0 chip")


def test_width_lock_leaves_missing_glyphs_alone():
    """An emoji flag has no glyph in Calibri: its predicted width is a guess, so no tracking."""
    text = "\U0001F1F8\U0001F1EC Singapore"
    predicted = text_engine.text_width_px(text, "Calibri", 14.667, 400, False)
    paragraph = {"align": "left", "lineHeightPx": 18, "spaceBeforePx": 0, "spaceAfterPx": 0, "bullet": None,
                 "lines": [{"box": {"x": 0, "y": 0, "w": predicted - 8.0, "h": 18},
                            "runs": [{"text": text, "font": "Calibri", "sizePx": 14.667, "weight": 400,
                                      "italic": False}]}]}
    layout = text_engine.plan_cell({"paragraphs": [paragraph]}, {}, SimpleNamespace(options=EmitOptions()))
    line = layout.paragraphs[0].lines[0]
    assert line.predicted == line.width and line.tracking_px == 0.0
    assert any("no glyph" in warning for warning in layout.warnings)
    calibri = text_engine.face_for("Calibri")
    assert not text_engine._has_glyphs(calibri, "\U0001F1F8")
    assert text_engine._has_glyphs(calibri, "Singapore‍")


# --------------------------------------------------------------------------------------- charts


def test_chart_is_a_native_part_with_a_workbook(emitted):
    out, report, _, _, _ = emitted["emit-charts"]
    frame = by_name(out, "growth")
    assert frame.has_chart
    chart = frame.chart
    assert [c for c in chart.plots[0].categories] == ["2019", "2021", "2022", "2023"]
    assert [s.name for s in chart.series] == ["Internal", "Partner"]
    assert list(chart.series[0].values) == [40, 52, 57, 71]
    assert chart.part.chart_workbook.xlsx_part is not None
    assert report.charts[0]["id"] == "e1"
    assert report.charts[0]["origin"] == "authored"
    assert report.charts[0]["type"] == "column_stacked"


def test_plot_rect_becomes_the_chart_spec_plot_area(emitted):
    """The design's plot rectangle is handed to the chart engine as fractions of the frame.

    Turning `options.plotArea` into `c:manualLayout` is WP3a's half of the job (the vendored
    StageFlow `add` does not write one yet), so what is checked here is the half WP4 owns: the
    fractions the emitter derives from `plotRect`, reported so the two halves can be compared.
    """
    _, report, ir, _, _ = emitted["emit-charts"]
    element = ir.elements[0]
    assert report.charts[0]["plotArea"] == {
        "x": round((element.plotRect.x - element.box.x) / element.box.w, 4),
        "y": round((element.plotRect.y - element.box.y) / element.box.h, 4),
        "w": round(element.plotRect.w / element.box.w, 4),
        "h": round(element.plotRect.h / element.box.h, 4),
    }


def test_chart_overlay_is_painted_after_the_chart(emitted):
    out, _, _, _, _ = emitted["emit-charts"]
    names = [s.name for s in shapes_of(out)]
    assert names.index("growth") < names.index("callout")


# -------------------------------------------------------------------------------- emit options


def test_line_lock_off_lets_powerpoint_reflow(tmp_path):
    """The A/B arm: no `a:br`, no exact spacing, and the words still read as one sentence."""
    out, _, _, _, _ = emit_fixture("emit-text", tmp_path,
                                   options=EmitOptions(line_lock=False))
    paragraph = by_name(out, "paragraphs").text_frame.paragraphs[0]
    assert paragraph._p.findall(f"{{{A}}}br") == []
    assert paragraph._p.get_or_add_pPr().find(f"{{{A}}}lnSpc") is None
    assert "".join(r.text for r in paragraph.runs) == \
        "A paragraph the browser broke across exactly two lines."
    # Unlocked, PowerPoint must be allowed to break the paragraph itself — except a one-line frame,
    # which a narrow measure would stack one character per line (bad2f9b; `text._wraps`).
    assert _body_properties(by_name(out, "paragraphs"))["wrap"] == "square"
    assert _body_properties(by_name(out, "line-height-1"))["wrap"] == "none"


def test_group_svg_off_flattens_groups(tmp_path):
    out, report, _, _, _ = emit_fixture("emit-shapes", tmp_path,
                                        options=EmitOptions(group_svg=False))
    assert "group" not in report.shapes_by_kind
    assert all(not hasattr(s, "shapes") for s in shapes_of(out))
    assert report.shapes_by_kind["shape"] == 19


def test_name_shapes_off_leaves_powerpoints_own_names(tmp_path):
    out, _, _, _, _ = emit_fixture("emit-shapes", tmp_path,
                                   options=EmitOptions(name_shapes=False))
    assert "donut" not in [s.name for s in shapes_of(out)]


def test_theme_fonts_off_writes_the_literal_family(tmp_path):
    ir = load_ir("emit-text")
    master, manifest = master_of(ir)
    manifest.masters[0].theme["fonts"] = {"major": "Arial", "minor": "Arial"}
    out = tmp_path / "literal.pptx"
    emitter.emit([map_placeholders(ir, manifest.layout(ir.slide.layoutId))], manifest, master, out,
                 EmitOptions(theme_fonts=False))
    latin = by_name(out, "paragraphs").text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr() \
        .find(f"{{{A}}}latin")
    assert latin.get("typeface") == "Arial"


def test_finalize_off_still_keeps_the_template_intact(tmp_path):
    """Determinism is optional; not touching the customer's master never is."""
    out, report, _, master, _ = emit_fixture("emit-shapes", tmp_path,
                                             options=EmitOptions(finalize=False))
    assert template_module.verify_template(out, master) == []
    assert report.warnings == []


# ---------------------------------------------------------------------------------- determinism


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


@pytest.mark.parametrize("name", ["emit-shapes", "emit-charts", "emit-images", "emit-tables-rich"])
def test_two_emits_are_byte_identical(name, tmp_path):
    first = emit_fixture(name, tmp_path / "a")[0]
    second = emit_fixture(name, tmp_path / "b")[0]
    assert _sha256(first) == _sha256(second), f"{name} is not reproducible"


def test_every_zip_entry_is_stamped_at_the_zip_epoch(emitted):
    out, _, _, _, _ = emitted["emit-charts"]
    with zipfile.ZipFile(out) as archive:
        dates = {info.date_time for info in archive.infolist()}
        assert dates == {emitter.FIXED_ZIP_DATE}
        core = archive.read("docProps/core.xml").decode("utf-8")
        assert emitter.FIXED_TIMESTAMP in core
        workbooks = [n for n in archive.namelist() if n.startswith("ppt/embeddings/")]
        assert workbooks, "the chart has no embedded workbook"
        for name in workbooks:
            with zipfile.ZipFile(__import__("io").BytesIO(archive.read(name))) as book:
                assert {i.date_time for i in book.infolist()} == {emitter.FIXED_ZIP_DATE}


def test_png_metadata_is_stripped_only_from_our_own_media():
    source = (b"\x89PNG\r\n\x1a\n"
              + b"\x00\x00\x00\x04tIME" + b"\x07\xe8\x01\x01" + b"\x00\x00\x00\x00"
              + b"\x00\x00\x00\x00IEND" + b"\xae\x42\x60\x82")
    stripped = emitter.strip_png_metadata(source)
    assert b"tIME" not in stripped and stripped.endswith(b"IEND\xae\x42\x60\x82")
    assert emitter.strip_png_metadata(b"not a png") == b"not a png"


# -------------------------------------------------------------------------------- the whole file


@pytest.mark.validator
def test_the_validator_is_happy_with_every_fixture(emitted):
    """`validate.py deck.pptx --original master.pptx` — the pptx skill's own OOXML check."""
    assert config.VALIDATE_PY.exists(), f"validator missing at {config.VALIDATE_PY}"
    for name, (out, _, _, master, _) in emitted.items():
        completed = subprocess.run(
            [sys.executable, str(config.VALIDATE_PY), str(out), "--original", str(master)],
            capture_output=True, text=True, timeout=300,
        )
        assert completed.returncode == 0, f"{name}:\n{completed.stdout}\n{completed.stderr}"


@pytest.mark.renderer_full
def test_the_renderer_reports_no_warnings_of_ours(emitted, tmp_path):
    """Renderer warnings on our decks minus the warnings an empty slide already produces."""
    manifest = Manifest.load(config.MASTERS_FIXTURE / "test-16x9.manifest.json")
    master = config.MASTERS_FIXTURE / "test-16x9.pptx"

    baseline_irs = []
    for index, layout_id in enumerate(sorted({load_ir(n).slide.layoutId for n in WIDE_FIXTURES})):
        baseline_irs.append(IR(canvas=Canvas(1280, 720),
                               slide=Slide(id=f"empty-{index}", layoutId=layout_id)))
    baseline = tmp_path / "baseline.pptx"
    emitter.emit(baseline_irs, manifest, master, baseline, EmitOptions())
    renderer.render(baseline, 1280, tmp_path / "baseline-png")
    expected = set(json.loads((tmp_path / "baseline-png" / "warnings.json")
                              .read_text(encoding="utf-8"))["warnings"])

    for name in WIDE_FIXTURES:
        out = emitted[name][0]
        renderer.render(out, 1280, tmp_path / name)
        warnings = json.loads((tmp_path / name / "warnings.json").read_text(encoding="utf-8"))
        ours = set(warnings["warnings"]) - expected
        assert not ours, f"{name}: {sorted(ours)}"
        assert not warnings["truncated"], f"{name}: the renderer hid {warnings['hidden']} warnings"


def test_four_by_three_is_the_same_engine(emitted):
    """Size generality: nothing in the emitter knows how wide a slide is."""
    out, report, ir, master, _ = emitted["emit-sizes-4x3"]
    presentation = Presentation(str(out))
    assert presentation.slide_width == 9144000 and presentation.slide_height == 6858000
    assert report.warnings == []
    assert report.placeholders_used == [{"slide": 1, "type": "title", "idx": 0}]
    donut = by_name(out, "donut-4x3")
    adjustments = {gd.get("name"): int(gd.get("fmla").split()[1])
                   for gd in donut._element.findall(f".//{{{A}}}gd")}
    assert adjustments["adj3"] == round((80 - 48) / 160 * 100000)
    assert template_module.verify_template(out, master) == []
