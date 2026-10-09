"""A merged cell's interior rule reaches its whole length in PowerPoint, in collapse tables too.

PowerPoint draws a merged cell's interior `lnB` / `lnR` along its first column / row only; under an
`hMerge` continuation it draws the lower cell's `lnT`, beside a `vMerge` one the right cell's `lnL`
— and that is what PowerPoint itself writes when a merged cell gets a bottom or right border
(measured through COM, 2026-09-29). `page.js: mergedEdges` gives those neighbours the rule;
`dedupeSharedEdges` still writes every other interior line once, so a table without a merged cell
that has an interior bottom or right rule is unchanged.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from app.engine.emit import pptx as emitter
from app.engine.ir import IR, Element
from app.engine.manifest import Manifest
from app.engine.reports import EmitOptions
from tests.engine.helpers import extract_html

BLUE = {"width": 2, "color": "1A9AFA", "dash": "solid"}
RED = {"width": 3, "color": "D64545", "dash": "solid"}
GREEN = {"width": 2, "color": "1E9E5A", "dash": "solid"}

#: A header spanning two columns over single cells, a rowspan beside single cells, a spanning row
#: over another spanning row, and a control table with no merged cell.
HTML = (
    '<!doctype html><html><body style="margin:0;font-family:Arial;font-size:12px">'
    '<table data-name="merged" style="position:absolute;left:40px;top:40px;border-collapse:collapse;'
    'table-layout:fixed;width:420px">'
    '<tr><th colspan="2" style="border-bottom:2px solid #1A9AFA">Market</th>'
    '<th style="border-bottom:2px solid #1A9AFA">2030</th></tr>'
    '<tr><td rowspan="2" style="border-right:3px solid #D64545">Partner</td><td>Upsell</td><td>31.0</td></tr>'
    '<tr><td>Business</td><td>17.5</td></tr>'
    '<tr><td colspan="2" style="border-bottom:2px solid #1E9E5A">Subtotal</td><td>48.5</td></tr>'
    '<tr><td colspan="2">Total</td><td>48.5</td></tr></table>'
    '<table data-name="plain" style="position:absolute;left:40px;top:300px;border-collapse:collapse">'
    '<tr><td style="border-bottom:2px solid #1A9AFA">A</td><td style="border-bottom:2px solid #1A9AFA">B</td></tr>'
    '<tr><td style="border-right:1px solid #999">C</td><td>D</td></tr></table>'
    '</body></html>'
)


@pytest.fixture(scope="module")
def merged(tmp_path_factory, sample_manifest: Manifest, sample_dir: Path) -> tuple[IR, Path]:
    out = tmp_path_factory.mktemp("merged")
    html = out / "merged.html"
    html.write_text(HTML, encoding="utf-8")
    ir = extract_html(html, sample_manifest, "layout-05", sample_dir / "assets", slide_id="merged")
    deck = out / "merged.pptx"
    emitter.emit([ir], sample_manifest, sample_dir / "master.pptx", deck, EmitOptions())
    return ir, deck


def _table(ir: IR, name: str) -> Element:
    return next(e for e in ir.elements if e.kind == "table" and e.name == name)


def _cells(table: Element) -> dict[tuple[int, int], dict]:
    return {(c["r"], c["c"]): c for c in table.cells}


def test_the_cells_beyond_a_merge_carry_its_rule(merged):
    ir, _ = merged
    cells = _cells(_table(ir, "merged"))
    assert cells[(0, 0)]["colSpan"] == 2 and cells[(0, 0)]["borders"]["bottom"] == BLUE
    assert cells[(1, 1)]["borders"]["top"] == BLUE, "under the header's continuation column"
    assert cells[(1, 0)]["borders"]["top"] is None, "the origin's own column is its lnB"
    assert cells[(1, 2)]["borders"]["top"] is None, "a single upper cell's line stays on its lnB"
    assert cells[(1, 0)]["rowSpan"] == 2 and cells[(1, 0)]["borders"]["right"] == RED
    assert cells[(2, 1)]["borders"]["left"] == RED, "beside the rowspan's continuation row"
    assert cells[(1, 1)]["borders"]["left"] is None, "the origin's own row is its lnR"
    # A spanning row over a spanning row: PowerPoint writes the rule on the lower origin's lnT.
    assert cells[(3, 0)]["borders"]["bottom"] == GREEN and cells[(4, 0)]["borders"]["top"] == GREEN


def test_a_table_without_merged_interior_rules_writes_each_line_once(merged):
    ir, _ = merged
    cells = _cells(_table(ir, "plain"))
    assert cells[(0, 0)]["borders"]["bottom"] == BLUE and cells[(0, 1)]["borders"]["bottom"] == BLUE
    assert all(cells[(1, c)]["borders"]["top"] is None for c in (0, 1))
    assert cells[(1, 1)]["borders"]["left"] is None


def test_the_rule_reaches_the_file(merged):
    _, deck = merged
    table = next(s for s in Presentation(str(deck)).slides[0].shapes if s.has_table and s.name == "merged").table

    def lines(r: int, c: int) -> set[str]:
        tcPr = table.cell(r, c)._tc.find(qn("a:tcPr"))
        return {child.tag.split("}")[1] for child in tcPr} & {"lnL", "lnR", "lnT", "lnB"}

    assert table.cell(0, 0)._tc.get("gridSpan") == "2" and "lnB" in lines(0, 0)
    assert "lnT" in lines(1, 1) and "lnL" in lines(2, 1) and "lnT" in lines(4, 0)
    assert "lnT" not in lines(1, 0) and "lnT" not in lines(1, 2)
