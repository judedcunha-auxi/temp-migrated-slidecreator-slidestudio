"""The duplicate-parts fix in `strip_slides` / `_rename_new_slides` (app/engine/emit/pptx.py).

Two kinds of master used to produce a `.pptx` with two zip entries of the same name, which
PowerPoint answers with a repair prompt (and a repair drops shapes):

* a master with a **custom show**: `p:custShowLst` refers to slides by the same relationship id as
  the slide list, so `drop_rel` (which keeps a relationship referenced twice) kept the old slide part
  and the new slide took its name;
* a master whose **layout links to a slide**: the old slide part stays reachable, and python-pptx
  names the new slide after the slide count, colliding with it.

Every master here is synthetic: python-pptx's default template with invented slides.
"""

from __future__ import annotations

import collections
import re
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from app.engine.emit.pptx import _rename_new_slides, emit, strip_slides
from app.engine.importer import import_master
from app.engine.ir import IR, Box, Canvas, Element, Slide
from app.engine.renderer import BlankLayoutsRenderer

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _master_with_slides(path: Path, count: int = 3, *, custom_show: bool = False,
                        layout_link: bool = False) -> Path:
    prs = Presentation()
    for number in range(1, count + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Synthetic slide {number}"
    if custom_show:
        root = prs.part._element
        rids = [entry.get(f"{{{R}}}id") for entry in prs.slides._sldIdLst]
        shows = etree.Element(f"{{{P}}}custShowLst")
        root.insert(list(root).index(root.find(f"{{{P}}}notesSz")) + 1, shows)
        show = etree.SubElement(shows, f"{{{P}}}custShow", name="Short", id="0")
        slides = etree.SubElement(show, f"{{{P}}}sldLst")
        for rid in rids[:2]:
            etree.SubElement(slides, f"{{{P}}}sld").set(f"{{{R}}}id", rid)
    if layout_link:
        # A "back to agenda" button: the first layout relates to the first slide.
        prs.slide_layouts[0].part.relate_to(prs.slides[0].part, RT.SLIDE)
    prs.save(str(path))
    return path


def _duplicates(path: Path) -> list[str]:
    names = zipfile.ZipFile(path).namelist()
    return sorted(name for name, seen in collections.Counter(names).items() if seen > 1)


def _slide_entries(path: Path) -> list[str]:
    return sorted(n for n in zipfile.ZipFile(path).namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))


def _strip_and_add(master: Path, out: Path, new: int = 3) -> Path:
    prs = Presentation(str(master))
    strip_slides(prs)
    slides = [prs.slides.add_slide(prs.slide_layouts[0]) for _ in range(new)]
    _rename_new_slides(prs, slides)
    prs.save(str(out))
    return out


def test_custom_show_master_gives_no_duplicate_parts(tmp_path: Path):
    master = _master_with_slides(tmp_path / "custom-show.pptx", custom_show=True)
    out = _strip_and_add(master, tmp_path / "out.pptx")
    assert _duplicates(out) == []
    assert _slide_entries(out) == [f"ppt/slides/slide{n}.xml" for n in (1, 2, 3)]
    presentation = zipfile.ZipFile(out).read("ppt/presentation.xml")
    assert b"custShowLst" not in presentation          # the stale custom show is gone


def test_strip_slides_drops_the_slide_relationships_even_when_referenced_twice(tmp_path: Path):
    master = _master_with_slides(tmp_path / "custom-show.pptx", custom_show=True)
    prs = Presentation(str(master))
    assert strip_slides(prs) == 3
    slide_rels = [rel for rel in prs.part.rels.values() if rel.reltype == RT.SLIDE]
    assert slide_rels == []
    assert len(prs.slides) == 0


def test_layout_link_master_gives_no_duplicate_parts_and_keeps_the_link(tmp_path: Path):
    master = _master_with_slides(tmp_path / "layout-link.pptx", layout_link=True)
    out = _strip_and_add(master, tmp_path / "out.pptx")
    assert _duplicates(out) == []
    entries = _slide_entries(out)
    # The old slide the layout links to survives under its own name; the three new ones go around it.
    assert entries == [f"ppt/slides/slide{n}.xml" for n in (1, 2, 3, 4)]
    # (Read from the saved rels: python-pptx renames slide parts in memory when `prs.slides` is
    # first touched, so a reopened Presentation would not show the saved names.)
    presentation_rels = zipfile.ZipFile(out).read("ppt/_rels/presentation.xml.rels").decode()
    assert re.findall(r'Target="slides/(slide\d+)\.xml"', presentation_rels) == ["slide2", "slide3", "slide4"]
    layout_rels = zipfile.ZipFile(out).read("ppt/slideLayouts/_rels/slideLayout1.xml.rels")
    assert b"slides/slide1.xml" in layout_rels         # the template's own link is untouched


def test_without_the_fix_the_same_masters_do_duplicate(tmp_path: Path):
    """The regression the fix answers, reproduced: `drop_rel` and python-pptx's own naming."""
    master = _master_with_slides(tmp_path / "custom-show.pptx", custom_show=True)
    prs = Presentation(str(master))
    for entry in list(prs.slides._sldIdLst):
        prs.part.drop_rel(entry.get(f"{{{R}}}id"))
        prs.slides._sldIdLst.remove(entry)
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(tmp_path / "old.pptx"))
    assert "ppt/slides/slide1.xml" in _duplicates(tmp_path / "old.pptx")


def _text_ir(slide_id: str, layout_id: str) -> IR:
    run = {"text": "Synthetic text", "font": "Arial", "sizePx": 20.0, "weight": 400, "italic": False,
           "color": "2E2E38", "alpha": 1, "underline": False, "strike": False, "letterSpacingPx": 0,
           "baseline": "normal", "baselineShiftPx": 0, "caps": False, "href": None}
    paragraph = {"align": "left", "lineHeightPx": 24.0, "spaceBeforePx": 0.0, "spaceAfterPx": 0.0,
                 "bullet": None, "lines": [{"box": {"x": 100, "y": 300, "w": 150, "h": 24}, "runs": [run]}]}
    element = Element(kind="text", box=Box(100, 300, 400, 24), paragraphs=[paragraph], id="e1", z=1)
    return IR(canvas=Canvas(960, 720), slide=Slide(id=slide_id, title=None, layoutId=layout_id),
              elements=[element])


def test_emit_on_masters_with_custom_shows_and_layout_links(tmp_path: Path):
    """End to end through `emit`: the saved deck has unique parts and only the new slides listed."""
    for flag in ("custom_show", "layout_link"):
        master = _master_with_slides(tmp_path / f"{flag}.pptx", **{flag: True})
        manifest = import_master(master, tmp_path / f"{flag}-import", renderer=BlankLayoutsRenderer())
        layout = manifest.layouts[1].id
        out = tmp_path / f"{flag}-out.pptx"
        emit([_text_ir("s1", layout), _text_ir("s2", layout)], manifest, master, out)
        assert _duplicates(out) == [], flag
        assert len(Presentation(str(out)).slides) == 2
