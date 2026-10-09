"""WP3a — the chart emitter: every Path A type, every option in the authoring contract.

What these tests hold down, in the order the acceptance gate asks for it (master brief §9, Charts):

* every `data-chart` type Path A claims emits a real chart part python-pptx can reopen;
* the **embedded workbook** holds exactly the spec's categories and values (`expected_workbook`
  says what "exactly" means for a waterfall, whose base series is derived);
* the pptx-skill validator accepts the deck and PptxRender renders it without chart warnings —
  except for the two types the renderer itself cannot draw, which the test names;
* `c:manualLayout` puts the plot rectangle where `options.plotArea` / the IR's `plotRect` said.

Nothing here skips: a missing fixture or a missing validator is a failure (master brief §10.7).
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

from app.config import engine as config
from app.engine import charts_spec
from app.engine.emit import charts
from app.engine.emit.pptx import open_master, resolve_layout, strip_slides
from app.engine.ir import IR
from app.engine.manifest import Manifest
from tests.engine import helpers as renderer

C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

#: Where the fixture charts are drawn on the 16 × 9 test master, in inches.
BOX = (0.8, 0.8, 6.0, 3.2)
STYLE = {"font": "Calibri", "sizePx": 13.0, "color": "516467"}

#: Path A types PptxRender's ChartPainter cannot draw. It paints Bar / HorizontalBar / Line / Area /
#: Pie / Doughnut / Scatter and treats everything else as unsupported — a **renderer** gap, not an
#: emitter defect (the files are valid and PowerPoint draws them). Named here so the day the
#: renderer grows `ofPieChart`, or the day these types leave Path A, a test says so.
RENDERER_CANNOT_DRAW = {"pie_of_pie", "bar_of_pie"}


# ------------------------------------------------------------------------------------- helpers


def _deck(master: Path, manifest: Manifest, layout_id: str = "layout-07"):
    """An empty deck on the test master, plus the layout every fixture chart sits on."""
    presentation = open_master(master)
    strip_slides(presentation)
    return presentation, resolve_layout(presentation, manifest, manifest.layout(layout_id))


def _chart_xml(chart) -> etree._Element:
    return chart._chartSpace


def _find(element, path: str):
    return element.find(path.replace("c:", f"{{{C}}}").replace("a:", f"{{{A}}}"))


def _findall(element, path: str):
    return element.findall(path.replace("c:", f"{{{C}}}").replace("a:", f"{{{A}}}"))


def _val(element, path: str) -> str | None:
    found = _find(element, path)
    return None if found is None else found.get("val")


def primary_value_axis(chart):
    """The value axis the chart's own groups are plotted on. python-pptx's `chart.value_axis` takes
    the *second* `c:valAx` whenever there are several, which on a waterfall is the in-chart
    connectors' deleted x axis (WP-C C3)."""
    from pptx.chart.axis import ValueAxis

    space = _chart_xml(chart)
    groups = [g for g in _findall(space, ".//c:plotArea/*") if etree.QName(g).localname.endswith("Chart")]
    has_categories = _find(space, ".//c:plotArea/c:catAx") is not None
    plotted = {e.get("val") for g in groups
               if not (has_categories and etree.QName(g).localname == "scatterChart")
               for e in _findall(g, "c:axId")}
    for axis in _findall(space, ".//c:plotArea/c:valAx"):
        if _val(axis, "c:axId") in plotted:
            return ValueAxis(axis)
    return chart.value_axis


def bar_series_names(chart) -> list[str]:
    """The names of the chart's first plot group — a waterfall's bars, without its connector series."""
    return [series.name for series in chart.plots[0].series]


def connector_series(chart) -> list[tuple[list, list]]:
    """A waterfall's in-chart connectors: `(x values, y values)` per series, from the literal caches."""
    def literal(series, tag: str) -> list:
        holder = _find(series, f"c:{tag}/c:numLit")
        count = int(_val(holder, "c:ptCount"))
        points = {int(p.get("idx")): float(_find(p, "c:v").text) for p in _findall(holder, "c:pt")}
        return [points.get(i) for i in range(count)]

    return [(literal(series, "xVal"), literal(series, "yVal"))
            for series in _findall(_chart_xml(chart), ".//c:scatterChart/c:ser")]


def embedded_workbook(chart) -> dict:
    """The chart's own xlsx, read as `{"categories": [...], "series": [{"name", "values"}]}`.

    Read from the workbook rather than the XML caches on purpose: the workbook is what opens when
    someone double-clicks the chart, and "native, editable" is the whole point of Path A.
    """
    blob = chart.part.chart_workbook.xlsx_part.blob
    with zipfile.ZipFile(BytesIO(blob)) as archive:
        sheet = etree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        strings = [
            "".join(node.itertext())
            for node in etree.fromstring(archive.read("xl/sharedStrings.xml"))
        ]

    cells: dict[str, str | float | None] = {}
    for row in sheet.iter(f"{{{S}}}row"):
        for cell in row.iter(f"{{{S}}}c"):
            value = cell.find(f"{{{S}}}v")
            reference = cell.get("r")
            if value is None or value.text is None:
                cells[reference] = None
            elif cell.get("t") == "s":
                cells[reference] = strings[int(value.text)]
            else:
                cells[reference] = float(value.text)

    rows = max((int(re.sub(r"[A-Z]+", "", ref)) for ref in cells), default=1)
    columns = sorted({re.sub(r"\d+", "", ref) for ref in cells})
    categories = [cells.get(f"A{row}") for row in range(2, rows + 1)]
    series = []
    for column in columns[1:]:
        series.append({
            "name": cells.get(f"{column}1"),
            "values": [cells.get(f"{column}{row}") for row in range(2, rows + 1)],
        })
    return {"categories": categories, "series": series}


def xy_cache(chart) -> list[dict]:
    """`[{"name", "x", "y"}]` from an XY chart's own numeric caches."""
    series = []
    for element in _findall(_chart_xml(chart), ".//c:ser"):
        name = _find(element, "c:tx/c:strRef/c:strCache/c:pt/c:v")
        series.append({
            "name": None if name is None else name.text,
            "x": [float(point.text) for point in
                  _findall(element, "c:xVal/c:numRef/c:numCache/c:pt/c:v")],
            "y": [float(point.text) for point in
                  _findall(element, "c:yVal/c:numRef/c:numCache/c:pt/c:v")],
        })
    return series


def spec_for(kind: str) -> dict:
    """A minimal, valid spec for any Path A type — enough to prove the type itself emits."""
    if kind in charts_spec.XY_TYPES:
        return {"type": kind,
                "series": [{"name": "A", "x": [1, 2, 3], "y": [2, 4, 3]},
                           {"name": "B", "x": [1, 2, 3], "y": [1, 3, 5]}],
                "colors": ["#1A9AFB", "#E4573D"]}
    if kind in charts_spec.PIE_TYPES:
        return {"type": kind, "categories": ["a", "b", "c"],
                "series": [{"name": "S", "values": [5, 3, 2]}],
                "colors": ["#1A9AFB", "#E4573D", "#516467"],
                "options": {"holeSize": 62} if kind.startswith("doughnut") else {}}
    if kind == "waterfall":
        return {"type": "waterfall", "categories": ["start", "up", "down", "end"],
                "series": [{"name": "Steps", "values": [100, 40, -15, 125]}],
                "colors": ["#1A9AFB", "#E4573D", "#516467"],
                "options": {"numberFormat": "0",
                            "plotArea": {"x": 0.1, "y": 0.06, "w": 0.86, "h": 0.76}}}
    return {"type": kind, "categories": ["a", "b", "c"],
            "series": [{"name": "One", "values": [3, 5, 4]}, {"name": "Two", "values": [2, 1, 6]}],
            "colors": ["#1A9AFB", "#E4573D"], "options": {"numberFormat": "0"}}


PATH_A_KINDS = sorted(charts_spec.TYPES - charts_spec.ROUND_2_TYPES)


# ------------------------------------------------------------------------------------ fixtures


@pytest.fixture(scope="session")
def test_master(masters_dir: Path) -> Path:
    master = masters_dir / "test-16x9.pptx"
    assert master.exists(), f"{master} is missing — WP0b's test master is a WP3a input"
    return master


@pytest.fixture(scope="session")
def test_manifest(masters_dir: Path) -> Manifest:
    path = masters_dir / "test-16x9.manifest.json"
    assert path.exists(), f"{path} is missing — WP0b's test manifest is a WP3a input"
    return Manifest.load(path)


@pytest.fixture(scope="session")
def every_type_deck(test_master: Path, test_manifest: Manifest, tmp_path_factory) -> Path:
    """One slide per Path A chart type, in a fixed order — built once, inspected by many tests."""
    presentation, layout = _deck(test_master, test_manifest)
    for kind in PATH_A_KINDS:
        slide = presentation.slides.add_slide(layout)
        charts.add_chart(slide, spec_for(kind), BOX, dict(STYLE, name=kind))
    path = tmp_path_factory.mktemp("charts") / "every-type.pptx"
    presentation.save(str(path))
    return path


@pytest.fixture(scope="session")
def every_type_warnings(every_type_deck: Path, tmp_path_factory) -> dict:
    out = tmp_path_factory.mktemp("render")
    pngs = renderer.render(every_type_deck, 1280, out)
    assert len(pngs) == len(PATH_A_KINDS)
    return json.loads((out / "warnings.json").read_text(encoding="utf-8"))


@pytest.fixture
def one_chart(test_master: Path, test_manifest: Manifest, tmp_path: Path):
    """`one_chart(spec, **style)` → the reopened chart, plus the slide's shapes and diagnostics."""

    def build(spec: dict, *, box: tuple = BOX, **style):
        assert charts_spec.validate(spec) == [], charts_spec.validate(spec)
        presentation, layout = _deck(test_master, test_manifest)
        slide = presentation.slides.add_slide(layout)
        diagnostics: list[dict] = []
        charts.add_chart(slide, spec, box, dict(STYLE, diagnostics=diagnostics, **style))
        shape_names = [shape.name for shape in slide.shapes]
        path = tmp_path / "chart.pptx"
        presentation.save(str(path))
        reopened = Presentation(str(path))
        chart = next(s for s in reopened.slides[0].shapes if s.has_chart).chart
        return chart, shape_names, diagnostics, path

    return build


# --------------------------------------------------------------------------- types and workbook


def test_every_path_a_type_emits_a_chart_part(every_type_deck: Path):
    presentation = Presentation(str(every_type_deck))
    assert len(presentation.slides) == len(PATH_A_KINDS)
    for kind, slide in zip(PATH_A_KINDS, presentation.slides, strict=False):
        frames = [shape for shape in slide.shapes if shape.has_chart]
        assert len(frames) == 1, f"{kind}: expected exactly one chart frame"
        assert frames[0].name == kind
        plot_area = _find(_chart_xml(frames[0].chart), ".//c:plotArea")
        tags = {etree.QName(child).localname for child in plot_area}
        assert charts.target_element(kind) in tags, f"{kind}: plot area holds {sorted(tags)}"


@pytest.mark.parametrize("kind", PATH_A_KINDS)
def test_embedded_workbook_equals_the_spec(kind: str, every_type_deck: Path):
    presentation = Presentation(str(every_type_deck))
    slide = presentation.slides[PATH_A_KINDS.index(kind)]
    chart = next(s for s in slide.shapes if s.has_chart).chart
    expected = charts.expected_workbook(spec_for(kind))

    if expected["categories"] is None:
        # An XY workbook stacks each series as its own block of rows rather than as a column pair,
        # so the numbers are checked in the chart's own caches — which is what PowerPoint plots.
        actual_xy = xy_cache(chart)
        assert [s["name"] for s in actual_xy] == [s["name"] for s in expected["series"]]
        for got, want in zip(actual_xy, expected["series"], strict=False):
            assert got["x"] == pytest.approx(want["x"])
            assert got["y"] == pytest.approx(want["y"])
        return

    actual = embedded_workbook(chart)
    assert [c for c in actual["categories"]] == [
        None if c == "" else c for c in expected["categories"]]
    assert [s["name"] for s in actual["series"]] == [s["name"] for s in expected["series"]]
    for got, want in zip(actual["series"], expected["series"], strict=False):
        assert got["values"] == want["values"], f"{kind}/{want['name']}"


@pytest.mark.validator
def test_the_validator_accepts_every_type(every_type_deck: Path, test_master: Path):
    assert config.VALIDATE_PY.exists(), (
        f"the pptx-skill validator is missing at {config.VALIDATE_PY}; set SLIDE_ENGINE_VALIDATE_PY"
    )
    result = subprocess.run(
        [sys.executable, str(config.VALIDATE_PY), str(every_type_deck),
         "--original", str(test_master)],
        capture_output=True, text=True, timeout=300,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "PASSED" in result.stdout, result.stdout


@pytest.mark.renderer_full
def test_the_renderer_draws_every_type_it_can(every_type_warnings: dict):
    """The only chart warnings are the renderer's own gap, and we can name every one of them."""
    assert not every_type_warnings["truncated"], every_type_warnings
    warned = set()
    for warning in every_type_warnings["warnings"]:
        match = re.match(r"slide (\d+):", warning)
        assert match, f"unparsed renderer warning: {warning}"
        warned.add(PATH_A_KINDS[int(match.group(1)) - 1])
    assert warned == RENDERER_CANNOT_DRAW, every_type_warnings["warnings"]
    assert {k for k in PATH_A_KINDS if not charts.previewable(k)} == RENDERER_CANNOT_DRAW


def test_a_type_the_renderer_cannot_draw_says_so(one_chart):
    _chart, _shapes, diagnostics, _path = one_chart(spec_for("pie_of_pie"))
    assert any("PptxRender cannot draw" in d["message"] for d in diagnostics), diagnostics


# ------------------------------------------------------------------------------------ waterfall


def test_waterfall_plan_floats_the_bars_on_an_invisible_base():
    spec = {"type": "waterfall", "categories": ["Open", "up", "down", "", "Close"],
            "series": [{"name": "Partners", "values": [180, 95, -40, None, 235]}]}
    plan = charts.waterfall_plan(spec)
    assert plan["roles"] == ["total", "increase", "decrease", "gap", "total"]
    assert plan["base"] == [0.0, 180.0, 235.0, None, 0.0]
    assert plan["delta"] == [180.0, 95.0, 40.0, None, 235.0]
    assert plan["levels"] == [180.0, 275.0, 235.0, 235.0, 235.0]
    assert plan["totals"] == [0, 4], "the closing bar equals the running total, so it is a total"


def test_waterfall_totals_can_be_declared():
    spec = {"type": "waterfall", "categories": ["a", "b", "c"],
            "series": [{"name": "s", "values": [10, 20, 30]}],
            "options": {"totals": [0]}}
    plan = charts.waterfall_plan(spec)
    assert plan["roles"] == ["total", "increase", "increase"]
    assert plan["levels"] == [10.0, 30.0, 60.0]


def test_waterfall_emits_base_plus_steps_with_role_colours(one_chart):
    spec = {"type": "waterfall", "categories": ["Open", "up", "down", "", "Close"],
            "series": [{"name": "Partners", "values": [180, 95, -40, None, 235]}],
            "colors": ["#1A9AFB", "#E4573D", "#516467"],
            "options": {"numberFormat": "0", "totals": [0, 4],
                        "plotArea": {"x": 0.1, "y": 0.06, "w": 0.86, "h": 0.76}}}
    chart, shapes, _diagnostics, _path = one_chart(spec, name="Chart")

    assert bar_series_names(chart) == [charts.WATERFALL_BASE_SERIES, "Partners"]
    assert embedded_workbook(chart)["series"] == [
        {"name": "Base", "values": [0.0, 180.0, 235.0, None, 0.0]},
        {"name": "Partners", "values": [180.0, 95.0, 40.0, None, 235.0]},
    ]

    series = _findall(_chart_xml(chart), ".//c:ser")
    assert _find(series[0], "c:spPr/a:noFill") is not None, "the base series must not paint"
    colours = {int(_val(point, "c:idx")): _val(point, "c:spPr/a:solidFill/a:srgbClr")
               for point in _findall(series[1], "c:dPt")}
    assert colours == {0: "516467", 1: "1A9AFB", 2: "E4573D", 4: "516467"}

    formats = {int(_val(label, "c:idx")): _find(label, "c:numFmt").get("formatCode")
               for label in _findall(series[1], "c:dLbls/c:dLbl")
               if _find(label, "c:numFmt") is not None}
    # Two sections, so the sign is right on either side of zero (a rise below zero has a negative
    # cell, and the negative section prints its magnitude with the literal sign): WP-C §3.4.
    assert formats[1] == '"+"0;"+"0' and formats[2] == '"-"0;"-"0', "a step has to read as up or down"

    assert shapes == ["Chart"], "the connectors are a series of the chart, not shapes on top of it"
    pairs = connector_series(chart)
    assert [ys for _xs, ys in pairs] == [[180.0, 180.0], [275.0, 275.0]], \
        "one connector per consecutive pair, none touching the gap"
    assert [xs for xs, _ys in pairs] == [pytest.approx([0.8125, 1.1875]), pytest.approx([1.8125, 2.1875])]
    names = [_find(s, "c:tx/c:v").text for s in _findall(_chart_xml(chart), ".//c:scatterChart/c:ser")]
    assert names == ["Connector 1", "Connector 2"]


def test_waterfall_pins_its_value_axis_so_the_connectors_can_be_placed(one_chart):
    spec = {"type": "waterfall", "categories": ["a", "b", "c"],
            "series": [{"name": "s", "values": [100, 40, 140]}],
            "options": {"plotArea": {"x": 0.1, "y": 0.06, "w": 0.86, "h": 0.76}}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    axis = primary_value_axis(chart)
    assert axis.minimum_scale == 0.0
    # 0..140 rounds out to a tick-friendly 0..150 in steps of 50 — deterministic, and the same
    # range every export, which is what the connectors' own y axis is pinned to as well.
    assert axis.maximum_scale == pytest.approx(150.0)
    assert axis.major_unit == pytest.approx(50.0)
    secondary = _findall(_chart_xml(chart), ".//c:plotArea/c:valAx")[2]
    assert (_val(secondary, "c:scaling/c:min"), _val(secondary, "c:scaling/c:max")) == ("0.0", "150.0")


def test_waterfall_connector_sits_at_the_running_total():
    spec = {"type": "waterfall", "categories": ["a", "b"],
            "series": [{"name": "s", "values": [100, 40]}],
            "options": {"gapWidth": 50, "valueAxis": {"min": 0, "max": 200},
                        "plotArea": {"x": 0.0, "y": 0.0, "w": 1.0, "h": 1.0}}}
    (line,) = charts.overlay_geometry(spec, (0.0, 0.0, 4.0, 2.0))
    assert line["kind"] == "connector"
    # value 100 of 0..200 is half way up a 2 in plot whose top is y = 0 → y = 1.0 in.
    assert line["y1"] == pytest.approx(1.0) and line["y2"] == pytest.approx(1.0)
    # Two categories in a 4 in plot: slots [0,2] and [2,4], bars 4/3 in wide, centred.
    assert line["x1"] == pytest.approx(1.0 + 2.0 / 3.0)
    assert line["x2"] == pytest.approx(3.0 - 2.0 / 3.0)


# -------------------------------------------------------------------------------------- options


def test_gap_slot_keeps_the_bar_positions(one_chart):
    spec = {"type": "column_stacked", "categories": ["FY24", "FY25", "", "FY27E"],
            "series": [{"name": "D", "values": [57, 71, None, 88]},
                       {"name": "I", "values": [21, 18, None, 55]}],
            "colors": ["#B9C7C9", "#1A9AFB"]}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    workbook = embedded_workbook(chart)
    assert workbook["categories"] == ["FY24", "FY25", None, "FY27E"]
    assert workbook["series"][0]["values"] == [57.0, 71.0, None, 88.0]
    first = _findall(_chart_xml(chart), ".//c:ser")[0]
    points = _findall(first, "c:val/c:numRef/c:numCache/c:pt")
    assert int(_val(first, "c:val/c:numRef/c:numCache/c:ptCount")) == 4
    assert [p.get("idx") for p in points] == ["0", "1", "3"], "the gap keeps its slot, unplotted"


def test_plot_area_becomes_a_manual_layout(one_chart):
    spec = dict(spec_for("column"),
                options={"plotArea": {"x": 0.08, "y": 0.05, "w": 0.9, "h": 0.78}})
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    manual = _find(_chart_xml(chart), ".//c:plotArea/c:layout/c:manualLayout")
    assert _val(manual, "c:layoutTarget") == "inner"
    assert _val(manual, "c:xMode") == "edge" and _val(manual, "c:yMode") == "edge"
    assert [float(_val(manual, f"c:{tag}")) for tag in "xywh"] == [0.08, 0.05, 0.9, 0.78]


def test_plot_rect_from_the_ir_lands_within_four_px(one_chart):
    """The IR's `plotRect` (canvas px) is the design's plot rectangle; WP5 allows 4 px."""
    box = (1.0, 1.0, 4.0, 2.5)                                   # inches
    plot_rect = {"x": 96 + 30.0, "y": 96 + 18.0, "w": 330.0, "h": 190.0}   # canvas px
    chart, _shapes, _diagnostics, _path = one_chart(
        spec_for("column"), box=box, plotRect=plot_rect)
    manual = _find(_chart_xml(chart), ".//c:plotArea/c:layout/c:manualLayout")
    fractions = {tag: float(_val(manual, f"c:{tag}")) for tag in "xywh"}
    emitted = {
        "x": (box[0] + fractions["x"] * box[2]) * config.PX_PER_IN,
        "y": (box[1] + fractions["y"] * box[3]) * config.PX_PER_IN,
        "w": fractions["w"] * box[2] * config.PX_PER_IN,
        "h": fractions["h"] * box[3] * config.PX_PER_IN,
    }
    for key, value in plot_rect.items():
        assert abs(emitted[key] - value) < 4.0, f"{key}: {emitted[key]} vs {value}"
        assert abs(emitted[key] - value) <= abs(value) * 0.01, "and within 1 % (brief §WP3a)"


def test_reference_line_is_drawn_at_the_axis_scaled_position(one_chart):
    spec = dict(spec_for("column"),
                options={"valueAxis": {"min": 0, "max": 10},
                         "plotArea": {"x": 0.0, "y": 0.0, "w": 1.0, "h": 1.0},
                         "referenceLines": [{"value": 5, "color": "#516467", "dash": [4, 3],
                                             "label": "target"}]})
    _chart, shapes, diagnostics, path = one_chart(spec, box=(0.0, 0.0, 4.0, 2.0), name="Bars")
    assert diagnostics == []
    assert shapes == ["Bars", "Bars referenceLine 1", "Bars referenceLine 1 label"]
    presentation = Presentation(str(path))
    line = presentation.slides[0].shapes[1]
    assert line.top == pytest.approx(config.EMU_PER_IN * 1.0, rel=1e-3)   # half way up 2 in
    assert line.width == pytest.approx(config.EMU_PER_IN * 4.0, rel=1e-3)
    assert str(line.line.color.rgb) == "516467"
    assert _find(line.line._get_or_add_ln(), "a:custDash") is not None


def test_overlay_is_dropped_loudly_when_the_value_range_is_unknown(one_chart):
    """No pinned value axis means the line's height would be a guess — and a guess is worse.

    Before WP-C a missing plot rectangle dropped the line too. A chart that carries an overlay is
    now pinned to `chart_model.layout`'s default rect instead (see
    `test_reference_line_chart_pins_the_layout_rect`), so what is left to be unknown is the value
    range of an axis PowerPoint lays out itself.
    """
    spec = dict(spec_for("column"), options={"referenceLines": [{"value": 5}]})
    _chart, shapes, diagnostics, _path = one_chart(spec, name="Bars")
    assert shapes == ["Bars"]
    loud = [d for d in diagnostics if d["level"] != "info"]
    assert [d["level"] for d in loud] == ["warn"]
    assert "value-axis range" in loud[0]["message"]


def test_data_labels_can_be_set_per_series(one_chart):
    spec = dict(spec_for("column"), options={"dataLabels": [True, False]})
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    series = _findall(_chart_xml(chart), ".//c:ser")
    assert _val(series[0], "c:dLbls/c:showVal") == "1"
    assert _val(series[1], "c:dLbls/c:showVal") == "0"


def test_label_position_is_clamped_to_what_the_chart_type_allows(one_chart):
    """`outEnd` on a line chart is not a style choice, it is a repair prompt."""
    spec = dict(spec_for("line"), options={"dataLabels": True, "labelPosition": "outEnd"})
    chart, _shapes, diagnostics, _path = one_chart(spec)
    assert any("not available on a line chart" in d["message"] for d in diagnostics)
    assert _val(_find(_chart_xml(chart), ".//c:ser"), "c:dLbls/c:dLblPos") == "t"

    stacked = dict(spec_for("column_stacked"),
                   options={"dataLabels": True, "labelPosition": "outEnd"})
    chart, _shapes, diagnostics, _path = one_chart(stacked)
    assert any("stacked" in d["message"] or "column_stacked" in d["message"] for d in diagnostics)
    # A stacked chart's labels are centred by default (WP-C §0.9, Peter #8): PptxRender's own
    # undeclared default is `ctr` and the preview draws `center`, so file, render and app agree.
    assert _val(_find(_chart_xml(chart), ".//c:ser"), "c:dLbls/c:dLblPos") == "ctr"


def test_point_colours_are_written_per_point(one_chart):
    spec = dict(spec_for("column"), options={"pointColors": {"1": {"2": "#7FC6FF"}}})
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    second = _findall(_chart_xml(chart), ".//c:ser")[1]
    points = {int(_val(p, "c:idx")): _val(p, "c:spPr/a:solidFill/a:srgbClr")
              for p in _findall(second, "c:dPt")}
    assert points == {2: "7FC6FF"}


def test_hole_size_and_slice_colours(one_chart):
    spec = {"type": "doughnut", "categories": ["a", "b", "c"],
            "series": [{"name": "Share", "values": [61, 27, 12]}],
            "colors": ["#1A9AFB", "#B9C7C9", "#516467"],
            "options": {"holeSize": 62, "legend": "topRight"}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    assert _val(_chart_xml(chart), ".//c:doughnutChart/c:holeSize") == "62"
    points = {int(_val(p, "c:idx")): _val(p, "c:spPr/a:solidFill/a:srgbClr")
              for p in _findall(_chart_xml(chart), ".//c:ser/c:dPt")}
    assert points == {0: "1A9AFB", 1: "B9C7C9", 2: "516467"}
    assert _val(_chart_xml(chart), ".//c:legend/c:legendPos") == "tr"


def test_axis_options_reach_the_axis(one_chart):
    spec = dict(spec_for("column"), options={
        "valueAxis": {"min": 0, "max": 20, "majorUnit": 5, "format": "0.0", "title": "Orders"},
        "categoryAxis": {"reverse": True, "majorUnit": 2},
        "gridlines": {"color": "#E4E9F0", "width": 0.75, "dash": "sysDash"}})
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    assert (chart.value_axis.minimum_scale, chart.value_axis.maximum_scale) == (0.0, 20.0)
    assert chart.value_axis.major_unit == 5.0
    assert chart.value_axis.tick_labels.number_format == "0.0"
    assert chart.value_axis.has_title and chart.value_axis.axis_title.text_frame.text == "Orders"
    category = chart.category_axis._element
    assert _val(category, "c:scaling/c:orientation") == "maxMin"
    assert _val(category, "c:tickLblSkip") == "2" and _val(category, "c:tickMarkSkip") == "2"
    gridlines = _find(chart.value_axis._element, "c:majorGridlines/c:spPr/a:ln")
    assert _val(gridlines, "a:prstDash") == "sysDash"
    assert _find(gridlines, "a:solidFill/a:srgbClr").get("val") == "E4E9F0"


def test_series_dash_and_markers(one_chart):
    spec = {"type": "line_markers", "categories": ["a", "b", "c"],
            "series": [{"name": "Actual", "values": [1, 2, 3]},
                       {"name": "Plan", "values": [2, 2, 2], "dash": [6, 3], "lineWidth": 1.5}],
            "colors": ["#1A9AFB", "#B9C7C9"],
            "options": {"markers": "circle", "markerSize": 6}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    second = _findall(_chart_xml(chart), ".//c:ser")[1]
    assert _find(second, "c:spPr/a:ln/a:custDash") is not None
    assert _val(second, "c:marker/c:symbol") == "circle"
    assert _val(second, "c:marker/c:size") == "6"


def test_bar_geometry_and_slice_options(one_chart):
    spec = dict(spec_for("column"), options={"gapWidth": 40, "overlap": -10})
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    assert _val(_chart_xml(chart), ".//c:barChart/c:gapWidth") == "40"
    assert _val(_chart_xml(chart), ".//c:barChart/c:overlap") == "-10"

    pie = {"type": "pie_exploded", "categories": ["a", "b"],
           "series": [{"name": "s", "values": [6, 4], "explosion": 8}],
           "colors": ["#1A9AFB", "#B9C7C9"], "options": {"varyColors": True}}
    chart, _shapes, _diagnostics, _path = one_chart(pie)
    assert _val(_chart_xml(chart), ".//c:ser/c:explosion") == "8"
    assert _val(_chart_xml(chart), ".//c:pieChart/c:varyColors") == "1"


def test_the_ir_style_sets_the_chart_font(one_chart):
    """`style` is the design's own type: px in the IR, points in the file, no guessing in between."""
    chart, _shapes, _diagnostics, _path = one_chart(
        spec_for("column"), font="Arial", sizePx=12.0, color="516467")
    text_properties = _find(_chart_xml(chart), "c:txPr/a:p/a:pPr/a:defRPr")
    assert text_properties.get("sz") == "900", "12 px = 9 pt, at 0.01 pt precision"
    assert _find(text_properties, "a:latin").get("typeface") == "Arial"
    assert _find(text_properties, "a:solidFill/a:srgbClr").get("val") == "516467"

    # options.fontSize (px, like everything else in the contract) overrides the element's style.
    chart, _shapes, _diagnostics, _path = one_chart(
        dict(spec_for("column"), options={"fontSize": 16, "fontColor": "#12233B"}), sizePx=12.0)
    text_properties = _find(_chart_xml(chart), "c:txPr/a:p/a:pPr/a:defRPr")
    assert text_properties.get("sz") == "1200"
    assert _find(text_properties, "a:solidFill/a:srgbClr").get("val") == "12233B"


def test_a_hidden_axis_says_so_with_an_attribute(one_chart):
    """`<c:delete/>` alone is true by default and a reader that wants @val draws the axis anyway."""
    spec = dict(spec_for("column"), options={"valueAxis": {"visible": False}})
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    assert _val(chart.value_axis._element, "c:delete") == "1"
    assert _val(chart.category_axis._element, "c:delete") == "0"


def test_scatter_has_no_connecting_line_but_scatter_lines_does(one_chart):
    chart, _shapes, _diagnostics, _path = one_chart(spec_for("scatter"))
    assert _find(_findall(_chart_xml(chart), ".//c:ser")[0], "c:spPr/a:ln/a:noFill") is not None
    chart, _shapes, _diagnostics, _path = one_chart(spec_for("scatter_lines"))
    assert _find(_findall(_chart_xml(chart), ".//c:ser")[0], "c:spPr/a:ln/a:solidFill") is not None


def test_unknown_option_is_reported_not_ignored(one_chart, test_master, test_manifest, tmp_path):
    """`charts_spec.validate` refuses it up front; `add_chart` says so again if one gets through."""
    assert any("unknown option" in p for p in charts_spec.validate(
        dict(spec_for("column"), options={"gapWith": 60})))
    presentation, layout = _deck(test_master, test_manifest)
    slide = presentation.slides.add_slide(layout)
    diagnostics: list[dict] = []
    charts.add_chart(slide, dict(spec_for("column"), options={"gapWith": 60}), BOX,
                     dict(STYLE, diagnostics=diagnostics))
    assert any("unknown chart option 'gapWith'" in d["message"] for d in diagnostics)


def test_two_emissions_of_one_spec_are_identical(test_master, test_manifest, tmp_path):
    """Determinism starts here: the same spec must produce the same chart XML, byte for byte."""
    def chart_xml() -> bytes:
        presentation, layout = _deck(test_master, test_manifest)
        slide = presentation.slides.add_slide(layout)
        charts.add_chart(slide, _fixture_specs()[0], BOX, dict(STYLE, name="Chart"))
        return etree.tostring(slide.shapes[0].chart._chartSpace)

    assert chart_xml() == chart_xml()


# ------------------------------------------------------------------------- the torture fixture


def _fixture_specs() -> list[dict]:
    html = (config.TORTURE_FIXTURE / "charts.html").read_text(encoding="utf-8")
    return [json.loads(match) for match in re.findall(r"data-chart='(.*?)'", html, re.S)]


def test_the_torture_fixture_and_its_ir_hold_the_same_specs():
    ir = IR.load(config.TORTURE_FIXTURE / "ir" / "charts.json")
    assert ir.validate() == []
    from_ir = [element.spec for element in ir.of_kind("chart")]
    assert from_ir == _fixture_specs(), "charts.html and ir/charts.json have drifted apart"


def test_the_torture_fixture_matches_its_expect_file():
    expect = json.loads((config.TORTURE_FIXTURE / "charts.expect.json").read_text(encoding="utf-8"))
    ir = IR.load(config.TORTURE_FIXTURE / "ir" / "charts.json")
    assert ir.counts() == expect["kinds"]
    assert ir.slide.layoutId == expect["layoutId"]
    assert expect["charts"] == {"authored": len(ir.of_kind("chart")), "recognised": 0}
    assert all(element.origin == "authored" for element in ir.of_kind("chart"))


@pytest.mark.parametrize("index", range(8))
def test_every_fixture_spec_validates(index: int):
    assert charts_spec.validate(_fixture_specs()[index]) == []


@pytest.mark.renderer_full
@pytest.mark.validator
def test_the_fixture_ir_emits_every_chart_natively(test_master, test_manifest, tmp_path):
    """Acceptance: `fixtures/torture/charts.html`'s IR emits every spec as a native chart."""
    ir = IR.load(config.TORTURE_FIXTURE / "ir" / "charts.json")
    presentation, layout = _deck(test_master, test_manifest, ir.slide.layoutId)
    slide = presentation.slides.add_slide(layout)
    diagnostics: list[dict] = []
    for element in ir.of_kind("chart"):
        charts.add_chart(
            slide,
            element.spec,
            (element.box.x / config.PX_PER_IN, element.box.y / config.PX_PER_IN,
             element.box.w / config.PX_PER_IN, element.box.h / config.PX_PER_IN),
            dict(element.style, name=element.name, plotRect=element.plotRect,
                 diagnostics=diagnostics),
        )
    path = tmp_path / "torture-charts.pptx"
    presentation.save(str(path))

    reopened = Presentation(str(path))
    frames = [shape for shape in reopened.slides[0].shapes if shape.has_chart]
    assert len(frames) == 8
    assert [f.name for f in frames] == [e.name for e in ir.of_kind("chart")]
    assert [d for d in diagnostics if d["level"] != "info"] == []

    for frame, element in zip(frames, ir.of_kind("chart"), strict=False):
        expected = charts.expected_workbook(element.spec)
        actual = embedded_workbook(frame.chart)
        if expected["categories"] is None:
            continue
        assert [s["name"] for s in actual["series"]] == [s["name"] for s in expected["series"]]

    # The one chart carrying a plotRect instead of options.plotArea must still be laid out by hand.
    with_rect = next(e for e in ir.of_kind("chart") if e.plotRect is not None)
    frame = frames[[e.name for e in ir.of_kind("chart")].index(with_rect.name)]
    assert _find(_chart_xml(frame.chart), ".//c:plotArea/c:layout/c:manualLayout") is not None

    out = tmp_path / "render"
    renderer.render(path, 1280, out)
    warnings = json.loads((out / "warnings.json").read_text(encoding="utf-8"))
    assert warnings["warnings"] == [], "no chart warnings on the torture fixture"

    result = subprocess.run(
        [sys.executable, str(config.VALIDATE_PY), str(path), "--original", str(test_master)],
        capture_output=True, text=True, timeout=300,
    )
    assert result.returncode == 0 and "PASSED" in result.stdout, result.stdout + result.stderr


# ------------------------------------------------------------------- the spec, extended by WP3a


def test_waterfall_options_are_validated():
    base = {"type": "waterfall", "categories": ["a", "b", "c"],
            "series": [{"name": "s", "values": [10, 20, 30]}]}
    assert charts_spec.validate(dict(base, options={"totals": [0, 2]})) == []
    assert any("totals indices" in p for p in charts_spec.validate(
        dict(base, options={"totals": [0, 9]})))
    assert any("totals must be a list" in p for p in charts_spec.validate(
        dict(base, options={"totals": "first"})))
    assert any("connectors" in p for p in charts_spec.validate(
        dict(base, options={"connectors": {"colour": "#000000"}})))
    assert charts_spec.validate(dict(base, options={"connectors": False})) == []
    assert any("only has meaning on a waterfall" in p for p in charts_spec.validate(
        {"type": "column", "categories": ["a"], "series": [{"values": [1]}],
         "options": {"totals": [0]}}))
    # A second series makes a stacked waterfall (WP-C §3.2) — valid, unless one bar mixes signs.
    assert charts_spec.validate(dict(base, series=[{"name": "s", "values": [1, 2, 3]},
                                                   {"name": "t", "values": [1, 2, 3]}])) == []
    assert any("mixes rises and falls" in p for p in charts_spec.validate(
        dict(base, series=[{"values": [10, 20, 30]}, {"values": [1, -2, 3]}], options={"totals": [0]})))
    assert any("mixes category indices" in p for p in charts_spec.validate(
        dict(base, options={"totals": [0, True]})))


def test_dashes_markers_and_series_keys_are_validated():
    spec = {"type": "line", "categories": ["a", "b"],
            "series": [{"name": "s", "values": [1, 2]}]}
    assert charts_spec.validate(dict(spec, options={"dash": "sysDash"})) == []
    assert any("must be one of" in p for p in charts_spec.validate(
        dict(spec, options={"dash": "wiggly"})))
    assert any("dash" in p for p in charts_spec.validate(dict(spec, options={"dash": [0, -1]})))
    assert any("unknown markers" in p for p in charts_spec.validate(
        dict(spec, options={"markers": "asterisk"})))
    assert any("unknown series[0] key 'valeus'" in p for p in charts_spec.validate(
        {"type": "line", "categories": ["a"], "series": [{"name": "s", "valeus": [1]}]}))
    assert any("plotArea runs outside" in p for p in charts_spec.validate(
        dict(spec, options={"plotArea": {"x": 0.5, "y": 0, "w": 0.8, "h": 0.5}})))
    assert any("round 2" in p for p in charts_spec.validate(
        dict(spec, options={"secondaryValueAxis": {"min": 0}})))


# ================================================================ WP-C: one model, split stacks
#
# docs/archive/engine/fidelity/12-WPC-charts.md §3–§5, §8.4. The emitter now reads `chart_model.normalise`
# and nothing else: waterfalls are cut at zero into split stacks with blank cells, a bar that
# crosses zero carries a frozen label, a stacked waterfall's legend lists its authored series,
# negative bars are filled and their category labels sit low, `c:manualLayout` is written only for
# a pinned chart, and a spec the model cannot draw leaves a labelled placeholder — never half a chart.

from pptx.oxml.ns import qn  # noqa: E402

from app.engine import chart_model  # noqa: E402
from app.engine.emit import pptx as emitter_module  # noqa: E402
from app.engine.ir import Box, Canvas, Element, Fonts, Slide  # noqa: E402
from app.engine.reports import EmitOptions  # noqa: E402
from app.engine.verify.coverage import chart_checks, coverage  # noqa: E402

MINUS = chart_model.MINUS
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

#: p05-D (probe `p05_waterfall`): a bridge that crosses zero.
P05_D = {"type": "waterfall", "categories": ["FY23", "Price", "Volume", "Mix", "Cost out", "FY24"],
         "series": [{"name": "EBITDA", "values": [20, -35, 10, -8, 25, 12]}],
         "colors": ["#1E9E5A", "#D64545", "#0B2545"],
         "options": {"dataLabels": True, "numberFormat": "0", "gridlines": False, "legend": False,
                     "totals": [0, 5]}}

#: WP-C §3.2's worked example (fixture J).
STACKED_CROSSING = {"type": "waterfall", "categories": ["FY23", "Q1", "FY24"],
                    "series": [{"name": "Internal", "values": [30, -25, -10]},
                               {"name": "Partner", "values": [0, -15, 0]}],
                    "colors": ["#0B2545", "#1A9AFA"],
                    "options": {"numberFormat": "0", "gridlines": False, "totals": [0, 2]}}


def _labels_of(series_element) -> dict[int, object]:
    """`{point: "delete" | ("text", t) | ("fmt", code) | ("plain",)}` for a series' `c:dLbl`s."""
    found: dict[int, object] = {}
    for label in _findall(series_element, "c:dLbls/c:dLbl"):
        point = int(_val(label, "c:idx"))
        if _find(label, "c:delete") is not None:
            found[point] = "delete"
        elif _find(label, "c:tx/c:rich") is not None:
            found[point] = ("text", "".join(t.text or "" for t in _findall(label, "c:tx/c:rich/a:p/a:r/a:t")))
        elif _find(label, "c:numFmt") is not None:
            found[point] = ("fmt", _find(label, "c:numFmt").get("formatCode"))
        else:
            found[point] = ("plain",)
    return found


def _placeholder_deck(test_master, test_manifest, spec, tmp_path, *, box=BOX, name="Bars"):
    presentation, layout = _deck(test_master, test_manifest)
    slide = presentation.slides.add_slide(layout)
    diagnostics: list[dict] = []
    shape = charts.add_chart(slide, spec, box, dict(STYLE, diagnostics=diagnostics, name=name))
    path = tmp_path / f"placeholder-{name}.pptx"
    presentation.save(str(path))
    return shape, diagnostics, path


def test_crossing_zero_waterfall_emits_split_stacks(one_chart):
    chart, _shapes, diagnostics, _path = one_chart(P05_D, name="Bridge")
    assert [d for d in diagnostics if d["level"] != "info"] == []
    assert bar_series_names(chart) == ["Base", "EBITDA", f"Base{MINUS}", f"EBITDA{MINUS}"]
    workbook = embedded_workbook(chart)
    price = {s["name"]: s["values"][1] for s in workbook["series"]}
    assert price == {"Base": 0.0, "EBITDA": 20.0, f"Base{MINUS}": 0.0, f"EBITDA{MINUS}": -15.0}
    series = _findall(_chart_xml(chart), ".//c:ser")
    # blanks where §3.2 says: no c:pt at all, so PowerPoint draws and labels nothing there
    ebitda_points = [p.get("idx") for p in _findall(series[1], "c:val/c:numRef/c:numCache/c:pt")]
    assert ebitda_points == ["0", "1", "4", "5"], "Volume and Mix lie wholly below zero"
    minus_points = [p.get("idx") for p in _findall(series[3], "c:val/c:numRef/c:numCache/c:pt")]
    assert minus_points == ["1", "2", "3", "4"], "FY23 and FY24 lie wholly above zero"
    assert _labels_of(series[3])[1] == ("text", "-35"), "the part holding the fall's end reads the step"
    assert _labels_of(series[1])[1] == "delete", "the other part of a crossing bar has no label"
    assert _labels_of(series[1])[4] == ("text", "+25")
    assert _labels_of(series[3])[4] == "delete"
    assert _labels_of(series[3])[2] == ("fmt", '"+"0;"+"0'), "a rise below zero reads +10 natively"
    for base in (series[0], series[2]):
        assert _find(base, "c:spPr/a:noFill") is not None
    assert primary_value_axis(chart).minimum_scale == -20.0 and primary_value_axis(chart).maximum_scale == 20.0


def test_non_negative_waterfall_still_emits_two_series(one_chart):
    """WP3a's two-series workbook, byte for byte, for every waterfall that never dips below zero."""
    from pptx.chart.data import CategoryChartData

    spec = next(s for s in _fixture_specs() if s["type"] == "waterfall")
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    assert bar_series_names(chart) == ["Base", "Partners"]
    # WP3a's formula, restated: base = the floor a bar stands on, delta = its height.
    values = spec["series"][0]["values"]
    totals = set(spec["options"]["totals"])
    base, delta, running = [], [], 0.0
    for index, value in enumerate(values):
        if value is None:
            base.append(None)
            delta.append(None)
            continue
        if index in totals:
            floor, height, running = 0.0, float(value), float(value)
        elif value >= 0:
            floor, height, running = running, float(value), running + value
        else:
            floor, height, running = running + value, float(-value), running + value
        base.append(floor)
        delta.append(height)
    expected = CategoryChartData()
    expected.categories = [str(c) for c in spec["categories"]]
    expected.add_series("Base", base, spec["options"]["numberFormat"])
    expected.add_series("Partners", delta, spec["options"]["numberFormat"])
    got = emitter_module._fix_embedded_workbook(chart.part.chart_workbook.xlsx_part.blob)
    assert got == emitter_module._fix_embedded_workbook(expected.xlsx_blob)


def test_totals_spellings_emit_identical_xml(test_master, test_manifest):
    base = {"type": "waterfall", "categories": ["Standalone", "Cost syn.", "Revenue syn.", "Subtotal",
                                                "Integration", "Tax", "Combined"],
            "colors": ["#1E9E5A", "#D64545", "#0B2545"],
            "options": {"dataLabels": True, "numberFormat": "0.0", "gridlines": False, "legend": False}}
    values = [8.1, 2.4, 1.6, 12.1, -0.9, 1.2, 12.4]
    spellings = [
        dict(base, series=[{"name": "Value", "values": values}],
             options=dict(base["options"], totals=[0, 3, 6])),
        dict(base, series=[{"name": "Value", "values": values, "types": [
            "total", "increase", "increase", "total", "decrease", "increase", "total"]}]),
        dict(base, series=[{"name": "Value", "values": values}],
             options=dict(base["options"], totals=[True, False, False, True, False, False, True])),
    ]

    def xml(spec) -> bytes:
        presentation, layout = _deck(test_master, test_manifest)
        slide = presentation.slides.add_slide(layout)
        charts.add_chart(slide, spec, BOX, dict(STYLE, name="Chart"))
        return etree.tostring(slide.shapes[0].chart._chartSpace)

    first, second, third = (xml(spec) for spec in spellings)
    assert first == second == third


def test_stacked_waterfall_keeps_its_series_and_deletes_extra_legend_entries(one_chart):
    stacked = {"type": "waterfall", "categories": ["FY23", "Price", "Volume", "Mix", "FY24"],
               "series": [{"name": "Internal", "values": [36, 9, 14, -11, 48]},
                          {"name": "Partner", "values": [13, 5, 6, -4, 20]}],
               "colors": ["#0B2545", "#1A9AFA"], "options": {"numberFormat": "0", "totals": [0, 4]}}
    chart, _shapes, diagnostics, _path = one_chart(stacked)
    assert [d for d in diagnostics if d["level"] != "info"] == []
    assert bar_series_names(chart) == ["Base", "Internal", "Partner"]
    legend = _find(_chart_xml(chart), ".//c:legend")
    assert _val(legend, "c:legendPos") == "b", "several authored series: a legend by default"
    deleted = [int(_val(entry, "c:idx")) for entry in _findall(legend, "c:legendEntry")
               if _val(entry, "c:delete") == "1"]
    assert deleted == [0, 3, 4, 5, 6], "neither the invisible base nor the four connectors are legend entries"
    children = [etree.QName(child).localname for child in legend]
    assert children.index("legendEntry") == children.index("legendPos") + 1, "CT_Legend order"
    series = _findall(_chart_xml(chart), ".//c:ser")
    assert _val(series[1], "c:spPr/a:solidFill/a:srgbClr") == "0B2545"
    assert _val(series[2], "c:spPr/a:solidFill/a:srgbClr") == "1A9AFA"

    crossing, _shapes, _diagnostics, _path = one_chart(STACKED_CROSSING)
    legend = _find(_chart_xml(crossing), ".//c:legend")
    assert [int(_val(e, "c:idx")) for e in _findall(legend, "c:legendEntry")] == [0, 3, 4, 5, 6, 7]
    series = _findall(_chart_xml(crossing), ".//c:ser")
    assert _val(series[4], "c:spPr/a:solidFill/a:srgbClr") == "0B2545", "Internal− is Internal's colour"


def test_stacked_crossing_bar_matches_the_plan(one_chart):
    chart, _shapes, _diagnostics, _path = one_chart(STACKED_CROSSING)
    plan = chart_model.normalise(STACKED_CROSSING)["waterfall"]
    workbook = embedded_workbook(chart)
    assert workbook["series"] == [{"name": s["name"], "values": s["values"]} for s in plan["series"]]
    series = _findall(_chart_xml(chart), ".//c:ser")
    assert _labels_of(series[4])[1] == ("text", "-25"), "Internal's crossing segment, frozen on Internal−"
    assert _labels_of(series[1])[1] == "delete"
    assert _labels_of(series[2]).get(1) == ("fmt", '"-"0;"-"0'), "Partner reads -15 natively"
    assert charts.expected_workbook(STACKED_CROSSING)["series"] == workbook["series"]


def test_frozen_label_xml_is_in_schema_order(one_chart):
    """`CT_DLbl`: idx, tx, spPr, dLblPos, show* — with c:tx there is no numFmt and no txPr."""
    chart, _shapes, _diagnostics, _path = one_chart(P05_D)
    all_series = _findall(_chart_xml(chart), ".//c:ser")
    minus = all_series[3]
    frozen = next(label for label in _findall(minus, "c:dLbls/c:dLbl") if _val(label, "c:idx") == "1")
    assert [etree.QName(child).localname for child in frozen] == [
        "idx", "tx", "spPr", "dLblPos", "showLegendKey", "showVal", "showCatName", "showSerName",
        "showPercent", "showBubbleSize"]
    rich = _find(frozen, "c:tx/c:rich")
    assert [etree.QName(child).localname for child in rich] == ["bodyPr", "lstStyle", "p"]
    run_properties = _find(rich, "a:p/a:r/a:rPr")
    assert run_properties.get("lang") == "en-US"
    assert run_properties.get("sz") == str(int(round(STYLE["sizePx"] * 0.75 * 0.95 * 100)))
    assert [etree.QName(child).localname for child in run_properties] == ["solidFill", "latin"]
    holder = _find(minus, "c:dLbls")
    tags = [etree.QName(child).localname for child in holder]
    assert tags[: tags.count("dLbl")] == ["dLbl"] * tags.count("dLbl"), "every dLbl before the group"
    deleted = next(label for label in _findall(all_series[1], "c:dLbls/c:dLbl") if _val(label, "c:idx") == "1")
    assert [etree.QName(child).localname for child in deleted] == ["idx", "delete"]


def test_point_colors_list_colours_every_visible_series_carrying_the_point(one_chart):
    spec = dict(P05_D, options=dict(P05_D["options"], pointColors=[
        "#C4C4CD", "#7FC6FF", "#C4C4CD", "#C4C4CD", "#7FC6FF", "#C4C4CD"]))
    chart, _shapes, diagnostics, _path = one_chart(spec)
    assert any("pointColors is a list" in d["message"] for d in diagnostics)
    series = _findall(_chart_xml(chart), ".//c:ser")

    def fills(element):
        return {int(_val(p, "c:idx")): _val(p, "c:spPr/a:solidFill/a:srgbClr") for p in _findall(element, "c:dPt")}

    assert fills(series[1])[1] == fills(series[3])[1] == "7FC6FF", "Price on both sides of zero"
    assert fills(series[1])[4] == fills(series[3])[4] == "7FC6FF"
    assert fills(series[3])[2] == "C4C4CD"


def test_bar_series_never_invert_and_labels_sit_low(one_chart):
    negative = {"type": "bar", "categories": ["Bear case", "Base case", "Bull case"],
                "series": [{"name": "FY2 EPS impact", "values": [-2.0, 2.3, 6.2]}], "colors": ["#0B3D2E"],
                "options": {"dataLabels": True, "numberFormat": '0.0"%"', "legend": False}}
    chart, _shapes, _diagnostics, _path = one_chart(negative)
    for series in _findall(_chart_xml(chart), ".//c:ser"):
        assert _val(series, "c:invertIfNegative") == "0"
        children = [etree.QName(child).localname for child in series]
        assert children.index("invertIfNegative") == children.index("spPr") + 1, "CT_BarSer order"
    assert _val(chart.category_axis._element, "c:tickLblPos") == "low"
    positive, _shapes, _diagnostics, _path = one_chart(spec_for("column"))
    assert _val(positive.category_axis._element, "c:tickLblPos") == "nextTo", "no negatives, no move"
    for series in _findall(_chart_xml(positive), ".//c:ser"):
        assert _val(series, "c:invertIfNegative") == "0"
    waterfall, _shapes, _diagnostics, _path = one_chart(P05_D)
    assert _val(waterfall.category_axis._element, "c:tickLblPos") == "low"


def test_unpinned_chart_has_no_manual_layout(one_chart):
    for spec in (spec_for("column"), spec_for("line"), spec_for("doughnut")):
        chart, _shapes, _diagnostics, _path = one_chart(spec)
        assert _find(_chart_xml(chart), ".//c:plotArea/c:layout/c:manualLayout") is None, spec["type"]


def test_overlay_chart_pins_the_layout_rect(one_chart):
    """`outEnd` labels are text boxes on top of the frame, so the chart pins its plot rect; its
    connectors are a series of the chart (C3's fallback) and pin nothing themselves."""
    spec = {"type": "waterfall", "categories": ["a", "b", "c"],
            "series": [{"name": "s", "values": [100, 40, 140]}],
            "options": {"gapWidth": 50, "labelPosition": "outEnd", "numberFormat": "0"}}
    box = (0.0, 0.0, 4.0, 2.5)
    chart, shapes, _diagnostics, path = one_chart(spec, box=box, name="Bridge")
    geometry = chart_model.layout(chart_model.normalise(spec), 384.0, 240.0, STYLE["sizePx"])
    assert geometry["pinned"]
    manual = _find(_chart_xml(chart), ".//c:plotArea/c:layout/c:manualLayout")
    assert {tag: float(_val(manual, f"c:{tag}")) for tag in "xywh"} == geometry["plotArea"]
    assert shapes == ["Bridge", "Bridge label 1", "Bridge label 2", "Bridge label 3"]
    label = Presentation(str(path)).slides[0].shapes[1]
    area = geometry["plotArea"]
    slot = area["w"] * 4.0 / 3
    centre = area["x"] * 4.0 + slot / 2
    assert label.left + label.width / 2 == pytest.approx(centre * config.EMU_PER_IN, abs=2)
    assert [xs for xs, _ys in connector_series(chart)] == [pytest.approx([0.5 + 1 / 3, 1.5 - 1 / 3]),
                                                           pytest.approx([1.5 + 1 / 3, 2.5 - 1 / 3])]
    body = _find(_chart_xml(chart), ".//c:catAx/c:txPr/a:bodyPr")
    assert (body.get("rot"), body.get("vert"), body.get("wrap")) == ("0", "horz", "none"), \
        "a pinned chart's category labels stay horizontal and unwrapped"
    unpinned, _shapes, _diagnostics, _path = one_chart(dict(spec, options={"gapWidth": 50}))
    assert _find(_chart_xml(unpinned), ".//c:plotArea/c:layout/c:manualLayout") is None, \
        "connectors alone pin nothing: PowerPoint lays the bridge out"


def test_reference_line_chart_pins_the_layout_rect(one_chart):
    spec = dict(spec_for("column"), options={"valueAxis": {"min": 0, "max": 10},
                                              "referenceLines": [{"value": 5, "label": "target"}]})
    box = (0.0, 0.0, 4.0, 2.0)
    chart, shapes, diagnostics, path = one_chart(spec, box=box, name="Bars")
    assert [d for d in diagnostics if d["level"] != "info"] == []
    geometry = chart_model.layout(chart_model.normalise(spec), 384.0, 192.0, STYLE["sizePx"])
    manual = _find(_chart_xml(chart), ".//c:plotArea/c:layout/c:manualLayout")
    assert {tag: float(_val(manual, f"c:{tag}")) for tag in "xywh"} == geometry["plotArea"]
    line = Presentation(str(path)).slides[0].shapes[1]
    area = geometry["plotArea"]
    assert line.top == pytest.approx((area["y"] + area["h"] / 2) * 2.0 * config.EMU_PER_IN, abs=2)
    assert line.left == pytest.approx(area["x"] * 4.0 * config.EMU_PER_IN, abs=2)
    assert shapes == ["Bars", "Bars referenceLine 1", "Bars referenceLine 1 label"]


def test_waterfall_out_end_labels_become_text_boxes(one_chart):
    spec = {"type": "waterfall", "categories": ["FY23", "Up", "Down", "", "FY24"],
            "series": [{"name": "s", "values": [100, 40, -15, None, 125]}],
            "options": {"labelPosition": "outEnd", "numberFormat": "0", "totals": [0, 4]}}
    chart, _shapes, diagnostics, path = one_chart(spec, name="Bridge")
    assert [d for d in diagnostics if d["level"] != "info"] == []
    for series in _findall(_chart_xml(chart), ".//c:ser")[1:]:
        assert _val(series, "c:dLbls/c:delete") == "1", "no native label next to a text box"
    shapes = list(Presentation(str(path)).slides[0].shapes)
    labels = [s for s in shapes if " label " in s.name]
    assert [s.name for s in labels] == ["Bridge label 1", "Bridge label 2", "Bridge label 3", "Bridge label 5"]
    assert [s.text_frame.text for s in labels] == ["100", "+40", "-15", "125"]
    geometry = chart_model.layout(chart_model.normalise(spec), BOX[2] * 96.0, BOX[3] * 96.0, STYLE["sizePx"])
    axis = chart_model.normalise(spec)["axis"]
    area = geometry["plotArea"]
    level_140 = (BOX[1] + (area["y"] + area["h"] * (1 - (140 - axis["min"]) / (axis["max"] - axis["min"])))
                 * BOX[3]) * config.EMU_PER_IN
    up, down = labels[1], labels[2]
    assert up.top + up.height <= level_140, "a rise's label sits above its bar"
    assert down.top >= level_140, "a fall's label sits below its bar"


def test_unknown_type_leaves_a_placeholder_not_a_chart(test_master, test_manifest, tmp_path):
    spec = {"type": "grouped_bar", "categories": ["Strategy", "Talent"],
            "series": [{"name": "UAE", "values": [72.4, 68]}]}
    shape, diagnostics, path = _placeholder_deck(test_master, test_manifest, spec, tmp_path)
    assert not shape.has_chart
    assert shape.name == "Bars (chart not built)"
    message = "unknown chart type 'grouped_bar' (did you mean 'bar'?)"
    assert shape.text_frame.text == message
    assert [d["level"] for d in diagnostics] == ["error"] and diagnostics[0]["message"] == message
    element = shape._element
    assert element.find(qn("p:style")) is None, "no theme fill"
    shape_properties = element.find(f"{{{P}}}spPr")
    assert _find(shape_properties, "a:noFill") is not None
    line = _find(shape_properties, "a:ln")
    assert _val(line, "a:prstDash") == "dash"
    assert _find(line, "a:solidFill/a:srgbClr").get("val") == "E5484D"
    with zipfile.ZipFile(path) as archive:
        assert not [n for n in archive.namelist() if n.startswith(("ppt/charts/", "ppt/embeddings/"))]
    tiny, _diagnostics, _path = _placeholder_deck(test_master, test_manifest, spec, tmp_path,
                                                  box=(0.5, 0.5, 1.0, 0.3), name="Tiny")
    assert tiny.text_frame.text == "" and tiny.name == f"Tiny (chart not built: {message})"


def test_radar_and_combo_leave_the_same_placeholder(test_master, test_manifest, tmp_path):
    for kind, words in (("radar", "not a chart object on Path A"), ("combo", "round 2")):
        spec = {"type": kind, "categories": ["a", "b", "c"], "series": [{"name": "s", "values": [1, 2, 3]}]}
        shape, diagnostics, _path = _placeholder_deck(test_master, test_manifest, spec, tmp_path, name=kind)
        assert not shape.has_chart and shape.name == f"{kind} (chart not built)", kind
        assert words in shape.text_frame.text and diagnostics[-1]["level"] == "error", kind


def _chart_ir(spec, *, name="Bridge") -> IR:
    ir = IR(canvas=Canvas(1280, 720), fonts=Fonts(forced={"major": "Calibri", "minor": "Calibri"}, used=["Calibri"]),
            slide=Slide(id="s1", title="t", layoutId="layout-07"),
            elements=[Element(kind="chart", box=Box(96, 96, 576, 300), name=name, spec=spec,
                              style=dict(STYLE), origin="authored", confidence=1.0, overlay=[],
                              source={"path": f"body > [data-name={name}]", "tag": "div", "svg": None})])
    ir.elements[0].id, ir.elements[0].z = "e1", 0
    return ir


def test_a_failure_mid_styling_rolls_the_chart_back(test_master, test_manifest, tmp_path, monkeypatch):
    def broken(*_args, **_kwargs):
        raise RuntimeError("styling broke")

    monkeypatch.setattr(charts, "_color_points", broken)
    ir = _chart_ir(P05_D)
    decks = []
    for attempt in ("a", "b"):
        deck = tmp_path / f"rolled-back-{attempt}.pptx"
        report = emitter_module.emit([ir], test_manifest, test_master, deck, EmitOptions())
        decks.append(deck)
    assert not any("failed to emit" in w for w in report.warnings), "the rollback, not the outer catch"
    assert report.shapes_by_kind == {"shape": 1}
    [entry] = report.charts
    assert entry["failed"] is True and entry["error"] == "chart could not be built: RuntimeError: styling broke"
    shapes = list(Presentation(str(decks[0])).slides[0].shapes)
    assert [s.name for s in shapes] == ["Bridge (chart not built)"], "no frame, no connector left behind"
    with zipfile.ZipFile(decks[0]) as archive:
        names = archive.namelist()
    assert not [n for n in names if n.startswith(("ppt/charts/", "ppt/embeddings/"))], names
    assert decks[0].read_bytes() == decks[1].read_bytes()
    cover = coverage([ir], report, [], [], pptx=decks[0])
    [check] = cover.charts
    assert check["failed"] and check["problems"][0] == f"chart not built: {entry['error']}"


def test_a_refused_type_is_named_in_the_emit_report(test_master, test_manifest, tmp_path):
    ir = _chart_ir({"type": "radar", "categories": ["a"], "series": [{"values": [1]}]}, name="Spider")
    report = emitter_module.emit([ir], test_manifest, test_master, tmp_path / "radar.pptx", EmitOptions())
    [entry] = report.charts
    assert entry["failed"] and "not a chart object on Path A" in entry["error"]
    assert report.shapes_by_kind == {"shape": 1}
    assert any("not a chart object on Path A" in w for w in report.warnings)


# ------------------------------------------------------------- WP-C torture families, end to end


def _family_specs(family: str) -> list[dict]:
    html = (config.TORTURE_FIXTURE / f"{family}.html").read_text(encoding="utf-8")
    return [json.loads(match) for match in re.findall(r"data-chart='(.*?)'", html, re.S)]


@pytest.mark.parametrize("family", ["charts-hardening", "charts-negative"])
def test_the_wpc_families_and_their_irs_hold_the_same_specs(family: str):
    ir = IR.load(config.TORTURE_FIXTURE / "ir" / f"{family}.json")
    assert ir.validate() == []
    assert [e.spec for e in ir.of_kind("chart")] == _family_specs(family), f"{family} drifted from its IR"


@pytest.mark.renderer_full
@pytest.mark.validator
@pytest.mark.parametrize("family", ["charts-hardening", "charts-negative"])
def test_charts_hardening_and_negative_families_verify_end_to_end(family: str, test_master, test_manifest,
                                                                   tmp_path):
    """Every spec a native chart part with the expected workbook, no renderer warning, the
    validator passes, every structural check is ok, and two emits are byte-identical."""
    ir = IR.load(config.TORTURE_FIXTURE / "ir" / f"{family}.json")
    first, second = tmp_path / "a.pptx", tmp_path / "b.pptx"
    report = emitter_module.emit([ir], test_manifest, test_master, first, EmitOptions())
    emitter_module.emit([ir], test_manifest, test_master, second, EmitOptions())
    assert first.read_bytes() == second.read_bytes()
    expected = json.loads((config.TORTURE_FIXTURE / f"{family}.expect.json").read_text(encoding="utf-8"))
    assert report.shapes_by_kind == expected["kinds"], report.shapes_by_kind
    assert report.warnings == []
    assert not any(entry.get("failed") for entry in report.charts)

    frames = [s for s in Presentation(str(first)).slides[0].shapes if s.has_chart]
    assert [f.name for f in frames] == [e.name for e in ir.of_kind("chart")]
    for frame, element in zip(frames, ir.of_kind("chart"), strict=False):
        wanted = charts.expected_workbook(element.spec)
        got = embedded_workbook(frame.chart)
        assert [s["name"] for s in got["series"]] == [s["name"] for s in wanted["series"]], element.name
        assert [s["values"] for s in got["series"]] == [s["values"] for s in wanted["series"]], element.name

    checks = chart_checks(first, [ir])
    assert [c["element"] for c in checks if not c["ok"]] == [], [c["problems"] for c in checks]

    out = tmp_path / "render"
    renderer.render(first, 1280, out)
    assert json.loads((out / "warnings.json").read_text(encoding="utf-8"))["warnings"] == []
    result = subprocess.run(
        [sys.executable, str(config.VALIDATE_PY), str(first), "--original", str(test_master)],
        capture_output=True, text=True, timeout=300,
    )
    assert result.returncode == 0 and "PASSED" in result.stdout, result.stdout + result.stderr


def test_probe_p08_reports_the_unknown_type_and_the_strings(tmp_path):
    """Probe p08 end to end: grouped_bar and radar are named failures, D's strings are noted.

    The probe is the `chart-specs` torture family since G-2 deleted `probes/` and `run_probes.py`
    (plan §16 #16): extracted, classified and emitted here as the torture runner does."""
    from app.engine.classify.charts import recognise_charts
    from app.engine.classify.placeholders import map_placeholders
    from app.engine.emit import pptx as emitter
    from app.engine.reports import EmitOptions
    from app.engine.verify.fixtures import context_for
    from tests.engine.helpers import extract_html

    html = config.TORTURE_FIXTURE / "chart-specs.html"
    context = context_for(html)
    ir = extract_html(html, context.manifest, context.layout_id, context.assets_dir,
                      slide_id=html.stem, title=html.stem)
    ir = map_placeholders(recognise_charts(ir), context.layout)
    report = emitter.emit([ir], context.manifest, context.master, tmp_path / "chart-specs.pptx", EmitOptions())
    emitted = json.loads(json.dumps({"warnings": report.warnings, "charts": report.charts}, default=str))
    assert not [w for w in emitted["warnings"] if "failed to emit" in w]
    by_type = {entry["type"]: entry for entry in emitted["charts"]}
    assert by_type["grouped_bar"]["failed"] and "did you mean 'bar'" in by_type["grouped_bar"]["error"]
    assert by_type["radar"]["failed"] and "not a chart object on Path A" in by_type["radar"]["error"]
    notes = [n for entry in emitted["charts"] for n in entry.get("notes", []) if "values[" in n]
    assert notes == [f'series[0].values[{i}]: "{v}" read as {v}' for i, v in enumerate(["4.2", "4.9", "5.6", "6.2"])]
    assert not by_type["column_clustered"].get("failed")


@pytest.mark.validator
def test_no_spec_leaves_a_half_built_chart(test_master, test_manifest, tmp_path):
    """Malformed specs: a chart part or a labelled placeholder, never an exception or half a chart."""
    from app.engine.tests.test_chart_model import MALFORMED

    presentation, layout = _deck(test_master, test_manifest)
    for spec in MALFORMED:
        slide = presentation.slides.add_slide(layout)
        diagnostics: list[dict] = []
        shape = charts.add_chart(slide, spec, BOX, dict(STYLE, diagnostics=diagnostics, name="C"))
        assert not any("could not be built" in d["message"] for d in diagnostics), (spec, diagnostics)
        if shape.has_chart:
            gap = _val(_chart_xml(shape.chart), ".//c:barChart/c:gapWidth")
            assert gap is None or 0 <= int(gap) <= 500
        else:
            assert shape.name.startswith("C (chart not built")
    path = tmp_path / "malformed.pptx"
    presentation.save(str(path))
    result = subprocess.run(
        [sys.executable, str(config.VALIDATE_PY), str(path), "--original", str(test_master)],
        capture_output=True, text=True, timeout=300,
    )
    assert result.returncode == 0 and "PASSED" in result.stdout, result.stdout + result.stderr


# --------------------------------------------------------- the C1 review's findings, emitter side


#: The review's minimal specs (out/review-c/cases.json), one per finding they reproduce.
K_BELOW = {"type": "waterfall", "categories": ["FY23", "Opex", "FY24"],
           "series": [{"name": "EBIT", "values": [-20, -10, -30]}], "options": {"numberFormat": "0", "totals": [0, 2]}}
NEGATIVE_POINT = {"type": "column", "categories": ["A", "B", "C"], "series": [{"name": "Delta", "values": [4, -3, 2]}],
                  "colors": ["#1A9AFA"], "options": {"pointColors": {"0": {"1": "#E5484D"}}, "numberFormat": "0"}}
FLAT_BELOW = {"type": "waterfall", "categories": ["Start", "Flat", "Up", "End"],
              "series": [{"name": "Cash", "values": [-20, 0, 5, -15]}], "options": {"numberFormat": "0", "totals": [0, 3]}}


def _chart_schema():
    """`dml-chart.xsd` from the pptx skill's validator — which itself skips chart parts, on purpose
    (`validators/pptx.py: _get_schema_path` returns None under `charts/`), so a chart part written out
    of schema order passes "validator PASSED" (C1 review #3)."""
    path = config.VALIDATE_PY.parent / "schemas" / "ISO-IEC29500-4_2016" / "dml-chart.xsd"
    assert path.exists(), f"{path} is missing — the pptx skill's schemas are a test input"
    return etree.XMLSchema(etree.parse(str(path)))


def chart_schema_errors(deck: Path, schema=None) -> dict[str, list[str]]:
    """`{part: [message]}` for every chart part of `deck` that `dml-chart.xsd` rejects.

    python-pptx's templates write negative `c:axId`/`c:crossAx` ids, which fail `xs:unsignedInt`
    and which PowerPoint opens without a word; those two messages are the only ones filtered.
    """
    schema = schema or _chart_schema()
    found: dict[str, list[str]] = {}
    with zipfile.ZipFile(deck) as archive:
        for name in sorted(n for n in archive.namelist() if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)):
            schema.validate(etree.fromstring(archive.read(name)))
            errors = [e.message for e in schema.error_log
                      if not (("}axId'" in e.message or "}crossAx'" in e.message) and "xs:unsignedInt" in e.message)]
            if errors:
                found[name] = errors
    return found


@pytest.mark.validator
def test_every_fixture_chart_part_is_schema_valid(every_type_deck, test_master, test_manifest, tmp_path):
    """Every chart part the fixtures emit, against the chart schema the validator never applies:
    the `charts`, `charts-hardening`, `charts-negative`, `emit-charts` and `sizes-4x3` IRs, one slide
    per Path A type (3-D and of-pie included), the review's cases and the malformed sweep."""
    from app.engine.tests.test_chart_model import MALFORMED

    schema = _chart_schema()
    decks = [every_type_deck]
    for family in ("charts", "charts-hardening", "charts-negative", "sizes-4x3"):
        ir = IR.load(config.TORTURE_FIXTURE / "ir" / f"{family}.json")
        master = config.MASTERS_FIXTURE / ("test-4x3.pptx" if family == "sizes-4x3" else "test-16x9.pptx")
        manifest = Manifest.load(master.with_suffix(".manifest.json"))
        deck = tmp_path / f"{family}.pptx"
        emitter_module.emit([ir], manifest, master, deck, EmitOptions())
        decks.append(deck)
    deck = tmp_path / "emit-charts.pptx"
    emitter_module.emit([IR.load(config.TORTURE_FIXTURE / "ir" / "emit-charts.json")], test_manifest,
                        test_master, deck, EmitOptions())
    decks.append(deck)
    presentation, layout = _deck(test_master, test_manifest)
    extra = [P05_D, STACKED_CROSSING, K_BELOW, NEGATIVE_POINT, FLAT_BELOW,
             dict(STACKED_CROSSING, options=dict(STACKED_CROSSING["options"], legend="right")),
             dict(P05_D, options=dict(P05_D["options"], categoryAxis={"reverse": True}, legend="bottom")),
             {"type": "pie_exploded", "categories": ["a", "b", "c"], "series": [{"values": [5, 3, 2]}],
              "options": {"explosion": 8}}] + MALFORMED
    for spec in extra:
        charts.add_chart(presentation.slides.add_slide(layout), spec, BOX, dict(STYLE, name="C"))
    cases = tmp_path / "cases.pptx"
    presentation.save(str(cases))
    decks.append(cases)
    bad = {deck.name: chart_schema_errors(deck, schema) for deck in decks}
    assert {name: errors for name, errors in bad.items() if errors} == {}


def test_bar_points_with_a_fill_never_invert(one_chart):
    """C1 review #1: every `c:dPt` of a bar series says `invertIfNegative 0`, straight after its idx —
    without it PowerPoint draws that point hollow when it is negative, whatever the series says."""
    for spec in (P05_D, K_BELOW, NEGATIVE_POINT):
        chart, _shapes, _diagnostics, _path = one_chart(spec)
        points = _findall(_chart_xml(chart), ".//c:barChart/c:ser/c:dPt")
        assert points, spec["type"]
        for point in points:
            children = [etree.QName(child).localname for child in point]
            assert children[:2] == ["idx", "invertIfNegative"], children
            assert _val(point, "c:invertIfNegative") == "0"
    pie, _shapes, _diagnostics, _path = one_chart(spec_for("pie"))
    assert _findall(_chart_xml(pie), ".//c:dPt/c:invertIfNegative") == [], "a pie point has no such element"


def test_exploded_pie_writes_one_explosion(one_chart):
    """C1 review #2: python-pptx's exploded template already has `c:explosion`; a second one is a
    schema error PowerPoint answers with its repair prompt."""
    spec = {"type": "pie_exploded", "categories": ["a", "b", "c"], "series": [{"values": [5, 3, 2]}],
            "options": {"explosion": 8}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    explosions = _findall(_chart_xml(chart), ".//c:ser/c:explosion")
    assert [e.get("val") for e in explosions] == ["8"]


@pytest.mark.parametrize(("legend", "position", "deleted"), [
    ("bottom", "b", [0, 3, 4, 5, 6, 7]), ("top", "t", [0, 3, 4, 5, 6, 7]),
    ("right", None, [0, 1, 2, 5, 6, 7]), ("left", "l", [0, 1, 2, 5, 6, 7]), ("topRight", "tr", [0, 1, 2, 5, 6, 7]),
])
def test_legend_entries_count_as_the_legend_displays(one_chart, legend, position, deleted):
    """C1 review #4: `c:legendEntry/c:idx` is the entry's place in the legend as drawn, and a vertical
    legend lists a stacked column top of the stack first and the connectors' scatter group after it
    (read in PowerPoint, `out/c3/ppt3/legends.json`). Series order: Base, Internal, Partner, Base−,
    Internal−, Partner−, Connector 1, Connector 2; the kept entries are Internal and Partner."""
    spec = dict(STACKED_CROSSING, options=dict(STACKED_CROSSING["options"], legend=legend))
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    node = _find(_chart_xml(chart), ".//c:legend")
    assert _val(node, "c:legendPos") == position
    assert [int(_val(e, "c:idx")) for e in _findall(node, "c:legendEntry")] == deleted


@pytest.mark.validator
def test_a_spec_that_validates_builds_a_chart(test_master, test_manifest, tmp_path):
    """C1 review #6, the converse invariant: `validate(spec) == []` means a chart part with no warn
    or error — the review's table, then 150 validated specs from the agreement fuzz (fixed seed)."""
    import random

    from app.engine.tests.test_chart_model import fuzz_spec

    column = {"type": "column", "categories": ["a", "b"], "series": [{"values": [1, 2]}]}
    waterfall = {"type": "waterfall", "categories": ["a", "b", "c"], "series": [{"values": [100, -30, 70]}],
                 "options": {"totals": [0, 2]}}
    table = [
        dict(column, options={"valueAxis": {"minorUnit": 0.5}}),
        {"type": "scatter", "series": [{"x": [1, 2], "y": [3, 4]}], "options": {"categoryAxis": {"minorUnit": 1}}},
        dict(column, options={"gridlines": {"width": 1584}}),
        dict(column, options={"valueAxis": {"min": 0, "max": 5}, "referenceLines": [{"value": 2, "width": 0}]}),
        dict(waterfall, options=dict(waterfall["options"], connectors={"width": 0.25, "dash": [1000, 0.5]})),
    ]
    rng = random.Random(7)
    fuzzed: list[dict] = []
    while len(fuzzed) < 150:
        spec = fuzz_spec(rng)
        if charts_spec.validate(spec) == []:
            fuzzed.append(spec)
    presentation, layout = _deck(test_master, test_manifest)
    for spec in table + fuzzed:
        assert charts_spec.validate(spec) == [], spec
        diagnostics: list[dict] = []
        shape = charts.add_chart(presentation.slides.add_slide(layout), spec, BOX,
                                 dict(STYLE, diagnostics=diagnostics, name="T"))
        loud = [d["message"] for d in diagnostics if d["level"] in ("warn", "error")]
        assert shape.has_chart and loud == [], (spec, loud)
    deck = tmp_path / "validated.pptx"
    presentation.save(str(deck))
    assert chart_schema_errors(deck) == {}


def test_a_flat_step_below_zero_is_labelled_at_its_level(one_chart):
    """C1 review #9: no 0 cell below zero (PowerPoint stacks a 0 at the zero line); the `+0` is a
    frozen label at the inside end of `Base−`'s bar — its end is the level — in the chart's ink."""
    chart, _shapes, diagnostics, _path = one_chart(FLAT_BELOW)
    assert [d for d in diagnostics if d["level"] != "info"] == []
    series = _findall(_chart_xml(chart), ".//c:ser")
    assert bar_series_names(chart) == ["Base", "Cash", f"Base{MINUS}", f"Cash{MINUS}"]
    for element in series:
        cells = {p.get("idx"): p.find(f"{{{C}}}v").text for p in _findall(element, "c:val/c:numRef/c:numCache/c:pt")}
        assert cells.get("1") in (None, "-20.0", "-20"), "only Base− carries a cell at the flat step"
    base_minus = series[2]
    label = next(lbl for lbl in _findall(base_minus, "c:dLbls/c:dLbl") if _val(lbl, "c:idx") == "1")
    assert _labels_of(base_minus)[1] == ("text", "+0")
    assert _val(label, "c:dLblPos") == "inEnd"
    assert _val(label, "c:showVal") == "1"
    assert _find(label, "c:tx/c:rich/a:p/a:r/a:rPr/a:solidFill/a:srgbClr").get("val") == "516467", "the chart's ink"
    above, _shapes, _diagnostics, _path = one_chart({"type": "waterfall", "categories": list("abcd"),
                                                     "series": [{"values": [20, 0, -20, 5]}],
                                                     "options": {"numberFormat": "0", "totals": [0]}})
    flat = _findall(_chart_xml(above), ".//c:ser")[1]
    ink = next(lbl for lbl in _findall(flat, "c:dLbls/c:dLbl") if _val(lbl, "c:idx") == "1")
    assert _find(ink, "c:txPr/a:p/a:pPr/a:defRPr/a:solidFill/a:srgbClr").get("val") == "516467", \
        "a zero-height label sits on the slide, not on its bar"


def test_reversed_categories_keep_the_value_axis_where_it_was(one_chart):
    """C1 review #10: with `categoryAxis.reverse` the value axis crosses at the last category, which
    is now on the left, so the tick gutter the pinned rect keeps is still the one used."""
    spec = dict(P05_D, options=dict(P05_D["options"], categoryAxis={"reverse": True}))
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    value_axis = _find(_chart_xml(chart), ".//c:valAx")
    assert _val(value_axis, "c:crosses") == "max"
    assert _val(_find(_chart_xml(chart), ".//c:catAx"), "c:scaling/c:orientation") == "maxMin"
    forward, _shapes, _diagnostics, _path = one_chart(P05_D)
    assert _val(_find(_chart_xml(forward), ".//c:valAx"), "c:crosses") == "autoZero"


def test_placeholder_names_and_messages_are_xml_safe(test_master, test_manifest, tmp_path):
    """C1 review #13: a control character in a type or a name never reaches python-pptx's setters,
    so `_placeholder` cannot raise and `_emit_slide`'s outer catch is never reached."""
    spec = {"type": "bar\x01chart", "categories": ["a"], "series": [{"values": [1]}]}
    shape, diagnostics, _path = _placeholder_deck(test_master, test_manifest, spec, tmp_path,
                                                  box=(0.5, 0.5, 1.0, 0.3), name="Tiny")
    assert not shape.has_chart and "\x01" not in shape.name and shape.name.startswith("Tiny (chart not built")
    valid = {"type": "column", "categories": ["a"], "series": [{"values": [1]}]}
    presentation, layout = _deck(test_master, test_manifest)
    diagnostics: list[dict] = []
    frame = charts.add_chart(presentation.slides.add_slide(layout), valid, BOX,
                             dict(STYLE, diagnostics=diagnostics, name="Revenue\x0bchart"))
    assert frame.has_chart and frame.name == "Revenuechart"
    assert [d for d in diagnostics if d["level"] != "info"] == []
    for chart_spec, failed in ((spec, True), (valid, False)):
        ir = _chart_ir(chart_spec, name="Rev\x0benue")
        report = emitter_module.emit([ir], test_manifest, test_master, tmp_path / "unsafe.pptx", EmitOptions())
        assert not any("failed to emit" in w for w in report.warnings), report.warnings
        assert bool(report.charts[0].get("failed")) is failed
    frame = next(s for s in Presentation(str(tmp_path / "unsafe.pptx")).slides[0].shapes if s.has_chart)
    assert frame.name == "Revenue", "the chart frame is named by add_chart, XML-safe"


def test_waterfall_legend_keys_and_label_ink_read_real_colours(one_chart, test_master):
    """C1 review #14: a single-series waterfall's legend key is its rise colour (the series fill),
    and a stacked waterfall on theme accents takes its label ink from the theme's accent RGB."""
    single = dict(P05_D, options=dict(P05_D["options"], legend="bottom"))
    chart, _shapes, _diagnostics, _path = one_chart(single)
    series = _findall(_chart_xml(chart), ".//c:ser")
    assert _val(series[1], "c:spPr/a:solidFill/a:srgbClr") == "1E9E5A", "EBITDA's legend key is the rise colour"
    assert _val(series[3], "c:spPr/a:solidFill/a:srgbClr") == "1E9E5A"

    presentation = Presentation(str(test_master))
    accents = charts._theme_accents(presentation.slides.add_slide(presentation.slide_layouts[0]))
    assert sorted(accents) == [1, 2, 3, 4, 5, 6]
    themed = {"type": "waterfall", "categories": ["FY23", "Q1", "Q2", "FY24"],
              "series": [{"name": "North", "values": [10, -18, 6, -2]}, {"name": "South", "values": [5, -4, 3, -1]}],
              "options": {"numberFormat": "0", "totals": [0, 3]}}
    chart, _shapes, _diagnostics, _path = one_chart(themed)
    for index, source in ((1, 0), (2, 1)):
        element = _findall(_chart_xml(chart), ".//c:ser")[index]
        assert _val(element, "c:spPr/a:solidFill/a:schemeClr") == f"accent{source + 1}"
        ink = _val(element, "c:dLbls/c:txPr/a:p/a:pPr/a:defRPr/a:solidFill/a:srgbClr")
        assert ink == charts._contrast(accents[source + 1], dark="516467"), (source, ink)


# ------------------------------------------------------------------ WP-C C3: connectors in the chart


def _plot_area_children(chart) -> list[str]:
    return [etree.QName(child).localname for child in _find(_chart_xml(chart), ".//c:plotArea")]


def test_connectors_are_a_scatter_group_on_a_deleted_axis_pair(one_chart):
    """WP-C §9's fallback (C3): the connectors are a `c:scatterChart` of the chart, after the bars and
    before the axes, on a secondary `c:valAx` pair nobody sees — x 0..count in category units at the
    top, y the pinned value axis on the right — so PowerPoint moves them with the plot it lays out."""
    spec = dict(P05_D, options=dict(P05_D["options"], connectors={"color": "#12233B", "width": 1.5,
                                                                  "dash": [4, 2]}))
    chart, shapes, diagnostics, _path = one_chart(spec)
    assert [d for d in diagnostics if d["level"] != "info"] == []
    assert len(shapes) == 1, "no connector shapes on the slide"
    children = _plot_area_children(chart)
    assert children[children.index("barChart") + 1] == "scatterChart", children
    assert children.index("scatterChart") < children.index("catAx")
    assert children.count("valAx") == 3
    space = _chart_xml(chart)
    group = _find(space, ".//c:scatterChart")
    x_id, y_id = charts.CONNECTOR_AXIS_IDS
    assert [e.get("val") for e in _findall(group, "c:axId")] == [str(x_id), str(y_id)]
    primary = primary_value_axis(chart)
    axes = {_val(axis, "c:axId"): axis for axis in _findall(space, ".//c:plotArea/c:valAx")}
    x_axis, y_axis = axes[str(x_id)], axes[str(y_id)]
    for axis, position, cross in ((x_axis, "t", y_id), (y_axis, "r", x_id)):
        assert _val(axis, "c:delete") == "1" and _val(axis, "c:axPos") == position
        assert _val(axis, "c:crossAx") == str(cross) and _val(axis, "c:tickLblPos") == "none"
        assert _val(axis, "c:scaling/c:orientation") == "minMax"
    assert (float(_val(x_axis, "c:scaling/c:min")), float(_val(x_axis, "c:scaling/c:max"))) == (0.0, 6.0)
    assert (float(_val(y_axis, "c:scaling/c:min")), float(_val(y_axis, "c:scaling/c:max"))) == \
        (primary.minimum_scale, primary.maximum_scale), "the connectors' y is the bars' own value axis"
    series = _findall(group, "c:ser")
    assert len(series) == 5, "one series per connector: PowerPoint plots only category-count points of one"
    for element in series:
        line = _find(element, "c:spPr/a:ln")
        assert line.get("w") == str(round(1.5 * 12700))
        assert _val(line, "a:solidFill/a:srgbClr") == "12233B"
        assert _find(line, "a:custDash") is not None or _find(line, "a:prstDash") is not None
        assert _val(element, "c:marker/c:symbol") == "none"
        assert _val(element, "c:dLbls/c:delete") == "1"
        assert _find(element, "c:xVal/c:numLit") is not None and _find(element, "c:yVal/c:numLit") is not None
    assert bar_series_names(chart) == ["Base", "EBITDA", f"Base{MINUS}", f"EBITDA{MINUS}"]


def test_connectors_follow_a_reversed_axis_and_can_be_switched_off(one_chart):
    """Reversed categories reverse the connectors' x axis (slot i is still slot i of the category
    axis); a reversed value axis reverses their y; `connectors: false` writes no group and no axes."""
    spec = dict(P05_D, options=dict(P05_D["options"], categoryAxis={"reverse": True},
                                    valueAxis={"reverse": True}))
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    axes = {_val(axis, "c:axId"): axis for axis in _findall(_chart_xml(chart), ".//c:plotArea/c:valAx")}
    x_id, y_id = charts.CONNECTOR_AXIS_IDS
    assert _val(axes[str(x_id)], "c:scaling/c:orientation") == "maxMin"
    assert _val(axes[str(y_id)], "c:scaling/c:orientation") == "maxMin"

    off, shapes, _diagnostics, _path = one_chart(dict(P05_D, options=dict(P05_D["options"], connectors=False)))
    assert "scatterChart" not in _plot_area_children(off) and _plot_area_children(off).count("valAx") == 1
    assert connector_series(off) == [] and len(shapes) == 1

    column, _shapes, _diagnostics, _path = one_chart({
        "type": "column", "categories": ["a", "b"], "series": [{"values": [1, 2]}],
        "options": {"valueAxis": {"min": 0, "max": 4}, "referenceLines": [{"value": 2}]}})
    assert "scatterChart" not in _plot_area_children(column), "only a waterfall's connectors go in the chart"


def test_coverage_counts_the_connector_series(test_master, test_manifest, tmp_path, monkeypatch):
    """`expected_series` is what the chart must hold: the workbook's series plus one literal series
    per connector. A deck whose connectors went missing fails `chart_checks`; one that has them passes."""
    ir = _chart_ir(STACKED_CROSSING)
    deck = tmp_path / "bridge.pptx"
    emitter_module.emit([ir], test_manifest, test_master, deck, EmitOptions())
    [check] = chart_checks(deck, [ir])
    assert check["ok"] and check["problems"] == [], check["problems"]
    expected = charts.expected_series(STACKED_CROSSING)
    assert [s["name"] for s in expected[-2:]] == ["Connector 1", "Connector 2"]
    assert len(expected) == len(charts.expected_workbook(STACKED_CROSSING)["series"]) + 2

    monkeypatch.setattr(charts, "_in_chart_connectors", lambda _chart, _model: [])
    missing = tmp_path / "no-connectors.pptx"
    emitter_module.emit([ir], test_manifest, test_master, missing, EmitOptions())
    [check] = chart_checks(missing, [ir])
    assert not check["ok"]
    assert any("series" in problem for problem in check["problems"]), check["problems"]


# ------------------------------------------------ render check 2026-09-29 (rendercheck/00-PLAN.md §2)


def _scaling(axis_element) -> dict:
    return {etree.QName(child).localname: child.get("val") for child in _find(axis_element, "c:scaling")}


def _rich_labels(series_element) -> dict[int, str]:
    return {int(_val(label, "c:idx")): "".join(t.text or "" for t in _findall(label, "c:tx/c:rich/a:p/a:r/a:t"))
            for label in _findall(series_element, "c:dLbls/c:dLbl") if _find(label, "c:tx/c:rich") is not None}


def _deleted_labels(series_element) -> set[int]:
    return {int(_val(label, "c:idx")) for label in _findall(series_element, "c:dLbls/c:dLbl")
            if _find(label, "c:delete") is not None}


@pytest.mark.parametrize("axis", [{"min": 0, "max": 100, "majorUnit": 25, "minorUnit": 5},
                                  {"min": 0, "max": 1, "majorUnit": 0.25, "minorUnit": 0.05}],
                         ids=["percent", "fractions"])
def test_percent_axis_is_written_on_powerpoints_0_to_1_scale(one_chart, axis):
    """Render check #3 (d2_s07): `max: 100` on a `percentStacked` axis is 100 × the whole bar, so
    every bar was a 1 % sliver. The spec states the axis in percent; the file holds fractions."""
    spec = {"type": "bar_stacked_100", "categories": ["Mix"],
            "series": [{"name": "Stock", "values": [27.6]}, {"name": "Debt", "values": [72.4]}],
            "options": {"numberFormat": '0.0"%"', "valueAxis": dict(axis, visible=True)}}
    chart, _shapes, diagnostics, _path = one_chart(spec)
    value_axis = _find(_chart_xml(chart), ".//c:valAx")
    assert _scaling(value_axis) == {"max": "1.0", "min": "0.0"}
    assert _val(value_axis, "c:majorUnit") == "0.25" and _val(value_axis, "c:minorUnit") == "0.05"
    # the ticks are fractions, so the data's `0.0"%"` would print "0.3%"; PowerPoint's own 0% instead
    assert _find(value_axis, "c:numFmt").get("formatCode") == "0%"
    assert embedded_workbook(chart)["series"][0]["values"] == [27.6], "the data keep their units"
    if axis["max"] == 1:
        assert sum("read as a fraction" in d["message"] for d in diagnostics) == 3
    # a stacked (not 100 %) chart keeps the author's units and number format
    plain, _s, _d, _p = one_chart({"type": "column_stacked", "categories": ["a"], "series": [{"values": [3]}],
                                   "options": {"numberFormat": "0.0", "valueAxis": {"max": 5, "visible": True}}})
    plain_axis = _find(_chart_xml(plain), ".//c:valAx")
    assert _scaling(plain_axis)["max"] == "5.0" and _find(plain_axis, "c:numFmt").get("formatCode") == "0.0"


def test_series_label_texts_are_rich_text_per_point(one_chart):
    """Render check #6 (d2_s07): `series[k].dataLabels` strings were silently dropped."""
    spec = {"type": "bar_stacked_100", "categories": ["Mix", "Peers"],
            "series": [{"name": "Equity", "values": [31.4, 40], "dataLabels": ["$48.2B (31.4%)", None]},
                       {"name": "Cash", "values": [68.6, 60], "dataLabels": ["", "$9B"]}],
            "colors": ["#00AECF", "#021D44"],
            "options": {"labelPosition": "center", "numberFormat": '0.0"%"'}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    stock, cash = _findall(_chart_xml(chart), ".//c:ser")
    assert _rich_labels(stock) == {0: "$48.2B (31.4%)"}, "point 1 keeps its formatted value"
    assert _rich_labels(cash) == {1: "$9B"} and _deleted_labels(cash) == {0}, '"" hides that label'
    frozen = next(label for label in _findall(stock, "c:dLbls/c:dLbl") if _val(label, "c:idx") == "0")
    series_ink = _find(stock, "c:dLbls/c:txPr/a:p/a:pPr/a:defRPr/a:solidFill/a:srgbClr").get("val")
    assert _find(frozen, "c:tx/c:rich/a:p/a:r/a:rPr/a:solidFill/a:srgbClr").get("val") == series_ink, \
        "the text sits on the bar, so it takes the series' contrast ink"
    assert _val(frozen, "c:dLblPos") == "ctr"
    # XY series carry their texts too (series-level dLbls written by hand)
    xy, _s, _d, _p = one_chart({"type": "scatter", "series": [{"name": "Cos", "x": [1, 2], "y": [3, 4],
                                                                "dataLabels": ["Acme", None]}]})
    assert _rich_labels(_find(_chart_xml(xy), ".//c:ser")) == {0: "Acme"}
    pie, _s, _d, _p = one_chart({"type": "doughnut", "categories": ["In", "Out"],
                                 "series": [{"name": "S", "values": [39, 61], "dataLabels": ["39% in", None]}]})
    assert _rich_labels(_find(_chart_xml(pie), ".//c:ser")) == {0: "39% in"}


def test_waterfall_label_texts_replace_the_step(one_chart):
    spec = {"type": "waterfall", "categories": ["Start", "Up", "Down", "End"],
            "series": [{"name": "S", "values": [10, 5, -3, 12], "dataLabels": ["$10m", None, "", "$12m"]}],
            "options": {"totals": [0, 3]}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    steps = _findall(_chart_xml(chart), ".//c:barChart/c:ser")[1]
    assert _rich_labels(steps) == {0: "$10m", 3: "$12m"}
    assert 2 in _deleted_labels(steps)


def test_one_colour_waterfall_totals_take_the_themes_dark_ink(one_chart, test_master):
    """Render check #5a (d2_s08): one colour → rises and falls in it, totals in the theme's `tx2`
    when that is dark (the test master's dk2 is 1F497D)."""
    spec = {"type": "waterfall", "categories": ["Base", "Sourcing", "Overhead", "One-offs", "Pro-Forma Net"],
            "series": [{"name": "Savings", "values": [11.6, 3.2, 2.4, -4.5, 12.7]}], "colors": ["#00AECF"]}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    steps = _findall(_chart_xml(chart), ".//c:barChart/c:ser")[1]
    colours = {int(_val(point, "c:idx")): _val(point, "c:spPr/a:solidFill/a:srgbClr")
               for point in _findall(steps, "c:dPt")}
    assert colours == {0: "1F497D", 1: "00AECF", 2: "00AECF", 3: "00AECF", 4: "1F497D"}, \
        "the fall is not red and the totals are not the preview's grey"


def _master_with_light_tx2(test_master: Path, target: Path) -> Path:
    """The synthetic test master with its theme's dk2 (tx2 through the colour map) set to yellow,
    as some corporate themes do."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    presentation = Presentation(str(test_master))
    theme = presentation.slide_master.part.part_related_by(RT.THEME)
    theme._blob = re.sub(r"<a:dk2>.*?</a:dk2>", '<a:dk2><a:srgbClr val="FFE600"/></a:dk2>',
                         theme.blob.decode("utf-8"), flags=re.S).encode("utf-8")
    presentation.save(str(target))
    return target


def test_the_chart_ink_is_the_themes_tx2_only_when_it_is_dark(test_master, test_manifest, tmp_path):
    presentation, layout = _deck(test_master, test_manifest)
    assert charts._theme_ink(presentation.slides.add_slide(layout)) == "1F497D"
    from app.engine.importer import import_master

    light = _master_with_light_tx2(test_master, tmp_path / "light-tx2.pptx")
    presentation, layout = _deck(light, import_master(light, tmp_path / "import", render=False), "layout-01")
    slide = presentation.slides.add_slide(layout)
    assert charts._theme_ink(slide) is None, "this theme's tx2 is yellow FFE600: the model's #1F2937 stands"
    model = charts.chart_model.normalise({"type": "waterfall", "categories": ["a", "b", "c"],
                                          "series": [{"values": [5, 2, 7]}], "colors": ["#00AECF"]})
    charts._ink_the_totals(model, None)
    assert model["waterfall"]["colors"]["total"] == "1F2937"
    charts._ink_the_totals(model, "021D44")
    assert model["waterfall"]["colors"]["total"] == "021D44"


def test_recognised_bars_export_top_down(test_master, test_manifest, tmp_path):
    """Render check #4: a recognised bar chart lists its categories top-down, as the design reads,
    so the axis is reversed and the value axis crosses at the far end — it used to export upside down."""
    from app.engine.classify.charts import recognise_charts

    ir = recognise_charts(IR.load(config.TORTURE_FIXTURE / "ir" / "svg-charts.json"))
    bar = next(e for e in ir.elements if e.kind == "chart" and e.spec["type"] == "bar")
    assert bar.spec["categories"] == ["Renewal", "Referral", "Upsell", "Business"]
    assert bar.spec["options"]["categoryAxis"] == {"visible": True, "reverse": True}
    presentation, layout = _deck(test_master, test_manifest)
    charts.add_chart(presentation.slides.add_slide(layout), bar.spec, BOX, dict(STYLE))
    path = tmp_path / "bars.pptx"
    presentation.save(str(path))
    chart = next(s for s in Presentation(str(path)).slides[0].shapes if s.has_chart).chart
    space = _chart_xml(chart)
    assert _val(_find(space, ".//c:catAx"), "c:scaling/c:orientation") == "maxMin"
    assert _val(_find(space, ".//c:valAx"), "c:crosses") == "max"


# ------------------------------------------ G24: label styles, line breaks, middle subtotals (render check §9)


def _paragraphs(label) -> list[str]:
    """The lines of a `c:dLbl`'s rich text, one per `a:p`."""
    return ["".join(t.text or "" for t in _findall(p, "a:r/a:t")) for p in _findall(label, "c:tx/c:rich/a:p")]


def test_label_style_bold_writes_b_on_every_label_run(one_chart):
    """`options.labelStyle.bold`: `b="1"` on the series labels' `txPr` and on each own-text run —
    authors overlaid HTML on pinned plots to get bold labels (d2_s07, d1_s06)."""
    spec = {"type": "column_stacked", "categories": ["FY23", "FY24"],
            "series": [{"name": "A", "values": [3, 4], "dataLabels": ["$3B\n(43%)", None]},
                       {"name": "B", "values": [4, 5]}],
            "colors": ["#021D44", "#DDEAF7"], "options": {"labelStyle": {"bold": True}}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    first, second = _findall(_chart_xml(chart), ".//c:ser")
    for series in (first, second):
        assert _find(series, "c:dLbls/c:txPr/a:p/a:pPr/a:defRPr").get("b") == "1"
    own = next(label for label in _findall(first, "c:dLbls/c:dLbl") if _val(label, "c:idx") == "0")
    assert {r.get("b") for r in _findall(own, "c:tx/c:rich/a:p/a:r/a:rPr")} == {"1"}
    assert {r.get("b") for r in _findall(own, "c:tx/c:rich/a:p/a:pPr/a:defRPr")} == {"1"}
    # not asked for, nothing written: the XML of a plain chart is unchanged
    plain, _s, _d, _p = one_chart(dict(spec, options={}))
    assert all(p.get("b") is None for p in _findall(_chart_xml(plain), ".//c:dLbls//a:defRPr"))
    assert all(p.get("b") is None for p in _findall(_chart_xml(plain), ".//c:dLbls//a:rPr"))


def test_a_line_break_in_a_label_text_is_a_second_paragraph(one_chart):
    """`series[k].dataLabels` with `\\n` ("$48.2B\\n(31.4%)") becomes one `a:p` per line, as
    PowerPoint writes a label broken with Enter; one-line texts keep one paragraph."""
    spec = {"type": "bar_stacked_100", "categories": ["Mix"],
            "series": [{"name": "Equity", "values": [31.4], "dataLabels": ["$48.2B\n(31.4%)"]},
                       {"name": "Cash", "values": [68.6], "dataLabels": ["$105B (68.6%)"]}],
            "options": {"labelPosition": "center"}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    stock, cash = _findall(_chart_xml(chart), ".//c:ser")
    two = _find(stock, "c:dLbls/c:dLbl")
    assert _paragraphs(two) == ["$48.2B", "(31.4%)"]
    assert _paragraphs(_find(cash, "c:dLbls/c:dLbl")) == ["$105B (68.6%)"]
    # CRLF reads as one break, and a waterfall's own texts break the same way
    bridge, _s, _d, _p = one_chart({"type": "waterfall", "categories": ["Start", "Up", "End"],
                                    "series": [{"name": "S", "values": [10, 5, 15],
                                                "dataLabels": ["$10m\r\nbase", None, "$15m\n(+50%)"]}]})
    steps = _findall(_chart_xml(bridge), ".//c:barChart/c:ser")[1]
    texts = {int(_val(label, "c:idx")): _paragraphs(label) for label in _findall(steps, "c:dLbls/c:dLbl")
             if _find(label, "c:tx") is not None}
    assert texts == {0: ["$10m", "base"], 2: ["$15m", "(+50%)"]}


def test_label_style_colour_and_size_replace_the_contrast_pick(one_chart):
    """`labelStyle.color` is the labels' ink on every bar (no per-point contrast overrides), and
    `labelStyle.fontSize` is their size in px (0.75 pt each)."""
    spec = {"type": "column", "categories": ["a", "b"], "series": [{"name": "A", "values": [3, 4]}],
            "colors": ["#021D44"], "options": {"labelPosition": "inEnd", "pointColors": {"0": {"1": "#F2F2F2"}},
                                               "labelStyle": {"color": "#FFC000", "fontSize": 16}}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    series = _find(_chart_xml(chart), ".//c:ser")
    properties = _find(series, "c:dLbls/c:txPr/a:p/a:pPr/a:defRPr")
    assert _val(properties, "a:solidFill/a:srgbClr") == "FFC000" and properties.get("sz") == "1200"
    assert _findall(series, "c:dLbls/c:dLbl") == [], "no contrast override on the pale point"
    # without it the pale point takes the dark ink and the dark bar white, as before
    plain, _s, _d, _p = one_chart(dict(spec, options=dict(spec["options"], labelStyle={})))
    assert len(_findall(_find(_chart_xml(plain), ".//c:ser"), "c:dLbls/c:dLbl")) == 1


def test_axis_bold_writes_b_on_the_tick_labels(one_chart):
    spec = {"type": "bar", "categories": ["Supply\nchain", "SG&A"], "series": [{"values": [4.1, 2.8]}],
            "options": {"valueAxis": {"bold": True}, "categoryAxis": {"bold": True}}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    space = _chart_xml(chart)
    assert _find(space, ".//c:catAx/c:txPr/a:p/a:pPr/a:defRPr").get("b") == "1"
    assert _find(primary_value_axis(chart)._element, "c:txPr/a:p/a:pPr/a:defRPr").get("b") == "1"
    plain, _s, _d, _p = one_chart(dict(spec, options={}))
    assert _find(_chart_xml(plain), ".//c:catAx/c:txPr/a:p/a:pPr/a:defRPr").get("b") is None


def test_a_line_break_in_a_category_is_kept_in_the_cache_and_the_workbook(one_chart):
    """PowerPoint draws a category's `\\n` as a line break (checked in PowerPoint, G24), so the
    cached text and the workbook keep it as written."""
    spec = {"type": "column", "categories": ["FY23\nActual", "FY24\nBudget"], "series": [{"values": [1, 2]}]}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    cached = [v.text for v in _findall(_chart_xml(chart), ".//c:ser/c:cat//c:pt/c:v")]
    assert cached == ["FY23\nActual", "FY24\nBudget"]
    assert embedded_workbook(chart)["categories"] == ["FY23\nActual", "FY24\nBudget"]


def test_out_end_waterfall_labels_take_the_label_style_and_break_into_paragraphs(one_chart):
    spec = {"type": "waterfall", "categories": ["Opening", "Wins", "Losses", "Closing"],
            "series": [{"name": "Cash", "values": [100, 30, -10, 120],
                        "dataLabels": ["$100m\n(base)", None, None, "$120m\n(+20%)"]}],
            "options": {"labelPosition": "outEnd", "labelStyle": {"bold": True, "color": "#C00000"}}}
    chart, _shapes, diagnostics, path = one_chart(spec, name="Bridge")
    assert [d for d in diagnostics if d["level"] != "info"] == []
    labels = [s for s in Presentation(str(path)).slides[0].shapes if " label " in s.name]
    assert [[p.text for p in s.text_frame.paragraphs] for s in labels] == \
        [["$100m", "(base)"], ["+30"], ["-10"], ["$120m", "(+20%)"]]
    runs = [r for s in labels for p in s.text_frame.paragraphs for r in p.runs]
    assert {r.font.bold for r in runs} == {True} and {str(r.font.color.rgb) for r in runs} == {"C00000"}
    assert labels[0].height > labels[1].height * 1.5, "a two-line box is two lines tall"


def test_a_middle_subtotal_is_inferred_and_stands_on_the_axis(one_chart):
    """d1_s03: "2025 EBITDA" equals the bridge so far and names a subtotal, so with nothing declared
    it stands on the axis (`Base` 0) instead of floating as a step, and the final bar follows it."""
    spec = {"type": "waterfall",
            "categories": ["2020 EBITDA", "Volume", "Input costs", "Overhead", "Pricing", "2025 EBITDA", "Uplift",
                           "Pro-forma EBITDA"],
            "series": [{"name": "EBITDA ($B)", "values": [12.6, -0.8, -2.1, -1.5, -1.2, 7.0, 8.4, 15.4]}],
            "options": {"numberFormat": "0.0", "valueAxis": {"min": 0, "max": 20}}}
    chart, _shapes, _diagnostics, _path = one_chart(spec)
    workbook = embedded_workbook(chart)["series"]
    base, steps = workbook[0]["values"], workbook[1]["values"]
    assert base[5] in (0, None) and steps[5] == 7.0, "the subtotal stands on the axis"
    assert base[6] == pytest.approx(7.0) and steps[6] == 8.4, "the uplift climbs from the subtotal"
    assert base[7] in (0, None) and steps[7] == 15.4
    declared, _s, _d, _p = one_chart(dict(spec, options=dict(spec["options"], totals=[0, 7])))
    assert embedded_workbook(declared)["series"][0]["values"][5] == pytest.approx(7.0), "declared totals win"
