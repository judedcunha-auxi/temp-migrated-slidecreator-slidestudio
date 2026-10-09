"""The pipeline runs end to end: project -> measured IR -> `.pptx` -> report.

This proves plumbing, not fidelity: the browser launches, assets resolve, fonts are forced, the right
layout is chosen, the deck saves and reopens, and the suite writes its two files. The project is the
synthetic sample (`tests/engine/sample_project.py`); every later change that breaks the path itself
fails here.
"""
from __future__ import annotations

import ast
import json
from pathlib import Path

import pytest
from PIL import Image

from app.config import engine as config
from app.core.browser_pool import BrowserPool
from app.engine.extract.html import font_forcing_css, installed_families, render_reference
from app.engine.ir import IR
from app.engine.manifest import Manifest
from app.engine.pipeline import PipelineError, export_deck, load_project
from app.engine.reports import ExportOptions
from tests.engine import helpers
from tests.engine.helpers import extract_html

REPO = Path(__file__).resolve().parents[2]
SLIDE_IDS = ["sld_sample_01", "sld_sample_02"]


def test_config_declares_every_key_the_engine_reads():
    """The engine reads these as module attributes of `app.config.engine`; a missing one is a crash."""
    for key in [
        "PROJECT_ROOT", "TORTURE_FIXTURE", "MASTERS_FIXTURE", "EXTRA_MASTERS", "RENDERER_URL",
        "VALIDATE_PY", "EDGE_CHANNEL", "FONT_DIRS", "CHART_PREVIEW_JS",
        "TEXT_DY_BY_FONT", "WIDTH_LOCK_FACTOR", "TEXT_BOX_WIDTH_SLACK",
        "GATE_BLUR", "GATE_DELTA_E", "GATE_MIN_AREA", "GATE_SLACK", "GATE_MASK_PLACEHOLDERS",
        "GATE_TARGET_PX2_1280", "PX_PER_IN", "EMU_PER_PX", "EMU_PER_IN", "PT_PER_PX", "FONT_SZ_PER_PX",
        "RENDERER_TIMEOUT_S", "TEXT_POSITION_TOL_PX", "TEXT_ROW_SHARE", "TEXT_ALIGN_TOL_PX",
        "REMOTE_RESOURCES",
    ]:
        assert hasattr(config, key), f"config.{key} is missing"

    assert config.PX_PER_IN == 96 and config.EMU_PER_PX == 9525 and config.FONT_SZ_PER_PX == 75
    assert config.CHART_PREVIEW_JS.exists(), "the chart preview ships inside app/engine"


def test_gate_target_scales_with_canvas_area(monkeypatch: pytest.MonkeyPatch):
    """`T(canvas) = T_1280 x (w*h)/(1280*720)`: a threshold in pixels must follow the canvas."""
    monkeypatch.setattr(config, "GATE_TARGET_PX2_1280", 20000.0)
    assert config.gate_target_px2(1280, 720) == 20000.0
    assert config.gate_target_px2(640, 360) == 5000.0          # quarter the area, quarter the budget
    assert config.gate_target_px2(960, 720) == 15000.0         # the 4:3 master

    monkeypatch.setattr(config, "GATE_TARGET_PX2_1280", None)
    assert config.gate_target_px2(1280, 720) is None, "no target means the row is measured, not judged"


def test_the_calibration_pass_has_set_a_gate_target():
    """An unset target silently turns the gate row into a no-op."""
    assert config.GATE_TARGET_PX2_1280 == 6000.0


def test_font_forcing_css_targets_headings_and_the_title_placeholder():
    css = font_forcing_css({"major": "Lexend", "minor": "Arial"})
    assert '--engine-font-major: "Lexend"' in css and '--engine-font-minor: "Arial"' in css
    assert '[data-placeholder="title"]' in css
    assert css.count("!important") >= 2


def _an_installed_family() -> str:
    """The family name of some font file on this machine (Arial on Windows, DejaVu/Liberation on Linux)."""
    from fontTools.ttLib import TTFont

    for path in config.font_files():
        if path.suffix.lower() not in (".ttf", ".otf"):
            continue
        try:
            name = TTFont(str(path), lazy=True)["name"].getDebugName(1)
        except Exception:  # noqa: BLE001 - an unreadable font file is just not the one we pick
            continue
        if name:
            return str(name)
    pytest.fail("no readable .ttf/.otf font found in config.FONT_DIRS")


def test_installed_families_uses_the_font_files_not_the_browser():
    """An installed family comes back True; a made-up one False (this is the `fonts.check` trap)."""
    real = _an_installed_family()
    result = installed_families([real, "Definitely Not A Real Font 91723"])
    assert result[real] is True
    assert result["Definitely Not A Real Font 91723"] is False


def test_load_project_reads_the_sample(sample_dir: Path):
    project = load_project(sample_dir)
    assert project.master == sample_dir / "master.pptx"
    assert project.id == "prj_sample"
    assert [s["slideId"] for s in project.slides()] == SLIDE_IDS


def test_load_project_slide_selection_and_errors(sample_dir: Path, tmp_path: Path):
    source = load_project(sample_dir)
    assert len(source.slides(["sld_sample_02"])) == 1
    with pytest.raises(PipelineError, match="no such slide"):
        source.slides(["sld_nope"])
    with pytest.raises(PipelineError, match="no project.json"):
        load_project(tmp_path / "nowhere")
    (tmp_path / "bare").mkdir()
    (tmp_path / "bare" / "project.json").write_text('{"slides": []}', encoding="utf-8")
    with pytest.raises(PipelineError, match="no manifest.json"):
        load_project(tmp_path / "bare")


def test_extract_produces_a_valid_ir_with_the_slide_identity(sample_dir: Path, sample_manifest: Manifest):
    ir = extract_html(
        sample_dir / "slide-01.html", sample_manifest, "layout-06", sample_dir / "assets",
        slide_id="sld_sample_01", title="Growth by channel",
    )
    assert ir.validate() == []
    assert (ir.canvas.w, ir.canvas.h) == (1280, 720)
    assert ir.slide.id == "sld_sample_01" and ir.slide.layoutId == "layout-06"
    assert ir.fonts.forced == {"major": "Arial", "minor": "Arial"}
    assert "Arial" in ir.fonts.used

    # Native content came back, and every picture still has a reason and a diagnostic.
    assert ir.counts()["text"] > 0 and ir.counts()["shape"] > 0
    for raster in ir.of_kind("raster"):
        assert Path(raster.src).exists() and raster.reason
        assert any(d.elementId == raster.id for d in ir.diagnostics)


def test_extraction_is_deterministic(sample_dir: Path, sample_manifest: Manifest):
    """Two extractions of one slide must produce identical JSON: the base of the determinism row."""
    first = extract_html(sample_dir / "slide-02.html", sample_manifest, "layout-01", sample_dir / "assets",
                         slide_id="sld_sample_02")
    second = extract_html(sample_dir / "slide-02.html", sample_manifest, "layout-01", sample_dir / "assets",
                          slide_id="sld_sample_02")
    assert first.dumps() == second.dumps()


def test_asset_urls_are_routed_to_the_assets_directory(sample_dir: Path, sample_manifest: Manifest, tmp_path: Path):
    """`/api/projects/<id>/assets/<file>` 404s under `file://` unless it is routed at the filesystem.

    Both halves: a present asset resolves and loads, a missing one becomes an `error` diagnostic
    instead of a silently blank measurement.
    """
    slide = tmp_path / "assets-slide.html"
    slide.write_text(
        '<!doctype html><html><body style="margin:0">'
        '<img id="ok" src="/api/projects/prj_x/assets/asset-logo.png">'
        '<img id="gone" src="/api/projects/prj_x/assets/not-here.png">'
        "</body></html>",
        encoding="utf-8",
    )

    ir = extract_html(slide, sample_manifest, "layout-06", sample_dir / "assets", slide_id="sld_assets")
    errors = [d for d in ir.diagnostics if d.level == "error" and d.source.startswith("asset:")]
    assert [d.source for d in errors] == ["asset:not-here.png"], (
        "the missing asset must be reported and the present one must not"
    )
    assert ir.counts().get("image", 0) >= 1, "the present asset loaded and became an image"


def test_derived_images_land_in_the_callers_workspace(sample_manifest: Manifest, sample_dir: Path, tmp_path: Path):
    """A raster goes under `<workspace>/derived/<slide id>/`, never into the project's assets."""
    slide = tmp_path / "raster.html"
    slide.write_text(
        '<!doctype html><html><body style="margin:0">'
        '<div style="position:absolute;left:10px;top:10px;width:200px;height:100px;'
        'background:conic-gradient(red,blue)"></div></body></html>',
        encoding="utf-8",
    )
    workspace = tmp_path / "ws"
    before = sorted(p.name for p in (sample_dir / "assets").iterdir())
    ir = extract_html(slide, sample_manifest, "layout-06", sample_dir / "assets",
                      workspace=workspace, slide_id="sld_raster")
    rasters = ir.of_kind("raster")
    assert rasters, "a conic gradient has no DrawingML equivalent"
    for raster in rasters:
        assert Path(raster.src).resolve().is_relative_to((workspace / "derived" / "sld_raster").resolve())
    assert sorted(p.name for p in (sample_dir / "assets").iterdir()) == before


def test_render_reference_composites_layout_and_slide(sample_dir: Path, sample_manifest: Manifest, out_dir: Path):
    layout = sample_manifest.layout("layout-06")
    assert layout.background
    output = render_reference(
        sample_dir / "slide-01.html",
        sample_dir / "layouts" / layout.background,
        out_dir / "reference.png",
        assets_dir=sample_dir / "assets",
        fonts=sample_manifest.fonts,
    )
    with Image.open(output) as image:
        assert image.size == (sample_manifest.canvas_w, sample_manifest.canvas_h)
    assert not list(out_dir.glob("*.composite.html")), "the temporary composite must be cleaned up"


def test_export_deck_runs_end_to_end_on_the_sample(sample_dir: Path, sample_manifest: Manifest, out_dir: Path):
    """Both sample slides, on their own layouts, into one deck that reopens."""
    from pptx import Presentation

    events: list[dict[str, object]] = []
    result = export_deck(sample_dir, ExportOptions(out_dir=out_dir, on_event=events.append))

    assert result.pptx.exists() and result.cost == 0.0
    assert result.slide_ids == SLIDE_IDS
    assert len(result.irs) == 2 and all(ir.validate() == [] for ir in result.irs)
    assert {e["type"] for e in events} >= {"status", "usage", "export", "done"}
    assert (out_dir / "work").is_dir(), "the default workspace is <out_dir>/work"

    presentation = Presentation(str(result.pptx))
    assert len(presentation.slides) == 2
    # The right layout, identified by part, not by name.
    assert [str(s.slide_layout.part.partname) for s in presentation.slides] == [
        sample_manifest.layout("layout-06").partName,
        sample_manifest.layout("layout-01").partName,
    ]
    # The cover's photo and logo arrived as pictures, and slide 1's title went into its placeholder.
    cover_pictures = [shape for shape in presentation.slides[1].shapes if shape.shape_type == 13]
    assert len(cover_pictures) >= 2
    titles = [shape for shape in presentation.slides[0].placeholders if shape.placeholder_format.idx == 0]
    assert titles and "Growth by channel" in titles[0].text_frame.text


def test_export_runs_on_a_browser_pool_thread(sample_dir: Path, tmp_path: Path):
    """The service runs exports on `BrowserPool` threads, each with its own Chromium."""
    with BrowserPool(1, name="test-engine") as pool:
        result = pool.run(lambda: export_deck(sample_dir, ExportOptions(
            out_dir=tmp_path / "pooled", slide_ids=["sld_sample_02"])), timeout=300)
    assert result.pptx.exists() and result.slide_ids == ["sld_sample_02"]


def test_export_deck_produces_a_stable_filename(sample_dir: Path, tmp_path: Path):
    """No timestamp in the name: a run must overwrite its predecessor for byte comparison to mean anything."""
    first = export_deck(sample_dir, ExportOptions(out_dir=tmp_path / "a", slide_ids=["sld_sample_02"]))
    second = export_deck(sample_dir, ExportOptions(out_dir=tmp_path / "b", slide_ids=["sld_sample_02"]))
    assert first.pptx.name == second.pptx.name == "Synthetic-sample-deck.pptx"


def test_export_result_rerun_reproduces_the_export(sample_dir: Path, tmp_path: Path):
    """The determinism check needs this hook."""
    result = export_deck(sample_dir, ExportOptions(out_dir=tmp_path / "one", slide_ids=["sld_sample_01"]))
    assert result.rerun is not None
    again = result.rerun(tmp_path / "two")
    assert again.pptx.exists() and again.pptx != result.pptx
    assert [ir.dumps() for ir in again.irs] == [ir.dumps() for ir in result.irs]


@pytest.mark.validator
def test_suite_writes_both_report_files(sample_dir: Path, out_dir: Path):
    from app.engine.verify.suite import run_suite

    result = export_deck(sample_dir, ExportOptions(out_dir=out_dir / "export", slide_ids=["sld_sample_01"]))
    report = run_suite(result, out_dir, project_dir=sample_dir)

    assert (out_dir / "report.json").exists() and (out_dir / "report.md").exists()
    data = json.loads((out_dir / "report.json").read_text(encoding="utf-8"))
    assert data["project_id"] == sample_dir.name
    # The sample's charts and donut are SVG the engine expands natively: no raster.
    assert report.coverage is not None and report.coverage.raster_count == 0
    # No renderer is passed, so the gate row cannot have run and must not pretend it did.
    assert "gate" not in report.gate_results


@pytest.mark.parametrize(
    "stem, canvas",
    [("test-16x9", (1280, 720)), ("test-4x3", (960, 720))],
)
def test_generated_test_masters_are_usable(masters_dir: Path, stem: str, canvas: tuple[int, int]):
    """The size-generality fixtures: two canvases, real `idx`/`partName`, boxes the gate can mask."""
    deck = masters_dir / f"{stem}.pptx"
    manifest = Manifest.load(masters_dir / f"{stem}.manifest.json")
    assert deck.exists()
    assert (manifest.canvas_w, manifest.canvas_h) == canvas
    assert manifest.validate() == []
    assert manifest.importer == "deterministic" and manifest.has_part_names

    title_layout = next(lay for lay in manifest.layouts if lay.placeholder("title"))
    title = title_layout.placeholder("title")
    assert title is not None
    assert title.idx is not None, "generated manifests must carry real placeholder idx values"
    assert title.w > 0 and title.h > 0, "placeholder geometry must be inherited, not left at zero"

    masked = [lay for lay in manifest.layouts if lay.masked_boxes()]
    assert masked, "at least one layout must carry sldNum/dt/ftr so the gate's masking is exercised"

    backgrounds = masters_dir / stem / "layouts"
    assert len(list(backgrounds.glob("layout-*.png"))) == len(manifest.layouts)
    first = manifest.layouts[0].background
    assert first
    with Image.open(backgrounds / first) as image:
        assert image.size == canvas


def _imports(path: Path) -> list[str]:
    tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
    names: list[str] = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            names += [alias.name for alias in node.names]
        elif isinstance(node, ast.ImportFrom):
            names.append(node.module or "")
    return names


def test_engine_never_imports_the_api_layer():
    """Architectural rule (docs/architecture.md): the engine is driven, it never drives the API."""
    offenders = sorted(
        str(path.relative_to(REPO))
        for path in (REPO / "app" / "engine").rglob("*.py")
        if any(name == "app.api" or name.startswith(("app.api.", "app.main", "server")) for name in _imports(path))
    )
    assert offenders == [], f"app.engine must not import the API layer: {offenders}"


def test_engine_imports_only_config_and_the_browser_pool_from_the_service():
    allowed = ("app.engine", "app.config", "app.core.browser_pool")
    offenders = sorted(
        f"{path.relative_to(REPO)}: {name}"
        for path in (REPO / "app" / "engine").rglob("*.py")
        for name in _imports(path)
        if name.startswith("app.") and not name.startswith(allowed)
    )
    assert offenders == []


@pytest.mark.renderer_full
def test_torture_gate_row_on_one_family(tmp_path: Path):
    """The torture runner's gate leg on the hand-written `text` IR, no browser in the loop.

    The row carries the gate fields; a `gateMaxPx2` below the measured area fails the family, `0`
    is honoured as a real ceiling (`is not None`), and a missing or stale reference is a problem,
    never a skip.
    """
    import shutil
    from dataclasses import replace

    from app.engine.emit import pptx as emitter
    from app.engine.reports import EmitOptions
    from app.engine.verify.fixtures import context_for
    from app.engine.verify.torture import gate_row, judge_family

    renderer = helpers.renderer()
    html = config.TORTURE_FIXTURE / "text.html"
    context = context_for(html)
    # the absent-ceiling case, whatever the family file carries
    context = replace(context, expect={k: v for k, v in (context.expect or {}).items() if k != "gateMaxPx2"})
    ir = IR.from_json(json.loads((config.TORTURE_FIXTURE / "ir" / "text.json").read_text(encoding="utf-8")))
    assert ir.validate() == []
    deck = tmp_path / "text.pptx"
    emitter.emit([ir], context.manifest, context.master, deck, EmitOptions())

    row, area = gate_row(html, context, ir, deck, tmp_path / "gate", renderer=renderer)
    assert row["referenceFresh"] is True and row["problems"] == []
    for field in ("gateMaxPx2", "nonChartArea", "defectArea", "components"):
        assert field in row, field
    assert area is not None and area == row["nonChartArea"] and area > 0
    assert row["gateMaxPx2"] == config.gate_target_px2(1280, 720), "absent gateMaxPx2 -> the canvas target"

    def gate_problems(ceiling: float) -> list[str]:
        judged = replace(context, expect={**(context.expect or {}), "gateMaxPx2": ceiling})
        return [p for p in judge_family(judged, ir, ir, deck, None, None, gate_px2=area) if p.startswith("gate:")]

    assert gate_problems(area - 1) and not gate_problems(area)
    assert gate_problems(0), "gateMaxPx2: 0 means pixel-clean, not 'no ceiling'"

    # A copy of the family with no reference beside it, then with a reference of different HTML.
    family = tmp_path / "torture"
    family.mkdir()
    shutil.copy(html, family / "text.html")
    shutil.copy(html.with_suffix(".expect.json"), family / "text.expect.json")
    copied = context_for(family / "text.html")
    missing, missing_area = gate_row(family / "text.html", copied, ir, deck, tmp_path / "missing", renderer=renderer)
    assert missing_area is None and missing["referenceFresh"] is False
    assert any("reference missing" in p for p in missing["problems"])

    (family / "reference").mkdir()
    shutil.copy(config.TORTURE_FIXTURE / "reference" / "text.png", family / "reference" / "text.png")
    shutil.copy(config.TORTURE_FIXTURE / "reference" / "manifest.json", family / "reference" / "manifest.json")
    fresh, _ = gate_row(family / "text.html", copied, ir, deck, tmp_path / "fresh", renderer=renderer)
    assert fresh["referenceFresh"] is True
    with (family / "text.html").open("a", encoding="utf-8") as handle:
        handle.write("\n<!-- edited after the reference was rendered -->\n")
    stale, stale_area = gate_row(family / "text.html", copied, ir, deck, tmp_path / "stale", renderer=renderer)
    assert stale_area is None and any("reference stale" in p for p in stale["problems"])
