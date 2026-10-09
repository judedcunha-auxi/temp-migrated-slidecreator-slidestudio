"""Exports (app/core/exports.py) on the real engine: a design deck's export, and `stitch_deck`.

`stitch_deck` gets its parts the way the jobs do: two slides each designed by `design_and_export`
(fake provider), packed as deck bundles, and unpacked into the stitch job's workspace.
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import pytest
from pptx import Presentation

from app.core import engine_service, exports, preview
from app.core.design import bundle
from app.core.design.deck import DesignDeck
from app.core.slides import pipeline
from app.engine.renderer import BlankLayoutsRenderer
from tests.conftest import build_ai_settings
from tests.core.design.conftest import content_layout, make_deck, slide_html, stub_services
from tests.core.slides.test_pipeline import SPEC, designer
from tests.fakes.llm import resolver


@pytest.fixture(scope="module", autouse=True)
def close_browsers() -> Iterator[None]:
    yield
    engine_service.shutdown()
    preview.shutdown()


def test_a_design_deck_exports_every_slide(tmp_path: Path):
    deck = make_deck(tmp_path)
    for title in ("Revenue grew 23%", "Costs fell 4%"):
        deck.save_slide(slide_html(title), title, content_layout(deck), "seed")
    outcome = exports.export_deck(deck, tmp_path / "job")
    assert outcome.pptx == tmp_path / "job" / "out" / "deck.pptx"
    assert len(Presentation(str(outcome.pptx)).slides) == 2
    assert outcome.element_count > 0 and isinstance(outcome.review_flag_count, int)
    assert (deck.dir / "diagnostics.json").is_file()
    report = engine_service.readiness(deck, deck.slides()[0]["id"])
    assert report["blocking"] is False and report["version"] == 1


def test_an_empty_deck_is_refused(tmp_path: Path):
    with pytest.raises(exports.StitchError):
        exports.export_deck(DesignDeck.create(tmp_path / "d", "Empty"), tmp_path / "job")


def designed(tmp_path: Path, name: str, title: str) -> tuple[bytes, str]:
    spec = {**SPEC, "slide": {**SPEC["slide"], "title": title}}  # type: ignore[dict-item]
    result = pipeline.design_and_export(pipeline.DesignRequest(spec=spec), tmp_path / name,
                                        settings=build_ai_settings(), resolve=resolver(designer()),
                                        services=stub_services(), renderer=BlankLayoutsRenderer())
    return (result.pptx.parent / "deck.zip").read_bytes(), result.slide_id


def test_stitch_deck_assembles_designed_slides_in_order(tmp_path: Path):
    first, second = designed(tmp_path, "a", "Revenue grew 23%"), designed(tmp_path, "b", "Costs fell 4%")
    work = tmp_path / "stitch"
    parts = []
    for i, (data, sid) in enumerate((first, second)):
        deck = bundle.unpack(data, work / f"part-{i}")
        parts.append(exports.SlidePart(deck.dir, sid))
    result = exports.stitch_deck(parts, work, title="Growth review")

    prs = Presentation(str(result.pptx))
    assert result.slide_count == len(prs.slides) == 2
    texts = [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in prs.slides]
    assert "Revenue" in texts[0] and "Costs" in texts[1], "the parts keep their order"
    assert result.element_count > 0 and result.moved == []


def test_stitch_refuses_nothing_and_unknown_slides(tmp_path: Path):
    with pytest.raises(exports.StitchError):
        exports.stitch_deck([], tmp_path / "w")
    deck = make_deck(tmp_path)
    with pytest.raises(exports.StitchError):
        exports.stitch_deck([exports.SlidePart(deck.dir, "sld_0000000000")], tmp_path / "w2")
