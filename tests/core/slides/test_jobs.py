"""The pipeline's background jobs, run by the real worker on fakeredis and the General service fake.

`design_and_export` (fake provider) -> its deck bundle -> `stitch_deck`, then `brand.extract` and
`masters.render_layouts`: each job's inputs come through the storage port into a job workspace, and
its outputs go back through the port. Nothing starts a worker but the test itself.
"""

from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Iterator
from pathlib import Path

import fakeredis
import pytest
from pptx import Presentation

from app.core import engine_service, preview
from app.core.brand import extract as brand_extract
from app.core.jobs.queue import JobQueue
from app.core.jobs.registry import JobRegistry
from app.core.jobs.worker import Worker
from app.core.redis_client import RedisClient
from app.core.slides import jobs
from app.core.storage.models import CallerContext, JobRecord
from tests.conftest import build_ai_settings
from tests.core.design.conftest import SYNTHETIC_MASTER, stub_services
from tests.core.slides.test_pipeline import SPEC, designer
from tests.fakes.general_service import FakeGeneralService
from tests.fakes.llm import resolver
from tests.fakes.renderer import FakeRenderer

pytestmark = pytest.mark.asyncio


@pytest.fixture(scope="module", autouse=True)
def close_browsers() -> Iterator[None]:
    yield
    engine_service.shutdown()
    preview.shutdown()


async def setup(tmp_path: Path) -> tuple[FakeGeneralService, JobQueue, Worker, CallerContext]:
    store = FakeGeneralService()
    ctx, _ = await store.make_user("alice")
    deps = jobs.PipelineDeps(blobs=store.blobs, scratch_root=tmp_path, settings=build_ai_settings(),
                             resolve=resolver(designer()), services=stub_services(), renderer_factory=FakeRenderer)
    registry = JobRegistry()
    jobs.register(registry, deps)
    queue = JobQueue(RedisClient(fakeredis.FakeAsyncRedis(decode_responses=True)), store)
    worker = Worker(queue, registry.install(queue))
    return store, queue, worker, ctx


async def finished(store: FakeGeneralService, worker: Worker, ctx: CallerContext, job_id: str) -> JobRecord:
    assert await worker.run_once() is True
    record = await store.jobs.get(ctx, job_id)
    assert record.status == "done", record.error
    return record


async def test_design_then_stitch_through_the_worker(tmp_path: Path):
    store, queue, worker, ctx = await setup(tmp_path)
    captured = brand_extract.extract(SYNTHETIC_MASTER.read_bytes())["capturedFurniture"]
    furniture = (await store.blobs.put(ctx, json.dumps(captured).encode(), "application/json")).ref

    refs = []
    for title in ("Revenue grew 23%", "Costs fell 4%"):
        spec = {**SPEC, "slide": {**SPEC["slide"], "title": title}}  # type: ignore[dict-item]
        job = await queue.enqueue(ctx, "slides.design_and_export",
                                  {"spec": spec, "furnitureRef": furniture, "brand": {"primary_color": "#1F2A44"}})
        record = await finished(store, worker, ctx, job.id)
        result = record.result
        assert result["elementCount"] > 0 and isinstance(result["reviewFlagCount"], int)
        assert record.cost_usd == pytest.approx(result["costUsd"]) and record.cost_usd > 0
        pptx = await store.blobs.get(ctx, result["pptxRef"])
        assert len(Presentation(io.BytesIO(pptx.data)).slides) == 1
        refs.append({"deckRef": result["deckRef"], "slideId": result["slideId"]})

    job = await queue.enqueue(ctx, "exports.stitch_deck", {"parts": refs, "title": "Growth"})
    record = await finished(store, worker, ctx, job.id)
    deck = await store.blobs.get(ctx, record.result["pptxRef"])
    assert len(Presentation(io.BytesIO(deck.data)).slides) == 2 == record.result["slideCount"]
    assert not list((tmp_path / "jobs").iterdir()), "every job workspace is cleaned up"


async def test_brand_extract_and_render_layouts_jobs(tmp_path: Path):
    store, queue, worker, ctx = await setup(tmp_path)
    master = (await store.blobs.put(ctx, SYNTHETIC_MASTER.read_bytes(), "application/octet-stream")).ref

    job = await queue.enqueue(ctx, "brand.extract", {"pptxRef": master, "captureAllLayouts": True})
    record = await finished(store, worker, ctx, job.id)
    brand = json.loads((await store.blobs.get(ctx, record.result["brandRef"])).data)
    assert brand["capturedFurniture"]["layoutsByIndex"]

    job = await queue.enqueue(ctx, "masters.render_layouts", {"pptxRef": master, "width": 640})
    record = await finished(store, worker, ctx, job.id)
    layouts = await store.blobs.get(ctx, record.result["layoutsRef"])
    assert len(zipfile.ZipFile(io.BytesIO(layouts.data)).namelist()) == 11


async def test_bad_input_fails_the_job_permanently_with_a_safe_message(tmp_path: Path):
    store, queue, worker, ctx = await setup(tmp_path)
    job = await queue.enqueue(ctx, "slides.design_and_export", {"spec": {"slide": {"title": ""}}})
    assert await worker.run_once() is True
    record = await store.jobs.get(ctx, job.id)
    assert record.status == "error" and "title" in (record.error or "")
    job = await queue.enqueue(ctx, "slides.design_and_export",
                              {"spec": SPEC, "brand": {"primary_color": "#000", "secret_field": 1}})
    assert await worker.run_once() is True
    record = await store.jobs.get(ctx, job.id)
    assert record.status == "error" and "unknown brand fields" in (record.error or "")


async def test_the_job_types_join_the_shared_registry():
    registry = JobRegistry()
    jobs.register(registry, jobs.PipelineDeps(blobs=FakeGeneralService().blobs, scratch_root=Path("."),
                                              settings=build_ai_settings()))
    assert set(registry.types) == {"slides.design_and_export", "exports.stitch_deck", "brand.extract",
                                   "masters.render_layouts"}
    assert registry.types["slides.design_and_export"].expensive
    with pytest.raises(ValueError):
        jobs.register(registry, jobs.PipelineDeps(blobs=FakeGeneralService().blobs, scratch_root=Path("."),
                                                  settings=build_ai_settings()))
