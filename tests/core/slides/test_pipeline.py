"""`design_and_export` end to end on the fake provider, with the real engine export (Chromium).

The flows: spec -> HTML -> PPTX, image -> HTML -> PPTX, debrand, the layout template background,
the workzone enforced by a repair turn, and the $0 stub. Each checks the `.pptx` opens, that
`element_count` and `review_flag_count` are filled (never null), and that the cost is the turn's.
The design review is stubbed here except where the workzone needs the static check;
`tests/core/design/test_e2e_browser.py` runs the real one.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from app.core import engine_service, preview
from app.core.brand import extract as brand_extract
from app.core.brand.workzone import Workzone
from app.core.design.context import DesignServices
from app.core.llm import history
from app.core.llm.types import RoundRequest
from app.core.slides import pipeline
from app.core.slides.spec import BrandOptions
from app.engine.renderer import BlankLayoutsRenderer
from tests.conftest import build_ai_settings
from tests.core.design.conftest import SYNTHETIC_MASTER, png_bytes, slide_html, stub_services
from tests.fakes.llm import ReactiveProvider, Round, first_user_text, last_tool_results, resolver, text, tool


@pytest.fixture(scope="module", autouse=True)
def close_browsers() -> Iterator[None]:
    yield
    engine_service.shutdown()
    preview.shutdown()


@pytest.fixture(scope="module")
def furniture() -> dict[str, Any]:
    return dict(brand_extract.extract(SYNTHETIC_MASTER.read_bytes())["capturedFurniture"])


def layout_from(request: RoundRequest) -> str:
    said = first_user_text(request)
    marker = 'Save it on the layout with layout_id "'
    return said.split(marker, 1)[1].split('"', 1)[0]


def designer(html_top: int = 150, captured: list[RoundRequest] | None = None) -> ReactiveProvider:
    def respond(request: RoundRequest) -> Round:
        if captured is not None:
            captured.append(request)
        if last_tool_results(request):
            return Round(text("Done."))
        said = first_user_text(request)
        title = said.split("The slide's title is: ", 1)[1].split(". Emit", 1)[0] if "title is: " in said else "Slide"
        return Round(tool("save_slide", {"title": title, "layout_id": layout_from(request),
                                         "html": slide_html(title, top=html_top)}))
    return ReactiveProvider(respond)


def run(tmp_path: Path, req: pipeline.DesignRequest, provider: Any, services: DesignServices | None = None,
        **settings: Any) -> pipeline.PipelineResult:
    return pipeline.design_and_export(req, tmp_path / "job", settings=build_ai_settings(**settings),
                                      resolve=resolver(provider), services=services or stub_services(),
                                      renderer=BlankLayoutsRenderer())


def check_pptx(result: pipeline.PipelineResult) -> Any:
    assert result.pptx.is_file() and result.pptx.name == "slide.pptx"
    prs = Presentation(str(result.pptx))
    assert len(prs.slides) == 1
    assert result.element_count > 0, "element_count is filled, not null"
    assert isinstance(result.review_flag_count, int)
    assert result.review_flag_count == len(result.review_flags)
    assert (result.pptx.parent / "deck.zip").is_file(), "the design deck is kept for stitching"
    return prs


SPEC = {"slide": {"title": "Revenue grew 23% across three markets", "type": "data", "framework": "column chart",
                  "bullets": ["Market A +31%", "Market B +22%", "Market C +16%"],
                  "chartData": {"type": "column", "categories": ["A", "B", "C"],
                                "series": [{"name": "Growth", "values": [31, 22, 16]}]}},
        "deck": {"presentationTitle": "Growth review", "audience": "board"}}


def test_spec_to_html_to_pptx(tmp_path: Path, furniture: dict[str, Any]):
    seen: list[RoundRequest] = []
    provider = designer(captured=seen)
    result = run(tmp_path, pipeline.DesignRequest(spec=SPEC, furniture=furniture), provider)

    prs = check_pptx(result)
    assert result.archetype == "content" and result.master_mode == "branded"
    instruction = first_user_text(seen[0])
    assert "storyline spec" in instruction and "Market A +31%" in instruction
    assert '"categories":["A","B","C"]' in instruction.replace(" ", ""), "the exact chart data is handed over"
    assert "Pre-filled brief" in instruction
    assert result.cost_usd > 0 and result.cost_usd == pytest.approx(sum(t.cost_usd for t in result.turns))
    title = next(s for s in prs.slides[0].shapes if s.is_placeholder and "Revenue" in (s.text_frame.text or ""))
    assert title is not None, "the title lands in the master's title placeholder"
    assert result.to_json()["elementCount"] == result.element_count


def test_image_to_html_to_pptx(tmp_path: Path, furniture: dict[str, Any]):
    seen: list[RoundRequest] = []
    image = png_bytes((640, 360), "#335577")
    req = pipeline.DesignRequest(image=image, image_name="slide.png", furniture=furniture,
                                 brand=BrandOptions(primary_color="#1F2A44", chart_data='{"type":"bar"}'))
    result = run(tmp_path, req, designer(captured=seen))

    check_pptx(result)
    first = seen[0]
    attachments = [b for m in first.messages for b in (m.get("content") or []) if isinstance(b, dict)
                   and b.get("type") == "attachment"]
    assert attachments and attachments[0]["kind"] == "image"
    wire = history.to_anthropic(first.messages, "claude-opus-5-5", first.attachments_root)
    images = [b for m in wire for b in m["content"] if isinstance(b, dict) and b.get("type") == "image"
              and b["source"].get("media_type") == "image/png"]
    assert images, "the picture reaches the model inline"
    said = first_user_text(first)
    assert "Recreate it as ONE editable HTML slide" in said and "#1F2A44" in said and '{"type":"bar"}' in said


def test_debrand_mode_carries_no_brand_furniture(tmp_path: Path, furniture: dict[str, Any]):
    seen: list[RoundRequest] = []
    req = pipeline.DesignRequest(image=png_bytes(), furniture=furniture,
                                 brand=BrandOptions(apply_brand_layout=False, primary_color="#FF0000"))
    result = run(tmp_path, req, designer(captured=seen))
    prs = check_pptx(result)
    assert result.master_mode == "debranded"
    layout = prs.slides[0].slide_layout
    assert all(s.is_placeholder for s in layout.shapes), "no furniture shapes on the layout"
    said = first_user_text(seen[0])
    assert "carries no brand furniture" in said and "#FF0000" not in said


def test_the_layout_template_is_painted_as_the_layout_background(tmp_path: Path, furniture: dict[str, Any]):
    req = pipeline.DesignRequest(spec=SPEC, furniture=furniture,
                                 brand=BrandOptions(layout_template=png_bytes((320, 180), "#AA0000")))
    result = run(tmp_path, req, designer())
    prs = check_pptx(result)
    bg = prs.slides[0].slide_layout.element.find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    assert bg is not None and "blipFill" in bg.xml.replace("a:", "")
    assert result.template_background and result.to_json()["templateBackground"] is True

    off = pipeline.DesignRequest(spec=SPEC, furniture=furniture,
                                 brand=BrandOptions(layout_template=png_bytes(), apply_template_bg=False))
    assert run(tmp_path / "off", off, designer()).template_background is False
    debrand = pipeline.DesignRequest(spec=SPEC, furniture=furniture,
                                     brand=BrandOptions(layout_template=png_bytes(), apply_brand_layout=False))
    assert run(tmp_path / "debrand", debrand, designer()).template_background is False, "never in debrand mode"


INLINE = ("<!doctype html><html><head><meta charset=\"utf-8\"><style>html,body{margin:0;width:1280px;height:720px;"
          "overflow:hidden;background:transparent}</style></head><body>"
          "<h1 data-placeholder=\"title\" style=\"position:absolute;left:64px;top:20px;width:1100px;height:50px;"
          "margin:0;font:700 26px Arial\">Revenue grew 23%</h1>"
          "<div style=\"position:absolute;left:64px;top:40px;width:1100px;height:300px;font:16px Arial\">"
          "<p>Market A grew 31% on volume.</p></div></body></html>")


def test_the_workzone_is_enforced_by_one_repair_turn(tmp_path: Path, furniture: dict[str, Any]):
    zone = Workzone(left=0.05, top=0.2, width=0.9, height=0.7)
    calls: list[RoundRequest] = []

    def respond(request: RoundRequest) -> Round:
        calls.append(request)
        said = first_user_text(request)
        results = last_tool_results(request)
        if "still breaks the brand's workzone" in said.split("<workspace>")[-1] and not results:
            sid = said.split("The slide you saved (", 1)[1].split(")", 1)[0]
            return Round(tool("edit_slide", {"slide_id": sid,
                                             "edits": [{"find": "top:40px", "replace": "top:160px"}]}))
        if results:
            return Round(text("Done."))
        title = "Revenue grew 23%"
        return Round(tool("save_slide", {"title": title, "layout_id": layout_from(request), "html": INLINE}))

    # no browser review here: the static check reads the inline boxes (the measured one is in
    # tests/core/design/test_e2e_browser.py)
    services = DesignServices(lint=engine_service.readiness, review=None, render=lambda _d, _s: png_bytes())
    req = pipeline.DesignRequest(spec=SPEC, furniture=furniture, brand=BrandOptions(workzone=zone))
    result = run(tmp_path, req, ReactiveProvider(respond), services=services)

    assert len(result.turns) == 2, "one repair turn after the first save broke the workzone"
    assert "WORKZONE CONSTRAINT" in first_user_text(calls[0])
    deck_slide = (result.deck_dir / "slides" / result.slide_id / "v002.html").read_text(encoding="utf-8")
    assert "top:160px" in deck_slide
    assert not [f for f in result.review_flags if f.get("rule") in ("workzone-overflow", "header-band-overlap")]
    check_pptx(result)


def test_the_stub_path_makes_no_model_call(tmp_path: Path, furniture: dict[str, Any]):
    provider = ReactiveProvider(lambda _r: pytest.fail("the stub path must not call a model"))  # type: ignore[arg-type,return-value]
    result = run(tmp_path, pipeline.DesignRequest(spec=SPEC, furniture=furniture), provider, design_stub=True)
    check_pptx(result)
    assert result.cost_usd == 0 and result.turns == []


def test_bad_requests_are_refused_before_any_call(tmp_path: Path):
    provider = ReactiveProvider(lambda _r: pytest.fail("no call"))  # type: ignore[arg-type,return-value]
    for req in (pipeline.DesignRequest(), pipeline.DesignRequest(spec=SPEC, image=b"x"),
                pipeline.DesignRequest(spec={"slide": {"title": ""}}),
                pipeline.DesignRequest(image=b"x", image_name="slide.gif"),
                pipeline.DesignRequest(spec=SPEC, brand=BrandOptions(layout_archetype="poster"))):
        with pytest.raises(pipeline.PipelineError):
            run(tmp_path / str(id(req)), req, provider)


def test_a_cover_spec_uses_the_cover_layout_and_a_short_instruction(tmp_path: Path, furniture: dict[str, Any]):
    seen: list[RoundRequest] = []
    spec = {"slide": {"title": "Growth review 2026", "type": "title"}}
    result = run(tmp_path, pipeline.DesignRequest(spec=spec, furniture=furniture,
                                                  brand=BrandOptions(date="October 2026")), designer(captured=seen))
    assert result.archetype == "cover"
    said = first_user_text(seen[0])
    assert "cover" in said and "Date line: October 2026." in said
    deck = io.BytesIO(result.pptx.read_bytes())
    assert Presentation(deck).slides[0].slide_layout.name == "Cover"
