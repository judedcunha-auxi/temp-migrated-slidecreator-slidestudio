"""core/brand/synth_master + furniture_inject: capturedFurniture -> a master the engine imports.

Round trip on synthetic masters: extract -> capturedFurniture -> synthesize_master -> python-pptx
opens it at the captured size, the archetype layouts carry their names, furniture and theme, and the
engine's importer reads it (no renderer: `render=False`).
"""

from __future__ import annotations

import base64
import io
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from app.config.engine import MASTERS_FIXTURE
from app.core.brand import extract as b
from app.core.brand import furniture_inject as fi
from app.core.brand.synth_master import ARCHETYPE_NAME, synthesize_master
from app.engine.importer import import_master
from tests.core.brand.helpers import add_furniture_picture, png_bytes

MASTERS = [MASTERS_FIXTURE / "test-16x9.pptx", MASTERS_FIXTURE / "test-4x3.pptx",
           MASTERS_FIXTURE / "extra" / "synthetic-widescreen.pptx"]


def _deck_with_furniture() -> bytes:
    """Default template with a logo picture on the master and a band on the content layout."""
    prs: Any = Presentation()
    w, h = prs.slide_width, prs.slide_height
    add_furniture_picture(prs, prs.slide_master, png_bytes((10, 120, 200)), (0.88, 0.03, 0.1, 0.05))
    content = prs.slide_master.slide_layouts[1]
    content.shapes._spTree.append(_band(w, h))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _band(w: int, h: int) -> Any:
    """A filled rectangle (a footer band) as raw `<p:sp>` XML: layouts have no add_shape."""
    from lxml import etree

    xml = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:nvSpPr><p:cNvPr id="90" name="Band"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="0" y="{int(h * 0.95)}"/><a:ext cx="{w}" cy="{int(h * 0.05)}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="0A78C8"/></a:solidFill>'
        '</p:spPr></p:sp>'
    )
    return etree.fromstring(xml)


@pytest.mark.parametrize("master", MASTERS, ids=lambda p: p.stem)
def test_round_trip_opens_at_the_captured_size_and_imports(master: Path, tmp_path: Path) -> None:
    captured = b.extract_file(master)["capturedFurniture"]
    out = synthesize_master(captured, tmp_path / "synth" / "master.pptx")
    assert out == tmp_path / "synth" / "master.pptx" and out.is_file()

    prs = Presentation(str(out))
    assert prs.slide_width == captured["slideWidthEmu"] and prs.slide_height == captured["slideHeightEmu"]
    names = [lay.name for lay in prs.slide_master.slide_layouts]
    assert {"Cover", "Divider", "Content"} <= set(names)

    manifest = import_master(out, tmp_path / "import", render=False)
    assert {lay.name for lay in manifest.layouts} >= {"Cover", "Divider", "Content"}
    content = next(lay for lay in manifest.layouts if lay.name == "Content")
    assert any(ph.type == "title" for ph in content.placeholders)
    # The brand theme came across: the manifest's fonts are the master's.
    assert manifest.theme["fonts"]["major"] == b.extract_file(master)["headingFont"]


def test_heading_box_moves_the_title_placeholder(tmp_path: Path) -> None:
    captured = b.extract_file(MASTERS[0])["capturedFurniture"]
    heading = {"left": 0.1, "top": 0.2, "width": 0.5, "height": 0.1, "font": "Georgia", "size": 30,
               "color": "#112233"}
    captured["layouts"]["content"]["placeholders"] = {"heading": heading}
    prs = Presentation(str(synthesize_master(captured, tmp_path / "m.pptx")))
    layout = next(lay for lay in prs.slide_master.slide_layouts if lay.name == "Content")
    title = next(ph for ph in layout.placeholders if ph.placeholder_format.type in (1, 3))
    assert abs(title.left / prs.slide_width - 0.1) < 0.001
    assert abs(title.top / prs.slide_height - 0.2) < 0.001
    font = title.text_frame.paragraphs[0].font
    assert font.name == "Georgia" and font.size.pt == 30 and str(font.color.rgb) == "112233"


def test_furniture_and_media_are_injected(tmp_path: Path) -> None:
    captured = b.extract(_deck_with_furniture())["capturedFurniture"]
    assert captured["media"], "the master's logo picture is captured as media"
    out = synthesize_master(captured, tmp_path / "m.pptx")
    prs = Presentation(str(out))
    master_pictures = [s for s in prs.slide_master.shapes if s.shape_type == 13]
    assert len(master_pictures) == 1
    assert master_pictures[0].image.blob[:8] == b"\x89PNG\r\n\x1a\n"
    content = next(lay for lay in prs.slide_master.slide_layouts if lay.name == "Content")
    assert any(not s.is_placeholder for s in content.shapes)


def test_a_capture_without_layouts_gives_a_clean_master(tmp_path: Path) -> None:
    captured = {"slideWidthEmu": 12192000, "slideHeightEmu": 6858000, "layouts": {}}
    prs = Presentation(str(synthesize_master(captured, tmp_path / "plain.pptx")))
    assert prs.slide_width == 12192000
    assert not set(ARCHETYPE_NAME.values()) & {lay.name for lay in prs.slide_master.slide_layouts}


def test_inject_scales_shapes_to_another_canvas() -> None:
    captured = b.extract(_deck_with_furniture())["capturedFurniture"]
    blob = next(s for s in captured["layouts"]["content"]["shapes"] if s.get("origin") == "layout")
    prs: Any = Presentation()
    layout = prs.slide_master.slide_layouts[6]
    before = len(layout.shapes)
    w, h = captured["slideWidthEmu"], captured["slideHeightEmu"]
    added = fi.inject_furniture_into_part(layout.part, layout.shapes._spTree, [blob], captured["media"],
                                          w, h, w // 2, h // 2)
    assert added == 1 and len(layout.shapes) == before + 1
    shape = list(layout.shapes)[-1]
    assert abs(shape.width - w // 2) <= 1


def test_a_bad_shape_blob_is_skipped_not_fatal() -> None:
    prs: Any = Presentation()
    layout = prs.slide_master.slide_layouts[6]
    bad = {"xmlB64": base64.b64encode(b"<not-xml").decode(), "rIdMap": {}}
    assert fi.inject_furniture_into_part(layout.part, layout.shapes._spTree, [bad], {}, 1, 1, 1, 1) == 0


def test_entities_in_captured_xml_are_not_expanded() -> None:
    prs: Any = Presentation()
    layout = prs.slide_master.slide_layouts[6]
    hostile = (b'<?xml version="1.0"?><!DOCTYPE p [<!ENTITY x "EXPANDED">]>'
               b'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">&x;</p:sp>')
    blob = {"xmlB64": base64.b64encode(hostile).decode(), "rIdMap": {}}
    fi.inject_furniture_into_part(layout.part, layout.shapes._spTree, [blob], {}, 1, 1, 1, 1)
    assert b"EXPANDED" not in layout.part.blob


def test_resolve_layout_falls_back_to_content_then_first() -> None:
    assert fi.resolve_layout({"layouts": {}}, "cover") is None
    assert fi.resolve_layout({"layouts": {"content": {"a": 1}}}, "cover") == {"a": 1}
    assert fi.resolve_layout({"layouts": {"divider": {"b": 2}}}, "cover") == {"b": 2}
    assert fi.resolve_layout({"layouts": {"cover": {"c": 3}}}, "cover") == {"c": 3}


def test_shape_identity_depends_on_xml_and_media_only() -> None:
    one = {"xmlB64": "AAAA", "rIdMap": {"rId2": "b", "rId1": "a"}, "origin": "layout"}
    two = {"xmlB64": "AAAA", "rIdMap": {"rId9": "a", "rId8": "b"}, "origin": "master"}
    assert fi.shape_identity(one) == fi.shape_identity(two)
    assert fi.shape_identity(one) != fi.shape_identity({**one, "xmlB64": "BBBB"})


def test_theme_is_installed_from_the_capture() -> None:
    captured = b.extract_file(MASTERS[2])["capturedFurniture"]
    prs: Any = Presentation()
    assert fi.apply_brand_theme(prs, captured)
    theme = prs.slide_masters[0].part.part_related_by(fi._THEME_RELTYPE)
    assert theme.blob == base64.b64decode(captured["theme"]["xmlB64"])
    assert not fi.apply_brand_theme(Presentation(), {"theme": None})
