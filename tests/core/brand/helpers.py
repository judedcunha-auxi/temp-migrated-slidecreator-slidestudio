"""Synthetic decks for the brand tests: furniture pictures on masters and layouts.

python-pptx only adds pictures to slides, so a picture is added to a scratch slide, and its `<p:pic>`
and image relationship are moved onto the master or layout part.
"""

from __future__ import annotations

import copy
import io
from typing import Any

from PIL import Image
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu

_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def png_bytes(colour: tuple[int, int, int] = (200, 30, 30), size: tuple[int, int] = (40, 20)) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", size, colour).save(out, format="PNG")
    return out.getvalue()


def add_furniture_picture(prs: Any, target: Any, png: bytes, box: tuple[float, float, float, float]) -> None:
    """Fix a picture into `target` (a slide master or layout) at `box` (fractions of the slide)."""
    w, h = prs.slide_width, prs.slide_height
    left, top, width, height = box
    scratch = prs.slides.add_slide(prs.slide_layouts[6])
    pic = scratch.shapes.add_picture(io.BytesIO(png), Emu(int(w * left)), Emu(int(h * top)),
                                     Emu(int(w * width)), Emu(int(h * height)))
    image_part = scratch.part.related_part(pic._element.find(".//" + qn("a:blip")).get(_R_EMBED))
    element = copy.deepcopy(pic._element)
    element.find(".//" + qn("a:blip")).set(_R_EMBED, target.part.relate_to(image_part, RT.IMAGE))
    target.shapes._spTree.append(element)
    # drop the scratch slide again
    slide_ids = prs.slides._sldIdLst
    rid = slide_ids[-1].get(qn("r:id"))
    slide_ids.remove(slide_ids[-1])
    prs.part.drop_rel(rid)
