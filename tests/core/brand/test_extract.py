"""core/brand/extract: the brand read out of a master, its errors, the role pass and per-layout capture.

Built on the synthetic masters under tests/engine/fixtures/masters and on python-pptx's default
template; no client material and no model calls (the role pass gets a fake annotator).
Ported from Slide Studio `server/tests/test_brand_extract_layouts.py` (the route test is dropped:
there is no `/v1/brand-extract` route any more).
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from app.config.engine import MASTERS_FIXTURE
from app.core.brand import extract as b
from tests.core.brand.helpers import add_furniture_picture, png_bytes

MASTER_16X9 = MASTERS_FIXTURE / "test-16x9.pptx"
MASTER_4X3 = MASTERS_FIXTURE / "test-4x3.pptx"
WIDESCREEN = MASTERS_FIXTURE / "extra" / "synthetic-widescreen.pptx"


def _default_bytes() -> bytes:
    buf = io.BytesIO()
    Presentation().save(buf)
    return buf.getvalue()


def _layout(prs: Any, name: str) -> Any:
    return next(lay for lay in prs.slide_master.slide_layouts if lay.name == name)


def _retype_placeholders_as_body(layout: Any) -> None:
    """Mimic templates that declare every placeholder as BODY (no <p:ph type="title"/>).

    Geometry is pinned first: default-template layout placeholders inherit their position from the
    master BY TYPE, so retyping would otherwise move them.
    """
    for ph in layout.placeholders:
        ph.left, ph.top, ph.width, ph.height = ph.left, ph.top, ph.width, ph.height
    for ph in layout.placeholders:
        ph._element.find(".//" + qn("p:ph")).set("type", "body")
        if ph.has_text_frame:
            ph.text_frame.text = ""


def _heading(layout: Any, prs: Any) -> dict[str, Any]:
    heading: dict[str, Any] = b._capture_heading_placeholders(
        layout, prs.slide_width, prs.slide_height, "content", {})["heading"]
    return heading


# ------------------------------------------------------------------------------ the brand itself


@pytest.mark.parametrize("master", [MASTER_16X9, MASTER_4X3, WIDESCREEN], ids=lambda p: p.stem)
def test_extract_reads_theme_layouts_and_workzone(master: Path) -> None:
    brand = b.extract_file(master)
    prs = Presentation(str(master))
    assert brand["headingFont"] and brand["bodyFont"]
    assert re.fullmatch(r"#[0-9A-F]{6}", brand["primaryColor"])
    assert re.fullmatch(r"#[0-9A-F]{6}", brand["accentColor"])
    assert {c["role"] for c in brand["allColors"]} >= {"Dark 1", "Light 1", "Accent 1"}
    assert [lay["name"] for lay in brand["layouts"]] == [lay.name for lay in prs.slide_master.slide_layouts]
    first = brand["layouts"][0]
    assert first["placeholders"] and all(0 <= p["left"] <= 1 for p in first["placeholders"])
    zone = brand["workzone"]
    assert zone is not None and 0 <= zone["left"] < zone["right"] <= 1 and 0 <= zone["top"] < zone["bottom"] <= 1
    assert brand["typographyScale"]["title"]["sizePt"]
    assert brand["masterPng"] is None
    furniture = brand["capturedFurniture"]
    assert furniture["slideWidthEmu"] == prs.slide_width and furniture["slideHeightEmu"] == prs.slide_height
    assert set(furniture["layouts"]) == {"cover", "divider", "content"}
    assert furniture["theme"]["xmlB64"] and "layoutsByIndex" not in furniture


def test_extract_adds_layouts_by_index_only_when_asked() -> None:
    content = _default_bytes()
    default = b.extract(content)
    assert "layoutsByIndex" not in default["capturedFurniture"]

    full = b.extract(content, capture_all_layouts=True)
    furniture = full["capturedFurniture"]
    assert len(furniture["layoutsByIndex"]) == len(Presentation(io.BytesIO(content)).slide_master.slide_layouts)
    assert "layoutsByIndex" in furniture["background"]
    # The archetype-keyed capture is unchanged alongside it.
    assert set(furniture["layouts"]) == set(default["capturedFurniture"]["layouts"])


def test_extract_is_json_serialisable() -> None:
    json.dumps(b.extract_file(WIDESCREEN, capture_all_layouts=True))


# ----------------------------------------------------------------------------------- bad inputs


def test_empty_file_is_refused() -> None:
    with pytest.raises(b.BrandExtractError, match="empty"):
        b.extract(b"")


def test_a_file_over_the_cap_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    content = MASTER_16X9.read_bytes()
    monkeypatch.setattr(b, "MAX_BYTES", len(content) - 1)
    with pytest.raises(b.BrandExtractError, match="too large"):
        b.extract(content)


def test_a_file_that_is_not_a_zip_is_refused() -> None:
    with pytest.raises(b.BrandExtractError, match="not a .pptx"):
        b.extract(b"%PDF-1.7 not a deck")


def test_a_zip_that_is_not_a_deck_is_refused() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("hello.txt", "not a presentation")
    with pytest.raises(b.BrandExtractError, match="could not be opened") as caught:
        b.extract(buf.getvalue())
    assert "hello" not in str(caught.value)  # the message names no internals


def test_brand_extract_error_is_a_value_error() -> None:
    assert issubclass(b.BrandExtractError, ValueError)


# ------------------------------------------------------------------------- the role annotation


def _deck_with_logo(layouts: int = 3) -> bytes:
    """Default template with the same small picture fixed in the same corner of `layouts` layouts."""
    prs: Any = Presentation()
    png = png_bytes()
    for layout in list(prs.slide_master.slide_layouts)[:layouts]:
        add_furniture_picture(prs, layout, png, (0.85, 0.9, 0.1, 0.05))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


class FakeAnnotator:
    """Labels the TITLE placeholders "title" and every decoration picture "logo"."""

    def __init__(self) -> None:
        self.prompts: list[str] = []

    def __call__(self, prompt: str) -> str:
        self.prompts.append(prompt)
        shapes = json.loads(prompt.split("Shapes:\n", 1)[1])
        roles = []
        for shape in shapes:
            if shape["pptxType"] == "TITLE":
                roles.append({"id": shape["id"], "role": "title"})
            elif shape["isDecoration"] and shape["pptxType"] == "PICTURE":
                roles.append({"id": shape["id"], "role": "logo"})
        return json.dumps({"roles": roles})


def test_no_annotator_leaves_roles_unset() -> None:
    brand = b.extract(_deck_with_logo())
    roles = [p["role"] for lay in brand["layouts"] for p in lay["placeholders"] + lay["decorations"]]
    assert roles and all(role is None for role in roles)


def test_an_annotator_labels_placeholders_and_finds_the_logo() -> None:
    annotate = FakeAnnotator()
    brand = b.extract(_deck_with_logo(), annotate_roles=annotate)
    assert len(annotate.prompts) == 1
    assert annotate.prompts[0].startswith(b._LAYOUT_ROLE_PROMPT)
    titled = [p for lay in brand["layouts"] for p in lay["placeholders"] if p["type"] == "TITLE"]
    assert titled and all(p["role"] == "title" for p in titled)
    logos = [d for lay in brand["layouts"] for d in lay["decorations"] if d["type"] == "PICTURE"]
    assert len(logos) == 3 and all(d["role"] == "logo" for d in logos)
    # The master has no pictures, so the layout-derived logo is the brand's logo.
    assert len(brand["logoShapes"]) == 1
    assert brand["logoShapes"][0]["role"] in ("logo", "logo_mark")
    assert set(brand["logoShapes"][0]) == {"left", "top", "width", "height", "role"}


@pytest.mark.parametrize("reply", [RuntimeError("provider down"), "not json", '{"roles": "nope"}', '[1, 2]'])
def test_a_failing_annotator_leaves_the_extraction_intact(reply: object) -> None:
    def annotate(_prompt: str) -> str:
        if isinstance(reply, Exception):
            raise reply
        return str(reply)

    content = _deck_with_logo()
    plain = b.extract(content)
    brand = b.extract(content, annotate_roles=annotate)
    assert brand["layouts"] == plain["layouts"]
    assert brand["logoShapes"] == plain["logoShapes"]
    assert brand["capturedFurniture"] == plain["capturedFurniture"]


# ------------------------------------------------------------------- heading placeholder choice


def test_declared_title_still_wins() -> None:
    prs = Presentation()
    lay = _layout(prs, "Title and Content")
    title = next(ph for ph in lay.placeholders if ph.placeholder_format.type.name == "TITLE")
    got = _heading(lay, prs)
    assert round(got["top"], 3) == round(title.top / prs.slide_height, 3)


def test_all_body_layout_picks_top_wide_box_not_the_larger_body() -> None:
    prs = Presentation()
    lay = _layout(prs, "Title and Content")
    title_top = min(ph.top for ph in lay.placeholders) / prs.slide_height
    _retype_placeholders_as_body(lay)
    got = _heading(lay, prs)
    # The largest box (the content body) used to win; the title strip is the top-most wide box.
    assert round(got["top"], 3) == round(title_top, 3)


def test_narrow_or_short_boxes_fall_back_to_largest() -> None:
    prs = Presentation()
    lay = _layout(prs, "Title and Content")
    _retype_placeholders_as_body(lay)
    for ph in lay.placeholders:  # shrink everything below the positional thresholds
        ph.width = int(prs.slide_width * 0.3)
    biggest = max(lay.placeholders, key=lambda ph: ph.width * ph.height)
    got = _heading(lay, prs)
    assert round(got["top"], 3) == round(biggest.top / prs.slide_height, 3)


def test_heading_inherits_theme_font_and_title_style() -> None:
    brand = b.extract_file(MASTER_16X9)
    heading = brand["capturedFurniture"]["layouts"]["content"]["placeholders"]["heading"]
    assert heading["font"] == brand["headingFont"]
    assert heading["size"] == brand["typographyScale"]["title"]["sizePt"]


def test_capture_all_layouts_keys_every_layout_and_shares_media() -> None:
    prs = Presentation()
    media: dict[str, Any] = {}
    by_index, backgrounds = b._capture_all_layouts(prs, media, {})
    n = len(prs.slide_master.slide_layouts)
    assert sorted(by_index, key=int) == [str(i) for i in range(n)]
    assert set(backgrounds) <= set(by_index)
    entry = by_index["1"]
    assert entry["sourceIndex"] == 1
    assert entry["sourceName"] == prs.slide_master.slide_layouts[1].name
    assert "heading" in entry["placeholders"]


def test_archetypes_map_by_name_then_by_position() -> None:
    prs = Presentation(str(WIDESCREEN))
    mapping = b._map_layouts_to_archetypes(prs)
    names = [lay.name for lay in prs.slide_master.slide_layouts]
    assert names[mapping["cover"]] == "Title"
    assert names[mapping["divider"]] == "Divider"
    assert len(set(mapping.values())) == 3  # content never reuses the divider's layout
