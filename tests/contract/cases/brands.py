"""Contract cases: Darwin's brand routes (data/api-brands.json, api-brand-*.json).

Cast: `alice` owns personal brands; `admin1` is an admin who created the org brand "Acme Org", mapped to
acme.com; `bob` is a member of that domain (verified email), so he may READ it but not edit it; `mallory`
is nobody's member.
"""

from __future__ import annotations

import base64
import io
import json
from typing import Any

from pptx import Presentation

from app.core.brand.kit import DEFAULT_STYLE_TEMPLATE
from app.core.darwin.brand_jobs import BRAND_PPTX_JOB, BrandJobDeps, brand_pptx_handler
from tests.contract.cases.common import add_auth, broken
from tests.contract.harness import Probe, RouteCases
from tests.fakes.darwin import DarwinEnv, tiny_pptx
from tests.fakes.image_gen import tiny_png
from tests.fakes.renderer import FakeRenderer
from tests.fakes.storyline_model import reply

UNKNOWN = "3fa85f64-5717-4562-b3fc-2c963f66afa6"
PNG = tiny_png()
PDF = b"%PDF-1.4\n%fake guidelines\n"
ORG_DOMAIN = "acme.com"

#: A synthetic per-layout furniture capture (the extractor's `capturedFurniture` with layoutsByIndex).
CAPTURE: dict[str, Any] = {
    "media": {"sha-a": {"b64": "AAAA"}, "sha-b": {"b64": "BBBB"}, "sha-unused": {"b64": "CCCC"}},
    "layouts": {"content": {"shapes": [], "placeholders": {"heading": {"left": 0.1, "top": 0.1, "width": 0.8,
                                                                       "height": 0.1}}}},
    "layoutsByIndex": {
        "0": {"sourceIndex": 0, "shapes": [{"rIdMap": {"rId1": "sha-a"}}],
              "placeholders": {"heading": {"left": 0.05, "top": 0.3, "width": 0.9, "height": 0.2, "font": "Georgia",
                                           "color": "#112233", "size": 40},
                               "date": {"left": 0.05, "top": 0.6, "width": 0.3, "height": 0.05}}},
        "2": {"sourceIndex": 2, "shapes": [], "placeholders": {"heading": {"left": 0.1, "top": 0.4, "width": 0.8,
                                                                           "height": 0.15}}},
        "5": {"sourceIndex": 5, "shapes": [{"rIdMap": {"rId9": "sha-b"}}],
              "placeholders": {"heading": {"left": 0.04, "top": 0.05, "width": 0.92, "height": 0.1},
                               "date": {"left": 0.0, "top": 0.0, "width": 0.1, "height": 0.1}}},
    },
    "background": {"master": {"fill": "white"}, "layoutsByIndex": {"2": {"fill": "navy"}}},
}
PREVIEWS = {f"layout-preview-{i}": PNG for i in (0, 2, 5)}


def j(value: Any) -> bytes:
    return json.dumps(value).encode("utf-8")


# ---------------------------------------------------------------------------------------- setup
def org_brand(env: DarwinEnv, **kw: Any) -> str:
    return env.seed_brand("admin1", name="Acme Org", org=True, domain=ORG_DOMAIN, **kw)


def bob(env: DarwinEnv) -> dict[str, str]:
    return env.member("bob", f"bob@{ORG_DOMAIN}")


def send(method: str, path: str, *, who: str = "alice", **kwargs: Any) -> Any:
    return lambda env: env.client.request(method, path, headers=env.auth(who), **kwargs)


def with_brand(build: Any, **seed: Any) -> Any:
    """A probe that seeds alice's brand first; `build(env, brand_id)` sends the request."""
    def run(env: DarwinEnv) -> Any:
        return build(env, env.seed_brand("alice", **seed))
    return run


def with_org(build: Any, **seed: Any) -> Any:
    def run(env: DarwinEnv) -> Any:
        return build(env, org_brand(env, **seed))
    return run


def exactly(expected: Any) -> Any:
    def check(_env: DarwinEnv, response: Any) -> None:
        assert response.json() == expected, response.text
    return check


def error_is(text: str) -> Any:
    return exactly({"error": text})


# ------------------------------------------------------------------------------------- /brands
brands = RouteCases("/api/brands")
add_auth(brands, "GET", "/api/brands")


def _listed(env: DarwinEnv, response: Any) -> None:
    body = response.json()["brands"]
    assert [b["name"] for b in body] == ["Acme Org", "Newer", "Older"]  # org first, then newest first
    assert [b["isOrg"] for b in body] == [True, False, False]
    assert body[1]["kit"] == {"primaryColor": "#1A2B3C", "junk": None}  # the raw stored kit, not normalised


def _list_three(env: DarwinEnv) -> Any:
    org_brand(env)
    env.seed_brand("bob", name="Older")
    env.store.clock.advance(seconds=5)
    env.seed_brand("bob", name="Newer", kit={"primaryColor": "#1A2B3C", "junk": None})
    return env.client.get("/api/brands", headers=bob(env))


brands.add("R0",
           Probe("org brand of the verified domain first, then own brands newest first", _list_three, _listed),
           Probe("no brands", send("GET", "/api/brands"), exactly({"brands": []})),
           Probe("an UNVERIFIED email does not reach the org brand", lambda env: (
               org_brand(env), env.user("eve", email="eve@acme.com"),
               env.client.get("/api/brands", headers=env.auth("eve", email="eve@acme.com", email_verified=False)))[2],
               exactly({"brands": []})))


def _created(env: DarwinEnv, response: Any) -> None:
    body = response.json()
    assert body["name"] == "Acme" and body["kit"] == {"primaryColor": "#1A2B3C", "layouts": [1], "company": None}
    assert env.brand("alice", body["id"]).kit == body["kit"]  # a null is stored literally on create


brands.add("R1",
           Probe("name trimmed, kit validated, nulls kept", send("POST", "/api/brands", json={
               "name": "  Acme  ", "kit": {"primaryColor": "#1A2B3C", "layouts": [1], "company": None}}), _created),
           Probe("no kit", send("POST", "/api/brands", json={"name": "Plain"})),
           Probe("an unknown key with a null value is accepted", send("POST", "/api/brands", json={
               "name": "x", "kit": {"furnitureArchetype": None}})))


def _patched(env: DarwinEnv, response: Any) -> None:
    body = response.json()
    owner = env.brand("alice", body["id"]).owner_id
    assert body["kit"] == {"accentColor": "#445566", "primaryColor": "#222222",
                           "masterKey": f"brand/{owner}/{body['id']}/master.png"}
    assert body["isOrg"] is False and body["name"] == "Renamed"


def _admin_patch_org(env: DarwinEnv, response: Any) -> None:
    body = response.json()
    owner = env.brand("admin1", body["id"]).owner_id
    assert body["isOrg"] is True and body["kit"]["logoKey"] == f"brand/{owner}/{body['id']}/logo.png"


brands.add("R2",
           Probe("shallow merge, null deletes, asset keys canonical", with_brand(
               lambda env, b: env.client.patch("/api/brands", headers=env.auth("alice"), json={
                   "brandId": b, "name": "Renamed",
                   "kit": {"primaryColor": "#222222", "company": None, "masterKey": True}}),
               kit={"primaryColor": "#111111", "accentColor": "#445566", "company": "Acme"}), _patched),
           Probe("kit {} counts as an update", with_brand(
               lambda env, b: env.client.patch("/api/brands", headers=env.auth("alice"), json={"brandId": b, "kit": {}}))),
           Probe("another admin edits an org brand: keys under the brand OWNER", lambda env: (
               b := org_brand(env), env.user("admin2", admin=True),
               env.client.patch("/api/brands", headers=env.auth("admin2"), json={"brandId": b, "kit": {"logoKey": 1}}))[2],
               _admin_patch_org))


def _deleted(env: DarwinEnv, response: Any) -> None:
    assert response.json() == {"ok": True}
    assert env.client.get("/api/brands", headers=env.auth("alice")).json() == {"brands": []}


brands.add("R3",
           Probe("own brand", with_brand(lambda env, b: env.client.delete(f"/api/brands?id={b}", headers=env.auth("alice"))),
                 _deleted),
           Probe("missing id: still ok", send("DELETE", f"/api/brands?id={UNKNOWN}"), exactly({"ok": True})),
           Probe("someone else's brand: a silent no-op", lambda env: (
               b := env.seed_brand("bob"), env.client.delete(f"/api/brands?id={b}", headers=env.auth("alice")),
               env.brand("bob", b))[1]),
           Probe("an org brand, even by an admin: a no-op", lambda env: (
               b := org_brand(env), env.client.delete(f"/api/brands?id={b}", headers=env.auth("admin1")),
               env.brand("admin1", b))[1]))
brands.add("E400 Invalid JSON body",
           *(Probe(f"POST {name}", send("POST", "/api/brands", content=raw)) for name, raw in (
               ("malformed", b"{nope"), ("empty", b""), ("null", b"null"), ("0", b"0"), ("false", b"false"),
               ('""', b'""'))),
           Probe("PATCH malformed", send("PATCH", "/api/brands", content=b"{nope")))
brands.add("E400 name is required (1-120 characters)",
           Probe("missing", send("POST", "/api/brands", json={"kit": {}})),
           Probe("blank", send("POST", "/api/brands", json={"name": "   "})),
           Probe("121 characters", send("POST", "/api/brands", json={"name": "x" * 121})),
           Probe("a truthy non-object body", send("POST", "/api/brands", json=5)),
           Probe("PATCH name null", with_brand(lambda env, b: env.client.patch(
               "/api/brands", headers=env.auth("alice"), json={"brandId": b, "name": None}))))
brands.add("E400 <key> cannot be set on create — upload the asset first, then PATCH",
           Probe("masterKey", send("POST", "/api/brands", json={"name": "x", "kit": {"masterKey": True}}),
                 error_is("masterKey cannot be set on create — upload the asset first, then PATCH")),
           Probe("furnitureKey even when null", send("POST", "/api/brands", json={"name": "x", "kit": {"furnitureKey": None}}),
                 error_is("furnitureKey cannot be set on create — upload the asset first, then PATCH")))
brands.add("E400 kit must be an object",
           *(Probe(f"kit {v!r}", send("POST", "/api/brands", json={"name": "x", "kit": v})) for v in (None, [1], "x", 3)))
brands.add("E400 kit too large (max 256KB)",
           Probe("262145 characters of JSON", send("POST", "/api/brands", json={
               "name": "x", "kit": {"styleTemplate": "x" * (256 * 1024)}})))
brands.add("E400 <key> must be a truthy value or null",
           *(Probe(f"logoKey {v!r}", with_brand(lambda env, b, v=v: env.client.patch(
               "/api/brands", headers=env.auth("alice"), json={"brandId": b, "kit": {"logoKey": v}})),
               error_is("logoKey must be a truthy value or null")) for v in (False, 0, "")))
brands.add("E400 furnitureKey must be a truthy value or null",
           Probe("furnitureKey false", with_brand(lambda env, b: env.client.patch(
               "/api/brands", headers=env.auth("alice"), json={"brandId": b, "kit": {"furnitureKey": False}}))))
brands.add('E400 furnitureArchetypes must be an array of "cover" | "divider" | "content"',
           *(Probe(f"{v!r}", send("POST", "/api/brands", json={"name": "x", "kit": {"furnitureArchetypes": v}}))
             for v in (["cover", "hero"], "cover", {})))
brands.add("E400 <key> must be a hex color like #1A2B3C",
           Probe("short hex", send("POST", "/api/brands", json={"name": "x", "kit": {"primaryColor": "#123"}}),
                 error_is("primaryColor must be a hex color like #1A2B3C")),
           Probe("neutralColor a number", send("POST", "/api/brands", json={"name": "x", "kit": {"neutralColor": 5}}),
                 error_is("neutralColor must be a hex color like #1A2B3C")))
brands.add("E400 <key> must be a string",
           Probe("headingFont a number", send("POST", "/api/brands", json={"name": "x", "kit": {"headingFont": 5}}),
                 error_is("headingFont must be a string")))
brands.add("E400 Unknown kit field: <key>",
           Probe("the first unknown key", send("POST", "/api/brands", json={
               "name": "x", "kit": {"primaryColor": "#111111", "layoutPreviews": [], "zzz": 1}}),
               error_is("Unknown kit field: layoutPreviews")))
brands.add("E400 brandId is required",
           *(Probe(f"brandId {v!r}", send("PATCH", "/api/brands", json={"brandId": v, "name": "x"})) for v in (None, "", 5)),
           Probe("a truthy non-object body", send("PATCH", "/api/brands", json=[1])))
brands.add("E404 Brand not found",
           Probe("unknown", send("PATCH", "/api/brands", json={"brandId": UNKNOWN, "name": "x"})),
           Probe("not a uuid", send("PATCH", "/api/brands", json={"brandId": "../etc", "name": "x"})),
           Probe("someone else's, even with a bad kit (access is checked first)", lambda env: env.client.patch(
               "/api/brands", headers=env.auth("alice"), json={"brandId": env.seed_brand("bob"), "kit": "bad"})),
           Probe("an org brand of a domain the caller is not in", with_org(lambda env, b: env.client.patch(
               "/api/brands", headers=env.member("mallory", "mallory@evil.com"), json={"brandId": b, "name": "x"}))))
brands.add("E403 Organization brands can only be edited by an admin",
           Probe("a domain member", with_org(lambda env, b: env.client.patch(
               "/api/brands", headers=bob(env), json={"brandId": b, "name": "x"}))))
brands.add("E400 Nothing to update — provide name and/or kit",
           Probe("neither", with_brand(lambda env, b: env.client.patch("/api/brands", headers=env.auth("alice"),
                                                                       json={"brandId": b}))))
brands.add("E400 id required",
           Probe("no id", send("DELETE", "/api/brands")), Probe("empty id", send("DELETE", "/api/brands?id=")))
brands.add("E405 GET, POST, PATCH or DELETE only",
           *(Probe(m, send(m, "/api/brands")) for m in ("PUT", "OPTIONS")))
brands.add("E500 Internal error",
           Probe("DELETE with a non-uuid id", send("DELETE", "/api/brands?id=not-a-uuid")),
           Probe("the store is unreachable", lambda env: (broken(env, "brands", "list_visible"),
                                                          send("GET", "/api/brands")(env))[1]))

# -------------------------------------------------------------------------------- /brand-asset
asset = RouteCases("/api/brand-asset")
add_auth(asset, "POST", "/api/brand-asset", json={"kind": "logo"})


def upload(body: dict[str, Any], who: str = "alice") -> Any:
    return with_brand(lambda env, b: env.client.post("/api/brand-asset", headers=env.auth(who),
                                                     json={"brandId": b, **body}), assets=PREVIEWS)


def _stored_master(env: DarwinEnv, response: Any) -> None:
    key = response.json()["key"]
    _, owner, brand_id, name = key.split("/")
    assert name == "master.png" and env.asset("alice", brand_id, "master") == PNG
    assert env.brand("alice", brand_id).owner_id == owner


def _admin_upload_org(env: DarwinEnv) -> Any:
    b = org_brand(env)
    env.user("admin2", admin=True)
    return env.client.post("/api/brand-asset", headers=env.auth("admin2"),
                           json={"brandId": b, "kind": "titleMaster", "b64": base64.b64encode(PNG).decode()})


def _under_owner(env: DarwinEnv, response: Any) -> None:
    _, owner, brand_id, _ = response.json()["key"].split("/")
    assert owner == env.brand("admin1", brand_id).owner_id  # admin1 created it, admin2 uploaded


asset.add("R0",
          Probe("b64 upload", upload({"kind": "master", "b64": base64.b64encode(PNG).decode()}), _stored_master),
          Probe("lenient base64 (whitespace, url-safe, data after '=')", upload(
              {"kind": "master", "b64": " ".join(base64.urlsafe_b64encode(PNG).decode()) + "=junk"}), _stored_master),
          Probe("copy a rendered layout", upload({"kind": "master", "fromLayoutIndex": 2}), _stored_master),
          Probe("fromLayoutIndex 2.0 is an integer", upload({"kind": "master", "fromLayoutIndex": 2.0})),
          Probe("an admin uploads to an org brand: stored under the owner", _admin_upload_org, _under_owner))
asset.add("R1",
          Probe("style-default", send("GET", "/api/brand-asset?kind=style-default"),
                exactly({"template": DEFAULT_STYLE_TEMPLATE})),
          Probe("any non-POST method is a GET", send("DELETE", "/api/brand-asset?kind=style-default")))


def _png_check(env: DarwinEnv, response: Any) -> None:
    assert response.content == PNG


asset.add("R2",
          Probe("index 2", with_brand(lambda env, b: env.client.get(
              f"/api/brand-asset?kind=layout-preview&brandId={b}&index=2", headers=env.auth("alice")), assets=PREVIEWS),
              _png_check),
          Probe("no index means index 0 (Number(null) === 0)", with_brand(lambda env, b: env.client.get(
              f"/api/brand-asset?kind=layout-preview&brandId={b}", headers=env.auth("alice")), assets=PREVIEWS)),
          Probe("a domain member reads an org brand's preview", with_org(lambda env, b: env.client.get(
              f"/api/brand-asset?kind=layout-preview&brandId={b}&index=5", headers=bob(env)), assets=PREVIEWS)))
asset.add("R3",
          Probe("logo", with_brand(lambda env, b: env.client.get(
              f"/api/brand-asset?kind=logo&brandId={b}", headers=env.auth("alice")), assets={"logo": PNG}), _png_check),
          Probe("a member reads an org brand's master (stored under the owner)", with_org(
              lambda env, b: env.client.put(f"/api/brand-asset?kind=master&brandId={b}", headers=bob(env)),
              assets={"master": PNG}), _png_check))
asset.add("E500 Internal error",
          Probe("POST malformed JSON", send("POST", "/api/brand-asset", content=b"{x")),
          Probe("POST empty body", send("POST", "/api/brand-asset", content=b"")),
          Probe("POST JSON null", send("POST", "/api/brand-asset", content=b"null")),
          Probe("the asset store is unreachable", with_brand(lambda env, b: (
              broken(env, "brands", "get_asset"),
              env.client.get(f"/api/brand-asset?kind=logo&brandId={b}", headers=env.auth("alice")))[1])))
asset.add("E400 kind must be one of logo, master, titleMaster, dividerMaster",
          Probe("POST bad kind (checked before brandId)", send("POST", "/api/brand-asset", json={"kind": "favicon"})),
          Probe("GET no kind", send("GET", "/api/brand-asset")),
          Probe("GET bad kind", send("GET", f"/api/brand-asset?kind=nope&brandId={UNKNOWN}")))
asset.add("E400 brandId required",
          Probe("POST no brandId", send("POST", "/api/brand-asset", json={"kind": "logo", "b64": "x"})),
          Probe("POST numeric brandId", send("POST", "/api/brand-asset", json={"kind": "logo", "brandId": 5})),
          Probe("GET layout-preview no brandId", send("GET", "/api/brand-asset?kind=layout-preview&index=0")),
          Probe("GET logo empty brandId", send("GET", "/api/brand-asset?kind=logo&brandId=")))
asset.add("E404 Brand not found",
          Probe("unknown", send("GET", f"/api/brand-asset?kind=logo&brandId={UNKNOWN}")),
          Probe("someone else's (no storage touched)", lambda env: env.client.post(
              "/api/brand-asset", headers=env.auth("alice"), json={"kind": "logo", "brandId": env.seed_brand("bob"),
                                                                   "b64": base64.b64encode(PNG).decode()})),
          Probe("someone else's, GET", lambda env: env.client.get(
              f"/api/brand-asset?kind=logo&brandId={env.seed_brand('bob')}", headers=env.auth("alice"))))
asset.add("E403 Organization brands can only be edited by an admin",
          Probe("a member uploads", with_org(lambda env, b: env.client.post(
              "/api/brand-asset", headers=bob(env), json={"kind": "logo", "brandId": b,
                                                           "b64": base64.b64encode(PNG).decode()}))))
asset.add("E400 fromLayoutIndex only valid for master, titleMaster, dividerMaster",
          Probe("logo", upload({"kind": "logo", "fromLayoutIndex": 0})))
asset.add("E400 fromLayoutIndex must be a non-negative integer",
          *(Probe(f"{v!r}", upload({"kind": "master", "fromLayoutIndex": v})) for v in (None, -1, 1.5, "2", True)))
asset.add("E404 No rendered layout at that index",
          Probe("index 3 was never rendered", upload({"kind": "master", "fromLayoutIndex": 3})),
          Probe("a huge index", upload({"kind": "master", "fromLayoutIndex": 10**30})))
asset.add("E422 Rendered layout is not a PNG",
          Probe("the stored preview is not a PNG", with_brand(lambda env, b: env.client.post(
              "/api/brand-asset", headers=env.auth("alice"), json={"brandId": b, "kind": "master", "fromLayoutIndex": 0}),
              assets={"layout-preview-0": b"GIF89a-not-a-png"})))
asset.add("E400 b64 image data required",
          *(Probe(f"b64 {v!r}", upload({"kind": "logo", **({} if v is ... else {"b64": v})})) for v in (..., "", 5)))
asset.add("E400 Image must be a PNG up to 4MB",
          Probe("decodes to nothing", upload({"kind": "logo", "b64": "!!!"})),
          Probe("over 4 MiB", upload({"kind": "logo", "b64": base64.b64encode(PNG + b"\0" * (4 * 1024 * 1024)).decode()})))
asset.add("E400 Image must be a PNG",
          Probe("a JPEG", upload({"kind": "logo", "b64": base64.b64encode(b"\xff\xd8\xff\xe0jpeg").decode()})))
asset.add("E400 index must be a non-negative integer",
          *(Probe(f"index={v}", with_brand(lambda env, b, v=v: env.client.get(
              f"/api/brand-asset?kind=layout-preview&brandId={b}&index={v}", headers=env.auth("alice"))))
            for v in ("-1", "1.5", "abc")))
asset.add("E404 No layout preview at that index",
          Probe("index 9", with_brand(lambda env, b: env.client.get(
              f"/api/brand-asset?kind=layout-preview&brandId={b}&index=9", headers=env.auth("alice")), assets=PREVIEWS)),
          Probe("before any import", with_brand(lambda env, b: env.client.get(
              f"/api/brand-asset?kind=layout-preview&brandId={b}&index=0", headers=env.auth("alice")))))
asset.add("E404 No asset uploaded",
          Probe("nothing uploaded", with_brand(lambda env, b: env.client.get(
              f"/api/brand-asset?kind=dividerMaster&brandId={b}", headers=env.auth("alice")))))

# --------------------------------------------------------------------------------- /brand-pptx
pptx = RouteCases("/api/brand-pptx")
PPTX = tiny_pptx(1)


def post_pptx(*, data: bytes = PPTX, brand: Any = None, who: str = "alice", name: str = "file") -> Any:
    def run(env: DarwinEnv) -> Any:
        fields = {"brandId": brand(env) if callable(brand) else brand} if brand is not None else {}
        return env.client.post("/api/brand-pptx", headers=env.auth(who), data=fields,
                               files={name: ("template.pptx", data, "application/octet-stream")})
    return run


def _queued_pptx(env: DarwinEnv, response: Any) -> None:
    record = env.job("alice", response.json()["jobId"])
    assert record.type == BRAND_PPTX_JOB and record.status == "queued"
    assert set(record.inputs) == {"pptxRef", "brandId"}


add_auth(pptx, "POST", "/api/brand-pptx", files={"file": ("t.pptx", PPTX)})
pptx.add("R0",
         Probe("with an editable brand", post_pptx(brand=lambda env: env.seed_brand("alice")), _queued_pptx),
         Probe("no brand", post_pptx()),
         Probe("an empty brandId is no brand", post_pptx(brand="")))
pptx.add("E405 POST only",
         Probe("GET (after auth)", send("GET", "/api/brand-pptx")))
pptx.add("E415 Expected multipart/form-data",
         Probe("JSON", send("POST", "/api/brand-pptx", json={"file": "x"})),
         Probe("no body", send("POST", "/api/brand-pptx")))
pptx.add("E400 Could not parse form data",
         Probe("multipart without a boundary", lambda env: env.client.post(
             "/api/brand-pptx", content=b"garbage", headers={**env.auth("alice"),
                                                             "content-type": "multipart/form-data"})))
pptx.add("E400 file is required",
         Probe("another field name", post_pptx(name="upload")),
         Probe("file as a plain field", lambda env: env.client.post(
             "/api/brand-pptx", headers=env.auth("alice"), data={"file": "not a file"},
             files={"other": ("x.txt", b"x", "text/plain")})))
pptx.add("E404 Brand not found",
         Probe("someone else's brand", post_pptx(brand=lambda env: env.seed_brand("bob"))),
         Probe("a member of an org brand (404, not 403)", lambda env: (b := org_brand(env), env.client.post(
             "/api/brand-pptx", headers=bob(env), data={"brandId": b},
             files={"file": ("t.pptx", PPTX, "application/octet-stream")}))[1]),
         Probe("checked BEFORE the size", post_pptx(data=b"", brand=UNKNOWN)))
pptx.add("E400 Template must be a .pptx file up to 5 MB",
         Probe("empty", post_pptx(data=b"")),
         Probe("over 5 MiB", post_pptx(data=b"PK\x03\x04" + b"\0" * (5 * 1024 * 1024))))
pptx.add("E400 File must be a valid .pptx",
         Probe("a PDF", post_pptx(data=PDF)))
pptx.add("E500 Internal error",
         Probe("the blob store is unreachable", lambda env: (broken(env, "blobs", "put"), post_pptx()(env))[1]))

# -------------------------------------------------------------------------- /brand-pptx-status
pptx_status = RouteCases("/api/brand-pptx-status")
add_auth(pptx_status, "GET", f"/api/brand-pptx-status?jobId={UNKNOWN}")


def fake_layout_renderer(env: DarwinEnv) -> None:
    """Run brand-pptx jobs with the fake PptxRender (white PNGs named as the real service names them)."""
    env.runtime.handlers[BRAND_PPTX_JOB] = brand_pptx_handler(BrandJobDeps(
        storage=env.store, renderer_factory=FakeRenderer, guidelines_model=env.model, storyline=env.runtime.storyline))


def imported(env: DarwinEnv, *, brand: bool = True, data: bytes | None = None, who: str = "alice") -> str:
    fake_layout_renderer(env)
    template = data if data is not None else _template()
    fields = {"brandId": env.seed_brand("alice")} if brand else {}
    r = env.client.post("/api/brand-pptx", headers=env.auth(who), data=fields,
                        files={"file": ("template.pptx", template, "application/octet-stream")})
    assert r.status_code == 202, r.text
    env.run_jobs()
    return str(r.json()["jobId"])


def _template() -> bytes:
    prs = Presentation()
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def poll_pptx(prepare: Any, who: str = "alice") -> Any:
    def run(env: DarwinEnv) -> Any:
        return env.client.get(f"/api/brand-pptx-status?jobId={prepare(env)}", headers=env.auth(who))
    return run


def _fields(env: DarwinEnv, response: Any) -> None:
    body = response.json()
    fields = body["fields"]
    assert body["status"] == "done" and "error" not in body
    assert {"headingFont", "bodyFont", "primaryColor", "accentColor", "layouts", "workzone", "allColors",
            "furnitureKey", "furnitureArchetypes", "headingPlaceholders", "layoutPreviews"} <= set(fields)
    assert "masterKey" not in fields and "logoKey" not in fields  # no master is produced; no logo role here
    assert set(fields["workzone"]) == {"left", "top", "width", "height"}
    names = {p["index"]: p["name"] for p in fields["layoutPreviews"]}
    assert names[0] == "Title Slide" and len(names) == 11  # re-keyed to the extractor's indices by name
    brand_id = fields["furnitureKey"].split("/")[2]
    assert env.asset("alice", brand_id, "layout-preview-10") is not None
    assert env.asset("alice", brand_id, "furniture") is not None
    assert env.asset("alice", brand_id, "furniture-all") is not None


def _fields_without_brand(env: DarwinEnv, response: Any) -> None:
    fields = response.json()["fields"]
    assert not {"furnitureKey", "layoutPreviews", "logoKey", "headingPlaceholders"} & set(fields)


def _error_text(env: DarwinEnv, response: Any) -> None:
    body = response.json()
    assert body["status"] == "error" and body["error"] == "The file could not be opened as a PowerPoint deck."
    assert "fields" not in body


pptx_status.add("R0",
                Probe("unknown id", send("GET", f"/api/brand-pptx-status?jobId={UNKNOWN}"), exactly({"status": "pending"})),
                Probe("someone ELSE's unknown id", send("GET", "/api/brand-pptx-status?jobId=zzz", who="mallory"),
                      exactly({"status": "pending"})))
pptx_status.add("R1",
                Probe("done, with a brand: furniture, heading and layout previews stored", poll_pptx(imported), _fields),
                Probe("done, no brand: extracted fields only", poll_pptx(lambda env: imported(env, brand=False)),
                      _fields_without_brand),
                Probe("a zip that is not a deck: status error", poll_pptx(lambda env: imported(
                    env, data=b"PK\x03\x04" + b"\0" * 64)), _error_text),
                Probe("queued: pending", poll_pptx(lambda env: env.client.post(
                    "/api/brand-pptx", headers=env.auth("alice"),
                    files={"file": ("t.pptx", PPTX, "application/octet-stream")}).json()["jobId"]),
                    exactly({"status": "pending"})))
pptx_status.add("E400 jobId is required",
                Probe("missing", send("GET", "/api/brand-pptx-status")),
                Probe("empty", send("GET", "/api/brand-pptx-status?jobId=")))
pptx_status.add("E403 Not your job",
                Probe("bob polls alice's job", poll_pptx(lambda env: imported(env, brand=False), who="bob")))
pptx_status.add("E500 Internal error",
                Probe("the job store is unreachable", lambda env: (broken(env, "jobs", "get"), send(
                    "GET", f"/api/brand-pptx-status?jobId={UNKNOWN}")(env))[1]))

# --------------------------------------------------------------------------- /brand-archetypes
arche = RouteCases("/api/brand-archetypes")
LAYOUTS = {"cover": 0, "divider": 2, "content": 5}


def apply_layouts(body: Any = None, *, who: str = "alice", **seed: Any) -> Any:
    seed = seed or {"assets": {**PREVIEWS, "furniture-all": j(CAPTURE)}}

    def run(env: DarwinEnv) -> Any:
        b = env.seed_brand("alice", **seed)
        payload = {"brandId": b, "layouts": LAYOUTS} if body is None else body(b)
        return env.client.post("/api/brand-archetypes", headers=env.auth(who), json=payload)
    return run


def _applied(env: DarwinEnv, response: Any) -> None:
    body = response.json()
    kit = body["kit"]
    brand_id = kit["masterKey"].split("/")[2]
    owner = env.brand("alice", brand_id).owner_id
    assert body["furnitureRebuilt"] is True
    assert kit["archetypeLayouts"] == LAYOUTS
    assert kit["titleMasterKey"] == f"brand/{owner}/{brand_id}/titleMaster.png"
    assert kit["furnitureArchetypes"] == ["cover", "divider", "content"]
    assert kit["headingPlaceholders"]["cover"] == {"left": 0.05, "top": 0.3, "width": 0.9, "height": 0.2,
                                                   "font": "Georgia", "color": "#112233", "sizePt": 40}
    furniture = json.loads(env.asset("alice", brand_id, "furniture") or b"{}")
    assert "layoutsByIndex" not in furniture and "layoutsByIndex" not in furniture["background"]
    assert set(furniture["media"]) == {"sha-a", "sha-b"}  # unreferenced media pruned
    assert set(furniture["layouts"]["divider"]["placeholders"]) == {"heading"}  # only a cover keeps the date
    assert set(furniture["layouts"]["cover"]["placeholders"]) == {"heading", "date"}
    assert furniture["background"]["layouts"] == {"divider": {"fill": "navy"}}
    stored = env.brand("alice", brand_id).kit
    assert stored["keep"] == "me" and stored["archetypeLayouts"] == LAYOUTS  # merged over the stored kit


def _legacy_import(env: DarwinEnv, response: Any) -> None:
    body = response.json()
    assert body["furnitureRebuilt"] is False
    assert set(body["kit"]) == {"archetypeLayouts", "titleMasterKey", "dividerMasterKey", "masterKey"}


arche.add("R0",
          Probe("masters copied, furniture rebuilt, kit merged",
                apply_layouts(kit={"keep": "me"}, assets={**PREVIEWS, "furniture-all": j(CAPTURE)}), _applied))
arche.add("E405 POST only",
          Probe("GET with no token: 405 before 401", lambda env: env.client.get("/api/brand-archetypes")),
          Probe("PUT", send("PUT", "/api/brand-archetypes", json={})))
add_auth(arche, "POST", "/api/brand-archetypes", json={})
arche.add("E400 brandId required",
          Probe("malformed JSON", send("POST", "/api/brand-archetypes", content=b"{x")),
          Probe("no brandId", send("POST", "/api/brand-archetypes", json={"layouts": LAYOUTS})),
          Probe("JSON null", send("POST", "/api/brand-archetypes", content=b"null")))
arche.add("E400 layouts must map cover, divider and content to layout indices",
          *(Probe(f"{v!r}", apply_layouts(lambda b, v=v: {"brandId": b, "layouts": v}))
            for v in (None, {"cover": 0, "divider": 1}, {"cover": 0, "divider": -1, "content": 1},
                      {"cover": 0, "divider": 1.5, "content": 1}, {"cover": "0", "divider": 1, "content": 1})))
arche.add("E404 Brand not found",
          Probe("unknown brand", send("POST", "/api/brand-archetypes", json={"brandId": UNKNOWN, "layouts": LAYOUTS})))
arche.add("E403 Organization brands can only be edited by an admin",
          Probe("a member", with_org(lambda env, b: env.client.post(
              "/api/brand-archetypes", headers=bob(env), json={"brandId": b, "layouts": LAYOUTS}), assets=PREVIEWS)))


def _nothing_written(env: DarwinEnv, response: Any) -> None:
    assert response.json() == {"error": "No rendered layout at index 2 — re-import the template"}
    brand_id = env.client.get("/api/brands", headers=env.auth("alice")).json()["brands"][0]["id"]
    assert env.asset("alice", brand_id, "titleMaster") is None  # read all before any write


arche.add("E404 No rendered layout at index <n> — re-import the template",
          Probe("divider's preview is missing: nothing written", apply_layouts(
              assets={"layout-preview-0": PNG, "layout-preview-5": PNG}), _nothing_written),
          Probe("before any import", apply_layouts(assets={}), error_is(
              "No rendered layout at index 0 — re-import the template")))
arche.add("E422 Rendered layout is not a PNG",
          Probe("a stored preview is not a PNG", apply_layouts(assets={**PREVIEWS, "layout-preview-5": b"GIF89a"})))
arche.add("E500 Internal error",
          Probe("furniture-all is not JSON", apply_layouts(assets={**PREVIEWS, "furniture-all": b"{oops"})),
          Probe("furniture-all is JSON null", apply_layouts(assets={**PREVIEWS, "furniture-all": b"null"})))

# ------------------------------------------------------------------------------ /brand-heading
heading = RouteCases("/api/brand-heading")
BOX = {"left": 0.04, "top": 0.05, "width": 0.92, "height": 0.1}
FURNITURE = {"layouts": {"content": {"placeholders": {"heading": {"left": 0.1, "top": 0.1, "width": 0.5,
                                                                  "height": 0.1, "size": 28, "font": "Perpetua"}}}}}


def move_title(body: Any = None, *, who: str = "alice", **seed: Any) -> Any:
    def run(env: DarwinEnv) -> Any:
        b = env.seed_brand("alice", **seed)
        payload = {"brandId": b, "archetype": "content", "box": BOX} if body is None else body(b)
        return env.client.post("/api/brand-heading", headers=env.auth(who), json=payload)
    return run


STYLED = {"headingPlaceholders": {
    "cover": {"left": 0.03, "top": 0.14, "width": 0.78, "height": 0.37, "junk": 1},
    "content": {"left": 0.1, "top": 0.1, "width": 0.5, "height": 0.1, "font": "Perpetua", "color": "#1F1F1F",
                "size": 28}}}


def _moved(env: DarwinEnv, response: Any) -> None:
    body = response.json()
    assert body == {"furnitureUpdated": True, "headingPlaceholders": {
        "cover": {"left": 0.03, "top": 0.14, "width": 0.78, "height": 0.37, "junk": 1},  # passed through as stored
        "content": {**BOX, "font": "Perpetua", "color": "#1F1F1F", "sizePt": 28}}}
    brand_id = env.client.get("/api/brands", headers=env.auth("alice")).json()["brands"][0]["id"]
    furniture = json.loads(env.asset("alice", brand_id, "furniture") or b"{}")
    assert furniture["layouts"]["content"]["placeholders"]["heading"] == {**BOX, "size": 28, "font": "Perpetua"}


heading.add("R0",
            Probe("preview box and furniture moved, styling kept",
                  move_title(kit=STYLED, assets={"furniture": j(FURNITURE)}), _moved),
            Probe("an empty brand: kit only", move_title(), exactly({
                "headingPlaceholders": {"content": BOX}, "furnitureUpdated": False})),
            Probe("edges within the 0.001 tolerance", move_title(lambda b: {
                "brandId": b, "archetype": "cover", "box": {"left": 0.5, "top": 0.5, "width": 0.5005, "height": 0.5}})))
heading.add("E405 POST only", Probe("GET with no token", lambda env: env.client.get("/api/brand-heading")))
add_auth(heading, "POST", "/api/brand-heading", json={})
heading.add("E400 brandId required",
            Probe("malformed JSON", send("POST", "/api/brand-heading", content=b"{")),
            Probe("brandId missing", send("POST", "/api/brand-heading", json={"archetype": "cover", "box": BOX})))
heading.add("E400 archetype must be cover, divider or content",
            Probe("hero", move_title(lambda b: {"brandId": b, "archetype": "hero", "box": BOX})))
heading.add("E400 box must be {left, top, width, height} as slide fractions",
            *(Probe(f"{v!r}", move_title(lambda b, v=v: {"brandId": b, "archetype": "cover", "box": v}))
              for v in (None, {**BOX, "left": -0.1}, {**BOX, "width": 0}, {**BOX, "left": 0.5, "width": 0.6},
                        {**BOX, "top": "0.1"}, {**BOX, "height": True})))
heading.add("E404 Brand not found",
            Probe("unknown", send("POST", "/api/brand-heading", json={"brandId": UNKNOWN, "archetype": "cover", "box": BOX})))
heading.add("E403 Organization brands can only be edited by an admin",
            Probe("a member", with_org(lambda env, b: env.client.post(
                "/api/brand-heading", headers=bob(env), json={"brandId": b, "archetype": "cover", "box": BOX}))))
heading.add("E500 Internal error",
            Probe("furniture is not JSON", move_title(assets={"furniture": b"{nope"})),
            Probe("the brand store is unreachable", lambda env: (
                b := env.seed_brand("alice"), broken(env, "brands", "update"),
                env.client.post("/api/brand-heading", headers=env.auth("alice"),
                                json={"brandId": b, "archetype": "cover", "box": BOX}))[2]))

# ------------------------------------------------------------------------------ /brand-extract
extract = RouteCases("/api/brand-extract")
B64_PDF = base64.b64encode(PDF).decode()


def post_extract(body: Any = None, *, raw: bytes | None = None, who: str = "alice") -> Any:
    def run(env: DarwinEnv) -> Any:
        payload = body(env) if callable(body) else body
        if raw is not None:
            return env.client.post("/api/brand-extract", headers=env.auth(who), content=raw)
        return env.client.post("/api/brand-extract", headers=env.auth(who), json=payload)
    return run


def _guidelines_kept(env: DarwinEnv, response: Any) -> None:
    record = env.job("alice", response.json()["jobId"])
    assert set(record.inputs) == {"brandId"}
    assert env.asset("alice", record.inputs["brandId"], "guidelines") == PDF


def _parked(env: DarwinEnv, response: Any) -> None:
    assert set(env.job("alice", response.json()["jobId"]).inputs) == {"pdfRef"}


add_auth(extract, "POST", "/api/brand-extract", json={"b64": B64_PDF})
extract.add("R0",
            Probe("with a brand: kept at the brand's guidelines slot", post_extract(
                lambda env: {"b64": B64_PDF, "brandId": env.seed_brand("alice")}), _guidelines_kept),
            Probe("no brand: parked for the job", post_extract({"b64": B64_PDF}), _parked))
extract.add("E405 POST only", Probe("GET (after auth)", send("GET", "/api/brand-extract")))
extract.add("E400 b64 PDF data required",
            Probe("malformed JSON is treated as {}", post_extract(raw=b"{bad")),
            Probe("empty body", post_extract(raw=b"")),
            Probe("no b64", post_extract({"brandId": UNKNOWN})),
            Probe("b64 a number", post_extract({"b64": 7})),
            Probe("a JSON string body", post_extract("abc")))
extract.add("E400 Guidelines must be a PDF up to 4MB",
            Probe("decodes to nothing", post_extract({"b64": "@@@@"})),
            Probe("over 4 MiB", post_extract({"b64": base64.b64encode(PDF + b"\0" * (4 * 1024 * 1024)).decode()})))
extract.add("E400 Guidelines must be a PDF",
            Probe("a PNG (checked before the brand)", post_extract({"b64": base64.b64encode(PNG).decode(),
                                                                     "brandId": UNKNOWN})))
extract.add("E404 Brand not found",
            *(Probe(f"brandId {v!r}", post_extract({"b64": B64_PDF, "brandId": v})) for v in (None, "", 5, UNKNOWN)),
            Probe("a member of an org brand", lambda env: (b := org_brand(env), env.client.post(
                "/api/brand-extract", headers=bob(env), json={"b64": B64_PDF, "brandId": b}))[1]))
extract.add("E500 Internal error",
            Probe("JSON null body", post_extract(raw=b"null")),
            Probe("the blob store is unreachable", lambda env: (broken(env, "blobs", "put"),
                                                                post_extract({"b64": B64_PDF})(env))[1]))

# ------------------------------------------------------------------------- /brand-extract-status
extract_status = RouteCases("/api/brand-extract-status")
add_auth(extract_status, "GET", f"/api/brand-extract-status?jobId={UNKNOWN}")


def extracted(env: DarwinEnv, step: Any = None) -> str:
    env.model.steps.append(step if step is not None else {
        "company": "Acme", "primaryColor": "#0B2D4F", "accentColor": "red", "headingFont": "Georgia",
        "bodyFont": "", "styleTemplate": "Use {primaryColor} for structure."})
    r = env.client.post("/api/brand-extract", headers=env.auth("alice"), json={"b64": B64_PDF})
    assert r.status_code == 202, r.text
    env.run_jobs()
    return str(r.json()["jobId"])


def poll_extract(prepare: Any, who: str = "alice") -> Any:
    return lambda env: env.client.get(f"/api/brand-extract-status?jobId={prepare(env)}", headers=env.auth(who))


extract_status.add("R0", Probe("unknown id", send("GET", f"/api/brand-extract-status?jobId={UNKNOWN}"),
                               exactly({"status": "pending"})))
extract_status.add("R1",
                   Probe("done: invalid fields dropped one by one", poll_extract(extracted), exactly({
                       "status": "done", "fields": {"company": "Acme", "primaryColor": "#0B2D4F",
                                                    "headingFont": "Georgia",
                                                    "styleTemplate": "Use {primaryColor} for structure."}})),
                   Probe("the model returned no fields: Darwin's error", poll_extract(
                       lambda env: extracted(env, reply(None, stop_reason="refusal"))), exactly({
                           "status": "error", "error": "Claude did not return an extraction tool call"})))


extract_status.add("E400 jobId is required", Probe("missing", send("GET", "/api/brand-extract-status")))
extract_status.add("E403 Not your job", Probe("bob polls alice's job", poll_extract(extracted, who="bob")))
extract_status.add("E500 Internal error",
                   Probe("the job store is unreachable", lambda env: (broken(env, "jobs", "get"), send(
                       "GET", f"/api/brand-extract-status?jobId={UNKNOWN}")(env))[1]))

# ------------------------------------------------------------------------------ /brand-preview
preview = RouteCases("/api/brand-preview")
add_auth(preview, "POST", "/api/brand-preview", json={"brandId": UNKNOWN})
preview.add("R0",
            Probe("own brand", with_brand(lambda env, b: env.client.post(
                "/api/brand-preview", headers=env.auth("alice"), json={"brandId": b}))),
            Probe("any method with a JSON body", with_brand(lambda env, b: env.client.put(
                "/api/brand-preview", headers=env.auth("alice"), json={"brandId": b}))))
preview.add("E400 brandId required",
            *(Probe(f"brandId {v!r}", send("POST", "/api/brand-preview", json={"brandId": v})) for v in (None, "", 3)),
            Probe("a JSON array body", send("POST", "/api/brand-preview", json=[1])))
preview.add("E404 Brand not found",
            Probe("unknown", send("POST", "/api/brand-preview", json={"brandId": UNKNOWN})),
            Probe("a member of an org brand", with_org(lambda env, b: env.client.post(
                "/api/brand-preview", headers=bob(env), json={"brandId": b}))))
preview.add("E500 Internal error",
            Probe("malformed JSON", send("POST", "/api/brand-preview", content=b"{x")),
            Probe("bodyless GET", send("GET", "/api/brand-preview")),
            Probe("JSON null", send("POST", "/api/brand-preview", content=b"null")))

# ------------------------------------------------------------------------ /brand-preview-status
preview_status = RouteCases("/api/brand-preview-status")
add_auth(preview_status, "GET", f"/api/brand-preview-status?jobId={UNKNOWN}")


def previewed(env: DarwinEnv, *, run: bool = True, **seed: Any) -> str:
    b = env.seed_brand("alice", **seed)
    r = env.client.post("/api/brand-preview", headers=env.auth("alice"), json={"brandId": b})
    assert r.status_code == 202, r.text
    if run:
        env.run_jobs()
    return str(r.json()["jobId"])


def poll_preview(prepare: Any, who: str = "alice") -> Any:
    return lambda env: env.client.get(f"/api/brand-preview-status?jobId={prepare(env)}", headers=env.auth(who))


def _image(env: DarwinEnv, response: Any) -> None:
    assert base64.b64decode(response.json()["image"]) == tiny_png()
    assert len(env.images.calls) == 1 and env.images.calls[0].kind == "generate"


def _gone(env: DarwinEnv) -> str:
    job_id = previewed(env)
    record = env.job("alice", job_id)
    brand_id, name = record.result["brandId"], record.result["asset"]
    env.call(env.store.brands.delete_asset, env.ctx("alice"), brand_id, name)
    return job_id


def _master_missing(env: DarwinEnv, response: Any) -> None:
    assert response.json() == {"status": "error", "error": (
        "Brand master could not be loaded — re-upload it, then preview again.")}
    assert env.images.calls == []


preview_status.add("R0", Probe("unknown id", send("GET", f"/api/brand-preview-status?jobId={UNKNOWN}"),
                               exactly({"status": "pending"})))
preview_status.add("R1", Probe("done: the PNG as base64", poll_preview(previewed), _image))
preview_status.add("R2", Probe("done, but the image is gone", poll_preview(_gone), exactly(
    {"status": "error", "error": "Preview image is no longer available"})))
preview_status.add("R3",
                   Probe("queued", poll_preview(lambda env: previewed(env, run=False)), exactly({"status": "pending"})),
                   Probe("the kit names a master that is not stored", poll_preview(
                       lambda env: previewed(env, kit={"masterKey": "brand/x/y/master.png"})), _master_missing))
preview_status.add("E400 jobId is required", Probe("missing", send("GET", "/api/brand-preview-status?jobId=")))
preview_status.add("E403 Not your job", Probe("bob polls alice's job", poll_preview(previewed, who="bob")))
preview_status.add("E500 Internal error",
                   Probe("the job store is unreachable", lambda env: (broken(env, "jobs", "get"), send(
                       "GET", f"/api/brand-preview-status?jobId={UNKNOWN}")(env))[1]))

ROUTES = [brands, asset, pptx, pptx_status, arche, heading, extract, extract_status, preview, preview_status]
